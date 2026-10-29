"""Command-wide timeout and cancellation contract tests."""

from __future__ import annotations

import argparse
import json
import sys
import threading
from pathlib import Path
from typing import Any, cast

import pytest

from docwen_application.runtime_capability_catalog import RuntimeRoute

pytestmark = pytest.mark.contract


def _markdown_to_docx_route() -> RuntimeRoute:
    return RuntimeRoute(
        id="test:markdown:docx",
        operation="conversion",
        source="markdown",
        source_category="markdown",
        target="docx",
        action_name="",
        available=True,
        state="available",
        options=(),
    )


def test_execution_deadline_cancels_registered_application_operation() -> None:
    from docwen_cli.execution_deadline import ExecutionDeadline
    from docwen_core.models.request import ConversionRequest, FileRef

    cancelled = threading.Event()
    released: list[tuple[str, object]] = []
    reservation = object()

    class _Controller:
        @staticmethod
        def prepare_execution_cancellation(
            request: ConversionRequest, *, batch: bool = False, on_reserved=None
        ) -> object:
            del request
            assert batch is False
            if on_reserved is not None:
                on_reserved(reservation)
            return reservation

        @staticmethod
        def cancel(task_id: str, *, reservation: object | None = None) -> None:
            assert task_id == "deadline-task"
            assert reservation is not None
            cancelled.set()

        @staticmethod
        def release_execution_cancellation(task_id: str, reservation: object) -> None:
            released.append((task_id, reservation))

    request = ConversionRequest(
        request_id="deadline-task",
        input_refs=[FileRef(path="sample.md", format="md", category="text")],
        target_format="docx",
    )
    deadline = ExecutionDeadline(_Controller(), 0.02).start()
    token = deadline.register(request)
    try:
        assert cancelled.wait(1)
        assert deadline.timed_out is True
    finally:
        deadline.release(request, token)
        deadline.close()

    assert released == [("deadline-task", reservation)]


def test_timed_out_cancellation_is_reported_as_timeout(tmp_path: Path, capsys: pytest.CaptureFixture[str]) -> None:
    from docwen_cli.commands.convert import _execute_single
    from docwen_cli.exit_codes import ExitCode
    from docwen_core.models.result import ConversionErrorInfo, ConversionResult

    source = tmp_path / "sample.md"
    source.write_text("hello", encoding="utf-8")

    class _Controller:
        @staticmethod
        def execute_single(request):
            return ConversionResult(
                task_id=request.request_id,
                success=False,
                error=ConversionErrorInfo(error_type="cancelled", message="Uncommitted work cancelled"),
            )

    class _ExpiredDeadline:
        timed_out = True

        @staticmethod
        def register(_request, *, batch=False):
            return object()

        @staticmethod
        def release(_request, _reservation) -> None:
            return None

        @staticmethod
        def finish() -> None:
            return None

    args = argparse.Namespace(
        json=True,
        timing=False,
        quiet=True,
        verbose=False,
        command_path="convert",
        output_path=str(tmp_path / "sample.docx"),
        overwrite=False,
    )
    exit_code = _execute_single(
        _Controller(),
        "",
        str(source),
        "docx",
        {},
        args,
        json_mode=True,
        deadline=cast(Any, _ExpiredDeadline()),
        route=_markdown_to_docx_route(),
    )

    captured = capsys.readouterr()
    payload = json.loads(captured.out)
    assert exit_code == int(ExitCode.TIMEOUT)
    assert captured.err == ""
    assert payload["success"] is False
    assert payload["error"]["category"] == "timeout"
    assert payload["error"]["code"] == "operation_timeout"


def test_deadline_cancels_application_operation_and_publishes_no_artifact(
    tmp_path: Path, capsys: pytest.CaptureFixture[str]
) -> None:
    from docwen_application.controller import ApplicationController
    from docwen_cli.commands.convert import _execute_single
    from docwen_cli.execution_deadline import ExecutionDeadline
    from docwen_cli.exit_codes import ExitCode
    from docwen_core.models.result import ConversionErrorInfo, ConversionResult

    source = tmp_path / "source.md"
    output = tmp_path / "must-not-exist.docx"
    source.write_text("hello", encoding="utf-8")
    cancelled = threading.Event()

    class _Runtime:
        is_available = True

        @staticmethod
        def reserve_cancellation(task_id: str) -> None:
            del task_id
            return None

        @staticmethod
        def release_cancellation(task_id: str) -> None:
            del task_id
            return None

        @staticmethod
        def shutdown() -> None:
            return None

        @staticmethod
        def cancel(task_id: str) -> None:
            del task_id
            cancelled.set()

        @staticmethod
        def execute(request):
            assert cancelled.wait(1)
            return ConversionResult(
                task_id=request.request_id,
                success=False,
                error=ConversionErrorInfo(error_type="cancelled", message="cancelled"),
            )

    controller = ApplicationController(runtime_port=_Runtime())
    args = argparse.Namespace(
        json=True,
        timing=False,
        quiet=True,
        verbose=False,
        command_path="convert",
        output_path=str(output),
        overwrite=False,
    )
    deadline = ExecutionDeadline(controller, 0.02).start()
    try:
        exit_code = _execute_single(
            controller,
            "",
            str(source),
            "docx",
            {},
            args,
            json_mode=True,
            deadline=deadline,
            route=_markdown_to_docx_route(),
        )
    finally:
        deadline.close()

    payload = json.loads(capsys.readouterr().out)
    assert exit_code == int(ExitCode.TIMEOUT)
    assert payload["error"]["code"] == "operation_timeout"
    # A cold pre-admission inspection may observe the deadline before Runtime
    # starts.  If Runtime does start, its assertion above proves the cancel was
    # forwarded.  Both valid races must produce the same public timeout result.
    assert not output.exists()


def _read_cli_trace(path: Path) -> list[dict[str, Any]]:
    if not path.exists():
        return []
    records = []
    for line in path.read_text(encoding="utf-8").splitlines():
        try:
            records.append(json.loads(line))
        except json.JSONDecodeError:
            # The child may still be appending the final bounded trace line.
            continue
    return records


def _real_cli_child(tmp_path: Path, mode: str, *, interrupt: bool) -> tuple[int, dict[str, Any], list[dict[str, Any]]]:
    import os
    import signal
    import subprocess
    import sys
    import time

    from tests.support.subprocess_runner import append_subprocess_run_record

    for name in ("first", "second", "third"):
        (tmp_path / f"{name}.csv").write_bytes(f"Name,Amount\n{name},7\n".encode())
    (tmp_path / "published").mkdir()
    env = dict(os.environ)
    env.update(
        DOCWEN_DATA_DIR=str(tmp_path / "profile"),
        DOCWEN_CONFIG_DIR=str(tmp_path / "config"),
        DOCWEN_LOG_DIR=str(tmp_path / "logs"),
        DOCWEN_RUNTIME_ROOT=str(tmp_path / "runtime"),
        PYTHONPATH=os.pathsep.join(sys.path),
    )
    argv = [sys.executable, "-m", "tests.support.cli_deadline_process", mode, str(tmp_path)]
    start = time.monotonic()
    child = subprocess.Popen(argv, env=env, stdout=subprocess.PIPE, stderr=subprocess.PIPE, text=True, encoding="utf-8")
    timed_out = False
    stdout = stderr = ""
    try:
        if interrupt:
            end = time.monotonic() + 25
            while child.poll() is None and time.monotonic() < end:
                if any(row["event"] == "ready" for row in _read_cli_trace(tmp_path / "trace.jsonl")):
                    break
                time.sleep(0.01)
            trace = _read_cli_trace(tmp_path / "trace.jsonl")
            assert any(row["event"] == "ready" for row in trace), trace
            child.send_signal(signal.SIGINT)
            # Allow the actual Application cancellation call to run before
            # releasing an executor-start gate. Baseline failures stay bounded.
            end = time.monotonic() + 1
            while child.poll() is None and time.monotonic() < end:
                if any(
                    row["event"] == "cancelled-through-application" for row in _read_cli_trace(tmp_path / "trace.jsonl")
                ):
                    break
                time.sleep(0.01)
            (tmp_path / "release").write_text("parent released owned phase gate", encoding="utf-8")
        stdout, stderr = child.communicate(timeout=25)
    except subprocess.TimeoutExpired:
        timed_out = True
        raise
    finally:
        (tmp_path / "release").write_text("final release", encoding="utf-8")
        if child.poll() is None:
            child.kill()
            stdout, stderr = child.communicate(timeout=5)
        (tmp_path / "stdout.json").write_text(stdout, encoding="utf-8")
        (tmp_path / "stderr.txt").write_text(stderr, encoding="utf-8")
        append_subprocess_run_record(
            {
                "argv": argv,
                "returncode": child.returncode,
                "duration_seconds": time.monotonic() - start,
                "timed_out": timed_out,
                "timeout_seconds": 25,
                "text": True,
                "capture_output": True,
            }
        )
    assert stdout.strip(), stderr
    return child.returncode, json.loads(stdout), _read_cli_trace(tmp_path / "trace.jsonl")


def _assert_published_csv_matrix(path: str, name: str) -> None:
    from openpyxl import load_workbook

    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        assert workbook.sheetnames == ["Sheet1"]
        assert list(workbook.worksheets[0].iter_rows(values_only=True)) == [("Name", "Amount"), (name, "7")]
    finally:
        workbook.close()


@pytest.mark.integration
@pytest.mark.linux_only
@pytest.mark.skipif(sys.platform != "linux", reason="Real Linux SIGINT; not a Windows Ctrl+C claim")
@pytest.mark.parametrize("mode", ["before_worker", "during_conversion", "partial_interrupt"])
def test_sigint_stops_real_cli_before_executor_wait_and_preserves_commits(tmp_path: Path, mode: str) -> None:
    code, payload, trace = _real_cli_child(tmp_path, mode, interrupt=True)
    assert code == 130, payload
    assert payload["error"]["code"] == "operation_cancelled"
    assert payload["data"]["interrupted"] is True
    results = payload["data"]["results"]
    expected_names = (
        ["first.csv", "second.csv", "third.csv"] if mode == "partial_interrupt" else ["first.csv", "second.csv"]
    )
    assert [Path(item["input"]).name for item in results] == expected_names
    reserved = {row["name"]: row["task_id"] for row in trace if row["event"] == "reserved"}
    cancelled = {row["task_id"] for row in trace if row["event"] == "cancelled-through-application"}
    assert reserved["second.csv"] in cancelled
    if mode == "partial_interrupt":
        assert results[0]["success"] is True
        _assert_published_csv_matrix(results[0]["output"], "first")
        assert not results[1]["success"] and results[1]["output"] == ""
        assert results[2]["details"]["execution_state"] == "not_started"
        assert "third.csv" not in reserved
        assert [path.name for path in (tmp_path / "published").rglob("*.xlsx")] == ["first.xlsx"]
    else:
        assert all(not item["success"] and item["output"] == "" for item in results)
        assert reserved["first.csv"] in cancelled
        assert not list((tmp_path / "published").rglob("*.xlsx"))
        if mode == "before_worker":
            assert not any(row["event"] == "worker" for row in trace)
    for name in ("first", "second", "third"):
        assert (tmp_path / f"{name}.csv").read_bytes() == f"Name,Amount\n{name},7\n".encode()


@pytest.mark.integration
def test_timeout_retains_real_finalizer_success_and_does_not_start_following_input(tmp_path: Path) -> None:
    code, payload, trace = _real_cli_child(tmp_path, "partial_timeout", interrupt=False)
    assert code == 8, payload
    assert payload["error"]["code"] == "operation_timeout"
    first, second, third = payload["data"]["results"]
    assert first["success"] is True and first["error"] is None
    _assert_published_csv_matrix(first["output"], "first")
    assert second["success"] is False and second["output"] == ""
    assert second["error"]["code"] == "operation_timeout"
    assert third["details"]["execution_state"] == "not_started" and third["output"] == ""
    assert not any(row.get("name") == "third.csv" for row in trace)
    assert [path.name for path in (tmp_path / "published").rglob("*.xlsx")] == ["first.xlsx"]


@pytest.mark.integration
@pytest.mark.parametrize("mode", ["before_publish", "after_publish"])
def test_real_finalizer_commit_order_is_authoritative_at_deadline(tmp_path: Path, mode: str) -> None:
    code, payload, trace = _real_cli_child(tmp_path, mode, interrupt=False)
    assert code == 8, payload
    [item] = payload["data"]["results"]
    assert any(row["event"] == "observed-cancellation" and row["cancelled"] for row in trace)
    if mode == "after_publish":
        assert item["success"] is True and item["error"] is None
        _assert_published_csv_matrix(item["output"], "first")
    else:
        assert item["success"] is False and item["output"] == ""
        assert item["error"]["code"] == "operation_timeout"
        assert not list((tmp_path / "published").rglob("*.xlsx"))


@pytest.mark.integration
def test_bounded_drain_retains_cancelled_scope_until_delayed_real_worker_finishes(tmp_path, monkeypatch):
    from docwen_application.controller import ApplicationController
    from docwen_bundle.runtime_factory import create_runtime_port
    from docwen_cli import batch_execution
    from docwen_cli.commands.execution_request import file_ref_for_runtime
    from docwen_cli.execution_deadline import ExecutionDeadline
    from docwen_core.models.request import ConversionRequest, OutputPolicy
    from docwen_runtime.config import ConfigLoader

    source = tmp_path / "delayed.csv"
    source.write_bytes(b"Name,Amount\nlate,7\n")
    public = tmp_path / "published"
    public.mkdir()
    monkeypatch.setenv("DOCWEN_RUNTIME_ROOT", str(tmp_path / "runtime"))
    loader = ConfigLoader(user_dir=tmp_path / "config", runtime_overrides={"logger": {"console_enable": False}})
    controller = ApplicationController(runtime_port=create_runtime_port(config_loader=loader))
    controller.start()
    request = ConversionRequest(
        request_id="retained-worker-scope",
        input_refs=[file_ref_for_runtime(str(source))],
        target_format="xlsx",
        output_policy=OutputPolicy(output_dir=str(public)),
    )
    release = threading.Event()
    started = threading.Event()
    finished = threading.Event()
    results = []
    original_execute = controller.execute_single
    original_wait = batch_execution.wait
    injected = False

    def delayed_execute(actual):
        started.set()
        try:
            assert release.wait(5)
            result = original_execute(actual)
            results.append(result)
            return result
        finally:
            finished.set()

    def interrupt_once(*args, **kwargs):
        nonlocal injected
        if not injected:
            assert started.wait(2)
            injected = True
            raise KeyboardInterrupt
        return original_wait(*args, **kwargs)

    monkeypatch.setattr(controller, "execute_single", delayed_execute)
    monkeypatch.setattr(batch_execution, "wait", interrupt_once)
    monkeypatch.setattr(batch_execution, "_CANCELLATION_WAIT_SECONDS", 0.05)
    deadline = ExecutionDeadline(controller, 30).start()
    try:
        batch = batch_execution.execute_batch_requests(
            controller, [(0, request)], deadline=deadline, max_workers=1, stop_on_error=True
        )
        deadline.close()
        assert batch.interrupted
        assert batch.results[0].metrics.extra["execution_state"] == "cancellation_unconfirmed"
        release.set()
        assert finished.wait(5)
        [result] = results
        assert not result.success and result.error.error_type == "cancelled"
        assert not result.artifacts
        assert not list(public.iterdir())
        assert source.read_bytes() == b"Name,Amount\nlate,7\n"
    finally:
        release.set()
        finished.wait(5)
        deadline.close()
        controller.stop()


@pytest.mark.integration
@pytest.mark.linux_only
@pytest.mark.skipif(sys.platform != "linux", reason="Real Linux SIGINT; not a Windows Ctrl+C claim")
@pytest.mark.parametrize("mode", ["submit_after_enqueue", "submit_after_commit", "after_submit_return"])
def test_sigint_inside_submit_and_before_registration_keeps_owned_results(tmp_path: Path, mode: str) -> None:
    code, payload, trace = _real_cli_child(tmp_path, mode, interrupt=True)
    assert code == 130, payload
    assert payload["error"]["code"] == "operation_cancelled"
    first, second = payload["data"]["results"]
    assert Path(first["input"]).name == "first.csv"
    assert Path(second["input"]).name == "second.csv"
    assert second["details"]["execution_state"] == "not_started"
    assert not any(row.get("name") == "second.csv" for row in trace)
    if mode == "submit_after_enqueue":
        assert first["details"]["execution_state"] == "not_started"
        assert not any(row["event"] == "worker" for row in trace)
        assert not list((tmp_path / "published").rglob("*.xlsx"))
    else:
        assert first["success"] and first["error"] is None
        _assert_published_csv_matrix(first["output"], "first")
        assert [path.name for path in (tmp_path / "published").rglob("*.xlsx")] == ["first.xlsx"]


@pytest.mark.integration
@pytest.mark.linux_only
@pytest.mark.skipif(sys.platform != "linux", reason="Real Linux SIGINT; not a Windows Ctrl+C claim")
@pytest.mark.parametrize("mode", ["prepare_return", "deadline_return", "alias_published"])
def test_reservation_return_interrupt_releases_exact_owner_and_controller_is_reusable(
    tmp_path: Path, mode: str
) -> None:
    code, payload, trace = _real_cli_child(tmp_path, mode, interrupt=True)
    assert code == 130, payload
    assert payload["error"]["code"] == "operation_cancelled"
    first, second = payload["data"]["results"]
    assert first["details"]["execution_state"] == "not_started"
    assert second["details"]["execution_state"] == "not_started"
    assert not list((tmp_path / "published").iterdir())
    assert not any(row["event"] == "worker" for row in trace)
    [reused] = [row for row in trace if row["event"] == "controller-reused"]
    assert reused["before"] == reused["after"] == []
    assert reused["success"]
    [output] = reused["outputs"]
    _assert_published_csv_matrix(output, "first")
    assert (tmp_path / "first.csv").read_bytes() == b"Name,Amount\nfirst,7\n"


def test_expired_rejected_registration_does_not_cancel_or_release_another_owner(tmp_path, monkeypatch):
    from docwen_application.controller import ApplicationController, ControllerError
    from docwen_cli.execution_deadline import ExecutionDeadline
    from docwen_core.models.request import ConversionRequest, FileRef

    source = tmp_path / "source.csv"
    source.write_bytes(b"Name,Amount\nfirst,7\n")
    request = ConversionRequest(
        request_id="existing-owner",
        input_refs=[FileRef(path=str(source), format="csv", category="spreadsheet")],
        target_format="xlsx",
    )
    controller = ApplicationController()
    original = controller.prepare_execution_cancellation(request)
    original_prepare = controller.prepare_execution_cancellation
    deadline = ExecutionDeadline(controller, 0.02)

    def delayed_prepare(actual, *, batch=False, on_reserved=None):
        assert deadline._timed_out.wait(2)
        deadline._timer.join(timeout=2)
        assert not deadline._timer.is_alive()
        return original_prepare(actual, batch=batch, on_reserved=on_reserved)

    monkeypatch.setattr(controller, "prepare_execution_cancellation", delayed_prepare)
    deadline.start()
    try:
        with pytest.raises(ControllerError, match="already owned"):
            deadline.register(request)
        deadline.close()
        assert controller._cancellation_scopes[request.request_id] is original
        assert original.retain_until_release and not original.token.is_cancelled and not original.claimed
        assert deadline._reservations == {} and deadline._request_ids == set()
    finally:
        deadline.close()
        controller.release_execution_cancellation(request.request_id, original)
    assert controller._cancellation_scopes == {}


def test_deadline_during_capture_is_delivered_after_alias_publication(tmp_path):
    from docwen_application.controller import ApplicationController
    from docwen_cli.execution_deadline import ExecutionDeadline
    from docwen_core.models.request import ConversionRequest, FileRef

    source = tmp_path / "source.csv"
    source.write_bytes(b"Name,Amount\nfirst,7\n")
    request = ConversionRequest(
        request_id="capture-expiry",
        input_refs=[FileRef(path=str(source), format="csv", category="spreadsheet")],
        target_format="xlsx",
    )
    controller = ApplicationController()
    captured = []
    deadline = ExecutionDeadline(controller, 0.02).start()

    def capture(handle):
        captured.append(handle)
        assert controller._cancellation_scopes == {}
        # The actual timer is allowed to expire while Application holds its
        # alias-publication lock. No cancellation may re-enter that lock here.
        assert deadline._timed_out.wait(2)

    reservation = None
    try:
        reservation = deadline.register(request, on_reserved=capture)
        assert captured == [reservation]
        scope = controller._cancellation_scopes[request.request_id]
        assert scope is reservation and scope.token.is_cancelled
    finally:
        if reservation is not None:
            deadline.release(request, reservation)
        deadline.close()
    assert controller._cancellation_scopes == {}
