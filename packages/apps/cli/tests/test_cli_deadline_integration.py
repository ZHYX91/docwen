"""Real composed CLI cancellation, bounded drain and final publication regressions."""

from __future__ import annotations

import json
import sys
import threading
from pathlib import Path
from typing import Any

import pytest

pytestmark = [pytest.mark.integration, pytest.mark.release_gate]


def _read_cli_trace(path: Path) -> list[dict[str, Any]]:
    if not path.exists():
        return []
    records = []
    for line in path.read_text(encoding="utf-8").splitlines():
        try:
            records.append(json.loads(line))
        except json.JSONDecodeError:
            # The child may still be appending the final bounded trace line.
            continue
    return records


def _real_cli_child(tmp_path: Path, mode: str, *, interrupt: bool) -> tuple[int, dict[str, Any], list[dict[str, Any]]]:
    import os
    import signal
    import subprocess
    import sys
    import time

    from tests.support.subprocess_runner import append_subprocess_run_record

    for name in ("first", "second", "third"):
        (tmp_path / f"{name}.csv").write_bytes(f"Name,Amount\n{name},7\n".encode())
    (tmp_path / "published").mkdir()
    env = dict(os.environ)
    env.update(
        DOCWEN_DATA_DIR=str(tmp_path / "profile"),
        DOCWEN_CONFIG_DIR=str(tmp_path / "config"),
        DOCWEN_LOG_DIR=str(tmp_path / "logs"),
        DOCWEN_RUNTIME_ROOT=str(tmp_path / "runtime"),
        PYTHONPATH=os.pathsep.join(sys.path),
    )
    argv = [sys.executable, "-m", "tests.support.cli_deadline_process", mode, str(tmp_path)]
    start = time.monotonic()
    child = subprocess.Popen(argv, env=env, stdout=subprocess.PIPE, stderr=subprocess.PIPE, text=True, encoding="utf-8")
    timed_out = False
    stdout = stderr = ""
    try:
        if interrupt:
            end = time.monotonic() + 25
            while child.poll() is None and time.monotonic() < end:
                if any(row["event"] == "ready" for row in _read_cli_trace(tmp_path / "trace.jsonl")):
                    break
                time.sleep(0.01)
            trace = _read_cli_trace(tmp_path / "trace.jsonl")
            assert any(row["event"] == "ready" for row in trace), trace
            child.send_signal(signal.SIGINT)
            # Allow the actual Application cancellation call to run before
            # releasing an executor-start gate. Baseline failures stay bounded.
            end = time.monotonic() + 1
            while child.poll() is None and time.monotonic() < end:
                if any(
                    row["event"] == "cancelled-through-application" for row in _read_cli_trace(tmp_path / "trace.jsonl")
                ):
                    break
                time.sleep(0.01)
            (tmp_path / "release").write_text("parent released owned phase gate", encoding="utf-8")
        stdout, stderr = child.communicate(timeout=25)
    except subprocess.TimeoutExpired:
        timed_out = True
        raise
    finally:
        (tmp_path / "release").write_text("final release", encoding="utf-8")
        if child.poll() is None:
            child.kill()
            stdout, stderr = child.communicate(timeout=5)
        (tmp_path / "stdout.json").write_text(stdout, encoding="utf-8")
        (tmp_path / "stderr.txt").write_text(stderr, encoding="utf-8")
        append_subprocess_run_record(
            {
                "argv": argv,
                "returncode": child.returncode,
                "duration_seconds": time.monotonic() - start,
                "timed_out": timed_out,
                "timeout_seconds": 25,
                "text": True,
                "capture_output": True,
            }
        )
    assert stdout.strip(), stderr
    return child.returncode, json.loads(stdout), _read_cli_trace(tmp_path / "trace.jsonl")


def _assert_published_csv_matrix(path: str, name: str) -> None:
    from openpyxl import load_workbook

    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        assert workbook.sheetnames == ["Sheet1"]
        assert list(workbook.worksheets[0].iter_rows(values_only=True)) == [("Name", "Amount"), (name, "7")]
    finally:
        workbook.close()


@pytest.mark.integration
@pytest.mark.linux_only
@pytest.mark.skipif(sys.platform != "linux", reason="Real Linux SIGINT; not a Windows Ctrl+C claim")
@pytest.mark.parametrize("mode", ["before_worker", "during_conversion", "partial_interrupt"])
def test_sigint_stops_real_cli_before_executor_wait_and_preserves_commits(tmp_path: Path, mode: str) -> None:
    code, payload, trace = _real_cli_child(tmp_path, mode, interrupt=True)
    assert code == 130, payload
    assert payload["error"]["code"] == "operation_cancelled"
    assert payload["data"]["interrupted"] is True
    results = payload["data"]["results"]
    expected_names = (
        ["first.csv", "second.csv", "third.csv"] if mode == "partial_interrupt" else ["first.csv", "second.csv"]
    )
    assert [Path(item["input"]).name for item in results] == expected_names
    reserved = {row["name"]: row["task_id"] for row in trace if row["event"] == "reserved"}
    cancelled = {row["task_id"] for row in trace if row["event"] == "cancelled-through-application"}
    assert reserved["second.csv"] in cancelled
    if mode == "partial_interrupt":
        assert results[0]["success"] is True
        _assert_published_csv_matrix(results[0]["output"], "first")
        assert not results[1]["success"] and results[1]["output"] == ""
        assert results[2]["details"]["execution_state"] == "not_started"
        assert "third.csv" not in reserved
        assert [path.name for path in (tmp_path / "published").rglob("*.xlsx")] == ["first.xlsx"]
    else:
        assert all(not item["success"] and item["output"] == "" for item in results)
        assert reserved["first.csv"] in cancelled
        assert not list((tmp_path / "published").rglob("*.xlsx"))
        if mode == "before_worker":
            assert not any(row["event"] == "worker" for row in trace)
    for name in ("first", "second", "third"):
        assert (tmp_path / f"{name}.csv").read_bytes() == f"Name,Amount\n{name},7\n".encode()


@pytest.mark.integration
def test_timeout_retains_real_finalizer_success_and_does_not_start_following_input(tmp_path: Path) -> None:
    code, payload, trace = _real_cli_child(tmp_path, "partial_timeout", interrupt=False)
    assert code == 8, payload
    assert payload["error"]["code"] == "operation_timeout"
    first, second, third = payload["data"]["results"]
    assert first["success"] is True and first["error"] is None
    _assert_published_csv_matrix(first["output"], "first")
    assert second["success"] is False and second["output"] == ""
    assert second["error"]["code"] == "operation_timeout"
    assert third["details"]["execution_state"] == "not_started" and third["output"] == ""
    assert not any(row.get("name") == "third.csv" for row in trace)
    assert [path.name for path in (tmp_path / "published").rglob("*.xlsx")] == ["first.xlsx"]


@pytest.mark.integration
@pytest.mark.parametrize("mode", ["before_publish", "after_publish"])
def test_real_finalizer_commit_order_is_authoritative_at_deadline(tmp_path: Path, mode: str) -> None:
    code, payload, trace = _real_cli_child(tmp_path, mode, interrupt=False)
    assert code == 8, payload
    [item] = payload["data"]["results"]
    assert any(row["event"] == "observed-cancellation" and row["cancelled"] for row in trace)
    if mode == "after_publish":
        assert item["success"] is True and item["error"] is None
        _assert_published_csv_matrix(item["output"], "first")
    else:
        assert item["success"] is False and item["output"] == ""
        assert item["error"]["code"] == "operation_timeout"
        assert not list((tmp_path / "published").rglob("*.xlsx"))


@pytest.mark.integration
def test_bounded_drain_retains_cancelled_scope_until_delayed_real_worker_finishes(tmp_path, monkeypatch):
    from docwen_application.controller import ApplicationController
    from docwen_bundle.runtime_factory import create_runtime_port
    from docwen_cli import batch_execution
    from docwen_cli.commands.execution_request import file_ref_for_runtime
    from docwen_cli.execution_deadline import ExecutionDeadline
    from docwen_core.models.request import ConversionRequest, OutputPolicy
    from docwen_runtime.config import ConfigLoader

    source = tmp_path / "delayed.csv"
    source.write_bytes(b"Name,Amount\nlate,7\n")
    public = tmp_path / "published"
    public.mkdir()
    monkeypatch.setenv("DOCWEN_RUNTIME_ROOT", str(tmp_path / "runtime"))
    loader = ConfigLoader(user_dir=tmp_path / "config", runtime_overrides={"logger": {"console_enable": False}})
    controller = ApplicationController(runtime_port=create_runtime_port(config_loader=loader))
    controller.start()
    request = ConversionRequest(
        request_id="retained-worker-scope",
        input_refs=[file_ref_for_runtime(str(source))],
        target_format="xlsx",
        output_policy=OutputPolicy(output_dir=str(public)),
    )
    release = threading.Event()
    started = threading.Event()
    finished = threading.Event()
    results = []
    original_execute = controller.execute_single
    original_wait = batch_execution.wait
    injected = False

    def delayed_execute(actual):
        started.set()
        try:
            assert release.wait(5)
            result = original_execute(actual)
            results.append(result)
            return result
        finally:
            finished.set()

    def interrupt_once(*args, **kwargs):
        nonlocal injected
        if not injected:
            assert started.wait(2)
            injected = True
            raise KeyboardInterrupt
        return original_wait(*args, **kwargs)

    monkeypatch.setattr(controller, "execute_single", delayed_execute)
    monkeypatch.setattr(batch_execution, "wait", interrupt_once)
    monkeypatch.setattr(batch_execution, "_CANCELLATION_WAIT_SECONDS", 0.05)
    deadline = ExecutionDeadline(controller, 30).start()
    try:
        batch = batch_execution.execute_batch_requests(
            controller, [(0, request)], deadline=deadline, max_workers=1, stop_on_error=True
        )
        deadline.close()
        assert batch.interrupted
        assert batch.results[0].metrics.extra["execution_state"] == "cancellation_unconfirmed"
        release.set()
        assert finished.wait(5)
        [result] = results
        assert not result.success and result.error.error_type == "cancelled"
        assert not result.artifacts
        assert not list(public.iterdir())
        assert source.read_bytes() == b"Name,Amount\nlate,7\n"
    finally:
        release.set()
        finished.wait(5)
        deadline.close()
        controller.stop()


@pytest.mark.integration
@pytest.mark.linux_only
@pytest.mark.skipif(sys.platform != "linux", reason="Real Linux SIGINT; not a Windows Ctrl+C claim")
@pytest.mark.parametrize("mode", ["submit_after_enqueue", "submit_after_commit", "after_submit_return"])
def test_sigint_inside_submit_and_before_registration_keeps_owned_results(tmp_path: Path, mode: str) -> None:
    code, payload, trace = _real_cli_child(tmp_path, mode, interrupt=True)
    assert code == 130, payload
    assert payload["error"]["code"] == "operation_cancelled"
    first, second = payload["data"]["results"]
    assert Path(first["input"]).name == "first.csv"
    assert Path(second["input"]).name == "second.csv"
    assert second["details"]["execution_state"] == "not_started"
    assert not any(row.get("name") == "second.csv" for row in trace)
    if mode == "submit_after_enqueue":
        assert first["details"]["execution_state"] == "not_started"
        assert not any(row["event"] == "worker" for row in trace)
        assert not list((tmp_path / "published").rglob("*.xlsx"))
    else:
        assert first["success"] and first["error"] is None
        _assert_published_csv_matrix(first["output"], "first")
        assert [path.name for path in (tmp_path / "published").rglob("*.xlsx")] == ["first.xlsx"]


@pytest.mark.integration
@pytest.mark.linux_only
@pytest.mark.skipif(sys.platform != "linux", reason="Real Linux SIGINT; not a Windows Ctrl+C claim")
@pytest.mark.parametrize("mode", ["prepare_return", "deadline_return", "alias_published"])
def test_reservation_return_interrupt_releases_exact_owner_and_controller_is_reusable(
    tmp_path: Path, mode: str
) -> None:
    code, payload, trace = _real_cli_child(tmp_path, mode, interrupt=True)
    assert code == 130, payload
    assert payload["error"]["code"] == "operation_cancelled"
    first, second = payload["data"]["results"]
    assert first["details"]["execution_state"] == "not_started"
    assert second["details"]["execution_state"] == "not_started"
    assert not list((tmp_path / "published").iterdir())
    assert not any(row["event"] == "worker" for row in trace)
    [reused] = [row for row in trace if row["event"] == "controller-reused"]
    assert reused["before"] == reused["after"] == []
    assert reused["success"]
    [output] = reused["outputs"]
    _assert_published_csv_matrix(output, "first")
    assert (tmp_path / "first.csv").read_bytes() == b"Name,Amount\nfirst,7\n"
