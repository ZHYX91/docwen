"""Focused GUI coverage for template-management integration."""

from __future__ import annotations

import pytest

from docwen_gui.widgets.template_selector import TemplateItemDetails, TemplateSelector
from docwen_gui.widgets.template_selector_tabbed import TabbedTemplateSelector

pytestmark = pytest.mark.gui


def test_tabbed_selector_preserves_registry_order(qapp) -> None:
    selector = TabbedTemplateSelector()
    selector.load_templates(
        "docx",
        ["Third", "First", "Second"],
    )
    page = selector.get_selector("docx")
    assert page is not None
    assert [page._list.item(index).text() for index in range(page._list.count())] == [
        "Third",
        "First",
        "Second",
    ]
    selector.deleteLater()


def test_identity_survives_display_rename_and_same_names(qapp):
    selector = TabbedTemplateSelector()
    selector.load_templates(
        "docx",
        ["id-a", "id-b"],
        details={
            "id-a": TemplateItemDetails(display_name="Report", resource_id="id-a"),
            "id-b": TemplateItemDetails(display_name="Report", resource_id="id-b"),
        },
    )
    page = selector.get_selector("docx")
    assert page is not None
    page.select_template("id-b", selection_source="user")
    selector.load_templates(
        "docx",
        ["id-b", "id-a"],
        details={
            "id-a": TemplateItemDetails(display_name="Report", resource_id="id-a"),
            "id-b": TemplateItemDetails(display_name="Renamed", resource_id="id-b"),
        },
    )
    assert selector.get_selected_template_resource() == ("docx", "id-b")
    assert page._list.currentItem().text() == "Renamed"
    selector.deleteLater()


def test_disabled_current_template_never_falls_back_on_refresh_or_tab_switch(qapp):
    selector = TabbedTemplateSelector()
    selector.set_defaults({"docx": "other"})
    selector.load_templates("docx", ["chosen", "other"])
    page = selector.get_selector("docx")
    assert page is not None
    page.select_template("chosen", selection_source="user")
    selector.load_templates("docx", ["other"])
    assert selector.get_selected_template_resource() is None
    assert not selector.ensure_preferred_selection("docx")
    selector.load_templates("docx", ["other", "new"])
    assert selector.get_selected_template_resource() is None
    selector.deleteLater()


@pytest.fixture
def template_vm(tmp_path):
    from docx import Document

    from docwen_gui.view_models.template_vm import TemplateViewModel
    from docwen_runtime.templates import TemplateManager, TemplateRegistry, TemplateStateStore

    builtin = tmp_path / "builtin"
    user = tmp_path / "data" / "templates"
    builtin.mkdir()
    user.mkdir(parents=True)
    Document().save(str(builtin / "Report.docx"))
    state = TemplateStateStore(tmp_path / "data" / "template-state.json", user_dir=user)
    registry = TemplateRegistry(builtin, extra_paths=[user], state_store=state, managed_user_dir=user)
    return TemplateViewModel(manager=TemplateManager(registry, state_store=state, user_dir=user))


def test_management_page_toggles_and_copies_through_shared_model(qapp, template_vm):
    from PySide6.QtCore import Qt

    from docwen_gui.widgets.settings.templates_tab import TemplatesTab

    page = TemplatesTab(view_model=template_vm)
    widget = page._lists["docx"]
    original = template_vm.templates[0]
    widget.item(0).setCheckState(Qt.CheckState.Unchecked)
    assert not template_vm.enabled[original.id]
    copied = template_vm.manager.copy_builtin_as_custom(original.id, custom_name="Mine")
    template_vm.refresh()
    assert widget.count() == 2
    page.focus_template("docx", copied.id)
    assert page._edit_button.isEnabled()
    assert page._copy_button.isHidden()
    page._set_default()
    assert template_vm.defaults["docx"] == copied.id
    page.deleteLater()


def test_settings_has_shared_template_page_and_no_text_selector(qapp, template_vm):
    from docwen_gui.widgets.settings.dialog import SettingsDialog
    from docwen_gui.widgets.settings.templates_tab import TemplatesTab

    dialog = SettingsDialog(template_view_model=template_vm)
    assert isinstance(dialog._tabs["templates"], TemplatesTab)
    assert dialog.activate_section("templates")
    assert not dialog._tabs["text"].findChildren(TabbedTemplateSelector)
    dialog.focus_template("docx", template_vm.templates[0].id)
    assert dialog._tabs["templates"]._lists["docx"].currentRow() == 0
    dialog.close()


def test_template_selector_exposes_management_actions(qapp) -> None:
    selector = TemplateSelector(template_type="docx")
    assert selector._empty_manage_button.text()
    selector.add_templates(["Default"])
    assert selector._manage_button.text()
    assert selector._manage_button.isVisible() or not selector.isVisible()
    selector.deleteLater()


def test_runtime_discovery_before_gui_still_initializes_default(qapp, template_vm):
    discovered = template_vm.manager.list_templates()
    assert template_vm.manager.state_store.path.exists()
    template_vm.refresh()
    assert template_vm.defaults["docx"] == discovered[0].id
    template_vm.manager.set_enabled(discovered[0].id, False)
    template_vm.refresh()
    assert template_vm.defaults["docx"] is None


def test_template_readiness_updates_existing_generation_button(qapp):
    from docwen_gui.view_models.action_area_vm import ActionAreaViewModel
    from docwen_gui.widgets.action_area import ActionArea

    vm = ActionAreaViewModel()
    area = ActionArea(view_model=vm)
    vm.setup_for_md_to_document("sample.md")
    button = area.convert_docx_button
    assert button is not None
    vm.set_template_ready(False)
    assert area.convert_docx_button is button
    assert not button.isEnabled()
    assert button.toolTip()
    area.deleteLater()


def test_text_output_format_remains_editable_without_template_selector(qapp, template_vm):
    from docwen_gui.widgets.settings.dialog import SettingsDialog
    from docwen_gui.widgets.settings.text_tab import TextTab

    dialog = SettingsDialog(template_view_model=template_vm)
    page = dialog._tabs["text"]
    assert isinstance(page, TextTab)
    combo = page._output_format_combo
    combo.setCurrentIndex(combo.findData("xlsx"))
    assert page._vm.config.gui.md_default_template == "xlsx"
    combo.setCurrentIndex(combo.findData("docx"))
    assert page._vm.config.gui.md_default_template == "docx"
    dialog.close()


def test_template_import_reports_partial_success_and_cancelled_conflict(qapp, template_vm, tmp_path, monkeypatch):
    from docx import Document
    from PySide6.QtWidgets import QFileDialog

    from docwen_gui.dialogs import feedback
    from docwen_gui.i18n import t
    from docwen_gui.widgets.settings.templates_tab import TemplatesTab

    template_vm.refresh()
    existing = template_vm.manager.copy_builtin_as_custom(template_vm.templates[0].id, custom_name="Existing")
    existing_bytes = existing.path.read_bytes()
    valid = tmp_path / "New.docx"
    conflict = tmp_path / "Existing.docx"
    invalid = tmp_path / "invalid.docx"
    Document().save(str(valid))
    Document().save(str(conflict))
    invalid.write_text("not an Office package", encoding="utf-8")
    page = TemplatesTab(view_model=template_vm)
    monkeypatch.setattr(
        QFileDialog, "getOpenFileNames", lambda *args: ([str(p) for p in (valid, invalid, conflict)], "")
    )
    reports = []
    choices = []

    def cancel_conflict(*args, **kwargs):
        choices.append(kwargs)
        return "cancel"

    monkeypatch.setattr(feedback, "choose", cancel_conflict)
    monkeypatch.setattr(feedback, "warn", lambda title, message, **kwargs: reports.append((message, kwargs)))
    page._import_button.click()

    summary = t("settings.templates.import_summary", succeeded=1, failed=1, cancelled=1)
    assert page._import_summary.text() == summary
    assert not page._import_summary.isHidden()
    assert len(reports) == 1
    assert reports[0][0] == summary
    assert "invalid.docx" in reports[0][1]["details"]
    assert "New.docx" not in reports[0][1]["details"]
    assert reports[0][1]["copyable"] is True
    assert reports[0][1]["parent"] is page
    assert choices[0]["default"] == "keep"
    assert existing.path.read_bytes() == existing_bytes
    imported = [item for item in template_vm.templates if item.name == "New"]
    assert len(imported) == 1
    assert imported[0].path.read_bytes() == valid.read_bytes()
    assert page._selected_id() == imported[0].id
    page.deleteLater()


@pytest.mark.parametrize("valid_input", [True, False])
def test_template_import_success_is_inline_and_total_failure_has_details(
    qapp, template_vm, tmp_path, monkeypatch, valid_input
):
    from docx import Document
    from PySide6.QtWidgets import QFileDialog

    from docwen_gui.dialogs import feedback
    from docwen_gui.i18n import t
    from docwen_gui.widgets.settings.templates_tab import TemplatesTab

    source = tmp_path / "Import.docx"
    if valid_input:
        Document().save(str(source))
    else:
        source.write_text("invalid", encoding="utf-8")
    page = TemplatesTab(view_model=template_vm)
    reports = []
    monkeypatch.setattr(QFileDialog, "getOpenFileNames", lambda *args: ([str(source)], ""))
    monkeypatch.setattr(feedback, "error", lambda title, message, **kwargs: reports.append(kwargs))
    page._import_button.click()
    assert page._import_summary.text() == t(
        "settings.templates.import_summary", succeeded=int(valid_input), failed=int(not valid_input), cancelled=0
    )
    assert len(reports) == int(not valid_input)
    if not valid_input:
        assert "Import.docx" in reports[0]["details"]
        assert reports[0]["copyable"] is True
    page.deleteLater()
