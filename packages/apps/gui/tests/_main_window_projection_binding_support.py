"""Tests for MainWindow projection binding.

Validates that the MainWindow correctly binds ``ui_projection_changed``
to widget visibility and right-panel stack switching under the current GUI
behavior contract.

These tests use a real MainWindow with an offscreen QApplication so the
widget tree is fully constructed.  Projection changes are simulated by
emitting ``ui_projection_changed`` directly (bypassing file add/select
which requires runtime services).
"""

from __future__ import annotations

from pathlib import Path
from types import SimpleNamespace

import pytest
from PySide6.QtCore import QMimeData, QPointF, Qt, QUrl
from PySide6.QtGui import QDropEvent
from PySide6.QtWidgets import QApplication, QGridLayout, QStackedWidget, QWidget
from tests.support.gui_vm_fakes import FakeController, optimization_capability_projection

from docwen_core.models.file_ref import FileRef
from docwen_gui.view_models.input_area_vm import _BATCH_SCAN_LIMIT
from docwen_gui.view_models.interaction import (
    ConversionContext,
    MainWindowUiProjection,
    RightPanelSlot,
    TemplateContext,
)

pytestmark = pytest.mark.gui


@pytest.fixture
def window(qapp):
    from docwen_gui.main_window import MainWindow
    from docwen_gui.view_models.main_window_vm import MainWindowViewModel

    vm = MainWindowViewModel(controller=FakeController())  # type: ignore[arg-type]
    w = MainWindow(view_model=vm)
    w.setup_ui()
    yield w
    w.close()


@pytest.fixture
def right_frame(window):
    return window.findChild(QWidget, "rightPanelFrame")


@pytest.fixture
def left_frame(window):
    return window.findChild(QWidget, "leftPanelFrame")


@pytest.fixture
def right_stack(window, right_frame):
    return right_frame.findChild(QStackedWidget) if right_frame else None


def _emit(window, projection: MainWindowUiProjection) -> None:
    window._view_model.ui_projection_changed.emit(projection)


def _root_grid(window) -> QGridLayout:
    central = window.findChild(QWidget, "centralContainer")
    assert central is not None
    layout = central.layout()
    assert isinstance(layout, QGridLayout)
    return layout


class _FakeConfigPort:
    def __init__(self, values: dict[str, object]) -> None:
        self._values = values

    def get(self, key: str, default: object = None) -> object:
        return self._values.get(key, default)


class _FakeController:
    def __init__(self, values: dict[str, object]) -> None:
        self.config_port = _FakeConfigPort(values)

    def stop(self) -> None:
        pass


def _make_window_with_config(qapp, values: dict[str, object]):
    from docwen_gui.main_window import MainWindow
    from docwen_gui.view_models.main_window_vm import MainWindowViewModel

    vm = MainWindowViewModel(controller=_FakeController(values))  # type: ignore[arg-type]
    w = MainWindow(view_model=vm)
    w.setup_ui()
    return w


_DOCX_TEMPLATE_ID = f"template.docx.{'a' * 64}"

_XLSX_TEMPLATE_ID = f"template.xlsx.{'b' * 64}"


def _load_request_templates(window, *, docx: bool = True, xlsx: bool = True) -> None:
    from docwen_gui.widgets.template_selector import TemplateItemDetails

    names = {
        "docx": ["Corporate Report"] if docx else [],
        "xlsx": ["Budget"] if xlsx else [],
    }
    details = {
        "docx": {
            "Corporate Report": TemplateItemDetails(resource_id=_DOCX_TEMPLATE_ID),
        }
        if docx
        else {},
        "xlsx": {
            "Budget": TemplateItemDetails(resource_id=_XLSX_TEMPLATE_ID),
        }
        if xlsx
        else {},
    }
    window._template_selector.load_all_templates(names, details=details)


_PROJECTION_HIDDEN = MainWindowUiProjection(
    left_panel_visible=False,
    right_panel_visible=False,
    right_panel_slot=RightPanelSlot.NONE,
    center_action_visible=False,
    info_area_visible=True,
    conversion_context=None,
    template_context=None,
)

_PROJECTION_TEMPLATE = MainWindowUiProjection(
    left_panel_visible=False,
    right_panel_visible=True,
    right_panel_slot=RightPanelSlot.TEMPLATE,
    center_action_visible=True,
    info_area_visible=True,
    conversion_context=None,
    template_context=TemplateContext(file_path="/tmp/test.md"),
)

_PROJECTION_CONVERSION = MainWindowUiProjection(
    left_panel_visible=False,
    right_panel_visible=True,
    right_panel_slot=RightPanelSlot.CONVERSION,
    center_action_visible=True,
    info_area_visible=True,
    conversion_context=ConversionContext(category="document", current_format="docx", file_path="/tmp/test.docx"),
    template_context=None,
)

_PROJECTION_BATCH_HIDDEN_RIGHT = MainWindowUiProjection(
    left_panel_visible=True,
    right_panel_visible=False,
    right_panel_slot=RightPanelSlot.NONE,
    center_action_visible=False,
    info_area_visible=True,
    conversion_context=None,
    template_context=None,
)


def _file_ref(path: str, category: str, fmt: str) -> FileRef:
    return FileRef(path=path, category=category, format=fmt)


def _write_format_fixture(path: Path, fmt: str) -> None:
    """Write structurally valid content for formats with strict containers."""

    if fmt == "docx":
        from docx import Document

        Document().save(str(path))
        return
    if fmt == "xlsx":
        from openpyxl import Workbook

        Workbook().save(str(path))
        return
    if fmt == "pptx":
        from pptx import Presentation

        Presentation().save(str(path))
        return
    if fmt == "epub":
        from ebooklib import epub

        book = epub.EpubBook()
        book.set_identifier("docwen-gui-request-contract")
        book.set_title("DocWen GUI request contract")
        book.set_language("en")
        chapter = epub.EpubHtml(title="Probe", file_name="probe.xhtml", lang="en")
        chapter.content = "<h1>Probe</h1>"
        book.add_item(chapter)
        book.spine = ["nav", chapter]
        book.add_item(epub.EpubNcx())
        book.add_item(epub.EpubNav())
        epub.write_epub(str(path), book)
        return
    if fmt == "html":
        path.write_text("<html><body><h1>Probe</h1></body></html>", encoding="utf-8")
        return
    if fmt == "tsv":
        path.write_text("name\tvalue\nprobe\t1\n", encoding="utf-8")
        return
    if fmt == "pdf":
        path.write_bytes(b"%PDF-1.4\n% deterministic request fixture\n")
        return
    if fmt == "png":
        from PIL import Image

        Image.new("RGB", (2, 2), "white").save(path)
        return
    if fmt == "jpeg":
        from PIL import Image

        Image.new("RGB", (2, 2), "white").save(path, format="JPEG")
        return
    path.write_text("focused request-binding fixture", encoding="utf-8")


def _bind_admitted_ref(window, path: Path, category: str, fmt: str) -> FileRef:
    """Bind a concrete upstream-admitted ref for request-only GUI tests.

    These tests exercise option and request projection, not Core inspection.
    Container-admission behavior is covered separately, so the request builder
    receives the same concrete ``FileRef`` it would get after successful
    ingress instead of reconstructing a route from the filename suffix.
    """

    from docwen_core.models import (
        FILE_INSPECTION_METADATA_KEY,
        AdmissionDecision,
        DetectionConfidence,
        DetectionMethod,
        FileInspection,
        FormatRelation,
        StructureStatus,
    )
    from docwen_gui.main_window import _normalize_path

    stat = path.stat()
    container_formats = {"doc", "docx", "epub", "pptx", "wps", "xlsx"}
    text_formats = {"html", "markdown", "tsv", "txt"}
    method = DetectionMethod.CONTAINER if fmt in container_formats else DetectionMethod.SIGNATURE
    if fmt in text_formats:
        method = DetectionMethod.TEXT_SNIFF
    inspection = FileInspection(
        file_path=str(path.resolve()),
        size_bytes=stat.st_size,
        mtime_ns=stat.st_mtime_ns,
        extension=path.suffix.lower(),
        declared_format=fmt,
        declared_category=category,
        detected_format=fmt,
        detected_category=category,
        workflow_category=category,
        detection_method=method,
        confidence=DetectionConfidence.CERTAIN,
        structure_status=StructureStatus.VALID if fmt in container_formats else StructureStatus.NOT_APPLICABLE,
        relation=FormatRelation.EXACT_MATCH,
        decision=AdmissionDecision.ALLOW,
        declared_supported=True,
        detected_supported=True,
    )
    ref = FileRef(
        path=str(path),
        category=category,
        format=fmt,
        size_bytes=stat.st_size,
        metadata={FILE_INSPECTION_METADATA_KEY: inspection.to_dict()},
    )
    window._view_model._files = [ref]
    window._file_contexts = {_normalize_path(str(path)): (fmt, category)}
    return ref


__all__ = (
    "_BATCH_SCAN_LIMIT",
    "_DOCX_TEMPLATE_ID",
    "_PROJECTION_BATCH_HIDDEN_RIGHT",
    "_PROJECTION_CONVERSION",
    "_PROJECTION_HIDDEN",
    "_PROJECTION_TEMPLATE",
    "_XLSX_TEMPLATE_ID",
    "ConversionContext",
    "FakeController",
    "FileRef",
    "MainWindowUiProjection",
    "Path",
    "QApplication",
    "QDropEvent",
    "QMimeData",
    "QPointF",
    "QUrl",
    "QWidget",
    "Qt",
    "RightPanelSlot",
    "SimpleNamespace",
    "_bind_admitted_ref",
    "_emit",
    "_file_ref",
    "_load_request_templates",
    "_make_window_with_config",
    "_root_grid",
    "_write_format_fixture",
    "left_frame",
    "optimization_capability_projection",
    "pytest",
    "pytestmark",
    "right_frame",
    "right_stack",
    "window",
)
