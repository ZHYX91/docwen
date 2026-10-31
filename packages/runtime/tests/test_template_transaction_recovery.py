"""Template transaction cleanup warnings, compensation death and durable rollback."""

from __future__ import annotations

import zipfile
from pathlib import Path

import pytest

from docwen_runtime.templates import (
    TemplateManager,
    TemplateRegistry,
    TemplateStateStore,
)

from ._template_manager_support import (
    _manager,
)

pytestmark = pytest.mark.integration


def _saved_template_docx(path: Path, text: str) -> bytes:
    from xml.etree import ElementTree

    from docx import Document

    document = Document()
    document.add_paragraph(text)
    path.parent.mkdir(parents=True, exist_ok=True)
    document.save(str(path))
    assert [paragraph.text for paragraph in Document(str(path)).paragraphs] == [text]
    with zipfile.ZipFile(path) as package:
        root = ElementTree.fromstring(package.read("word/document.xml"))
    assert [node.text for node in root.iter("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}t")] == [
        text
    ]
    return path.read_bytes()


def _reopen_template_manager(builtin_dir: Path, user_dir: Path) -> TemplateManager:
    state = TemplateStateStore(user_dir.parent / "template-state.json", user_dir=user_dir)
    return TemplateManager(
        TemplateRegistry(
            builtin_dir,
            extra_paths=[user_dir],
            state_store=state,
            managed_user_dir=user_dir,
            builtin_dir=builtin_dir,
        ),
        state_store=state,
        user_dir=user_dir,
    )


@pytest.mark.parametrize("boundary", ["unlink", "directory_sync"])
def test_committed_import_cleanup_warning_preserves_file_identity_and_does_not_reimport(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch, caplog: pytest.LogCaptureFixture, boundary: str
) -> None:
    import hashlib
    import json
    import logging

    from tests.support.transaction_faults import fail_committed_journal_cleanup_once

    manager, builtin_dir, user_dir = _manager(tmp_path)
    earlier = tmp_path / "incoming" / "Earlier.docx"
    _saved_template_docx(earlier, "Existing default content")
    first = manager.import_template(earlier)
    manager.set_default(first.id)
    source = tmp_path / "incoming" / "Report.docx"
    source_bytes = _saved_template_docx(source, "Committed report body")
    journal = user_dir.parent / ".docwen-config.transaction.json"
    before = manager.state_store.load()
    with caplog.at_level(logging.WARNING):
        fault = fail_committed_journal_cleanup_once(monkeypatch, journal, boundary=boundary)
        imported = manager.import_template(source)
    assert fault.injected is True
    assert fault.committed_envelope is not None
    assert json.loads(fault.committed_envelope)["payload"]["state"] == "COMMITTED"
    assert fault.removed is (boundary == "directory_sync")
    assert journal.exists() is (boundary == "unlink")
    assert imported.name == "Report"
    assert imported.path == user_dir / "Report.docx"
    assert hashlib.sha256(imported.path.read_bytes()).digest() == hashlib.sha256(source_bytes).digest()
    assert source.read_bytes() == source_bytes
    saved = manager.state_store.load()
    assert saved["defaults"] == before["defaults"]
    assert saved["enabled"][imported.id] is True
    assert saved["order"]["docx"] == [first.id, imported.id]
    assert saved["user_identities"]["report.docx"]["id"] == imported.id
    assert saved["user_identities"]["report.docx"]["filename"] == "Report.docx"
    warnings = manager.take_cleanup_warnings()
    assert len(warnings) == 1
    assert "committed journal" in warnings[0]
    assert manager.take_cleanup_warnings() == ()
    assert "Committed template journal cleanup deferred" in caplog.text
    # Retry cleanup via the catalog, without repeating the successful import.
    assert [(item.id, item.name) for item in manager.list_templates()] == [
        (first.id, "Earlier"),
        (imported.id, "Report"),
    ]
    assert not journal.exists()
    reopened = _reopen_template_manager(builtin_dir, user_dir)
    assert [(item.id, item.name) for item in reopened.list_templates()] == [
        (first.id, "Earlier"),
        (imported.id, "Report"),
    ]
    assert sorted(path.name for path in user_dir.iterdir()) == ["Earlier.docx", "Report.docx"]
    assert reopened.state_store.load()["defaults"]["docx"] == first.id
    assert reopened.state_store.load()["order"]["docx"] == [first.id, imported.id]
    assert imported.path.read_bytes() == source_bytes


@pytest.mark.parametrize("operation", ["rename", "delete"])
def test_committed_template_cleanup_failure_keeps_rename_or_delete_result(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch, operation: str
) -> None:
    from tests.support.transaction_faults import fail_committed_journal_cleanup_once

    manager, builtin_dir, user_dir = _manager(tmp_path)
    source = tmp_path / "source" / "Report.docx"
    source_bytes = _saved_template_docx(source, "Stable identity body")
    imported = manager.import_template(source)
    manager.set_default(imported.id)
    journal = user_dir.parent / ".docwen-config.transaction.json"
    # This owned local directory is a filesystem substitute only for the trash
    # service. It is not evidence of a native desktop recycle-bin integration.
    trash = tmp_path / "owned-trash"
    trash.mkdir()
    monkeypatch.setattr(
        "docwen_runtime.templates.manager._move_to_recycle_bin",
        lambda path: path.rename(trash / path.name),
    )
    fault = fail_committed_journal_cleanup_once(monkeypatch, journal)
    if operation == "rename":
        renamed = manager.rename_custom(imported.id, "Renamed")
        assert renamed.id == imported.id
        assert renamed.path == user_dir / "Renamed.docx"
        assert renamed.path.read_bytes() == source_bytes
    else:
        manager.delete_custom(imported.id)
        assert (trash / "Report.docx").read_bytes() == source_bytes
    assert fault.injected and fault.removed
    assert not journal.exists()
    assert len(manager.take_cleanup_warnings()) == 1
    reopened = _reopen_template_manager(builtin_dir, user_dir)
    templates = reopened.list_templates()
    state = reopened.state_store.load()
    if operation == "rename":
        assert [(item.id, item.name) for item in templates] == [(imported.id, "Renamed")]
        assert state["defaults"]["docx"] == imported.id
        assert state["order"]["docx"] == [imported.id]
        assert state["enabled"][imported.id] is True
    else:
        assert templates == []
        assert state["defaults"]["docx"] is None
        assert state["order"]["docx"] == []
        assert state["enabled"][imported.id] is False
        assert state["user_identities"] == {}
    assert not imported.path.exists()
    assert source.read_bytes() == source_bytes


@pytest.mark.parametrize("failure_phase", ["prepare", "mutation"])
def test_template_prepare_or_mutation_failure_is_not_committed_cleanup_success(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch, failure_phase: str
) -> None:
    import json

    from docwen_runtime import toml_io

    manager, _builtin_dir, user_dir = _manager(tmp_path)
    manager.list_templates()
    source = tmp_path / "incoming" / "Report.docx"
    original_source = _saved_template_docx(source, "Rejected operation")
    state_bytes = manager.state_store.path.read_bytes() if manager.state_store.path.exists() else None
    journal = user_dir.parent / ".docwen-config.transaction.json"
    original_replace = toml_io.os.replace
    failed = False

    def replace(staged, destination) -> None:
        nonlocal failed
        candidate = Path(destination)
        should_fail = candidate == user_dir / "Report.docx"
        if failure_phase == "prepare":
            should_fail = candidate == journal and json.loads(Path(staged).read_bytes())["payload"]["state"] == (
                "PREPARED"
            )
        if should_fail and not failed:
            failed = True
            raise OSError(f"injected {failure_phase} failure before replacement")
        original_replace(staged, destination)

    monkeypatch.setattr(toml_io.os, "replace", replace)
    with pytest.raises(OSError, match=f"injected {failure_phase}"):
        manager.import_template(source)
    assert failed
    assert list(user_dir.iterdir()) == []
    assert source.read_bytes() == original_source
    assert manager.take_cleanup_warnings() == ()
    assert (manager.state_store.path.read_bytes() if manager.state_store.path.exists() else None) == state_bytes
    assert manager.list_templates() == []
    assert not journal.exists()


def test_template_commit_barrier_failure_then_compensation_death_recovers_files_and_canonical_state(
    tmp_path: Path,
) -> None:
    import json
    import os
    import stat
    import sys
    import textwrap

    from tests.support.subprocess_runner import run_subprocess

    manager, builtin_dir, user_dir = _manager(tmp_path)
    source = tmp_path / "source" / "Report.docx"
    expected_bytes = _saved_template_docx(source, "Recover this whole template and identity")
    imported = manager.import_template(source)
    imported.path.chmod(0o640)
    expected_mode = stat.S_IMODE(imported.path.stat().st_mode)
    manager.set_default(imported.id)
    before_state = manager.state_store.path.read_bytes()
    script = """
        import json
        import os
        import sys
        from pathlib import Path
        from docwen_runtime import file_transactions as transaction, toml_io
        from docwen_runtime.templates import TemplateManager, TemplateRegistry, TemplateStateStore
        builtin, user = map(Path, sys.argv[1:])
        store = TemplateStateStore(user.parent / 'template-state.json', user_dir=user)
        manager = TemplateManager(
            TemplateRegistry(builtin, extra_paths=[user], state_store=store, managed_user_dir=user),
            state_store=store, user_dir=user,
        )
        identity = manager.list_templates()[0].id
        journal = user.parent / transaction.CONFIG_JOURNAL_NAME
        original_replace = toml_io.os.replace
        original_sync = toml_io._sync_directory
        original_restore = transaction.restore_user_file_preimage
        phase = {'pending': None, 'committed_replaced': False, 'rollback_durable': False}
        def replace(source, destination):
            original_replace(source, destination)
            if Path(destination) == journal:
                state = json.loads(journal.read_bytes())['payload']['state']
                if state == 'COMMITTED':
                    phase['committed_replaced'] = True
                    phase['pending'] = 'commit'
                    phase['committed_state'] = json.loads(store.path.read_bytes())
                    phase['committed_files'] = sorted(path.name for path in user.iterdir())
                elif phase['committed_replaced']:
                    phase['pending'] = 'rollback'
        def sync(path):
            pending = phase.pop('pending', None)
            if pending == 'commit':
                raise OSError('actual template commit replacement followed by failed barrier')
            original_sync(path)
            if pending == 'rollback':
                phase['rollback_durable'] = True
        def restore(before, *, operation, **kwargs):
            original_restore(before, operation=operation, **kwargs)
            phase['first_restored'] = before.path.name
            phase['journal'] = json.loads(journal.read_bytes())
            phase['files_at_death'] = sorted(path.name for path in user.iterdir())
            phase['state_at_death'] = json.loads(store.path.read_bytes())
            (user.parent / 'template-writer-observation.json').write_text(json.dumps(phase), encoding='utf-8')
            os._exit(96)
        toml_io.os.replace = replace
        toml_io._sync_directory = sync
        transaction.restore_user_file_preimage = restore
        manager.rename_custom(identity, 'Renamed')
        raise AssertionError('template compensation did not execute')
    """
    writer = run_subprocess(
        [sys.executable, "-c", textwrap.dedent(script), str(builtin_dir), str(user_dir)],
        env=dict(os.environ),
        timeout=20,
    )
    assert writer.returncode == 96, writer.stderr
    observation = json.loads((user_dir.parent / "template-writer-observation.json").read_text(encoding="utf-8"))
    assert observation["committed_replaced"] is True
    assert observation["committed_files"] == ["Renamed.docx"]
    assert observation["committed_state"]["user_identities"]["renamed.docx"]["id"] == imported.id
    assert observation["first_restored"] == "Renamed.docx"
    assert observation["files_at_death"] == []
    assert observation["state_at_death"]["user_identities"]["renamed.docx"]["id"] == imported.id
    recovery_script = """
        import json
        import stat
        import sys
        from pathlib import Path
        from docwen_runtime.templates import TemplateManager, TemplateRegistry, TemplateStateStore
        builtin, user = map(Path, sys.argv[1:])
        store = TemplateStateStore(user.parent / 'template-state.json', user_dir=user)
        # Inspect the recovered bytes before discovery updates inode cache facts.
        with store.locked():
            path = user / 'Report.docx'
            observed = {
                'state_bytes': store.path.read_bytes().hex(),
                'template_bytes': path.read_bytes().hex() if path.exists() else None,
                'mode': stat.S_IMODE(path.stat().st_mode) if path.exists() else None,
                'files': sorted(item.name for item in user.iterdir()),
                'journal': (user.parent / '.docwen-config.transaction.json').exists(),
            }
        manager = TemplateManager(
            TemplateRegistry(builtin, extra_paths=[user], state_store=store, managed_user_dir=user),
            state_store=store, user_dir=user,
        )
        observed['catalog'] = [(item.id, item.name) for item in manager.list_templates()]
        observed['state'] = store.load()
        (user.parent / 'template-reader-observation.json').write_text(json.dumps(observed), encoding='utf-8')
    """
    reader = run_subprocess(
        [sys.executable, "-c", textwrap.dedent(recovery_script), str(builtin_dir), str(user_dir)],
        env=dict(os.environ),
        timeout=20,
    )
    assert reader.returncode == 0, reader.stderr
    recovered = json.loads((user_dir.parent / "template-reader-observation.json").read_text(encoding="utf-8"))
    assert recovered["state_bytes"] == before_state.hex()
    assert recovered["template_bytes"] == expected_bytes.hex()
    assert recovered["mode"] == expected_mode
    assert recovered["files"] == ["Report.docx"]
    assert recovered["journal"] is False
    assert recovered["catalog"] == [[imported.id, "Report"]]
    assert recovered["state"]["defaults"]["docx"] == imported.id
    assert recovered["state"]["order"]["docx"] == [imported.id]
    assert recovered["state"]["enabled"][imported.id] is True
    assert observation["rollback_durable"] is True
    assert observation["journal"]["payload"]["state"] == "PREPARED"


@pytest.mark.parametrize("rollback_boundary", ["before_replace", "after_replace"])
def test_template_uncertain_commit_does_not_compensate_without_durable_rollback(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch, rollback_boundary: str
) -> None:
    import json

    from docwen_runtime import toml_io

    manager, builtin_dir, user_dir = _manager(tmp_path)
    source = tmp_path / "source" / "Report.docx"
    expected_bytes = _saved_template_docx(source, "Leave a complete generation when rollback cannot be armed")
    imported = manager.import_template(source)
    manager.set_default(imported.id)
    before_state = manager.state_store.path.read_bytes()
    journal = user_dir.parent / ".docwen-config.transaction.json"
    original_replace = toml_io.os.replace
    original_sync = toml_io._sync_directory
    phase = {"committed": False, "pending": None}

    def replace(staged, destination) -> None:
        if Path(destination) == journal:
            state = json.loads(Path(staged).read_bytes())["payload"]["state"]
            if state == "PREPARED" and phase["committed"] and rollback_boundary == "before_replace":
                raise OSError("rollback journal cannot be replaced")
        original_replace(staged, destination)
        if Path(destination) == journal:
            state = json.loads(journal.read_bytes())["payload"]["state"]
            if state == "COMMITTED":
                phase["committed"] = True
                phase["pending"] = "commit"
            elif phase["committed"]:
                phase["pending"] = "rollback"

    def sync(path) -> None:
        pending = phase.pop("pending", None)
        if pending:
            raise OSError(f"{pending} journal directory barrier failed")
        original_sync(path)

    with monkeypatch.context() as fault:
        fault.setattr(toml_io.os, "replace", replace)
        fault.setattr(toml_io, "_sync_directory", sync)
        with pytest.raises(OSError, match="commit journal directory barrier failed") as caught:
            manager.rename_custom(imported.id, "Renamed")
        assert any("no preimages were restored" in note for note in getattr(caught.value, "__notes__", ()))
        assert (user_dir / "Renamed.docx").read_bytes() == expected_bytes
        assert not (user_dir / "Report.docx").exists()
        new_state = manager.state_store.path.read_bytes()
        assert json.loads(new_state)["user_identities"]["renamed.docx"]["id"] == imported.id
        assert json.loads(journal.read_bytes())["payload"]["state"] == (
            "COMMITTED" if rollback_boundary == "before_replace" else "PREPARED"
        )
    reopened = _reopen_template_manager(builtin_dir, user_dir)
    # The next catalog owner must finish recovery before exposing the result.
    with reopened.state_store.locked():
        expected_name = "Renamed" if rollback_boundary == "before_replace" else "Report"
        assert sorted(path.name for path in user_dir.iterdir()) == [f"{expected_name}.docx"]
        assert (user_dir / f"{expected_name}.docx").read_bytes() == expected_bytes
        assert reopened.state_store.path.read_bytes() == (
            new_state if rollback_boundary == "before_replace" else before_state
        )
        assert not journal.exists()
    assert [(item.id, item.name) for item in reopened.list_templates()] == [(imported.id, expected_name)]
    assert reopened.state_store.default_id("docx") == imported.id
    assert reopened.state_store.load()["order"]["docx"] == [imported.id]
