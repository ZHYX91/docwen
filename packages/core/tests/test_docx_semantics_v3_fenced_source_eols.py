"""Hash-authenticated EOL compatibility and negative proofs."""

from __future__ import annotations

from pathlib import Path
from typing import Any, cast
from zipfile import ZipFile

import lxml.etree as etree
import pytest
from docx import Document
from docx.oxml.ns import qn

from docwen_core.docx_semantics_v3 import DocxSemanticsV3Error, DocxSemanticsV3Recovery, DocxSemanticsV3Session

from ._docx_semantics_v3_fenced_source_support import (
    _identity,
    _map_part,
    _mutate_document_xml,
    _mutate_xml_member,
    _write,
)

pytestmark = pytest.mark.contract


@pytest.mark.parametrize("host_shape", ["body-lf", "framing-lf", "both"])
@pytest.mark.parametrize("corruption", ["none", "body-text", "space", "line-count", "uuid", "info", "source-range"])
def test_host_eol_projection_requires_the_original_body_block_and_package_proof(
    tmp_path: Path, host_shape: str, corruption: str
) -> None:
    logical_body = "  alpha\t**literal**\r\nbeta\r\n"
    source = "```rust\r\n" + logical_body + "```\r\n"
    identity = _identity(
        source,
        logical_body,
        fence_character="`",
        opening_length=3,
        opening_prefix="",
        info="rust",
        opening_eol="\r\n",
        body_prefixes=("", ""),
        closing_state="present",
        closing_length=3,
        closing_prefix="",
        closing_suffix="",
        closing_eol="\r\n",
    )
    document = Document()
    paragraph = document.add_paragraph()
    session = DocxSemanticsV3Session(document, source_sha256=identity.source_sha256)
    session.bind_fenced_source(paragraph, identity, logical_body=logical_body)
    output = _write(session, document, tmp_path / "eol.docx")

    def change_body(root) -> None:
        if host_shape in {"body-lf", "both"}:
            for line_break in root.iter(qn("w:cr")):
                line_break.tag = qn("w:br")
                line_break.set(qn("w:type"), "textWrapping")
        if corruption in {"body-text", "space"}:
            text = next(root.iter(qn("w:t")))
            text.text = "tampered" if corruption == "body-text" else " " + (text.text or "")
        elif corruption == "line-count":
            content = next(root.iter(qn("w:sdtContent")))
            run = etree.SubElement(content, qn("w:r"))
            etree.SubElement(run, qn("w:br"))

    _mutate_document_xml(output, change_body)
    with ZipFile(output) as package:
        map_part = _map_part(package)
    if host_shape in {"framing-lf", "both"} or corruption in {"info", "source-range"}:

        def change_map(root) -> None:
            if host_shape in {"framing-lf", "both"}:
                root[0].set("opening_eol", "\n")
                root[0].set("closing_eol", "\n")
            if corruption == "info":
                root[0].set("info_b64", "cHl0aG9u")
            elif corruption == "source-range":
                root[0].set("source_end", str(identity.source_end + 1))

        _mutate_xml_member(output, map_part, change_map)
    if corruption == "uuid":
        _mutate_xml_member(
            output,
            map_part.replace("/item", "/itemProps"),
            lambda root: root.set(
                "{http://schemas.openxmlformats.org/officeDocument/2006/customXml}itemID",
                "{00000000-0000-0000-0000-000000000001}",
            ),
        )
    before = output.read_bytes()
    if corruption == "none":
        loaded = Document(str(output))
        recovery = DocxSemanticsV3Recovery.load(output, loaded)
        current = next(cast(Any, loaded.element).body.iter(qn("w:p")))
        assert recovery.render_fenced_source(current) == source
        assert recovery.fenced_source_identities == (identity,)
    else:
        with pytest.raises(DocxSemanticsV3Error):
            session.prove_package(output)
    assert output.read_bytes() == before


def test_uniform_host_eols_cannot_guess_an_original_mixed_body(tmp_path: Path) -> None:
    body = "alpha\r\nbeta\n"
    source = "```\n" + body + "```\n"
    identity = _identity(
        source,
        body,
        fence_character="`",
        opening_length=3,
        opening_prefix="",
        info="",
        opening_eol="\n",
        body_prefixes=("", ""),
        closing_state="present",
        closing_length=3,
        closing_prefix="",
        closing_suffix="",
        closing_eol="\n",
    )
    document = Document()
    session = DocxSemanticsV3Session(document, source_sha256=identity.source_sha256)
    session.bind_fenced_source(document.add_paragraph(), identity, logical_body=body)
    output = _write(session, document, tmp_path / "mixed.docx")

    def normalize_body(root) -> None:
        for line_break in root.iter(qn("w:cr")):
            line_break.tag = qn("w:br")

    _mutate_document_xml(output, normalize_body)
    with pytest.raises(DocxSemanticsV3Error, match="visible body hash"):
        session.prove_package(output)
