"""Shared comments preserve old XML while anchoring only the requested span."""

from io import BytesIO

import pytest
from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from lxml import etree

from docwen_core.docx_comments import append_docx_comment

pytestmark = pytest.mark.contract


def test_comments_preserve_legacy_package_and_anchor_multiple_runs():
    document = Document()
    old = document.add_comment(document.add_paragraph("original").runs, "old review")
    old._element.set(qn("w:id"), "malformed")
    # The old body's unchanged anchors reserve 0 despite its malformed definition.
    missing = OxmlElement("w:comment")
    document.comments._comments_elm.append(missing)
    for name, kind in (("CommentReference", WD_STYLE_TYPE.PARAGRAPH), ("CommentText", WD_STYLE_TYPE.CHARACTER)):
        document.styles.add_style(name, kind)
    header = document.sections[0].header.paragraphs[0]
    reference = OxmlElement("w:commentReference")
    reference.set(qn("w:id"), " \t+0001\r\n")
    header.add_run()._r.append(reference)
    old_nodes = [etree.tostring(node) for node in document.comments._comments_elm]
    old_styles = [etree.tostring(style.element) for style in document.styles]
    old_header = etree.tostring(header._p)
    old_relations = {key: (rel.reltype, rel.target_ref) for key, rel in document.part.rels.items()}
    paragraph = document.add_paragraph()
    paragraph.add_run("outside ")
    first = paragraph.add_run("selected")
    first.bold = True
    last = paragraph.add_run(" range")
    paragraph.add_run(" outside")
    assert append_docx_comment(document, [first, last], text="第一行\n第二行\t值") == 2
    assert [etree.tostring(node) for node in document.comments._comments_elm][:2] == old_nodes
    assert [etree.tostring(style.element) for style in document.styles][: len(old_styles)] == old_styles
    assert etree.tostring(header._p) == old_header
    assert {
        key: (document.part.rels[key].reltype, document.part.rels[key].target_ref) for key in old_relations
    } == old_relations
    assert first._r.getprevious().tag == qn("w:commentRangeStart")
    assert last._r.getnext().tag == qn("w:commentRangeEnd")
    assert first.bold is True
    stream = BytesIO()
    document.save(stream)
    stream.seek(0)
    reopened = Document(stream)
    assert list(reopened.comments)[2].text == "第一行\n第二行\t值"
    for node in reopened.element.iter(qn("w:rStyle")):
        style_id = node.get(qn("w:val"), "")
        if style_id.startswith("DocWenComment"):
            assert reopened.styles[style_id].type == WD_STYLE_TYPE.CHARACTER
    assert reopened.paragraphs[-1].text == "outside selected range outside"


def test_comment_text_is_literal_and_preserves_boundary_whitespace():
    document = Document()
    runs = document.add_paragraph("target").runs
    append_docx_comment(document, runs, text=" <review>& \n", author="A&B", initials="<> ")
    stream = BytesIO()
    document.save(stream)
    stream.seek(0)
    comment = next(iter(Document(stream).comments))
    assert comment.text == " <review>& \n"
    assert comment.author == "A&B"
    assert comment.initials == "<> "


def test_empty_comment_range_does_not_mutate_package():
    document = Document()
    relationships = dict(document.part.rels)
    styles = etree.tostring(document.styles.element)
    with pytest.raises(ValueError, match="nonempty"):
        append_docx_comment(document, [], text="review")
    assert dict(document.part.rels) == relationships
    assert etree.tostring(document.styles.element) == styles


@pytest.mark.parametrize("case", ["foreign", "header", "reverse", "duplicate", "detached", "invalid-text"])
def test_invalid_comment_request_preserves_document(case):
    document = Document()
    paragraph = document.add_paragraph()
    first, last = paragraph.add_run("first"), paragraph.add_run("last")
    choices = {
        "foreign": [Document().add_paragraph("foreign").runs[0]],
        "header": [document.sections[0].header.paragraphs[0].add_run("header")],
        "reverse": [last, first],
        "duplicate": [first, first],
        "detached": [first],
        "invalid-text": [first, last],
    }
    if case == "detached":
        paragraph._p.remove(first._r)
    relationships = dict(document.part.rels)
    body = etree.tostring(document.element)
    styles = etree.tostring(document.styles.element)
    with pytest.raises(ValueError):
        append_docx_comment(document, choices[case], text="invalid\0" if case == "invalid-text" else "review")
    assert etree.tostring(document.element) == body
    assert etree.tostring(document.styles.element) == styles
    assert dict(document.part.rels) == relationships


def test_anchor_commit_failure_rolls_back_new_comment_styles_and_part(monkeypatch):
    from docx.oxml.text.run import CT_R

    document = Document()
    run = document.add_paragraph("target").runs[0]
    relationships = dict(document.part.rels)
    related_parts = dict(document.part.related_parts)
    body = etree.tostring(document.element)
    styles = etree.tostring(document.styles.element)
    original = CT_R.addnext

    def fail_after_insert(self, element):
        original(self, element)
        if element.tag == qn("w:commentRangeEnd"):
            raise RuntimeError("injected anchor failure")

    monkeypatch.setattr(CT_R, "addnext", fail_after_insert)
    with pytest.raises(RuntimeError, match="anchor failure"):
        append_docx_comment(document, [run], text="review")
    assert etree.tostring(document.element) == body
    assert etree.tostring(document.styles.element) == styles
    assert dict(document.part.rels) == relationships
    assert document.part.related_parts == related_parts
    monkeypatch.setattr(CT_R, "addnext", original)
    # Reuse the released rId for an external hyperlink first. A stale internal
    # target mapping must not make that external relationship resolve to comments.
    from docx.opc.constants import RELATIONSHIP_TYPE as RT

    hyperlink_id = document.part.relate_to("https://example.test", RT.HYPERLINK, is_external=True)
    assert hyperlink_id not in document.part.related_parts
    assert append_docx_comment(document, [run], text="retry") == 0
    stream = BytesIO()
    document.save(stream)
    stream.seek(0)
    reopened = Document(stream)
    assert [comment.text for comment in reopened.comments] == ["retry"]
    assert reopened.part.rels[hyperlink_id].is_external
    assert reopened.part.rels[hyperlink_id].target_ref == "https://example.test"


@pytest.mark.parametrize("include_main", [False, True])
def test_textbox_story_cannot_receive_main_document_comment(include_main):
    from docx.text.paragraph import Paragraph

    document = Document()
    main = document.add_paragraph("main").runs[0]
    textbox = OxmlElement("w:txbxContent")
    paragraph_xml = OxmlElement("w:p")
    textbox.append(paragraph_xml)
    main._r.append(textbox)
    textbox_run = Paragraph(paragraph_xml, document).add_run("textbox")
    body = etree.tostring(document.element)
    relationships = dict(document.part.rels)
    related_parts = dict(document.part.related_parts)
    runs = [main, textbox_run] if include_main else [textbox_run]
    with pytest.raises(ValueError, match="main story"):
        append_docx_comment(document, runs, text="review")
    assert etree.tostring(document.element) == body
    assert dict(document.part.rels) == relationships
    assert document.part.related_parts == related_parts
