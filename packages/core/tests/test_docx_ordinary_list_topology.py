"""Fresh native-list packages retain recursive numbering topology refusals."""

from __future__ import annotations

import hashlib
from copy import deepcopy
from pathlib import Path
from typing import Any
from zipfile import ZipFile

import pytest
from docx import Document
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph
from lxml import etree

from docwen_core.docx_numbering_import import DocxNumberingImportError
from docwen_core.docx_resolved_numbering_recovery import ResolvedNumberingV4Recovery
from docwen_core.docx_semantics_v3 import (
    TARGET_MAP_NAMESPACE,
    DocxSemanticsV3Recovery,
    DocxSemanticsV3Session,
)

pytestmark = pytest.mark.contract

_TITLE = "Owned heading"
_OUTSIDE = "Outside."
_TARGET = "native-list-heading"
_MC = "http://schemas.openxmlformats.org/markup-compatibility/2006"
_BAD_TOPOLOGIES = [
    ("level_styleLink", "style-linked numbering"),
    ("level_numStyleLink", "style-linked numbering"),
    ("level_pStyle", "exact ancestry"),
    ("level_start", "exact ancestry"),
    ("level_mc", "markup-compatibility topology"),
    ("reference_pStyle", "exact ancestry"),
    ("reference_mc", "markup-compatibility topology"),
    ("override_styleLink", "style-linked numbering"),
]


def _value(parent: Any, tag: str, value: str) -> Any:
    child = etree.SubElement(parent, qn(tag))
    child.set(qn("w:val"), value)
    return child


def _write_native_heading(path: Path, *, num_id: int, number_format: str, start: int) -> Path:
    """Use the real v3 target/bookmark writer, with independent one-level XML."""
    document = Document()
    numbering = document.part.numbering_part.element
    for child in list(numbering):
        numbering.remove(child)
    abstract = etree.SubElement(numbering, qn("w:abstractNum"))
    abstract.set(qn("w:abstractNumId"), "0")
    _value(abstract, "w:multiLevelType", "hybridMultilevel")
    level = etree.SubElement(abstract, qn("w:lvl"))
    level.set(qn("w:ilvl"), "0")
    _value(level, "w:start", str(start))
    _value(level, "w:numFmt", number_format)
    _value(level, "w:lvlText", "•" if number_format == "bullet" else "%1.")
    _value(level, "w:suff", "space")
    # Legal presentation children are not themselves a reason to reject a list.
    r_pr = etree.SubElement(level, qn("w:rPr"))
    _value(r_pr, "w:color", "102030")
    instance = etree.SubElement(numbering, qn("w:num"))
    instance.set(qn("w:numId"), str(num_id))
    _value(instance, "w:abstractNumId", "0")
    paragraph = document.add_heading(_TITLE, level=1)
    num_pr = etree.SubElement(paragraph._p.get_or_add_pPr(), qn("w:numPr"))
    _value(num_pr, "w:ilvl", "0")
    _value(num_pr, "w:numId", str(num_id))
    document.add_paragraph(_OUTSIDE)
    authored = f"- # {_TITLE} ^{_TARGET}\n\n{_OUTSIDE}\n"
    session = DocxSemanticsV3Session(document, source_sha256=hashlib.sha256(authored.encode()).hexdigest())
    session.bind_heading(paragraph, {"kind": "heading", "id": _TARGET})
    session.finalize_document()
    document.save(str(path))
    session.write_package(path)
    session.prove_package(path)
    return path


def _numbering(path: Path) -> Any:
    with ZipFile(path) as package:
        return etree.fromstring(package.read("word/numbering.xml"))


def _assert_native_package(path: Path, *, num_id: int, number_format: str, start: int) -> None:
    """Check native content, list identity and bookmark/map bijection before reading."""
    with ZipFile(path) as package:
        body = etree.fromstring(package.read("word/document.xml"))
        numbering = etree.fromstring(package.read("word/numbering.xml"))
        maps = [
            etree.fromstring(package.read(name))
            for name in package.namelist()
            if name.startswith("customXml/item") and name.endswith(".xml")
        ]
    paragraphs = list(body.iter(qn("w:p")))
    assert ["".join(node.text or "" for node in p.iter(qn("w:t"))) for p in paragraphs] == [_TITLE, _OUTSIDE]
    heading = paragraphs[0]
    p_style = heading.find(f"{qn('w:pPr')}/{qn('w:pStyle')}")
    assert p_style is not None and p_style.get(qn("w:val")) == "Heading1"
    binding = heading.find(f"{qn('w:pPr')}/{qn('w:numPr')}")
    assert binding is not None
    assert [(node.tag, node.get(qn("w:val"))) for node in binding] == [
        (qn("w:ilvl"), "0"),
        (qn("w:numId"), str(num_id)),
    ]
    [abstract] = [node for node in numbering.findall(qn("w:abstractNum")) if node.get(qn("w:abstractNumId")) == "0"]
    [level] = abstract.findall(qn("w:lvl"))
    assert level.get(qn("w:ilvl")) == "0"
    for tag, expected in [
        ("w:start", str(start)),
        ("w:numFmt", number_format),
        ("w:lvlText", "•" if number_format == "bullet" else "%1."),
        ("w:suff", "space"),
    ]:
        node = level.find(qn(tag))
        assert node is not None and node.get(qn("w:val")) == expected
    [target_map] = [root for root in maps if etree.QName(root).namespace == TARGET_MAP_NAMESPACE]
    [target] = target_map.findall(f"{{{TARGET_MAP_NAMESPACE}}}targets/{{{TARGET_MAP_NAMESPACE}}}target")
    assert (target.get("kind"), target.get("source_id")) == ("heading", _TARGET)
    [bookmark] = list(heading.iter(qn("w:bookmarkStart")))
    [end] = list(heading.iter(qn("w:bookmarkEnd")))
    assert bookmark.get(qn("w:name")) == target.get("bookmark_name")
    assert bookmark.get(qn("w:id")) == end.get(qn("w:id"))
    assert not list(paragraphs[1].iter(qn("w:bookmarkStart")))


def _read_fresh(path: Path) -> tuple[str, str]:
    document = Document(str(path))
    assert ResolvedNumberingV4Recovery.load_if_present(path, document) is None
    recovery = DocxSemanticsV3Recovery.load(path, document)
    paragraphs = [node for node in recovery.logical_body_elements(document) if node.tag == qn("w:p")]
    assert len(paragraphs) == 2
    [heading, outside] = paragraphs
    anchor = recovery.source_anchor(heading)
    assert anchor is not None
    assert (anchor.owner_kind, anchor.block_kind, anchor.source_id) == ("semantic_target", "heading", _TARGET)
    assert recovery.source_anchor(outside) is None
    return Paragraph(heading, document).text, Paragraph(outside, document).text


def _add_bad_topology(abstract: Any, instance: Any, mutation: str) -> tuple[str, str]:
    [level] = abstract.findall(qn("w:lvl"))
    if mutation.startswith("level_"):
        parent = level.find(qn("w:rPr"))
        assert parent is not None
        kind = mutation.removeprefix("level_")
    elif mutation.startswith("reference_"):
        [parent] = instance.findall(qn("w:abstractNumId"))
        kind = mutation.removeprefix("reference_")
    elif mutation == "override_styleLink":
        override = etree.SubElement(instance, qn("w:lvlOverride"))
        override.set(qn("w:ilvl"), "0")
        parent = _value(override, "w:startOverride", "1")
        kind = "styleLink"
    else:
        raise AssertionError(mutation)
    if kind == "mc":
        bad = etree.SubElement(parent, f"{{{_MC}}}AlternateContent")
        etree.SubElement(bad, f"{{{_MC}}}Fallback")
    else:
        bad = _value(parent, "w:" + kind, "1" if kind == "start" else "Heading1")
    return parent.tag, bad.tag


def _mutated_package(source: Path, target: Path, mutation: str, *, unused: bool) -> tuple[str, str]:
    numbering = _numbering(source)
    [abstract] = numbering.findall(qn("w:abstractNum"))
    [instance] = numbering.findall(qn("w:num"))
    if unused:
        abstract = deepcopy(abstract)
        abstract.set(qn("w:abstractNumId"), "45")
        instance = deepcopy(instance)
        instance.set(qn("w:numId"), "65535")
        [reference] = instance.findall(qn("w:abstractNumId"))
        reference.set(qn("w:val"), "45")
        numbering.append(abstract)
        numbering.append(instance)
    location = _add_bad_topology(abstract, instance, mutation)
    data = etree.tostring(numbering, xml_declaration=True, encoding="UTF-8")
    with ZipFile(source) as original, ZipFile(target, "w") as result:
        for info in original.infolist():
            result.writestr(info, data if info.filename == "word/numbering.xml" else original.read(info.filename))
    with ZipFile(source) as original, ZipFile(target) as result:
        assert original.namelist() == result.namelist()
        for name in original.namelist():
            if name != "word/numbering.xml":
                assert original.read(name) == result.read(name)
    return location


def _assert_bad_node(path: Path, location: tuple[str, str], *, unused: bool) -> None:
    parent_tag, bad_tag = location
    numbering = _numbering(path)
    [bad] = [
        node for node in numbering.iter(bad_tag) if node.getparent() is not None and node.getparent().tag == parent_tag
    ]
    parent = bad.getparent()
    assert parent is not None and parent.tag == parent_tag
    if parent_tag == qn("w:rPr"):
        level = parent.getparent()
        assert level is not None and level.tag == qn("w:lvl") and level.get(qn("w:ilvl")) == "0"
    owner = next(node for node in bad.iterancestors() if node.tag in {qn("w:abstractNum"), qn("w:num")})
    owner_id = owner.get(qn("w:abstractNumId")) if owner.tag == qn("w:abstractNum") else owner.get(qn("w:numId"))
    assert owner_id in ({"45", "65535"} if unused else {"0", "1"})


@pytest.mark.parametrize("num_id", [1, 4097])
@pytest.mark.parametrize("number_format", ["decimal", "bullet"])
@pytest.mark.parametrize("start", [0, 1])
def test_ordinary_definition_with_real_ancestry_retains_native_content(
    tmp_path: Path, num_id: int, number_format: str, start: int
) -> None:
    path = _write_native_heading(tmp_path / "ordinary.docx", num_id=num_id, number_format=number_format, start=start)
    _assert_native_package(path, num_id=num_id, number_format=number_format, start=start)
    assert _read_fresh(path) == (_TITLE, _OUTSIDE)


@pytest.mark.parametrize(("mutation", "message"), _BAD_TOPOLOGIES)
def test_referenced_recursive_topology_cannot_receive_the_ordinary_exemption(
    tmp_path: Path, mutation: str, message: str
) -> None:
    original = _write_native_heading(tmp_path / "valid.docx", num_id=1, number_format="decimal", start=1)
    _assert_native_package(original, num_id=1, number_format="decimal", start=1)
    assert _read_fresh(original) == (_TITLE, _OUTSIDE)
    candidate = tmp_path / "invalid.docx"
    location = _mutated_package(original, candidate, mutation, unused=False)
    _assert_native_package(candidate, num_id=1, number_format="decimal", start=1)
    _assert_bad_node(candidate, location, unused=False)
    with pytest.raises(DocxNumberingImportError, match=message):
        _read_fresh(candidate)


@pytest.mark.parametrize(("mutation", "_message"), _BAD_TOPOLOGIES)
def test_unused_recursive_topology_does_not_revoke_the_referenced_native_binding(
    tmp_path: Path, mutation: str, _message: str
) -> None:
    original = _write_native_heading(tmp_path / "valid.docx", num_id=1, number_format="decimal", start=1)
    candidate = tmp_path / "unused.docx"
    location = _mutated_package(original, candidate, mutation, unused=True)
    _assert_native_package(candidate, num_id=1, number_format="decimal", start=1)
    _assert_bad_node(candidate, location, unused=True)
    assert _read_fresh(candidate) == (_TITLE, _OUTSIDE)
