"""List-item carriers keep their leading paragraph and complete native blocks."""

from __future__ import annotations

import hashlib
from pathlib import Path
from typing import Any
from zipfile import ZipFile

import lxml.etree as etree
import pytest
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from docwen_core.docx_semantics_v3 import DocxSemanticsV3Error, DocxSemanticsV3Recovery, DocxSemanticsV3Session

pytestmark = pytest.mark.contract


def _tag(source_id: str) -> str:
    digest = hashlib.sha256(f"docwen-anchor-map-v1\0anchor\0{source_id}".encode()).hexdigest()
    return "docwen-anchor-v1:" + digest[:32]


def _block(document: Any, kind: str, text: str) -> Any:
    if kind == "p":
        return document.add_paragraph(text)._p
    if kind == "tbl":
        table = document.add_table(rows=1, cols=1)
        table.cell(0, 0).text = text
        return table._tbl
    block = OxmlElement("w:sdt" if kind in {"foreign-sdt", "unmapped-sdt"} else "w:customXml")
    if kind in {"foreign-sdt", "unmapped-sdt"}:
        properties = OxmlElement("w:sdtPr")
        tag = OxmlElement("w:tag")
        tag.set(qn("w:val"), "docwen-anchor-v1:" + "0" * 32 if kind == "unmapped-sdt" else "foreign")
        properties.append(tag)
        content = OxmlElement("w:sdtContent")
        content.append(OxmlElement("w:p"))
        block.extend((properties, content))
    document.element.body.insert(len(document.element.body) - 1, block)
    return block


def _write(
    path: Path,
    kinds: tuple[str, ...],
    *,
    child: bool = False,
    field: str | None = None,
    field_in_table: bool = False,
) -> Path:
    document = Document()
    blocks = tuple(
        _block(document, kind, ("Lead", "Cell", "Continuation.")[index % 3]) for index, kind in enumerate(kinds)
    )
    if field is not None:
        paragraph = next(blocks[1].iter(qn("w:p"))) if field_in_table else blocks[0]
        if field == "bookmark":
            start = OxmlElement("w:bookmarkStart")
            start.set(qn("w:id"), "42")
            start.set(qn("w:name"), "Foreign")
            end = OxmlElement("w:bookmarkEnd")
            end.set(qn("w:id"), "42")
            paragraph.insert(0, start)
            paragraph.append(end)
        else:
            for field_type in ("begin", "separate", "end"):
                run = OxmlElement("w:r")
                marker = OxmlElement("w:fldChar")
                marker.set(qn("w:fldCharType"), field_type)
                run.append(marker)
                paragraph.append(run)
                if field_type == "begin":
                    instruction_run = OxmlElement("w:r")
                    instruction = OxmlElement("w:instrText")
                    instruction.text = f" {field} Foreign "
                    instruction_run.append(instruction)
                    paragraph.append(instruction_run)
    document.add_paragraph("Outside.")
    session = DocxSemanticsV3Session(document, source_sha256=hashlib.sha256(b"list-item carrier source").hexdigest())
    session.bind_ordinary_anchor(blocks, {"block_kind": "list_item", "id": "item"})
    if child:
        session.bind_ordinary_anchor(
            (blocks[-1],), {"block_kind": "paragraph", "id": "child"}, direct_parent_source_id="item"
        )
    session.finalize_document()
    document.save(str(path))
    session.write_package(path)
    return path


def test_item_carrier_accepts_a_table_between_its_lead_and_child_paragraph(tmp_path: Path) -> None:
    path = _write(tmp_path / "item-table.docx", ("p", "tbl", "p"), child=True)
    with ZipFile(path) as package:
        root = etree.fromstring(package.read("word/document.xml"))
    wrappers = {
        tag.get(qn("w:val")): wrapper
        for wrapper in root.iter(qn("w:sdt"))
        if (tag := wrapper.find(f"{qn('w:sdtPr')}/{qn('w:tag')}")) is not None
    }
    assert set(wrappers) == {_tag("item"), _tag("child")}
    parent_content = wrappers[_tag("item")].find(qn("w:sdtContent"))
    child_content = wrappers[_tag("child")].find(qn("w:sdtContent"))
    assert parent_content is not None and child_content is not None
    assert [block.tag for block in parent_content] == [qn("w:p"), qn("w:tbl"), qn("w:sdt")]
    assert [block.tag for block in child_content] == [qn("w:p")]
    assert parent_content[-1] is wrappers[_tag("child")]
    assert [node.text for node in parent_content.iter(qn("w:t"))] == ["Lead", "Cell", "Continuation."]
    outside = [p for p in root.iter(qn("w:p")) if [node.text for node in p.iter(qn("w:t"))] == ["Outside."]]
    assert len(outside) == 1 and outside[0] not in parent_content.iter()
    assert not list(root.iter(qn("w:bookmarkStart"))) and not list(root.iter(qn("w:instrText")))

    loaded = Document(str(path))
    recovery = DocxSemanticsV3Recovery.load(path, loaded)
    body = recovery.logical_body_elements(loaded)
    assert body[-1].tag == qn("w:sectPr")
    logical = body[:-1]
    assert [block.tag for block in logical] == [qn("w:p"), qn("w:tbl"), qn("w:p"), qn("w:p")]
    assert [tuple(group.anchor.source_id for group in recovery.ordinary_anchor_groups(block)) for block in logical] == [
        ("item",),
        ("item",),
        ("child", "item"),
        (),
    ]


@pytest.mark.parametrize(
    ("kinds", "message"),
    [
        (("tbl",), "wrong block kind or cardinality"),
        (("tbl", "p"), "wrong block kind or cardinality"),
        (("p", "unknown"), "wrong block kind or cardinality"),
        (("p", "foreign-sdt"), "wrong block kind or cardinality"),
        (("p", "unmapped-sdt"), "no authenticated map record"),
    ],
    ids=["table-only", "table-first", "unknown-block", "unproven-sdt", "unmapped-owned-sdt"],
)
def test_item_carrier_rejects_invalid_shapes_in_a_fresh_package(
    tmp_path: Path, kinds: tuple[str, ...], message: str
) -> None:
    path = _write(tmp_path / "bad-shape.docx", kinds)
    with pytest.raises(DocxSemanticsV3Error, match=message):
        DocxSemanticsV3Recovery.load(path, Document(str(path)))


@pytest.mark.parametrize("field", ["REF", "SEQ", "bookmark"])
@pytest.mark.parametrize("in_table", [False, True], ids=["lead-paragraph", "table-cell"])
def test_item_carrier_does_not_admit_target_fields_or_bookmarks(tmp_path: Path, field: str, in_table: bool) -> None:
    kinds = ("p", "tbl", "p") if in_table else ("p", "p")
    path = _write(tmp_path / "field.docx", kinds, field=field, field_in_table=in_table)
    with pytest.raises(DocxSemanticsV3Error, match="ordinary anchor must not contain a bookmark, SEQ, or REF"):
        DocxSemanticsV3Recovery.load(path, Document(str(path)))


def test_item_table_range_cannot_partially_overlap_another_anchor() -> None:
    document = Document()
    blocks = tuple(_block(document, kind, str(index)) for index, kind in enumerate(("p", "tbl", "p", "p")))
    session = DocxSemanticsV3Session(document, source_sha256=hashlib.sha256(b"crossing item").hexdigest())
    session.bind_ordinary_anchor(blocks[:3], {"block_kind": "list_item", "id": "item"})
    session.bind_ordinary_anchor(blocks[2:], {"block_kind": "block_quote", "id": "other"})
    with pytest.raises(DocxSemanticsV3Error, match="partially overlap"):
        session.finalize_document()
    assert not list(document.element.iter(qn("w:sdt")))
