"""Zero-character Word proofing anchors preserve authenticated fenced literals."""

from __future__ import annotations

from pathlib import Path
from typing import Any, cast
from zipfile import ZipFile

import lxml.etree as etree
import pytest
from docx import Document
from docx.oxml.ns import qn

from docwen_core.docx_semantics_v3 import DocxSemanticsV3Error, DocxSemanticsV3Recovery, DocxSemanticsV3Session

from ._docx_semantics_v3_fenced_source_support import (
    _identity,
    _map_part,
    _mutate_document_xml,
    _mutate_xml_member,
    _write,
)

pytestmark = pytest.mark.contract

_BODY = 'print("literal **code**")\n'
_SOURCE = "```python\n" + _BODY + "```\n"
_WORD_CONTENT = Path(__file__).resolve().parents[3] / "tests/fixtures/files/word-fenced-proofing-content.xml"


def _package(tmp_path: Path):
    identity = _identity(
        _SOURCE,
        _BODY,
        fence_character="`",
        opening_length=3,
        opening_prefix="",
        info="python",
        opening_eol="\n",
        body_prefixes=("",),
        closing_state="present",
        closing_length=3,
        closing_prefix="",
        closing_suffix="",
        closing_eol="\n",
    )
    document = Document()
    session = DocxSemanticsV3Session(document, source_sha256=identity.source_sha256)
    session.bind_fenced_source(document.add_paragraph(), identity, logical_body=_BODY)
    return session, _write(session, document, tmp_path / "proofing.docx"), identity


def _word_runs(root: etree._Element, marker_types: tuple[str, ...]) -> etree._Element:
    content = next(root.iter(qn("w:sdtContent")))
    for child in list(content):
        content.remove(child)
    captured = etree.fromstring(_WORD_CONTENT.read_bytes())
    for child in list(captured):
        content.append(child)
    first_marker, last_marker = content.findall(qn("w:proofErr"))
    first_marker.set(qn("w:type"), marker_types[0])
    if len(marker_types) == 1:
        content.remove(last_marker)
    else:
        last_marker.set(qn("w:type"), marker_types[1])
    return first_marker


@pytest.mark.parametrize(
    "marker_types",
    [
        ("gramStart", "gramEnd"),
        ("spellStart", "spellEnd"),
        ("gramStart",),
        ("gramEnd",),
        ("spellStart",),
        ("spellEnd",),
    ],
)
def test_empty_word_proofing_anchors_and_split_runs_keep_exact_literal(
    tmp_path: Path, marker_types: tuple[str, ...]
) -> None:
    session, output, identity = _package(tmp_path)
    _mutate_document_xml(output, lambda root: _word_runs(root, marker_types))
    before = output.read_bytes()
    session.prove_package(output)
    loaded = Document(str(output))
    recovery = DocxSemanticsV3Recovery.load(output, loaded)
    paragraph = next(cast(Any, loaded.element).body.iter(qn("w:p")))
    assert recovery.render_fenced_source(paragraph) == _SOURCE
    assert recovery.fenced_source_identities == (identity,)
    assert output.read_bytes() == before


@pytest.mark.parametrize(
    "corruption",
    [
        "extra-attribute",
        "missing-type",
        "unknown-type",
        "unqualified-type",
        "text",
        "tail",
        "child",
        "comment",
        "inside-run",
        "body-text",
        "uuid",
    ],
)
def test_proofing_compatibility_rejects_hidden_payload_and_broken_authentication(
    tmp_path: Path, corruption: str
) -> None:
    session, output, _identity_record = _package(tmp_path)

    def mutate(root: etree._Element) -> None:
        marker = _word_runs(root, ("gramStart", "gramEnd"))
        if corruption == "extra-attribute":
            marker.set(qn("w:val"), "hidden")
        elif corruption == "missing-type":
            del marker.attrib[qn("w:type")]
        elif corruption == "unknown-type":
            marker.set(qn("w:type"), "grammar")
        elif corruption == "unqualified-type":
            del marker.attrib[qn("w:type")]
            marker.set("type", "gramStart")
        elif corruption == "text":
            marker.text = "hidden"
        elif corruption == "tail":
            marker.tail = "hidden"
        elif corruption == "child":
            etree.SubElement(marker, qn("w:t")).text = "hidden"
        elif corruption == "comment":
            marker.append(etree.Comment("hidden"))
        elif corruption == "inside-run":
            next(root.iter(qn("w:r"))).append(marker)
        elif corruption == "body-text":
            next(root.iter(qn("w:t"))).text = "prInt("

    _mutate_document_xml(output, mutate)
    if corruption == "uuid":
        with ZipFile(output) as package:
            part = _map_part(package).replace("/item", "/itemProps")
        _mutate_xml_member(
            output,
            part,
            lambda root: root.set(
                "{http://schemas.openxmlformats.org/officeDocument/2006/customXml}itemID",
                "{00000000-0000-0000-0000-000000000001}",
            ),
        )
    before = output.read_bytes()
    with pytest.raises(DocxSemanticsV3Error):
        session.prove_package(output)
    assert output.read_bytes() == before


def test_valid_package_cannot_authorize_payload_in_an_in_memory_proofing_marker(tmp_path: Path) -> None:
    _session, output, _identity_record = _package(tmp_path)
    _mutate_document_xml(output, lambda root: _word_runs(root, ("gramStart", "gramEnd")))
    loaded = Document(str(output))
    marker = next(cast(Any, loaded.element).iter(qn("w:proofErr")))
    etree.SubElement(marker, qn("w:t")).text = "hidden"
    with pytest.raises(DocxSemanticsV3Error, match="proofing marker"):
        DocxSemanticsV3Recovery.load(output, loaded)
