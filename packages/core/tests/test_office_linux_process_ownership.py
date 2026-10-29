"""Real Linux native process ownership; not LibreOffice/Office format acceptance."""

from __future__ import annotations

import hashlib
import json
import multiprocessing
import os
import shutil
import signal
import subprocess
import sys
import threading
import time
from dataclasses import asdict
from pathlib import Path
from urllib.parse import unquote, urlparse

import pytest
from tests.support.subprocess_runner import append_subprocess_run_record, run_subprocess

from docwen_core import office_bridge, office_process

# Keep the Linux-only test module importable for other platform collectors.
_WNOHANG = getattr(os, "WNOHANG", 0)
_SIGKILL = getattr(signal, "SIGKILL", 0)

pytestmark = [
    pytest.mark.integration,
    pytest.mark.office,
    pytest.mark.linux_only,
    pytest.mark.skipif(sys.platform != "linux", reason="Linux wait-child/pidfd ownership"),
]

_NATIVE_SOURCE = r"""
#define _POSIX_C_SOURCE 200809L
#include <stdio.h>
#include <stdlib.h>
#include <string.h>
#include <signal.h>
#include <unistd.h>
#include <sys/types.h>
#include <time.h>
static void pause_tick(void) { struct timespec delay = {0, 50000000}; nanosleep(&delay, NULL); }
static void beat(const char *base, int index) {
    char name[8192]; snprintf(name, sizeof name, "%s.heartbeat", base);
    FILE *f = fopen(name, "w"); if (f) { fprintf(f, "%d\n", index); fclose(f); }
}
int main(int argc, char **argv) {
    if (argc == 3 && strcmp(argv[1], "--sentinel") == 0) {
        for (int i = 0; i < 400; i++) { beat(argv[2], i); pause_tick(); } return 0;
    }
    const char *base = getenv("DOCWEN_NATIVE_OFFICE_EVIDENCE");
    const char *mode = getenv("DOCWEN_NATIVE_OFFICE_MODE");
    const char *profile = ""; const char *out = "";
    if (!base || !mode || argc < 2) return 2;
    for (int i = 1; i < argc; i++) {
        if (strncmp(argv[i], "-env:UserInstallation=", 22) == 0) profile = argv[i] + 22;
        if (strcmp(argv[i], "--outdir") == 0 && i + 1 < argc) out = argv[i + 1];
    }
    pid_t child = fork(); if (child < 0) return 3; if (child > 0) return 0;
    if (setsid() < 0) return 4;
    if (strcmp(mode, "stubborn") == 0) signal(SIGTERM, SIG_IGN);
    char name[8192]; snprintf(name, sizeof name, "%s.identity", base);
    FILE *identity = fopen(name, "w"); if (!identity) return 5;
    fprintf(identity, "%ld\n%s\n%s\n%s\n", (long)getpid(), profile, out, argv[argc - 1]);
    char proc_path[256]; ssize_t proc_length = readlink("/proc/self", proc_path, sizeof proc_path - 1);
    if (proc_length <= 0) return 8; proc_path[proc_length] = 0; fprintf(identity, "%s\n", proc_path);
    FILE *status = fopen("/proc/self/status", "r"); if (!status) return 9;
    char status_line[1024]; int found = 0;
    while (fgets(status_line, sizeof status_line, status)) {
        if (strncmp(status_line, "NSpid:", 6) == 0) { fputs(status_line, identity); found = 1; break; }
    }
    fclose(status); fclose(identity); if (!found) return 10;
    for (int i = 0; i < 300; i++) {
        beat(base, i); pause_tick();
        if (strcmp(mode, "normal") == 0 && i >= 5) {
            snprintf(name, sizeof name, "%s.release", base);
            if (access(name, F_OK) != 0) continue;
            const char *source = argv[argc - 1]; const char *leaf = strrchr(source, '/'); leaf = leaf ? leaf + 1 : source;
            snprintf(name, sizeof name, "%s/%s", out, leaf);
            FILE *input = fopen(source, "rb"); FILE *output = fopen(name, "wb");
            if (!input || !output) return 6;
            char buffer[4096]; size_t count;
            while ((count = fread(buffer, 1, sizeof buffer, input)) > 0) fwrite(buffer, 1, count, output);
            fclose(input); fclose(output); return 0;
        }
    }
    return 7;
}
"""


@pytest.fixture
def native_office(tmp_path, monkeypatch):
    compiler = shutil.which("cc") or shutil.which("gcc")
    if compiler is None:
        pytest.skip("Native C compiler is required for the Linux lifecycle fixture")
    if not hasattr(os, "pidfd_open") or not hasattr(signal, "pidfd_send_signal"):
        pytest.skip("Linux pidfd support is required")
    source = tmp_path / "office-fixture.c"
    binary = tmp_path / "office-fixture"
    source.write_text(_NATIVE_SOURCE, encoding="utf-8")
    run_subprocess([compiler, "-std=c11", "-O2", str(source), "-o", str(binary)], check=True, timeout=30)
    actual_run = office_process.run_owned_office

    def recorded_owner(cmd, *, is_cancelled, timeout_s):
        started = time.perf_counter()
        result = None
        try:
            result = actual_run(cmd, is_cancelled=is_cancelled, timeout_s=timeout_s)
            return result
        finally:
            append_subprocess_run_record(
                {
                    "argv": cmd,
                    "command": "Owned native Office launcher: " + " ".join(cmd),
                    "returncode": result.launcher_exit_code if result is not None else None,
                    "owner_exitcode": result.owner_exit_code if result is not None else None,
                    "owner_pid": result.owner_pid if result is not None else None,
                    "reaped_pids": list(result.reaped_pids) if result is not None else [],
                    "cleanup_confirmed": result.cleanup_confirmed if result is not None else False,
                    "timed_out": result is not None and "timed out" in result.reason,
                    "timeout_seconds": timeout_s,
                    "duration_seconds": time.perf_counter() - started,
                    "text": False,
                    "capture_output": False,
                }
            )

    monkeypatch.setattr(office_process, "run_owned_office", recorded_owner)
    return binary


def _start_ticks(pid: int) -> int | None:
    try:
        text = Path(f"/proc/{pid}/stat").read_text()
    except FileNotFoundError:
        return None
    return int(text[text.rfind(")") + 2 :].split()[19])


def _active_evidence(base: Path, deadline: float = 5.0) -> tuple[int, int, list[str]]:
    limit = time.monotonic() + deadline
    while time.monotonic() < limit:
        try:
            lines = base.with_suffix(".identity").read_text().splitlines()
            if len(lines) == 6 and base.with_suffix(".heartbeat").read_text().strip():
                pid = int(lines[0])
                proc_pid = int(lines[4])
                namespace_ids = [int(value) for value in lines[5].split()[1:]]
                assert lines[5].startswith("NSpid:")
                assert namespace_ids[-1] == pid and namespace_ids[0] == proc_pid
                ticks = _start_ticks(proc_pid)
                if ticks is not None:
                    return pid, ticks, lines
        except (FileNotFoundError, ValueError):
            pass
        time.sleep(0.01)
    raise AssertionError("The native Office child was never observed active")


def _drop_receipt_worker(connection, stop, cmd, timeout_s):
    class DropFinalReceipt:
        def send(self, message):
            if not isinstance(message, office_process.OfficeProcessResult):
                connection.send(message)

        def close(self):
            connection.close()

    office_process._owner_worker(DropFinalReceipt(), stop, cmd, timeout_s)


def _abnormal_receipt_worker(connection, stop, cmd, timeout_s):
    class ExitAfterReceipt:
        def send(self, message):
            connection.send(message)
            if isinstance(message, office_process.OfficeProcessResult):
                os._exit(17)

        def close(self):
            connection.close()

    office_process._owner_worker(ExitAfterReceipt(), stop, cmd, timeout_s)


def _start_owner_with_abnormal_exit(cmd, timeout_s):
    context = multiprocessing.get_context("spawn")
    parent, child = context.Pipe(duplex=False)
    stop = context.Event()
    owner = context.Process(target=_abnormal_receipt_worker, args=(child, stop, cmd, timeout_s), daemon=True)
    owner.start()
    child.close()
    return owner, parent, stop


def _start_owner_losing_receipt(cmd, timeout_s):
    context = multiprocessing.get_context("spawn")
    parent, child = context.Pipe(duplex=False)
    stop = context.Event()
    owner = context.Process(target=_drop_receipt_worker, args=(child, stop, cmd, timeout_s), daemon=True)
    owner.start()
    child.close()
    return owner, parent, stop


def _discovery_fixture_worker(connection, stop, cmd, timeout_s):
    """Inject only proc discovery availability; all process/wait/pidfd work is real."""
    evidence = Path(os.environ["DOCWEN_NATIVE_OFFICE_EVIDENCE"])
    mode = os.environ["DOCWEN_NATIVE_OFFICE_DISCOVERY"]
    task_root = evidence.with_suffix(".task-view")
    # This is a test-provided children interface populated from the independent
    # native C identity. It does not claim that this host exposes /proc children.
    owner_proc_pid = int(os.readlink("/proc/self"))
    task = task_root / str(owner_proc_pid)
    task.mkdir(parents=True)
    if mode == "available":
        (task / "children").write_text("", encoding="ascii")
    office_process._TASK_ROOT = task_root
    actual_task_reader = office_process._task_child_proc_pids
    actual_status_reader = office_process._status_direct_children

    def record(method, values):
        with evidence.with_suffix(".discovery.jsonl").open("a", encoding="utf-8") as stream:
            stream.write(
                json.dumps({"method": method, "values": sorted(values) if values is not None else None}) + "\n"
            )

    def task_reader(proc_pid, deadline):
        if mode == "available":
            try:
                native = evidence.with_suffix(".identity").read_text().splitlines()
                if len(native) == 6:
                    (task / "children").write_text(native[4] + "\n", encoding="ascii")
            except FileNotFoundError:
                pass
        values = actual_task_reader(proc_pid, deadline)
        record("task", values)
        return values

    def status_reader(proc_pid, depth, deadline):
        values = actual_status_reader(proc_pid, depth, deadline)
        record("status", values)
        return values

    office_process._task_child_proc_pids = task_reader
    office_process._status_direct_children = status_reader
    office_process._owner_worker(connection, stop, cmd, timeout_s)


def _start_owner_with_discovery_fixture(cmd, timeout_s):
    context = multiprocessing.get_context("spawn")
    parent, child = context.Pipe(duplex=False)
    stop = context.Event()
    owner = context.Process(target=_discovery_fixture_worker, args=(child, stop, cmd, timeout_s), daemon=True)
    owner.start()
    child.close()
    return owner, parent, stop


@pytest.mark.parametrize(
    "mode,action,discovery",
    [
        ("normal", "success", "native"),
        ("waiting", "cancel", "native"),
        ("waiting", "timeout", "native"),
        ("stubborn", "cancel", "native"),
        ("stubborn", "timeout", "native"),
        ("waiting", "cancel", "available"),
        ("stubborn", "timeout", "unavailable"),
    ],
)
def test_native_launcher_exit_is_not_process_family_completion(
    native_office, tmp_path, monkeypatch, mode, action, discovery
):
    evidence = tmp_path / "native"
    sentinel_evidence = tmp_path / "unrelated"
    source = tmp_path / "input.docx"
    source.write_bytes(b"independent native fixture bytes")
    output = tmp_path / "private-output"
    output.mkdir()
    monkeypatch.setenv("DOCWEN_NATIVE_OFFICE_EVIDENCE", str(evidence))
    monkeypatch.setenv("DOCWEN_NATIVE_OFFICE_MODE", mode)
    if discovery != "native":
        monkeypatch.setenv("DOCWEN_NATIVE_OFFICE_DISCOVERY", discovery)
        monkeypatch.setattr(office_process, "_start_owner", _start_owner_with_discovery_fixture)
    cancel = threading.Event()
    observed = []
    sentinel_started = time.perf_counter()
    sentinel_command = [str(native_office), "--sentinel", str(sentinel_evidence)]
    sentinel = subprocess.Popen(sentinel_command)
    worker = threading.Thread(
        target=lambda: observed.append(
            office_process.run_owned_office(
                [str(native_office), "--outdir", str(output), str(source)],
                is_cancelled=cancel.is_set,
                timeout_s=0.8 if action == "timeout" else 5.0,
            )
        )
    )
    worker.start()
    try:
        pid, ticks, _lines = _active_evidence(evidence)
        getpgid = getattr(os, "getpgid", None)
        assert getpgid is not None and getpgid(pid) == pid  # Escaped launcher session still belongs to the owner.
        assert not observed
        (tmp_path / "active-process.json").write_text(
            json.dumps({"pid": pid, "startTicks": ticks, "cmd": _lines}), encoding="utf-8"
        )
        if action == "cancel":
            cancel.set()
        elif action == "success":
            evidence.with_suffix(".release").write_text("native observed", encoding="utf-8")
        worker.join(timeout=10)
        assert not worker.is_alive() and len(observed) == 1
        result = observed[0]
        assert result.cleanup_confirmed
        assert result.succeeded is (action == "success")
        assert pid in result.reaped_pids and result.launcher_pid in result.reaped_pids
        if discovery != "native":
            rows = [json.loads(line) for line in evidence.with_suffix(".discovery.jsonl").read_text().splitlines()]
            if discovery == "available":
                assert any(row["method"] == "task" and int(_lines[4]) in (row["values"] or []) for row in rows)
                assert not any(row["method"] == "status" for row in rows)
            else:
                assert any(row["method"] == "task" and row["values"] is None for row in rows)
                assert any(row["method"] == "status" and pid in row["values"] for row in rows)
        assert _start_ticks(int(_lines[4])) != ticks
        with pytest.raises(ChildProcessError):
            os.waitpid(result.owner_pid, _WNOHANG)
        final_beat = evidence.with_suffix(".heartbeat").read_bytes()
        sentinel_before = sentinel_evidence.with_suffix(".heartbeat").read_bytes()
        time.sleep(0.12)
        assert evidence.with_suffix(".heartbeat").read_bytes() == final_beat
        assert sentinel.poll() is None
        sentinel_after = sentinel_evidence.with_suffix(".heartbeat").read_bytes()
        assert sentinel_after != sentinel_before
        (tmp_path / "process-closeout.json").write_text(
            json.dumps(
                {
                    "result": asdict(result),
                    "nativeStartTicksAfter": _start_ticks(int(_lines[4])),
                    "ownerKernelWaitReaped": True,
                    "sentinelBefore": sentinel_before.decode(),
                    "sentinelAfter": sentinel_after.decode(),
                }
            ),
            encoding="utf-8",
        )
        if action == "success":
            assert (output / source.name).read_bytes() == source.read_bytes()
        else:
            assert not list(output.iterdir())
            assert "cancelled" in result.reason if action == "cancel" else "timed out" in result.reason
    finally:
        cancel.set()
        worker.join(timeout=10)
        try:
            sentinel.terminate()
            sentinel.wait(timeout=5)
        finally:
            append_subprocess_run_record(
                {
                    "argv": sentinel_command,
                    "command": " ".join(sentinel_command),
                    "returncode": sentinel.returncode,
                    "duration_seconds": time.perf_counter() - sentinel_started,
                    "timed_out": sentinel.returncode is None,
                    "timeout_seconds": 20,
                    "text": False,
                    "capture_output": False,
                }
            )


@pytest.mark.parametrize(
    "mode,lose_receipt",
    [("normal", False), ("stubborn", False), ("normal", True), ("stubborn", True), ("normal", "abnormal_exit")],
)
def test_bridge_native_process_receipt_controls_all_private_cleanup_and_final_publication(
    native_office,
    tmp_path,
    monkeypatch,
    mode,
    lose_receipt,
):
    from docx import Document

    source = tmp_path / "source.docx"
    document = Document()
    document.add_paragraph("Immutable saved and reopened input")
    document.save(source)
    assert Document(source).paragraphs[0].text == "Immutable saved and reopened input"
    source_hash = hashlib.sha256(source.read_bytes()).hexdigest()
    requested = tmp_path / "final.docx"
    evidence = tmp_path / "bridge-native"
    monkeypatch.setenv("DOCWEN_NATIVE_OFFICE_EVIDENCE", str(evidence))
    monkeypatch.setenv("DOCWEN_NATIVE_OFFICE_MODE", mode)
    monkeypatch.setattr(office_bridge, "find_soffice_path", lambda: str(native_office))
    if lose_receipt:
        # The native owner really runs and reaps its child. Only the final IPC
        # proof is dropped: this is a deterministic transport fault, not D-state.
        fault_start = (
            _start_owner_with_abnormal_exit if lose_receipt == "abnormal_exit" else _start_owner_losing_receipt
        )
        monkeypatch.setattr(office_process, "_start_owner", fault_start)
    results = []
    worker = threading.Thread(
        target=lambda: results.append(
            office_bridge.convert_with_backend_priority(
                str(source),
                str(requested),
                source_format="docx",
                backend_priority=["libreoffice", "unused"],
                com_candidates={},
                libreoffice_format="docx",
                libreoffice_timeout_s=0.8 if mode == "stubborn" else 5.0,
            )
        )
    )
    worker.start()
    child, ticks, facts = _active_evidence(evidence)
    profile = Path(unquote(urlparse(facts[1]).path))
    conversion = Path(facts[2])
    private_input = Path(facts[3])
    (tmp_path / "active-process.json").write_text(
        json.dumps({"pid": child, "startTicks": ticks, "cmd": facts}), encoding="utf-8"
    )
    assert private_input != source and private_input.read_bytes() == source.read_bytes()
    from xml.etree import ElementTree

    policy = ElementTree.parse(profile / "user/registrymodifications.xcu")
    props = [
        node
        for node in policy.iter("prop")
        if node.get("{http://openoffice.org/2001/registry}name") == "DisableMacrosExecution"
    ]
    assert len(props) == 1 and props[0].findtext("value") == "true"
    if mode == "normal":
        evidence.with_suffix(".release").write_text("native and macro profile observed", encoding="utf-8")
    worker.join(timeout=10)
    assert not worker.is_alive() and len(results) == 1
    result = results[0]
    (tmp_path / "bridge-closeout.json").write_text(
        json.dumps(
            {
                "result": asdict(result),
                "nativeStartTicksAfter": _start_ticks(int(facts[4])),
                "sourceSha256": hashlib.sha256(source.read_bytes()).hexdigest(),
                "receiptFault": lose_receipt,
            }
        ),
        encoding="utf-8",
    )
    assert hashlib.sha256(source.read_bytes()).hexdigest() == source_hash
    assert _start_ticks(int(facts[4])) != ticks
    assert result.success is (mode == "normal" and not lose_receipt)
    assert requested.exists() is result.success
    if result.success:
        assert hashlib.sha256(requested.read_bytes()).hexdigest() == source_hash
        assert Document(requested).paragraphs[0].text == "Immutable saved and reopened input"
    if lose_receipt:
        assert result.cleanup_failed and result.error_code == "OFFICE_PROCESS_CLEANUP_UNCONFIRMED"
        assert result.attempted_backend_ids == ("libreoffice",)
        assert profile.is_dir() and conversion.is_dir() and private_input.is_file()
        assert str(profile) in result.cleanup_message and str(private_input.parent) in result.cleanup_message
        # Test-owned cleanup is allowed only after independent identity absence.
        assert _start_ticks(int(facts[4])) != ticks
        shutil.rmtree(profile)
        shutil.rmtree(private_input.parent)
    else:
        assert not profile.exists() and not conversion.exists() and not private_input.parent.exists()


@pytest.mark.parametrize("lose_receipt", [False, True])
def test_runtime_cleanup_cannot_remove_unconfirmed_native_office_process_state(
    native_office,
    tmp_path,
    monkeypatch,
    lose_receipt,
):
    from docx import Document

    from docwen_bundle.runtime_factory import create_runtime_port
    from docwen_core.models.file_ref import FileRef
    from docwen_core.models.request import ConversionRequest, OutputPolicy
    from docwen_runtime.config import ConfigLoader

    source = tmp_path / "runtime-source.docx"
    document = Document()
    document.add_paragraph("Runtime native ownership source")
    document.save(source)
    assert Document(source).paragraphs[0].text == "Runtime native ownership source"
    original = source.read_bytes()
    runtime_root = tmp_path / "runtime"
    output = tmp_path / "published"
    output.mkdir()
    evidence = tmp_path / "runtime-native"
    monkeypatch.setenv("DOCWEN_RUNTIME_ROOT", str(runtime_root))
    monkeypatch.setenv("DOCWEN_NATIVE_OFFICE_EVIDENCE", str(evidence))
    monkeypatch.setenv("DOCWEN_NATIVE_OFFICE_MODE", "waiting")
    monkeypatch.setattr(office_bridge, "find_soffice_path", lambda: str(native_office))
    if lose_receipt:
        monkeypatch.setattr(office_process, "_start_owner", _start_owner_losing_receipt)
    loader = ConfigLoader(base_dir=Path(__file__).resolve().parents[3] / "configs", user_dir=tmp_path / "configs")
    runtime = create_runtime_port(config_loader=loader)
    request = ConversionRequest(
        request_id="native-office-runtime",
        input_refs=[FileRef(str(source), "docx", "document", size_bytes=len(original))],
        target_format="pdf",
        output_policy=OutputPolicy(output_dir=str(output)),
        config_snapshot={"software": {"special_conversions": {"document_to_pdf": ["libreoffice"]}}},
    )
    observed = []
    worker = threading.Thread(target=lambda: observed.append(runtime.execute(request)))
    worker.start()
    child, ticks, facts = _active_evidence(evidence)
    profile = Path(unquote(urlparse(facts[1]).path))
    conversion = Path(facts[2])
    private_input = Path(facts[3])
    (tmp_path / "runtime-active-process.json").write_text(
        json.dumps(
            {
                "pid": child,
                "startTicks": ticks,
                "profile": str(profile),
                "input": str(private_input),
                "conversion": str(conversion),
                "runtimeRoot": str(runtime_root),
            }
        ),
        encoding="utf-8",
    )
    assert private_input.read_bytes() == original
    assert not private_input.is_relative_to(runtime_root)
    runtime.cancel(request.request_id)
    worker.join(timeout=10)
    assert not worker.is_alive() and len(observed) == 1
    result = observed[0]
    assert not result.success and not result.artifacts
    assert not list(output.iterdir())
    assert not list(runtime_root.iterdir())  # Real TaskManager finally/WorkspaceManager cleanup ran.
    assert source.read_bytes() == original
    assert _start_ticks(int(facts[4])) != ticks
    diagnostics = [diagnostic.message for diagnostic in result.diagnostics]
    (tmp_path / "runtime-closeout.json").write_text(
        json.dumps(
            {
                "success": result.success,
                "diagnostics": diagnostics,
                "nativeStartTicksAfter": _start_ticks(int(facts[4])),
                "runtimeWorkspaceEmpty": not list(runtime_root.iterdir()),
                "receiptFault": lose_receipt,
            }
        ),
        encoding="utf-8",
    )
    if lose_receipt:
        assert profile.is_dir() and conversion.is_dir() and private_input.is_file()
        assert any(str(profile) in message and str(private_input) in message for message in diagnostics)
        # The owner really reaped the child; this test hid only its receipt.
        shutil.rmtree(profile)
        shutil.rmtree(private_input.parent)
    else:
        assert not profile.exists() and not conversion.exists() and not private_input.exists()
    runtime.shutdown()


@pytest.mark.parametrize("failure", ["repeated_interrupt", "broken_poll", "join_interrupt", "close_interrupt"])
def test_owner_control_plane_fault_never_grants_cleanup_permission(monkeypatch, failure):
    """Control-plane fault injection; native lifecycle coverage is separate above."""
    from types import SimpleNamespace

    stopped = []
    polls = []
    clock = [0.0]

    def tick():
        clock[0] += 0.2
        return clock[0]

    class Connection:
        def poll(self, _timeout=0):
            polls.append(True)
            if failure == "repeated_interrupt":
                raise KeyboardInterrupt("repeated interrupt")
            if failure == "broken_poll":
                raise OSError("broken pipe polling")
            return True

        def recv(self):
            return office_process.OfficeProcessResult(True, True, owner_pid=123)

        def close(self):
            if failure == "close_interrupt":
                raise KeyboardInterrupt("close interrupted")

    def join(*, timeout):
        assert timeout == office_process._OWNER_JOIN_S
        if failure == "join_interrupt":
            raise KeyboardInterrupt("join interrupted")

    owner = SimpleNamespace(pid=123, exitcode=0, is_alive=lambda: False, join=join, close=lambda: None)
    monkeypatch.setattr(office_process.time, "monotonic", tick)
    monkeypatch.setattr(
        office_process,
        "_start_owner",
        lambda *_args: (owner, Connection(), SimpleNamespace(set=lambda: stopped.append(True))),
    )
    with pytest.raises(office_process.OfficeProcessCleanupUnconfirmed):
        office_process.run_owned_office(["injected-control-plane"], is_cancelled=lambda: False, timeout_s=0.8)
    assert len(polls) < 40
    if failure in {"repeated_interrupt", "broken_poll"}:
        assert stopped


def _deny_kill_then_rescue_worker(connection, stop, cmd, timeout_s):
    """Real owner/native child; deny only production KILL, then test-owned rescue."""
    evidence = Path(os.environ["DOCWEN_NATIVE_OFFICE_EVIDENCE"])
    actual_open = getattr(os, "pidfd_open", None)
    actual_send = getattr(signal, "pidfd_send_signal", None)
    waitid = getattr(os, "waitid", None)
    p_pid = getattr(os, "P_PID", None)
    wexited = getattr(os, "WEXITED", None)
    wnowait = getattr(os, "WNOWAIT", None)
    assert actual_open is not None and actual_send is not None and waitid is not None
    assert p_pid is not None and wexited is not None and wnowait is not None
    opened: dict[int, int] = {}
    held: list[tuple[int, int]] = []
    denied = 0

    def open_pidfd(pid, flags=0):
        descriptor = actual_open(pid, flags)
        opened[descriptor] = pid
        return descriptor

    def send_pidfd(descriptor, sig, *args):
        nonlocal denied
        if sig == _SIGKILL:
            denied += 1
            native = evidence.with_suffix(".identity").read_text().splitlines()
            target = opened[descriptor]
            assert len(native) == 6 and int(native[0]) == target
            if not held:
                # Keep the original kernel identity alive across production's
                # fd close; rescue does not rediscover a numeric PID.
                held.append((target, os.dup(descriptor)))
            raise PermissionError("Test fault: production KILL is refused")
        return actual_send(descriptor, sig, *args)

    patch = pytest.MonkeyPatch()
    patch.setattr(os, "pidfd_open", open_pidfd)
    patch.setattr(signal, "pidfd_send_signal", send_pidfd)
    office_process._owner_worker(connection, stop, cmd, timeout_s)
    assert held and denied
    native_pid, rescue_fd = held[0]
    flags = wexited | _WNOHANG | wnowait
    live_observation = waitid(p_pid, native_pid, flags)
    assert live_observation is None or live_observation.si_pid == 0
    evidence.with_suffix(".kill-denied.json").write_text(
        json.dumps(
            {
                "nativePid": native_pid,
                "productionKillAttempts": denied,
                "nativeKernelWaitChildLiveAfterProductionReturn": True,
                "productionKillFault": "PermissionError",
                "rescueReleased": False,
            }
        ),
        encoding="utf-8",
    )
    deadline = time.monotonic() + 12.0
    while not evidence.with_suffix(".rescue-release").exists() and time.monotonic() < deadline:
        time.sleep(0.02)
    assert evidence.with_suffix(".rescue-release").exists(), "Parent did not authorize fixture rescue"
    # The parent has independently asserted continued activity and retained
    # resources. This rescue is test closeout, never a production stop receipt.
    live_observation = waitid(p_pid, native_pid, flags)
    assert live_observation is None or live_observation.si_pid == 0
    actual_send(rescue_fd, _SIGKILL)
    reaped = []
    complete = False
    deadline = time.monotonic() + 2.0
    while time.monotonic() < deadline:
        try:
            child, status = os.waitpid(-1, _WNOHANG)
        except ChildProcessError:
            complete = True
            break
        if child:
            reaped.append({"pid": child, "exitCode": os.waitstatus_to_exitcode(status)})
        else:
            time.sleep(0.01)
    os.close(rescue_fd)
    evidence.with_suffix(".rescue.json").write_text(
        json.dumps(
            {
                "nativePid": native_pid,
                "originalPidfdUsed": True,
                "kernelWaitChildLiveBeforeRescue": True,
                "reaped": reaped,
                "ECHILD": complete,
            }
        ),
        encoding="utf-8",
    )
    assert complete and any(row["pid"] == native_pid and row["exitCode"] == -_SIGKILL for row in reaped)


def test_bridge_kill_refusal_retains_active_native_process_and_all_private_state(native_office, tmp_path, monkeypatch):
    """Permission fault after genuine TERM/KILL attempts, not D-state or lost receipt."""
    from docx import Document

    source = tmp_path / "kill-refusal.docx"
    document = Document()
    document.add_paragraph("Native activity must survive until authorized fixture rescue")
    document.save(source)
    assert Document(source).paragraphs[0].text == "Native activity must survive until authorized fixture rescue"
    original = source.read_bytes()
    requested = tmp_path / "never-published.docx"
    evidence = tmp_path / "kill-refusal-native"
    monkeypatch.setenv("DOCWEN_NATIVE_OFFICE_EVIDENCE", str(evidence))
    monkeypatch.setenv("DOCWEN_NATIVE_OFFICE_MODE", "stubborn")
    monkeypatch.setattr(office_bridge, "find_soffice_path", lambda: str(native_office))
    owners = []

    def start_owner(cmd, timeout_s):
        context = multiprocessing.get_context("spawn")
        parent, child = context.Pipe(duplex=False)
        stop = context.Event()
        owner = context.Process(target=_deny_kill_then_rescue_worker, args=(child, stop, cmd, timeout_s), daemon=True)
        owner.start()
        child.close()
        owners.append(owner)
        return owner, parent, stop

    monkeypatch.setattr(office_process, "_start_owner", start_owner)
    results = []
    started = time.perf_counter()
    worker = threading.Thread(
        target=lambda: results.append(
            office_bridge.convert_with_backend_priority(
                str(source),
                str(requested),
                source_format="docx",
                backend_priority=["libreoffice", "unused"],
                com_candidates={},
                libreoffice_format="docx",
                libreoffice_timeout_s=0.8,
            )
        )
    )
    worker.start()
    facts = None
    profile = None
    private_input = None
    try:
        native_pid, ticks, facts = _active_evidence(evidence)
        profile = Path(unquote(urlparse(facts[1]).path))
        conversion = Path(facts[2])
        private_input = Path(facts[3])
        worker.join(timeout=10)
        assert not worker.is_alive() and len(results) == 1
        assert time.perf_counter() - started < 8.0
        result = results[0]
        assert not result.success and result.cleanup_failed
        assert result.error_code == "OFFICE_PROCESS_CLEANUP_UNCONFIRMED"
        assert result.attempted_backend_ids == ("libreoffice",)
        assert not requested.exists() and not list(conversion.iterdir())
        assert source.read_bytes() == original and private_input.read_bytes() == original
        assert profile.is_dir() and conversion.is_dir()
        assert str(profile) in result.message and str(private_input) in result.message
        assert _start_ticks(int(facts[4])) == ticks
        assert owners[0].is_alive()
        first_heartbeat = None
        advanced = False
        deadline = time.monotonic() + 0.5
        while time.monotonic() < deadline:
            text = evidence.with_suffix(".heartbeat").read_text().strip()
            if text:
                if first_heartbeat is None:
                    first_heartbeat = int(text)
                elif int(text) > first_heartbeat:
                    advanced = True
                    break
            time.sleep(0.02)
        assert advanced and _start_ticks(int(facts[4])) == ticks
        fault = json.loads(evidence.with_suffix(".kill-denied.json").read_text())
        assert fault["nativePid"] == native_pid and fault["productionKillAttempts"] > 0
        assert fault["nativeKernelWaitChildLiveAfterProductionReturn"] and not fault["rescueReleased"]
        (tmp_path / "kill-refusal-active-closeout.json").write_text(
            json.dumps(
                {
                    "result": asdict(result),
                    "nativeSourceFacts": facts,
                    "nativeStartTicks": ticks,
                    "nativeStillSameIdentity": True,
                    "heartbeatAdvancedAfterFailedConversion": True,
                    "productionKillFault": fault,
                    "profile": str(profile),
                    "privateInput": str(private_input),
                    "privateOutput": str(conversion),
                    "sourceSha256": hashlib.sha256(original).hexdigest(),
                }
            ),
            encoding="utf-8",
        )
    finally:
        evidence.with_suffix(".rescue-release").write_text(
            "Parent assertions finished; fixture may reap its child", encoding="utf-8"
        )
        worker.join(timeout=10)
        if owners:
            owner = owners[0]
            owner.join(timeout=5)
            append_subprocess_run_record(
                {
                    "argv": ["test-owned-original-pidfd-rescue"],
                    "command": "Native fixture rescue after KILL-refusal assertions",
                    "returncode": owner.exitcode,
                    "duration_seconds": time.perf_counter() - started,
                    "timed_out": owner.is_alive(),
                    "timeout_seconds": 18,
                    "text": False,
                    "capture_output": False,
                }
            )
            assert not owner.is_alive() and owner.exitcode == 0
            owner_pid = owner.pid
            owner.close()
            with pytest.raises(ChildProcessError):
                os.waitpid(owner_pid, _WNOHANG)
        if facts is not None:
            assert _start_ticks(int(facts[4])) != ticks
            rescue = json.loads(evidence.with_suffix(".rescue.json").read_text())
            assert rescue["originalPidfdUsed"] and rescue["kernelWaitChildLiveBeforeRescue"] and rescue["ECHILD"]
            assert any(row["pid"] == int(facts[0]) and row["exitCode"] == -_SIGKILL for row in rescue["reaped"])
            if profile is not None:
                shutil.rmtree(profile)
            if private_input is not None:
                shutil.rmtree(private_input.parent)
