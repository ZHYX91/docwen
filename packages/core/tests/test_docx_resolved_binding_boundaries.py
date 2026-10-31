"""Source-plan completeness and immutable post-render object bindings."""

from __future__ import annotations

import hashlib
from pathlib import Path
from zipfile import ZipFile

import pytest
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from lxml import etree

from docwen_core._docx_semantics_v3_fenced import derive_fenced_source_identity_v3
from docwen_core._docx_semantics_v3_fenced_map import fenced_source_mapping_v3
from docwen_core.docx_resolved_numbering import ResolvedNumberingDocxError, ResolvedNumberingDocxSession
from docwen_core.models.resolved_numbering import NumberingTarget, ResolvedDocumentTarget

from .test_docx_resolved_source_carriers_v4 import _caption_bindings, _port, _session, _sha

pytestmark = pytest.mark.contract


@pytest.mark.parametrize("carrier", ["ordinary_anchor", "fenced_source"])
def test_authenticated_source_inventory_cannot_succeed_with_an_empty_binding_plan(tmp_path: Path, carrier: str) -> None:
    source = "ordinary ^ordinary\n\n````text\nliteral\n````\n"
    document = Document()
    session = ResolvedNumberingDocxSession(
        document,
        _port(source),
        heading_style_ids={1: "Heading1"},
        heading_style_names={"heading_1": "Heading 1"},
        caption_style_bindings=_caption_bindings(document),
        expected_ordinary_anchor_ids=("ordinary",) if carrier == "ordinary_anchor" else None,
        expected_fenced_source_tags=("docwen-fenced-source-v1:" + "a" * 32,) if carrier == "fenced_source" else None,
    )
    document.add_paragraph("ordinary")
    document.add_paragraph("literal")
    output = tmp_path / "candidate.docx"
    with pytest.raises(ResolvedNumberingDocxError, match="inventory differs from the authenticated source plan"):
        session.write_package(output)
    assert not output.exists()


@pytest.mark.parametrize("mutation", ["decoration", "payload"])
def test_legitimate_prebind_decoration_does_not_allow_postbind_mutation(tmp_path: Path, mutation: str) -> None:
    source = "Table: Cap\n\n````text\nliteral\n````\n"
    end = len(source) - 1
    target = ResolvedDocumentTarget(0, end, _sha(source[:end]), "table", None, None, "Cap")
    plan_target = NumberingTarget(0, end, "table", False, None, None, None)
    document = Document()
    session = _session(document, _port(source, (target,), (plan_target,)))
    caption = document.add_paragraph(style="Table Caption V4")
    code = document.add_paragraph("literal")
    fence_start = source.index("````")
    identity = derive_fenced_source_identity_v3(
        source_sha256=_sha(source),
        source_start=fence_start,
        source_end=len(source),
        block_sha256=_sha(source[fence_start:]),
        body_sha256=_sha("literal\n"),
        fence_character="`",
        opening_length=4,
        opening_prefix="",
        info="text",
        opening_eol="\n",
        body_prefixes=("",),
        closing_state="present",
        closing_length=4,
        closing_prefix="",
        closing_suffix="",
        closing_eol="\n",
    )
    session.bind_fenced_source(code, fenced_source_mapping_v3(identity), logical_body="literal\n")
    border = OxmlElement("w:pBdr")
    left = OxmlElement("w:left")
    left.set(qn("w:val"), "single")
    left.set(qn("w:sz"), "24")
    border.append(left)
    code._p.get_or_add_pPr().append(border)
    session.bind_caption(caption, (code._p,), source_start=0, source_end=end, kind="table")
    output = tmp_path / "verified.docx"
    session.write_package(output)
    original = hashlib.sha256(output.read_bytes()).hexdigest()
    with ZipFile(output) as package:
        native = etree.fromstring(package.read("word/document.xml"))
    assert [item.text for item in native.iter(qn("w:t"))] == ["Cap", "literal"]
    native_left = native.find(f".//{qn('w:pBdr')}/{qn('w:left')}")
    assert native_left is not None and native_left.get(qn("w:sz")) == "24"
    if mutation == "decoration":
        left.set(qn("w:sz"), "48")
    else:
        next(code._p.iter(qn("w:t"))).text = "changed"
    with pytest.raises((ResolvedNumberingDocxError, ValueError)):
        session.write_package(output)
    assert hashlib.sha256(output.read_bytes()).hexdigest() == original
