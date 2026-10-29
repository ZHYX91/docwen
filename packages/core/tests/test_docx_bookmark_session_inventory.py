"""Bounded allocation work with native content and fresh collision proofs."""

from __future__ import annotations

import hashlib
from pathlib import Path
from typing import Any
from zipfile import ZipFile

import pytest
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from lxml import etree

from docwen_core.docx_bookmarks import build_docx_bookmark_inventory, prove_bookmark_name
from docwen_core.docx_resolved_numbering import ResolvedNumberingDocxSession
from docwen_core.docx_semantics_v3 import DocxSemanticsV3Error, DocxSemanticsV3Session
from docwen_core.models.resolved_numbering import NumberingTarget, ResolvedDocumentTarget

from .test_docx_resolved_source_carriers_v4 import _port, _session

pytestmark = pytest.mark.contract


def _bookmark(paragraph: Any, name: str, raw_id: str) -> None:
    start = OxmlElement("w:bookmarkStart")
    start.set(qn("w:id"), raw_id)
    start.set(qn("w:name"), name)
    end = OxmlElement("w:bookmarkEnd")
    end.set(qn("w:id"), raw_id)
    paragraph._p.append(start)
    paragraph._p.append(end)


def _bound_headings(document: Any, route: str, count: int) -> DocxSemanticsV3Session | ResolvedNumberingDocxSession:
    lines = [f"# Heading {index} ^heading-{index}\n" for index in range(count)]
    source = "\n".join(lines)
    if route == "v3":
        session = DocxSemanticsV3Session(document, source_sha256=hashlib.sha256(source.encode()).hexdigest())
        for index in range(count):
            paragraph = document.add_heading(f"Heading {index}", level=1)
            session.bind_heading(paragraph, {"kind": "heading", "id": f"heading-{index}"})
        return session
    targets = []
    plan_targets = []
    start = 0
    for index, line in enumerate(lines):
        end = start + len(line) - 1
        targets.append(
            ResolvedDocumentTarget(
                start,
                end,
                hashlib.sha256(source[start:end].encode()).hexdigest(),
                "heading",
                f"heading-{index}",
                1,
                f"Heading {index}",
            )
        )
        plan_targets.append(NumberingTarget(start, end, "heading", False, f"heading-{index}", None, None))
        start += len(line) + 1
    resolved = _session(document, _port(source, tuple(targets), tuple(plan_targets)))
    for index, target in enumerate(targets):
        paragraph = document.add_heading(f"Heading {index}", level=1)
        resolved.bind_heading(paragraph, source_start=target.source_start, source_end=target.source_end)
    return resolved


def _write(session: DocxSemanticsV3Session | ResolvedNumberingDocxSession, document: Any, output: Path) -> None:
    if isinstance(session, DocxSemanticsV3Session):
        session.finalize_document()
        document.save(str(output))
    session.write_package(output)


@pytest.mark.parametrize("route", ["v3", "v4"])
@pytest.mark.parametrize("count", [1, 64])
def test_binding_inventory_cost_is_bounded_and_final_ids_are_globally_unique(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch, route: str, count: int
) -> None:
    import docwen_core._docx_semantics_v3_session as v3_module
    import docwen_core.docx_resolved_numbering as v4_module

    document = Document()
    header = document.sections[0].header.paragraphs[0]
    _bookmark(header, "TemplateRange", "+0000")
    orphan = OxmlElement("w:bookmarkEnd")
    orphan.set(qn("w:id"), "1")
    header._p.append(orphan)
    inspected = []

    def inspect(actual: Any):
        inspected.append(actual)
        return build_docx_bookmark_inventory(actual)

    monkeypatch.setattr(v3_module if route == "v3" else v4_module, "build_docx_bookmark_inventory", inspect)
    session = _bound_headings(document, route, count)
    assert inspected == [document]
    output = tmp_path / "headings.docx"
    _write(session, document, output)
    session.prove_package(output)
    with ZipFile(output) as package:
        main = etree.fromstring(package.read("word/document.xml"))
        header_xml = etree.fromstring(package.read("word/header1.xml"))
    starts = tuple(main.iter(qn("w:bookmarkStart")))
    ends = tuple(main.iter(qn("w:bookmarkEnd")))
    assert len(starts) == len(ends) == count
    assert [item.get(qn("w:id")) for item in starts] == [str(index + 2) for index in range(count)]
    assert [item.get(qn("w:id")) for item in ends] == [item.get(qn("w:id")) for item in starts]
    assert [item.text for item in main.iter(qn("w:t"))] == [f"Heading {index}" for index in range(count)]
    assert [item.get(qn("w:id")) for item in header_xml.iter(qn("w:bookmarkStart"))] == ["+0000"]
    reopened = Document(str(output))
    fresh = build_docx_bookmark_inventory(reopened)
    assert all(prove_bookmark_name(fresh, item.get(qn("w:name"), "")).valid for item in starts)


@pytest.mark.parametrize("route", ["v3", "v4"])
@pytest.mark.parametrize("collision", ["casefold_name", "numeric_id", "orphan_end"])
def test_allocation_snapshot_never_replaces_fresh_package_collision_proof(
    tmp_path: Path, route: str, collision: str
) -> None:
    document = Document()
    session = _bound_headings(document, route, 2)
    inventory = build_docx_bookmark_inventory(document)
    start = inventory.starts[0]
    header = document.sections[0].header.paragraphs[0]
    if collision == "casefold_name":
        assert start.name is not None
        _bookmark(header, start.name.swapcase(), "999")
    elif collision == "numeric_id":
        _bookmark(header, "LateUnrelated", f"+00{start.raw_id}")
    else:
        orphan = OxmlElement("w:bookmarkEnd")
        orphan.set(qn("w:id"), start.raw_id)
        header._p.append(orphan)
    output = tmp_path / "collision.docx"
    with pytest.raises((DocxSemanticsV3Error, ValueError)):
        _write(session, document, output)
        session.prove_package(output)
