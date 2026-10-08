"""Physical binding and native-edit boundaries of table role recovery."""

from copy import deepcopy
from pathlib import Path
from typing import Any

import pytest
from docx import Document
from docx.oxml.ns import qn

from docwen_core.docx_parsing.document_semantics import extract_semantic_table_metadata
from docwen_core.docx_semantics import apply_semantic_table_roles
from docwen_core.docx_table_roles import inject_table_roles, prepare_table_roles, recover_table_roles

pytestmark = pytest.mark.unit


def _candidate(tmp_path: Path):
    document: Any = Document()
    for header_rows, header_columns in ((2, 1), (1, 2)):
        table = document.add_table(rows=4, cols=3)
        table.cell(2, 0).text = "Editable body"
        apply_semantic_table_roles(table, header_rows=header_rows, header_columns=header_columns, repeat_header="never")
    payload = prepare_table_roles(document)
    path = tmp_path / "candidate.docx"
    document.save(str(path))
    inject_table_roles(path, payload)
    return path


def _normalize(document):
    for table in document.tables:
        for marker in list(table._tbl.iter(qn("w:cnfStyle"))):
            marker.getparent().remove(marker)


def test_reordered_equal_shape_tables_keep_own_roles_and_edited_text(tmp_path: Path) -> None:
    path = _candidate(tmp_path)
    document: Any = Document(str(path))
    _normalize(document)
    document.tables[0].cell(2, 0).text = "Current edited body"
    second = document.tables[1]._tbl
    document.element.body.remove(second)
    document.element.body.insert(0, second)
    document.save(str(path))
    readback: Any = Document(str(path))
    assert recover_table_roles(path, readback) == ()
    roles = [extract_semantic_table_metadata(table._tbl) for table in readback.tables]
    assert [(item.header_rows, item.header_columns, item.repeat_header) for item in roles] == [
        (1, 2, "never"),
        (2, 1, "never"),
    ]
    assert readback.tables[1].cell(2, 0).text == "Current edited body"


@pytest.mark.parametrize("edit", ["delete", "add_row", "merge"])
def test_stale_records_warn_without_transferring_roles(tmp_path: Path, edit: str) -> None:
    path = _candidate(tmp_path)
    document: Any = Document(str(path))
    _normalize(document)
    first = document.tables[0]
    if edit == "delete":
        document.element.body.remove(first._tbl)
    elif edit == "add_row":
        first.add_row()
    else:
        first.cell(2, 0).merge(first.cell(2, 1))
    document.save(str(path))
    readback: Any = Document(str(path))
    assert len(recover_table_roles(path, readback)) == 1
    unaffected = extract_semantic_table_metadata(readback.tables[-1]._tbl)
    assert (unaffected.header_rows, unaffected.header_columns) == (1, 2)
    if edit != "delete":
        stale = extract_semantic_table_metadata(readback.tables[0]._tbl)
        assert (stale.header_rows, stale.header_columns) == (1, 0)


@pytest.mark.parametrize("edit", ["duplicate_name", "duplicate_id", "missing_end", "cross_table"])
def test_ambiguous_bookmarks_are_rejected(tmp_path: Path, edit: str) -> None:
    path = _candidate(tmp_path)
    document: Any = Document(str(path))
    starts = list(document.element.iter(qn("w:bookmarkStart")))
    ends = list(document.element.iter(qn("w:bookmarkEnd")))
    if edit == "duplicate_name":
        starts[1].set(qn("w:name"), starts[0].get(qn("w:name")))
    elif edit == "duplicate_id":
        starts[1].set(qn("w:id"), starts[0].get(qn("w:id")))
    elif edit == "missing_end":
        ends[0].getparent().remove(ends[0])
    else:
        ends[0].getparent().remove(ends[0])
        starts[1].getparent().append(ends[0])
    document.save(str(path))
    with pytest.raises(ValueError, match="bookmark"):
        recover_table_roles(path, Document(str(path)))


def test_native_disabling_roles_and_repeat_policy_is_respected(tmp_path: Path) -> None:
    path = _candidate(tmp_path)
    document: Any = Document(str(path))
    _normalize(document)
    table = document.tables[0]
    look = table._tbl.tblPr.find(qn("w:tblLook"))
    look.set(qn("w:firstRow"), "0")
    look.set(qn("w:firstColumn"), "0")
    document.save(str(path))
    readback: Any = Document(str(path))
    assert recover_table_roles(path, readback) == ()
    roles = extract_semantic_table_metadata(readback.tables[0]._tbl)
    assert (roles.header_rows, roles.header_columns) == (0, 0)


def test_repeated_preparation_is_stable_and_duplicate_injection_is_rejected(tmp_path: Path) -> None:
    path = _candidate(tmp_path)
    document: Any = Document(str(path))
    payload = prepare_table_roles(document)
    assert payload == prepare_table_roles(document)
    document.save(str(path))
    with pytest.raises(ValueError, match="UUID collides"):
        inject_table_roles(path, payload)
    assert recover_table_roles(path, Document(str(path))) == ()


def test_copied_binding_never_selects_first_equal_shape_table(tmp_path: Path) -> None:
    path = _candidate(tmp_path)
    document: Any = Document(str(path))
    document.element.body.insert(0, deepcopy(document.tables[0]._tbl))
    document.save(str(path))
    with pytest.raises(ValueError, match="bookmark"):
        recover_table_roles(path, Document(str(path)))


@pytest.mark.parametrize("mutation", ["foreign_name", "nonempty_range", "outside_first_cell"])
def test_ordinary_anchor_exception_does_not_allow_arbitrary_bookmarks(tmp_path: Path, mutation: str) -> None:
    from docwen_core._docx_semantics_v3_topology import prove_ordinary_anchor_group

    path = _candidate(tmp_path)
    document: Any = Document(str(path))
    table = document.tables[0]._tbl
    import lxml.etree as etree

    from docwen_core.docx_table_roles import prove_table_role_bookmarks

    payload = prepare_table_roles(document)
    assert payload is not None
    role_map = etree.fromstring(payload)
    proven = prove_table_role_bookmarks(document, role_map)
    assert prove_ordinary_anchor_group((table,), "table", proven_table_role_nodes=proven) == (table,)
    start = next(table.iter(qn("w:bookmarkStart")))
    end = next(table.iter(qn("w:bookmarkEnd")))
    if mutation == "foreign_name":
        start.set(qn("w:name"), "UnrelatedBookmark")
    elif mutation == "nonempty_range":
        from docx.oxml import OxmlElement

        start.addnext(OxmlElement("w:r"))
    else:
        parent = document.tables[0].cell(2, 0).paragraphs[0]._p
        parent.append(start)
        parent.append(end)
    with pytest.raises(ValueError, match="bookmark"):
        proven = prove_table_role_bookmarks(document, role_map)
        prove_ordinary_anchor_group((table,), "table", proven_table_role_nodes=proven)
