"""Native list bindings distinguish v3 headings from proven resolved-v4 facts."""

from __future__ import annotations

import hashlib
import json
from copy import deepcopy
from pathlib import Path
from typing import Any
from zipfile import ZipFile

import pytest
from docx import Document
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph
from lxml import etree

from docwen_core.docx_numbering_import import DocxNumberingImportError
from docwen_core.docx_resolved_numbering import ResolvedNumberingDocxSession
from docwen_core.docx_resolved_numbering_recovery import ResolvedNumberingV4Recovery
from docwen_core.docx_semantics_v3 import (
    CAPTION_STYLE_BINDING_MAP_NAMESPACE,
    TARGET_MAP_NAMESPACE,
    DocxSemanticsV3Error,
    DocxSemanticsV3Recovery,
    DocxSemanticsV3Session,
)
from docwen_core.models.resolved_numbering import (
    ResolvedNumberingPort,
    canonicalize_numbering_plan,
    load_resolved_numbering_bytes,
)

from .test_docx_resolved_numbering_v4 import _caption_bindings

pytestmark = pytest.mark.contract

_TITLE = "2.3 Authored heading"
_TARGET = "heading-target"


def _sha(value: str) -> str:
    return hashlib.sha256(value.encode("utf-8")).hexdigest()


def _value(parent: Any, tag: str, value: str) -> Any:
    child = etree.SubElement(parent, qn(tag))
    child.set(qn("w:val"), value)
    return child


def _native_list_heading(
    path: Path, *, level: int, num_id: int, number_format: str = "decimal", start: int = 1
) -> Path:
    """A real v3 session over an independent ordinary-list OOXML fixture."""
    document = Document()
    numbering = document.part.numbering_part.element
    for child in list(numbering):
        numbering.remove(child)
    abstract = etree.SubElement(numbering, qn("w:abstractNum"))
    abstract.set(qn("w:abstractNumId"), "0")
    _value(abstract, "w:multiLevelType", "hybridMultilevel")
    for depth in range(9):
        native_level = etree.SubElement(abstract, qn("w:lvl"))
        native_level.set(qn("w:ilvl"), str(depth))
        _value(native_level, "w:start", str(start))
        _value(native_level, "w:numFmt", number_format)
        _value(native_level, "w:lvlText", "•" if number_format == "bullet" else f"%{depth + 1}.")
        _value(native_level, "w:suff", "space")
    instance = etree.SubElement(numbering, qn("w:num"))
    instance.set(qn("w:numId"), str(num_id))
    _value(instance, "w:abstractNumId", "0")

    paragraph = document.add_heading(_TITLE, level=level)
    num_pr = etree.SubElement(paragraph._p.get_or_add_pPr(), qn("w:numPr"))
    _value(num_pr, "w:ilvl", "0")
    _value(num_pr, "w:numId", str(num_id))
    session = DocxSemanticsV3Session(document, source_sha256=_sha(f"- {'#' * level} {_TITLE} ^{_TARGET}\n"))
    session.bind_heading(paragraph, {"kind": "heading", "id": _TARGET})
    session.finalize_document()
    document.save(str(path))
    session.write_package(path)
    session.prove_package(path)
    return path


def _heading_only_port(level: int) -> ResolvedNumberingPort:
    source = f"{'#' * level} {_TITLE} ^{_TARGET}\n"
    end = len(source) - 1
    target = {
        "source_start": 0,
        "source_end": end,
        "source_slice_sha256": _sha(source[:end]),
        "kind": "heading",
        "target_id": _TARGET,
        "heading_level": level,
        "authored_text": _TITLE,
    }
    plan = {
        "heading_definitions": [
            {
                "definition_id": "headings",
                "levels": [
                    {
                        "level": index,
                        "start": 1,
                        "number_format": "arabic_half",
                        "display": [{"counter": {"level": index, "number_format": "arabic_half"}}],
                        "suffix": "space",
                        "restart_after_level": None,
                    }
                    for index in range(1, 10)
                ],
            }
        ],
        "heading_instances": [{"instance_id": "document", "definition_id": "headings", "starts": []}],
        "targets": [
            {
                "source_start": 0,
                "source_end": end,
                "kind": "heading",
                "enabled": True,
                "target_id": _TARGET,
                "derived_number": "1",
                "materialization": {
                    "type": "heading_list",
                    "definition_id": "headings",
                    "instance_id": "document",
                    "level": level,
                },
            }
        ],
    }
    digest = hashlib.sha256(canonicalize_numbering_plan(plan)).hexdigest()
    identity = {"input_id": "heading-reader", "source_sha256": _sha(source), "plan_sha256": digest}
    neutral = {
        "$schema": "urn:docwen:schema:resolved-document:v1",
        "schema": "docwen.resolved_document.v1",
        **identity,
        "document": {
            "authored_markdown": source,
            "targets": [target],
            "references": [],
            "resource_occurrences": [],
            "citations": [],
            "resources": [],
        },
    }
    envelope = {
        "$schema": "urn:docwen:schema:numbering-export-plan:v1",
        "schema": "docwen.numbering_export_plan.v1",
        **identity,
        "plan": plan,
    }
    return load_resolved_numbering_bytes(json.dumps(neutral).encode(), json.dumps(envelope).encode())


def _resolved_heading(path: Path, *, level: int) -> Path:
    document = Document()
    # Do not make a high numId an accidental routing signal: let the real v4
    # writer allocate from an empty native numbering part. The default styles
    # still reserve IDs 1..3, so its independently expected fresh ID is 4.
    numbering = document.part.numbering_part.element
    for child in list(numbering):
        numbering.remove(child)
    session = ResolvedNumberingDocxSession(
        document,
        _heading_only_port(level),
        heading_style_ids={index: f"Heading{index}" for index in range(1, 10)},
        heading_style_names={f"heading_{index}": f"Heading {index}" for index in range(1, 10)},
        caption_style_bindings=_caption_bindings(document),
    )
    paragraph = document.add_heading(_TITLE, level=level)
    target = session.port.document.targets[0]
    session.bind_heading(paragraph, source_start=target.source_start, source_end=target.source_end)
    session.write_package(path)
    session.prove_package(path)
    return path


def _package_evidence(path: Path) -> tuple[Any, Any, dict[str, Any]]:
    with ZipFile(path) as package:
        document = etree.fromstring(package.read("word/document.xml"))
        numbering = etree.fromstring(package.read("word/numbering.xml"))
        maps = {}
        for name in package.namelist():
            if name.startswith("customXml/item") and name.endswith(".xml"):
                root = etree.fromstring(package.read(name))
                maps[etree.QName(root).namespace] = root
    assert CAPTION_STYLE_BINDING_MAP_NAMESPACE not in maps
    [target] = maps[TARGET_MAP_NAMESPACE].findall(f"{{{TARGET_MAP_NAMESPACE}}}targets/{{{TARGET_MAP_NAMESPACE}}}target")
    assert (target.get("kind"), target.get("source_id")) == ("heading", _TARGET)
    return document, numbering, maps


def _heading_element(root: Any) -> Any:
    paragraphs = [
        paragraph
        for paragraph in root.iter(qn("w:p"))
        if "".join(node.text or "" for node in paragraph.iter(qn("w:t"))) == _TITLE
    ]
    assert len(paragraphs) == 1
    return paragraphs[0]


@pytest.mark.parametrize("level", [1, 7, 8, 9])
@pytest.mark.parametrize("num_id", [1, 4097])
@pytest.mark.parametrize("number_format", ["decimal", "bullet"])
def test_v3_list_heading_keeps_ordinary_list_routing(
    tmp_path: Path, level: int, num_id: int, number_format: str
) -> None:
    path = _native_list_heading(tmp_path / "ordinary.docx", level=level, num_id=num_id, number_format=number_format)
    document, numbering, _maps = _package_evidence(path)
    [definition] = numbering.findall(qn("w:abstractNum"))
    assert definition.find(qn("w:multiLevelType")).get(qn("w:val")) == "hybridMultilevel"
    assert not list(definition.iter(qn("w:pStyle")))
    paragraph = _heading_element(document)
    assert paragraph.find(f"{qn('w:pPr')}/{qn('w:pStyle')}").get(qn("w:val")) == f"Heading{level}"
    assert paragraph.find(f"{qn('w:pPr')}/{qn('w:numPr')}/{qn('w:numId')}").get(qn("w:val")) == str(num_id)
    reopened = Document(str(path))
    assert ResolvedNumberingV4Recovery.load_if_present(path, reopened) is None
    recovery = DocxSemanticsV3Recovery.load(path, reopened)
    [logical] = [node for node in recovery.logical_body_elements(reopened) if node.tag == qn("w:p")]
    assert Paragraph(logical, reopened).text == _TITLE
    anchor = recovery.source_anchor(logical)
    assert anchor is not None
    assert (anchor.owner_kind, anchor.block_kind, anchor.source_id) == ("semantic_target", "heading", _TARGET)


@pytest.mark.parametrize("level", [1, 7, 8, 9])
def test_heading_only_v4_with_no_caption_registry_proves_real_numbering(tmp_path: Path, level: int) -> None:
    path = _resolved_heading(tmp_path / "resolved.docx", level=level)
    document, numbering, _maps = _package_evidence(path)
    [definition] = numbering.findall(qn("w:abstractNum"))
    assert definition.find(qn("w:multiLevelType")).get(qn("w:val")) == "multilevel"
    levels = definition.findall(qn("w:lvl"))
    assert [item.find(qn("w:pStyle")).get(qn("w:val")) for item in levels] == [
        f"Heading{index}" for index in range(1, 10)
    ]
    paragraph = _heading_element(document)
    num_pr = paragraph.find(f"{qn('w:pPr')}/{qn('w:numPr')}")
    assert num_pr.find(qn("w:numId")).get(qn("w:val")) == "4"
    assert num_pr.find(qn("w:ilvl")).get(qn("w:val")) == str(level - 1)
    assert not list(paragraph.iter(qn("w:instrText")))
    reopened = Document(str(path))
    recovery = ResolvedNumberingV4Recovery.load_if_present(path, reopened)
    assert recovery is not None and recovery.is_resolved_v4
    [logical] = [node for node in recovery.logical_body_elements(reopened) if node.tag == qn("w:p")]
    imported = recovery.heading_import(Paragraph(logical, reopened))
    assert imported.authored_text == _TITLE
    assert imported.diagnostics == ()
    assert imported.numbering is not None
    fact = imported.numbering
    assert (fact.num_id, fact.abstract_num_id, fact.level, fact.number_format) == (4, 0, level, "decimal")
    assert (fact.level_text, fact.start, fact.suffix, fact.style_id) == (f"%{level}", 1, "space", f"Heading{level}")
    assert recovery.recovered_captions == ()


def _tamper(path: Path, mutation: str) -> None:
    with ZipFile(path) as package:
        infos = package.infolist()
        payloads = {info.filename: package.read(info.filename) for info in infos}
    if mutation.startswith("numPr"):
        root = etree.fromstring(payloads["word/document.xml"])
        num_pr = _heading_element(root).find(f"{qn('w:pPr')}/{qn('w:numPr')}")
        assert num_pr is not None
        if mutation == "numPr_missing_level":
            level = num_pr.find(qn("w:ilvl"))
            assert level is not None
            num_pr.remove(level)
        elif mutation == "numPr_duplicate_id":
            num_id = num_pr.find(qn("w:numId"))
            assert num_id is not None
            num_pr.append(deepcopy(num_id))
        elif mutation == "numPr_missing_instance":
            num_id = num_pr.find(qn("w:numId"))
            assert num_id is not None
            num_id.set(qn("w:val"), "2147483647")
        else:
            raise AssertionError(mutation)
        payloads["word/document.xml"] = etree.tostring(root, xml_declaration=True, encoding="UTF-8")
    elif mutation in {"definition_full_override", "definition_styleLink", "definition_numStyleLink"}:
        root = etree.fromstring(payloads["word/numbering.xml"])
        if mutation == "definition_full_override":
            instance = root.find(qn("w:num"))
            assert instance is not None
            override = etree.SubElement(instance, qn("w:lvlOverride"))
            override.set(qn("w:ilvl"), "0")
            original_level = root.find(f"{qn('w:abstractNum')}/{qn('w:lvl')}")
            assert original_level is not None
            level = deepcopy(original_level)
            _value(level, "w:pStyle", "Heading9")
            override.append(level)
        else:
            definition = root.find(qn("w:abstractNum"))
            assert definition is not None
            _value(definition, "w:" + mutation.removeprefix("definition_"), "UnprovenStyle")
        payloads["word/numbering.xml"] = etree.tostring(root, xml_declaration=True, encoding="UTF-8")
    elif mutation in {"definition_missing_format", "definition_duplicate_text", "definition_invalid_suffix"}:
        root = etree.fromstring(payloads["word/numbering.xml"])
        level = root.find(f"{qn('w:abstractNum')}/{qn('w:lvl')}")
        assert level is not None
        if mutation == "definition_missing_format":
            number_format = level.find(qn("w:numFmt"))
            assert number_format is not None
            level.remove(number_format)
        elif mutation == "definition_duplicate_text":
            level_text = level.find(qn("w:lvlText"))
            assert level_text is not None
            level.append(deepcopy(level_text))
        else:
            suffix = level.find(qn("w:suff"))
            assert suffix is not None
            suffix.set(qn("w:val"), "newline")
        payloads["word/numbering.xml"] = etree.tostring(root, xml_declaration=True, encoding="UTF-8")
    elif mutation == "definition_invalid_start":
        root = etree.fromstring(payloads["word/numbering.xml"])
        start = root.find(f"{qn('w:abstractNum')}/{qn('w:lvl')}/{qn('w:start')}")
        assert start is not None
        start.set(qn("w:val"), "0")
        payloads["word/numbering.xml"] = etree.tostring(root, xml_declaration=True, encoding="UTF-8")
    elif mutation == "target_map_hash":
        name = next(
            name
            for name, data in payloads.items()
            if name.startswith("customXml/item") and name.endswith(".xml") and TARGET_MAP_NAMESPACE.encode() in data
        )
        root = etree.fromstring(payloads[name])
        target = root.find(f"{{{TARGET_MAP_NAMESPACE}}}targets/{{{TARGET_MAP_NAMESPACE}}}target")
        assert target is not None
        target.set("sha256", "0" * 64)
        payloads[name] = etree.tostring(root, xml_declaration=True, encoding="UTF-8")
    else:
        raise AssertionError(mutation)
    with ZipFile(path, "w") as package:
        for info in infos:
            package.writestr(info, payloads[info.filename])


@pytest.mark.parametrize(
    "mutation",
    [
        "numPr_missing_level",
        "numPr_duplicate_id",
        "numPr_missing_instance",
        "definition_invalid_start",
        "definition_missing_format",
        "definition_duplicate_text",
        "definition_invalid_suffix",
        "target_map_hash",
    ],
)
def test_v4_routing_does_not_bypass_invalid_heading_evidence(tmp_path: Path, mutation: str) -> None:
    path = _resolved_heading(tmp_path / "tampered.docx", level=7)
    _package_evidence(path)
    _tamper(path, mutation)
    with pytest.raises((DocxSemanticsV3Error, DocxNumberingImportError)):
        ResolvedNumberingV4Recovery.load_if_present(path, Document(str(path)))


@pytest.mark.parametrize(
    "mutation",
    [
        "numPr_missing_level",
        "numPr_duplicate_id",
        "numPr_missing_instance",
        "definition_full_override",
        "definition_styleLink",
        "definition_numStyleLink",
        "definition_missing_format",
        "definition_duplicate_text",
        "definition_invalid_suffix",
    ],
)
def test_malformed_ordinary_binding_does_not_gain_the_v3_exemption(tmp_path: Path, mutation: str) -> None:
    path = _native_list_heading(tmp_path / "tampered-ordinary.docx", level=9, num_id=4097)
    _package_evidence(path)
    _tamper(path, mutation)
    with pytest.raises((DocxSemanticsV3Error, DocxNumberingImportError)):
        ResolvedNumberingV4Recovery.load_if_present(path, Document(str(path)))


@pytest.mark.parametrize("number_format", ["decimal", "bullet"])
def test_ordinary_zero_start_remains_valid_without_weakening_resolved_numbering(
    tmp_path: Path, number_format: str
) -> None:
    path = _native_list_heading(
        tmp_path / "ordinary-zero.docx", level=7, num_id=4097, number_format=number_format, start=0
    )
    _document, numbering, _maps = _package_evidence(path)
    levels = numbering.findall(f"{qn('w:abstractNum')}/{qn('w:lvl')}")
    assert len(levels) == 9
    assert [level.find(qn("w:start")).get(qn("w:val")) for level in levels] == ["0"] * 9
    assert [level.find(qn("w:numFmt")).get(qn("w:val")) for level in levels] == [number_format] * 9
    reopened = Document(str(path))
    assert ResolvedNumberingV4Recovery.load_if_present(path, reopened) is None
    recovery = DocxSemanticsV3Recovery.load(path, reopened)
    [logical] = [node for node in recovery.logical_body_elements(reopened) if node.tag == qn("w:p")]
    assert Paragraph(logical, reopened).text == _TITLE
    anchor = recovery.source_anchor(logical)
    assert anchor is not None and anchor.source_id == _TARGET
