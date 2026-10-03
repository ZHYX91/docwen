"""Resolved-v4 standalone caption occurrence authority gates."""

from __future__ import annotations

import hashlib

import lxml.etree as etree
import pytest
from docx import Document
from docx.enum.style import WD_STYLE_TYPE

from docwen_core._docx_semantics_v3_model import DocxSemanticsV3Error
from docwen_core.docx_numbering_ooxml import append_complex_field
from docwen_core.docx_standalone_caption_occurrence import (
    derive_standalone_caption_occurrence,
    parse_standalone_caption_occurrence_map,
    prove_standalone_caption_occurrence_sdt,
    standalone_caption_occurrence_map_xml,
    wrap_standalone_caption_occurrence,
)

pytestmark = pytest.mark.unit

SOURCE_SHA = hashlib.sha256(b"standalone captions").hexdigest()
PLAN_SHA = hashlib.sha256(b"plan").hexdigest()


def _identity(*, enabled: bool, derived_number: str | None):
    return derive_standalone_caption_occurrence(
        source_sha256=SOURCE_SHA,
        source_start=3,
        source_end=29,
        kind="table",
        plan_sha256=PLAN_SHA,
        enabled=enabled,
        derived_number=derived_number,
    )


def _wrapped_caption(*, enabled: bool):
    document = Document()
    style = document.styles.add_style("Standalone Table Caption", WD_STYLE_TYPE.PARAGRAPH)
    caption = document.add_paragraph(style=style)
    if enabled:
        caption.add_run("Table ")
        append_complex_field(
            caption,
            instruction=" SEQ Table \\* ARABIC ",
            cached_result="3",
        )
        caption.add_run(" Results")
    else:
        caption.add_run("Results")
    identity = _identity(enabled=enabled, derived_number="3" if enabled else None)
    wrap_standalone_caption_occurrence(caption._p, identity)
    [wrapper] = list(document.element.body)[:-1]
    return document, wrapper, caption, style.style_id, identity


@pytest.mark.parametrize(
    ("enabled", "derived_number"),
    ((False, None), (True, "3")),
)
def test_map_round_trips_enabled_and_disabled_occurrences(
    enabled: bool,
    derived_number: str | None,
) -> None:
    identity = _identity(enabled=enabled, derived_number=derived_number)
    data = standalone_caption_occurrence_map_xml([identity])
    root = etree.fromstring(data)

    assert parse_standalone_caption_occurrence_map(root) == [identity]
    assert f'enabled="{"true" if enabled else "false"}"'.encode() in data
    assert f'derived_number="{derived_number or ""}"'.encode() in data


@pytest.mark.parametrize("enabled", [False, True])
def test_single_caption_sdt_is_exact_and_has_no_invented_target(enabled: bool) -> None:
    _document, wrapper, _caption, style_id, identity = _wrapped_caption(enabled=enabled)

    caption = prove_standalone_caption_occurrence_sdt(
        wrapper,
        identity,
        caption_style_id=style_id,
    )

    assert caption.tag.endswith("}p")
    assert not list(caption.iter("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}bookmarkStart"))


def test_disabled_occurrence_rejects_unowned_seq_field() -> None:
    _document, wrapper, caption, style_id, identity = _wrapped_caption(enabled=False)
    append_complex_field(
        caption,
        instruction=" SEQ Table \\* ARABIC ",
        cached_result="9",
    )

    with pytest.raises(DocxSemanticsV3Error, match="unowned fields"):
        prove_standalone_caption_occurrence_sdt(
            wrapper,
            identity,
            caption_style_id=style_id,
        )


def test_enabled_and_derived_number_must_agree() -> None:
    with pytest.raises(DocxSemanticsV3Error, match="requires a derived number"):
        _identity(enabled=True, derived_number=None)
    with pytest.raises(DocxSemanticsV3Error, match="must not carry a derived number"):
        _identity(enabled=False, derived_number="1")
