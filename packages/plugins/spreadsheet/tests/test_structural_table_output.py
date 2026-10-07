"""Spreadsheet→Structural Tables output regressions."""

from __future__ import annotations

from pathlib import Path

import openpyxl
import pytest

from docwen_plugin_markdown.document_semantics import analyze_document_semantics
from docwen_plugin_markdown.mistune_extensions import parse_markdown_text

from ._xlsx_to_md_golden_support import _build_fake_context

pytestmark = pytest.mark.contract


def _convert_structural(tmp_path: Path, workbook: openpyxl.Workbook) -> str:
    source = tmp_path / "structural.xlsx"
    workbook.save(source)
    workbook.close()
    staging = tmp_path / "stage"
    staging.mkdir()

    from docwen_plugin_spreadsheet.to_markdown.converter import SpreadsheetToMarkdownConverter

    context = _build_fake_context(
        str(source),
        str(staging),
        options={"markdown_extensions": {"output": {"structural_tables": True}}},
        source_format="xlsx",
    )
    result = SpreadsheetToMarkdownConverter().convert(context)

    assert result.success, result.error
    return Path(result.artifacts[0].staging_path).read_text(encoding="utf-8")


def _table_source(markdown: str) -> str:
    rows = [line for line in markdown.splitlines() if line.startswith("|")]
    return "\n".join(rows)


def test_literal_merge_markers_are_escaped_in_structural_output(tmp_path: Path) -> None:
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet["A1"] = "Kind"
    sheet["B1"] = "Value"
    sheet["A2"] = "<"
    sheet["B2"] = "^"

    markdown = _convert_structural(tmp_path, workbook)

    assert "| \\< | \\^ |" in markdown
    analysis = analyze_document_semantics(parse_markdown_text(_table_source(markdown)), current_v3=True)
    assert not analysis.has_errors
    metadata = analysis.ast[0]["_document_semantics_table"]
    assert all(anchor["row_span"] == anchor["column_span"] == 1 for anchor in metadata["anchors"])
    text = [
        child["raw"] for anchor in metadata["anchors"] for child in anchor["children"] if child.get("type") == "text"
    ]
    assert "<" in text
    assert "^" in text


def test_merge_crossing_first_row_forces_zero_header_output(tmp_path: Path) -> None:
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet["A1"] = "Merged"
    sheet["B1"] = "X"
    sheet["B2"] = "Y"
    sheet.merge_cells("A1:A2")

    markdown = _convert_structural(tmp_path, workbook)
    table = _table_source(markdown)
    lines = table.splitlines()

    assert lines[0] == "| --- | --- |"
    assert lines[1] == "| Merged | X |"
    assert lines[2] == "| ^ | Y |"

    analysis = analyze_document_semantics(parse_markdown_text(table), current_v3=True)
    assert not analysis.has_errors
    metadata = analysis.ast[0]["_document_semantics_table"]
    assert metadata["header_rows"] == 0
    assert any(
        anchor["row"] == 0 and anchor["column"] == 0 and anchor["row_span"] == 2 and anchor["column_span"] == 1
        for anchor in metadata["anchors"]
    )
