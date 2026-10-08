"""Spreadsheet→Markdown value-fidelity regressions."""

from __future__ import annotations

from pathlib import Path

import openpyxl
import pytest

from ._xlsx_to_md_golden_support import _build_fake_context

pytestmark = pytest.mark.contract


@pytest.mark.parametrize("anchor_value", [0, False])
def test_merged_falsy_anchor_value_is_not_treated_as_empty(anchor_value) -> None:
    from docwen_plugin_spreadsheet.to_markdown.converter import _worksheet_to_dataframe

    workbook = openpyxl.Workbook()
    worksheet = workbook.active
    assert worksheet is not None
    worksheet["A1"] = anchor_value
    worksheet.merge_cells("A1:B1")

    frame = _worksheet_to_dataframe(worksheet, table_merge_strategy="fill")

    expected = str(anchor_value)
    assert frame.iat[0, 0] == expected
    assert frame.iat[0, 1] == expected
    workbook.close()


@pytest.mark.parametrize(
    ("source_format", "suffix"),
    [("xlsx", ".xlsx"), ("csv", ".csv")],
)
def test_numeric_looking_text_keeps_full_authored_precision(
    tmp_path: Path,
    source_format: str,
    suffix: str,
) -> None:
    from docwen_plugin_spreadsheet.to_markdown.converter import SpreadsheetToMarkdownConverter

    value = "0.1234567890123456789"
    source = tmp_path / f"precision{suffix}"
    if source_format == "xlsx":
        workbook = openpyxl.Workbook()
        worksheet = workbook.active
        assert worksheet is not None
        worksheet["A1"] = "Value"
        worksheet["A2"] = value
        worksheet["A2"].data_type = "s"
        workbook.save(source)
        workbook.close()
    else:
        source.write_text(f"Value\n{value}\n", encoding="utf-8")

    staging = tmp_path / f"stage-{source_format}"
    staging.mkdir()
    context = _build_fake_context(
        str(source),
        str(staging),
        source_format=source_format,
    )

    result = SpreadsheetToMarkdownConverter().convert(context)

    assert result.success, result.error
    markdown = Path(result.artifacts[0].staging_path).read_text(encoding="utf-8")
    assert value in markdown
    assert "0.123457" not in markdown
