"""Sparse validation visits stored formulas, not the enclosing worksheet rectangle."""

from __future__ import annotations

from zipfile import ZipFile

import openpyxl
import pytest

from docwen_plugin_spreadsheet.csv_xlsx.converter import _find_unavailable_formula_caches, _load_xlsx_views
from docwen_plugin_spreadsheet.csv_xlsx.formula_cache import visit_formula_caches

pytestmark = pytest.mark.contract


@pytest.mark.parametrize("formula", [False, True])
def test_sparse_hidden_worksheet_does_not_materialize_empty_cells(tmp_path, formula):
    path = tmp_path / "sparse.xlsx"
    book = openpyxl.Workbook()
    sheet = book.create_sheet("Sparse")
    sheet.sheet_state = "hidden"
    sheet.cell(400, 400, "=1+2" if formula else "only stored value")
    book.save(path)
    book.close()
    values, formulas = _load_xlsx_views(str(path))
    try:
        sizes = [len(view["Sparse"]._cells) for view in (values, formulas)]
        result = _find_unavailable_formula_caches(values, formulas, input_path=str(path))
        assert result == ((1, ["Sparse!OJ400"]) if formula else (0, []))
        assert [len(view["Sparse"]._cells) for view in (values, formulas)] == sizes == [1, 1]
        assert _find_unavailable_formula_caches(values, formulas, input_path=str(path), sheet_names=["Sheet"]) == (
            0,
            [],
        )
    finally:
        values.close()
        formulas.close()


def test_scan_handles_implicit_coordinates_and_shared_formula_elements(tmp_path):
    path = tmp_path / "shared.xlsx"
    book = openpyxl.Workbook()
    sheet = book.active
    assert sheet is not None
    sheet.append(["one", "two"])
    sheet.append(["=1+2", "=1+2"])
    book.save(path)
    book.close()
    with ZipFile(path) as package:
        parts = {name: package.read(name) for name in package.namelist()}
    xml = parts["xl/worksheets/sheet1.xml"].decode()
    xml = xml.replace('<row r="2">', "<row>").replace('<c r="B2">', "<c>")
    xml = xml.replace("<f>1+2</f>", '<f t="shared" si="0" ref="A2:B2">1+2</f>', 1)
    xml = xml.replace("<f>1+2</f>", '<f t="shared" si="0"/>', 1)
    parts["xl/worksheets/sheet1.xml"] = xml.encode()
    with ZipFile(path, "w") as package:
        for name, value in parts.items():
            package.writestr(name, value)
    records = []
    visit_formula_caches(str(path), lambda *record: records.append(record))
    assert records == [("Sheet", "A2", False), ("Sheet", "B2", False)]


def test_cancelled_scan_closes_package_and_propagates(tmp_path, monkeypatch):
    path = tmp_path / "cancel.xlsx"
    book = openpyxl.Workbook()
    sheet = book.active
    assert sheet is not None
    sheet["A1"] = "=1+2"
    book.save(path)
    book.close()
    opened = []
    streams = []

    class TrackedZip(ZipFile):
        def __enter__(self):
            opened.append(self)
            return super().__enter__()

        def open(self, *args, **kwargs):
            stream = super().open(*args, **kwargs)
            streams.append(stream)
            return stream

    monkeypatch.setattr("docwen_plugin_spreadsheet.csv_xlsx.formula_cache.ZipFile", TrackedZip)

    def fail_during_visit(*_):
        raise RuntimeError("cancelled")

    with pytest.raises(RuntimeError, match="cancelled"):
        visit_formula_caches(str(path), fail_during_visit)
    assert opened and opened[0].fp is None
    assert streams and all(stream.closed for stream in streams)
