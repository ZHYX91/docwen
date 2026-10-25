"""Real XLSX merge inputs must not publish truncated cell text."""

from __future__ import annotations

import hashlib
from pathlib import Path

import openpyxl
import pytest

from docwen_core.models.file_ref import FileRef
from docwen_core.models.request import ConversionRequest, OutputPolicy
from docwen_core.models.result import ConversionResult
from docwen_plugin_spreadsheet.plugin import SpreadsheetPlugin
from docwen_runtime.engine.route_resolver import RouteResolver
from docwen_runtime.engine.task_manager import TaskManager
from docwen_runtime.output.finalizer import OutputFinalizer
from docwen_runtime.plugin_registry.registry import PluginRegistry
from docwen_runtime.workspace.manager import WorkspaceManager

pytestmark = [pytest.mark.integration, pytest.mark.pr_gate, pytest.mark.release_gate]


def _merge(tmp_path: Path, values: list[str]) -> tuple[ConversionResult, Path]:
    paths = []
    for index, value in enumerate(values):
        path = tmp_path / f"input-{index}.xlsx"
        book = openpyxl.Workbook()
        sheet = book.active
        assert sheet is not None
        sheet.title = "Text"
        sheet.append(["ID", "Value"])
        sheet.append(["row", value])
        book.save(path)
        book.close()
        paths.append(path)

    hashes = {path: hashlib.sha256(path.read_bytes()).hexdigest() for path in paths}
    output = tmp_path / "out"
    output.mkdir()
    # An unrelated existing output must survive a rejected merge as well.
    existing = output / "existing.xlsx"
    existing.write_bytes(b"existing output")
    registry = PluginRegistry()
    registry.register(SpreadsheetPlugin())
    workspace = WorkspaceManager(str(tmp_path / "workspace"))
    task_manager = TaskManager(registry, RouteResolver(registry), workspace, OutputFinalizer())
    request = ConversionRequest(
        request_id="merge-text-limit",
        input_refs=[FileRef(path=str(path), format="xlsx", category="spreadsheet") for path in paths],
        target_format="xlsx",
        action_name="merge_tables",
        options={"merge_mode": "cell", "offset_range": 0},
        output_policy=OutputPolicy(output_dir=str(output)),
    )
    try:
        result = task_manager.execute_single(request)
    finally:
        workspace.cleanup_all()
    assert {path: hashlib.sha256(path.read_bytes()).hexdigest() for path in paths} == hashes
    assert existing.read_bytes() == b"existing output"
    return result, output


def _text(units: int, kind: str) -> str:
    if kind == "non-bmp":
        return "😀" * (units // 2) + "x" * (units % 2)
    unit = {"ascii": "x", "cjk": "中", "multiline": "x\n"}[kind]
    return (unit * units)[:units]


@pytest.mark.parametrize("length", [32766, 32767, 32768])
@pytest.mark.parametrize("kind", ["ascii", "cjk", "non-bmp", "multiline"])
def test_final_merged_text_obeys_utf16_boundary(tmp_path: Path, length: int, kind: str) -> None:
    first = _text(16000, kind)
    second = _text(length - 16001, kind)
    expected = f"{first},{second}"
    assert len(expected.encode("utf-16-le")) // 2 == length

    result, output = _merge(tmp_path, [first, second])

    if length <= 32767:
        assert result.success, result.error
        saved = openpyxl.load_workbook(result.artifacts[0].staging_path)
        try:
            assert saved["Text"]["B2"].value == expected
            assert saved["Text"]["B2"].data_type == "s"
        finally:
            saved.close()
    else:
        assert not result.success
        assert not result.artifacts
        assert result.error is not None
        assert result.error.diagnostic_code == "MERGE-XLSX-CELL-TEXT-TOO-LONG"
        assert "Text!B2" in result.error.message
        assert "32768" in result.error.message
        assert "32767" in result.error.message
        assert "input-0.xlsx" in result.error.message
        assert "input-1.xlsx" in result.error.message
        assert result.diagnostics[0].location == "Text!B2"
        assert sorted(path.name for path in output.iterdir()) == ["existing.xlsx"]


@pytest.mark.parametrize("parts", [[20000, 20000], [16000, 16000, 8000]])
def test_oversized_accumulated_merge_has_no_final_artifact(tmp_path: Path, parts: list[int]) -> None:
    result, output = _merge(tmp_path, [chr(ord("A") + index) * size for index, size in enumerate(parts)])

    assert not result.success
    assert not result.artifacts
    assert result.error is not None
    assert result.error.diagnostic_code == "MERGE-XLSX-CELL-TEXT-TOO-LONG"
    assert str(sum(parts) + len(parts) - 1) in result.error.message
    assert all(f"input-{index}.xlsx" in result.error.message for index in range(len(parts)))
    assert sorted(path.name for path in output.iterdir()) == ["existing.xlsx"]


def test_identical_boundary_text_is_deduplicated_without_false_overflow(tmp_path: Path) -> None:
    value = _text(32767, "non-bmp")
    result, _output = _merge(tmp_path, [value, value])

    assert result.success, result.error
    saved = openpyxl.load_workbook(result.artifacts[0].staging_path)
    try:
        assert saved["Text"]["B2"].value == value
    finally:
        saved.close()
