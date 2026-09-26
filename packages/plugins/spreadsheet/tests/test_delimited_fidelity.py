"""Delimited inputs retain text semantics across spreadsheet routes."""

import csv
from pathlib import Path

import openpyxl
import pytest

from docwen_plugin_spreadsheet.csv_xlsx.converter import (
    DelimitedCellTextTooLongError,
    DelimitedWorkbookDimensionError,
    _build_delimited_workbook,
    _check_delimited_dimensions,
)
from docwen_plugin_spreadsheet.delimited import decoded_samples
from docwen_plugin_spreadsheet.to_markdown.converter import _read_csv_flexible

pytestmark = [pytest.mark.golden, pytest.mark.contract]


@pytest.mark.parametrize("sep", [",", "\t"])
def test_delimited_fidelity_persisted_cells(tmp_path: Path, sep: str) -> None:
    values = ["1234567890123456789", "00123", "0.1234567890123456789", "=1+1", " 42 ", "#N/A"]
    source = tmp_path / "input.txt"
    source.write_text(sep.join(values), encoding="utf-8")
    workbook, rows = _build_delimited_workbook(str(source), sep=sep)
    output = tmp_path / "output.xlsx"
    workbook.save(output)
    workbook.close()
    loaded = openpyxl.load_workbook(output)
    try:
        assert rows == 1
        assert loaded.active is not None
        assert [cell.value for cell in loaded.active[1]] == values
        assert all(cell.data_type == "s" for cell in loaded.active[1])
    finally:
        loaded.close()
    frame = _read_csv_flexible(str(source), "tsv" if sep == "\t" else "csv")
    assert frame.iloc[0].tolist() == values


@pytest.mark.parametrize("encoding", ["utf-8", "utf-8-sig", "utf-16"])
def test_delimited_fidelity_sample_boundary(tmp_path: Path, encoding: str) -> None:
    source = tmp_path / "boundary.csv"
    # Short cells keep the workbook within Excel's cell length limit.
    content = "a,b\n" * 16383 + "a,中文\n"
    source.write_bytes(content.encode(encoding))
    assert decoded_samples(str(source))[0][0] == ("utf-16" if encoding == "utf-16" else "utf-8-sig")
    workbook, rows = _build_delimited_workbook(str(source), sep=",")
    try:
        assert rows == 16384
        assert workbook.active.cell(rows, 2).value == "中文"
    finally:
        workbook.close()
    assert _read_csv_flexible(str(source), "csv").iloc[-1, 1] == "中文"


@pytest.mark.parametrize("sep", [",", "\t"])
def test_delimited_fidelity_gbk_after_ascii_sample(tmp_path: Path, sep: str) -> None:
    source = tmp_path / "late-chinese.txt"
    values = ["中文", "00123", "1234567890123456789", "=1+1"]
    content = ((sep.join(["a"] * 4) + "\n") * 9000) + sep.join(values) + "\n"
    source.write_bytes(content.encode("gbk"))
    assert decoded_samples(str(source))[0][0] == "utf-8-sig"
    workbook, rows = _build_delimited_workbook(str(source), sep=sep)
    output = tmp_path / "late-chinese.xlsx"
    try:
        assert rows == 9001
        workbook.save(output)
    finally:
        workbook.close()
    loaded = openpyxl.load_workbook(output)
    try:
        assert loaded.active is not None
        assert loaded.active.max_row == rows
        assert [cell.value for cell in loaded.active[rows]] == values
        assert all(cell.data_type == "s" for cell in loaded.active[rows])
    finally:
        loaded.close()
    frame = _read_csv_flexible(str(source), "tsv" if sep == "\t" else "csv")
    assert frame.iloc[-1].tolist() == values


def test_delimited_bom_decode_error_is_not_reinterpreted(tmp_path: Path) -> None:
    source = tmp_path / "invalid-utf8.csv"
    source.write_bytes(b"\xef\xbb\xbf" + b"a,b\n" * 17000 + "中文".encode("gbk"))
    with pytest.raises(UnicodeError):
        _build_delimited_workbook(str(source), sep=",")


def test_delimited_cancellation_is_not_retried_as_encoding(tmp_path: Path) -> None:
    source = tmp_path / "cancel.csv"
    source.write_text("a,b\n" * 2000, encoding="utf-8")
    checks = 0

    def cancel() -> None:
        nonlocal checks
        checks += 1
        if checks == 2:
            raise RuntimeError("cancelled")

    with pytest.raises(RuntimeError, match="cancelled"):
        _build_delimited_workbook(str(source), sep=",", cancel_check=cancel)
    assert checks == 2


@pytest.mark.parametrize("sep", [",", "\t"])
@pytest.mark.parametrize("length", [32766, 32767])
@pytest.mark.parametrize("kind", ["ascii", "cjk", "non-bmp", "newline"])
def test_delimited_xlsx_cell_text_limit_accepts_exact_boundary(
    tmp_path: Path,
    sep: str,
    length: int,
    kind: str,
) -> None:
    source = tmp_path / "boundary.txt"
    unit = {"ascii": "x", "cjk": "中", "non-bmp": "😀", "newline": "a\n"}[kind]
    value = (unit * length)[:length]
    if kind == "non-bmp":
        value = "😀" * (length // 2) + "a" * (length % 2)
    assert len(value.encode("utf-16-le")) // 2 == length
    with source.open("w", encoding="utf-8", newline="") as handle:
        csv.writer(handle, delimiter=sep).writerow([value])
    workbook, rows = _build_delimited_workbook(str(source), sep=sep)
    try:
        assert rows == 1
        assert workbook.active is not None
        assert workbook.active.cell(1, 1).value == value
        output = tmp_path / "boundary.xlsx"
        workbook.save(output)
    finally:
        workbook.close()
    loaded = openpyxl.load_workbook(output)
    try:
        assert loaded.active is not None
        assert loaded.active.cell(1, 1).value == value
    finally:
        loaded.close()


@pytest.mark.parametrize("sep", [",", "\t"])
@pytest.mark.parametrize("value_kind", ["ascii", "cjk", "emoji", "newline"])
@pytest.mark.parametrize("length", [32768, 40000])
def test_delimited_xlsx_cell_text_limit_rejects_without_truncation(
    tmp_path: Path,
    sep: str,
    value_kind: str,
    length: int,
) -> None:
    values = {
        "ascii": "x" * length,
        "cjk": "中" * length,
        "emoji": "😀" * (length // 2) + "a" * (length % 2),
        "newline": ("line\n" * length)[:length],
    }
    value = values[value_kind]
    source = tmp_path / "too-long.txt"
    with source.open("w", encoding="utf-8", newline="") as handle:
        csv.writer(handle, delimiter=sep).writerow([value])

    with pytest.raises(DelimitedCellTextTooLongError) as rejected:
        _build_delimited_workbook(str(source), sep=sep)

    assert rejected.value.row == 1
    assert rejected.value.column == 1
    assert rejected.value.length == length
    assert rejected.value.limit == 32767
    assert value[:32] not in str(rejected.value)


@pytest.mark.parametrize("sep", [",", "\t"])
@pytest.mark.parametrize("columns", [16383, 16384])
def test_delimited_xlsx_column_limit_accepts_native_boundary(
    tmp_path: Path,
    sep: str,
    columns: int,
) -> None:
    source = tmp_path / "wide-ok.txt"
    values = [f"v{index}" for index in range(columns)]
    with source.open("w", encoding="utf-8", newline="") as handle:
        csv.writer(handle, delimiter=sep).writerow(values)

    workbook, rows = _build_delimited_workbook(str(source), sep=sep)
    try:
        assert rows == 1
        assert workbook.active is not None
        assert workbook.active.max_column == columns
        assert workbook.active.cell(1, columns).value == values[-1]
    finally:
        workbook.close()


@pytest.mark.parametrize("sep", [",", "\t"])
def test_delimited_xlsx_column_limit_rejects_16385_columns(
    tmp_path: Path,
    sep: str,
) -> None:
    source = tmp_path / "wide-rejected.txt"
    with source.open("w", encoding="utf-8", newline="") as handle:
        csv.writer(handle, delimiter=sep).writerow(["x"] * 16385)

    with pytest.raises(DelimitedWorkbookDimensionError) as rejected:
        _build_delimited_workbook(str(source), sep=sep)

    assert rejected.value.axis == "column"
    assert rejected.value.actual == 16385
    assert rejected.value.limit == 16384
    assert rejected.value.row == 1


@pytest.mark.parametrize(("row", "columns"), [(1048576, 16384), (1, 16384)])
def test_delimited_xlsx_dimension_check_accepts_native_limits(row: int, columns: int) -> None:
    _check_delimited_dimensions(row=row, columns=columns)


def test_delimited_xlsx_dimension_check_rejects_row_1048577() -> None:
    with pytest.raises(DelimitedWorkbookDimensionError) as rejected:
        _check_delimited_dimensions(row=1048577, columns=1)

    assert rejected.value.axis == "row"
    assert rejected.value.actual == 1048577
    assert rejected.value.limit == 1048576
    assert rejected.value.row is None
