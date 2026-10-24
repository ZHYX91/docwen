"""Content invariants for real XLSX merges, independent of historical output."""

from __future__ import annotations

import re
from collections.abc import Sequence
from pathlib import Path
from typing import Any
from zipfile import ZipFile

import openpyxl
import pytest
from tests.support.config import FakeConfigView
from tests.support.execution import FakeExecutionContext
from tests.support.logging import FakePluginLogger
from tests.support.progress import FakeProgressSink
from tests.support.workspace import FakeWorkspaceHandle

from docwen_core.cancellation import CancellationToken
from docwen_core.models.file_ref import FileRef
from docwen_core.models.request import ConversionRequest
from docwen_plugin_spreadsheet.table_merger.converter import TableMergerConverter
from docwen_plugin_spreadsheet.table_merger.semantics import value_identity

pytestmark = pytest.mark.contract


def _write(path: Path, rows: Sequence[Sequence[Any]], mode: str = "row") -> None:
    book = openpyxl.Workbook()
    sheet = book.active
    assert sheet is not None
    for row in zip(*rows, strict=True) if mode == "col" else rows:
        sheet.append(tuple(row))
    book.save(path)
    book.close()


def _run(tmp_path: Path, mode: str = "row", offset_range: int = 0):
    paths = [tmp_path / "base.xlsx", tmp_path / "collect.xlsx"]
    staging = tmp_path / "staging"
    staging.mkdir(exist_ok=True)
    context = FakeExecutionContext(
        request=ConversionRequest(
            request_id="merge-semantics",
            input_refs=[FileRef(path=str(p), format="xlsx", category="spreadsheet") for p in paths],
            target_format="xlsx",
            action_name="merge_tables",
            options={"merge_mode": mode, "offset_range": offset_range},
        ),
        workspace=FakeWorkspaceHandle(str(paths[0]), str(staging)),
        config=FakeConfigView(),
        progress=FakeProgressSink(),
        cancellation=CancellationToken(),
        logger=FakePluginLogger(),
    )
    return TableMergerConverter().convert(context)


def _records(path: str, mode: str) -> Sequence[tuple[object, ...]]:
    book = openpyxl.load_workbook(path, data_only=True)
    try:
        sheet = book.active
        assert sheet is not None
        rows = list(sheet.values)
        return list(zip(*rows, strict=True)) if mode == "col" else rows
    finally:
        book.close()


@pytest.mark.parametrize("mode", ["row", "col"])
@pytest.mark.parametrize(
    ("base", "collect", "expected", "decision"),
    [
        (["Alice", 7, None], ["Alice", 7, None], [("Alice", 7, None)], "skipped"),
        (["Alice", 7, None], ["Alice", None, None], [("Alice", 7, None)], "skipped"),
        (["Alice", None, None], ["Alice", 7, None], [("Alice", 7, None)], "replaced"),
        (["Alice", 7, None], ["Alice", None, 9], [("Alice", 7, None), ("Alice", None, 9)], "added"),
        (["Alice", 7, None], ["Alice", 8, None], [("Alice", 7, None), ("Alice", 8, None)], "added"),
    ],
)
def test_each_nonempty_field_survives_coverage_decisions(tmp_path, mode, base, collect, expected, decision):
    header = ["Name", "Old", "New"]
    _write(tmp_path / "base.xlsx", [header, base], mode)
    _write(tmp_path / "collect.xlsx", [header, collect], mode)
    originals = {p: p.read_bytes() for p in tmp_path.glob("*.xlsx")}
    result = _run(tmp_path, mode)
    assert result.success
    assert _records(result.artifacts[0].staging_path, mode) == [tuple(header), *expected]
    summary = result.metrics.extra["merge_summaries"][0]
    assert summary["counts"][decision] == (2 if decision == "skipped" else 1)
    assert summary["locations"][decision][-1] == ("row 2" if mode == "row" else "column 2")
    assert all(path.read_bytes() == before for path, before in originals.items())


@pytest.mark.parametrize("mode", ["row", "col"])
@pytest.mark.parametrize(
    ("first", "second"),
    [
        ("110101199001010011", "110101199001010012"),
        ("001", "1"),
        ("1e3", "1000"),
        ("1", 1),
        (True, 1),
        (" A", "A"),
        ("alpha", "beta"),
    ],
)
def test_distinct_identities_are_never_implicitly_deduplicated(tmp_path, mode, first, second):
    _write(tmp_path / "base.xlsx", [["ID", "Value"], [first, "A"]], mode)
    _write(tmp_path / "collect.xlsx", [["ID", "Value"], [second, "A"]], mode)
    result = _run(tmp_path, mode)
    assert result.success
    records = _records(result.artifacts[0].staging_path, mode)
    assert len(records) == 3
    assert {value_identity(row[0]) for row in records[1:]} == {value_identity(first), value_identity(second)}


def test_numeric_identity_does_not_round_integers_before_comparison():
    assert value_identity(2**53 + 1) != value_identity(float(2**53))
    assert value_identity(7) == value_identity(7.0)


@pytest.mark.parametrize("mode", ["row", "col", "cell"])
@pytest.mark.parametrize("formula_source", ["base", "collect"])
def test_missing_formula_cache_stops_before_creating_output(tmp_path, mode, formula_source):
    for name in ("base", "collect"):
        _write(tmp_path / f"{name}.xlsx", [["ID", "Value"], ["Alice", "=1+2" if name == formula_source else 4]])
    result = _run(tmp_path, mode)
    assert not result.success
    assert result.artifacts == []
    assert result.error.diagnostic_code == "MERGE-FORMULA-CACHE-MISSING"
    assert f"{formula_source}.xlsx" in result.error.message
    assert "Sheet!B2" in result.diagnostics[0].location
    assert "1 formula" in result.error.message
    assert not list((tmp_path / "staging").iterdir())


@pytest.mark.parametrize(("cache_type", "cache", "expected"), [("n", "3", 7), ("str", "", 4)])
def test_cached_number_and_cached_empty_string_are_distinguished_from_unknown(tmp_path, cache_type, cache, expected):
    base = tmp_path / "base.xlsx"
    formula = '""' if cache_type == "str" else "1+2"
    _write(base, [["ID", "Value"], ["Alice", f"={formula}"]])
    with ZipFile(base) as archive:
        parts = {name: archive.read(name) for name in archive.namelist()}
    xml = parts["xl/worksheets/sheet1.xml"].decode()
    xml, count = re.subn(
        r'<c r="B2"[^>]*><f>.*?</f><v></v></c>', f'<c r="B2" t="{cache_type}"><f>{formula}</f><v>{cache}</v></c>', xml
    )
    assert count == 1
    parts["xl/worksheets/sheet1.xml"] = xml.encode()
    with ZipFile(base, "w") as archive:
        for name, content in parts.items():
            archive.writestr(name, content)
    _write(tmp_path / "collect.xlsx", [["ID", "Value"], ["Alice", 4]])
    result = _run(tmp_path, "cell")
    assert result.success
    assert _records(result.artifacts[0].staging_path, "row")[1][1] == expected


@pytest.mark.parametrize("mode", ["row", "col"])
@pytest.mark.parametrize("leading_source", ["base", "collect"])
def test_alignment_preserves_leading_fields_and_complementary_values(tmp_path, mode, leading_source):
    base = [["ID", "Old", "New"], ["Alice", 7, None]]
    collect = [["ID", "Old", "New"], ["Alice", None, 9]]
    leading = base if leading_source == "base" else collect
    leading[0].insert(0, "Extra")
    leading[1].insert(0, "keep")
    _write(tmp_path / "base.xlsx", base, mode)
    _write(tmp_path / "collect.xlsx", collect, mode)
    result = _run(tmp_path, mode, offset_range=2)
    assert result.success
    values = [value for row in _records(result.artifacts[0].staging_path, mode) for value in row]
    assert all(value in values for value in (7, 9, "keep", "Extra"))


@pytest.mark.parametrize(("source", "success"), [("base", False), ("collect", True)])
def test_cache_policy_follows_the_sheets_retained_in_the_output(tmp_path, source, success):
    for name in ("base", "collect"):
        path = tmp_path / f"{name}.xlsx"
        _write(path, [["ID", "Value"], ["Alice", 3]])
        if name == source:
            book = openpyxl.load_workbook(path)
            extra = book.create_sheet("Extra")
            extra["A1"] = "=1+2"
            book.save(path)
            book.close()
    result = _run(tmp_path, "cell")
    assert result.success is success
    if success:
        book = openpyxl.load_workbook(result.artifacts[0].staging_path)
        try:
            assert book.sheetnames == ["Sheet"]
        finally:
            book.close()
    else:
        assert "Extra!A1" in result.error.message
