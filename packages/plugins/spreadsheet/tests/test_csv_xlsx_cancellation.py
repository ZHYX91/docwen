"""Cancellation and large-shape contracts for delimited/XLSX conversion."""

from __future__ import annotations

import csv
from pathlib import Path

import openpyxl
import pytest

from docwen_core.errors import CancellationRequested
from docwen_plugin_spreadsheet.csv_xlsx.converter import (
    CsvToXlsxConverter,
    TsvToXlsxConverter,
    XlsxToCsvConverter,
    XlsxToTsvConverter,
    _build_delimited_workbook,
)

from ._csv_xlsx_support import _build_fake_context

pytestmark = [pytest.mark.golden, pytest.mark.contract]


class _CancelAfterChecks:
    def __init__(self, fail_at: int) -> None:
        self.fail_at = fail_at
        self.checks = 0

    @property
    def is_cancelled(self) -> bool:
        return self.checks >= self.fail_at

    def check(self) -> None:
        self.checks += 1
        if self.checks >= self.fail_at:
            raise CancellationRequested("test cancellation")


@pytest.mark.parametrize(
    ("converter_type", "suffix", "sep"),
    [
        (CsvToXlsxConverter, ".csv", ","),
        (TsvToXlsxConverter, ".tsv", "\t"),
    ],
)
def test_delimited_to_xlsx_propagates_mid_parse_cancellation(
    tmp_path: Path,
    converter_type: type,
    suffix: str,
    sep: str,
) -> None:
    source = tmp_path / f"large{suffix}"
    with source.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.writer(handle, delimiter=sep)
        for index in range(2500):
            writer.writerow([index, f"value-{index}"])

    staging = tmp_path / "staging"
    staging.mkdir()
    context = _build_fake_context(str(source), str(staging), target_format="xlsx")
    cancellation = _CancelAfterChecks(fail_at=3)
    context._cancellation = cancellation

    with pytest.raises(CancellationRequested, match="test cancellation"):
        converter_type().convert(context)

    assert cancellation.checks == 3
    assert context.workspace.registered_artifacts == []
    assert not list(staging.glob("*.xlsx"))


@pytest.mark.parametrize(
    ("converter_type", "target_format", "suffix"),
    [
        (XlsxToCsvConverter, "csv", ".csv"),
        (XlsxToTsvConverter, "tsv", ".tsv"),
    ],
)
def test_xlsx_to_delimited_propagates_mid_write_cancellation_without_registering_partial_artifact(
    tmp_path: Path,
    converter_type: type,
    target_format: str,
    suffix: str,
) -> None:
    source = tmp_path / "large.xlsx"
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    assert sheet is not None
    for index in range(2500):
        sheet.append([index, f"value-{index}"])
    workbook.save(source)
    workbook.close()

    staging = tmp_path / "staging"
    staging.mkdir()
    context = _build_fake_context(str(source), str(staging), target_format=target_format)

    class _CancelAfterOutputStarts(_CancelAfterChecks):
        def check(self) -> None:
            if any(path.stat().st_size > 0 for path in staging.glob(f"*{suffix}")):
                super().check()

    cancellation = _CancelAfterOutputStarts(fail_at=1)
    context._cancellation = cancellation

    with pytest.raises(CancellationRequested, match="test cancellation"):
        converter_type().convert(context)

    assert cancellation.checks == 1
    assert context.workspace.registered_artifacts == []
    partials = list(staging.glob(f"*{suffix}"))
    assert len(partials) == 1
    assert partials[0].stat().st_size > 0


@pytest.mark.parametrize("sep", [",", "\t"])
def test_single_wide_row_can_cancel_before_completing(tmp_path: Path, sep: str) -> None:
    source = tmp_path / "wide.txt"
    source.write_text(sep.join(["value"] * 16000), encoding="utf-8")
    cancellation = _CancelAfterChecks(2)
    with pytest.raises(CancellationRequested):
        _build_delimited_workbook(str(source), sep=sep, cancel_check=cancellation.check)
    assert cancellation.checks == 2


@pytest.mark.parametrize("converter", [CsvToXlsxConverter, TsvToXlsxConverter])
def test_cancel_after_xlsx_save_closes_workbook_and_does_not_register(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch, converter
) -> None:
    source = tmp_path / "input.csv"
    source.write_text("literal", encoding="utf-8")
    staging = tmp_path / "staging"
    staging.mkdir()
    context = _build_fake_context(str(source), str(staging), target_format="xlsx")
    cancellation = _CancelAfterChecks(1)
    context._cancellation = cancellation
    saved = False
    closed = False
    original_save = openpyxl.Workbook.save
    original_close = openpyxl.Workbook.close

    def save(workbook, path):
        nonlocal saved
        original_save(workbook, path)
        saved = True

    def close(workbook):
        nonlocal closed
        closed = True
        original_close(workbook)

    def check():
        if saved:
            raise CancellationRequested("after save")

    monkeypatch.setattr(cancellation, "check", check)
    monkeypatch.setattr(openpyxl.Workbook, "save", save)
    monkeypatch.setattr(openpyxl.Workbook, "close", close)
    with pytest.raises(CancellationRequested, match="after save"):
        converter().convert(context)
    assert saved and closed
    assert context.workspace.registered_artifacts == []
    partials = list(staging.glob("*.xlsx"))
    assert len(partials) == 1
    assert partials[0].stat().st_size > 0


@pytest.mark.parametrize("sep", [",", "\t"])
def test_delimited_wide_row_preserves_all_columns_without_timing_assumptions(
    tmp_path: Path,
    sep: str,
) -> None:
    source = tmp_path / "wide.txt"
    values = [f"c{index:04d}" for index in range(1024)]
    with source.open("w", encoding="utf-8", newline="") as handle:
        csv.writer(handle, delimiter=sep).writerow(values)

    workbook, row_count = _build_delimited_workbook(str(source), sep=sep)
    try:
        assert row_count == 1
        sheet = workbook.active
        assert sheet is not None
        assert sheet.max_column == len(values)
        assert sheet.cell(1, 1).value == values[0]
        assert sheet.cell(1, len(values)).value == values[-1]
    finally:
        workbook.close()


def test_formula_cache_scan_cancellation_closes_both_workbook_views(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    from docwen_plugin_spreadsheet.csv_xlsx import converter as converter_module

    source = tmp_path / "input.xlsx"
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet["A1"] = "Value"
    sheet["A2"] = "=1+1"
    workbook.save(source)
    workbook.close()
    original_source = source.read_bytes()
    staging = tmp_path / "staging"
    staging.mkdir()
    context = _build_fake_context(str(source), str(staging), target_format="csv")
    loaded: list[openpyxl.Workbook] = []
    closed: set[int] = set()
    load_views = converter_module._load_xlsx_views

    def load_then_cancel(path: str):
        values, formulas = load_views(path)
        loaded.extend((values, formulas))
        assert values.worksheets[0]["A2"].value is None
        assert formulas.worksheets[0]["A2"].value == "=1+1"
        for view in (values, formulas):
            original_close = view.close
            view_identity = id(view)

            def close(*, close_view=original_close, identity=view_identity) -> None:
                closed.add(identity)
                close_view()

            monkeypatch.setattr(view, "close", close)
        context.cancellation.cancel("scan cancellation")
        return values, formulas

    monkeypatch.setattr(converter_module, "_load_xlsx_views", load_then_cancel)

    # The real cache scanner receives the cancelled token after both native
    # views are open. No placeholder ZIP or replacement scanner can hide leaks.
    with pytest.raises(CancellationRequested, match="Task has been cancelled"):
        XlsxToCsvConverter().convert(context)

    assert len(loaded) == 2
    assert closed == {id(view) for view in loaded}
    assert context.workspace.registered_artifacts == []
    assert not list(staging.glob("*.csv"))
    assert source.read_bytes() == original_source
