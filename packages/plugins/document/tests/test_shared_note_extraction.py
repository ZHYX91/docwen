from zipfile import ZipFile

import mistune
import pytest
from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from lxml import etree

from docwen_core.docx_parsing.format_features import DocxMarkdownSyntaxConfig, StyleDetectorConfig
from docwen_plugin_document.shared.note_extraction import (
    NoteExtractor,
    _extract_note_content,
    build_note_definitions,
)

pytestmark = pytest.mark.unit

WML_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
XML_NS = "http://www.w3.org/XML/1998/namespace"


@pytest.mark.parametrize("ref_tag", ["footnoteRef", "endnoteRef"])
def test_note_content_removes_only_structural_reference_separator(ref_tag: str):
    note = etree.Element(f"{{{WML_NS}}}note")
    first = etree.SubElement(note, f"{{{WML_NS}}}p")
    marker_run = etree.SubElement(first, f"{{{WML_NS}}}r")
    etree.SubElement(marker_run, f"{{{WML_NS}}}{ref_tag}")
    separator_run = etree.SubElement(first, f"{{{WML_NS}}}r")
    separator = etree.SubElement(separator_run, f"{{{WML_NS}}}t")
    separator.set(f"{{{XML_NS}}}space", "preserve")
    separator.text = " "
    authored_run = etree.SubElement(first, f"{{{WML_NS}}}r")
    authored = etree.SubElement(authored_run, f"{{{WML_NS}}}t")
    authored.set(f"{{{XML_NS}}}space", "preserve")
    authored.text = "  Authored leading whitespace"

    second = etree.SubElement(note, f"{{{WML_NS}}}p")
    continuation_run = etree.SubElement(second, f"{{{WML_NS}}}r")
    continuation = etree.SubElement(continuation_run, f"{{{WML_NS}}}t")
    continuation.set(f"{{{XML_NS}}}space", "preserve")
    continuation.text = " Continuation leading whitespace"

    assert _extract_note_content(note, WML_NS, ref_tag) == (
        "  Authored leading whitespace\n Continuation leading whitespace"
    )


def test_note_content_preserves_breaks_and_basic_inline_formatting():
    note = etree.Element(f"{{{WML_NS}}}note")
    paragraph = etree.SubElement(note, f"{{{WML_NS}}}p")

    first = etree.SubElement(paragraph, f"{{{WML_NS}}}r")
    etree.SubElement(first, f"{{{WML_NS}}}t").text = "First line"

    break_run = etree.SubElement(paragraph, f"{{{WML_NS}}}r")
    etree.SubElement(break_run, f"{{{WML_NS}}}br")

    bold_run = etree.SubElement(paragraph, f"{{{WML_NS}}}r")
    bold_props = etree.SubElement(bold_run, f"{{{WML_NS}}}rPr")
    etree.SubElement(bold_props, f"{{{WML_NS}}}b")
    etree.SubElement(bold_run, f"{{{WML_NS}}}t").text = "Bold line"

    assert _extract_note_content(note, WML_NS, "footnoteRef") == "First line\n**Bold line**"


def test_note_content_preserves_codespan_with_embedded_backtick():
    note = etree.Element(f"{{{WML_NS}}}note")
    paragraph = etree.SubElement(note, f"{{{WML_NS}}}p")
    run = etree.SubElement(paragraph, f"{{{WML_NS}}}r")
    props = etree.SubElement(run, f"{{{WML_NS}}}rPr")
    fonts = etree.SubElement(props, f"{{{WML_NS}}}rFonts")
    fonts.set(f"{{{WML_NS}}}ascii", "Consolas")
    fonts.set(f"{{{WML_NS}}}hAnsi", "Consolas")
    shading = etree.SubElement(props, f"{{{WML_NS}}}shd")
    shading.set(f"{{{WML_NS}}}fill", "D9D9D9")
    etree.SubElement(run, f"{{{WML_NS}}}t").text = "a`b"

    assert _extract_note_content(note, WML_NS, "footnoteRef") == "``a`b``"


@pytest.mark.parametrize("ref_tag", ["footnoteRef", "endnoteRef"])
@pytest.mark.parametrize("value", ["value", "a`b", "`value", "value`", "`value`", " value ", "   "])
def test_note_codespan_reparses_with_exact_literal_value(ref_tag, value):
    note = etree.Element(f"{{{WML_NS}}}note")
    paragraph = etree.SubElement(note, f"{{{WML_NS}}}p")
    # Word can split a code span across otherwise identical runs when saving.
    for text in [value[:1], value[1:]]:
        run = etree.SubElement(paragraph, f"{{{WML_NS}}}r")
        props = etree.SubElement(run, f"{{{WML_NS}}}rPr")
        fonts = etree.SubElement(props, f"{{{WML_NS}}}rFonts")
        fonts.set(f"{{{WML_NS}}}ascii", "Consolas")
        shading = etree.SubElement(props, f"{{{WML_NS}}}shd")
        shading.set(f"{{{WML_NS}}}fill", "D9D9D9")
        etree.SubElement(run, f"{{{WML_NS}}}t").text = text
    markdown = _extract_note_content(note, WML_NS, ref_tag)
    ast = mistune.create_markdown(renderer="ast")(markdown)
    assert ast == [{"type": "paragraph", "children": [{"type": "codespan", "raw": value}]}]
    assert _extract_note_content(note, WML_NS, ref_tag, preserve_formatting=False) == value


@pytest.mark.parametrize("ref_tag", ["footnoteRef", "endnoteRef"])
@pytest.mark.parametrize("preserve", [False, True])
def test_note_content_uses_selected_formatting_and_syntax(ref_tag, preserve):
    note = etree.Element(f"{{{WML_NS}}}note")
    paragraph = etree.SubElement(note, f"{{{WML_NS}}}p")
    for property_name, text in [("b", "Bold"), ("i", "Italic"), ("strike", "Strike")]:
        run = etree.SubElement(paragraph, f"{{{WML_NS}}}r")
        properties = etree.SubElement(run, f"{{{WML_NS}}}rPr")
        etree.SubElement(properties, f"{{{WML_NS}}}{property_name}")
        etree.SubElement(run, f"{{{WML_NS}}}t").text = text
    actual = _extract_note_content(
        note,
        WML_NS,
        ref_tag,
        preserve_formatting=preserve,
        syntax_config=DocxMarkdownSyntaxConfig(bold="underscore", italic="underscore", strikethrough="html"),
    )
    assert actual == ("__Bold___Italic_<del>Strike</del>" if preserve else "BoldItalicStrike")


@pytest.mark.parametrize("note_tag", ["footnote", "endnote"])
@pytest.mark.parametrize("style_name", ["Inline Code", "行内代码", "LiteralSnippet"])
@pytest.mark.parametrize("preserve", [False, True])
def test_note_character_style_from_real_document_uses_body_policy(tmp_path, note_tag, style_name, preserve):
    document = Document()
    style = document.styles.add_style(style_name, WD_STYLE_TYPE.CHARACTER)
    body_run = document.add_paragraph().add_run("`value`")
    body_run.style = style_name
    path = tmp_path / "styled-notes.docx"
    document.save(path)
    root = etree.Element(f"{{{WML_NS}}}{note_tag}s", nsmap={"w": WML_NS})
    note = etree.SubElement(root, f"{{{WML_NS}}}{note_tag}")
    note.set(f"{{{WML_NS}}}id", "1")
    paragraph = etree.SubElement(note, f"{{{WML_NS}}}p")
    run = etree.SubElement(paragraph, f"{{{WML_NS}}}r")
    properties = etree.SubElement(run, f"{{{WML_NS}}}rPr")
    applied_style = etree.SubElement(properties, f"{{{WML_NS}}}rStyle")
    applied_style.set(f"{{{WML_NS}}}val", style.style_id)
    etree.SubElement(run, f"{{{WML_NS}}}t").text = "`value`"
    with ZipFile(path, "a") as package:
        package.writestr(f"word/{note_tag}s.xml", etree.tostring(root))
    extractor = NoteExtractor(
        Document(path),
        str(path),
        preserve_formatting=preserve,
        style_detector_config=StyleDetectorConfig(code_character_style_names=frozenset({"LiteralSnippet"})),
    )
    notes = extractor.footnotes if note_tag == "footnote" else extractor.endnotes
    assert notes == {1: "`` `value` ``" if preserve else "`value`"}


def test_build_note_definitions_formats_multiline_content():
    notes = {5: "第一行\n第二行"}
    assert build_note_definitions(notes, {5: "1"}) == "[^1]: 第一行\n    第二行"


def test_endnote_definitions_use_prefix():
    notes = {9: "尾注"}
    assert build_note_definitions(notes, {9: "endnote:1"}) == "[^endnote:1]: 尾注"


def test_note_extractor_footnotes_default_empty():
    extractor = NoteExtractor.__new__(NoteExtractor)
    extractor.footnotes = {}
    extractor.endnotes = {}
    extractor.footnote_id_map = {1: "1"}
    extractor.endnote_id_map = {1: "endnote:1"}

    assert extractor.get_reference_text("footnote", 1) == "[^1]"
    assert extractor.get_reference_text("endnote", 1) == "[^endnote:1]"
    assert extractor.build_definitions_block() == ""


def test_note_extractor_builds_definitions_block():
    extractor = NoteExtractor.__new__(NoteExtractor)
    extractor.footnotes = {7: "脚注内容"}
    extractor.endnotes = {8: "尾注内容"}
    extractor.footnote_id_map = {7: "1"}
    extractor.endnote_id_map = {8: "endnote:1"}

    block = extractor.build_definitions_block()
    assert "[^1]: 脚注内容" in block
    assert "[^endnote:1]: 尾注内容" in block


def test_note_extractor_numbers_each_domain_by_first_reference():
    extractor = NoteExtractor.__new__(NoteExtractor)
    extractor.footnotes = {40: "later Word ID", 3: "earlier Word ID"}
    extractor.endnotes = {90: "later Word ID", 2: "earlier Word ID"}
    extractor.footnote_id_map = {}
    extractor.endnote_id_map = {}

    assert extractor.get_reference_text("footnote", 40) == "[^1]"
    assert extractor.get_reference_text("footnote", 3) == "[^2]"
    assert extractor.get_reference_text("footnote", 40) == "[^1]"
    assert extractor.get_reference_text("endnote", 90) == "[^endnote:1]"
    assert extractor.get_reference_text("endnote", 2) == "[^endnote:2]"

    block = extractor.build_definitions_block()
    assert block.index("[^1]: later Word ID") < block.index("[^2]: earlier Word ID")
    assert block.index("[^endnote:1]: later Word ID") < block.index("[^endnote:2]: earlier Word ID")
