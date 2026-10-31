"""Real image parts preserve occurrence identity and extraction cleanup."""

from __future__ import annotations

import hashlib
from io import BytesIO
from pathlib import Path

import pytest
from docx import Document
from PIL import Image

from docwen_core.docx_parsing.image_extraction import extract_images_from_element

pytestmark = pytest.mark.unit


def _png(color: tuple[int, int, int]) -> bytes:
    with BytesIO() as buffer, Image.new("RGB", (3, 2), color) as image:
        image.save(buffer, format="PNG")
        return buffer.getvalue()


@pytest.mark.parametrize("preoccupied", [False, True])
def test_independent_element_extractions_cannot_reuse_an_existing_source_name(
    tmp_path: Path, preoccupied: bool
) -> None:
    document = Document()
    originals = [_png(color) for color in ((210, 20, 20), (20, 210, 20), (20, 20, 210), (50, 50, 50))]
    for index, original in enumerate(originals):
        path = tmp_path / f"source-{index}.png"
        path.write_bytes(original)
        document.add_paragraph().add_run().add_picture(str(path))._inline.docPr.set("title", "photo.png")
    output = tmp_path / "extracted"
    output.mkdir()
    if preoccupied:
        (output / "photo_1.png").write_bytes(b"preexisting output must survive")
    infos = [
        extract_images_from_element(paragraph._p, document.part.related_parts, str(output), prefer_source_names=True)[0]
        for paragraph in document.paragraphs
    ]
    assert [hashlib.sha256(Path(info.path).read_bytes()).digest() for info in infos] == [
        hashlib.sha256(data).digest() for data in originals
    ]
    assert len({info.path for info in infos}) == len(originals)
    if preoccupied:
        assert (output / "photo_1.png").read_bytes() == b"preexisting output must survive"


def test_failed_extraction_removes_only_its_new_files_and_preserves_preoccupied_suffix(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    document = Document()
    original = _png((50, 90, 130))
    image = tmp_path / "original.png"
    image.write_bytes(original)
    paragraph = document.add_paragraph()
    for _ in range(3):
        paragraph.add_run().add_picture(str(image))._inline.docPr.set("title", "photo.png")
    output = tmp_path / "extracted"
    output.mkdir()
    occupied = output / "photo_1.png"
    occupied.write_bytes(b"keep this original")
    real_write = Path.write_bytes

    def fail_new_suffix(path: Path, content: bytes) -> int:
        if path == output / "photo_2.png":
            real_write(path, b"partial new image")
            raise OSError("injected resource write failure")
        return real_write(path, content)

    monkeypatch.setattr(Path, "write_bytes", fail_new_suffix)
    with pytest.raises(OSError, match="injected resource write failure"):
        extract_images_from_element(paragraph._p, document.part.related_parts, str(output), prefer_source_names=True)
    assert occupied.read_bytes() == b"keep this original"
    assert image.read_bytes() == original
    assert [path.name for path in output.iterdir()] == ["photo_1.png"]
