"""Word-native numbering contracts for Gongwen heading extraction."""

from __future__ import annotations

from typing import Any

import pytest

pytestmark = pytest.mark.contract


def _install_numbering_level(
    doc: Any,
    *,
    abstract_id: str,
    num_id: str,
    num_fmt: str,
    level_text: str,
    paragraph_style: str | None = None,
    start_at: int = 1,
    list_level: int = 0,
) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    abstract_num = OxmlElement("w:abstractNum")
    abstract_num.set(qn("w:abstractNumId"), abstract_id)
    level = OxmlElement("w:lvl")
    level.set(qn("w:ilvl"), str(list_level))
    start = OxmlElement("w:start")
    start.set(qn("w:val"), str(start_at))
    level.append(start)
    fmt = OxmlElement("w:numFmt")
    fmt.set(qn("w:val"), num_fmt)
    level.append(fmt)
    if paragraph_style:
        style = OxmlElement("w:pStyle")
        style.set(qn("w:val"), paragraph_style)
        level.append(style)
    text = OxmlElement("w:lvlText")
    text.set(qn("w:val"), level_text)
    level.append(text)
    suffix = OxmlElement("w:suff")
    suffix.set(qn("w:val"), "space")
    level.append(suffix)
    abstract_num.append(level)

    numbering_root = doc.part.numbering_part.element
    first_num_index = next(
        (index for index, child in enumerate(numbering_root) if child.tag == qn("w:num")),
        len(numbering_root),
    )
    numbering_root.insert(first_num_index, abstract_num)
    num = OxmlElement("w:num")
    num.set(qn("w:numId"), num_id)
    abstract_ref = OxmlElement("w:abstractNumId")
    abstract_ref.set(qn("w:val"), abstract_id)
    num.append(abstract_ref)
    numbering_root.append(num)


def _inject_numpr(paragraph: Any, *, num_id: str, list_level: int = 0) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    num_properties = OxmlElement("w:numPr")
    level = OxmlElement("w:ilvl")
    level.set(qn("w:val"), str(list_level))
    identifier = OxmlElement("w:numId")
    identifier.set(qn("w:val"), num_id)
    num_properties.extend((level, identifier))
    paragraph._p.get_or_add_pPr().append(num_properties)


def test_native_chinese_numbering_becomes_a_retained_gongwen_heading(tmp_path) -> None:
    from docx import Document

    from docwen_core.docx_parsing.numbering_index import NumberingIndex
    from docwen_plugin_optimizer_gongwen.extraction.paragraph_reader import read_paragraphs
    from docwen_plugin_optimizer_gongwen.models import GongwenMetadata
    from docwen_plugin_optimizer_gongwen.rendering.markdown_renderer import render

    document = Document()
    _install_numbering_level(
        document,
        abstract_id="93001",
        num_id="94001",
        num_fmt="chineseCounting",
        level_text="%1、",
    )
    first = document.add_paragraph("公文写作核心原则")
    document.add_paragraph("公文写作必须坚持准确、简洁和规范。")
    second = document.add_paragraph("工作部署")
    document.add_paragraph("各部门应当结合职责落实具体工作。")
    _inject_numpr(first, num_id="94001")
    _inject_numpr(second, num_id="94001")
    source = tmp_path / "native-chinese-heading.docx"
    document.save(source)

    reopened = Document(source)
    features = read_paragraphs(reopened, numbering_index=NumberingIndex(reopened))
    markdown = render(
        GongwenMetadata.default(),
        [item.text for item in features],
        feature_map=dict(enumerate(features)),
        remove_numbering=False,
    )

    headings = [item for item in features if item.heading_level]
    assert [(item.heading_level, item.heading_numbering_text) for item in headings] == [
        (1, "一、 "),
        (1, "二、 "),
    ]
    assert "# 一、 公文写作核心原则" in markdown
    assert "# 二、 工作部署" in markdown


def test_style_inherited_native_numbering_is_recognised(tmp_path) -> None:
    from docx import Document
    from docx.enum.style import WD_STYLE_TYPE

    from docwen_core.docx_parsing.numbering_index import NumberingIndex
    from docwen_plugin_optimizer_gongwen.extraction.paragraph_reader import read_paragraphs

    document = Document()
    style = document.styles.add_style("Gongwen Native Heading", WD_STYLE_TYPE.PARAGRAPH)
    _install_numbering_level(
        document,
        abstract_id="93002",
        num_id="94002",
        num_fmt="chineseCountingThousand",
        level_text="%1、",
        paragraph_style=style.style_id,
    )
    document.add_paragraph("继承样式标题", style=style.name)
    source = tmp_path / "style-native-heading.docx"
    document.save(source)

    reopened = Document(source)
    feature = read_paragraphs(reopened, numbering_index=NumberingIndex(reopened))[0]

    assert (feature.heading_level, feature.heading_numbering_text) == (1, "一、 ")


@pytest.mark.parametrize(
    ("num_fmt", "level_text"),
    [("decimal", "%1."), ("bullet", "•")],
)
def test_ordinary_word_lists_fail_closed_as_body_text(tmp_path, num_fmt: str, level_text: str) -> None:
    from docx import Document

    from docwen_core.docx_parsing.numbering_index import NumberingIndex
    from docwen_plugin_optimizer_gongwen.extraction.paragraph_reader import read_paragraphs

    document = Document()
    _install_numbering_level(
        document,
        abstract_id="93003",
        num_id="94003",
        num_fmt=num_fmt,
        level_text=level_text,
    )
    paragraph = document.add_paragraph("普通列表内容")
    _inject_numpr(paragraph, num_id="94003")
    source = tmp_path / f"ordinary-{num_fmt}.docx"
    document.save(source)

    reopened = Document(source)
    feature = read_paragraphs(reopened, numbering_index=NumberingIndex(reopened))[0]

    assert feature.heading_level == 0
    assert feature.heading_numbering_text == ""
    assert feature.text == "普通列表内容"


def test_ordinary_chinese_numbered_sequence_remains_body_text(tmp_path) -> None:
    from docx import Document

    from docwen_core.docx_parsing.numbering_index import NumberingIndex
    from docwen_plugin_optimizer_gongwen.extraction.paragraph_reader import read_paragraphs

    document = Document()
    _install_numbering_level(
        document,
        abstract_id="93004",
        num_id="94004",
        num_fmt="chineseCounting",
        level_text="%1、",
    )
    for text in ("准备材料", "核对数据", "提交结果"):
        paragraph = document.add_paragraph(text)
        _inject_numpr(paragraph, num_id="94004")
    source = tmp_path / "ordinary-chinese-list.docx"
    document.save(source)

    reopened = Document(source)
    features = read_paragraphs(reopened, numbering_index=NumberingIndex(reopened))

    assert [feature.heading_level for feature in features] == [0, 0, 0]
    assert [feature.heading_numbering_text for feature in features] == ["", "", ""]


@pytest.mark.parametrize("remove_numbering", [False, True])
@pytest.mark.parametrize("num_fmt,level_text", [("decimal", "%1."), ("bullet", "•"), ("chineseCounting", "%1、")])
def test_ordinary_lists_keep_markers_in_gongwen_output(tmp_path, remove_numbering, num_fmt, level_text) -> None:
    from docx import Document

    from docwen_plugin_optimizer_gongwen.pipeline import convert_docx_to_md_gongwen

    document = Document()
    _install_numbering_level(document, abstract_id="93005", num_id="94005", num_fmt=num_fmt, level_text=level_text)
    for text in ("第一项检查", "第二项检查"):
        paragraph = document.add_paragraph(text)
        _inject_numpr(paragraph, num_id="94005")
    source = tmp_path / "ordinary-list.docx"
    document.save(source)

    result = convert_docx_to_md_gongwen(Document(source), str(source), {"remove_numbering": remove_numbering})

    first, second = ("-", "-") if num_fmt == "bullet" else ("1.", "2.")
    assert f"{first} 第一项检查" in result["markdown"]
    assert f"{second} 第二项检查" in result["markdown"]
    assert "# 第一项检查" not in result["markdown"]


def test_list_start_restart_nesting_and_disabled_numbering_survive(tmp_path) -> None:
    from docx import Document

    from docwen_plugin_optimizer_gongwen.pipeline import convert_docx_to_md_gongwen

    document = Document()
    for abstract_id, num_id, start_at, level in (("93006", "94006", 4, 0), ("93007", "94007", 1, 1)):
        _install_numbering_level(
            document,
            abstract_id=abstract_id,
            num_id=num_id,
            num_fmt="decimal",
            level_text=f"%{level + 1}.",
            start_at=start_at,
            list_level=level,
        )
    _inject_numpr(document.add_paragraph("父项"), num_id="94006")
    _inject_numpr(document.add_paragraph("子项\n续行"), num_id="94007", list_level=1)
    _inject_numpr(document.add_paragraph("父项继续"), num_id="94006")
    _inject_numpr(document.add_paragraph("普通段落"), num_id="0")
    _install_numbering_level(document, abstract_id="93008", num_id="94008", num_fmt="decimal", level_text="%1.")
    _inject_numpr(document.add_paragraph("重新开始"), num_id="94008")
    source = tmp_path / "list-boundaries.docx"
    document.save(source)

    result = convert_docx_to_md_gongwen(Document(source), str(source), {})

    assert "4. 父项\n" in result["markdown"]
    assert "    1. 子项\n       续行" in result["markdown"]
    assert "5. 父项继续" in result["markdown"]
    assert "\n\n普通段落\n" in result["markdown"]
    assert "1. 重新开始" in result["markdown"]
