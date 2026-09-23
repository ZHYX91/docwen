"""Copy identifiers need metadata context inside tables, not numeric cells alone."""

from __future__ import annotations

import pytest
from docx import Document

from docwen_plugin_optimizer_gongwen.pipeline import convert_docx_to_md_gongwen

pytestmark = pytest.mark.contract


@pytest.mark.parametrize("with_body", [False, True])
def test_numeric_data_table_does_not_supply_copy_id(with_body: bool) -> None:
    document = Document()
    if with_body:
        document.add_paragraph("WPS 综合测试", style="Heading 1")
        document.add_paragraph("正文内容。")
    table = document.add_table(rows=3, cols=3)
    rows = (("项目", "数量", "备注"), ("材料", "12", "需复核"), ("档案", "3", "正常"))
    for cells, values in zip(table.rows, rows, strict=True):
        for cell, value in zip(cells.cells, values, strict=True):
            cell.text = value

    result = convert_docx_to_md_gongwen(document, "data-table.docx", {})

    assert result["yaml_info"]["份号"] == ""
    assert "| 材料 | 12 | 需复核 |" in result["markdown"]
    assert "| 档案 | 3 | 正常 |" in result["markdown"]


@pytest.mark.parametrize("placement", ["paragraph", "labelled_table", "header_layout"])
def test_legitimate_copy_identifier_keeps_leading_zeroes(placement: str) -> None:
    document = Document()
    if placement == "paragraph":
        document.add_paragraph("000123")
    elif placement == "labelled_table":
        table = document.add_table(rows=1, cols=2)
        table.cell(0, 0).text = "份号："
        table.cell(0, 1).text = "000123"
    else:
        table = document.add_table(rows=2, cols=1)
        table.cell(0, 0).text = "000123"
        table.cell(1, 0).text = "秘密★一年"
    document.add_paragraph("关于开展测试工作的通知", style="Title")
    document.add_paragraph("请按要求办理。")

    result = convert_docx_to_md_gongwen(document, "official-header.docx", {})

    assert result["yaml_info"]["份号"] == "000123"


def test_unlabelled_numeric_cell_after_metadata_is_still_body_data() -> None:
    document = Document()
    document.add_paragraph("秘密★一年")
    document.add_paragraph("关于开展测试工作的通知", style="Title")
    document.add_paragraph("以下为本次统计结果。")
    table = document.add_table(rows=1, cols=1)
    table.cell(0, 0).text = "12"

    result = convert_docx_to_md_gongwen(document, "numeric-body.docx", {})

    assert result["yaml_info"]["份号"] == ""
    assert "| 12 |" in result["markdown"]
