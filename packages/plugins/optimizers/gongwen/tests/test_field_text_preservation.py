"""Imported organizations retain punctuation; signers are never silently discarded."""

from __future__ import annotations

import pytest

from docwen_plugin_optimizer_gongwen.models import ParagraphFeature, RecognitionCandidate, RecognitionResult
from docwen_plugin_optimizer_gongwen.recognition.signals import collect_structured_signals
from docwen_plugin_optimizer_gongwen.recognition.yaml_builder import build_yaml
from docwen_plugin_optimizer_gongwen.utils import extract_signers_from_text, process_copy_to

pytestmark = pytest.mark.contract


def _build(texts: list[str], kinds: list[str]):
    features = [ParagraphFeature(index=index, text=text, source_index=index) for index, text in enumerate(texts)]
    result = RecognitionResult(
        candidates={
            index: RecognitionCandidate(element_type=kind, score=100, para_index=index)
            for index, kind in enumerate(kinds)
        },
        yaml_info={},
        skip_indices=[],
        review_signals=[],
        missing_required=[],
        validation_finding_count=0,
    )
    return features, build_yaml(None, features, result)


@pytest.mark.parametrize(
    "text,expected",
    [
        ("抄送：各县（市、区）人民政府，省直各单位。", "各县（市、区）人民政府，省直各单位"),
        ("报送：甲局、乙局；丙单位。", "甲局、乙局；丙单位"),
        ("分送: 甲（乙,丙）单位，丁单位.", "甲（乙,丙）单位，丁单位"),
        ("甲局、乙局，丙单位", "甲局、乙局，丙单位"),
    ],
)
def test_copy_recipients_keep_the_whole_authored_expression(text, expected) -> None:
    assert process_copy_to(text) == expected
    _features, result = _build([text], ["copy_to"])
    assert result.yaml_info["抄送机关"] == expected


@pytest.mark.parametrize(
    "text,expected",
    [
        ("张三、李四", ["张三", "李四"]),
        ("张三，李四;王五", ["张三", "李四", "王五"]),
        ("张三 李四", ["张三", "李四"]),
        ("张　三", ["张　三"]),
        ("张 三　李四", ["张 三　李四"]),
        ("𠮷田、阿里·艾力", ["𠮷田", "阿里·艾力"]),
        ("John Smith\t王五", ["John Smith", "王五"]),
        ("签发人：张三、张三", ["张三", "张三"]),
        ("单字名：甲", ["单字名：甲"]),
        ("签发人：", []),
    ],
)
def test_signers_preserve_unusual_names_ambiguous_spacing_and_duplicates(text, expected) -> None:
    assert extract_signers_from_text(text) == expected


def test_signer_paragraphs_are_processed_in_source_order_without_deduplication() -> None:
    features, result = _build(["签发人：张三", "李四、张三"], ["signer", "signer_following"])
    result.candidates = dict(reversed(list(result.candidates.items())))
    result = build_yaml(None, features, result)
    assert result.yaml_info["签发人"] == ["张三", "李四", "张三"]
    assert set(result.skip_indices) == {0, 1}


def test_empty_signer_extraction_keeps_the_source_paragraph_and_requires_review() -> None:
    features, result = _build(["签发人："], ["signer"])
    assert result.yaml_info["签发人"] == []
    assert 0 not in result.skip_indices
    assert "signer_text_preserved" in result.candidates[0].trace

    class Scorer:
        def score_round(self, _feature, *, round_group):
            return []

    signals = collect_structured_signals(result, Scorer(), features)
    assert signals["recognition_summary"]["needs_review"]
    assert any(warning["code"] == "GW003" for warning in signals["gongwen_warnings"])


def test_ambiguous_spaced_name_is_retained_in_yaml_and_marked_for_review() -> None:
    _features, result = _build(["签发人：张　三"], ["signer"])
    assert result.yaml_info["签发人"] == ["张　三"]
    assert "signer_text_preserved" in result.candidates[0].trace


def test_recipient_continuation_keeps_mixed_punctuation_and_consumes_both_paragraphs() -> None:
    _features, result = _build(
        ["各县（市、区）人民政府，", "省直各单位："],
        ["body", "recipient"],
    )
    assert result.yaml_info["主送机关"] == "各县（市、区）人民政府，省直各单位"
    assert set(result.skip_indices) == {0, 1}


@pytest.mark.parametrize("earlier", ["请通知各单位，", "会议要求各单位，", "普通正文。", "未明确的前文，"])
def test_recipient_does_not_absorb_ordinary_or_uncertain_previous_text(earlier) -> None:
    _features, result = _build([earlier, "省直各单位："], ["body", "recipient"])
    assert result.yaml_info["主送机关"] == "省直各单位"
    assert 0 not in result.skip_indices


def test_recipient_does_not_cross_a_blank_source_gap_or_table_boundary() -> None:
    features, result = _build(["各市人民政府，", "省直各单位："], ["body", "recipient"])
    features[1].source_index = 3
    result = build_yaml(None, features, result)
    assert result.yaml_info["主送机关"] == "省直各单位"
    assert 0 not in result.skip_indices
    features[1].source_index = 1
    features[0].source = "table"
    result = build_yaml(None, features, result)
    assert result.yaml_info["主送机关"] == "省直各单位"
    assert 0 not in result.skip_indices


@pytest.mark.parametrize(
    "attribute", ["has_image", "has_formula", "has_page_break", "has_section_break", "heading_level", "list_marker"]
)
def test_recipient_does_not_consume_objects_breaks_headings_or_lists(attribute) -> None:
    features, result = _build(["各市人民政府，", "省直各单位："], ["body", "recipient"])
    setattr(features[0], attribute, True)
    result = build_yaml(None, features, result)
    assert result.yaml_info["主送机关"] == "省直各单位"
    assert 0 not in result.skip_indices


@pytest.mark.parametrize(
    "earlier",
    ["承办单位：省财政厅，", "办理联系人为省财政厅，", "工作由省财政厅，", "省财政厅，", "各单位应联系省财政厅，"],
)
def test_recipient_keeps_uncertain_agency_suffix_prose_and_requests_review(earlier: str) -> None:
    from docwen_plugin_optimizer_gongwen.recognition.rounds import run_rounds
    from docwen_plugin_optimizer_gongwen.recognition.scorer import ElementScorer

    texts = ["关于测试的通知", earlier, "各市人民政府：", "正文内容。"]
    features = [
        ParagraphFeature(
            index=i,
            text=text,
            source_index=i,
            font_name="小标宋" if i == 0 else "仿宋",
            font_size_pt=22 if i == 0 else 16,
        )
        for i, text in enumerate(texts)
    ]
    scorer = ElementScorer()
    result = build_yaml(scorer, features, run_rounds(scorer, features))

    assert result.yaml_info["主送机关"] == "各市人民政府"
    assert 1 not in result.skip_indices
    signals = collect_structured_signals(result, scorer, features)
    assert any(warning["code"] == "GW004" for warning in signals["gongwen_warnings"])


@pytest.mark.parametrize("names", ["𠮷田", "John Smith", "𠮷田、John Smith、张三、张三"])
def test_labelled_composite_signers_survive_the_full_docx_pipeline(names: str) -> None:
    from docx import Document
    from docx.shared import Pt

    from docwen_plugin_optimizer_gongwen.pipeline import convert_docx_to_md_gongwen

    document = Document()
    first = document.add_paragraph(f"某发〔2026〕1号\t签发人：{names}")
    first.runs[0].font.name = "仿宋"
    first.runs[0].font.size = Pt(16)
    title = document.add_paragraph("关于测试的通知")
    title.runs[0].font.name = "小标宋"
    title.runs[0].font.size = Pt(22)
    document.add_paragraph("各市人民政府：")
    document.add_paragraph("正文内容。")

    result = convert_docx_to_md_gongwen(document, "signers.docx", options={})

    assert result["success"]
    assert result["yaml_info"]["发文字号"] == "某发〔2026〕1号"
    assert result["yaml_info"]["签发人"] == names.split("、")
    assert "正文内容。" in result["markdown"]


@pytest.mark.parametrize("text", ["某发〔2026〕1号 正文内容", "正文内容 签发人：John Smith"])
def test_unlabelled_or_invalid_document_number_prose_is_not_a_composite_signer(text: str) -> None:
    from docwen_plugin_optimizer_gongwen.recognition.scorer import ElementScorer

    assert not ElementScorer().has_doc_number_and_signer(ParagraphFeature(index=0, text=text))


@pytest.mark.parametrize(
    "earlier",
    ["各县（市、区）人民政府，", "省直各单位，", "各应急管理局，", "各自由贸易试验区办公室，", "各教育工作委员会，"],
)
def test_clear_collective_address_continuations_keep_legitimate_agency_names(earlier: str) -> None:
    _features, result = _build([earlier, "各市人民政府："], ["body", "recipient"])
    assert result.yaml_info["主送机关"] == earlier + "各市人民政府"
    assert 0 in result.skip_indices


def test_uncertain_recipient_prose_remains_in_the_full_docx_markdown_output() -> None:
    from docx import Document
    from docx.shared import Pt

    from docwen_plugin_optimizer_gongwen.pipeline import convert_docx_to_md_gongwen

    document = Document()
    title = document.add_paragraph("关于测试的通知")
    title.runs[0].font.name = "小标宋"
    title.runs[0].font.size = Pt(22)
    document.add_paragraph("承办单位：省财政厅，")
    document.add_paragraph("各市人民政府：")
    document.add_paragraph("正文内容。")

    result = convert_docx_to_md_gongwen(document, "recipient.docx", options={})

    assert result["yaml_info"]["主送机关"] == "各市人民政府"
    assert "承办单位：省财政厅，" in result["markdown"]
    assert "正文内容。" in result["markdown"]
    signals = result["metadata"]["recognition_review_signals"]
    assert any(warning["code"] == "GW004" for warning in signals["gongwen_warnings"])
