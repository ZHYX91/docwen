"""Custom secrecy labels use a complete retention marker in the leading header."""

from __future__ import annotations

from pathlib import Path

import pytest
from docx import Document

from docwen_plugin_optimizer_gongwen.models import ParagraphFeature
from docwen_plugin_optimizer_gongwen.pipeline import convert_docx_to_md_gongwen
from docwen_plugin_optimizer_gongwen.recognition.scorer import ElementScorer

pytestmark = pytest.mark.contract


def _body(markdown: str) -> str:
    return markdown.split("\n---\n", 1)[1] if markdown.startswith("---\n") else markdown


@pytest.mark.parametrize(
    "text",
    ["测试★1年", "演练★五年", "★1年", "Custom★12个月", "示例★30天", "测试★长期", "测试　★　１年", "测试★3月"],
)
@pytest.mark.parametrize("table_header", [False, True])
def test_custom_retention_line_moves_intact_to_metadata(tmp_path: Path, text: str, table_header: bool) -> None:
    document = Document()
    if table_header:
        table = document.add_table(rows=2, cols=1)
        table.cell(0, 0).text = "000123"
        table.cell(1, 0).text = text
    else:
        document.add_paragraph("000123")
        document.add_paragraph(text)
    document.add_paragraph("关于开展测试工作的通知", style="Title")
    document.add_paragraph("正文内容保持完整。")
    source = tmp_path / "retention.docx"
    document.save(str(source))
    reopened = Document(str(source))
    original_xml = reopened.element.xml

    result = convert_docx_to_md_gongwen(reopened, str(source), {})

    assert result["success"]
    assert result["yaml_info"]["密级和保密期限"] == text
    assert result["yaml_info"]["份号"] == "000123"
    body = _body(result["markdown"])
    assert text not in body
    assert "正文内容保持完整。" in body
    assert reopened.element.xml == original_xml


@pytest.mark.parametrize(
    "text",
    [
        "产品★1年保修",
        "评分★",
        "测试☆1年",
        "测试★1年，另有说明",
        "预计★1.5年",
        "测试★1",
        "x★1年+y",
        "正文提到测试★1年。",
    ],
)
def test_incomplete_or_inline_star_text_is_not_metadata(text: str) -> None:
    feature = ParagraphFeature(index=0, text=text)
    scorer = ElementScorer()
    scorer.reset_context([feature])
    assert not scorer.is_security_classification(feature)
    document = Document()
    document.add_paragraph("关于开展测试工作的通知", style="Title")
    document.add_paragraph(text)
    document.add_paragraph("正文内容保持完整。")

    result = convert_docx_to_md_gongwen(document, "ordinary-star.docx", {})

    assert result["yaml_info"]["密级和保密期限"] == ""
    assert text in _body(result["markdown"])


@pytest.mark.parametrize(
    "placement", ["prose", "title", "document_number", "heading", "list", "data_table", "heading_text"]
)
def test_exact_retention_shape_outside_header_stays_in_body(placement: str) -> None:
    document = Document()
    text = "测试★1年"
    if placement == "prose":
        document.add_paragraph("这是正文中的说明。")
    elif placement == "title":
        document.add_paragraph("关于开展测试工作的通知", style="Title")
    elif placement == "document_number":
        document.add_paragraph("甲发〔2026〕1号")
    elif placement == "heading":
        document.add_heading("工作说明", level=1)
    if placement == "list":
        document.add_paragraph(text, style="List Bullet")
    elif placement == "heading_text":
        document.add_heading(text, level=1)
    elif placement == "data_table":
        table = document.add_table(rows=1, cols=2)
        table.cell(0, 0).text = "事项"
        table.cell(0, 1).text = text
    else:
        document.add_paragraph(text)

    result = convert_docx_to_md_gongwen(document, "body-retention.docx", {})

    assert result["yaml_info"]["密级和保密期限"] == ""
    assert text in _body(result["markdown"])
