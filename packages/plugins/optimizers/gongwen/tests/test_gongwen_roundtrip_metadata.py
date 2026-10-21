"""Preserve authored attachment entries and distinguish letterheads from imprints."""

from __future__ import annotations

from pathlib import Path

import pytest
from docx import Document
from docx.shared import Pt

from docwen_plugin_optimizer_gongwen.models import ParagraphFeature
from docwen_plugin_optimizer_gongwen.pipeline import convert_docx_to_md_gongwen
from docwen_plugin_optimizer_gongwen.recognition.scorer import ElementScorer

pytestmark = pytest.mark.contract

_ATTACHMENTS = ["名单甲", "材料乙", "材料乙", "材料丙"] * 4


def _document(
    *, prefix_count: int, table_letterhead: bool, table_imprint: bool, document_number: str = "甲发〔2026〕1号"
):
    document = Document()
    for text in ["001", "测试★1年", "特急", "校核信息甲", "校核信息乙", "校核信息丙"][:prefix_count]:
        document.add_paragraph(text)
    if table_letterhead:
        document.add_table(rows=1, cols=1).cell(0, 0).text = "甲方服务中心"
    else:
        document.add_paragraph("甲方服务中心")
    document.add_paragraph(document_number)
    title = document.add_paragraph("关于测试的通知")
    title.runs[0].font.name = "小标宋"
    title.runs[0].font.size = Pt(22)
    document.add_paragraph("各市人民政府：")
    document.add_heading("工作要求", level=1)
    document.add_paragraph("正文保留重复附件的名称和顺序。")
    for index, name in enumerate(_ATTACHMENTS, 1):
        prefix = "附件：" if index == 1 else "　　"
        document.add_paragraph(f"{prefix}{index}. {name}")
    document.add_paragraph("乙方服务中心")
    document.add_paragraph("2026年10月20日")
    if table_imprint:
        imprint = document.add_table(rows=2, cols=2)
        imprint.cell(0, 0).merge(imprint.cell(0, 1)).text = "抄送：甲局、乙局。"
        imprint.cell(1, 0).text = "文印服务单位"
        imprint.cell(1, 1).text = "2026年10月21日印发"
    else:
        document.add_paragraph("抄送：甲局、乙局。")
        document.add_paragraph("文印服务单位")
        document.add_paragraph("2026年10月21日印发")
    return document


@pytest.mark.parametrize("prefix_count", [0, 3, 6])
@pytest.mark.parametrize("table_letterhead", [False, True])
@pytest.mark.parametrize("table_imprint", [False, True])
def test_actual_docx_preserves_duplicate_attachments_and_consumes_only_the_imprint(
    tmp_path: Path, prefix_count: int, table_letterhead: bool, table_imprint: bool
) -> None:
    document = _document(prefix_count=prefix_count, table_letterhead=table_letterhead, table_imprint=table_imprint)
    source = tmp_path / "official.docx"
    document.save(str(source))
    reopened = Document(str(source))
    original_xml = reopened.element.xml

    result = convert_docx_to_md_gongwen(reopened, str(source), {})

    assert result["success"]
    fields = result["yaml_info"]
    assert fields["附件说明"] == _ATTACHMENTS
    assert fields["密级和保密期限"] == ("测试★1年" if prefix_count else "")
    assert fields["发文机关标志"] == "甲方服务中心"
    assert fields["印发机关"] == "文印服务单位"
    assert fields["印发日期"] == "2026年10月21日"
    assert fields["抄送机关"] == "甲局、乙局"
    body = result["markdown"].split("\n---\n", 1)[1]
    assert "正文保留重复附件的名称和顺序。" in body
    assert "# 工作要求" in body
    assert "抄送：" not in body
    assert "文印服务单位" not in body
    assert "甲方服务中心" not in body
    assert reopened.element.xml == original_xml


def test_mixed_imprint_table_keeps_unrecognized_content() -> None:
    document = _document(prefix_count=3, table_letterhead=True, table_imprint=True)
    extra = document.tables[-1].add_row()
    extra.cells[0].text = "工作事项"
    extra.cells[1].text = "这段表格正文必须保留"

    result = convert_docx_to_md_gongwen(document, "mixed.docx", {})

    assert result["yaml_info"]["印发机关"] == "文印服务单位"
    assert result["yaml_info"]["附件说明"] == _ATTACHMENTS
    assert "| 工作事项 | 这段表格正文必须保留 |" in result["markdown"]
    assert result["attachment_documents"] == []


@pytest.mark.parametrize("table_imprint", [False, True])
def test_combined_printing_line_owns_its_authority(table_imprint: bool) -> None:
    document = _document(prefix_count=3, table_letterhead=True, table_imprint=table_imprint)
    if table_imprint:
        document.tables[-1].cell(1, 0).text = "业务办理说明"
        document.tables[-1].cell(1, 1).text = "实际文印单位　2026年10月21日印发"
    else:
        document.paragraphs[-2].text = "业务办理说明"
        document.paragraphs[-1].text = "实际文印单位　2026年10月21日印发"

    result = convert_docx_to_md_gongwen(document, "combined-imprint.docx", {})

    assert result["yaml_info"]["印发机关"] == "实际文印单位"
    assert result["yaml_info"]["印发日期"] == "2026年10月21日"
    if table_imprint:
        assert "业务办理说明" in result["markdown"]


@pytest.mark.parametrize("document_number", ["001 甲发〔2026〕1号", "甲发〔2026〕1号　签发人：𠮷田、John Smith"])
def test_letterhead_precedes_composite_document_number_beyond_fixed_line_limits(document_number: str) -> None:
    document = _document(prefix_count=6, table_letterhead=True, table_imprint=True, document_number=document_number)

    result = convert_docx_to_md_gongwen(document, "composite.docx", {})

    assert result["yaml_info"]["发文机关标志"] == "甲方服务中心"
    assert result["yaml_info"]["印发机关"] == "文印服务单位"
    assert result["yaml_info"]["发文字号"] == "甲发〔2026〕1号"
    if "签发人" in document_number:
        assert result["yaml_info"]["签发人"] == ["𠮷田", "John Smith"]


@pytest.mark.parametrize("boundary", ["body_to_table", "other_table", "other_row", "heading"])
def test_printing_authority_cannot_cross_unrelated_source_boundaries(boundary: str) -> None:
    first = ParagraphFeature(
        index=0, text="业务服务中心", source="table", table_index=0, table_row_index=0, table_cell_index=0
    )
    date = ParagraphFeature(
        index=1, text="2026年10月21日印发", source="table", table_index=0, table_row_index=0, table_cell_index=1
    )
    if boundary == "body_to_table":
        first.source = "body"
    elif boundary == "other_table":
        date.table_index = 1
    elif boundary == "other_row":
        date.table_row_index = 1
    else:
        first.source = date.source = "body"
        first.heading_level = 1
    scorer = ElementScorer()
    scorer.reset_context([first, date])
    scorer.update_context("printing_date", 1)

    candidates = scorer.score_round(first, round_group="round2")

    assert not any(candidate.element_type == "printing_authority" for candidate in candidates)
