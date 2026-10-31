from __future__ import annotations

import hashlib
import io
from pathlib import Path

import pytest
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt
from PIL import Image

from docwen_application.bundle_mapping import build_bundle_draft
from docwen_core.models import ConversionRequest, FileRef, OutputPolicy, validate_artifact_bundle_draft
from docwen_plugin_optimizer_gongwen.plugin import GongwenOptimizerPlugin
from docwen_runtime.engine.route_resolver import RouteResolver
from docwen_runtime.engine.task_manager import TaskManager
from docwen_runtime.output.finalizer import OutputFinalizer
from docwen_runtime.plugin_registry.registry import PluginRegistry
from docwen_runtime.workspace.manager import WorkspaceManager

pytestmark = pytest.mark.integration


def test_real_attachment_images_follow_their_document_owner(tmp_path: Path):
    def png(color):
        stream = io.BytesIO()
        Image.new("RGB", (8, 8), color).save(stream, format="PNG")
        return stream.getvalue()

    main_bytes, child_bytes, shared_bytes = png("red"), png("blue"), png("green")
    document = Document()
    title = document.add_paragraph("关于印发项目管理办法的通知")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.runs[0].font.name = "小标宋"
    title.runs[0].font.size = Pt(22)
    document.add_paragraph("国办发〔2024〕10号")
    document.add_paragraph("各省人民政府办公厅：")
    body = document.add_paragraph("为规范项目管理工作，现制定本办法。")
    body.add_run().add_picture(io.BytesIO(main_bytes))
    body.add_run().add_picture(io.BytesIO(shared_bytes))
    document.add_paragraph("附件：1. 项目管理办法")
    attachment = document.add_paragraph("附件内容：项目管理办法全文，含办理步骤及图示。")
    attachment.add_run().add_picture(io.BytesIO(child_bytes))
    attachment.add_run().add_picture(io.BytesIO(shared_bytes))
    source = tmp_path / "notice.docx"
    document.save(str(source))
    source_hash = hashlib.sha256(source.read_bytes()).hexdigest()
    registry = PluginRegistry()
    registry.register(GongwenOptimizerPlugin())
    manager = TaskManager(
        registry, RouteResolver(registry), WorkspaceManager(root_dir=str(tmp_path / "workspace")), OutputFinalizer()
    )
    result = manager.execute_single(
        ConversionRequest(
            request_id="real-resource-owners",
            input_refs=[
                FileRef(path=str(source), format="docx", category="document", size_bytes=source.stat().st_size)
            ],
            target_format="md",
            action_name="gongwen",
            options={
                "to_md_keep_images": True,
                "image_mode": "file",
                "image_link_style": "markdown_embed",
                "to_md_enable_ocr": False,
            },
            output_policy=OutputPolicy(output_dir=str(tmp_path / "output")),
        )
    )
    assert result.success, result.error
    main = next(item for item in result.artifacts if item.is_primary)
    child = next(
        item for item in result.artifacts if item.ownership and item.ownership.relation_type == "attachment_of"
    )
    resources = [item for item in result.artifacts if item.ownership and item.ownership.relation_type == "resource_of"]
    assert len(resources) == 4
    main_hashes = {hashlib.sha256(main_bytes).hexdigest(), hashlib.sha256(shared_bytes).hexdigest()}
    child_hashes = {hashlib.sha256(child_bytes).hexdigest(), hashlib.sha256(shared_bytes).hexdigest()}
    for owner, expected_hashes in ((main, main_hashes), (child, child_hashes)):
        owned = []
        for item in resources:
            assert item.ownership is not None
            if item.ownership.owner_artifact_id == owner.artifact_id:
                owned.append(item)
        assert {item.sha256 for item in owned} == expected_hashes
        content = Path(owner.staging_path).read_text(encoding="utf-8")
        for item in owned:
            assert Path(item.staging_path).parent == Path(owner.staging_path).parent
            assert item.suggested_name in content
            assert hashlib.sha256(Path(item.staging_path).read_bytes()).hexdigest() == item.sha256
        foreign_names = {item.suggested_name for item in resources if item.sha256 not in expected_hashes}
        assert all(name not in content for name in foreign_names)
    assert "../" not in Path(child.staging_path).read_text(encoding="utf-8")
    shared = [item for item in resources if item.sha256 == hashlib.sha256(shared_bytes).hexdigest()]
    assert len({item.artifact_id for item in shared}) == 2
    draft = build_bundle_draft(
        profile="document_with_resources", output_media_type="text/markdown", artifacts=result.artifacts
    )
    validate_artifact_bundle_draft(draft)
    assert {relation.source_artifact_id: relation.target_artifact_id for relation in draft.relations} == {
        item.artifact_id: item.ownership.owner_artifact_id for item in result.artifacts if item.ownership
    }
    assert hashlib.sha256(source.read_bytes()).hexdigest() == source_hash
