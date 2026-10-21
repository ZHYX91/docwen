"""Both user-selectable attachment layouts preserve content and template axes."""

from __future__ import annotations

from io import BytesIO
from typing import Literal

import pytest
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml.ns import qn
from docx.shared import Pt
from tests.support.gongwen_template import attachment_template

from docwen_plugin_markdown.template_filler import fill_template
from docwen_plugin_markdown.template_policy import template_attachment_alignment
from docwen_plugin_markdown.template_utils import scan_placeholders

from .test_gongwen_conversion import _conversion_context

pytestmark = pytest.mark.contract


@pytest.mark.parametrize("location", ["body", "cell"])
@pytest.mark.parametrize("count", [1, 10, 101])
@pytest.mark.parametrize("alignment", ["spaces", "tab"])
def test_saved_alignment_modes_preserve_names_fonts_and_first_line(
    location: Literal["body", "cell"], count: int, alignment: Literal["spaces", "tab"]
) -> None:
    document = attachment_template(location=location)
    anchor = document.paragraphs[1] if location == "body" else document.tables[0].cell(0, 0).paragraphs[0]
    anchor.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    anchor.paragraph_format.tab_stops.add_tab_stop(Pt(90), WD_TAB_ALIGNMENT.RIGHT)
    anchor.paragraph_format.tab_stops.add_tab_stop(Pt(140), WD_TAB_ALIGNMENT.RIGHT)
    styles_before = document.styles.element.xml
    names = [f"材料{index}：中文与 English 2026 混排长名称。" * 6 for index in range(1, count + 1)]

    fill_template(document, {"附件说明": names}, [], None, scan_placeholders(document), attachment_alignment=alignment)
    buffer = BytesIO()
    document.save(buffer)
    buffer.seek(0)
    reopened = Document(buffer)
    paragraphs = reopened.paragraphs[1:] if location == "body" else reopened.tables[0].cell(0, 0).paragraphs
    assert len(paragraphs) == count
    assert reopened.styles.element.xml == styles_before
    for paragraph, name in zip(paragraphs, names, strict=True):
        assert paragraph.text.endswith(name)
        assert paragraph.alignment == WD_ALIGN_PARAGRAPH.JUSTIFY
        assert paragraph._p.find(f"{qn('w:pPr')}/{qn('w:numPr')}") is None
        assert paragraph._p.findall(f".//{qn('w:fitText')}") == []
        left = paragraph.paragraph_format.left_indent
        first_line = paragraph.paragraph_format.first_line_indent
        assert left is not None and first_line is not None
        assert left + first_line == Pt(31.6)
        if alignment == "spaces":
            assert "\t" not in paragraph.text
        else:
            assert paragraph.text.count("\t") == 1
            assert paragraph.text.split("\t", 1)[1] == name
            stops = list(paragraph.paragraph_format.tab_stops)
            name_stop = next(stop for stop in stops if stop.alignment == WD_TAB_ALIGNMENT.LEFT)
            assert name_stop.position == paragraph.paragraph_format.left_indent
            assert any(stop.position == Pt(140) and stop.alignment == WD_TAB_ALIGNMENT.RIGHT for stop in stops)
            if count > 1:
                assert any(stop.position == Pt(90) and stop.alignment == WD_TAB_ALIGNMENT.CLEAR for stop in stops)


@pytest.mark.parametrize("route", ["source", "resolved"])
@pytest.mark.parametrize("alignment", [None, "spaces", "tab"])
def test_request_config_reaches_both_production_word_routes(tmp_path, route, alignment) -> None:
    from docwen_plugin_markdown.to_docx.converter import MdToDocxConverter

    template = attachment_template()
    path = tmp_path / "template.docx"
    template.save(path)
    context, _workspace = _conversion_context(tmp_path, route, path, True, alignment=alignment)
    result = MdToDocxConverter().convert(context)
    assert result.success, result.error
    output = Document(result.artifacts[0].staging_path)
    paragraphs = [p for p in output.paragraphs if "材料清单" in p.text or "办理依据" in p.text]
    assert len(paragraphs) == 2
    assert all(("\t" in p.text) == (alignment != "spaces") for p in paragraphs)


@pytest.mark.parametrize("value", ["unknown", "tabs", "", None, True])
def test_invalid_alignment_is_not_silently_changed(value) -> None:
    class Config:
        def get(self, key, default):
            return value

    with pytest.raises(ValueError, match="attachment alignment"):
        template_attachment_alignment(Config())


@pytest.mark.parametrize("inherited", [False, True])
def test_name_tab_preserves_visible_bar_and_clears_only_positioning_stops(inherited: bool) -> None:
    document = attachment_template()
    anchor = document.paragraphs[1]
    owner = anchor.style if inherited else anchor
    assert owner is not None
    owner.paragraph_format.tab_stops.add_tab_stop(Pt(24), WD_TAB_ALIGNMENT.BAR)
    owner.paragraph_format.tab_stops.add_tab_stop(Pt(90), WD_TAB_ALIGNMENT.RIGHT)
    styles_before = document.styles.element.xml
    fill_template(
        document, {"附件说明": ["材料一", "材料二"]}, [], None, scan_placeholders(document), attachment_alignment="tab"
    )
    buffer = BytesIO()
    document.save(buffer)
    buffer.seek(0)
    reopened = Document(buffer)
    assert reopened.styles.element.xml == styles_before
    for paragraph in reopened.paragraphs[1:]:
        stops = list(paragraph.paragraph_format.tab_stops)
        assert [stop.position for stop in stops] == sorted(stop.position for stop in stops)
        assert not any(stop.position == Pt(24) and stop.alignment == WD_TAB_ALIGNMENT.CLEAR for stop in stops)
        if not inherited:
            assert any(stop.position == Pt(24) and stop.alignment == WD_TAB_ALIGNMENT.BAR for stop in stops)
        assert any(stop.position == Pt(90) and stop.alignment == WD_TAB_ALIGNMENT.CLEAR for stop in stops)


def test_name_tab_retains_local_cancellation_of_an_inherited_stop() -> None:
    document = attachment_template()
    anchor = document.paragraphs[1]
    assert anchor.style is not None
    anchor.style.paragraph_format.tab_stops.add_tab_stop(Pt(94), WD_TAB_ALIGNMENT.RIGHT)
    anchor.paragraph_format.tab_stops.add_tab_stop(Pt(94), WD_TAB_ALIGNMENT.CLEAR)
    fill_template(
        document, {"附件说明": ["材料一", "材料二"]}, [], None, scan_placeholders(document), attachment_alignment="tab"
    )
    buffer = BytesIO()
    document.save(buffer)
    buffer.seek(0)
    for paragraph in Document(buffer).paragraphs[1:]:
        assert any(
            stop.position == Pt(94) and stop.alignment == WD_TAB_ALIGNMENT.CLEAR
            for stop in paragraph.paragraph_format.tab_stops
        )


def test_name_tab_does_not_share_the_position_of_an_authored_bar() -> None:
    document = attachment_template()
    anchor = document.paragraphs[1]
    anchor.paragraph_format.tab_stops.add_tab_stop(Pt(95.6), WD_TAB_ALIGNMENT.BAR)
    fill_template(
        document, {"附件说明": ["材料一", "材料二"]}, [], None, scan_placeholders(document), attachment_alignment="tab"
    )
    for paragraph in document.paragraphs[1:]:
        stops = list(paragraph.paragraph_format.tab_stops)
        assert any(stop.position == Pt(95.6) and stop.alignment == WD_TAB_ALIGNMENT.BAR for stop in stops)
        assert any(stop.position.twips == 1913 and stop.alignment == WD_TAB_ALIGNMENT.LEFT for stop in stops)
        left = paragraph.paragraph_format.left_indent
        first_line = paragraph.paragraph_format.first_line_indent
        assert left is not None and first_line is not None
        assert left.twips == 1913 and left + first_line == Pt(31.6)
