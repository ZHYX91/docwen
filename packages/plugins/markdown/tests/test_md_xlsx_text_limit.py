"""Excel text capacity must never result in a successful truncated export."""

from __future__ import annotations

from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

from .conftest import make_context, write_temp_md

pytestmark = pytest.mark.contract


@pytest.mark.parametrize(
    "content",
    ["x" * 32767 + "IMPORTANT-TAIL", "😀" * 16384],
    ids=["ascii-overflow", "utf16-overflow"],
)
def test_markdown_table_cell_over_capacity_fails_with_location(content: str) -> None:
    from docwen_plugin_markdown.to_spreadsheet.converter import MdToXlsxConverter

    source = write_temp_md(f"| Value |\n| --- |\n| {content} |\n")
    context, _workspace = make_context(source, target_format="xlsx")
    result = MdToXlsxConverter().convert(context)

    assert not result.success
    assert result.error is not None
    assert result.error.diagnostic_code == "MD2XLSX-CELL-TEXT-TOO-LONG"
    assert "row 2, column 1" in result.error.message
    assert "32767" in result.error.message
    assert not result.artifacts


@pytest.mark.parametrize("value", ["x" * 32767, "😀" * 16383 + "x"], ids=["ascii-limit", "utf16-limit"])
def test_exactly_maximum_cell_is_preserved(value: str) -> None:
    from docwen_plugin_markdown.to_spreadsheet.converter import MdToXlsxConverter

    source = write_temp_md(f"| Value |\n| --- |\n| {value} |\n")
    context, _workspace = make_context(source, target_format="xlsx")
    result = MdToXlsxConverter().convert(context)
    assert result.success, result.error
    workbook = load_workbook(result.artifacts[0].staging_path)
    try:
        sheet = workbook.active
        assert sheet is not None
        assert sheet["A2"].value == value
    finally:
        workbook.close()


@pytest.mark.parametrize("placeholder", ["{{↓Value}}", "{{→Value}}", "{{Value}}", "prefix {{Value}}"])
def test_template_fill_over_capacity_fails_before_openpyxl_truncates(tmp_path: Path, placeholder: str) -> None:
    from docwen_plugin_markdown.to_spreadsheet.converter import MdToXlsxConverter

    template = tmp_path / "template.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet["A1"] = placeholder
    workbook.save(template)
    workbook.close()

    value = "x" * (32766 if placeholder.startswith("prefix") else 32768)
    content = f"| Value |\n| --- |\n| {value} |\n"
    if "↓" not in placeholder and "→" not in placeholder:
        content = f'---\nValue: "{value}"\n---\n\n| Other |\n| --- |\n| row |\n'
    source = write_temp_md(content)
    context, _workspace = make_context(source, target_format="xlsx", options={"template_name": str(template)})
    result = MdToXlsxConverter().convert(context)
    assert not result.success
    assert result.error is not None
    assert result.error.diagnostic_code == "MD2XLSX-CELL-TEXT-TOO-LONG"
    assert result.diagnostics[0].location == "Sheet!R1C1"
    assert not result.artifacts


def test_oversized_header_is_rejected_before_workbook_publication() -> None:
    from docwen_plugin_markdown.to_spreadsheet.converter import MdToXlsxConverter

    source = write_temp_md(f"| {'h' * 32768} |\n| --- |\n| value |\n")
    context, workspace = make_context(source, target_format="xlsx")
    result = MdToXlsxConverter().convert(context)
    assert not result.success
    assert result.error is not None
    assert result.error.diagnostic_code == "MD2XLSX-CELL-TEXT-TOO-LONG"
    assert result.diagnostics[0].location == "Sheet1!R1C1"
    assert not result.artifacts
    assert not list(Path(workspace.staging_dir).glob("*.xlsx"))


def test_plain_csv_keeps_text_beyond_excel_cell_capacity() -> None:
    import csv

    from docwen_plugin_markdown.to_spreadsheet.converter import MdToCsvConverter

    value = "x" * 32768 + "UNIQUE-TAIL"
    source = write_temp_md(f"| Value |\n| --- |\n| {value} |\n")
    context, _workspace = make_context(source, target_format="csv")
    result = MdToCsvConverter().convert(context)
    assert result.success, result.error
    with Path(result.artifacts[0].staging_path).open(encoding="utf-8-sig", newline="") as handle:
        assert list(csv.reader(handle)) == [["Value"], [value]]


def test_sparse_template_fills_without_materializing_its_empty_rectangle(tmp_path: Path, monkeypatch) -> None:
    from openpyxl.worksheet.worksheet import Worksheet

    from docwen_plugin_markdown.to_spreadsheet.converter import MdToXlsxConverter

    template = tmp_path / "sparse.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet["XFD1048576"] = "far sentinel"
    sheet["Z101"] = "=1+2"
    sheet["D3"] = "{{→Other}}"
    sheet["A3"] = "{{↓Value}}"
    sheet["A1"] = "{{title}}"
    sheet["B2"] = "{{IMAGE:missing.png}}"
    workbook.save(template)
    workbook.close()

    original_iter_rows = Worksheet.iter_rows

    def bounded_rows(worksheet, *args, **kwargs):
        if worksheet.max_row * worksheet.max_column > 1_000_000:
            raise AssertionError("Sparse template scans must not expand the entire worksheet rectangle")
        return original_iter_rows(worksheet, *args, **kwargs)

    monkeypatch.setattr(Worksheet, "iter_rows", bounded_rows)
    source = write_temp_md(
        "---\ntitle: Workbook\n---\n\n| Value | Other |\n| --- | --- |\n| first | one |\n| second | two |\n"
    )
    context, _workspace = make_context(source, target_format="xlsx", options={"template_name": str(template)})
    result = MdToXlsxConverter().convert(context)
    assert result.success, result.error
    output = load_workbook(result.artifacts[0].staging_path)
    try:
        sheet = output.active
        assert sheet is not None
        assert sheet["A1"].value == "Workbook"
        assert [sheet["A3"].value, sheet["A4"].value] == ["first", "second"]
        assert [sheet["D3"].value, sheet["E3"].value] == ["one", "two"]
        assert sheet["B2"].value is None
        assert sheet["Z101"].value == "=1+2" and sheet["Z101"].data_type == "f"
        assert sheet["XFD1048576"].value == "far sentinel"
        assert len(sheet._cells) <= 8
    finally:
        output.close()
