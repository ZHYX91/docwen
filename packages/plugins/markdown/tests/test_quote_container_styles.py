"""Effective template borders win over optional quote child accents."""

from __future__ import annotations

from typing import Any

import pytest
from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from lxml import etree

from docwen_plugin_markdown.renderer_utils import apply_quote_container_accent

pytestmark = pytest.mark.contract


def _left(properties: Any, value: str) -> Any:
    borders = properties.find(qn("w:pBdr"))
    if borders is None:
        borders = OxmlElement("w:pBdr")
        properties.insert(0, borders)
    left = OxmlElement("w:left")
    for name, text in {"val": value, "sz": "10", "color": "ABCDEF", "space": "7", "shadow": "1"}.items():
        left.set(qn(f"w:{name}"), text)
    borders.append(left)
    return left


@pytest.mark.parametrize("child_source", ["direct", "style", "basedOn", "defaults"])
@pytest.mark.parametrize("value", ["none", "nil", "double"])
def test_quote_accent_preserves_every_effective_child_left_border(child_source: str, value: str) -> None:
    document: Any = Document()
    quote = document.styles.add_style("QuoteOwner", WD_STYLE_TYPE.PARAGRAPH)
    _left(quote._element.get_or_add_pPr(), "single")
    child = document.styles.add_style("ChildHeading", WD_STYLE_TYPE.PARAGRAPH)
    base = document.styles.add_style("ChildBase", WD_STYLE_TYPE.PARAGRAPH)
    child.base_style = base
    paragraph = document.add_paragraph("Current child", style=child)
    if child_source == "direct":
        properties = paragraph._p.get_or_add_pPr()
    elif child_source in {"style", "basedOn"}:
        properties = (child if child_source == "style" else base)._element.get_or_add_pPr()
    else:
        defaults = document.styles._element.find(qn("w:docDefaults"))
        properties = defaults.find(f"{qn('w:pPrDefault')}/{qn('w:pPr')}")
    _left(properties, value)
    before_styles = etree.tostring(document.styles._element)
    before_paragraph = etree.tostring(paragraph._p)
    apply_quote_container_accent(paragraph, style=quote)
    assert etree.tostring(paragraph._p) == before_paragraph
    assert etree.tostring(document.styles._element) == before_styles


def test_quote_accent_copies_complete_inherited_source_and_valid_element_order() -> None:
    document: Any = Document()
    base = document.styles.add_style("QuoteBase", WD_STYLE_TYPE.PARAGRAPH)
    source = _left(base._element.get_or_add_pPr(), "double")
    quote = document.styles.add_style("QuoteOwner", WD_STYLE_TYPE.PARAGRAPH)
    quote.base_style = base
    paragraph = document.add_paragraph("Heading child", style="Heading 2")
    properties = paragraph._p.get_or_add_pPr()
    properties.append(OxmlElement("w:contextualSpacing"))
    properties.append(OxmlElement("w:jc"))
    apply_quote_container_accent(paragraph, style=quote)
    assert [node.tag for node in properties] == [qn("w:pStyle"), qn("w:pBdr"), qn("w:contextualSpacing"), qn("w:jc")]
    assert properties.find(f"{qn('w:pBdr')}/{qn('w:left')}").attrib == source.attrib
    assert paragraph.style.style_id == "Heading2"


def test_left_accent_precedes_existing_bottom_right_between_edges() -> None:
    document: Any = Document()
    quote = document.styles.add_style("QuoteOwner", WD_STYLE_TYPE.PARAGRAPH)
    _left(quote._element.get_or_add_pPr(), "double")
    paragraph = document.add_paragraph("Child")
    borders = OxmlElement("w:pBdr")
    for name in ("top", "bottom", "right", "between", "bar"):
        edge = OxmlElement(f"w:{name}")
        edge.set(qn("w:val"), "none")
        borders.append(edge)
    paragraph._p.get_or_add_pPr().append(borders)
    apply_quote_container_accent(paragraph, style=quote)
    assert [etree.QName(node).localname for node in borders] == ["top", "left", "bottom", "right", "between", "bar"]


def test_list_continuation_uses_template_text_column_without_a_marker_or_hanging_indent() -> None:
    from docwen_core.docx_parsing.list_processing import detect_list_item
    from docwen_core.docx_parsing.numbering_index import NumberingIndex
    from docwen_plugin_markdown.mistune_extensions import parse_markdown_text
    from docwen_plugin_markdown.renderer import MdToDocxRenderer

    document: Any = Document()
    style = document.styles.add_style("TemplateList", WD_STYLE_TYPE.PARAGRAPH)
    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), "77777")
    level = OxmlElement("w:lvl")
    level.set(qn("w:ilvl"), "0")
    for name, value in (("start", "1"), ("numFmt", "bullet"), ("pStyle", style.style_id), ("lvlText", "•")):
        node = OxmlElement(f"w:{name}")
        node.set(qn("w:val"), value)
        level.append(node)
    properties = OxmlElement("w:pPr")
    indent = OxmlElement("w:ind")
    indent.set(qn("w:left"), "840")
    indent.set(qn("w:hanging"), "420")
    properties.append(indent)
    level.append(properties)
    abstract.append(level)
    root = document.part.numbering_part.element
    root.insert(0, abstract)
    number = OxmlElement("w:num")
    number.set(qn("w:numId"), "77777")
    reference = OxmlElement("w:abstractNumId")
    reference.set(qn("w:val"), "77777")
    number.append(reference)
    root.append(number)
    renderer = MdToDocxRenderer(document)
    renderer._list_block_style = style
    renderer.render(parse_markdown_text("- First.\n\n  Continuation one.\n  Continuation two.\n"))
    continuation = document.paragraphs[1]
    index = NumberingIndex(document)
    assert index.lookup_by_style_id(style.style_id) is not None
    assert detect_list_item(continuation, index) == (None, None, None)
    properties = continuation._p.find(qn("w:pPr"))
    indent = properties.find(qn("w:ind"))
    assert indent.get(qn("w:left")) == "840" and indent.get(qn("w:firstLine")) == "0"
    assert indent.get(qn("w:hanging")) is None
    assert len(properties.findall(qn("w:ind"))) == 1
