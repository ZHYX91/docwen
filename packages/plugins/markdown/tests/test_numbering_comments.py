"""Real DOCX packages prove ownership of added numbers and comment preservation."""

from __future__ import annotations

from io import BytesIO
from pathlib import Path
from zipfile import ZipFile

import pytest
from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from lxml import etree
from PIL import Image
from tests.support.numbering import repository_numbering_registry, repository_numbering_snapshot

from docwen_plugin_markdown.to_docx.converter import MdToDocxConverter
from docwen_plugin_markdown.to_docx.numbering_comments import (
    NUMBERING_REVIEW_TEXT,
    NumberingProvenance,
    add_numbering_review,
)
from docwen_runtime.numbering import NumberingSchemeRegistry

from .conftest import PROJECT_ROOT, make_context

pytestmark = pytest.mark.contract


def _convert(
    tmp_path,
    *,
    enabled: bool | None = True,
    add=True,
    mode="text",
    fields=False,
    scheme="gongwen_standard",
    source=None,
    conflicting_styles=False,
    registry=None,
    text_enabled=None,
    template_native=False,
    captions=False,
    auto_link=False,
):
    template = Document()
    template.styles.add_style("CommentReference", WD_STYLE_TYPE.CHARACTER)
    template.styles.add_style("CommentText", WD_STYLE_TYPE.PARAGRAPH)
    original = template.add_paragraph("Original comment anchor")
    comment = template.add_comment(original.runs[0], "Original review", author="Original author")
    image = BytesIO()
    Image.new("RGB", (2, 2), "red").save(image, format="PNG")
    image.seek(0)
    comment.paragraphs[0].add_run().add_picture(image)
    if conflicting_styles:
        # The old comment uses valid private styles. The canonical names are
        # occupied by valid user styles of the opposite type.
        for name in ("CommentReference", "CommentText"):
            style = template.styles[name]
            style.name = "User" + name
            assert style._element is not None
            style._element.set(qn("w:styleId"), "User" + name)
        template.styles.add_style("CommentReference", WD_STYLE_TYPE.PARAGRAPH)
        template.styles.add_style("CommentText", WD_STYLE_TYPE.CHARACTER)
        for root in (template.element, comment._element):
            for element in root.iter():
                if element.tag in {qn("w:rStyle"), qn("w:pStyle")}:
                    value = element.get(qn("w:val"))
                    if value in {"CommentReference", "CommentText"}:
                        element.set(qn("w:val"), "User" + value)
    # IDs without definitions must still be reserved in every template part.
    for paragraph, comment_id in (
        (original, 1),
        (template.sections[0].header.paragraphs[0], 2),
        (template.sections[0].footer.paragraphs[0], 3),
    ):
        reference = OxmlElement("w:commentReference")
        reference.set(qn("w:id"), str(comment_id))
        paragraph.add_run()._r.append(reference)
    if template_native:
        original.style = "Heading 1"
        style = template.styles["Heading 1"]._element
        assert style is not None
        style.get_or_add_pPr().get_or_add_numPr().get_or_add_numId().val = 1
    template.add_paragraph("{{正文}}")
    template_path = tmp_path / "template.docx"
    template.save(template_path)
    source = source if source is not None else "# 一、手写标题\n\n## **子标题**：\n合并正文[^a]。\n\n[^a]: 脚注正文\n"
    source_path = tmp_path / "source.md"
    source_path.write_text(source, encoding="utf-8")
    context, _workspace = make_context(
        str(source_path),
        options={
            "template_name": str(template_path),
            "add_numbering": add,
            "remove_numbering": False,
            "numbering_scheme": scheme,
            "heading_numbering_render_mode": mode,
            "heading_formatting_mode": "apply",
            "markdown_extensions": {"input": {"typed_endnotes": True, "captions_references": captions}},
        },
        config_values={
            "text": {"heading_numbering_render_mode": mode},
            "link": {"non_embed_links": {"auto_link_bare_url": auto_link}},
            "template_fill": {
                **({"gongwen_numbering_comments": enabled} if enabled is not None else {}),
                **({"gongwen_numbering_text": text_enabled} if text_enabled is not None else {}),
                "gongwen_fields_enabled": fields,
            },
        },
        numbering_registry=registry or repository_numbering_registry(),
    )
    result = MdToDocxConverter().convert(context)
    assert result.success, result.error
    assert context.config.get("text.heading_numbering_render_mode") == mode
    assert source_path.read_text(encoding="utf-8") == source
    output_path = Path(result.artifacts[0].staging_path)
    return result, template_path, output_path


def _anchors(document):
    anchored = {}
    for paragraph in document.element.iter(qn("w:p")):
        active = {}
        for element in paragraph.iter():
            if element.tag == qn("w:commentRangeStart"):
                active[element.get(qn("w:id"))] = []
            elif element.tag == qn("w:t"):
                for texts in active.values():
                    texts.append(element.text or "")
            elif element.tag == qn("w:commentRangeEnd"):
                key = element.get(qn("w:id"))
                if key in active:
                    anchored[key] = "".join(active.pop(key))
    return anchored


def _paragraph_num_id(paragraph):
    properties = paragraph._p.find(qn("w:pPr"))
    numbering = properties.find(qn("w:numPr")) if properties is not None else None
    number_id = numbering.find(qn("w:numId")) if numbering is not None else None
    return number_id.get(qn("w:val")) if number_id is not None else None


def _style_bytes(document, name):
    return etree.tostring(document.styles[name]._element)


@pytest.mark.parametrize("fields", [False, True])
@pytest.mark.parametrize(
    "scheme,expected", [("gongwen_standard", ["一、", "（一）"]), ("hierarchical_standard", ["1", "1.1"])]
)
def test_added_numbers_only_have_comments_and_existing_resources_survive(tmp_path, fields, scheme, expected):
    _result, original_path, output_path = _convert(tmp_path, fields=fields, scheme=scheme)
    document = Document(str(output_path))
    comments = list(document.comments)
    assert comments[0].text == "Original review"
    assert comments[0].author == "Original author"
    reviews = [comment for comment in comments if comment.text == NUMBERING_REVIEW_TEXT]
    assert [comment.comment_id for comment in reviews] == [4, 5]
    anchors = _anchors(document)
    assert [anchors[str(comment.comment_id)] for comment in reviews] == expected
    assert anchors["0"] == "Original comment anchor"
    assert any("一、手写标题" in p.text for p in document.paragraphs)
    assert any("合并正文" in p.text for p in document.paragraphs)
    with ZipFile(original_path) as original, ZipFile(output_path) as output:
        for name in original.namelist():
            if name.startswith("word/media/") or name == "word/_rels/comments.xml.rels":
                assert output.read(name) == original.read(name)
        assert "脚注正文" in output.read("word/footnotes.xml").decode("utf-8")
        for name, comment_id in (("word/header1.xml", "2"), ("word/footer1.xml", "3")):
            assert f'w:id="{comment_id}"' in output.read(name).decode("utf-8")


@pytest.mark.parametrize("enabled,add", [(False, True), (True, False)])
def test_switch_off_or_numbering_off_never_comments_existing_numbers(tmp_path, enabled, add):
    _result, _original, output_path = _convert(tmp_path, enabled=enabled, add=add)
    document = Document(str(output_path))
    assert [comment.text for comment in document.comments] == ["Original review"]


@pytest.mark.parametrize("text_enabled", [None, True, False])
@pytest.mark.parametrize("comments_enabled", [True, False])
def test_text_policy_overrides_only_new_native_numbers_and_enforces_comment_dependency(
    tmp_path, text_enabled, comments_enabled
):
    result, original_path, output_path = _convert(
        tmp_path,
        mode="word_native",
        source="# Title\n",
        text_enabled=text_enabled,
        enabled=comments_enabled,
        template_native=True,
    )
    document = Document(str(output_path))
    reviews = [comment for comment in document.comments if comment.text == NUMBERING_REVIEW_TEXT]
    assert not any(item.code == "MD2DOCX-NUMBERING-COMMENTS-NATIVE" for item in result.diagnostics)
    if text_enabled is not False:
        assert len(reviews) == int(comments_enabled)
        if comments_enabled:
            assert _anchors(document)[str(reviews[0].comment_id)] == "一、"
        heading = next(p for p in document.paragraphs if p.text == "一、Title")
        assert _paragraph_num_id(heading) == "0"
        original = Document(str(original_path))
        assert _style_bytes(document, "Heading 1") == _style_bytes(original, "Heading 1")
        with ZipFile(original_path) as before, ZipFile(output_path) as after:
            assert before.read("word/numbering.xml") == after.read("word/numbering.xml")
    else:
        assert not reviews  # external inconsistent comments=true is ignored
        assert any(p.text == "Title" for p in document.paragraphs)
        with ZipFile(output_path) as package:
            assert b"20000" in package.read("word/numbering.xml")


def test_missing_old_comment_key_now_defaults_on_without_enabling_numbering(tmp_path):
    _result, _original, output_path = _convert(tmp_path, enabled=None)
    assert len([c for c in Document(str(output_path)).comments if c.text == NUMBERING_REVIEW_TEXT]) == 2


@pytest.mark.parametrize(
    "source,expected,captions",
    [
        ("# Main\n\n![[sample.md]]\n", ["一、"], False),
        ("![[sample.md]]\n\n# Main\n\n![[same.md]]\n\n## Tail\n", ["一、", "（一）"], False),
        (
            "> # Quoted\n>\n> ![[sample.md]]\n\n- ## List\n\n  ![[same.md]]\n\nSetext\n======\n\n#\n",
            ["一、", "（一）", "二、", "三、"],
            False,
        ),
        ("![[sample.md]]\n\n# [Main](https://example.com)\n\n## [[Missing|Tail]]\n", ["一、", "（一）"], False),
        ("# Main ^main\n\n![[sample.md]]\n\nSee @[[#^main]].\n", ["一、"], True),
    ],
    ids=[
        "embed-after",
        "embed-before-between-identical",
        "containers-setext-empty",
        "changed-links",
        "semantic-carrier",
    ],
)
def test_embedded_headings_cannot_consume_original_numbering_provenance(tmp_path, source, expected, captions):
    fixture = PROJECT_ROOT / "tests/fixtures/files/sample.md"
    fixture_bytes = fixture.read_bytes()
    (tmp_path / "sample.md").write_bytes(fixture_bytes)
    (tmp_path / "same.md").write_text("# 一、Main\n", encoding="utf-8")
    result, original_path, output_path = _convert(
        tmp_path, source=source, mode="word_native", template_native=True, captions=captions
    )
    document = Document(str(output_path))
    reviews = [comment for comment in document.comments if comment.text == NUMBERING_REVIEW_TEXT]
    anchors = _anchors(document)
    assert [anchors[str(comment.comment_id)] for comment in reviews] == expected
    review_ids = {str(comment.comment_id) for comment in reviews}
    assert any(paragraph.text == "Sample" for paragraph in document.paragraphs)
    for paragraph in document.paragraphs:
        style = paragraph.style
        if style is None or not (style.name or "").lower().startswith("heading"):
            continue
        paragraph_reviews = {node.get(qn("w:id")) for node in paragraph._p.iter(qn("w:commentRangeStart"))}
        assert _paragraph_num_id(paragraph) == ("0" if paragraph_reviews.intersection(review_ids) else None)
    assert _style_bytes(document, "Heading 1") == _style_bytes(Document(str(original_path)), "Heading 1")
    assert fixture.read_bytes() == fixture_bytes
    with ZipFile(original_path) as before, ZipFile(output_path) as after:
        original_definitions = etree.fromstring(before.read("word/numbering.xml"))
        final_definitions = etree.fromstring(after.read("word/numbering.xml"))
        for definition in original_definitions:
            assert any(etree.tostring(definition) == etree.tostring(node) for node in final_definitions)
        for name in after.namelist():
            if name.endswith(".xml"):
                assert b"DWNUM" not in after.read(name), name
    assert all("DWNUM" not in diagnostic.message for diagnostic in result.diagnostics)


def test_provenance_scope_collision_retains_authored_marker_like_text(monkeypatch):
    from docwen_plugin_markdown.to_docx import numbering_comments

    tokens = iter(["a" * 48, "b" * 48])
    monkeypatch.setattr(numbering_comments.secrets, "token_hex", lambda _size: next(tokens))
    literal = "DWNUM" + "a" * 48 + "X0Z"
    provenance = NumberingProvenance(f"# {literal}\n")
    marker = provenance.mark("一、")
    clean, prefixes, missing = provenance.prepare_source(f"# {marker}一、{literal}\n")
    assert clean == f"# 一、{literal}\n"
    assert prefixes == ["一、"] and missing == 0
    assert provenance.strip_markers(marker + literal) == literal


def test_preparse_errors_cannot_expose_numbering_provenance_markers(tmp_path, monkeypatch):
    from docwen_plugin_markdown.to_docx import converter

    source = tmp_path / "failure.md"
    source.write_text("# Main\n", encoding="utf-8")

    def fail_with_source(content, *_args, **_kwargs):
        assert "DWNUM" in content
        raise ValueError(content)

    monkeypatch.setattr(converter, "process_markdown_links", fail_with_source)
    context, _workspace = make_context(
        str(source),
        options={"add_numbering": True, "numbering_scheme": "gongwen_standard"},
        numbering_registry=repository_numbering_registry(),
    )
    result = MdToDocxConverter().convert(context)
    assert not result.success and result.error is not None
    assert "一、Main" in result.error.message
    assert "DWNUM" not in result.error.message
    assert all("DWNUM" not in diagnostic.message for diagnostic in result.diagnostics)


@pytest.mark.parametrize(
    "format_string",
    [
        "{1.chinese_lower}、",
        " {1.arabic_half} ",
        "\t{1.arabic_half} ",
        "https://example.com/{1.arabic_half} ",
        "*{1.arabic_half}* ",
        "`{1.arabic_half}` ",
        "[{1.arabic_half}](https://example.com) ",
        "<b>{1.arabic_half}</b> ",
    ],
    ids=["standard", "space", "tab", "url", "emphasis", "code", "link", "html"],
)
def test_provenance_is_transformation_equivalent_for_custom_prefixes(format_string):
    from docwen_core.export_semantics import LinkRuntimeConfig
    from docwen_core.links import process_markdown_links
    from docwen_core.text.heading_numbering import HeadingFormatter
    from docwen_plugin_markdown.common_utils import _apply_numbering

    source = "# Main\n"
    scheme = {"level_1": {"format": format_string}}
    provenance = NumberingProvenance(source)
    marked = _apply_numbering(source, HeadingFormatter(scheme), prefix_marker=provenance.mark)
    ordinary = _apply_numbering(source, HeadingFormatter(scheme))
    config = LinkRuntimeConfig(auto_link_bare_url=True)
    control = process_markdown_links(ordinary, "", link_config=config, target_format="docx")
    transformed = process_markdown_links(marked, "", link_config=config, target_format="docx")
    clean, prefixes, missing = provenance.prepare_source(transformed)
    assert clean == control
    assert len(prefixes) == 1 and prefixes[0]
    assert missing == 0


@pytest.mark.parametrize(
    "format_string,visible,review",
    [
        (" {1.arabic_half} ", "1 Main", True),
        ("\t{1.arabic_half} ", "1 Main", True),
        ("https://example.com/{1.arabic_half} ", "https://example.com/1 Main", False),
    ],
    ids=["space", "tab", "url"],
)
def test_custom_prefix_boundaries_preserve_text_numbering_and_links(tmp_path, format_string, visible, review):
    snapshot = repository_numbering_snapshot()
    snapshot["numbering"]["add"]["schemes"]["custom"] = {
        "name": "Custom",
        "enabled": True,
        "locales": ["*"],
        "level_1": {"format": format_string},
    }
    registry = NumberingSchemeRegistry.from_config_snapshot(
        snapshot, locale_path=PROJECT_ROOT / "i18n/locales/zh_CN.toml"
    )
    result, _original, output_path = _convert(
        tmp_path,
        scheme="custom",
        registry=registry,
        source="# Main\n",
        template_native=True,
        auto_link=True,
    )
    document = Document(str(output_path))
    heading = next(paragraph for paragraph in document.paragraphs if paragraph.text == visible)
    assert _paragraph_num_id(heading) == "0"
    reviews = [comment for comment in document.comments if comment.text == NUMBERING_REVIEW_TEXT]
    assert len(reviews) == int(review)
    if review:
        assert _anchors(document)[str(reviews[0].comment_id)] == "1"
    else:
        assert any(rel.target_ref == "https://example.com/1" for rel in document.part.rels.values())
    assert any(item.code == "MD2DOCX-NUMBERING-COMMENTS-UNANCHORED" for item in result.diagnostics) is not review


@pytest.mark.parametrize("label,part", [("a", "footnotes"), ("endnote:a", "endnotes")])
@pytest.mark.parametrize(
    "note_heading", ["    # Note title\n", "    Note title\n    ==========\n"], ids=["atx", "setext"]
)
def test_deferred_note_headings_do_not_consume_body_numbering_provenance(tmp_path, label, part, note_heading):
    source = f"# Main\n\nref[^{label}].\n\n[^{label}]:\n{note_heading}\n    Note body.\n"
    result, _original, output_path = _convert(tmp_path, source=source, template_native=True)
    document = Document(str(output_path))
    heading = next(paragraph for paragraph in document.paragraphs if paragraph.text == "一、Main")
    assert _paragraph_num_id(heading) == "0"
    reviews = [comment for comment in document.comments if comment.text == NUMBERING_REVIEW_TEXT]
    assert len(reviews) == 1 and _anchors(document)[str(reviews[0].comment_id)] == "一、"
    assert not any(item.code == "MD2DOCX-NUMBERING-COMMENTS-UNANCHORED" for item in result.diagnostics)
    with ZipFile(output_path) as package:
        note = etree.fromstring(package.read(f"word/{part}.xml"))
        assert "Note title" in "".join(note.itertext())
        assert not list(note.iter(qn("w:commentRangeStart")))
        assert not list(note.iter(qn("w:numPr")))
        assert b"DWNUM" not in package.read(f"word/{part}.xml")


def test_nested_setext_empty_headings_and_code_do_not_shift_comment_ownership(tmp_path):
    source = "#\n\nTitle\n=====\n\n> ## Quote\n\n- ### List\n\n```md\n# Literal\n```\n"
    _result, _original, output_path = _convert(tmp_path, source=source)
    document = Document(str(output_path))
    reviews = [comment for comment in document.comments if comment.text == NUMBERING_REVIEW_TEXT]
    anchors = _anchors(document)
    assert [anchors[str(comment.comment_id)] for comment in reviews] == ["一、", "二、", "（一）", "1."]
    assert "Literal" not in "".join(anchors[str(comment.comment_id)] for comment in reviews)


def test_new_comment_styles_do_not_overwrite_conflicting_template_types(tmp_path):
    _result, template_path, output_path = _convert(tmp_path, conflicting_styles=True)
    output = Document(str(output_path))
    original = Document(str(template_path))
    assert output.styles["CommentReference"].type == WD_STYLE_TYPE.PARAGRAPH
    assert output.styles["CommentText"].type == WD_STYLE_TYPE.CHARACTER
    for name in ("CommentReference", "CommentText"):
        output_style = output.styles[name]._element
        original_style = original.styles[name]._element
        assert output_style is not None and original_style is not None
        assert etree.tostring(output_style) == etree.tostring(original_style)
    reviews = [comment for comment in output.comments if comment.text == NUMBERING_REVIEW_TEXT]
    style_types = {style.style_id: style.type for style in output.styles}
    for comment in reviews:
        for element in comment._element.iter(qn("w:pStyle")):
            assert style_types[element.get(qn("w:val"), "")] == WD_STYLE_TYPE.PARAGRAPH
        for element in comment._element.iter(qn("w:rStyle")):
            assert style_types[element.get(qn("w:val"), "")] == WD_STYLE_TYPE.CHARACTER
    for reference in output.element.iter(qn("w:commentReference")):
        if int(reference.get(qn("w:id"), "-1")) >= 4:
            run = reference.getparent()
            assert run is not None
            properties = run.find(qn("w:rPr"))
            assert properties is not None
            style = properties.find(qn("w:rStyle"))
            assert style is not None
            assert style_types[style.get(qn("w:val"), "")] == WD_STYLE_TYPE.CHARACTER


def test_custom_markdown_number_preserves_content_and_reports_unanchored_comments(tmp_path):
    snapshot = repository_numbering_snapshot()
    snapshot["numbering"]["add"]["schemes"]["custom"] = {
        "name": "Custom",
        "enabled": True,
        "locales": ["*"],
        "level_1": {"format": "**{1.arabic_half}** "},
    }
    registry = NumberingSchemeRegistry.from_config_snapshot(
        snapshot, locale_path=PROJECT_ROOT / "i18n/locales/zh_CN.toml"
    )
    result, _original, output_path = _convert(tmp_path, scheme="custom", registry=registry, source="# Title\n")
    output = Document(str(output_path))
    assert [comment.text for comment in output.comments] == ["Original review"]
    heading = next(p for p in output.paragraphs if p.text == "1 Title")
    assert not any(node.tag == qn("w:commentRangeStart") for node in heading._p)
    assert any(item.code == "MD2DOCX-NUMBERING-COMMENTS-UNANCHORED" for item in result.diagnostics)
    assert heading.runs[0].bold is True


def test_review_comments_coexist_with_repeated_footnotes_and_endnotes(tmp_path):
    source = (
        "# Title\n\nBody[^a] again[^a], end[^endnote:e] again[^endnote:e].\n\n[^a]: Footnote\n[^endnote:e]: Endnote\n"
    )
    _result, _original, output_path = _convert(tmp_path, source=source)
    output = Document(str(output_path))
    reviews = [comment for comment in output.comments if comment.text == NUMBERING_REVIEW_TEXT]
    assert len(reviews) == 1
    assert _anchors(output)[str(reviews[0].comment_id)] == "一、"
    with ZipFile(output_path) as package:
        assert "Footnote" in package.read("word/footnotes.xml").decode("utf-8")
        assert "Endnote" in package.read("word/endnotes.xml").decode("utf-8")
        body = package.read("word/document.xml").decode("utf-8")
        assert "NOTEREF" in body


def test_allocator_reserves_max_id_and_ambiguous_style_ids_without_overwriting(tmp_path):
    from copy import deepcopy

    document = Document()
    first = document.styles.add_style("CommentReference", WD_STYLE_TYPE.CHARACTER)
    assert first._element is not None
    document.styles._element.append(deepcopy(first._element))
    second = document.styles.add_style("CommentText", WD_STYLE_TYPE.PARAGRAPH)
    assert second._element is not None
    document.styles._element.append(deepcopy(second._element))
    original_styles = [
        etree.tostring(element)
        for element in document.styles._element
        if element.get(qn("w:styleId")) in {"CommentText", "CommentReference"}
    ]
    paragraph = document.add_paragraph("Existing")
    marker = OxmlElement("w:commentReference")
    marker.set(qn("w:id"), "2147483647")
    paragraph.add_run()._r.append(marker)
    number = document.add_paragraph().add_run("一、")
    add_numbering_review(document, number)
    assert next(iter(document.comments)).comment_id == 0
    assert [
        etree.tostring(element)
        for element in document.styles._element
        if element.get(qn("w:styleId")) in {"CommentText", "CommentReference"}
    ] == original_styles
    range_end = number._r.getnext()
    assert range_end is not None
    reference_run = range_end.getnext()
    assert reference_run is not None
    properties = reference_run.find(qn("w:rPr"))
    assert properties is not None
    style = properties.find(qn("w:rStyle"))
    assert style is not None
    reference_style = style.get(qn("w:val"), "")
    assert reference_style.startswith("DocWenNumberReview")
    output = tmp_path / "sparse-id.docx"
    document.save(output)
    assert next(iter(Document(str(output)).comments)).text == NUMBERING_REVIEW_TEXT


def test_malformed_legacy_comment_ids_are_preserved_without_using_library_allocator(tmp_path):
    document = Document()
    document.styles.add_style("CommentReference", WD_STYLE_TYPE.CHARACTER)
    document.styles.add_style("CommentText", WD_STYLE_TYPE.PARAGRAPH)
    old = document.add_comment(document.add_paragraph("Original").runs[0], "Original")
    old._element.set(qn("w:id"), "bad_id")
    malformed = OxmlElement("w:comment")
    malformed.set(qn("w:author"), "Original missing ID")
    document.comments._comments_elm.append(malformed)
    originals = [etree.tostring(old._element), etree.tostring(malformed)]
    add_numbering_review(document, document.add_paragraph().add_run("一、"))
    comments = list(document.comments)
    assert [etree.tostring(comment._element) for comment in comments[:2]] == originals
    assert comments[2].comment_id == 1  # original body anchor still reserves ID 0
    output = tmp_path / "legacy-comments.docx"
    document.save(output)
    with ZipFile(output) as package:
        root = etree.fromstring(package.read("word/comments.xml"))
        assert root[0].get(qn("w:id")) == "bad_id"
        assert root[1].get(qn("w:id")) is None
        assert root[2].get(qn("w:id")) == "1"


@pytest.mark.parametrize(
    "lexical,expected",
    [
        ("0" * 5000, 1),
        (" \t+" + "0" * 5000 + "\r\n", 1),
        ("0_0", 0),
        ("００", 0),
        ("٠", 0),
    ],
    ids=["long-zero", "xml-whitespace", "underscore", "fullwidth", "arabic-digit"],
)
def test_xml_integer_comment_identity_reserves_long_zero_and_rejects_non_xml_spellings(tmp_path, lexical, expected):
    document = Document()
    document.styles.add_style("CommentReference", WD_STYLE_TYPE.CHARACTER)
    document.styles.add_style("CommentText", WD_STYLE_TYPE.PARAGRAPH)
    old = document.add_comment(document.add_paragraph("Original").runs[0], "Original")
    for root in (document.element, old._element):
        for node in root.iter():
            if node.tag in {
                qn("w:comment"),
                qn("w:commentRangeStart"),
                qn("w:commentRangeEnd"),
                qn("w:commentReference"),
            }:
                node.set(qn("w:id"), lexical)
    add_numbering_review(document, document.add_paragraph().add_run("一、"))
    assert list(document.comments)[1].comment_id == expected
    assert old._element.get(qn("w:id")) == lexical
    output = tmp_path / "lexical-comment.docx"
    document.save(output)
    with ZipFile(output) as package:
        comments = etree.fromstring(package.read("word/comments.xml"))
        assert comments[0].get(qn("w:id")) == lexical
        assert comments[1].get(qn("w:id")) == str(expected)
