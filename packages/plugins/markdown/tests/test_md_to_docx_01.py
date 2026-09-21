"""Focused tests split from test_md_to_docx.py."""

from __future__ import annotations

from docx import Document

from ._md_to_docx_support import (
    MdToDocxConverter,
    Path,
    make_context,
    pytest,
    write_temp_md,
)

pytestmark = pytest.mark.contract


def test_temp_helpers_release_file_handles_and_own_distinct_staging_paths() -> None:
    source = Path(write_temp_md("# Temp ownership\n"))
    renamed = source.with_name(f"{source.stem}-renamed{source.suffix}")

    source.rename(renamed)
    renamed.unlink()
    source.write_text("# Recreated after closed descriptor\n", encoding="utf-8")

    _first_context, first_workspace = make_context(str(source), target_format="docx")
    _second_context, second_workspace = make_context(str(source), target_format="docx")
    staging_paths = (Path(first_workspace.staging_dir), Path(second_workspace.staging_dir))

    assert staging_paths[0] != staging_paths[1]
    assert all(path.is_dir() for path in staging_paths)


def test_success_outcome_is_complete_before_artifact_registration(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    source = tmp_path / "atomic-success.md"
    source.write_text("# Heading ^heading\n", encoding="utf-8")
    context, workspace = make_context(str(source), target_format="docx")

    def reject_terminal_progress(percent: float, _message: str = "") -> None:
        if percent == 100.0:
            raise RuntimeError("synthetic terminal progress failure")

    monkeypatch.setattr(context.progress, "report_progress", reject_terminal_progress)
    result = MdToDocxConverter().convert(context)

    assert not result.success
    assert result.artifacts == []
    assert workspace.registered_artifacts == []


def test_unsupported_source_eol_fails_semantics_v3_with_zero_artifacts(tmp_path: Path) -> None:
    source = tmp_path / "unsupported-eol.md"
    source.write_bytes(b"```rust\rbody\r```\r")
    context, workspace = make_context(str(source), target_format="docx")

    result = MdToDocxConverter().convert(context)

    assert not result.success
    assert result.artifacts == []
    assert workspace.registered_artifacts == []
    assert list(Path(workspace.staging_dir).iterdir()) == []
    assert result.error is not None
    assert result.error.error_type == "invalid_document_semantics"
    assert result.error.diagnostic_code == "MD2DOCX-SEMANTICS-V3-UNSUPPORTED"
    assert "only LF and CRLF" in result.error.message
    assert [item.code for item in result.diagnostics] == ["MD2DOCX-SEMANTICS-V3-UNSUPPORTED"]


def test_custom_template_without_body_placeholder_omits_markdown_body(tmp_path: Path) -> None:
    source = tmp_path / "metadata-only.md"
    source.write_text(
        "---\ntitle: Metadata title\n---\n\n# Body heading\n\nBody paragraph that must not be appended.\n",
        encoding="utf-8",
    )
    template_path = tmp_path / "metadata-only-template.docx"
    template = Document()
    template.add_paragraph("Title: {{title}}")
    template.add_paragraph("Template suffix")
    template.save(str(template_path))

    context, _workspace = make_context(
        str(source),
        target_format="docx",
        options={"template_name": str(template_path)},
    )
    result = MdToDocxConverter().convert(context)

    assert result.success is True, result.error
    output = Path(result.artifacts[0].staging_path)
    reopened = Document(str(output))
    visible_text = "\n".join(paragraph.text for paragraph in reopened.paragraphs)

    assert "Metadata title" in visible_text
    assert "Template suffix" in visible_text
    assert "Body heading" not in visible_text
    assert "Body paragraph that must not be appended." not in visible_text
