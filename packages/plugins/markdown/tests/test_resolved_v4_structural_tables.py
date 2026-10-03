"""Resolved exact-two Structural Tables rendering contract."""

from __future__ import annotations

import hashlib
import json
from pathlib import Path
from typing import Any

import pytest
from docx import Document
from docx.oxml.ns import qn

from docwen_core.docx_parsing.document_semantics import extract_semantic_table_metadata
from docwen_core.docx_parsing.table_extraction import build_docx_table_semantic_grid
from docwen_core.models.resolved_numbering import canonicalize_numbering_plan
from docwen_plugin_markdown.to_docx.converter import MdToDocxConverter

from .test_resolved_v4_to_docx import _context, _refs, _sha_text

pytestmark = pytest.mark.contract


def test_exact_two_structural_table_preserves_roles_and_merges(tmp_path: Path) -> None:
    source = """| Region | Sales | < |
| Quarter | Q1 | Q2 |
| --- || --- | --- |
| North | 10 | 12 |
| ^ | 8 | 11 |
"""
    plan_body = {"heading_definitions": [], "heading_instances": [], "targets": []}
    source_sha256 = _sha_text(source)
    plan_sha256 = hashlib.sha256(canonicalize_numbering_plan(plan_body)).hexdigest()
    neutral_payload = {
        "$schema": "urn:docwen:schema:resolved-document:v1",
        "schema": "docwen.resolved_document.v1",
        "input_id": "structural-exact-two",
        "source_sha256": source_sha256,
        "plan_sha256": plan_sha256,
        "document": {
            "authored_markdown": source,
            "targets": [],
            "references": [],
            "resource_occurrences": [],
            "citations": [],
            "resources": [],
        },
    }
    plan_payload = {
        "$schema": "urn:docwen:schema:numbering-export-plan:v1",
        "schema": "docwen.numbering_export_plan.v1",
        "input_id": "structural-exact-two",
        "source_sha256": source_sha256,
        "plan_sha256": plan_sha256,
        "plan": plan_body,
    }
    neutral = tmp_path / "structural-neutral.json"
    plan = tmp_path / "structural-plan.json"
    neutral.write_text(json.dumps(neutral_payload, separators=(",", ":")), encoding="utf-8")
    plan.write_text(json.dumps(plan_payload, separators=(",", ":")), encoding="utf-8")
    context, _workspace = _context(
        tmp_path / "run",
        _refs(neutral, plan),
        options={
            "locale": "zh_CN",
            "heading_merge_mode": "never",
            "markdown_extensions": {"input": {"structural_tables": True}},
        },
    )

    result = MdToDocxConverter().convert(context)

    assert result.success, result.error
    document = Document(result.artifacts[0].staging_path)
    assert len(document.tables) == 1
    table = document.tables[0]
    metadata = extract_semantic_table_metadata(table._tbl)
    assert (metadata.header_rows, metadata.header_columns) == (2, 1)

    def cell_text(cell: Any, _row: int, _column: int) -> str:
        return "".join(node.text or "" for node in cell.iter(qn("w:t")))

    grid = build_docx_table_semantic_grid(table._tbl, cell_text_resolver=cell_text)
    anchors = [cell for row in grid for cell in row if not cell.is_covered]
    assert any(cell.anchor_text == "Sales" and cell.colspan == 2 for cell in anchors)
    assert any(cell.anchor_text == "North" and cell.rowspan == 2 for cell in anchors)
