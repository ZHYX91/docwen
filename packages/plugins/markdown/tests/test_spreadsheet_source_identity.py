"""User-visible spreadsheet names must survive private input staging."""

from __future__ import annotations

import csv
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

from docwen_plugin_markdown.to_spreadsheet.converter import MdToCsvConverter, MdToXlsxConverter

from .conftest import make_context

pytestmark = pytest.mark.contract


@pytest.mark.parametrize("target", ["xlsx", "csv"])
@pytest.mark.parametrize("title", [None, "Authored title"])
def test_template_title_uses_declared_source_name(tmp_path: Path, target: str, title: str | None) -> None:
    source = tmp_path / "input-0000.md"
    source.write_text(f"---\ntitle: {title}\n---\n" if title else "# Ordinary note\n", encoding="utf-8")
    template = tmp_path / "template.xlsx"
    workbook = Workbook()
    assert workbook.active is not None
    workbook.active["A1"] = "{{title}}"
    workbook.save(template)
    workbook.close()
    context, _ = make_context(str(source), target_format=target, options={"template_name": str(template)})
    context.request.input_refs[0].logical_path = "Folder/原始 笔记.md"

    result = (MdToXlsxConverter() if target == "xlsx" else MdToCsvConverter()).convert(context)

    assert result.success, result.error
    artifact = result.artifacts[0]
    assert artifact.suggested_name.startswith("原始 笔记")
    expected = title or "原始 笔记"
    if target == "xlsx":
        loaded = load_workbook(artifact.staging_path, read_only=True)
        assert loaded.active is not None
        assert loaded.active["A1"].value == expected
        loaded.close()
    else:
        with Path(artifact.staging_path).open(encoding="utf-8-sig", newline="") as handle:
            assert next(csv.reader(handle)) == [expected]


@pytest.mark.parametrize("target", ["xlsx", "csv"])
def test_plain_table_export_uses_declared_source_name(tmp_path: Path, target: str) -> None:
    source = tmp_path / "input-0000.md"
    source.write_text("| Item | Value |\n| --- | --- |\n| Unicode | 中文数据 |\n", encoding="utf-8")
    context, _ = make_context(str(source), target_format=target)
    context.request.input_refs[0].logical_path = "Folder/原始 笔记.md"
    result = (MdToXlsxConverter() if target == "xlsx" else MdToCsvConverter()).convert(context)
    assert result.success, result.error
    assert result.artifacts[0].suggested_name == f"原始 笔记.{target}"
