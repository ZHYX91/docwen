"""Literal attachment numbers keep a compact separator at both digit boundaries."""

from __future__ import annotations

from io import BytesIO

import pytest
from docx import Document
from docx.enum.text import WD_TAB_ALIGNMENT
from docx.oxml.ns import qn
from docx.shared import Pt
from tests.support.gongwen_template import attachment_template

from docwen_plugin_markdown.template_filler import fill_template
from docwen_plugin_markdown.template_utils import scan_placeholders

pytestmark = pytest.mark.contract


@pytest.mark.parametrize("location", ["body", "cell"])
@pytest.mark.parametrize("count", [9, 10, 99, 100, 101])
def test_saved_literal_padding_crosses_ten_and_hundred_without_tabs(location, count) -> None:
    document = attachment_template(location=location)
    styles_before = document.styles.element.xml
    anchor = document.paragraphs[1] if location == "body" else document.tables[0].cell(0, 0).paragraphs[0]
    # Authored stops stay untouched and are never used for the generated list.
    anchor.paragraph_format.tab_stops.add_tab_stop(Pt(140), WD_TAB_ALIGNMENT.RIGHT)
    names = [f"材料{i}长名称需要续行时仍保持名称首字对齐" * 4 for i in range(1, count + 1)]
    fill_template(document, {"附件说明": names}, [], None, scan_placeholders(document))
    saved = BytesIO()
    document.save(saved)
    saved.seek(0)
    reopened = Document(saved)
    assert reopened.styles.element.xml == styles_before
    paragraphs = reopened.paragraphs[1:] if location == "body" else reopened.tables[0].cell(0, 0).paragraphs
    expected_prefixes = {
        1: "附件：1. ",
        2: "\u3000\u3000\u30002. ",
        9: "\u3000\u3000\u30009. ",
        10: "\u3000\u3000\u200210. ",
        11: "\u3000\u3000\u200211. ",
        99: "\u3000\u3000\u200299. ",
        100: "\u3000\u3000100. ",
        101: "\u3000\u3000101. ",
    }
    assert len(paragraphs) == count
    for index, prefix in expected_prefixes.items():
        if index <= count:
            assert paragraphs[index - 1].text == prefix + names[index - 1]
    for paragraph in paragraphs:
        assert not paragraph._p.findall(f"{qn('w:r')}/{qn('w:tab')}")
        assert paragraph._p.find(f"{qn('w:pPr')}/{qn('w:numPr')}") is None
        assert ". " in paragraph.text and ".  " not in paragraph.text
        assert paragraph.paragraph_format.left_indent == Pt(48)
        assert paragraph.paragraph_format.first_line_indent == Pt(-64)
        prefix, name = paragraph.runs
        assert prefix._r.rPr is not None
        kern = prefix._r.rPr.find(qn("w:kern"))
        assert kern is not None and kern.get(qn("w:val")) == "33"
        assert name._r.rPr is None or name._r.rPr.find(qn("w:kern")) is None
        stops = list(paragraph.paragraph_format.tab_stops)
        assert len(stops) == 1 and stops[0].position == Pt(140) and stops[0].alignment == WD_TAB_ALIGNMENT.RIGHT
