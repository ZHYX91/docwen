"""Request-owned ordinary Markdown prose paragraph boundaries."""

from __future__ import annotations

import pytest
from docx import Document

from docwen_plugin_markdown.ast_transforms import split_plain_paragraph_lines
from docwen_plugin_markdown.mistune_extensions import parse_markdown_text
from docwen_plugin_markdown.renderer import MdToDocxRenderer
from docwen_plugin_markdown.to_docx.converter import MdToDocxConverter

from .conftest import make_context

pytestmark = pytest.mark.unit


def test_off_by_default_keeps_soft_breaks() -> None:
    nodes = parse_markdown_text("First\nSecond\nThird")
    doc = Document()
    MdToDocxRenderer(doc).render(nodes)
    assert [p.text for p in doc.paragraphs] == ["First\nSecond\nThird"]


def test_each_line_creates_independent_word_paragraphs() -> None:
    nodes = parse_markdown_text("First **bold**\nSecond `code`\nThird")
    doc = Document()
    MdToDocxRenderer(doc).render(split_plain_paragraph_lines(nodes))
    assert [p.text for p in doc.paragraphs] == [
        "First bold",
        "Second code",
        "Third",
    ]


def test_hard_break_stays_inside_one_word_paragraph() -> None:
    nodes = parse_markdown_text("First  \nSecond")
    assert split_plain_paragraph_lines(nodes) == nodes


def test_lines_inside_nested_emphasis_keep_formatting() -> None:
    nodes = parse_markdown_text("Before **bold *first\nsecond* end** after")
    document = Document()
    MdToDocxRenderer(document).render(split_plain_paragraph_lines(nodes))
    assert [paragraph.text for paragraph in document.paragraphs] == ["Before bold first", "second end after"]
    for paragraph, word in zip(document.paragraphs, ["first", "second"], strict=True):
        run = next(run for run in paragraph.runs if word in run.text)
        assert run.bold and run.italic


def test_nested_semantic_owner_preserves_paragraph_boundary() -> None:
    node = {
        "type": "paragraph",
        "children": [
            {
                "type": "strong",
                "children": [
                    {"type": "text", "raw": "First"},
                    {"type": "softbreak"},
                    {"type": "semantic_citation", "raw": "[@smith2020]"},
                ],
            }
        ],
    }
    assert split_plain_paragraph_lines([node]) == [node]


def test_nested_containers_are_unchanged() -> None:
    source = "> Quote one\n> Quote two\n\n- Item one\n- Item two"
    nodes = parse_markdown_text(source)
    assert split_plain_paragraph_lines(nodes) == nodes


def test_semantic_owner_does_not_split() -> None:
    node = {
        "type": "paragraph",
        "children": [
            {"type": "text", "raw": "First"},
            {"type": "softbreak"},
            {"type": "text", "raw": "Second"},
        ],
        "_docwen_v3_ordinary_anchor": {"id": "p"},
    }
    assert split_plain_paragraph_lines([node]) == [node]


@pytest.mark.parametrize(
    ("enabled", "expected"),
    [
        (False, ["First\nSecond\nThird"]),
        (True, ["First", "Second", "Third"]),
    ],
)
def test_exported_docx_respects_request_paragraph_policy(tmp_path, enabled: bool, expected: list[str]) -> None:
    source = tmp_path / "three-lines.md"
    source.write_text("First\nSecond\nThird", encoding="utf-8")
    template = tmp_path / "body.docx"
    document = Document()
    document.add_paragraph("{{body}}")
    document.save(template)

    context, _workspace = make_context(
        str(source),
        options={"template_name": str(template), "each_line_paragraph": enabled},
    )
    result = MdToDocxConverter().convert(context)
    assert result.success, result.error
    rendered = Document(result.artifacts[0].staging_path)
    assert [paragraph.text for paragraph in rendered.paragraphs] == expected


def test_protected_inline_semantics_are_not_split() -> None:
    children = [
        {"type": "text", "raw": "Before"},
        {"type": "softbreak"},
        {"type": "semantic_citation", "raw": "[@smith2020]"},
    ]
    node = {"type": "paragraph", "children": children}
    assert split_plain_paragraph_lines([node]) == [node]
