"""Schema order and resolved foundation identity across both style tables."""

from __future__ import annotations

from io import BytesIO
from typing import Any
from zipfile import ZipFile

import lxml.etree as etree
import pytest
from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml.ns import qn
from docx.styles.style import ParagraphStyle

from docwen_plugin_markdown.mistune_extensions import parse_markdown_text
from docwen_plugin_markdown.renderer import MdToDocxRenderer
from docwen_plugin_markdown.to_docx.managed_styles import complete_managed_styles
from docwen_runtime.config.document_styles import build_document_style_catalog

from .conftest import PROJECT_ROOT

pytestmark = pytest.mark.contract

_STYLE_CHILDREN = (
    "name",
    "aliases",
    "basedOn",
    "next",
    "link",
    "autoRedefine",
    "hidden",
    "uiPriority",
    "semiHidden",
    "unhideWhenUsed",
    "qFormat",
    "locked",
    "personal",
    "personalCompose",
    "personalReply",
    "rsid",
    "pPr",
    "rPr",
    "tblPr",
    "trPr",
    "tcPr",
    "tblStylePr",
)


def _catalog(locale: str = "en_US"):
    return build_document_style_catalog(
        {"gui": {"language": {"locale": locale}}},
        locales_dir=PROJECT_ROOT / "i18n" / "locales",
    )


def _style_tables(document: Any) -> dict[str, etree._Element]:
    buffer = BytesIO()
    document.save(buffer)
    with ZipFile(buffer) as archive:
        return {
            name: etree.fromstring(archive.read(name)) for name in ("word/styles.xml", "word/stylesWithEffects.xml")
        }


def _rewrite_effects(document: Any, transform: Any) -> None:
    part = next(part for part in document.part.package.parts if str(part.partname) == "/word/stylesWithEffects.xml")
    root = etree.fromstring(part.blob)
    transform(root)
    part._blob = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)


@pytest.mark.parametrize("locale", ["en_US", "zh_CN"])
def test_managed_style_ui_children_obey_ct_style_order(locale: str) -> None:
    template = Document()

    def remove_styles(root: Any) -> None:
        for style in root.findall(qn("w:style")):
            root.remove(style)

    remove_styles(template.styles.element)
    _rewrite_effects(template, remove_styles)
    # This synthetic empty registry also removes unused built-in numbering
    # references, so it does not ask completion to invent out-of-registry styles.
    for reference in list(template.part.numbering_part.element.iter(qn("w:pStyle"))):
        parent = reference.getparent()
        assert parent is not None
        parent.remove(reference)
    document, bindings = complete_managed_styles(template, _catalog(locale))
    required_ids = {style_id for _, style_id in (*bindings.foundation_style_ids, *bindings.style_ids)}
    rank = {qn(f"w:{name}"): index for index, name in enumerate(_STYLE_CHILDREN)}
    for part_name, root in _style_tables(document).items():
        by_id = {item.get(qn("w:styleId")): item for item in root.findall(qn("w:style"))}
        assert set(by_id) == required_ids
        assert len(required_ids) == 46
        for style_id in required_ids:
            style = by_id[style_id]
            ordered = [rank[child.tag] for child in style]
            assert ordered == sorted(ordered), (part_name, style_id)
            priorities = style.findall(qn("w:uiPriority"))
            assert len(priorities) <= 1
            if any(style.find(qn(f"w:{name}")) is not None for name in ("qFormat", "semiHidden", "unhideWhenUsed")):
                assert len(priorities) == 1, (part_name, style_id)
        for key, flags in (
            ("normal", ["uiPriority", "qFormat"]),
            ("default_paragraph_font", ["uiPriority", "semiHidden", "unhideWhenUsed"]),
            ("table_normal", ["uiPriority", "semiHidden", "unhideWhenUsed"]),
        ):
            style = by_id[dict(bindings.foundation_style_ids)[key]]
            assert [child.tag for child in style if child.tag in {qn(f"w:{name}") for name in flags}] == [
                qn(f"w:{name}") for name in flags
            ]


@pytest.mark.parametrize("normal_id", ["HostNormal42", "0"])
def test_host_specific_normal_id_uses_resolved_foundation_identity(normal_id: str) -> None:
    template = Document()
    normal = template.styles["Normal"]
    assert isinstance(normal, ParagraphStyle)
    normal_element = normal._element
    assert normal_element is not None
    normal_element.set(qn("w:styleId"), normal_id)
    for reference in template.styles.element.iter(qn("w:basedOn")):
        if reference.get(qn("w:val")) == "Normal":
            reference.set(qn("w:val"), normal_id)

    def rename_effects_normal(root: Any) -> None:
        for style in root.findall(qn("w:style")):
            if style.get(qn("w:styleId")) == "Normal":
                style.set(qn("w:styleId"), "EffectsNormal")
        for reference in root.iter(qn("w:basedOn")):
            if reference.get(qn("w:val")) == "Normal":
                reference.set(qn("w:val"), "EffectsNormal")

    _rewrite_effects(template, rename_effects_normal)
    template, bindings = complete_managed_styles(template, _catalog())
    assert dict(bindings.foundation_style_ids)["normal"] == normal_id
    actual_normal = next(style for style in template.styles if style.style_id == normal_id)
    assert isinstance(actual_normal, ParagraphStyle)
    placeholder = template.add_paragraph("{{body}}", style=actual_normal)
    rendered = MdToDocxRenderer(template, body_style=placeholder.style, managed_styles=bindings).render(
        parse_markdown_text("Body text.\n")
    )
    assert rendered[0].style.style_id == bindings.style_id("body_paragraph")
    for part_name, root in _style_tables(template).items():
        normal_styles = [item for item in root.findall(qn("w:style")) if item.get(qn("w:styleId")) == normal_id]
        assert len(normal_styles) == 1, part_name
        assert normal_styles[0].get(qn("w:type")) == "paragraph"
        assert not any(item.get(qn("w:styleId")) in {"Normal", "EffectsNormal"} for item in root.findall(qn("w:style")))
        body = next(
            item
            for item in root.findall(qn("w:style"))
            if item.get(qn("w:styleId")) == bindings.style_id("body_paragraph")
        )
        based_on = body.find(qn("w:basedOn"))
        assert based_on is not None and based_on.get(qn("w:val")) == normal_id

    custom = template.styles.add_style("CustomBody", WD_STYLE_TYPE.PARAGRAPH)
    assert isinstance(custom, ParagraphStyle)
    custom.base_style = actual_normal
    rendered_custom = MdToDocxRenderer(template, body_style=custom, managed_styles=bindings).render(
        parse_markdown_text("Custom body text.\n")
    )
    assert rendered_custom[0].style.style_id == custom.style_id
