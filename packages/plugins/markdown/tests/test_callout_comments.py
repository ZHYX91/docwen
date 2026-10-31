"""Saved callouts keep editable quote content and native source-type comments."""

from __future__ import annotations

from pathlib import Path

import pytest
from docx import Document
from docx.oxml.ns import qn

from docwen_plugin_markdown.mistune_extensions import parse_markdown_text
from docwen_plugin_markdown.renderer import MdToDocxRenderer

pytestmark = pytest.mark.contract


@pytest.mark.parametrize(
    ("source", "visible", "comment"),
    [
        (
            "> [!warning]- **Check first**\n> Save the file.\n",
            ["Check first", "Save the file."],
            "Markdown callout type: warning\nOriginal state: collapsed (-)",
        ),
        ("> [!custom-kind]\n> Body\n", ["Body"], "Markdown callout type: custom-kind"),
        ("> [!提示]\n> Body\n", ["Body"], "Markdown callout type: 提示"),
        ("> [!NOTE]+ Title only\n", ["Title only"], "Markdown callout type: NOTE\nOriginal state: expanded (+)"),
        ("> [!info]\n", [], "Markdown callout type: info"),
        ("> [!tip]\n>\n> - First\n> - Second\n", ["First", "Second"], "Markdown callout type: tip"),
    ],
)
def test_saved_callout_content_and_native_comment(
    tmp_path: Path, source: str, visible: list[str], comment: str
) -> None:
    document = Document()
    MdToDocxRenderer(document).render(parse_markdown_text(source))
    path = tmp_path / "callout.docx"
    document.save(str(path))
    restored = Document(str(path))
    assert [p.text for p in restored.paragraphs if p.text] == visible
    assert [item.text for item in restored.comments] == [comment]
    assert len(list(restored.element.iter(qn("w:commentRangeStart")))) == 1
    assert len(list(restored.element.iter(qn("w:commentRangeEnd")))) == 1
    assert "[!" not in "".join(p.text for p in restored.paragraphs)
    if "**Check first**" in source:
        assert any(run.bold and run.text == "Check first" for run in restored.paragraphs[0].runs)


@pytest.mark.parametrize(
    "source",
    [
        "> \\[!warning] literal\n",
        "> `[!warning]` literal\n",
        "[!warning] ordinary\n",
        "> ordinary [!warning]\n",
        "> ```\n> [!warning]\n> ```\n",
        "> [!warning]suffix\n",
    ],
)
def test_literal_markers_do_not_become_comments(source: str) -> None:
    document = Document()
    MdToDocxRenderer(document).render(parse_markdown_text(source))
    assert "[!warning]" in "\n".join(p.text for p in document.paragraphs)
    assert len(document.comments) == 0


def test_nested_callouts_and_table_only_body_keep_separate_anchors(tmp_path: Path) -> None:
    document = Document()
    MdToDocxRenderer(document).render(
        parse_markdown_text(
            "> [!note] Outer\n>\n> > [!warning] Inner\n> > text\n\n> [!tip]\n>\n> | A | B |\n> |---|---|\n> | 1 | 2 |\n"
        )
    )
    path = tmp_path / "nested.docx"
    document.save(str(path))
    restored = Document(str(path))
    assert sorted(item.text for item in restored.comments) == [
        "Markdown callout type: note",
        "Markdown callout type: tip",
        "Markdown callout type: warning",
    ]
    assert len(restored.tables) == 1
    assert restored.tables[0].cell(0, 0).text == "A"
    assert len(list(restored.element.iter(qn("w:commentRangeStart")))) == 3


def test_note_story_callout_uses_readable_type_with_located_diagnostic() -> None:
    from docwen_plugin_markdown.to_docx.notes import process_md_body_with_notes

    _body, notes = process_md_body_with_notes("Text[^1]\n\n[^1]: > [!warning] Note title\n    > body\n")
    children = notes._footnote_children["1"]
    text = "".join(str(item.get("raw", "")) for paragraph in children for item in paragraph)
    assert text == "Warning: Note title\nbody"
    renderer = MdToDocxRenderer(Document(), note_ctx=notes)
    assert renderer.warnings == (
        ("CALLOUT-NOTE-TEXT", "Callouts in notes (1) use readable type labels instead of Word comments."),
    )


def test_untitled_nested_callouts_can_share_first_paragraph(tmp_path: Path) -> None:
    document = Document()
    MdToDocxRenderer(document).render(parse_markdown_text("> [!note]\n>\n> > [!warning]\n> > Body\n"))
    path = tmp_path / "shared-anchor.docx"
    document.save(str(path))
    restored = Document(str(path))
    assert [p.text for p in restored.paragraphs if p.text] == ["Body"]
    assert sorted(item.text for item in restored.comments) == [
        "Markdown callout type: note",
        "Markdown callout type: warning",
    ]
    for tag in ("commentRangeStart", "commentRangeEnd", "commentReference"):
        assert sorted(element.get(qn("w:id")) or "" for element in restored.element.iter(qn(f"w:{tag}"))) == ["0", "1"]
