"""Both production MD→DOCX routes share built-in Gongwen template semantics."""

from __future__ import annotations

import hashlib
import json
from copy import deepcopy
from pathlib import Path

import pytest
from docx import Document
from docx.oxml.ns import qn
from docx.shared import Pt
from tests.support.config import FakeConfigView
from tests.support.gongwen_template import attachment_template

from docwen_core.models.resolved_numbering import canonicalize_numbering_plan
from docwen_plugin_markdown.to_docx.converter import MdToDocxConverter
from docwen_runtime.config.document_styles import build_document_style_catalog

from .conftest import PROJECT_ROOT, make_context
from .test_resolved_v4_to_docx import _context, _refs

pytestmark = pytest.mark.contract

_SOURCE = """---
aliases: [公文别名]
附件说明:
  - 1. 材料清单
  - （二）办理依据
抄送机关: [单位甲, 单位乙]
附注: （公开发布）
成文日期: 2026-06-27
印发日期: 2026/6/28
普通列表: [甲, 乙]
first: '{{second}}'
second: 实际值
---

正文保留字面量 {{附件说明}}。
"""


def _conversion_context(tmp_path: Path, route: str, template: Path, enabled: bool | None, source_text: str = _SOURCE):
    config = {
        "template_fill": {"list_separator": " / "},
        "gui": {"language": {"locale": "en_US"}},
        "conversion": {"output_locale": "en_US"},
    }
    if enabled is not None:
        config["template_fill"]["gongwen_fields_enabled"] = enabled
    options = {"template_name": str(template), "locale": "en_US"}
    style_catalog = build_document_style_catalog(
        config, request_options=options, locales_dir=PROJECT_ROOT / "i18n" / "locales"
    )
    if route == "source":
        source = tmp_path / "source.md"
        source.write_text(source_text, encoding="utf-8")
        return make_context(str(source), options=options, config_values=config, document_style_catalog=style_catalog)

    plan_body = {"heading_definitions": [], "heading_instances": [], "targets": []}
    source_hash = hashlib.sha256(source_text.encode("utf-8")).hexdigest()
    plan_hash = hashlib.sha256(canonicalize_numbering_plan(plan_body)).hexdigest()
    envelope = {"input_id": "gongwen-fields", "source_sha256": source_hash, "plan_sha256": plan_hash}
    neutral_payload = {
        **envelope,
        "$schema": "urn:docwen:schema:resolved-document:v1",
        "schema": "docwen.resolved_document.v1",
        "document": {
            "authored_markdown": source_text,
            "targets": [],
            "references": [],
            "resource_occurrences": [],
            "citations": [],
            "resources": [],
        },
    }
    plan_payload = {
        **envelope,
        "$schema": "urn:docwen:schema:numbering-export-plan:v1",
        "schema": "docwen.numbering_export_plan.v1",
        "plan": plan_body,
    }
    neutral_path = tmp_path / "neutral.json"
    plan_path = tmp_path / "plan.json"
    neutral_path.write_text(json.dumps(neutral_payload, ensure_ascii=False), encoding="utf-8")
    plan_path.write_text(json.dumps(plan_payload), encoding="utf-8")
    context, workspace = _context(tmp_path, _refs(neutral_path, plan_path), options=options, style_locale="en_US")
    context._config = FakeConfigView(config)
    context._document_style_catalog = style_catalog
    return context, workspace


@pytest.mark.parametrize("route", ["source", "resolved"])
@pytest.mark.parametrize("enabled", [None, True, False])
def test_real_docx_fields_share_switch_and_preserve_yaml_and_literal_values(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
    route: str,
    enabled: bool | None,
) -> None:
    from docwen_plugin_markdown.to_docx import converter, resolved_v4_route

    module = converter if route == "source" else resolved_v4_route
    original_extract = module.extract_yaml_front_matter
    observed = []

    def observe_extract(content):
        values, body = original_extract(content)
        observed.append((values, deepcopy(values)))
        return values, body

    monkeypatch.setattr(module, "extract_yaml_front_matter", observe_extract)
    template = Document()
    template.add_paragraph("{{标题}}")
    template.add_paragraph("{{正文}}")
    template.add_paragraph("{{附件说明}}")
    template.add_paragraph("{{抄送机关}}|{{附注}}|{{成文日期}}|{{印发日期}}")
    template.add_paragraph("{{普通列表}}")
    template.add_paragraph("{{first}} / {{second}}")
    template.add_paragraph("{{主送机关}}：")
    table = template.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "{{附件说明}}"
    table.cell(0, 1).add_table(rows=1, cols=1).cell(0, 0).text = "{{附件说明}}"
    template_path = tmp_path / "fields.docx"
    template.save(str(template_path))
    context, workspace = _conversion_context(tmp_path, route, template_path, enabled)

    result = MdToDocxConverter().convert(context)

    assert result.success, result.error
    assert len(workspace.registered_artifacts) == 1
    assert observed and all(values == original for values, original in observed)
    output = Document(result.artifacts[0].staging_path)
    paragraphs = [p.text for p in output.paragraphs]
    assert paragraphs[0] == "公文别名"
    assert output.core_properties.title == "公文别名"
    assert "正文保留字面量 {{附件说明}}。" in paragraphs
    assert "甲 / 乙" in paragraphs
    assert "{{second}} / 实际值" in paragraphs
    cells = [output.tables[0].cell(0, 0), output.tables[0].cell(0, 1).tables[0].cell(0, 0)]
    if enabled is False:
        assert "1. 材料清单 / （二）办理依据" in paragraphs
        assert "单位甲 / 单位乙|（公开发布）|2026-06-27|2026/6/28" in paragraphs
        assert "：" in paragraphs
        for cell in cells:
            assert [p.text for p in cell.paragraphs] == ["1. 材料清单 / （二）办理依据"]
    else:
        assert "附件：1. 材料清单" in paragraphs
        assert "\u3000\u3000\u30002. 办理依据" in paragraphs
        assert "单位甲、单位乙|公开发布|2026年6月27日|2026年6月28日印发" in paragraphs
        assert "：" not in paragraphs
        for cell in cells:
            assert [p.text for p in cell.paragraphs] == ["附件：1. 材料清单", "\u3000\u3000\u30002. 办理依据"]


@pytest.mark.parametrize("route", ["source", "resolved"])
def test_invalid_attachment_placement_returns_explicit_failure_without_artifacts(tmp_path: Path, route: str) -> None:
    template = Document()
    template.add_paragraph("{{正文}}")
    template.sections[0].header.paragraphs[0].text = "{{附件说明}}"
    template_path = tmp_path / "unsupported.docx"
    template.save(str(template_path))
    context, workspace = _conversion_context(tmp_path, route, template_path, True)

    result = MdToDocxConverter().convert(context)

    assert not result.success
    assert result.error is not None
    assert result.error.diagnostic_code == "MD2DOCX-GONGWEN-FIELD-PLACEMENT"
    assert "附件说明占位符必须独占" in result.error.message
    assert not result.artifacts
    assert not workspace.registered_artifacts


@pytest.mark.parametrize("enabled,empty", [(True, False), (True, True), (False, False), (False, True)])
def test_release_fixture_checks_real_conversion_layout(tmp_path: Path, enabled: bool, empty: bool) -> None:
    from scripts.release.gongwen_fields_fixture import verify_gongwen_fields_output, write_gongwen_fields_cases

    case = next(
        case
        for case in write_gongwen_fields_cases(tmp_path / "cases")
        if (case.enabled, case.empty) == (enabled, empty)
    )
    config = {
        "template_fill": {"list_separator": " / ", "gongwen_fields_enabled": enabled},
        "conversion": {"output_locale": "en_US"},
    }
    options = {"template_name": str(case.template_path), "locale": "en_US"}
    catalog = build_document_style_catalog(config, request_options=options, locales_dir=PROJECT_ROOT / "i18n/locales")
    context, _workspace = make_context(
        str(case.source_path), options=options, config_values=config, document_style_catalog=catalog
    )
    result = MdToDocxConverter().convert(context)
    assert result.success, result.error
    verify_gongwen_fields_output(Path(result.artifacts[0].staging_path), case=case)


@pytest.mark.parametrize("route", ["source", "resolved"])
@pytest.mark.parametrize("count", [1, 2, 9, 10, 99, 100, 101])
def test_generated_template_character_indents_survive_production_routes(tmp_path: Path, route: str, count: int) -> None:
    template = attachment_template()
    attachment_style = template.styles["Attachment"].element.xml
    template_path = tmp_path / "generated.docx"
    template.save(str(template_path))
    original_hash = hashlib.sha256(template_path.read_bytes()).hexdigest()
    names = [f"材料{i:02d}" for i in range(1, count + 1)]
    markdown = "---\n附件说明:\n" + "".join(f"  - {name}\n" for name in names) + "---\n\n正文。\n"
    context, _workspace = _conversion_context(tmp_path, route, template_path, True, markdown)

    result = MdToDocxConverter().convert(context)

    assert result.success, result.error
    output = Document(result.artifacts[0].staging_path)
    expected = (
        ["附件：" + names[0]]
        if count == 1
        else [
            f"{'附件：' if i == 1 else '　' * ((7 - len(str(i))) // 2) + ' ' * ((7 - len(str(i))) % 2)}{i}. {name}"
            for i, name in enumerate(names, 1)
        ]
    )
    attachments = [p for p in output.paragraphs if p.text in expected]
    assert [p.text for p in attachments] == expected
    hanging_points = 48 if count == 1 else 64
    for paragraph in attachments:
        assert paragraph.paragraph_format.left_indent == Pt(-16 + hanging_points)
        assert paragraph.paragraph_format.first_line_indent == Pt(-hanging_points)
        assert not list(paragraph._p.iter(qn("w:tab")))
        assert paragraph._p.pPr is not None and paragraph._p.pPr.ind is not None
        assert paragraph._p.pPr.ind.get(qn("w:leftChars")) == "0"
        assert paragraph._p.pPr.ind.get(qn("w:hangingChars")) == "0"
    assert output.styles["Attachment"].element.xml == attachment_style
    assert hashlib.sha256(template_path.read_bytes()).hexdigest() == original_hash
