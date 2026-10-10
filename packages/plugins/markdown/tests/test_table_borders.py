"""Physical header borders respect spans and the selected template rule."""

from __future__ import annotations

from functools import cache
from typing import Any

import pytest
from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from lxml import etree

from docwen_plugin_markdown.document_semantics import analyze_document_semantics
from docwen_plugin_markdown.mistune_extensions import parse_markdown_text
from docwen_plugin_markdown.renderer import MdToDocxRenderer
from docwen_plugin_markdown.table_borders import apply_multirow_header_borders, table_header_separator
from docwen_plugin_markdown.to_docx.managed_styles import complete_managed_styles
from docwen_runtime.config.document_styles import build_document_style_catalog

from .conftest import PROJECT_ROOT

pytestmark = pytest.mark.contract

_TWO = """| Region | Site | Scores | < |
| ^ | ^ | A | B |
| --- | --- || --- | --- |
| North | Alpha | 10 | < |
| ^ | Beta | ^ | ^ |
"""
_THREE = """| Region | All | < | < | < |
| ^ | West | < | East | < |
| ^ | W1 | W2 | E1 | E2 |
| --- || --- | --- | --- | --- |
| North | 10 | < | 20 | 21 |
| ^ | ^ | ^ | 22 | 23 |
"""
_STAGGERED = """| Stub | Block | < | Tier |
| ^ | ^ | ^ | Sub |
| ^ | B1 | B2 | Leaf |
| --- || --- | --- | --- |
| North | 10 | < | 20 |
| ^ | ^ | ^ | 21 |
"""


@cache
def _catalog():
    return build_document_style_catalog(
        {"gui": {"language": {"locale": "en_US"}}}, locales_dir=PROJECT_ROOT / "i18n" / "locales"
    )


def _render(source: str, *, document: Any = None, key: str | None = "three_line_table"):
    analysis = analyze_document_semantics(parse_markdown_text(source), current_v3=True)
    assert not analysis.has_errors, analysis.diagnostics
    document, bindings = complete_managed_styles(document or Document(), _catalog())
    MdToDocxRenderer(
        document,
        managed_styles=bindings,
        table_style_key=key,
        table_style_name=bindings.get(key or "three_line_table").name,
    ).render(analysis.ast)
    return document, analysis.ast[0]["_document_semantics_table"]


def _physical(table) -> dict[tuple[int, int, int], Any]:
    result = {}
    for row_index, row in enumerate(table._tbl.findall(qn("w:tr"))):
        column = 0
        for cell in row.findall(qn("w:tc")):
            grid_span = cell.find(f"{qn('w:tcPr')}/{qn('w:gridSpan')}")
            span = int(grid_span.get(qn("w:val"))) if grid_span is not None else 1
            result[(row_index, column, span)] = cell
            column += span
    return result


def _bottoms(table) -> dict[tuple[int, int, int], dict[str, str]]:
    result = {}
    for key, cell in _physical(table).items():
        bottom = cell.find(f"{qn('w:tcPr')}/{qn('w:tcBorders')}/{qn('w:bottom')}")
        if bottom is not None:
            result[key] = {name.rsplit("}", 1)[-1]: value for name, value in bottom.attrib.items()}
    return result


@pytest.mark.parametrize(
    "source, expected",
    [
        (
            _TWO,
            {(0, 2, 2): "single", (1, 0, 1): "single", (1, 1, 1): "single", (1, 2, 1): "single", (1, 3, 1): "single"},
        ),
        (
            _THREE,
            {
                (0, 1, 4): "single",
                (1, 1, 2): "single",
                (1, 3, 2): "single",
                **{(2, col, 1): "single" for col in range(5)},
            },
        ),
        (
            _STAGGERED,
            {(0, 3, 1): "nil", (1, 1, 2): "single", (1, 3, 1): "nil", **{(2, col, 1): "single" for col in range(4)}},
        ),
    ],
    ids=["two", "nested-three", "staggered-three"],
)
@pytest.mark.parametrize("repeat", [None, "true", "false"])
def test_multirow_headers_have_main_and_group_lines_on_real_bottom_edges(source, expected, repeat):
    document, metadata = _render(source + (f"{{repeat-header={repeat}}}\n" if repeat is not None else ""))
    table = document.tables[0]
    bottoms = _bottoms(table)
    assert {key: value["val"] for key, value in bottoms.items()} == expected
    assert all(value["sz"] == "4" for value in bottoms.values() if value["val"] == "single")
    # Last-row continuation cells carry the main border, not row.cells' restart aliases.
    boundary = metadata["header_rows"] - 1
    continuation = _physical(table)[(boundary, 0, 1)]
    assert continuation.find(f"{qn('w:tcPr')}/{qn('w:vMerge')}") is not None
    assert continuation is not table.rows[boundary].cells[0]._tc
    assert not any(key[0] >= metadata["header_rows"] for key in bottoms)
    markers = list(table._tbl.iter(qn("w:tblHeader")))
    assert len(markers) == (metadata["header_rows"] if repeat is not None else 0)
    assert all(marker.get(qn("w:val")) == ("1" if repeat == "true" else "0") for marker in markers)


@pytest.mark.parametrize("key", ["table_grid", None])
def test_other_style_or_name_alone_does_not_authorize_header_projection(key):
    document, _ = _render(_THREE, key=key)
    assert _bottoms(document.tables[0]) == {}


def test_known_docwen_direct_fallback_also_projects_merged_header_edges():
    analysis = analyze_document_semantics(parse_markdown_text(_STAGGERED), current_v3=True)
    assert not analysis.has_errors
    document = Document()
    MdToDocxRenderer(document, table_style_key="three_line_table", table_style_name="Three Line Table").render(
        analysis.ast
    )
    bottoms = _bottoms(document.tables[0])
    assert {key for key, value in bottoms.items() if value["val"] == "single"} == {
        (1, 1, 2),
        *((2, col, 1) for col in range(4)),
    }
    assert {key for key, value in bottoms.items() if value["val"] == "nil"} == {(0, 3, 1), (1, 3, 1)}


@pytest.mark.parametrize("source", ["| --- | --- |\n| A | B |\n", "| A | B |\n| --- | --- |\n| 1 | 2 |\n"])
def test_zero_and_single_header_keep_managed_style_behavior(source):
    document, _ = _render(source)
    assert _bottoms(document.tables[0]) == {}


def _rule(style, **attributes):
    rule = OxmlElement("w:tblStylePr")
    rule.set(qn("w:type"), "firstRow")
    properties = OxmlElement("w:tcPr")
    borders = OxmlElement("w:tcBorders")
    bottom = OxmlElement("w:bottom")
    bottom.attrib.update({qn(f"w:{name}"): value for name, value in attributes.items()})
    borders.append(bottom)
    properties.append(borders)
    rule.append(properties)
    style.element.append(rule)
    return bottom


@pytest.mark.parametrize("inherit", [False, True])
def test_template_separator_attributes_and_based_on_are_preserved(inherit):
    template = Document()
    style = template.styles.add_style("Three Line Table", WD_STYLE_TYPE.TABLE)
    parent = template.styles.add_style("Parent Table", WD_STYLE_TYPE.TABLE) if inherit else style
    if inherit:
        style.base_style = parent
    attributes = {"val": "double", "sz": "8", "color": "AB1234", "space": "1", "shadow": "1"}
    original = _rule(parent, **attributes)
    original_xml = etree.tostring(original)
    table_properties = OxmlElement("w:tblPr")
    outer = OxmlElement("w:tblBorders")
    for edge, value in (("top", "dashed"), ("bottom", "nil")):
        border = OxmlElement(f"w:{edge}")
        border.attrib.update({qn("w:val"): value, qn("w:sz"): "6", qn("w:color"): "445566"})
        outer.append(border)
    table_properties.append(outer)
    parent.element.append(table_properties)
    outer_before = [(edge.tag, dict(edge.attrib)) for edge in outer]
    document, _ = _render(_STAGGERED, document=template)
    assert etree.tostring(original) == original_xml
    assert table_header_separator(document.tables[0].style) == {qn(f"w:{k}"): v for k, v in attributes.items()}
    assert all(value == attributes for value in _bottoms(document.tables[0]).values() if value["val"] != "nil")
    active_parent = document.styles["Parent Table"] if inherit else document.tables[0].style
    actual_outer = active_parent.element.find(f"{qn('w:tblPr')}/{qn('w:tblBorders')}")
    assert [(edge.tag, dict(edge.attrib)) for edge in actual_outer] == outer_before


@pytest.mark.parametrize("value", [None, "none", "nil"])
def test_template_missing_or_explicitly_disabled_separator_is_not_replaced(value):
    template = Document()
    parent = template.styles.add_style("Parent Table", WD_STYLE_TYPE.TABLE)
    _rule(parent, val="double", sz="8", color="AB1234")
    style = template.styles.add_style("Three Line Table", WD_STYLE_TYPE.TABLE)
    if value is not None:
        style.base_style = parent
        original = _rule(style, val=value)
        original_xml = etree.tostring(original)
    document, _ = _render(_THREE, document=template)
    assert _bottoms(document.tables[0]) == {}
    if value is not None:
        assert etree.tostring(original) == original_xml
        assert (
            document.tables[0]
            .style.element.find(f"{qn('w:tblStylePr')}/{qn('w:tcPr')}/{qn('w:tcBorders')}/{qn('w:bottom')}")
            .get(qn("w:val"))
            == value
        )


@pytest.mark.parametrize("value", ["none", "nil", "double"])
def test_authored_direct_bottom_wins_and_projection_is_idempotent(value):
    document, metadata = _render(_STAGGERED, key=None)
    table = document.tables[0]
    boundary_cell = _physical(table)[(2, 0, 1)]
    properties = boundary_cell.find(qn("w:tcPr"))
    borders = OxmlElement("w:tcBorders")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), value)
    borders.append(bottom)
    properties.append(borders)
    properties_before = table._tbl.tblPr.xml
    kwargs = {
        "header_rows": metadata["header_rows"],
        "anchors": metadata["anchors"],
        "separator": table_header_separator(table.style),
    }
    apply_multirow_header_borders(table._tbl, **kwargs)
    first_xml = table._tbl.xml
    apply_multirow_header_borders(table._tbl, **kwargs)
    assert table._tbl.xml == first_xml
    assert table._tbl.tblPr.xml == properties_before
    assert _bottoms(table)[(2, 0, 1)] == {"val": value}
