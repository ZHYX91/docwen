"""Saved converter output preserves the effective ordinary-body run contract."""

from __future__ import annotations

from hashlib import sha256
from pathlib import Path
from typing import Any

import pytest
from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor
from docx.text.run import Run

from docwen_plugin_markdown.template_utils import extract_body_run_format, resolve_template

from ._md_to_docx_support import MdToDocxConverter, make_context

pytestmark = pytest.mark.contract


def _convert(tmp_path: Path, template: Any = None, *, body: str) -> Any:
    source = tmp_path / "source.md"
    source.write_text(body, encoding="utf-8")
    options = {}
    template_path = tmp_path / "template.docx"
    if template is not None:
        template.save(template_path)
        before = sha256(template_path.read_bytes()).hexdigest()
        options["template_name"] = str(template_path)
    context, _ = make_context(str(source), options=options)
    result = MdToDocxConverter().convert(context)
    assert result.success, result.error
    if template is not None:
        assert sha256(template_path.read_bytes()).hexdigest() == before
    return Document(result.artifacts[0].staging_path)


def _runs(paragraph) -> list[Run]:
    return [Run(element, paragraph) for element in paragraph._p.iter(qn("w:r")) if Run(element, paragraph).text]


def _fonts(run) -> dict[str, str]:
    properties = run._r.rPr
    fonts = properties.find(qn("w:rFonts")) if properties is not None else None
    return {name.rsplit("}", 1)[-1]: value for name, value in fonts.attrib.items()} if fonts is not None else {}


def _property(run, name: str) -> Any:
    properties = run._r.rPr
    return properties.find(qn(f"w:{name}")) if properties is not None else None


@pytest.mark.parametrize("marker_carrier", ["run", "split", "hyperlink"])
@pytest.mark.parametrize("code_size", [None, 9])
def test_authored_body_properties_reach_links_and_merged_body_per_attribute(
    tmp_path: Path, marker_carrier: str, code_size: int | None
) -> None:
    template: Any = Document()
    body_style = template.styles.add_style("User Body", WD_STYLE_TYPE.PARAGRAPH)
    body_style.base_style = template.styles["Normal"]
    body_style.font.name = "StyleLatin"
    body_style.font.size = Pt(12)
    hyperlink_style = template.styles.add_style("Hyperlink", WD_STYLE_TYPE.CHARACTER)
    hyperlink_style.font.color.rgb = RGBColor.from_string("0563C1")
    hyperlink_style.font.underline = True
    code_style = template.styles.add_style("行内代码", WD_STYLE_TYPE.CHARACTER)
    code_style.style_id = "DocWenInlineCode"
    code_style.font.name = "CodeLatin"
    code_style._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "CodeCJK")
    if code_size is not None:
        code_style.font.size = Pt(code_size)
    marker = template.add_paragraph(style=body_style)
    authored = marker.add_run("{{" if marker_carrier == "split" else "{{body}}")
    authored.font.name = "AuthoredLatin"
    authored.font.size = Pt(17)
    authored.font.bold = False
    authored.font.italic = True
    authored.font.underline = False
    authored.font.color.rgb = RGBColor.from_string("112233")
    authored._r.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "AuthoredCJK")
    complex_size = OxmlElement("w:szCs")
    complex_size.set(qn("w:val"), "38")
    authored._r.rPr.append(complex_size)
    if marker_carrier == "split":
        marker.add_run("body").bold = True
        marker.add_run("}}")
    elif marker_carrier == "hyperlink":
        hyperlink = OxmlElement("w:hyperlink")
        hyperlink.set(qn("r:id"), template.part.relate_to("https://example.com/marker", RT.HYPERLINK, is_external=True))
        hyperlink.append(authored._r)
        marker._p.append(hyperlink)
    paragraph = "Plain **strong** [link](https://example.com) [`code label`](https://example.com/code) `code`."
    output = _convert(
        tmp_path,
        template,
        body=f"# [Heading title](https://example.com/heading):\n{paragraph}\n\n{paragraph}\n",
    )
    mixed = next(p for p in output.paragraphs if p.style.style_id == "Heading1")
    ordinary = next(p for p in output.paragraphs if p.style.style_id == "UserBody")
    title_link = next(run for run in _runs(mixed) if run.text == "Heading title")
    assert _fonts(title_link).get("ascii") != "AuthoredLatin"
    assert title_link.font.size is None

    for body in (mixed, ordinary):
        runs = _runs(body)
        plain = next(run for run in runs if run.text == "Plain ")
        assert plain.font.name == "AuthoredLatin"
        assert plain.font.size == Pt(17)
        assert _fonts(plain)["eastAsia"] == "AuthoredCJK"
        assert _property(plain, "szCs").get(qn("w:val")) == "38"
        assert plain.bold is False and plain.italic is True and plain.underline is False
        assert next(run for run in runs if run.text == "strong").bold is True
        link = next(run for run in runs if run.text == "link")
        assert link.style.style_id == "Hyperlink"
        assert link.font.name == "AuthoredLatin" and link.font.size == Pt(17)
        assert _fonts(link)["eastAsia"] == "AuthoredCJK"
        assert _property(link, "color") is None  # Hyperlink's authored color wins.
        assert _property(link, "u") is None  # Hyperlink's underline wins.
        code = next(run for run in runs if run.text == "code")
        assert code.style.style_id == "DocWenInlineCode"
        assert code.font.name is None  # The code character style owns its fonts.
        assert code.font.size == (None if code_size is not None else Pt(17))
        code_link = next(run for run in runs if run.text == "code label")
        assert code_link.style.style_id == "Hyperlink"
        assert _fonts(code_link)["ascii"] == "CodeLatin"
        assert _fonts(code_link)["eastAsia"] == "CodeCJK"
        assert code_link.font.size == Pt(code_size or 17)


@pytest.mark.parametrize("source", ["fallback", "normal", "inherited", "document_defaults"])
def test_merged_body_uses_the_same_font_source_as_ordinary_body(tmp_path: Path, source: str) -> None:
    template: Any = None
    if source != "fallback":
        template = Document()
        style = template.styles["Normal"]
        if source == "normal":
            # Normal anchors render using the locale's managed body style.
            style.font.name = "UnusedNormalFont"
            style.font.size = Pt(10.5)
        else:
            base = template.styles.add_style("Body Base", WD_STYLE_TYPE.PARAGRAPH)
            base.base_style = style
            style = template.styles.add_style("Body Leaf", WD_STYLE_TYPE.PARAGRAPH)
            style.base_style = base
            if source == "inherited":
                base.font.name = "InheritedLatin"
                base.font.size = Pt(13.5)
                base._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "InheritedCJK")
            else:
                defaults = template.styles.element.find(f"{qn('w:docDefaults')}/{qn('w:rPrDefault')}/{qn('w:rPr')}")
                fonts = defaults.find(qn("w:rFonts"))
                fonts.attrib.clear()
                fonts.set(qn("w:asciiTheme"), "minorHAnsi")
                fonts.set(qn("w:hAnsiTheme"), "minorHAnsi")
                fonts.set(qn("w:eastAsiaTheme"), "minorEastAsia")
                defaults.find(qn("w:sz")).set(qn("w:val"), "28")
                defaults.find(qn("w:szCs")).set(qn("w:val"), "28")
        template.add_paragraph("{{body}}", style=style)
        assert extract_body_run_format(template) is None
    else:
        assert extract_body_run_format(resolve_template()) is None

    output = _convert(tmp_path, template, body="# Heading:\nBody.\n\nBody.\n")
    mixed = next(p for p in output.paragraphs if p.style.style_id == "Heading1")
    ordinary = next(p for p in output.paragraphs if p.style.style_id != "Heading1" and p.text == "Body.")
    mixed_body = next(run for run in _runs(mixed) if run.text == "Body.")
    plain = next(run for run in _runs(ordinary) if run.text == "Body.")
    assert plain.font.name is None and plain.font.size is None
    assert mixed_body.bold is False
    assert _property(mixed_body, "color").get(qn("w:val")) == "auto"
    if source in {"fallback", "normal"}:
        assert ordinary.style.style_id == "DocWenBodyParagraph"
        assert ordinary.style.font.size == mixed_body.font.size == Pt(16)
        assert mixed_body.font.name == ordinary.style.font.name == "Times New Roman"
        assert _fonts(mixed_body)["eastAsia"] == "仿宋_GB2312"
    elif source == "inherited":
        assert mixed_body.font.name == "InheritedLatin"
        assert mixed_body.font.size == Pt(13.5)
        assert _fonts(mixed_body)["eastAsia"] == "InheritedCJK"
    else:
        assert mixed_body.font.name is None  # Preserve the template's theme source.
        assert mixed_body.font.size == Pt(14)
        assert _fonts(mixed_body)["asciiTheme"] == "minorHAnsi"
        assert _fonts(mixed_body)["eastAsiaTheme"] == "minorEastAsia"


def test_style_toggles_and_character_theme_fonts_survive_heading_merge(tmp_path: Path) -> None:
    template: Any = Document()
    template.styles["Heading 1"].font.strike = True
    template.styles["Heading 1"].font.double_strike = True
    base = template.styles.add_style("Body Base", WD_STYLE_TYPE.PARAGRAPH)
    base.base_style = template.styles["Normal"]
    base.font.bold = True
    base.font.italic = True
    leaf = template.styles.add_style("Body Leaf", WD_STYLE_TYPE.PARAGRAPH)
    leaf.base_style = base
    leaf.font.bold = True  # Two style toggles yield False.
    leaf.font.italic = False  # Style False leaves the inherited True unchanged.
    leaf.font.size = Pt(15)
    template.add_paragraph("{{body}}", style=leaf)
    hyperlink = template.styles.add_style("Hyperlink", WD_STYLE_TYPE.CHARACTER)
    hyperlink.font.bold = True
    hyperlink.font.italic = False
    hyperlink._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:asciiTheme"), "majorHAnsi")

    paragraph = "Plain [link](https://example.com)."
    output = _convert(tmp_path, template, body=f"# Heading:\n{paragraph}\n\n{paragraph}\n")
    mixed = next(p for p in output.paragraphs if p.style.style_id == "Heading1")
    mixed_plain = next(run for run in _runs(mixed) if run.text == "Plain ")
    mixed_link = next(run for run in _runs(mixed) if run.text == "link")
    assert mixed_plain.bold is False and mixed_plain.italic is True
    assert mixed_plain.font.strike is False and mixed_plain.font.double_strike is False
    assert mixed_link.bold is True and mixed_link.italic is True
    assert mixed_link.font.strike is False and mixed_link.font.double_strike is False
    assert "ascii" not in _fonts(mixed_link) and "asciiTheme" not in _fonts(mixed_link)
    style_element = mixed_link.style._element
    assert style_element is not None and style_element.rPr is not None
    assert style_element.rPr.rFonts is not None
    assert style_element.rPr.rFonts.get(qn("w:asciiTheme")) == "majorHAnsi"
