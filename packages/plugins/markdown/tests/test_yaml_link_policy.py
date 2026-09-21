"""YAML front-matter ordinary-link parity for Markdown -> DOCX."""

from __future__ import annotations

from pathlib import Path

import pytest
from docx import Document

from ._link_processing_routes_support import _convert_docx, _link_config

pytestmark = pytest.mark.contract


def _gongwen_config(*, wiki_mode: str, markdown_mode: str) -> dict:
    config = _link_config(wiki_mode=wiki_mode, markdown_mode=markdown_mode)
    config["gui"] = {"language": {"locale": "zh_CN"}}
    config["field_processors"] = {
        "settings": {"order": ["gongwen"]},
        "processors": {
            "gongwen": {
                "module": "docwen_plugin_markdown.field_processors.gongwen",
                "locales": ["zh_CN"],
                "enabled": True,
            }
        },
    }
    return config


def _template(tmp_path: Path) -> Path:
    path = tmp_path / "yaml-links-template.docx"
    doc = Document()
    doc.add_paragraph("抄送：{{抄送机关}}")
    doc.add_paragraph("网站：{{site}}")
    doc.add_paragraph("{{正文}}")
    doc.save(str(path))
    return path


def _source(tmp_path: Path) -> Path:
    (tmp_path / "guide.md").write_text("target", encoding="utf-8")
    source = tmp_path / "yaml-links.md"
    source.write_text(
        "---\n"
        "抄送机关:\n"
        '  - "[[guide|喵喵]]"\n'
        "  - 钱钱钱\n"
        'site: "[项目主页](https://example.com/project)"\n'
        '份号: "001"\n'
        "enabled: false\n"
        "count: 0\n"
        "---\n\n"
        "正文。\n",
        encoding="utf-8",
    )
    return source


@pytest.mark.parametrize("mode", ["keep", "extract_text", "remove", "hyperlink"])
def test_yaml_links_follow_request_scoped_ordinary_link_policy(tmp_path: Path, mode: str) -> None:
    source = _source(tmp_path)
    template = _template(tmp_path)
    observation = _convert_docx(
        source,
        _gongwen_config(wiki_mode=mode, markdown_mode=mode),
        options={"template_name": str(template)},
    )

    if mode == "keep":
        assert "[[guide|喵喵]]" in observation.text
        assert "[项目主页](https://example.com/project)" in observation.text
        assert not observation.hyperlink_targets
    elif mode == "extract_text":
        assert "抄送：喵喵，钱钱钱" in observation.text
        assert "网站：项目主页" in observation.text
        assert "[[guide|喵喵]]" not in observation.text
        assert "[项目主页](" not in observation.text
        assert not observation.hyperlink_targets
    elif mode == "remove":
        assert "抄送：钱钱钱" in observation.text
        assert "喵喵" not in observation.text
        assert "项目主页" not in observation.text
        assert not observation.hyperlink_targets
    else:
        assert "抄送：喵喵，钱钱钱" in observation.text
        assert "网站：项目主页" in observation.text
        assert "https://example.com/project" in observation.hyperlink_targets
        assert any(
            target.startswith("file:") and target.endswith("/guide.md") for target in observation.hyperlink_targets
        )
        assert "<w:hyperlink" in observation.document_xml


@pytest.mark.parametrize(
    ("wiki_mode", "markdown_mode", "wiki_text", "markdown_text"),
    [
        ("extract_text", "keep", "抄送：喵喵，钱钱钱", "[项目主页](https://example.com/project)"),
        ("keep", "extract_text", "[[guide|喵喵]]", "网站：项目主页"),
    ],
)
def test_yaml_wiki_and_markdown_link_modes_are_independent(
    tmp_path: Path,
    wiki_mode: str,
    markdown_mode: str,
    wiki_text: str,
    markdown_text: str,
) -> None:
    observation = _convert_docx(
        _source(tmp_path),
        _gongwen_config(wiki_mode=wiki_mode, markdown_mode=markdown_mode),
        options={"template_name": str(_template(tmp_path))},
    )

    assert wiki_text in observation.text
    assert markdown_text in observation.text


@pytest.mark.parametrize("wiki_mode", ["keep", "extract_text", "remove", "hyperlink"])
@pytest.mark.parametrize("markdown_mode", ["keep", "extract_text", "remove", "hyperlink"])
def test_all_independent_modes_do_not_reactivate_kept_links(tmp_path: Path, wiki_mode: str, markdown_mode: str) -> None:
    observation = _convert_docx(
        _source(tmp_path),
        _gongwen_config(wiki_mode=wiki_mode, markdown_mode=markdown_mode),
        options={"template_name": str(_template(tmp_path))},
    )
    assert ("[项目主页](https://example.com/project)" in observation.text) == (markdown_mode == "keep")
    assert ("[[guide|喵喵]]" in observation.text) == (wiki_mode == "keep")
    assert ("https://example.com/project" in observation.hyperlink_targets) == (markdown_mode == "hyperlink")
    assert any(target.startswith("file:") for target in observation.hyperlink_targets) == (wiki_mode == "hyperlink")


def test_split_placeholders_and_special_fields_preserve_template_runs(tmp_path: Path) -> None:
    source = _source(tmp_path)
    source.write_text(
        source.read_text(encoding="utf-8").replace(
            '份号: "001"', '附件说明:\n  - "[附件](https://example.com/attachment)"\n份号: "001"'
        ),
        encoding="utf-8",
    )
    template = Document()
    p = template.add_paragraph()
    r = p.add_run("[模板原文](https://example.com/literal) {{si")
    r.bold = True
    r.add_break()
    # A separate valid split marker must not stitch across the structural break.
    p.add_run("te}} / {{si").italic = True
    p.add_run("te}} 后缀").italic = True
    p = template.add_paragraph()
    r = p.add_run("{{site}}")
    r.bold = True
    r.add_break()
    r.add_text("{{site}}")
    template.add_paragraph("{{附件说明}}")
    template.add_paragraph("{{份号}}/{{enabled}}/{{count}}")
    template.add_paragraph("{{正文}}")
    path = tmp_path / "complex.docx"
    template.save(str(path))
    observation = _convert_docx(
        source,
        _gongwen_config(wiki_mode="keep", markdown_mode="hyperlink"),
        options={"template_name": str(path)},
    )
    assert "[模板原文](https://example.com/literal)" in observation.text
    assert "https://example.com/literal" not in observation.hyperlink_targets
    assert "https://example.com/attachment" in observation.hyperlink_targets
    assert "https://example.com/project" in observation.hyperlink_targets
    assert " 后缀" in observation.text
    assert "001/False/0" in observation.text
    assert observation.document_xml.count("<w:br") == 2
    assert "DOCWENYAMLLINK" not in observation.document_xml
    assert "<w:b" in observation.document_xml and "<w:i" in observation.document_xml


def test_unused_yaml_links_never_probe_files(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    from docwen_plugin_markdown.yaml_links import YamlLinkProjection

    source = tmp_path / "unused.md"
    source.write_text('---\nunused: "[[missing]]"\ntitle: Plain title\n---\nBody', encoding="utf-8")
    template = Document()
    template.add_paragraph("{{title}}")
    template.add_paragraph("{{body}}")
    path = tmp_path / "unused.docx"
    template.save(str(path))
    original = YamlLinkProjection.project

    def checked(self, value):
        assert value != "[[missing]]"
        return original(self, value)

    monkeypatch.setattr(YamlLinkProjection, "project", checked)
    observation = _convert_docx(source, _link_config(), options={"template_name": str(path)})
    assert "Plain title" in observation.text


def test_declared_yaml_links_reject_local_lookup_before_resolution() -> None:
    from docwen_core.export_semantics import LinkRuntimeConfig
    from docwen_core.links import DeclaredResourceError
    from docwen_plugin_markdown.yaml_links import YamlLinkProjection

    projection = YamlLinkProjection("source.md", LinkRuntimeConfig(), declared_inputs=True)
    with pytest.raises(DeclaredResourceError):
        projection.project("[[local|alias]]")
