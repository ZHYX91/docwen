"""Regression coverage for request-owned DOCX foundation-style completion."""

from __future__ import annotations

import hashlib
from io import BytesIO
from zipfile import ZipFile

import lxml.etree as etree
import pytest
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from docwen_plugin_markdown.to_docx.managed_styles import (
    ManagedStyleCompletionError,
    complete_managed_styles,
    validate_managed_style_package,
)
from docwen_runtime.config.document_styles import build_document_style_catalog

from .conftest import PROJECT_ROOT

pytestmark = pytest.mark.unit


def _catalog(locale="zh_CN"):
    return build_document_style_catalog(
        {"gui": {"language": {"locale": locale}}},
        locales_dir=PROJECT_ROOT / "i18n" / "locales",
    )


def _remove_styles(document, *style_ids: str) -> None:
    root = document.styles.element
    wanted = set(style_ids)
    for element in list(root.findall(qn("w:style"))):
        if element.get(qn("w:styleId"), "") in wanted:
            root.remove(element)


def _style_ids(document) -> set[str]:
    return {style.style_id for style in document.styles}


@pytest.mark.unit
def test_missing_foundations_are_completed_but_title_is_not_required_or_injected() -> None:
    document = Document()
    _remove_styles(document, "Normal", "DefaultParagraphFont", "TableNormal", "Title", "Subtitle")

    completed, bindings = complete_managed_styles(document, _catalog())
    style_ids = _style_ids(completed)

    assert {"Normal", "DefaultParagraphFont", "TableNormal"}.issubset(style_ids)
    assert "Title" not in style_ids
    assert "Subtitle" not in style_ids
    assert len(bindings.styles) == 43


@pytest.mark.unit
def test_existing_title_is_preserved_without_becoming_a_managed_dependency() -> None:
    document = Document()
    title = document.styles["Title"]
    original_id = title.style_id

    completed, _bindings = complete_managed_styles(document, _catalog())

    assert completed.styles[original_id].style_id == original_id


def _package_snapshot(document) -> tuple[tuple[str, str], ...]:
    snapshot: list[tuple[str, str]] = []
    for part in document.part.package.parts:
        try:
            payload = etree.tostring(etree.fromstring(part.blob), method="c14n", exclusive=True)
        except etree.XMLSyntaxError:
            payload = part.blob
        snapshot.append((str(part.partname), hashlib.sha256(payload).hexdigest()))
    return tuple(sorted(snapshot))


def _document_bytes(document) -> bytes:
    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _rewrite_xml_member(blob: bytes, member_name: str, transform) -> bytes:
    source = BytesIO(blob)
    target = BytesIO()
    with ZipFile(source, "r") as archive_in, ZipFile(target, "w", allowZip64=True) as archive_out:
        for item in archive_in.infolist():
            payload = archive_in.read(item.filename)
            if item.filename == member_name:
                root = etree.fromstring(payload)
                transform(root)
                payload = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
            archive_out.writestr(item, payload)
    return target.getvalue()


@pytest.mark.parametrize("part_name", ("word/styles.xml", "word/stylesWithEffects.xml"))
@pytest.mark.parametrize(
    ("foundation_id", "kind"),
    (("Normal", "paragraph"), ("DefaultParagraphFont", "character"), ("TableNormal", "table")),
)
@pytest.mark.parametrize("default_value", ("1", "true", "on"))
def test_missing_foundation_does_not_replace_a_template_owned_default(
    part_name: str, foundation_id: str, kind: str, default_value: str
) -> None:
    def customize(root: etree._Element) -> None:
        foundation = next(item for item in root.findall(qn("w:style")) if item.get(qn("w:styleId")) == foundation_id)
        root.remove(foundation)
        custom = OxmlElement("w:style")
        custom.set(qn("w:type"), kind)
        custom.set(qn("w:styleId"), "TemplateDefault")
        custom.set(qn("w:default"), default_value)
        name = OxmlElement("w:name")
        name.set(qn("w:val"), "Template-owned default")
        custom.append(name)
        custom.append(OxmlElement("w:locked"))
        root.append(custom)

    document = Document(BytesIO(_rewrite_xml_member(_document_bytes(Document()), part_name, customize)))
    before = _package_snapshot(document)
    completed, bindings = complete_managed_styles(document, _catalog("en_US"))
    assert _package_snapshot(document) == before
    blob = _document_bytes(completed)
    validate_managed_style_package(blob, _catalog("en_US"), bindings)
    with ZipFile(BytesIO(blob)) as archive:
        root = etree.fromstring(archive.read(part_name))
    by_id = {item.get(qn("w:styleId")): item for item in root.findall(qn("w:style"))}
    assert by_id[foundation_id].get(qn("w:default")) is None
    assert by_id["TemplateDefault"].get(qn("w:default")) == default_value
    assert by_id["TemplateDefault"].find(qn("w:locked")) is not None
    assert [
        item.get(qn("w:styleId"))
        for item in root.findall(qn("w:style"))
        if item.get(qn("w:type")) == kind and item.get(qn("w:default")) in {"1", "true", "on"}
    ] == ["TemplateDefault"]


@pytest.mark.parametrize("part_name", ("word/styles.xml", "word/stylesWithEffects.xml"))
@pytest.mark.parametrize("foundation_id", ("Normal", "DefaultParagraphFont", "TableNormal"))
def test_final_package_rejects_damaged_foundation_identity(part_name: str, foundation_id: str) -> None:
    completed, bindings = complete_managed_styles(Document(), _catalog("en_US"))

    def damage(root: etree._Element) -> None:
        foundation = next(item for item in root.findall(qn("w:style")) if item.get(qn("w:styleId")) == foundation_id)
        foundation.set(qn("w:type"), "numbering")

    damaged = _rewrite_xml_member(_document_bytes(completed), part_name, damage)
    with pytest.raises(ManagedStyleCompletionError, match="style") as error:
        validate_managed_style_package(damaged, _catalog("en_US"), bindings)
    assert error.value.error_type == "conversion_failed"
