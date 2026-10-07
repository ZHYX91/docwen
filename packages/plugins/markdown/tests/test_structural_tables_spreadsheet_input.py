"""Structural Tables spreadsheet-input regression coverage."""

from __future__ import annotations

import csv

import pytest
from openpyxl import load_workbook

from ._md_to_spreadsheet_support import MdToCsvConverter, MdToXlsxConverter, Path, make_context, write_temp_md

pytestmark = pytest.mark.contract


@pytest.mark.parametrize(
    ("source", "expected"),
    [
        ("| - | - |\n| Alice | 10 |\n| Bob | 20 |", [["Alice", "10"], ["Bob", "20"]]),
        (
            "| Region | Sales |\n| Quarter | Q1 |\n| --- || --- |\n| North | 10 |",
            [["Region", "Sales"], ["Quarter", "Q1"], ["North", "10"]],
        ),
    ],
)
def test_csv_preserves_all_structural_content_rows(source: str, expected: list[list[str]]) -> None:
    context, _workspace = make_context(
        write_temp_md(source),
        target_format="csv",
        options={"markdown_extensions": {"input": {"structural_tables": True}}},
    )
    result = MdToCsvConverter().convert(context)
    assert result.success, result.error
    with Path(result.artifacts[0].staging_path).open(encoding="utf-8", newline="") as handle:
        assert list(csv.reader(handle)) == expected


def test_csv_disabled_structural_tables_keeps_headerless_input_unrecognized() -> None:
    context, _workspace = make_context(
        write_temp_md("| - | - |\n| Alice | 10 |"),
        target_format="csv",
        options={"markdown_extensions": {"input": {"structural_tables": False}}},
    )
    result = MdToCsvConverter().convert(context)
    assert not result.success


@pytest.mark.parametrize(
    ("cell", "expected"),
    [("`C:\\`", "`C:\\`"), ("`unclosed", "`unclosed"), ("``a`|b``", "``a`|b``")],
)
def test_spreadsheet_table_code_delimiters_preserve_column_boundaries(cell: str, expected: str) -> None:
    from docwen_plugin_markdown.common_utils import parse_md_tables

    tables = parse_md_tables(f"| --- | --- |\n| {cell} | right |", structural_tables=True)
    assert tables[0]["all_rows"] == [[expected, "right"]]


def _convert_xlsx(tmp_path: Path, source: str):
    md_path = write_temp_md(source)
    context, _workspace = make_context(
        md_path,
        target_format="xlsx",
        options={"markdown_extensions": {"input": {"structural_tables": True}}},
    )
    result = MdToXlsxConverter().convert(context)
    assert result.success, result.error
    workbook = load_workbook(Path(result.artifacts[0].staging_path))
    return result, workbook


def test_xlsx_preserves_zero_header_structural_table_rows(tmp_path: Path) -> None:
    _result, workbook = _convert_xlsx(
        tmp_path,
        "| - | - |\n| Alice | 10 |\n| Bob | 20 |",
    )

    sheet = workbook.active
    assert sheet is not None
    assert sheet["A1"].value == "Alice"
    assert sheet["B1"].value == "10"
    assert sheet["A2"].value == "Bob"
    assert sheet["B2"].value == "20"
    workbook.close()


def test_xlsx_preserves_multi_header_row_header_and_merge_geometry(tmp_path: Path) -> None:
    source = "| Region | Sales | < |\n| Quarter | Q1 | Q2 |\n| --- || --- | --- |\n| North | 10 | 12 |\n| ^ | 8 | 11 |"
    result, workbook = _convert_xlsx(tmp_path, source)

    sheet = workbook.active
    assert sheet is not None
    assert [sheet.cell(row=2, column=column).value for column in range(1, 4)] == ["Quarter", "Q1", "Q2"]
    assert [sheet.cell(row=3, column=column).value for column in range(1, 4)] == ["North", "10", "12"]
    assert {str(merged) for merged in sheet.merged_cells.ranges} == {"B1:C1", "A3:A4"}
    assert result.artifacts[0].metadata["table_count"] == 1
    workbook.close()


def test_xlsx_preserves_escaped_literal_merge_markers(tmp_path: Path) -> None:
    source = "| Group | \\< |\n| Name | Value |\n| --- | --- |\n| A | \\^ |"
    _result, workbook = _convert_xlsx(tmp_path, source)

    sheet = workbook.active
    assert sheet is not None
    assert sheet["B1"].value == "<"
    assert sheet["B3"].value == "^"
    assert not sheet.merged_cells.ranges
    workbook.close()


def test_merge_planner_accepts_equivalent_rectangles_and_empty_anchor() -> None:
    from docwen_plugin_markdown.to_spreadsheet.template_xlsx import _plan_rectangular_merges

    canonical, canonical_warnings = _plan_rectangular_merges([["A", "<"], ["^", "^"]])
    equivalent, equivalent_warnings = _plan_rectangular_merges([["A", "<"], ["^", "<"]])
    empty_anchor, empty_warnings = _plan_rectangular_merges([["", "<"], ["^", "<"]])

    assert canonical == equivalent == empty_anchor == [(0, 0, 1, 1)]
    assert canonical_warnings == equivalent_warnings == empty_warnings == 0
