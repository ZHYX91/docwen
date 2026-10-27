"""Internal image markers never consume the final XLSX text capacity."""

from __future__ import annotations

from pathlib import Path
from urllib.parse import quote
from zipfile import ZipFile

import pytest
from openpyxl import Workbook, load_workbook

from docwen_plugin_markdown.to_spreadsheet.converter import MdToXlsxConverter

from ._md_to_spreadsheet_support import _png_bytes
from .conftest import make_context

pytestmark = pytest.mark.contract


def _drawing_count(sheet) -> int:
    drawings = vars(sheet).get("_images")
    assert isinstance(drawings, list)
    return len(drawings)


@pytest.mark.parametrize("route", ["template", "table", "column", "row"])
@pytest.mark.parametrize("filename", ["p.png", "图 像%{name}.png"])
def test_image_scope_and_path_overhead_never_enter_xlsx_text(
    tmp_path: Path, route: str, filename: str, monkeypatch
) -> None:
    from openpyxl.cell.cell import Cell

    original_check = Cell.check_string

    def bounded_string(cell, value):
        if isinstance(value, str):
            assert len(value.encode("utf-16-le")) // 2 <= 32767
        return original_check(cell, value)

    monkeypatch.setattr(Cell, "check_string", bounded_string)
    payload = _png_bytes()
    image = tmp_path / filename
    image.write_bytes(payload)
    text = "x" * 32730
    source = tmp_path / "source.md"
    options = {}
    if route == "template":
        value = text + "{{IMAGE:" + filename + "}}"
        markdown = "| Value |\n| --- |\n| control |\n"
    else:
        value = "{{↓Value}}" if route == "column" else "{{→Value}}"
        markdown = f"| Value |\n| --- |\n| {text}![photo]({quote(filename)}) |\n"
    if route != "table":
        template = tmp_path / "template.xlsx"
        workbook = Workbook()
        sheet = workbook.active
        assert sheet is not None
        sheet["A1"] = value
        workbook.save(template)
        workbook.close()
        options["template_name"] = str(template)
    source.write_text(markdown, encoding="utf-8")
    context, _workspace = make_context(str(source), target_format="xlsx", options=options)
    result = MdToXlsxConverter().convert(context)
    assert result.success, result.error
    path = Path(result.artifacts[0].staging_path)
    output = load_workbook(path)
    try:
        sheet = output.active
        assert sheet is not None
        assert sheet["A2" if route == "table" else "A1"].value == text
        assert _drawing_count(sheet) == 1
    finally:
        output.close()
    with ZipFile(path) as archive:
        images = [name for name in archive.namelist() if name.startswith("xl/media/")]
        assert len(images) == 1 and archive.read(images[0]) == payload
    assert image.read_bytes() == payload


@pytest.mark.parametrize("units", [32767, 32768])
@pytest.mark.parametrize("nonbmp", [False, True], ids=["ascii", "nonbmp"])
def test_final_image_cell_capacity_is_checked_after_removing_internal_markers(
    tmp_path: Path, units: int, nonbmp: bool
) -> None:
    text = ("😀" * (units // 2) + "x" * (units % 2)) if nonbmp else "x" * units
    (tmp_path / "p.png").write_bytes(_png_bytes())
    template = tmp_path / "template.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet["A1"] = "{{label}}{{IMAGE:p.png}}"
    workbook.save(template)
    workbook.close()
    source = tmp_path / "source.md"
    source.write_text(f'---\nlabel: "{text}"\n---\n', encoding="utf-8")
    context, workspace = make_context(str(source), target_format="xlsx", options={"template_name": str(template)})
    result = MdToXlsxConverter().convert(context)
    if units > 32767:
        assert not result.success and not result.artifacts
        assert result.error is not None and result.error.diagnostic_code == "MD2XLSX-CELL-TEXT-TOO-LONG"
        assert not list(Path(workspace.staging_dir).glob("*.xlsx"))
        return
    assert result.success, result.error
    output = load_workbook(result.artifacts[0].staging_path)
    try:
        sheet = output.active
        assert sheet is not None and sheet["A1"].value == text and _drawing_count(sheet) == 1
    finally:
        output.close()


def test_authored_markers_never_gain_image_trust_even_when_a_scope_would_collide(tmp_path: Path, monkeypatch) -> None:
    from docwen_plugin_markdown.to_spreadsheet import converter, template_xlsx

    scope = "s" * 32
    monkeypatch.setattr(converter.secrets, "token_urlsafe", lambda _size: scope)
    payload = _png_bytes()
    (tmp_path / "p.png").write_bytes(payload)
    (tmp_path / "secret.png").write_bytes(payload)
    observed = []
    original_check = template_xlsx._is_image_content

    def observe(path):
        observed.append(path.name)
        return original_check(path)

    monkeypatch.setattr(template_xlsx, "_is_image_content", observe)
    template = tmp_path / "template.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet["A1"] = "{{label}}"
    sheet["A2"] = "{{↓Value}}"
    literal = "{{IMAGE@" + scope + "_:secret.png}}"
    sheet["A3"] = literal
    sheet["A4"] = "{{IMAGE:p.png}}"
    workbook.save(template)
    workbook.close()
    raw = "{{IMAGE:secret.png}} {{IMAGE@" + scope + ":secret.png}}"
    source = tmp_path / "source.md"
    source.write_text(
        '---\nlabel: "{{IMAGE:secret.png}}"\n---\n\n| Value |\n| --- |\n| ' + raw + " |\n", encoding="utf-8"
    )
    context, _workspace = make_context(str(source), target_format="xlsx", options={"template_name": str(template)})
    result = MdToXlsxConverter().convert(context)
    assert result.success, result.error
    output = load_workbook(result.artifacts[0].staging_path)
    try:
        sheet = output.active
        assert sheet is not None
        assert sheet["A1"].value == "{{IMAGE:secret.png}}"
        assert sheet["A2"].value == raw and sheet["A3"].value == literal
        assert sheet["A4"].value is None and _drawing_count(sheet) == 1
        assert observed == ["p.png"]
    finally:
        output.close()


def test_template_drafts_preserve_placeholder_order_overwrites_and_protected_images(tmp_path: Path) -> None:
    from openpyxl.styles import Protection

    (tmp_path / "p.png").write_bytes(_png_bytes())
    template = tmp_path / "template.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet["A1"] = "{{↓Name}}"
    sheet["A2"] = "{{IMAGE:p.png}}"
    sheet["A3"] = "{{IMAGE:p.png}}"
    sheet["B1"] = "{{↓Name}}{{IMAGE:p.png}}"
    sheet["C1"] = "=1+2"
    sheet["E1"] = "{{→Name}}"
    sheet["F1"] = "{{IMAGE:p.png}}"
    sheet["G1"] = "{{IMAGE:p.png}}"
    sheet["XFD1048576"] = "far sentinel"
    sheet.protection.sheet = True
    for cell in ["A1", "A3", "E1", "F1"]:
        sheet[cell].protection = Protection(locked=False)
    workbook.save(template)
    workbook.close()
    source = tmp_path / "source.md"
    source.write_text("| Name |\n| --- |\n| first |\n| second |\n| third |\n", encoding="utf-8")
    context, _workspace = make_context(str(source), target_format="xlsx", options={"template_name": str(template)})
    result = MdToXlsxConverter().convert(context)
    assert result.success, result.error
    output = load_workbook(result.artifacts[0].staging_path)
    try:
        sheet = output.active
        assert sheet is not None
        assert sheet["A1"].value == sheet["E1"].value == "first"
        assert sheet["A3"].value == "third" and sheet["F1"].value == "second"
        assert sheet["A2"].value is sheet["G1"].value is sheet["B1"].value is None
        assert _drawing_count(sheet) == 3
        assert sheet["C1"].value == "=1+2" and sheet["C1"].data_type == "f"
        assert sheet["XFD1048576"].value == "far sentinel" and len(vars(sheet).get("_cells", {})) == 9
    finally:
        output.close()


def test_duplicate_template_markers_keep_each_drawing_and_csv_drops_internal_transport(tmp_path: Path) -> None:
    from docwen_plugin_markdown.to_spreadsheet.converter import MdToCsvConverter

    (tmp_path / "p.png").write_bytes(_png_bytes())
    template = tmp_path / "template.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet["A1"] = "before{{IMAGE:p.png}}{{IMAGE:p.png}}after"
    workbook.save(template)
    workbook.close()
    source = tmp_path / "source.md"
    source.write_text("| Value |\n| --- |\n| control |\n", encoding="utf-8")
    for kind, converter in [("xlsx", MdToXlsxConverter()), ("csv", MdToCsvConverter())]:
        context, _workspace = make_context(str(source), target_format=kind, options={"template_name": str(template)})
        result = converter.convert(context)
        assert result.success, result.error
        if kind == "csv":
            assert Path(result.artifacts[0].staging_path).read_text(encoding="utf-8-sig").strip() == "beforeafter"
        else:
            output = load_workbook(result.artifacts[0].staging_path)
            try:
                sheet = output.active
                assert sheet is not None and sheet["A1"].value == "beforeafter" and _drawing_count(sheet) == 2
            finally:
                output.close()


def test_yaml_decoding_cannot_turn_authored_marker_text_into_an_image(tmp_path: Path, monkeypatch) -> None:
    from docwen_plugin_markdown.to_spreadsheet import converter

    scope = "s" * 32
    monkeypatch.setattr(converter.secrets, "token_urlsafe", lambda _size: scope)
    (tmp_path / "secret.png").write_bytes(_png_bytes())
    template = tmp_path / "template.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet["A1"] = "{{label}}"
    workbook.save(template)
    workbook.close()
    literal = "{{IMAGE@" + scope + ":secret.png}}"
    escaped = literal.replace("{", r"\u007b").replace("s", r"\u0073")
    source = tmp_path / "source.md"
    source.write_text('---\nlabel: "' + escaped + '"\n---\n', encoding="utf-8")
    context, _workspace = make_context(str(source), target_format="xlsx", options={"template_name": str(template)})
    result = MdToXlsxConverter().convert(context)
    assert result.success, result.error
    output = load_workbook(result.artifacts[0].staging_path)
    try:
        sheet = output.active
        assert sheet is not None and sheet["A1"].value == literal and _drawing_count(sheet) == 0
    finally:
        output.close()


def test_yaml_fields_cannot_assemble_an_image_capability_around_a_trusted_slot(tmp_path: Path, monkeypatch) -> None:
    from docwen_plugin_markdown.to_spreadsheet import converter

    scope = "s" * 32
    monkeypatch.setattr(converter.secrets, "token_urlsafe", lambda _size: scope)
    (tmp_path / "secret.png").write_bytes(_png_bytes())
    (tmp_path / "p.png").write_bytes(_png_bytes())
    template = tmp_path / "template.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet["A1"] = "{{a}}{{b}}{{IMAGE:p.png}}"
    workbook.save(template)
    workbook.close()
    literal = "{{IMAGE@" + scope + ":secret.png}}"
    source = tmp_path / "source.md"
    escaped_scope = scope.replace("s", r"\u0073")
    source.write_text('---\na: "{{IMAGE@"\nb: "' + escaped_scope + ':secret.png}}"\n---\n', encoding="utf-8")
    context, _workspace = make_context(str(source), target_format="xlsx", options={"template_name": str(template)})
    result = MdToXlsxConverter().convert(context)
    assert result.success, result.error
    output = load_workbook(result.artifacts[0].staging_path)
    try:
        sheet = output.active
        assert sheet is not None and sheet["A1"].value == literal and _drawing_count(sheet) == 1
    finally:
        output.close()


def test_table_html_decoding_preserves_an_authored_image_marker(tmp_path: Path, monkeypatch) -> None:
    from docwen_plugin_markdown.to_spreadsheet import converter

    scope = "s" * 32
    monkeypatch.setattr(converter.secrets, "token_urlsafe", lambda _size: scope)
    (tmp_path / "secret.png").write_bytes(_png_bytes())
    template = tmp_path / "template.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet["A1"] = "{{↓Value}}"
    workbook.save(template)
    workbook.close()
    literal = "{{IMAGE@" + scope + ":secret.png}}"
    encoded = literal.replace("{", "&#123;").replace("s", "&#115;")
    source = tmp_path / "source.md"
    source.write_text("| Value |\n| --- |\n| " + encoded + " |\n", encoding="utf-8")
    context, _workspace = make_context(str(source), target_format="xlsx", options={"template_name": str(template)})
    result = MdToXlsxConverter().convert(context)
    assert result.success, result.error
    output = load_workbook(result.artifacts[0].staging_path)
    try:
        sheet = output.active
        assert sheet is not None and sheet["A1"].value == literal and _drawing_count(sheet) == 0
    finally:
        output.close()
