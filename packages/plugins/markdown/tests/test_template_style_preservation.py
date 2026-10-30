"""Template-owned style defects stay local; generated formatting remains valid."""

from copy import deepcopy
from io import BytesIO
from pathlib import Path
from zipfile import ZipFile

import pytest
from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.opc.constants import CONTENT_TYPE as CT
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.opc.packuri import PackURI
from docx.opc.part import Part
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.styles.style import ParagraphStyle
from lxml import etree

from docwen_plugin_markdown.to_docx.converter import MdToDocxConverter
from docwen_plugin_markdown.to_docx.managed_styles import (
    ManagedStyleCompletionError,
    complete_managed_styles,
    validate_managed_style_package,
)

from .conftest import make_context
from .test_managed_styles import _attach_xml_part, _catalog, _document_bytes

pytestmark = pytest.mark.contract


def _bad_run(document, style_id: str = "MissingTemplateStyle"):
    run = document.paragraphs[0].add_run("Preserved old text")
    reference = OxmlElement("w:rStyle")
    reference.set(qn("w:val"), style_id)
    run._r.get_or_add_rPr().append(reference)
    return run._r


@pytest.mark.parametrize(
    "defect", ["missing_reference", "wrong_type_reference", "unused_duplicate", "wrong_type_note_name"]
)
@pytest.mark.parametrize("with_notes", [False, True])
def test_converter_preserves_unused_style_defects_without_blocking_plain_body(
    tmp_path: Path, defect: str, with_notes: bool
) -> None:
    document = Document()
    document.add_paragraph("Template literal")
    document.add_paragraph("{{正文}}")
    if defect == "missing_reference":
        original = _bad_run(document)
    elif defect == "wrong_type_reference":
        original = _bad_run(document, "Normal")
    elif defect == "unused_duplicate":
        style = document.styles.add_style("Unused template style", WD_STYLE_TYPE.CHARACTER)
        original = style._element
        assert original is not None
        original.set(qn("w:styleId"), "UnusedDuplicate")
        document.styles.element.append(deepcopy(original))
    else:
        style = document.styles.add_style("footnote text", WD_STYLE_TYPE.CHARACTER)
        original = style._element
        assert original is not None
        original.set(qn("w:styleId"), "UserFootnoteName")
    original_xml = etree.tostring(original, method="c14n", exclusive=True)
    template = tmp_path / "template.docx"
    document.save(str(template))
    source = tmp_path / "source.md"
    source.write_text(
        "New valid body.\n" + ("Footnote[^n].\n\n[^n]: New note body.\n" if with_notes else ""), encoding="utf-8"
    )
    context, _workspace = make_context(str(source), target_format="docx", options={"template_name": str(template)})
    result = MdToDocxConverter().convert(context)
    assert result.success, result.error
    with ZipFile(result.artifacts[0].staging_path) as archive:
        root = etree.fromstring(archive.read("word/document.xml" if "reference" in defect else "word/styles.xml"))
        assert any(
            etree.tostring(node, method="c14n", exclusive=True) == original_xml for node in root.iter(original.tag)
        )
        if with_notes:
            notes = etree.fromstring(archive.read("word/footnotes.xml"))
            styles = etree.fromstring(archive.read("word/styles.xml"))
            normal = next(note for note in notes if note.get(qn("w:type"), "normal") == "normal")
            style_ref = normal.find(".//" + qn("w:pStyle"))
            assert style_ref is not None
            style_id = style_ref.get(qn("w:val"))
            candidates = styles.findall(qn("w:style") + "[@" + qn("w:styleId") + f"='{style_id}']")
            assert len(candidates) == 1 and candidates[0].get(qn("w:type")) == "paragraph"
        if defect == "unused_duplicate":
            assert len(root.findall(qn("w:style") + "[@" + qn("w:styleId") + "='UnusedDuplicate']")) == 2
    assert "New valid body." in "\n".join(p.text for p in Document(result.artifacts[0].staging_path).paragraphs)


@pytest.mark.parametrize("same_owner", [False, True])
def test_new_invalid_same_id_reference_cannot_use_old_occurrence_allowance(same_owner: bool) -> None:
    document = Document()
    document.add_paragraph()
    old = _bad_run(document)
    completed, bindings = complete_managed_styles(document, _catalog("en_US"))
    validate_managed_style_package(_document_bytes(completed), _catalog("en_US"), bindings)
    if same_owner:
        completed.paragraphs[0]._p.append(deepcopy(old))
    else:
        completed.add_paragraph()
        new = _bad_run(completed)
        text = new.find(qn("w:t"))
        assert text is not None
        text.text = "New invalid text"
    with pytest.raises(ManagedStyleCompletionError, match="missing or ambiguous style"):
        validate_managed_style_package(_document_bytes(completed), _catalog("en_US"), bindings)


def test_old_invalid_run_survives_unrelated_sibling_insertion() -> None:
    document = Document()
    document.add_paragraph()
    _bad_run(document)
    completed, bindings = complete_managed_styles(document, _catalog("en_US"))
    completed.paragraphs[0].add_run("A newly inserted valid sibling")
    validate_managed_style_package(_document_bytes(completed), _catalog("en_US"), bindings)


def test_duplicate_managed_id_is_never_reused_for_new_heading() -> None:
    document = Document()
    heading = document.styles["Heading 1"]._element
    assert heading is not None
    document.styles.element.append(deepcopy(heading))
    completed, bindings = complete_managed_styles(document, _catalog("en_US"))
    assert bindings.style_id("heading_1") == "Heading1DocWen1"
    heading_style = bindings.get("heading_1")
    assert isinstance(heading_style, ParagraphStyle)
    completed.add_paragraph("Generated heading", style=heading_style)
    validate_managed_style_package(_document_bytes(completed), _catalog("en_US"), bindings)
    with ZipFile(BytesIO(_document_bytes(completed))) as archive:
        styles = etree.fromstring(archive.read("word/styles.xml"))
        assert len(styles.findall(qn("w:style") + "[@" + qn("w:styleId") + "='Heading1']")) == 2


@pytest.mark.parametrize("tag", ["pStyle", "rStyle", "tblStyle"])
def test_new_direct_reference_has_no_template_exemption(tag: str) -> None:
    completed, bindings = complete_managed_styles(Document(), _catalog("en_US"))
    reference = OxmlElement("w:" + tag)
    reference.set(qn("w:val"), "MissingNewStyle")
    if tag == "pStyle":
        completed.add_paragraph("new")._p.get_or_add_pPr().append(reference)
    elif tag == "rStyle":
        completed.add_paragraph().add_run("new")._r.get_or_add_rPr().append(reference)
    else:
        completed.add_table(rows=1, cols=1)._tbl.tblPr.append(reference)
    with pytest.raises(ManagedStyleCompletionError, match="missing or ambiguous style"):
        validate_managed_style_package(_document_bytes(completed), _catalog("en_US"), bindings)


@pytest.mark.parametrize("kind", ["paragraph", "table"])
def test_removed_invalid_owner_cannot_exempt_new_owner_with_same_properties(kind: str) -> None:
    document = Document()
    if kind == "paragraph":
        old = document.add_paragraph("Old template text")
        properties = old._p.get_or_add_pPr()
        ref = OxmlElement("w:pStyle")
    else:
        old = document.add_table(rows=1, cols=1)
        old.cell(0, 0).text = "Old template table"
        properties = old._tbl.tblPr
        ref = OxmlElement("w:tblStyle")
    ref.set(qn("w:val"), "MissingTemplateStyle")
    properties.append(ref)
    completed, bindings = complete_managed_styles(document, _catalog("en_US"))
    if kind == "paragraph":
        old_element = completed.paragraphs[0]._p
        parent = old_element.getparent()
        assert parent is not None
        parent.remove(old_element)
        replacement = completed.add_paragraph("New generated text")
        replacement._p.get_or_add_pPr().append(deepcopy(ref))
    else:
        old_element = completed.tables[0]._tbl
        parent = old_element.getparent()
        assert parent is not None
        parent.remove(old_element)
        replacement = completed.add_table(rows=1, cols=1)
        replacement.cell(0, 0).text = "New generated table"
        replacement._tbl.tblPr.append(deepcopy(ref))
    with pytest.raises(ManagedStyleCompletionError, match="missing or ambiguous style"):
        validate_managed_style_package(_document_bytes(completed), _catalog("en_US"), bindings)


@pytest.mark.parametrize("level", [1, 3])
def test_full_conversion_heading_uses_private_binding_before_wrong_type_name_lookup(tmp_path: Path, level: int) -> None:
    document = Document()
    document.add_paragraph("Template literal")
    document.add_paragraph("{{正文}}")
    conflicting = document.styles[f"Heading {level}"]._element
    assert conflicting is not None
    conflicting.set(qn("w:type"), "character")
    before = etree.tostring(conflicting, method="c14n", exclusive=True)
    template = tmp_path / "heading-conflict.docx"
    document.save(str(template))
    template_bytes = template.read_bytes()
    source = tmp_path / "heading-source.md"
    source.write_text("#" * level + " Fresh heading\n\nBody paragraph.\n", encoding="utf-8")
    context, _workspace = make_context(
        str(source), target_format="docx", options={"template_name": str(template), "heading_merge_mode": "never"}
    )
    result = MdToDocxConverter().convert(context)
    assert result.success, result.error
    output = result.artifacts[0].staging_path
    with ZipFile(output) as archive:
        root = etree.fromstring(archive.read("word/document.xml"))
        styles = etree.fromstring(archive.read("word/styles.xml"))
        old = next(style for style in styles.findall(qn("w:style")) if style.get(qn("w:styleId")) == f"Heading{level}")
        assert etree.tostring(old, method="c14n", exclusive=True) == before
        heading = next(p for p in root.iter(qn("w:p")) if "Fresh heading" in "".join(p.itertext()))
        reference = heading.find(qn("w:pPr") + "/" + qn("w:pStyle"))
        assert reference is not None
        assert reference.get(qn("w:val")) == f"Heading{level}DocWen1"
        generated = next(
            style for style in styles.findall(qn("w:style")) if style.get(qn("w:styleId")) == f"Heading{level}DocWen1"
        )
        assert generated.get(qn("w:type")) == "paragraph"
    assert template.read_bytes() == template_bytes


@pytest.mark.parametrize("kind", ["paragraph", "table"])
def test_yaml_fill_retains_original_bad_style_reference_on_edited_owner(tmp_path: Path, kind: str) -> None:
    document = Document()
    if kind == "paragraph":
        owner = document.add_paragraph("{{标题}}")
        properties = owner._p.get_or_add_pPr()
        reference = OxmlElement("w:pStyle")
    else:
        owner = document.add_table(rows=1, cols=1)
        owner.cell(0, 0).text = "{{标题}}"
        properties = owner._tbl.tblPr
        reference = OxmlElement("w:tblStyle")
    reference.set(qn("w:val"), "MissingTemplateStyle")
    properties.append(reference)
    document.add_paragraph("{{正文}}")
    template = tmp_path / "yaml-template.docx"
    document.save(str(template))
    original = template.read_bytes()
    source = tmp_path / "source.md"
    source.write_text("---\n标题: Filled title\n---\n\nNew body.\n", encoding="utf-8")
    context, _workspace = make_context(str(source), target_format="docx", options={"template_name": str(template)})
    result = MdToDocxConverter().convert(context)
    assert result.success, result.error
    with ZipFile(result.artifacts[0].staging_path) as archive:
        root = etree.fromstring(archive.read("word/document.xml"))
        preserved = next(root.iter(reference.tag))
        assert preserved.get(qn("w:val")) == "MissingTemplateStyle"
        assert "Filled title" in "".join(root.itertext())
        assert "{{标题}}" not in "".join(root.itertext())
    assert template.read_bytes() == original


@pytest.mark.parametrize("replacement", ["reference", "run", "paragraph", "table"])
def test_identical_xml_copy_cannot_replace_original_reference_provenance(replacement: str) -> None:
    document = Document()
    if replacement == "table":
        table = document.add_table(rows=1, cols=1)
        table.cell(0, 0).text = "Template table"
        reference = OxmlElement("w:tblStyle")
        reference.set(qn("w:val"), "MissingTemplateStyle")
        table._tbl.tblPr.append(reference)
    else:
        document.add_paragraph()
        _bad_run(document)
    completed, bindings = complete_managed_styles(document, _catalog("en_US"))
    validate_managed_style_package(_document_bytes(completed), _catalog("en_US"), bindings)
    if replacement == "table":
        original = completed.tables[0]._tbl
    elif replacement == "paragraph":
        original = completed.paragraphs[0]._p
    elif replacement == "run":
        original = completed.paragraphs[0].runs[0]._r
    else:
        original = next(completed.element.iter(qn("w:rStyle")))
    parent = original.getparent()
    assert parent is not None
    parent.replace(original, deepcopy(original))
    with pytest.raises(ManagedStyleCompletionError, match="missing or ambiguous style"):
        validate_managed_style_package(_document_bytes(completed), _catalog("en_US"), bindings)


def test_modified_original_bad_reference_does_not_retain_exemption() -> None:
    document = Document()
    document.add_paragraph()
    _bad_run(document)
    completed, bindings = complete_managed_styles(document, _catalog("en_US"))
    reference = next(completed.element.iter(qn("w:rStyle")))
    reference.set(qn("w:val"), "AnotherMissingStyle")
    with pytest.raises(ManagedStyleCompletionError, match="missing or ambiguous style"):
        validate_managed_style_package(_document_bytes(completed), _catalog("en_US"), bindings)


@pytest.mark.parametrize("replace_part", [False, True])
def test_blob_only_part_baseline_requires_original_reachable_part(replace_part: bool) -> None:
    document = Document()
    payload = (
        b'<w:footnotes xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
        b'<w:footnote w:id="1"><w:p><w:r><w:rPr><w:rStyle w:val="MissingTemplateStyle"/>'
        b"</w:rPr><w:t>Preserved note text</w:t></w:r></w:p></w:footnote></w:footnotes>"
    )
    _attach_xml_part(
        document,
        name="word/footnotes.xml",
        content_type=CT.WML_FOOTNOTES,
        relationship_type=RT.FOOTNOTES,
        payload=payload,
    )
    completed, bindings = complete_managed_styles(document, _catalog("en_US"))
    validate_managed_style_package(_document_bytes(completed), _catalog("en_US"), bindings)
    if not replace_part:
        completed.add_paragraph("Unrelated valid text")
        validate_managed_style_package(_document_bytes(completed), _catalog("en_US"), bindings)
        return
    relationship = next(rel for rel in completed.part.rels.values() if rel.reltype == RT.FOOTNOTES)
    old_part = relationship.target_part
    new_part = Part(PackURI("/word/footnotes.xml"), CT.WML_FOOTNOTES, old_part.blob, completed.part.package)
    completed.part.drop_rel(relationship.rId)
    completed.part.relate_to(new_part, RT.FOOTNOTES)
    with pytest.raises(ManagedStyleCompletionError, match="missing or ambiguous style"):
        validate_managed_style_package(_document_bytes(completed), _catalog("en_US"), bindings)


def test_valid_sibling_reference_cannot_donate_bad_reference_provenance() -> None:
    document = Document()
    document.add_paragraph()
    run = _bad_run(document)
    properties = run.find(qn("w:rPr"))
    assert properties is not None
    valid = OxmlElement("w:rStyle")
    valid.set(qn("w:val"), "Emphasis")
    properties.insert(0, valid)
    completed, bindings = complete_managed_styles(document, _catalog("en_US"))
    bad = next(ref for ref in completed.element.iter(qn("w:rStyle")) if ref.get(qn("w:val")) == "MissingTemplateStyle")
    parent = bad.getparent()
    assert parent is not None
    parent.replace(bad, deepcopy(bad))
    with pytest.raises(ManagedStyleCompletionError, match="missing or ambiguous style"):
        validate_managed_style_package(_document_bytes(completed), _catalog("en_US"), bindings)
