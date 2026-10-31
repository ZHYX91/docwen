"""Template image identity and protected-fill explanations survive XLSX save."""

from __future__ import annotations

from io import BytesIO
from pathlib import Path
from urllib.parse import quote
from zipfile import ZipFile

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Protection
from PIL import Image

from docwen_plugin_markdown.to_spreadsheet.converter import MdToXlsxConverter
from docwen_plugin_markdown.to_spreadsheet.template_xlsx import process_image_placeholders

from .conftest import make_context

pytestmark = pytest.mark.contract


def _png(color: tuple[int, int, int]) -> bytes:
    with BytesIO() as buffer:
        Image.new("RGB", (2, 2), color).save(buffer, format="PNG")
        return buffer.getvalue()


@pytest.mark.parametrize("cwd_image", ["valid-decoy", "invalid-decoy", "absent"])
@pytest.mark.parametrize("reference", ["relative", "absolute"])
def test_template_neighbor_precedes_working_directory(
    tmp_path: Path, monkeypatch, cwd_image: str, reference: str
) -> None:
    template_dir, cwd = tmp_path / "template", tmp_path / "cwd"
    template_dir.mkdir()
    cwd.mkdir()
    correct = _png((240, 0, 0))
    neighbor = template_dir / "图 像%{name}.png"
    neighbor.write_bytes(correct)
    if cwd_image != "absent":
        (cwd / neighbor.name).write_bytes(_png((0, 0, 240)) if cwd_image == "valid-decoy" else b"not an image")
    template = template_dir / "template.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet["B2"] = "{{IMAGE:" + (neighbor.name if reference == "relative" else neighbor.as_posix()) + "}}"
    workbook.save(template)
    workbook.close()
    original = template.read_bytes()
    source = tmp_path / "source.md"
    source.write_text("| Value |\n| --- |\n| control |\n", encoding="utf-8")
    monkeypatch.chdir(cwd)
    context, _ = make_context(str(source), target_format="xlsx", options={"template_name": str(template)})
    result = MdToXlsxConverter().convert(context)
    assert result.success, result.error
    output_path = Path(result.artifacts[0].staging_path)
    output = load_workbook(output_path)
    try:
        sheet = output.active
        assert sheet is not None and sheet["B2"].value is None
        images = vars(sheet).get("_images")
        assert isinstance(images, list) and len(images) == 1
        assert (images[0].anchor._from.col, images[0].anchor._from.row) == (1, 1)
    finally:
        output.close()
    with ZipFile(output_path) as archive:
        names = [name for name in archive.namelist() if name.startswith("xl/media/")]
        assert len(names) == 1 and archive.read(names[0]) == correct
    assert template.read_bytes() == original and neighbor.read_bytes() == correct


@pytest.mark.parametrize("direction", ["↓", "→"])
def test_protected_image_consumption_is_readable_and_does_not_insert(tmp_path: Path, direction: str) -> None:
    image = tmp_path / "图 像%{name}.png"
    image.write_bytes(_png((240, 0, 0)))
    template = tmp_path / "template.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet["A1"] = "{{" + direction + "Value}}"
    sheet["A1"].protection = Protection(locked=False)
    protected = "A2" if direction == "↓" else "B1"
    long_protected = "A3" if direction == "↓" else "C1"
    sheet[protected] = "keep"
    sheet[long_protected] = "keep long"
    sheet.protection.sheet = True
    workbook.save(template)
    workbook.close()
    source = tmp_path / "source.md"
    source.write_text(
        f"| Value |\n| --- |\n| first |\n| ![photo]({quote(image.name)}) |\n| {'😀' * 1000} |\n", encoding="utf-8"
    )
    context, _ = make_context(str(source), target_format="xlsx", options={"template_name": str(template)})
    result = MdToXlsxConverter().convert(context)
    assert result.success, result.error
    output = load_workbook(result.artifacts[0].staging_path)
    try:
        sheet = output.active
        assert sheet is not None and sheet[protected].value == "keep" and sheet["A1"].value == "first"
        assert sheet[protected].comment is not None
        assert sheet[protected].comment.text == "Protected cell skipped; consumed value: Image: " + image.name
        assert sheet[long_protected].value == "keep long"
        assert sheet[long_protected].comment is not None
        assert len(sheet[long_protected].comment.text) == 500
        assert sheet[long_protected].comment.text.endswith("…")
        assert vars(sheet).get("_images") == []
        assert result.artifacts[0].metadata["protected_cells_skipped"] == 2
    finally:
        output.close()


@pytest.mark.parametrize("neighbor", ["absent", "invalid"])
def test_working_directory_fallback_requires_an_absent_neighbor(tmp_path: Path, monkeypatch, neighbor: str) -> None:
    reference_dir, cwd = tmp_path / "reference", tmp_path / "cwd"
    reference_dir.mkdir()
    cwd.mkdir()
    (cwd / "p.png").write_bytes(_png((240, 0, 0)))
    if neighbor == "invalid":
        (reference_dir / "p.png").write_bytes(b"not an image")
    monkeypatch.chdir(cwd)
    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet["B2"] = "{{IMAGE:p.png}}"
    inserted = process_image_placeholders(workbook, reference_dir / "source.md")
    output_path = tmp_path / "fallback.xlsx"
    workbook.save(output_path)
    workbook.close()
    assert inserted == (1 if neighbor == "absent" else 0)
    with ZipFile(output_path) as archive:
        images = [name for name in archive.namelist() if name.startswith("xl/media/")]
        assert len(images) == inserted
        if images:
            assert archive.read(images[0]) == (cwd / "p.png").read_bytes()
