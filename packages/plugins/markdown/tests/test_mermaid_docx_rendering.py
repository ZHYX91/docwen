from __future__ import annotations

from io import BytesIO
from typing import cast

import pytest
from docx import Document
from docx.oxml.ns import qn
from docx.styles.style import ParagraphStyle
from docx.text.paragraph import Paragraph
from PIL import Image

from docwen_plugin_markdown.renderer import MdToDocxRenderer

pytestmark = pytest.mark.unit


def _png_bytes(size: tuple[int, int] = (200, 100)) -> bytes:
    stream = BytesIO()
    Image.new("RGB", size, "white").save(stream, format="PNG")
    return stream.getvalue()


def _mermaid_node() -> dict:
    return {
        "type": "block_code",
        "attrs": {"info": "mermaid"},
        "raw": "graph TD\nA-->B\n",
    }


def test_mermaid_image_mode_embeds_png_instead_of_visible_source() -> None:
    rendered_sources: list[str] = []
    renderer = MdToDocxRenderer(
        Document(),
        mermaid_mode="image",
        mermaid_render=lambda source: rendered_sources.append(source) or _png_bytes(),
    )

    paragraph = renderer._handle_block_code(_mermaid_node())
    assert isinstance(paragraph, Paragraph)

    assert rendered_sources == ["graph TD\nA-->B\n"]
    assert paragraph.text == ""
    assert paragraph._p.find(".//" + qn("w:drawing")) is not None
    assert renderer.warnings == ()


def test_mermaid_image_failure_falls_back_to_code_with_warning() -> None:
    def fail(_source: str) -> bytes:
        raise RuntimeError("renderer unavailable")

    renderer = MdToDocxRenderer(
        Document(),
        mermaid_mode="image",
        mermaid_render=fail,
    )

    paragraph = renderer._handle_block_code(_mermaid_node())
    assert isinstance(paragraph, Paragraph)

    assert paragraph.text == "graph TD\nA-->B"
    assert renderer.warnings[0][0] == "MD2DOCX-MERMAID-FALLBACK"
    assert "renderer unavailable" in renderer.warnings[0][1]


def test_non_mermaid_fence_never_calls_mermaid_renderer() -> None:
    def unexpected(_source: str) -> bytes:
        raise AssertionError("renderer must not be called")

    renderer = MdToDocxRenderer(
        Document(),
        mermaid_mode="image",
        mermaid_render=unexpected,
    )
    node = {"type": "block_code", "attrs": {"info": "python"}, "raw": "print('ok')\n"}

    paragraph = renderer._handle_block_code(node)
    assert isinstance(paragraph, Paragraph)

    assert paragraph.text == "print('ok')"
    assert renderer.warnings == ()


def test_invalid_png_falls_back_without_an_extra_empty_paragraph() -> None:
    doc = Document()
    renderer = MdToDocxRenderer(doc, mermaid_mode="image", mermaid_render=lambda _: b"invalid PNG")
    renderer.render([_mermaid_node()])
    assert [p.text for p in doc.paragraphs] == ["graph TD\nA-->B"]
    assert len(doc.inline_shapes) == 0
    assert "diagram 1" in renderer.warnings[0][1]


def test_cancellation_is_never_converted_to_successful_fallback() -> None:
    from docwen_core.cancellation import CancellationToken

    token = CancellationToken()

    def cancel(_source: str) -> bytes:
        token.cancel()
        token.check()
        return b""

    doc = Document()
    renderer = MdToDocxRenderer(doc, mermaid_mode="image", mermaid_render=cancel, cancellation=token)
    with pytest.raises(Exception, match=r"[Cc]ancel"):
        renderer.render([_mermaid_node()])
    assert renderer.warnings == ()
    assert not doc.paragraphs


@pytest.mark.parametrize("size", [(3000, 600), (600, 6000)])
@pytest.mark.parametrize("quoted", [False, True])
def test_diagram_fits_content_and_quote_indent_with_aspect_ratio(size, quoted) -> None:
    from docx.shared import Inches

    doc = Document()
    cast(ParagraphStyle, doc.styles["Normal"]).paragraph_format.right_indent = Inches(0.2)
    extent = (int(Inches(3)), int(Inches(5)))
    renderer = MdToDocxRenderer(
        doc, mermaid_mode="image", mermaid_render=lambda _: _png_bytes(size), content_extent=extent
    )
    node = _mermaid_node()
    renderer.render([{"type": "block_quote", "children": [node]} if quoted else node])
    paragraph = doc.paragraphs[0]
    shape = doc.inline_shapes[0]
    available = extent[0] - sum(renderer._paragraph_length(paragraph, key) for key in ("left_indent", "right_indent"))
    assert shape.width <= available
    assert shape.height <= extent[1]
    assert shape.width / shape.height == pytest.approx(size[0] / size[1], rel=1e-5)
    assert paragraph.paragraph_format.line_spacing == 1.0
    assert renderer.warnings[0][0] == "MD2DOCX-MERMAID-SCALED"


def test_body_extent_uses_placeholder_section_and_columns() -> None:
    from docx.enum.section import WD_SECTION_START
    from docx.shared import Inches

    from docwen_plugin_markdown.template_utils import body_content_extent

    doc = Document()
    first = doc.sections[0]
    first.page_width = Inches(8)
    first.left_margin = first.right_margin = Inches(1)
    columns = first._sectPr.find(qn("w:cols"))
    assert columns is not None
    columns.set(qn("w:num"), "2")
    columns.set(qn("w:space"), "720")
    placeholder = doc.add_paragraph("{{body}}")
    second = doc.add_section(WD_SECTION_START.NEW_PAGE)
    second.page_width = Inches(12)
    assert body_content_extent(doc, placeholder)[0] == int(Inches(2.75))
    assert body_content_extent(doc)[0] == int(Inches(4.75))
