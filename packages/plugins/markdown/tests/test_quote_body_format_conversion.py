"""Body anchor formatting does not override quote paragraphs in saved conversion."""

from __future__ import annotations

from hashlib import sha256
from pathlib import Path
from typing import Any

import pytest
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt

from ._md_to_docx_support import MdToDocxConverter, make_context

pytestmark = pytest.mark.contract


def _indent(paragraph: Any) -> dict[str, str] | None:
    element = paragraph._p.find(f"{qn('w:pPr')}/{qn('w:ind')}")
    return {name.rsplit("}", 1)[-1]: value for name, value in element.attrib.items()} if element is not None else None


@pytest.mark.parametrize(
    ("authored", "expected"),
    [
        ({"firstLineChars": "125", "firstLine": "400"}, {"firstLineChars": "125", "firstLine": "400"}),
        ({"firstLine": "400"}, {"firstLine": "400", "firstLineChars": "0"}),
        ({"firstLineChars": "0"}, {"firstLineChars": "0", "firstLine": "0"}),
    ],
)
def test_body_anchor_properties_leave_quote_and_quoted_list_ownership_intact(
    tmp_path: Path, authored: dict[str, str], expected: dict[str, str]
) -> None:
    template: Any = Document()
    marker = template.add_paragraph("{{body}}", style="Normal")
    marker.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    marker.paragraph_format.space_after = Pt(19)
    marker.runs[0].font.size = Pt(18)
    ind = marker._p.get_or_add_pPr().get_or_add_ind()
    for name, value in {**authored, "left": "320", "right": "160"}.items():
        ind.set(qn(f"w:{name}"), value)
    template_path = tmp_path / "template.docx"
    template.save(template_path)
    template_before = sha256(template_path.read_bytes()).hexdigest()
    source = tmp_path / "source.md"
    source.write_text(
        "Ordinary before.\n\n"
        "> Outer quote.\n>\n> > Nested quote.\n> >\n> > Nested continuation.\n>\n"
        "> Outer continuation.\n>\n> - Quoted item.\n>\n>   Quoted list continuation.\n>\n> - Quoted second.\n\n"
        "Ordinary after.\n",
        encoding="utf-8",
    )
    context, _ = make_context(str(source), options={"template_name": str(template_path)})
    result = MdToDocxConverter().convert(context)
    assert result.success, result.error
    assert sha256(template_path.read_bytes()).hexdigest() == template_before
    output = Document(result.artifacts[0].staging_path)
    paragraphs = {paragraph.text: paragraph for paragraph in output.paragraphs if paragraph.text}
    assert set(paragraphs) == {
        "Ordinary before.",
        "Ordinary after.",
        "Outer quote.",
        "Nested quote.",
        "Nested continuation.",
        "Outer continuation.",
        "Quoted item.",
        "Quoted list continuation.",
        "Quoted second.",
    }
    for text in ("Ordinary before.", "Ordinary after."):
        paragraph = paragraphs[text]
        style = paragraph.style
        assert style is not None and style.style_id == "DocWenBodyParagraph"
        assert _indent(paragraph) == {**expected, "left": "320", "right": "160"}
        assert paragraph.alignment == WD_ALIGN_PARAGRAPH.RIGHT
        assert paragraph.paragraph_format.space_after == Pt(19)
        assert paragraph.runs[0].font.size == Pt(18)
    for text, paragraph in paragraphs.items():
        if text.startswith("Ordinary"):
            continue
        level = 2 if text.startswith("Nested") else 1
        style = paragraph.style
        if text.startswith("Quoted"):
            assert style is not None and style.style_id == "DocWenListBlock"
            assert paragraph._p.find(f"{qn('w:pPr')}/{qn('w:pBdr')}/{qn('w:left')}") is not None
            assert paragraph.alignment is None and paragraph.paragraph_format.space_after is None
            indent = _indent(paragraph)
            assert indent is None or (
                "firstLineChars" not in indent and indent.get("firstLine", "0") == "0" and "right" not in indent
            )
            number = paragraph._p.find(f"{qn('w:pPr')}/{qn('w:numPr')}")
            if "continuation" in text:
                assert number is not None
                number_id = number.find(qn("w:numId"))
                assert number_id is not None and number_id.get(qn("w:val")) == "0"
            else:
                assert number is not None and number.find(qn("w:numId")) is not None
            continue
        assert style is not None and style.style_id == f"DocWenQuote{level}"
        assert _indent(paragraph) is None
        assert paragraph.alignment is None and paragraph.paragraph_format.space_after is None
        style_element = style._element
        assert style_element is not None
        style_ind = style_element.find(f"{qn('w:pPr')}/{qn('w:ind')}")
        assert style_ind is not None
        assert style_ind.get(qn("w:left")) == str(480 + (level - 1) * 240)
        assert style_ind.get(qn("w:right")) == "480" and style_ind.get(qn("w:firstLine")) == "0"
        assert all(run.font.size is None for run in paragraph.runs)
        number = paragraph._p.find(f"{qn('w:pPr')}/{qn('w:numPr')}")
        if text.startswith("Quoted"):
            assert number is not None and number.find(qn("w:numId")) is not None
        else:
            assert number is None
