"""Effective attachment origins retain Word's character cancellation and fallback rules."""

from __future__ import annotations

import pytest
from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt
from docx.styles.style import ParagraphStyle

from docwen_plugin_markdown.template_filler import fill_template

pytestmark = pytest.mark.contract


@pytest.mark.parametrize("in_table", [False, True])
@pytest.mark.parametrize(
    "inherited,direct,in_defaults,origin",
    [
        ({}, {"left": "240", "leftChars": "0"}, False, 12),
        ({"left": "400", "leftChars": "200"}, {"left": "240"}, False, 32),
        ({"firstLineChars": "200", "firstLine": "1280"}, {"hanging": "200"}, False, 32),
        (
            {"left": "1580", "leftChars": "200", "hanging": "948", "hangingChars": "300"},
            {"leftChars": "0", "firstLineChars": "0"},
            False,
            31.6,
        ),
        ({"left": "640", "firstLineChars": "100"}, {}, True, 48),
        ({"leftChars": "200", "left": "1580"}, {}, True, 32),
        ({"firstLine": "640"}, {"hanging": "0"}, False, 0),
        ({"hanging": "480"}, {"firstLine": "0"}, False, 0),
    ],
)
def test_attachment_origin_resolves_units_defaults_and_related_axes(
    inherited: dict[str, str], direct: dict[str, str], in_defaults: bool, origin: float, in_table: bool
) -> None:
    document = Document()
    paragraph = document.add_table(1, 1).cell(0, 0).paragraphs[0] if in_table else document.add_paragraph()
    paragraph.add_run("{{附件说明}}").font.size = Pt(16)
    if in_defaults:
        defaults = document.styles.element.find(f"{qn('w:docDefaults')}/{qn('w:pPrDefault')}/{qn('w:pPr')}")
        assert defaults is not None
        indent = OxmlElement("w:ind")
        defaults.append(indent)
    else:
        style = document.styles.add_style("Attachment source", WD_STYLE_TYPE.PARAGRAPH)
        assert isinstance(style, ParagraphStyle)
        paragraph.style = style
        properties = OxmlElement("w:pPr")
        indent = OxmlElement("w:ind")
        properties.append(indent)
        style.element.append(properties)
    for name, value in inherited.items():
        indent.set(qn(f"w:{name}"), value)
    authored = paragraph._p.get_or_add_pPr().get_or_add_ind()
    for name, value in direct.items():
        authored.set(qn(f"w:{name}"), value)
    styles_before = document.styles.element.xml

    fill_template(document, {"附件说明": ["甲", "乙"]}, [], None, {"附件说明": [paragraph]})

    paragraphs = document.tables[0].cell(0, 0).paragraphs if in_table else document.paragraphs
    assert len(paragraphs) == 2
    for output in paragraphs:
        left, first = output.paragraph_format.left_indent, output.paragraph_format.first_line_indent
        assert left is not None and first is not None
        assert left == Pt(origin + 64)
        assert left.twips + first.twips == round(origin * 20)
    assert document.styles.element.xml == styles_before
