"""Simple field-list rendering, precise Word comments and attachment anchors."""

from __future__ import annotations

from copy import deepcopy
from io import BytesIO

import pytest
from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt
from docx.styles.style import ParagraphStyle

from docwen_plugin_markdown.gongwen_fields import GongwenFieldPlacementError, prepare_gongwen_fields
from docwen_plugin_markdown.template_filler import fill_template

pytestmark = pytest.mark.contract
FIELDS = ("签发人", "主送机关", "抄送机关")


def _fill(document, values, placeholders, *, enabled=True, separator=" / ") -> None:
    fill_template(
        document,
        values,
        [],
        None,
        placeholders,
        gongwen_fields_enabled=enabled,
        list_separator=separator,
    )


def _roundtrip(document):
    stream = BytesIO()
    document.save(stream)
    stream.seek(0)
    return Document(stream)


def _anchors(document) -> dict[str, str]:
    active: dict[str, list[str]] = {}
    complete: dict[str, str] = {}
    for node in document.element.body.iter():
        if node.tag == qn("w:commentRangeStart"):
            active[node.get(qn("w:id"))] = []
        elif node.tag == qn("w:commentRangeEnd"):
            ident = node.get(qn("w:id"))
            complete[ident] = "".join(active.pop(ident))
        elif node.tag == qn("w:t"):
            for parts in active.values():
                parts.append(node.text or "")
    assert not active
    return complete


def test_field_reviews_preserve_malformed_comments_and_conflicting_styles():
    document = Document()
    old = document.add_comment(document.add_paragraph("original").runs, "old review")
    old._element.set(qn("w:id"), "bad-id")
    document.styles.add_style("CommentReference", WD_STYLE_TYPE.PARAGRAPH)
    document.styles.add_style("CommentText", WD_STYLE_TYPE.CHARACTER)
    paragraph = document.add_paragraph("前 {{签发人}} 后")
    _fill(document, {"签发人": ["甲", "乙"]}, {"签发人": [paragraph]})
    reopened = _roundtrip(document)
    comments = list(reopened.comments)
    assert comments[0]._element.get(qn("w:id")) == "bad-id"
    assert comments[1].comment_id == 1
    assert _anchors(reopened)["1"] == "甲、乙"
    assert reopened.styles["CommentReference"].type == WD_STYLE_TYPE.PARAGRAPH
    assert reopened.styles["CommentText"].type == WD_STYLE_TYPE.CHARACTER


@pytest.mark.parametrize("key", FIELDS)
@pytest.mark.parametrize("separator", [", ", " / ", "，", "", " "])
def test_multi_item_fields_override_generic_separator_and_annotate_only_value(key, separator) -> None:
    document = Document()
    paragraph = document.add_paragraph(f"前文 {{{{{key}}}}} 后文")
    paragraph.runs[0].bold = True
    normal = document.add_paragraph("{{普通列表}}")
    values = {key: ["甲", None, "", "乙", "甲"], "普通列表": ["丙", "丁"]}
    original = deepcopy(values)
    _fill(document, values, {key: [paragraph], "普通列表": [normal]}, separator=separator)
    assert values == original
    assert paragraph.text == "前文 甲、乙、甲 后文"
    assert normal.text == f"丙{separator}丁"
    reopened = _roundtrip(document)
    comments = list(reopened.comments)
    assert len(comments) == 1
    assert reopened.styles["CommentReference"].type == WD_STYLE_TYPE.CHARACTER
    assert reopened.styles["CommentText"].type == WD_STYLE_TYPE.PARAGRAPH
    assert comments[0].author == "DocWen"
    assert _anchors(reopened) == {str(comments[0].comment_id): "甲、乙、甲"}
    assert all(run.bold for run in paragraph.runs if run.text)
    if key == "签发人":
        assert "自行调整" in comments[0].text and "换行" in comments[0].text
    else:
        assert key in comments[0].text and "顿号" in comments[0].text and "逗号" in comments[0].text


@pytest.mark.parametrize("key", FIELDS)
@pytest.mark.parametrize(
    "value,expected",
    [
        ([], ""),
        (None, ""),
        (["甲"], "甲"),
        ([None, " ", "甲"], "甲"),
        (" 各县（市、区）人民政府，省直各单位 ", " 各县（市、区）人民政府，省直各单位 "),
    ],
)
def test_single_empty_and_authored_strings_are_not_punctuated_or_annotated(key, value, expected) -> None:
    document = Document()
    paragraph = document.add_paragraph(f"保留前文 {{{{{key}}}}} 保留后文")
    _fill(document, {key: value}, {key: [paragraph]})
    assert paragraph.text == f"保留前文 {expected} 保留后文"
    assert not list(document.comments)


@pytest.mark.parametrize("key", FIELDS)
def test_switch_off_uses_the_exact_generic_separator_without_comments(key) -> None:
    document = Document()
    paragraph = document.add_paragraph(f"{{{{{key}}}}}")
    _fill(document, {key: ["甲", "乙"]}, {key: [paragraph]}, enabled=False)
    assert paragraph.text == "甲 / 乙"
    assert not list(document.comments)


def test_multiple_fields_and_literal_markers_in_one_run_keep_original_spans() -> None:
    document = Document()
    paragraph = document.add_paragraph("{{主送机关}} / {{签发人}} / {{主送机关}}")
    _fill(
        document,
        {"主送机关": ["{{签发人}}", "乙单位"], "签发人": ["甲", "乙"]},
        {"主送机关": [paragraph], "签发人": [paragraph]},
    )
    assert paragraph.text == "{{签发人}}、乙单位 / 甲、乙 / {{签发人}}、乙单位"
    assert sorted(_anchors(_roundtrip(document)).values()) == sorted(
        [
            "{{签发人}}、乙单位",
            "甲、乙",
            "{{签发人}}、乙单位",
        ]
    )


def test_cross_run_placeholder_keeps_surrounding_controls_and_existing_comments() -> None:
    document = Document()
    paragraph = document.add_paragraph()
    first = paragraph.add_run("前文\t{{签")
    first.bold = True
    second = paragraph.add_run("发人}}\n后文")
    second.italic = True
    document.add_comment(first, text="原批注", author="原作者")
    _fill(document, {"签发人": ["张三", "李四"]}, {"签发人": [paragraph]})
    assert paragraph.text == "前文\t张三、李四\n后文"
    reopened = _roundtrip(document)
    assert len(reopened.element.findall(f".//{qn('w:tab')}")) == 1
    assert len(reopened.element.findall(f".//{qn('w:br')}")) == 1
    comments = list(reopened.comments)
    assert next(c.text for c in comments) == "原批注"
    inserted = next(c for c in comments if c.author == "DocWen")
    assert _anchors(reopened)[str(inserted.comment_id)] == "张三、李四"


def test_comments_inside_hyperlinks_and_table_cells_stay_in_their_owners() -> None:
    document = Document()
    table = document.add_table(rows=1, cols=1)
    paragraph = table.cell(0, 0).paragraphs[0]
    link = OxmlElement("w:hyperlink")
    link.set(qn("w:anchor"), "target")
    run = OxmlElement("w:r")
    text = OxmlElement("w:t")
    text.text = "前{{抄送机关}}后"
    run.append(text)
    link.append(run)
    paragraph._p.append(link)
    _fill(document, {"抄送机关": ["甲", "乙"]}, {"抄送机关": [paragraph]})
    assert paragraph.text == "前甲、乙后"
    assert "".join(node.text or "" for node in link.iter(qn("w:t"))) == "前甲、乙后"
    assert list(_anchors(_roundtrip(document)).values()) == ["甲、乙"]


@pytest.mark.parametrize("part", ["header", "footer"])
def test_required_comments_in_unsupported_parts_fail_before_mutation(part) -> None:
    document = Document()
    paragraph = getattr(document.sections[0], part).paragraphs[0]
    paragraph.text = "{{签发人}}"
    original = paragraph._p.xml
    with pytest.raises(GongwenFieldPlacementError, match="需要添加核对批注"):
        _fill(document, {"签发人": ["甲", "乙"]}, {"签发人": [paragraph]})
    assert paragraph._p.xml == original


@pytest.mark.parametrize("key", FIELDS)
@pytest.mark.parametrize("value", [123, False, {"姓名": "甲"}, [["甲"]], ["甲", 123]])
def test_list_fields_reject_invalid_types_instead_of_printing_python_containers(key, value) -> None:
    with pytest.raises(ValueError, match=key):
        prepare_gongwen_fields({key: value})


@pytest.mark.parametrize("key", ["附注", "成文日期", "印发日期"])
def test_scalar_fields_do_not_stringify_nonempty_lists(key) -> None:
    with pytest.raises(ValueError, match="不要填写列表或对象"):
        prepare_gongwen_fields({key: ["甲", "乙"]})


@pytest.mark.parametrize("count", [1, 2, 16, 100])
@pytest.mark.parametrize("mode", ["left", "first", "inherited", "characters", "hanging"])
def test_attachment_label_keeps_effective_first_line_origin(count, mode) -> None:
    document = Document()
    paragraph = document.add_paragraph("{{附件说明}}")
    paragraph.runs[0].font.size = Pt(16)
    paragraph.paragraph_format.right_indent = Pt(8)
    if mode == "left":
        paragraph.paragraph_format.left_indent = Pt(32)
    elif mode == "first":
        paragraph.paragraph_format.first_line_indent = Pt(32)
    elif mode == "inherited":
        style = document.styles.add_style("Attachment origin", WD_STYLE_TYPE.PARAGRAPH)
        assert isinstance(style, ParagraphStyle)
        style.paragraph_format.left_indent = Pt(16)
        style.paragraph_format.first_line_indent = Pt(16)
        paragraph.style = style
    elif mode == "characters":
        indent = paragraph._p.get_or_add_pPr().get_or_add_ind()
        indent.set(qn("w:firstLineChars"), "200")
    else:
        paragraph.paragraph_format.left_indent = Pt(48)
        paragraph.paragraph_format.first_line_indent = Pt(-16)
    _fill(document, {"附件说明": [f"附件{i}" for i in range(count)]}, {"附件说明": [paragraph]})
    paragraphs = document.paragraphs
    assert len(paragraphs) == count
    first = paragraphs[0].paragraph_format
    left, first_line = first.left_indent, first.first_line_indent
    assert left is not None and first_line is not None
    assert left + first_line == Pt(32)
    assert all(p.paragraph_format.right_indent == Pt(8) for p in paragraphs)
    if count > 1:
        assert len({p.paragraph_format.left_indent for p in paragraphs}) == 1
        assert all(p.paragraph_format.first_line_indent == Pt(-64) for p in paragraphs)
    assert all("\t" not in p.text for p in paragraphs)
