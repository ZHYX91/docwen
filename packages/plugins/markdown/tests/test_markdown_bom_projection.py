"""Only a complete source may consume a document encoding signature."""

from __future__ import annotations

from pathlib import Path
from zipfile import ZipFile

import pytest
from docx import Document
from docx.oxml.ns import qn
from lxml import etree

from docwen_core.export_semantics import LinkRuntimeConfig
from docwen_core.markdown_extensions import MarkdownExtensions
from docwen_plugin_markdown.mistune_extensions import parse_markdown_text
from docwen_plugin_markdown.renderer import MdToDocxRenderer
from docwen_plugin_markdown.to_docx.notes import (
    prepare_note_context_for_document,
    process_md_body_with_notes,
    write_notes_to_docx,
)
from docwen_plugin_markdown.yaml_links import YamlLinkProjection

pytestmark = pytest.mark.contract


@pytest.mark.parametrize("consume", [False, True])
def test_parser_requires_document_context_to_consume_a_leading_bom(consume: bool) -> None:
    source = "\ufeff# Title\ufeffinside\r\n"
    [node] = parse_markdown_text(source, extensions=MarkdownExtensions(), consume_initial_bom=consume)
    assert node["type"] == ("heading" if consume else "paragraph")
    assert "".join(child.get("raw", "") for child in node["children"]) == (
        "Title\ufeffinside" if consume else "\ufeff# Title\ufeffinside"
    )


def test_document_parser_consumes_at_most_one_signature() -> None:
    [node] = parse_markdown_text("\ufeff\ufeff# Literal\n", extensions=MarkdownExtensions())
    assert node["type"] == "paragraph"
    assert "".join(child.get("raw", "") for child in node["children"]) == "\ufeff# Literal"


def test_yaml_hyperlink_label_keeps_its_leading_literal_bom_in_the_package(tmp_path: Path) -> None:
    projection = YamlLinkProjection(str(tmp_path / "source.md"), LinkRuntimeConfig())
    token = projection.project("[\ufeff**Label**](https://example.test/bom)")
    assert projection.plain(token) == "\ufeffLabel"
    document = Document()
    document.add_paragraph(token)
    projection.materialize(document)
    output = tmp_path / "label.docx"
    document.save(str(output))

    with ZipFile(output) as package:
        body = etree.fromstring(package.read("word/document.xml"))
        relationships = etree.fromstring(package.read("word/_rels/document.xml.rels"))
    [link] = body.findall(".//" + qn("w:hyperlink"))
    assert "".join(node.text or "" for node in link.iter(qn("w:t"))) == "\ufeffLabel"
    [relationship] = [item for item in relationships if item.get("Id") == link.get(qn("r:id"))]
    assert (relationship.get("Target"), relationship.get("TargetMode")) == ("https://example.test/bom", "External")


def test_note_body_fragment_keeps_literal_boms_in_document_and_note_parts(tmp_path: Path) -> None:
    body = "\ufeff# Literal\n\nText[^n].\n\n[^n]: \ufeffNote text\n"
    nodes, note_context = process_md_body_with_notes(body)
    assert nodes[0]["type"] == "paragraph"
    document = Document()
    prepare_note_context_for_document(document, note_context)
    MdToDocxRenderer(document, note_ctx=note_context).render(nodes)
    output = tmp_path / "note.docx"
    document.save(str(output))
    write_notes_to_docx(str(output), note_context)

    with ZipFile(output) as package:
        main = etree.fromstring(package.read("word/document.xml"))
        notes = etree.fromstring(package.read("word/footnotes.xml"))
    [body_element] = main.findall(qn("w:body"))
    first = body_element.findall(qn("w:p"))[0]
    assert "".join(node.text or "" for node in first.iter(qn("w:t"))) == "\ufeff# Literal"
    [reference] = main.findall(".//" + qn("w:footnoteReference"))
    [note] = [item for item in notes if item.get(qn("w:id")) == reference.get(qn("w:id"))]
    assert "".join(node.text or "" for node in note.iter(qn("w:t"))) == " \ufeffNote text"
