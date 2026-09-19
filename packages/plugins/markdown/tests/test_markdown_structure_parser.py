"""Markdown block and inline boundaries survive parsing without source rewrites."""

from __future__ import annotations

import pytest
from docx import Document

from docwen_plugin_markdown.mistune_extensions import parse_markdown_text
from docwen_plugin_markdown.renderer import MdToDocxRenderer

pytestmark = pytest.mark.unit


def _nodes(source):
    return [node for node in parse_markdown_text(source) if node["type"] != "blank_line"]


def _descendants(nodes):
    for node in nodes:
        yield node
        yield from _descendants(node.get("children", []))


@pytest.mark.parametrize("newline", ["\n", "\r\n"])
@pytest.mark.parametrize("trailing", ["", "   ", "\t"])
def test_blank_separated_rule_is_not_an_empty_setext_heading(newline, trailing):
    nodes = _nodes(newline.join(("A", "", f"---{trailing}")))
    assert [node["type"] for node in nodes] == ["paragraph", "thematic_break"]
    assert nodes[1]["_hr_marker"] == "dash"
    assert not nodes[1].get("_attach_to_prev")


@pytest.mark.parametrize("newline", ["\n", "\r\n"])
@pytest.mark.parametrize("trailing", ["", "   ", "\t"])
@pytest.mark.parametrize("underline", ["-", "--", "---", "=", "==="])
def test_multiline_setext_is_one_heading_with_line_local_trailing_space(newline, trailing, underline):
    nodes = _nodes(newline.join(("First line", "Second line", underline + trailing, "Body")))
    assert [node["type"] for node in nodes] == ["heading", "paragraph"]
    assert nodes[0]["attrs"]["level"] == (1 if underline.startswith("=") else 2)
    assert nodes[0]["children"] == [
        {"type": "text", "raw": "First line"},
        {"type": "softbreak"},
        {"type": "text", "raw": "Second line"},
    ]


@pytest.mark.parametrize("underline", ["---", "==="])
def test_whitespace_only_predecessor_is_not_a_setext_heading(underline):
    assert all(node["type"] != "heading" for node in _nodes(f"A\r\n \t\r\n{underline}\r\nB"))


@pytest.mark.parametrize("prefix,kind", [("- Item", "list"), ("# ATX", "heading"), ("> Quote", "block_quote")])
def test_setext_rewrite_cannot_consume_a_different_block(prefix, kind):
    nodes = _nodes(prefix + "\n---\nBody")
    assert [node["type"] for node in nodes] == [kind, "thematic_break", "paragraph"]


@pytest.mark.parametrize("fence", ["```", "````", "~~~"])
def test_fenced_setext_and_html_examples_remain_literal(fence):
    literal = "---\nReportName: example\nUnit: office<br>\n---\n"
    nodes = _nodes(f"{fence}markdown\n{literal}{fence}\nReal heading\n---\n")
    assert nodes[0]["type"] == "block_code"
    assert nodes[0]["raw"] == literal
    assert nodes[1]["type"] == "heading"
    assert nodes[1]["children"] == [{"type": "text", "raw": "Real heading"}]


def test_indented_example_is_literal_and_does_not_consume_later_markers():
    nodes = _nodes("    Title\n    ---\n    <br>\n\n___\n")
    assert nodes[0]["type"] == "block_code"
    assert nodes[0]["raw"] == "Title\n---\n<br>"
    assert nodes[1]["_hr_marker"] == "underscore"


@pytest.mark.parametrize(
    "atom,kind", [("`code`", "codespan"), ("[link](https://example.com)", "link"), ("$x+y$", "inline_math")]
)
def test_setext_heading_keeps_inline_atoms_inside_one_heading(atom, kind):
    nodes = _nodes(f"Title {atom} tail\n---\n")
    assert len(nodes) == 1 and nodes[0]["type"] == "heading"
    assert [child["type"] for child in nodes[0]["children"]] == ["text", kind, "text"]
    assert nodes[0]["children"][-1]["raw"] == " tail"


@pytest.mark.parametrize("tag", ["<br>", "<br/>", "<BR />"])
def test_html_break_stays_inside_its_heading_and_table_cell(tag):
    ast = parse_markdown_text(f"# First{tag}Second\n\n| Head |\n| --- |\n| Left{tag}Right |\n")
    nodes = [node for node in ast if node["type"] != "blank_line"]
    assert [node["type"] for node in nodes] == ["heading", "table"]
    assert [child["type"] for child in nodes[0]["children"]] == ["text", "linebreak", "text"]
    document = Document()
    MdToDocxRenderer(document).render(ast)
    assert document.paragraphs[0].text == "First\nSecond"
    assert document.tables[0].cell(1, 0).text == "Left\nRight"


def test_html_break_does_not_enter_code_math_or_an_escaped_tag():
    ast = parse_markdown_text(r"Text<br>next and `<br>` and $x<br>y$ and \<br>.")
    nodes = list(_descendants(ast))
    assert sum(node["type"] == "linebreak" for node in nodes) == 1
    assert any(node["type"] == "codespan" and node["raw"] == "<br>" for node in nodes)
    assert any(node["type"] == "inline_math" and node["raw"] == "x<br>y" for node in nodes)
    assert "<br>" in "".join(node.get("raw", "") for node in nodes if node["type"] == "text")


def test_html_break_remains_in_link_text_and_preserves_other_html():
    nodes = _nodes("[Before<br/>after](https://example.com) <span>inside<br>text</span>")
    children = nodes[0]["children"]
    assert children[0]["type"] == "link"
    assert children[0]["attrs"]["url"] == "https://example.com"
    assert children[0]["children"][1]["type"] == "linebreak"
    assert [node["raw"] for node in children if node["type"] == "inline_html"] == ["<span>", "</span>"]


def test_rule_marker_is_bound_after_code_setext_and_nested_rules():
    source = "```md\n***\n```\n\nSetext\n---\n\n> ___\n\n* * * *\n\n_ _ _\n"
    ast = parse_markdown_text(source)
    assert [node["_hr_marker"] for node in ast if node["type"] == "thematic_break"] == ["asterisk", "underscore"]
    assert [node["_hr_marker"] for node in _descendants(ast) if node["type"] == "thematic_break"] == [
        "underscore",
        "asterisk",
        "underscore",
    ]


@pytest.mark.parametrize(
    "prefix,attached",
    [("Text", True), ("- Item", True), ("Text\n", False), ("# Heading", False), ("```\ncode\n```", False)],
)
def test_rule_attachment_uses_actual_adjacent_body_block(prefix, attached):
    rule = _nodes(prefix + "\n___")[-1]
    assert rule["type"] == "thematic_break"
    assert bool(rule.get("_attach_to_prev")) is attached
