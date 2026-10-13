"""Body anchor indentation projection preserves units and effective overrides."""

from __future__ import annotations

from dataclasses import replace
from pathlib import Path

import lxml.etree as etree
import pytest
from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt
from docx.styles.style import ParagraphStyle

from docwen_plugin_markdown.mistune_extensions import parse_markdown_text
from docwen_plugin_markdown.renderer import MdToDocxRenderer
from docwen_plugin_markdown.template_filler import fill_template
from docwen_plugin_markdown.template_utils import extract_body_paragraph_format
from docwen_plugin_markdown.to_docx.managed_styles import complete_managed_styles
from docwen_runtime.config.document_styles import build_document_style_catalog

from .conftest import PROJECT_ROOT

pytestmark = pytest.mark.contract


def _catalog(locale: str = "en_US"):
    return build_document_style_catalog(
        {"gui": {"language": {"locale": locale}}},
        locales_dir=PROJECT_ROOT / "i18n" / "locales",
    )


def _save_reopen(document, tmp_path: Path, name: str):
    path = tmp_path / name
    document.save(path)
    return path, Document(str(path))


@pytest.mark.parametrize(
    "indent_attributes",
    [
        {"hangingChars": "200"},
        {"hanging": "480", "hangingChars": "200", "left": "720", "leftChars": "300"},
        {"firstLineChars": "200", "firstLine": "0", "rightChars": "150"},
        {"firstLineChars": "0", "leftChars": "0"},
    ],
)
def test_body_placeholder_preserves_word_character_indent_units(
    tmp_path: Path, indent_attributes: dict[str, str]
) -> None:
    template = Document()
    custom = template.styles.add_style("UserBodyIndent", WD_STYLE_TYPE.PARAGRAPH)
    assert isinstance(custom, ParagraphStyle)
    custom.paragraph_format.first_line_indent = Inches(0.2)
    placeholder = template.add_paragraph("{{正文}}", style=custom)
    source_indent = placeholder._p.get_or_add_pPr().get_or_add_ind()
    for name, value in indent_attributes.items():
        source_indent.set(qn(f"w:{name}"), value)
    template, bindings = complete_managed_styles(template, _catalog())
    body_format = extract_body_paragraph_format(template)
    assert body_format is not None
    rendered = MdToDocxRenderer(
        template,
        body_style=next(p.style for p in template.paragraphs if p.text == "{{正文}}"),
        body_paragraph_format=body_format,
        managed_styles=bindings,
    ).render(parse_markdown_text("An ordinary paragraph.\n"))
    fill_template(template, {}, rendered, next(p for p in template.paragraphs if p.text == "{{正文}}"))
    _, reopened = _save_reopen(template, tmp_path, "chars-indent.docx")
    body = next(p for p in reopened.paragraphs if p.text == "An ordinary paragraph.")
    output_indent = body._p.find(f"{qn('w:pPr')}/{qn('w:ind')}")
    assert output_indent is not None
    assert dict(output_indent.attrib) == {qn(f"w:{name}"): value for name, value in indent_attributes.items()}
    assert isinstance(body.style, ParagraphStyle)
    assert body.style.style_id == custom.style_id


def test_body_indent_with_only_style_inheritance_keeps_output_unset(tmp_path: Path) -> None:
    template = Document()
    body_style = template.styles.add_style("StyleOnlyIndent", WD_STYLE_TYPE.PARAGRAPH)
    assert isinstance(body_style, ParagraphStyle)
    body_style.paragraph_format.left_indent = Inches(0.3)
    template.add_paragraph("{{正文}}", style=body_style)
    template, bindings = complete_managed_styles(template, _catalog())
    body_format = extract_body_paragraph_format(template)
    assert body_format is not None
    assert not body_format.indent_attributes
    placeholder = next(p for p in template.paragraphs if p.text == "{{正文}}")
    rendered = MdToDocxRenderer(
        template, body_style=placeholder.style, body_paragraph_format=body_format, managed_styles=bindings
    ).render(parse_markdown_text("Inherited indentation.\n"))
    fill_template(template, {}, rendered, placeholder)
    _, reopened = _save_reopen(template, tmp_path, "inherited-indent.docx")
    body = next(p for p in reopened.paragraphs if p.text == "Inherited indentation.")
    assert body._p.find(f"{qn('w:pPr')}/{qn('w:ind')}") is None
    assert isinstance(body.style, ParagraphStyle)
    assert body.style.paragraph_format.left_indent == Inches(0.3)


@pytest.mark.parametrize(
    ("authored", "inherited", "projection"),
    [
        ({"firstLine": "0"}, {}, {"firstLineChars": "0"}),
        ({"firstLine": "480"}, {}, {"firstLineChars": "0"}),
        ({"hanging": "360"}, {}, {"hangingChars": "0"}),
        ({"firstLineChars": "0"}, {}, {"firstLine": "0"}),
        ({"firstLineChars": "125", "firstLine": "400"}, {}, {}),
        ({"firstLineChars": "0"}, {"firstLine": "480"}, {"firstLine": "480"}),
        ({"firstLineChars": "0"}, {"firstLine": "640"}, {}),
        ({"firstLineChars": "0"}, {"firstLine": "400", "firstLineChars": "100"}, {"firstLine": "400"}),
        ({"firstLine": "0"}, {"firstLineChars": "100"}, {"firstLineChars": "100"}),
        ({"hanging": "360"}, {"hangingChars": "150"}, {"hangingChars": "150"}),
        ({"left": "720"}, {"leftChars": "300"}, {}),
        ({}, {"firstLine": "480"}, {}),
    ],
)
def test_normal_mapping_keeps_authored_indent_effective_without_rewriting_style(
    tmp_path: Path, authored: dict[str, str], inherited: dict[str, str], projection: dict[str, str]
) -> None:
    template = Document()
    normal = template.styles["Normal"]
    assert isinstance(normal, ParagraphStyle)
    normal_element = normal._element
    assert normal_element is not None
    normal_element.set(qn("w:styleId"), "HostNormal42")
    for reference in template.styles.element.iter(qn("w:basedOn")):
        if reference.get(qn("w:val")) == "Normal":
            reference.set(qn("w:val"), "HostNormal42")
    normal_indent = normal_element.get_or_add_pPr().get_or_add_ind()
    for name, value in inherited.items():
        normal_indent.set(qn(f"w:{name}"), value)
    placeholder = template.add_paragraph("{{body}}", style="Normal")
    if authored:
        source = placeholder._p.get_or_add_pPr().get_or_add_ind()
        for name, value in authored.items():
            source.set(qn(f"w:{name}"), value)
    template, bindings = complete_managed_styles(template, _catalog("zh_CN"))
    managed = bindings.get("body_paragraph")
    assert isinstance(managed, ParagraphStyle)
    managed_element = managed._element
    assert managed_element is not None
    style_before = managed_element.xml
    managed_indent = managed_element.find(f"{qn('w:pPr')}/{qn('w:ind')}")
    assert managed_indent is not None and managed_indent.get(qn("w:firstLineChars")) == "200"
    body_format = extract_body_paragraph_format(template)
    assert body_format is not None
    assert body_format.source_style_id == "HostNormal42"
    assert dict(body_format.indent_attributes) == {qn(f"w:{name}"): value for name, value in authored.items()}
    placeholder = next(p for p in template.paragraphs if p.text == "{{body}}")
    rendered = MdToDocxRenderer(
        template, body_style=placeholder.style, body_paragraph_format=body_format, managed_styles=bindings
    ).render(parse_markdown_text("Body one.\n\nBody two.\n"))
    fill_template(template, {}, rendered, placeholder)
    _, reopened = _save_reopen(template, tmp_path, "mapped-effective-indent.docx")
    for body in reopened.paragraphs:
        assert isinstance(body.style, ParagraphStyle)
        assert body.style.style_id == bindings.style_id("body_paragraph")
        output_indent = body._p.find(f"{qn('w:pPr')}/{qn('w:ind')}")
        if authored:
            assert output_indent is not None
            assert dict(output_indent.attrib) == {
                qn(f"w:{name}"): value for name, value in {**authored, **projection}.items()
            }
        else:
            assert output_indent is None
        # Unrelated font/spacing continue to come from the managed style.
        assert body._p.find(f"{qn('w:pPr')}/{qn('w:spacing')}") is None
        assert not any(run._r.find(qn("w:rPr")) is not None for run in body.runs)
    assert managed_element.xml == style_before


@pytest.mark.parametrize(
    ("authored", "inherited"),
    [
        ({"left": "720"}, {"leftChars": "300", "rightChars": "125"}),
        ({"right": "0"}, {"leftChars": "300", "rightChars": "125"}),
        ({"start": "240", "end": "0"}, {"startChars": "300", "endChars": "125"}),
        ({"left": "720", "leftChars": "0"}, {"leftChars": "300"}),
        ({"leftChars": "200", "left": "720"}, {"leftChars": "300"}),
        ({"firstLine": "360"}, {"hangingChars": "200"}),
        ({"hanging": "360"}, {"firstLineChars": "200"}),
        ({"firstLine": "360"}, {"firstLineChars": "0"}),
        ({"firstLine": "360"}, {"firstLineChars": "000"}),
        ({"hangingChars": "125"}, {"firstLineChars": "200"}),
    ],
)
def test_unmapped_custom_style_retains_original_attribute_inheritance(
    tmp_path: Path, authored: dict[str, str], inherited: dict[str, str]
) -> None:
    template = Document()
    ancestor = template.styles.add_style("IndentAncestor", WD_STYLE_TYPE.PARAGRAPH)
    custom = template.styles.add_style("IndentChild", WD_STYLE_TYPE.PARAGRAPH)
    assert isinstance(ancestor, ParagraphStyle) and isinstance(custom, ParagraphStyle)
    custom.base_style = ancestor
    custom.paragraph_format.space_after = Pt(17)
    custom.font.italic = True
    ancestor_element = ancestor._element
    assert ancestor_element is not None
    inherited_indent = ancestor_element.get_or_add_pPr().get_or_add_ind()
    for name, value in inherited.items():
        inherited_indent.set(qn(f"w:{name}"), value)
    placeholder = template.add_paragraph("{{body}}", style=custom)
    direct_indent = placeholder._p.get_or_add_pPr().get_or_add_ind()
    for name, value in authored.items():
        direct_indent.set(qn(f"w:{name}"), value)
    body_format = extract_body_paragraph_format(template)
    assert body_format is not None
    styles_before = template.styles.element.xml
    body = template.add_paragraph("Projected", style=custom)
    body_format.apply_to(body)
    _, reopened = _save_reopen(template, tmp_path, "inherited-effective-indent.docx")
    output = next(p for p in reopened.paragraphs if p.text == "Projected")
    indent = output._p.find(f"{qn('w:pPr')}/{qn('w:ind')}")
    assert indent is not None
    assert dict(indent.attrib) == {qn(f"w:{name}"): value for name, value in authored.items()}
    assert template.styles.element.xml == styles_before
    assert isinstance(output.style, ParagraphStyle)
    assert output.style.style_id == custom.style_id
    assert output._p.find(f"{qn('w:pPr')}/{qn('w:spacing')}") is None
    assert output.style.paragraph_format.space_after == Pt(17)
    assert output.style.font.italic is True


def test_nearer_character_zero_keeps_older_opposite_indent_cancelled() -> None:
    template = Document()
    ancestor = template.styles.add_style("EarlierCharacters", WD_STYLE_TYPE.PARAGRAPH)
    child = template.styles.add_style("CancelledCharacters", WD_STYLE_TYPE.PARAGRAPH)
    assert isinstance(ancestor, ParagraphStyle) and isinstance(child, ParagraphStyle)
    ancestor_element = ancestor._element
    child_element = child._element
    assert ancestor_element is not None and child_element is not None
    ancestor_element.get_or_add_pPr().get_or_add_ind().set(qn("w:hangingChars"), "200")
    child.base_style = ancestor
    child_element.get_or_add_pPr().get_or_add_ind().set(qn("w:firstLineChars"), "0")
    placeholder = template.add_paragraph("{{body}}", style=child)
    placeholder.paragraph_format.first_line_indent = Inches(0.2)
    body_format = extract_body_paragraph_format(template)
    assert body_format is not None
    body = template.add_paragraph("Projected", style=child)
    body_format.apply_to(body)
    assert dict(body._p.get_or_add_pPr().get_or_add_ind().attrib) == {qn("w:firstLine"): "288"}


def test_unmapped_body_keeps_document_defaults_inherited() -> None:
    template = Document()
    defaults = template.styles.element.find(qn("w:docDefaults"))
    assert defaults is not None
    parent = defaults
    for name in ("pPrDefault", "pPr", "ind"):
        child = parent.find(qn(f"w:{name}"))
        if child is None:
            child = OxmlElement(f"w:{name}")
            parent.append(child)
        parent = child
    parent.set(qn("w:leftChars"), "300")
    parent.set(qn("w:hangingChars"), "200")
    defaults_before = etree.tostring(defaults)
    placeholder = template.add_paragraph("{{body}}")
    placeholder.paragraph_format.left_indent = Pt(0)
    placeholder.paragraph_format.first_line_indent = Pt(24)
    body_format = extract_body_paragraph_format(template)
    assert body_format is not None
    body = template.add_paragraph("Projected")
    body_format.apply_to(body)
    assert dict(body._p.get_or_add_pPr().get_or_add_ind().attrib) == {
        qn("w:left"): "0",
        qn("w:firstLine"): "480",
    }
    assert etree.tostring(defaults) == defaults_before


def test_length_api_projection_preserves_explicit_zero_and_bounds_cyclic_style_lookup() -> None:
    template = Document()
    first = template.styles.add_style("FirstCycleStyle", WD_STYLE_TYPE.PARAGRAPH)
    second = template.styles.add_style("SecondCycleStyle", WD_STYLE_TYPE.PARAGRAPH)
    assert isinstance(first, ParagraphStyle) and isinstance(second, ParagraphStyle)
    first.base_style = second
    second.base_style = first
    second_element = second._element
    assert second_element is not None
    inherited = second_element.get_or_add_pPr().get_or_add_ind()
    inherited.set(qn("w:firstLineChars"), "200")
    inherited.set(qn("w:leftChars"), "300")
    template.add_paragraph("{{body}}", style=first)
    body_format = extract_body_paragraph_format(template)
    assert body_format is not None
    assert not body_format.indent_attributes
    body_format = replace(body_format, left_indent=Pt(0), first_line_indent=Pt(0))
    body = template.add_paragraph("Projected", style=first)
    body_format.apply_to(body)
    assert dict(body._p.get_or_add_pPr().get_or_add_ind().attrib) == {
        qn("w:left"): "0",
        qn("w:firstLine"): "0",
    }


@pytest.mark.parametrize(
    ("authored", "source_values", "target_values", "projection"),
    [
        ({"left": "720"}, {"leftChars": "300"}, {"leftChars": "200"}, {"leftChars": "300"}),
        ({"right": "0"}, {}, {"rightChars": "200"}, {"rightChars": "0"}),
        ({"startChars": "0"}, {"start": "240"}, {"start": "480"}, {"start": "240"}),
        ({"endChars": "0"}, {}, {"end": "720"}, {"end": "0"}),
        ({"left": "720"}, {"leftChars": "300"}, {"leftChars": "300"}, {}),
        ({"leftChars": "125"}, {"left": "100"}, {"left": "200"}, {}),
    ],
)
def test_mapping_protects_only_declared_side_axes(
    authored: dict[str, str], source_values: dict[str, str], target_values: dict[str, str], projection: dict[str, str]
) -> None:
    template = Document()
    normal = template.styles["Normal"]
    target_style = template.styles.add_style("Body Paragraph", WD_STYLE_TYPE.PARAGRAPH)
    assert isinstance(normal, ParagraphStyle) and isinstance(target_style, ParagraphStyle)
    normal_element = normal._element
    target_element = target_style._element
    assert normal_element is not None and target_element is not None
    target_style.base_style = normal
    for style, values in ((normal_element, source_values), (target_element, target_values)):
        indent = style.get_or_add_pPr().get_or_add_ind()
        for name, value in values.items():
            indent.set(qn(f"w:{name}"), value)
    placeholder = template.add_paragraph("{{body}}", style=normal)
    direct = placeholder._p.get_or_add_pPr().get_or_add_ind()
    for name, value in authored.items():
        direct.set(qn(f"w:{name}"), value)
    template, bindings = complete_managed_styles(template, _catalog())
    body_format = extract_body_paragraph_format(template)
    assert body_format is not None
    styles_before = template.styles.element.xml
    rendered = MdToDocxRenderer(
        template, body_style=template.styles["Normal"], body_paragraph_format=body_format, managed_styles=bindings
    ).render(parse_markdown_text("Projected.\n"))
    assert dict(rendered[0]._p.get_or_add_pPr().get_or_add_ind().attrib) == {
        qn(f"w:{name}"): value for name, value in {**authored, **projection}.items()
    }
    assert template.styles.element.xml == styles_before


def test_mapping_uses_source_document_default_character_units() -> None:
    template = Document()
    defaults = template.styles.element.find(f"{qn('w:docDefaults')}/{qn('w:pPrDefault')}/{qn('w:pPr')}")
    assert defaults is not None
    indent = OxmlElement("w:ind")
    indent.set(qn("w:hangingChars"), "150")
    defaults.append(indent)
    placeholder = template.add_paragraph("{{body}}", style="Normal")
    placeholder.paragraph_format.first_line_indent = Pt(24)
    template, bindings = complete_managed_styles(template, _catalog("zh_CN"))
    body_format = extract_body_paragraph_format(template)
    assert body_format is not None
    assert dict(body_format.source_indent_attributes)[qn("w:hangingChars")] == "150"
    rendered = MdToDocxRenderer(
        template, body_style=template.styles["Normal"], body_paragraph_format=body_format, managed_styles=bindings
    ).render(parse_markdown_text("Projected.\n"))
    assert dict(rendered[0]._p.get_or_add_pPr().get_or_add_ind().attrib) == {
        qn("w:firstLine"): "480",
        qn("w:hangingChars"): "150",
    }
