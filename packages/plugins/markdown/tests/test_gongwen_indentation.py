"""Character-unit inheritance must not displace generated attachment columns."""

from __future__ import annotations

from typing import Literal

import pytest
from docx.oxml.ns import qn
from docx.shared import Pt
from tests.support.gongwen_template import attachment_template

from docwen_plugin_markdown.template_filler import fill_template
from docwen_plugin_markdown.template_utils import scan_placeholders

pytestmark = pytest.mark.contract


@pytest.mark.parametrize("inheritance", ["style", "ancestor", "defaults", "direct"])
@pytest.mark.parametrize("location", ["body", "cell"])
@pytest.mark.parametrize("count", [1, 2, 9, 10, 12, 99, 100, 101])
def test_generated_attachment_columns_override_character_indents_without_editing_styles(
    inheritance: Literal["style", "ancestor", "defaults", "direct"],
    location: Literal["body", "cell"],
    count: int,
) -> None:
    document = attachment_template(inheritance=inheritance, location=location)
    original_styles = document.styles.element.xml
    names = [f"材料{index}" for index in range(1, count + 1)]

    fill_template(document, {"附件说明": names}, [], None, scan_placeholders(document))

    paragraphs = document.paragraphs[1:] if location == "body" else document.tables[0].cell(0, 0).paragraphs
    assert len(paragraphs) == count
    # Native Word places this paired-value template at 79 - 47.4 = 31.6 pt.
    # Font-size arithmetic on the character metadata would move it to -16 pt.
    # Ordinary text follows the period and one U+0020 directly;
    # leading padding shrinks as the number gains digits.
    hanging_points = 48 if count == 1 else 64
    text_points = 31.6 + hanging_points
    for paragraph in paragraphs:
        left = paragraph.paragraph_format.left_indent
        assert left is not None and left.twips == round(text_points * 20)
        assert paragraph.paragraph_format.first_line_indent == Pt(-hanging_points)
        indent = paragraph._p.get_or_add_pPr().get_or_add_ind()
        if inheritance != "direct":
            assert indent.get(qn("w:leftChars")) == "0"
            assert indent.get(qn("w:hangingChars")) == "0"
        else:
            assert indent.get(qn("w:leftChars")) is None
            assert indent.get(qn("w:hangingChars")) is None
        assert not list(paragraph._p.iter(qn("w:tab")))
        assert paragraph.style is not None and paragraph.style.name == "Attachment"
    assert document.styles.element.xml == original_styles


def test_direct_character_zero_uses_length_fallback_and_preserves_unrelated_formatting() -> None:
    document = attachment_template()
    target = document.paragraphs[1]
    target.paragraph_format.left_indent = Pt(12)
    target.paragraph_format.right_indent = Pt(8)
    target.paragraph_format.space_after = Pt(5)
    target._p.get_or_add_pPr().get_or_add_ind().set(qn("w:leftChars"), "0")
    target.runs[0].bold = True

    fill_template(document, {"附件说明": ["甲", "乙"]}, [], None, scan_placeholders(document))

    for paragraph in document.paragraphs[1:]:
        assert paragraph.paragraph_format.left_indent == Pt(28.6)
        assert paragraph.paragraph_format.right_indent == Pt(8)
        assert paragraph.paragraph_format.space_after == Pt(5)
        assert paragraph.runs[0].bold is True


def test_nearer_character_zero_cancels_ancestor_hanging_and_keeps_length_fallback() -> None:
    document = attachment_template(inheritance="ancestor")
    target = document.paragraphs[1]
    assert target.style is not None
    style_element = target.style._element
    assert style_element is not None
    style_element.get_or_add_pPr().get_or_add_ind().set(qn("w:firstLineChars"), "0")
    target._p.get_or_add_pPr().get_or_add_ind().set(qn("w:leftChars"), "0")

    fill_template(document, {"附件说明": ["甲", "乙"]}, [], None, scan_placeholders(document))

    for paragraph in document.paragraphs[1:]:
        assert paragraph.paragraph_format.left_indent == Pt(95.6)
        assert paragraph._p.pPr is not None and paragraph._p.pPr.ind is not None
        indent = paragraph._p.pPr.ind
        assert indent.get(qn("w:leftChars")) == "0"
        assert indent.get(qn("w:hangingChars")) is None


@pytest.mark.parametrize("first_line", ["firstLineChars", "hangingChars"])
def test_logical_start_and_inherited_first_line_characters_use_generated_columns(first_line: str) -> None:
    document = attachment_template()
    target = document.paragraphs[1]
    assert target.style is not None
    style_element = target.style._element
    assert style_element is not None
    inherited = style_element.get_or_add_pPr().get_or_add_ind()
    inherited.attrib.clear()
    for name, value in {"startChars": "200", "start": "1580", first_line: "300"}.items():
        inherited.set(qn(f"w:{name}"), value)

    fill_template(document, {"附件说明": ["甲", "乙"]}, [], None, scan_placeholders(document))

    for paragraph in document.paragraphs[1:]:
        text_points = 191 if first_line == "firstLineChars" else 95
        assert paragraph.paragraph_format.left_indent == Pt(text_points)
        assert paragraph._p.pPr is not None and paragraph._p.pPr.ind is not None
        indent = paragraph._p.pPr.ind
        assert indent.get(qn("w:start")) == str(Pt(text_points).twips)
        assert indent.get(qn("w:startChars")) == "0"
        assert indent.get(qn("w:hangingChars")) == "0"
