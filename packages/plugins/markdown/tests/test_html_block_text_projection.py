"""Raw HTML blocks must not disappear or trigger external resource access."""

from docx import Document

from docwen_plugin_markdown.html_text import project_html_block_text
from docwen_plugin_markdown.renderer import MdToDocxRenderer


def test_simple_html_block_keeps_visible_text() -> None:
    lines, lossy = project_html_block_text("<div>Hello <strong>there</strong></div>")
    assert lines == ["Hello there"]
    assert not lossy


def test_html_paragraphs_render_as_word_paragraphs() -> None:
    document = Document()
    renderer = MdToDocxRenderer(document)
    node = {"type": "block_html", "raw": "<p>First</p><p>Second</p>"}
    paragraphs = renderer._handle_block_html(node)
    assert paragraphs is not None
    assert [paragraph.text for paragraph in paragraphs] == ["First", "Second"]


def test_unsafe_script_and_external_image_are_not_processed() -> None:
    raw = "<div>Visible</div><script>secret()</script>"
    raw += "<img src='http://example.test' alt='Image'>"
    lines, lossy = project_html_block_text(raw)
    assert lines == ["Visible", "Image"]
    assert lossy


def test_table_flattening_reports_loss_without_dropping_cell_values() -> None:
    document = Document()
    renderer = MdToDocxRenderer(document)
    raw = "<table><tr><td>0012</td><td>North</td></tr></table>"
    renderer._handle_block_html({"type": "block_html", "raw": raw})
    visible = "\n".join(paragraph.text for paragraph in document.paragraphs)
    assert "0012" in visible
    assert "North" in visible
    assert any(code == "MD2DOCX-HTML-FLATTENED" for code, _ in renderer.warnings)


def test_pure_html_comment_has_no_visible_content() -> None:
    assert project_html_block_text("<!-- hidden -->") == ([], False)
