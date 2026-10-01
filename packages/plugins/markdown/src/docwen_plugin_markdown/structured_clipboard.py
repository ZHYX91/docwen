"""Converters for the internal recursive clipboard-document source format."""

from __future__ import annotations

import csv
from contextlib import suppress
from pathlib import Path
from typing import Any

from docwen_core.clipboard_table_associations import inject_clipboard_table_associations
from docwen_core.docx_semantics import apply_semantic_table_roles
from docwen_core.markdown_extensions import resolve_markdown_extensions
from docwen_core.models.artifact import ARTIFACT_KIND_AUXILIARY, ARTIFACT_KIND_PRIMARY, ArtifactManifest
from docwen_core.models.clipboard_document import (
    ClipboardBlock,
    ClipboardDocument,
    ClipboardHardBreak,
    ClipboardImageRef,
    ClipboardParagraph,
    ClipboardTable,
    ClipboardTableCell,
    ClipboardText,
    clipboard_table_header_shape,
    load_clipboard_document_bytes,
)
from docwen_core.models.result import ConversionDiagnostic, ConversionMetrics, ConversionResult
from docwen_plugin_markdown.structured_clipboard_images import (
    collect_image_occurrences,
    copy_bound_image_artifacts,
    prepare_markdown_images,
    resolve_bound_images,
)
from docwen_plugin_markdown.structured_clipboard_output_images import (
    add_xlsx_images,
    deduplicate_xlsx_png_media,
    image_placeholder,
    markdown_paragraph_blocks,
    missing_image_diagnostics,
    write_csv_image_semantics,
    write_docx_paragraph,
)
from docwen_plugin_markdown.structured_clipboard_projection import (
    has_html_header_associations,
    semantics_headers,
    structural_table_markdown,
    table_semantics_rows,
    table_semantics_text,
)
from docwen_plugin_markdown.structured_clipboard_xlsx import (
    create_table_sheet,
    materialize_order_rows,
    prepare_projection_workbook,
    write_string_matrix,
)

_DOCX_MEDIA = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
_XLSX_MEDIA = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def _load(context: Any) -> ClipboardDocument:
    return load_clipboard_document_bytes(Path(context.workspace.input_path).read_bytes())


def _paragraph_projection(paragraph: ClipboardParagraph, resources: dict[str, Any] | None = None) -> str:
    parts: list[str] = []
    resources = resources or {}
    for inline in paragraph.inlines:
        if isinstance(inline, ClipboardText):
            parts.append(inline.value)
        elif isinstance(inline, ClipboardHardBreak):
            parts.append("\n")
        elif inline.resource_id is not None and inline.resource_id in resources:
            parts.append(f"[Image: {resources[inline.resource_id].suggested_name}]")
        else:
            parts.append(image_placeholder(inline))
    return "".join(parts)


def _clear_docx_cell(cell: Any) -> None:
    tc = cell._tc
    for child in list(tc):
        if child.tag.rsplit("}", 1)[-1] != "tcPr":
            tc.remove(child)


def _remove_trailing_empty_paragraph(cell: Any) -> None:
    children = list(cell._tc)
    if not children:
        return
    last = children[-1]
    if last.tag.rsplit("}", 1)[-1] != "p":
        return
    text_nodes = [node.text or "" for node in last.iter() if node.tag.rsplit("}", 1)[-1] == "t"]
    if not "".join(text_nodes):
        cell._tc.remove(last)


def _set_docx_header_rows(table: Any, header_rows: int) -> None:
    if header_rows <= 0:
        return
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    for row in table.rows[:header_rows]:
        tr_pr = row._tr.get_or_add_trPr()
        marker = OxmlElement("w:tblHeader")
        marker.set(qn("w:val"), "true")
        tr_pr.append(marker)


def _render_docx_blocks(
    container: Any,
    blocks: tuple[ClipboardBlock, ...],
    resources: dict[str, Any],
) -> list[Any]:
    elements: list[Any] = []
    previous_table = False
    for block in blocks:
        if previous_table and hasattr(container, "_tc"):
            _remove_trailing_empty_paragraph(container)
        if isinstance(block, ClipboardParagraph):
            paragraph = write_docx_paragraph(container, block, resources)
            elements.append(paragraph._element)
            previous_table = False
            continue
        table = _render_docx_table(container, block, resources)
        elements.append(table._element)
        previous_table = True
    return elements


def _render_docx_table(container: Any, model: ClipboardTable, resources: dict[str, Any]) -> Any:
    table = container.add_table(rows=model.row_count, cols=model.column_count)
    with suppress(KeyError, ValueError):
        table.style = "Table Grid"

    for cell_model in model.cells:
        if cell_model.row_span > 1 or cell_model.column_span > 1:
            start = table.cell(cell_model.row, cell_model.column)
            end = table.cell(
                cell_model.row + cell_model.row_span - 1,
                cell_model.column + cell_model.column_span - 1,
            )
            start.merge(end)

    for cell_model in model.cells:
        cell = table.cell(cell_model.row, cell_model.column)
        _clear_docx_cell(cell)
        _render_docx_blocks(cell, cell_model.blocks, resources)
        if not list(cell._tc) or list(cell._tc)[-1].tag.rsplit("}", 1)[-1] != "p":
            cell.add_paragraph()
    header_rows, header_columns = clipboard_table_header_shape(model)
    apply_semantic_table_roles(
        table,
        header_rows=header_rows,
        header_columns=header_columns,
        repeat_header="always" if header_rows else "inherit",
    )
    return table


def _place_docx_body(document: Any, elements: list[Any], placeholder: Any | None) -> None:
    body = document.element.body
    for element in elements:
        if element.getparent() is body:
            body.remove(element)
    if placeholder is not None:
        placeholder_element = placeholder._element
        try:
            index = list(body).index(placeholder_element)
        except ValueError:
            index = len(body) - (1 if body.sectPr is not None else 0)
        for offset, element in enumerate(elements):
            body.insert(index + offset, element)
        if placeholder_element.getparent() is body:
            body.remove(placeholder_element)
        return
    section = body.sectPr
    index = body.index(section) if section is not None else len(body)
    for offset, element in enumerate(elements):
        body.insert(index + offset, element)


def convert_clipboard_document_to_docx(context: Any) -> ConversionResult:
    from docwen_plugin_markdown.template_filler import fill_template
    from docwen_plugin_markdown.template_utils import find_body_placeholder, resolve_template, scan_placeholders

    document_model = _load(context)
    resources = resolve_bound_images(context, document_model)
    context.cancellation.check()
    document = resolve_template(context.request.options.get("template_name"))
    placeholder = find_body_placeholder(document)
    # Clear template-only fields before authored clipboard blocks are appended,
    # so authored text resembling {{placeholders}} is never reinterpreted.
    template_placeholders = scan_placeholders(document)
    fill_template(document, {}, [], None, placeholder_map=template_placeholders)
    elements = _render_docx_blocks(document, document_model.blocks, resources)
    _place_docx_body(document, elements, placeholder)
    context.cancellation.check()
    output = context.workspace.create_artifact_path(ARTIFACT_KIND_PRIMARY, ".docx")
    document.save(output)
    inject_clipboard_table_associations(Path(output), document_model)
    artifact = ArtifactManifest(
        artifact_id="clipboard-document-docx",
        kind=ARTIFACT_KIND_PRIMARY,
        staging_path=output,
        suggested_name=f"{context.request.source_stem}.docx",
        media_type=_DOCX_MEDIA,
        is_primary=True,
    )
    return ConversionResult(
        task_id=context.request.request_id,
        success=True,
        artifacts=[artifact],
        diagnostics=missing_image_diagnostics(document_model),
        metrics=ConversionMetrics(input_bytes=Path(context.workspace.input_path).stat().st_size),
    )


class _OrderRecorder:
    def __init__(self) -> None:
        self.rows: list[tuple[str, str, str, str, str, str, str]] = []
        self._sequence = 0
        self._table_sequence = 0
        self.table_ids: dict[int, str] = {}
        self.tables: list[tuple[str, ClipboardTable, str, str]] = []
        self.paragraphs: dict[str, ClipboardParagraph] = {}

    def add(
        self, kind: str, text: str = "", sheet: str = "", parent: str = "", anchor: str = "", child: str = ""
    ) -> str:
        self._sequence += 1
        sequence = str(self._sequence)
        self.rows.append((sequence, kind, text, sheet, parent, anchor, child))
        return sequence

    def table_id(self, table: ClipboardTable) -> str:
        key = id(table)
        existing = self.table_ids.get(key)
        if existing is not None:
            return existing
        self._table_sequence += 1
        value = f"T{self._table_sequence}"
        self.table_ids[key] = value
        return value


def _record_order(
    blocks: tuple[ClipboardBlock, ...],
    recorder: _OrderRecorder,
    *,
    parent_table: str = "",
    anchor: str = "",
    resources: dict[str, Any] | None = None,
) -> None:
    for block in blocks:
        if isinstance(block, ClipboardParagraph):
            kind = "paragraph" if not parent_table else "cell_paragraph"
            sequence = recorder.add(
                kind,
                _paragraph_projection(block, resources),
                parent=parent_table,
                anchor=anchor,
            )
            recorder.paragraphs[sequence] = block
            continue
        table_id = recorder.table_id(block)
        sheet = f"Table {table_id[1:]}"
        recorder.tables.append((table_id, block, parent_table, anchor))
        recorder.add("table_ref", sheet=sheet, parent=parent_table, anchor=anchor, child=table_id)
        for cell in block.cells:
            cell_anchor = f"R{cell.row + 1}C{cell.column + 1}"
            _record_order(
                cell.blocks,
                recorder,
                parent_table=table_id,
                anchor=cell_anchor,
                resources=resources,
            )


def _xlsx_cell_value(
    cell: ClipboardTableCell,
    recorder: _OrderRecorder,
    table_titles: dict[str, str] | None = None,
    resources: dict[str, Any] | None = None,
) -> str:
    parts: list[str] = []
    titles = table_titles or {}
    for block in cell.blocks:
        if isinstance(block, ClipboardParagraph):
            parts.append(_paragraph_projection(block, resources))
        else:
            table_id = recorder.table_id(block)
            parts.append(f"[Nested table {titles.get(table_id, table_id)}]")
    return "\n".join(parts)


def _unique_sheet_name(workbook: Any, requested: str) -> str:
    base = requested[:31] or "Sheet"
    name = base
    index = 2
    while name in workbook.sheetnames:
        suffix = f" {index}"
        name = base[: 31 - len(suffix)] + suffix
        index += 1
    return name


def _new_projection_workbook(template_path: str | None) -> Any:
    from openpyxl import Workbook, load_workbook

    workbook = load_workbook(template_path) if template_path else Workbook()
    first = workbook.worksheets[0]
    for sheet in list(workbook.worksheets[1:]):
        workbook.remove(sheet)
    for merged in list(first.merged_cells.ranges):
        first.unmerge_cells(str(merged))
    for row in first.iter_rows():
        for cell in row:
            cell.value = None
    first.title = "Document Order"
    return workbook


def convert_clipboard_document_to_xlsx(context: Any) -> ConversionResult:
    document = _load(context)
    resources = resolve_bound_images(context, document)
    recorder = _OrderRecorder()
    _record_order(document.blocks, recorder, resources=resources)
    projection = prepare_projection_workbook(context.request.options.get("template_name"))

    table_sheets: dict[str, Any] = {}
    for table_id, table_model, _parent, _anchor in recorder.tables:
        table_sheets[table_id] = create_table_sheet(
            projection,
            table_id=table_id,
            requested_name=f"Table {table_id[1:]}",
            row_count=table_model.row_count,
            column_count=table_model.column_count,
        )

    for table_id, table_model, _parent, _anchor in recorder.tables:
        sheet = table_sheets[table_id]
        for cell_model in table_model.cells:
            output_cell = sheet.cell(
                cell_model.row + 1,
                cell_model.column + 1,
                _xlsx_cell_value(cell_model, recorder, projection.table_titles, resources),
            )
            output_cell.data_type = "s"
            output_cell.number_format = "@"
        for cell_model in table_model.cells:
            if cell_model.row_span > 1 or cell_model.column_span > 1:
                sheet.merge_cells(
                    start_row=cell_model.row + 1,
                    start_column=cell_model.column + 1,
                    end_row=cell_model.row + cell_model.row_span,
                    end_column=cell_model.column + cell_model.column_span,
                )

    order_headers = ("Sequence", "Kind", "Text", "Sheet", "Parent Table", "Anchor", "Child Table")
    write_string_matrix(
        projection.order_sheet,
        order_headers,
        materialize_order_rows(recorder.rows, projection.table_titles),
    )
    write_string_matrix(
        projection.semantics_sheet,
        semantics_headers(),
        table_semantics_rows(recorder.tables, projection.table_titles),
    )

    occurrences = collect_image_occurrences(document, recorder.table_id)
    add_xlsx_images(projection, occurrences, resources, table_sheets)
    diagnostics = missing_image_diagnostics(document)
    if any(occurrence.image.resource_id is not None for occurrence in occurrences):
        diagnostics.append(
            ConversionDiagnostic(
                level="warning",
                code="CLIPBOARD-XLSX-IMAGE-PROJECTION",
                message=(
                    "Clipboard images are preserved as worksheet drawings with visible Image Semantics; "
                    "XLSX does not provide native inline or nested-table-cell image semantics."
                ),
            )
        )
    if any(parent for _table_id, _table, parent, _anchor in recorder.tables) or any(
        isinstance(block, ClipboardParagraph) for block in document.blocks
    ):
        diagnostics.append(
            ConversionDiagnostic(
                level="warning",
                code="CLIPBOARD-XLSX-DOCUMENT-PROJECTION",
                message=(
                    "Document text and nested-table relationships are preserved in the Document Order worksheet; "
                    "XLSX has no native nested-table model."
                ),
            )
        )
    if any(has_html_header_associations(table) for _table_id, table, _parent, _anchor in recorder.tables):
        diagnostics.append(
            ConversionDiagnostic(
                level="warning",
                code="CLIPBOARD-XLSX-HEADER-ASSOCIATIONS-PROJECTED",
                message=(
                    "HTML table header scope/id/headers associations are preserved in the visible Table Semantics "
                    "worksheet because XLSX has no equivalent native association graph."
                ),
            )
        )

    output = context.workspace.create_artifact_path(ARTIFACT_KIND_PRIMARY, ".xlsx")
    from docwen_plugin_markdown.structured_clipboard_strings import save_workbook_preserving_empty_strings

    save_workbook_preserving_empty_strings(projection.workbook, output)
    deduplicate_xlsx_png_media(Path(output))
    artifact = ArtifactManifest(
        artifact_id="clipboard-document-xlsx",
        kind=ARTIFACT_KIND_PRIMARY,
        staging_path=output,
        suggested_name=f"{context.request.source_stem}.xlsx",
        media_type=_XLSX_MEDIA,
        is_primary=True,
    )
    return ConversionResult(
        task_id=context.request.request_id,
        success=True,
        artifacts=[artifact],
        diagnostics=diagnostics,
    )


def _markdown_fence(text: str) -> str:
    longest = max((len(match.group(0)) for match in __import__("re").finditer(r"`+", text)), default=0)
    fence = "`" * max(3, longest + 1)
    return f"{fence}text\n{text}\n{fence}"


def _markdown_projection(
    document: ClipboardDocument,
    *,
    structural_tables: bool,
    image_plan: Any,
    resources: dict[str, Any],
) -> tuple[str, list[ConversionDiagnostic]]:
    recorder = _OrderRecorder()
    _record_order(document.blocks, recorder, resources=resources)
    lines = ["# Clipboard document", ""]
    for sequence, kind, text, sheet, parent, anchor, child in recorder.rows:
        if kind in {"paragraph", "cell_paragraph"}:
            lines.extend([f"## {sequence}. {kind}", ""])
            lines.extend(
                markdown_paragraph_blocks(
                    recorder.paragraphs[sequence],
                    image_plan,
                    resources,
                    _markdown_fence,
                )
            )
        else:
            lines.extend(
                [
                    f"## {sequence}. table",
                    "",
                    _markdown_fence(f"table={child}\nsheet={sheet}\nparent={parent or '-'}\nanchor={anchor or '-'}"),
                    "",
                ]
            )

    fallback_reasons: list[str] = []
    association_projection = False
    for table_id, table, parent, anchor in recorder.tables:
        lines.extend([f"## Table {table_id}", ""])
        rendered = None
        reason = "extension_disabled"
        if structural_tables:
            rendered, reason = structural_table_markdown(table)
        if rendered is not None:
            lines.extend([rendered, ""])
        else:
            if structural_tables:
                fallback_reasons.append(reason)
            lines.append(
                _markdown_fence(
                    "\n".join(
                        [
                            (
                                f"rows={table.row_count} columns={table.column_count} "
                                f"parent={parent or '-'} anchor={anchor or '-'}"
                            ),
                            *[
                                (
                                    f"R{cell.row + 1}C{cell.column + 1} "
                                    f"span={cell.row_span}x{cell.column_span} "
                                    f"value={_xlsx_cell_value(cell, recorder, resources=resources)}"
                                )
                                for cell in table.cells
                            ],
                        ]
                    )
                )
            )
            lines.append("")
        lines.extend(
            [
                "### Table semantics",
                "",
                _markdown_fence(table_semantics_text(table_id, table)),
                "",
            ]
        )
        association_projection = association_projection or has_html_header_associations(table)

    diagnostics = missing_image_diagnostics(document)
    if not structural_tables:
        diagnostics.append(
            ConversionDiagnostic(
                level="warning",
                code="CLIPBOARD-MARKDOWN-STRUCTURE-PROJECTION",
                message=(
                    "Structural Tables output is disabled; recursive clipboard tables were emitted as visible "
                    "literal structure with their values, spans, and roles."
                ),
            )
        )
    elif fallback_reasons:
        diagnostics.append(
            ConversionDiagnostic(
                level="warning",
                code="CLIPBOARD-MARKDOWN-STRUCTURAL-FALLBACK",
                message=(
                    "Some clipboard tables cannot be represented losslessly by the selected Structural Tables "
                    f"dialect and were kept as literal structure: {', '.join(sorted(set(fallback_reasons)))}."
                ),
            )
        )
    if association_projection:
        diagnostics.append(
            ConversionDiagnostic(
                level="warning",
                code="CLIPBOARD-MARKDOWN-HEADER-ASSOCIATIONS-PROJECTED",
                message=(
                    "HTML scope/id/headers associations have no native Structural Tables spelling and are "
                    "preserved in each visible Table semantics block."
                ),
            )
        )
    return "\n".join(lines).rstrip() + "\n", diagnostics


def convert_clipboard_document_to_markdown(context: Any) -> ConversionResult:
    document = _load(context)
    resources = resolve_bound_images(context, document)
    image_plan = prepare_markdown_images(context, resources)
    extensions = resolve_markdown_extensions(
        context.request.options,
        context.config,
        direction="output",
    )
    text, diagnostics = _markdown_projection(
        document,
        structural_tables=extensions.structural_tables,
        image_plan=image_plan,
        resources=resources,
    )
    output = context.workspace.create_artifact_path(ARTIFACT_KIND_PRIMARY, ".md")
    Path(output).write_text(text, encoding="utf-8", newline="\n")
    artifact = ArtifactManifest(
        artifact_id="clipboard-document-markdown",
        kind=ARTIFACT_KIND_PRIMARY,
        staging_path=output,
        suggested_name=f"{context.request.source_stem}.md",
        media_type="text/markdown",
        is_primary=True,
    )
    return ConversionResult(
        task_id=context.request.request_id,
        success=True,
        artifacts=[artifact, *image_plan.artifacts],
        diagnostics=diagnostics,
    )


def _table_matrix(
    table: ClipboardTable,
    recorder: _OrderRecorder,
    resources: dict[str, Any] | None = None,
) -> list[list[str]]:
    matrix = [["" for _column in range(table.column_count)] for _row in range(table.row_count)]
    for cell in table.cells:
        matrix[cell.row][cell.column] = _xlsx_cell_value(cell, recorder, resources=resources)
    return matrix


def convert_clipboard_document_to_csv(context: Any) -> ConversionResult:
    document = _load(context)
    resources = resolve_bound_images(context, document)
    recorder = _OrderRecorder()
    _record_order(document.blocks, recorder, resources=resources)
    artifacts: list[ArtifactManifest] = []

    order_path = context.workspace.create_artifact_path(ARTIFACT_KIND_PRIMARY, ".csv")
    with Path(order_path).open("w", encoding="utf-8-sig", newline="") as stream:
        writer = csv.writer(stream)
        writer.writerow(("Sequence", "Kind", "Text", "Sheet", "Parent Table", "Anchor", "Child Table"))
        writer.writerows(recorder.rows)
    artifacts.append(
        ArtifactManifest(
            artifact_id="clipboard-document-order",
            kind=ARTIFACT_KIND_PRIMARY,
            staging_path=order_path,
            suggested_name=f"{context.request.source_stem}-document-order.csv",
            media_type="text/csv",
            is_primary=True,
        )
    )

    semantics_path = context.workspace.create_artifact_path(ARTIFACT_KIND_AUXILIARY, ".csv")
    with Path(semantics_path).open("w", encoding="utf-8-sig", newline="") as stream:
        writer = csv.writer(stream)
        writer.writerow(semantics_headers())
        writer.writerows(table_semantics_rows(recorder.tables))
    artifacts.append(
        ArtifactManifest(
            artifact_id="clipboard-table-semantics",
            kind=ARTIFACT_KIND_AUXILIARY,
            staging_path=semantics_path,
            suggested_name=f"{context.request.source_stem}-table-semantics.csv",
            media_type="text/csv",
        )
    )

    for table_id, table, _parent, _anchor in recorder.tables:
        table_path = context.workspace.create_artifact_path(ARTIFACT_KIND_AUXILIARY, ".csv")
        with Path(table_path).open("w", encoding="utf-8-sig", newline="") as stream:
            csv.writer(stream).writerows(_table_matrix(table, recorder, resources))
        artifacts.append(
            ArtifactManifest(
                artifact_id=f"clipboard-{table_id.lower()}",
                kind=ARTIFACT_KIND_AUXILIARY,
                staging_path=table_path,
                suggested_name=f"{context.request.source_stem}-{table_id}.csv",
                media_type="text/csv",
            )
        )

    occurrences = collect_image_occurrences(document, recorder.table_id)
    image_semantics_path = Path(
        context.workspace.create_artifact_path(ARTIFACT_KIND_AUXILIARY, ".csv")
    )
    write_csv_image_semantics(image_semantics_path, occurrences, resources)
    artifacts.append(
        ArtifactManifest(
            artifact_id="clipboard-image-semantics",
            kind=ARTIFACT_KIND_AUXILIARY,
            staging_path=str(image_semantics_path),
            suggested_name=f"{context.request.source_stem}-image-semantics.csv",
            media_type="text/csv",
        )
    )
    artifacts.extend(copy_bound_image_artifacts(context, resources))

    diagnostics = missing_image_diagnostics(document)
    if any(occurrence.image.resource_id is not None for occurrence in occurrences):
        diagnostics.append(
            ConversionDiagnostic(
                level="warning",
                code="CLIPBOARD-CSV-IMAGE-PROJECTION",
                message=(
                    "CSV cannot embed images; bound PNG resources are exported separately and "
                    "their occurrence locations are preserved in image-semantics.csv."
                ),
            )
        )
    diagnostics.append(
        ConversionDiagnostic(
            level="warning",
            code="CLIPBOARD-CSV-STRUCTURE-PROJECTION",
            message=(
                "CSV cannot natively represent merged/nested tables or HTML header associations; anchor values "
                "are exported per table and spans/roles/scope/id/headers are preserved in table-semantics.csv."
            ),
        )
    )
    return ConversionResult(
        task_id=context.request.request_id,
        success=True,
        artifacts=artifacts,
        diagnostics=diagnostics,
    )


__all__ = [
    "convert_clipboard_document_to_csv",
    "convert_clipboard_document_to_docx",
    "convert_clipboard_document_to_markdown",
    "convert_clipboard_document_to_xlsx",
]
