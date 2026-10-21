"""Built-in Gongwen field values and Word template layout.

YAML is source data.  Normalization returns a separate display projection;
attachment labels, numbering and indentation exist only in the Word output.
"""

from __future__ import annotations

import datetime as dt
import re
from collections.abc import Mapping, Sequence
from copy import deepcopy
from typing import Any, cast

from docx.enum.text import WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.oxml.text.paragraph import CT_P
from docx.shared import Emu, Pt, Twips
from docx.text.paragraph import Paragraph
from docx.text.run import Run

from docwen_plugin_markdown.template_utils import inherited_indent_attributes
from docwen_plugin_markdown.yaml_processor import is_empty_field_value

_PLACEHOLDER = re.compile(r"\{\{\s*([^{}\r\n]+?)\s*\}\}")
_ATTACHMENT = re.compile(r"\{\{\s*附件说明\s*\}\}")
_ATTACHMENT_NUMBER = re.compile(
    r"^(?:[一二三四五六七八九十百]+[、.]|"
    r"[（(](?:[一二三四五六七八九十百]+|\d+)[）)]|"
    r"\d+[.．、]|"
    r"[㈠㈡㈢㈣㈤㈥㈦㈧㈨㈩①②③④⑤⑥⑦⑧⑨⑩⑪⑫⑬⑭⑮⑯⑰⑱⑲⑳"
    r"㉑㉒㉓㉔㉕㉖㉗㉘㉙㉚㉛㉜㉝㉞㉟㊱㊲㊳㊴㊵㊶㊷㊸㊹㊺㊻㊼㊽㊾㊿"
    r"⓪⓵⓶⓷⓸⓹⓺⓻⓼⓽⓾⓿❶❷❸❹❺❻❼❽❾❿⓫⓬⓭⓮⓯⓰⓱⓲⓳⓴])\s*"
)
_PARAGRAPH_GROUPS = (
    ("密级和保密期限",),
    ("紧急程度",),
    ("发文字号",),
    ("公开方式",),
    ("主送机关",),
    ("签发人",),
    ("附注",),
    ("抄送机关",),
    ("附件说明",),
    ("份号", "发文字号"),
)
_ROW_GROUPS = (("抄送机关",), ("印发机关", "印发日期"))
_LABELS = {"抄送机关": ("抄送",), "印发机关": ("版记",), "附件说明": ("附件",)}
_DECORATION = re.compile(r"[\s:：;；,，、.。()（）\[\]【】/／—\-]*\Z")
_NON_TEXT_CONTENT = frozenset(
    qn(f"w:{name}")
    for name in (
        "drawing",
        "pict",
        "object",
        "fldSimple",
        "fldChar",
        "sym",
        "footnoteReference",
        "endnoteReference",
        "sectPr",
        "br",
        "cr",
    )
)


class GongwenFieldPlacementError(ValueError):
    """A Gongwen field cannot be safely rendered or reviewed at its location."""

    diagnostic_code = "MD2DOCX-GONGWEN-FIELD-PLACEMENT"


def prepare_gongwen_fields(source: Mapping[str, Any]) -> dict[str, Any]:
    """Build display values without changing the caller's YAML or nested data."""
    values = dict(source)
    if "附件说明" in source:
        values["附件说明"] = _attachment_names(source["附件说明"])
    for key in ("签发人", "主送机关", "抄送机关"):
        if key in source:
            value = source[key]
            values[key] = value if isinstance(value, str) else "、".join(_field_list_items(key, value))
    if "附注" in source:
        values["附注"] = _notes(source["附注"])
    for key, suffix in (("成文日期", ""), ("印发日期", "印发")):
        if key in source:
            values[key] = _date(source[key], suffix)
    return values


def _field_list_items(key: str, value: Any) -> tuple[str, ...]:
    """Accept a flat list of authored strings; never infer names or categories."""
    if value is None:
        return ()
    if not isinstance(value, (list, tuple)):
        raise ValueError(f"“{key}”请填写一段文字，或每项都是文字的列表。")
    items = []
    for item in value:
        if item is None:
            continue
        if not isinstance(item, str):
            raise ValueError(f"“{key}”的每一项都应是文字，不支持嵌套列表、数字或对象。")
        if item.strip():
            items.append(item.strip())
    return tuple(items)


def gongwen_review_notes(source: Mapping[str, Any]) -> dict[str, str]:
    """Explain only multi-item list rendering, not already-authored punctuation."""
    notes = {}
    for key in ("签发人", "主送机关", "抄送机关"):
        value = source.get(key)
        if isinstance(value, str) or len(_field_list_items(key, value)) < 2:
            continue
        if key == "签发人":
            notes[key] = (
                "这里有多位签发人，已暂用顿号连接。请在 Word 中自行调整姓名的换行和对齐，"
                "并核对姓名顺序；本程序没有自动排版。"
            )
        else:
            notes[key] = f"{key}有多项，已暂用顿号连接。请核对哪些地方应使用顿号、哪些应改成逗号，并检查机关排列顺序。"
    return notes


def validate_review_placeholders(doc: Any, placeholders: Mapping[str, Sequence[Any]], notes: Mapping[str, str]) -> None:
    """Do not silently drop required review comments in unsupported Word parts."""
    for key in notes:
        for paragraph in placeholders.get(key, ()):
            ancestor = paragraph._p.getparent()
            while ancestor is not None and ancestor is not doc.element.body:
                if ancestor.tag not in {qn("w:tc"), qn("w:tr"), qn("w:tbl")}:
                    break
                ancestor = ancestor.getparent()
            if ancestor is not doc.element.body:
                raise GongwenFieldPlacementError(
                    f"“{key}”有多个条目，需要添加核对批注。请将占位符放在正文或表格中，"
                    "或改用已核对标点和排版的单段文字。"
                )


def validate_attachment_placeholders(doc: Any, placeholders: Mapping[str, Sequence[Any]]) -> None:
    """Accept standalone, ordinary paragraphs in the main body or table cells."""
    for paragraph in placeholders.get("附件说明", ()):
        element = paragraph._p
        ancestor = element.getparent()
        allowed = ancestor is doc.element.body or (ancestor is not None and ancestor.tag == qn("w:tc"))
        while ancestor is not None and ancestor is not doc.element.body:
            if ancestor.tag not in {qn("w:tc"), qn("w:tr"), qn("w:tbl")}:
                allowed = False
                break
            ancestor = ancestor.getparent()
        # Restrict the anchor, not inserted YAML values or rendered Markdown.
        if (
            not allowed
            or ancestor is not doc.element.body
            or _ATTACHMENT.fullmatch(paragraph.text.strip()) is None
            or any(child.tag not in {qn("w:pPr"), qn("w:r")} for child in element)
            or any(node.tag in _NON_TEXT_CONTENT - {qn("w:sectPr")} for node in element.iter())
        ):
            raise GongwenFieldPlacementError(
                "附件说明占位符必须独占正文或表格单元格中的普通段落；"
                "不支持与其他内容混排，或放在页眉、页脚、文本框等位置。"
            )


def remove_empty_gongwen_fields(doc: Any, values: Mapping[str, Any], placeholders: Mapping[str, Sequence[Any]]) -> None:
    """Remove only empty field-only containers from the original template."""
    for group in _ROW_GROUPS:
        for row in _matching_containers(placeholders, group, "tr"):
            if not _removable_container(row, group, values):
                continue
            table = row.getparent()
            if table is None:
                continue
            table.remove(row)
            if not table.findall(qn("w:tr")):
                parent = table.getparent()
                if parent is not None:
                    parent.remove(table)
                    _ensure_cell_paragraph(parent)
    for group in _PARAGRAPH_GROUPS:
        for paragraph in _matching_containers(placeholders, group, "p"):
            if _removable_container(paragraph, group, values):
                _remove_paragraph(paragraph)


def render_attachment_placeholders(names: Sequence[str], placeholders: Mapping[str, Sequence[Any]]) -> None:
    """Expand each original attachment anchor in its own body/cell container."""
    for target in placeholders.get("附件说明", ()):
        parent = target._p.getparent()
        if parent is None:
            continue
        if not names:
            # A section-bearing empty paragraph must keep its section break.
            for child in list(target._p):
                if child.tag != qn("w:pPr"):
                    target._p.remove(child)
            _remove_paragraph(target._p)
            continue

        source_run = next((run for run in target.runs if "{{" in run.text or "{" in run.text), None)
        if source_run is None:
            source_run = next(iter(target.runs), None)
        char_width = _font_size(target, source_run)
        inherited = inherited_indent_attributes(target)
        base_indent = _left_indent(target, char_width, inherited)
        label_width = 3 * char_width
        number_width = int((len(str(len(names))) + 2) * char_width / 2) if len(names) > 1 else 0
        text_indent = base_indent + label_width + number_width
        insertion_index = parent.index(target._p)

        for index, name in enumerate(names, 1):
            element = cast(CT_P, OxmlElement("w:p"))
            if target._p.pPr is not None:
                element.append(deepcopy(target._p.pPr))
                # A section boundary belongs after the complete expansion.
                properties = element.get_or_add_pPr()
                section = properties.find(qn("w:sectPr"))
                if section is not None and index < len(names):
                    properties.remove(section)
            paragraph = Paragraph(element, target._parent)
            run = paragraph.add_run()
            if source_run is not None and source_run._r.rPr is not None:
                run._r.insert(0, deepcopy(source_run._r.rPr))
            if len(names) == 1:
                run.text = f"附件：{name}"
            else:
                run.text = f"{'附件：' if index == 1 else ''}{index}.\t{name}"
                _attachment_tab_stop(paragraph, target, text_indent)
            paragraph.paragraph_format.left_indent = text_indent
            paragraph.paragraph_format.first_line_indent = -(label_width + number_width if index == 1 else number_width)
            # Character-based indents would override the physical layout above.
            indent = paragraph._p.get_or_add_pPr().get_or_add_ind()
            for attribute in ("leftChars", "startChars", "firstLineChars", "hangingChars", "start"):
                key = qn(f"w:{attribute}")
                if key in indent.attrib:
                    del indent.attrib[key]
            # Removing a direct value exposes style/default character units.
            # Word requires a local zero to cancel those earlier declarations;
            # keep the template's shared styles and unrelated axes untouched.
            for attribute in ("leftChars", "startChars"):
                key = qn(f"w:{attribute}")
                if int(inherited.get(key, "0")):
                    indent.set(key, "0")
            if any(int(inherited.get(qn(f"w:{name}"), "0")) for name in ("firstLineChars", "hangingChars")):
                indent.set(qn("w:hangingChars"), "0")
            if any(qn(f"w:{name}") in inherited for name in ("start", "startChars")):
                indent.set(qn("w:start"), str(Emu(text_indent).twips))
            parent.insert(insertion_index + index - 1, element)
        parent.remove(target._p)


def _attachment_names(value: Any) -> tuple[str, ...]:
    names = []
    for item in _display_items(value):
        text = re.sub(r"^附件\s*[:：]\s*", "", item)
        cleaned = _ATTACHMENT_NUMBER.sub("", text, count=1).strip()
        names.append(cleaned or text)
    return tuple(names)


def _display_items(value: Any):
    if is_empty_field_value(value):
        return
    if isinstance(value, (list, tuple)):
        for item in value:
            yield from _display_items(item)
    else:
        yield str(value).strip()


def _notes(value: Any) -> str:
    if is_empty_field_value(value):
        return ""
    if isinstance(value, (Mapping, list, tuple, set)):
        raise ValueError("“附注”请填写一段文字，不要填写列表或对象。")
    text = str(value).strip()
    if len(text) < 2 or (text[0], text[-1]) not in {("（", "）"), ("(", ")")}:
        return text
    depth = 0
    for index, character in enumerate(text):
        if character == text[0]:
            depth += 1
        elif character == text[-1]:
            depth -= 1
        if depth == 0 and index < len(text) - 1:
            return text
    return text[1:-1] if depth == 0 else text


def _date(value: Any, suffix: str) -> str:
    if is_empty_field_value(value):
        return ""
    if isinstance(value, (Mapping, list, tuple, set)):
        raise ValueError("成文日期和印发日期应填写单个日期或文字，不要填写列表或对象。")
    if isinstance(value, dt.date):
        parsed = value
    else:
        parsed = None
        for format_string in ("%Y-%m-%d", "%Y/%m/%d", "%Y年%m月%d日", "%Y.%m.%d", "%Y年%m月%d号"):
            try:
                parsed = dt.datetime.strptime(str(value).strip(), format_string)
                break
            except ValueError:
                continue
    return f"{parsed.year}年{parsed.month}月{parsed.day}日{suffix}" if parsed is not None else str(value)


def _matching_containers(placeholders: Mapping[str, Sequence[Any]], group: tuple[str, ...], tag: str) -> set[Any]:
    common = None
    for key in group:
        containers = set()
        for paragraph in placeholders.get(key, ()):
            element = paragraph._p
            while element is not None and element.tag != qn(f"w:{tag}"):
                element = element.getparent()
            if element is not None:
                containers.add(element)
        common = containers if common is None else common & containers
    return common or set()


def _removable_container(container: Any, group: tuple[str, ...], values: Mapping[str, Any]) -> bool:
    if any(not is_empty_field_value(values.get(key)) for key in group):
        return False
    if any(node.tag in _NON_TEXT_CONTENT for node in container.iter()):
        return False
    # Never delete a containing row because of fields in its nested table.
    if container.tag == qn("w:tr") and any(node.tag == qn("w:tbl") for node in container.iter()):
        return False
    paragraphs = [container] if container.tag == qn("w:p") else container.iter(qn("w:p"))
    labels = set(group)
    labels.update(label for key in group for label in _LABELS.get(key, ()))
    for paragraph in paragraphs:
        text = "".join(node.text or "" for node in paragraph.iter(qn("w:t")))
        fields = {match.group(1).strip() for match in _PLACEHOLDER.finditer(text)}
        if any(not is_empty_field_value(values.get(key)) for key in fields):
            return False
        remaining = _PLACEHOLDER.sub("", text)
        for label in sorted(labels, key=len, reverse=True):
            remaining = remaining.replace(label, "")
        if _DECORATION.fullmatch(remaining) is None:
            return False
    return True


def _remove_paragraph(element: Any) -> None:
    parent = element.getparent()
    if parent is None or element.find(f"{qn('w:pPr')}/{qn('w:sectPr')}") is not None:
        return
    parent.remove(element)
    _ensure_cell_paragraph(parent)


def _ensure_cell_paragraph(element: Any) -> None:
    # Word requires a cell to finish with a paragraph, including after a table.
    if element.tag == qn("w:tc") and (not len(element) or element[-1].tag != qn("w:p")):
        element.append(OxmlElement("w:p"))


def _font_size(paragraph: Any, run: Run | None) -> int:
    if run is not None and run.font.size is not None:
        return int(run.font.size)
    styles = [run.style if run is not None else None, paragraph.style]
    for style in styles:
        while style is not None:
            if style.font.size is not None:
                return int(style.font.size)
            style = style.base_style
    return int(Pt(16))


def _left_indent(paragraph: Any, char_width: int, inherited: Mapping[str, str]) -> int:
    attributes = dict(inherited)
    direct = paragraph._p.pPr.ind if paragraph._p.pPr is not None else None
    if direct is not None:
        attributes.update(direct.attrib)
    axis = "start" if any(qn(f"w:{name}") in attributes for name in ("start", "startChars")) else "left"
    characters = int(attributes.get(qn(f"w:{axis}Chars"), "0"))
    if characters:
        return round(characters * char_width / 100)
    return int(Twips(int(attributes.get(qn(f"w:{axis}"), "0"))))


def _attachment_tab_stop(paragraph: Any, target: Any, position: int) -> None:
    stops = paragraph.paragraph_format.tab_stops
    stops.clear_all()
    properties = [target._p.pPr]
    style = target.style
    while style is not None:
        properties.append(style.element.pPr)
        style = style.base_style
    for properties_element in properties:
        if properties_element is None:
            continue
        for stop in properties_element.findall(f"{qn('w:tabs')}/{qn('w:tab')}"):
            stops.add_tab_stop(Twips(int(stop.get(qn("w:pos")))), WD_TAB_ALIGNMENT.CLEAR)
    stops.add_tab_stop(position)
