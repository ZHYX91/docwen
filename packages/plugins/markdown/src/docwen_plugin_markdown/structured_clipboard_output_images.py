"""Output-specific image rendering for structured clipboard documents."""

from __future__ import annotations

import csv
import hashlib
import os
import posixpath
import tempfile
from io import BytesIO
from pathlib import Path
from typing import Any
from zipfile import ZIP_DEFLATED, ZipFile, ZipInfo

from lxml import etree

from docwen_core.models.clipboard_document import (
    ClipboardDocument,
    ClipboardHardBreak,
    ClipboardImageRef,
    ClipboardParagraph,
    ClipboardText,
    iter_clipboard_inlines,
)
from docwen_core.models.result import ConversionDiagnostic
from docwen_plugin_markdown.structured_clipboard_images import (
    BoundClipboardImage,
    ClipboardImageOccurrenceInfo,
    MarkdownImagePlan,
)

_REL_NS = "http://schemas.openxmlformats.org/package/2006/relationships"


def image_resource_diagnostics(document: ClipboardDocument) -> list[ConversionDiagnostic]:
    missing_count = sum(
        isinstance(inline, ClipboardImageRef) and inline.resource_id is None
        for inline in iter_clipboard_inlines(document)
    )
    renderable_ids = {resource.resource_id for resource in document.resources if resource.media_type == "image/png"}
    not_rendered_count = sum(
        isinstance(inline, ClipboardImageRef)
        and inline.resource_id is not None
        and inline.resource_id not in renderable_ids
        for inline in iter_clipboard_inlines(document)
    )
    diagnostics = []
    if missing_count:
        diagnostics.append(
            ConversionDiagnostic(
                level="warning",
                code="CLIPBOARD-IMAGE-RESOURCE-UNAVAILABLE",
                message=(
                    f"{missing_count} clipboard image occurrence(s) had no verified resource bytes "
                    "and were kept as text placeholders."
                ),
            )
        )
    if not_rendered_count:
        diagnostics.append(
            ConversionDiagnostic(
                level="warning",
                code="CLIPBOARD-IMAGE-RESOURCE-NOT-RENDERED",
                message=(
                    f"{not_rendered_count} clipboard image occurrence(s) had verified opaque resources "
                    "without a supported PNG representation and were kept as text placeholders."
                ),
            )
        )
    return diagnostics


def image_placeholder(image: ClipboardImageRef) -> str:
    state = "not rendered" if image.resource_id is not None else "unavailable"
    return f"[Image {state}: {image.alt}]" if image.alt else f"[Image {state}]"


def write_docx_paragraph(
    container,
    paragraph: ClipboardParagraph,
    resources: dict[str, BoundClipboardImage],
):
    """Render text/break/image occurrences in original paragraph order."""

    from docx.shared import Emu

    output = container.add_paragraph()
    run = output.add_run()
    for inline in paragraph.inlines:
        if isinstance(inline, ClipboardText):
            # add_run converts literal tabs/newlines into native w:tab/w:br.
            # add_text would leave them inside w:t and lose the visible breaks.
            run = output.add_run(inline.value)
        elif isinstance(inline, ClipboardHardBreak):
            run.add_break()
        elif inline.resource_id is None:
            run.add_text(image_placeholder(inline))
        else:
            bound = resources.get(inline.resource_id)
            if bound is None:
                run.add_text(image_placeholder(inline))
                continue
            kwargs = {}
            if inline.extent_cx_emu is not None and inline.extent_cy_emu is not None:
                kwargs = {"width": Emu(inline.extent_cx_emu), "height": Emu(inline.extent_cy_emu)}
            run.add_picture(BytesIO(bound.path.read_bytes()), **kwargs)
    return output


def markdown_paragraph_blocks(
    paragraph: ClipboardParagraph,
    plan: MarkdownImagePlan,
    resources: dict[str, BoundClipboardImage],
    fence,
) -> list[str]:
    """Keep authored text literal while placing image references at occurrences."""

    output: list[str] = []
    text: list[str] = []

    def flush() -> None:
        value = "".join(text)
        text.clear()
        if value:
            output.extend([fence(value), ""])

    for inline in paragraph.inlines:
        if isinstance(inline, ClipboardText):
            text.append(inline.value)
        elif isinstance(inline, ClipboardHardBreak):
            text.append("\n")
        else:
            flush()
            output.extend([plan.reference(inline, resources), ""])
    flush()
    if not output:
        output.extend([fence(""), ""])
    return output


def add_xlsx_images(
    projection,
    occurrences: tuple[ClipboardImageOccurrenceInfo, ...],
    resources: dict[str, BoundClipboardImage],
    table_sheets: dict[str, Any],
) -> list[tuple[str, ...]]:
    """Add one drawing per bound occurrence and return visible semantics rows."""

    from openpyxl.drawing.image import Image as XlsxImage
    from openpyxl.drawing.spreadsheet_drawing import AnchorMarker, OneCellAnchor
    from openpyxl.drawing.xdr import XDRPositiveSize2D
    from openpyxl.utils import coordinate_to_tuple, get_column_letter

    semantics = projection.workbook.create_sheet("Image Semantics")
    headers = (
        "Ordinal",
        "State",
        "Resource ID",
        "Logical Path",
        "Table",
        "Cell Anchor",
        "Drawing Sheet",
        "Drawing Anchor",
        "Extent Cx Emu",
        "Extent Cy Emu",
    )
    rows: list[tuple[str, ...]] = []
    top_level_row = 2
    for occurrence in occurrences:
        image = occurrence.image
        if image.resource_id is None:
            rows.append(
                (
                    str(occurrence.ordinal),
                    "missing",
                    "",
                    "",
                    occurrence.table_id,
                    occurrence.cell_anchor,
                    "",
                    "",
                    str(image.extent_cx_emu or ""),
                    str(image.extent_cy_emu or ""),
                )
            )
            continue
        bound = resources.get(image.resource_id)
        if bound is None:
            rows.append(
                (
                    str(occurrence.ordinal),
                    "not_rendered",
                    image.resource_id,
                    "",
                    occurrence.table_id,
                    occurrence.cell_anchor,
                    "",
                    "",
                    str(image.extent_cx_emu or ""),
                    str(image.extent_cy_emu or ""),
                )
            )
            continue

        if occurrence.table_id:
            sheet = table_sheets[occurrence.table_id]
            raw = occurrence.cell_anchor.removeprefix("R")
            row_part, column_part = raw.split("C", 1)
            drawing_anchor = f"{get_column_letter(int(column_part))}{int(row_part)}"
        else:
            sheet = projection.order_sheet
            drawing_anchor = f"I{top_level_row}"
            top_level_row += 2

        drawing = XlsxImage(str(bound.path))
        if image.extent_cx_emu is not None and image.extent_cy_emu is not None:
            row, column = coordinate_to_tuple(drawing_anchor)
            drawing.anchor = OneCellAnchor(
                _from=AnchorMarker(col=column - 1, row=row - 1),
                ext=XDRPositiveSize2D(cx=image.extent_cx_emu, cy=image.extent_cy_emu),
            )
        elif bound.resource.pixel_width is not None and bound.resource.pixel_height is not None:
            drawing.width = bound.resource.pixel_width
            drawing.height = bound.resource.pixel_height
        if image.extent_cx_emu is not None and image.extent_cy_emu is not None:
            sheet.add_image(drawing)
        else:
            sheet.add_image(drawing, drawing_anchor)
        rows.append(
            (
                str(occurrence.ordinal),
                "bound",
                image.resource_id,
                bound.resource.logical_path,
                occurrence.table_id,
                occurrence.cell_anchor,
                sheet.title,
                drawing_anchor,
                str(image.extent_cx_emu or ""),
                str(image.extent_cy_emu or ""),
            )
        )

    for column, value in enumerate(headers, start=1):
        semantics.cell(1, column, value)
    for row_index, values in enumerate(rows, start=2):
        for column, value in enumerate(values, start=1):
            semantics.cell(row_index, column, value)
    return rows


def write_csv_image_semantics(
    path: Path,
    occurrences: tuple[ClipboardImageOccurrenceInfo, ...],
    resources: dict[str, BoundClipboardImage],
) -> None:
    """Write honest image occurrence/resource projection for CSV output."""

    with path.open("w", encoding="utf-8-sig", newline="") as stream:
        writer = csv.writer(stream)
        writer.writerow(
            (
                "Ordinal",
                "State",
                "Resource ID",
                "Logical Path",
                "Suggested Resource",
                "Table",
                "Cell Anchor",
                "Extent Cx Emu",
                "Extent Cy Emu",
            )
        )
        for occurrence in occurrences:
            image = occurrence.image
            bound = resources.get(image.resource_id or "")
            writer.writerow(
                (
                    occurrence.ordinal,
                    "bound" if bound is not None else ("not_rendered" if image.resource_id is not None else "missing"),
                    image.resource_id or "",
                    bound.resource.logical_path if bound is not None else "",
                    bound.suggested_name if bound is not None else "",
                    occurrence.table_id,
                    occurrence.cell_anchor,
                    image.extent_cx_emu or "",
                    image.extent_cy_emu or "",
                )
            )


def deduplicate_xlsx_png_media(path: Path) -> None:
    """Point duplicate drawing relationships at one PNG part and remove copies."""

    with ZipFile(path, "r") as archive:
        infos = archive.infolist()
        data = {info.filename: archive.read(info.filename) for info in infos if not info.is_dir()}

    media_names = sorted(name for name in data if name.startswith("xl/media/") and name.lower().endswith(".png"))
    canonical_by_hash: dict[str, str] = {}
    replacement: dict[str, str] = {}
    for name in media_names:
        digest = hashlib.sha256(data[name]).hexdigest()
        canonical = canonical_by_hash.setdefault(digest, name)
        if canonical != name:
            replacement[name] = canonical
    if not replacement:
        return

    parser = etree.XMLParser(resolve_entities=False, no_network=True)
    for name, payload in list(data.items()):
        if not name.startswith("xl/drawings/_rels/") or not name.endswith(".rels"):
            continue
        try:
            root = etree.fromstring(payload, parser)
        except etree.XMLSyntaxError as exc:
            raise ValueError("XLSX drawing relationships are invalid") from exc
        changed = False
        drawing_dir = posixpath.dirname(posixpath.dirname(name))
        for relation in root:
            if relation.tag != f"{{{_REL_NS}}}Relationship":
                continue
            target = relation.get("Target")
            if not target or relation.get("TargetMode") == "External":
                continue
            absolute = (
                posixpath.normpath(target).lstrip("/")
                if target.startswith("/")
                else posixpath.normpath(posixpath.join(drawing_dir, target))
            )
            canonical = replacement.get(absolute)
            if canonical is None:
                continue
            relation.set("Target", posixpath.relpath(canonical, drawing_dir))
            changed = True
        if changed:
            data[name] = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)

    for duplicate in replacement:
        data.pop(duplicate, None)

    descriptor, temp_name = tempfile.mkstemp(prefix=f".{path.name}.", suffix=".tmp", dir=path.parent)
    os.close(descriptor)
    Path(temp_name).unlink(missing_ok=True)
    temporary = Path(temp_name)
    try:
        original_by_name = {info.filename: info for info in infos}
        with ZipFile(temporary, "w") as output:
            for name, payload in data.items():
                info = original_by_name.get(name)
                if info is None:
                    info = ZipInfo(name, date_time=(1980, 1, 1, 0, 0, 0))
                    info.compress_type = ZIP_DEFLATED
                output.writestr(info, payload)
        temporary.replace(path)
    finally:
        temporary.unlink(missing_ok=True)


__all__ = [
    "add_xlsx_images",
    "deduplicate_xlsx_png_media",
    "image_placeholder",
    "image_resource_diagnostics",
    "markdown_paragraph_blocks",
    "write_csv_image_semantics",
    "write_docx_paragraph",
]
