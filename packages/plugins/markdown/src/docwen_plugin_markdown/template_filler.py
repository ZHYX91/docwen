"""Fill original Word template anchors once, preserving surrounding OOXML."""

from __future__ import annotations

import contextlib
import logging
import re
from collections.abc import Iterator, Mapping
from copy import deepcopy
from typing import Any, Literal, cast

from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.oxml.text.run import CT_R
from docx.text.run import Run

from docwen_core.docx_comments import append_docx_comment
from docwen_plugin_markdown.gongwen_fields import (
    gongwen_review_notes,
    prepare_gongwen_fields,
    remove_empty_gongwen_fields,
    render_attachment_placeholders,
    validate_attachment_placeholders,
    validate_review_placeholders,
)
from docwen_plugin_markdown.yaml_processor import is_empty_field_value

logger = logging.getLogger(__name__)
_PLACEHOLDER = re.compile(r"\{\{\s*([^{}\r\n]+?)\s*\}\}")


def fill_template(
    doc,
    yaml_dict: dict[str, Any],
    rendered_paragraphs: list,
    placeholder_para,
    placeholder_map: dict[str, list[Any]] | None = None,
    *,
    gongwen_fields_enabled: bool = True,
    attachment_alignment: Literal["spaces", "tab"] = "spaces",
    list_separator: str = ", ",
) -> None:
    """Fill the pre-render template snapshot using request-local display values.

    The caller scans placeholders before rendering Markdown. Inserted values
    and rendered body text are never interpreted as new template expressions.
    The Gongwen switch controls value formatting, attachment layout and empty
    field cleanup together; it is independent of the output/interface locale.
    """
    values = prepare_gongwen_fields(yaml_dict) if gongwen_fields_enabled else yaml_dict
    review_notes = gongwen_review_notes(yaml_dict) if gongwen_fields_enabled else {}
    if placeholder_map and gongwen_fields_enabled:
        validate_attachment_placeholders(doc, placeholder_map)
        validate_review_placeholders(doc, placeholder_map, review_notes)
        remove_empty_gongwen_fields(doc, values, placeholder_map)
        render_attachment_placeholders(values.get("附件说明", ()), placeholder_map, alignment=attachment_alignment)

    if rendered_paragraphs and placeholder_para is not None:
        inject_body_paragraphs(doc, placeholder_para, rendered_paragraphs)
    elif rendered_paragraphs:
        _append_rendered_batch(doc, rendered_paragraphs)

    if placeholder_map:
        fill_yaml_placeholders(
            doc,
            values,
            placeholder_map,
            skip_keys={"附件说明"} if gongwen_fields_enabled else None,
            list_separator=list_separator,
            review_notes=review_notes,
        )


def inject_body_paragraphs(doc, placeholder_para, rendered_paragraphs: list) -> None:
    """Batch-relocate rendered content to the original main-body anchor."""
    body = doc.element.body
    placeholder_elem = placeholder_para._element
    try:
        index = list(body).index(placeholder_elem)
    except ValueError:
        logger.warning("Placeholder element not found in document body")
        return
    elements = []
    for paragraph in rendered_paragraphs:
        element = paragraph._element
        with contextlib.suppress(ValueError):
            body.remove(element)
        elements.append(element)
    for offset, element in enumerate(elements):
        body.insert(index + offset, element)
    body.remove(placeholder_elem)


def fill_yaml_placeholders(
    doc,
    yaml_dict: Mapping[str, Any],
    placeholder_map: dict[str, list[Any]],
    skip_keys: set[str] | None = None,
    list_separator: str = ", ",
    *,
    review_notes: Mapping[str, str] | None = None,
) -> None:
    """Replace all authored keys in each paragraph in one nonrecursive pass.

    Missing/empty values replace only their original placeholder. Every
    replacement inherits the opening placeholder run's properties, without
    flattening surrounding runs, hyperlinks or paragraph properties.
    """
    skipped = skip_keys or set()
    paragraphs: dict[Any, tuple[Any, dict[str, str]]] = {}
    for key, locations in placeholder_map.items():
        if key in skipped:
            continue
        replacement = _format_placeholder_value(yaml_dict.get(key), list_separator=list_separator)
        for paragraph in locations:
            _, replacements = paragraphs.setdefault(paragraph._element, (paragraph, {}))
            replacements[key] = replacement
    for paragraph, replacements in paragraphs.values():
        _replace_placeholder_runs(paragraph, replacements, doc=doc, review_notes=review_notes)


def _replace_placeholder_runs(
    paragraph, replacements: Mapping[str, str], *, doc=None, review_notes: Mapping[str, str] | None = None
) -> None:
    characters: list[str] = []
    locations: list[tuple[Any, int] | None] = []
    for element in paragraph._element.iter():
        # Textbox paragraphs are separate locations, never part of the outer
        # paragraph that owns their drawing.
        owner = element.getparent()
        while owner is not None and owner.tag != qn("w:p"):
            owner = owner.getparent()
        if owner is not paragraph._element:
            continue
        if element.tag == qn("w:t"):
            for offset, character in enumerate(element.text or ""):
                characters.append(character)
                locations.append((element, offset))
        elif element.tag in {qn("w:br"), qn("w:cr"), qn("w:tab")}:
            characters.append("\0")
            locations.append(None)

    # Match all keys before editing any text. Reverse edits keep original
    # offsets valid even when values contain another key's {{placeholder}}.
    matches = list(_PLACEHOLDER.finditer("".join(characters)))
    for match in reversed(matches):
        key = match.group(1).strip()
        if key not in replacements or any(location is None for location in locations[match.start() : match.end()]):
            continue
        first = locations[match.start()]
        last = locations[match.end() - 1]
        if first is None or last is None:
            continue
        start_element, start_offset = first
        end_element, end_offset = last
        end_offset += 1
        replacement = replacements[key]
        if start_element is end_element:
            current = start_element.text or ""
            _set_word_text(start_element, current[:start_offset] + replacement + current[end_offset:])
        else:
            start_text = start_element.text or ""
            end_text = end_element.text or ""
            _set_word_text(start_element, start_text[:start_offset] + replacement)
            seen = {start_element, end_element}
            for location in locations[match.start() + 1 : match.end() - 1]:
                if location is not None and location[0] not in seen:
                    _set_word_text(location[0], "")
                    seen.add(location[0])
            _set_word_text(end_element, end_text[end_offset:])
        note = review_notes.get(key) if review_notes is not None else None
        if note and replacement and doc is not None:
            # Isolate only the inserted text. Right-to-left substitution keeps
            # earlier original offsets valid even after splitting this run.
            _split_text_run(start_element, start_offset + len(replacement))
            inserted = _split_text_run(start_element, start_offset)
            append_docx_comment(doc, [Run(cast(CT_R, inserted), paragraph)], text=note)


def _split_text_run(text_element, offset: int):
    """Split at a text-node offset without flattening controls or hyperlinks."""
    run = text_element.getparent()
    if run is None or run.tag != qn("w:r"):
        raise ValueError("无法为公文字段定位批注，请检查模板中的占位符。")
    text = text_element.text or ""
    if not 0 <= offset <= len(text):
        raise ValueError("公文字段批注位置无效，未保存转换结果。")
    following = OxmlElement("w:r")
    following.attrib.update(run.attrib)
    if run.rPr is not None:
        following.append(deepcopy(run.rPr))
    suffix = deepcopy(text_element)
    _set_word_text(suffix, text[offset:])
    _set_word_text(text_element, text[:offset])
    following.append(suffix)
    for child in list(run)[run.index(text_element) + 1 :]:
        following.append(child)
    run.addnext(following)
    return following


def _set_word_text(element, text: str) -> None:
    element.text = text
    if text[:1].isspace() or text[-1:].isspace():
        element.set(qn("xml:space"), "preserve")
    else:
        element.attrib.pop(qn("xml:space"), None)


def _append_rendered_batch(doc, paragraphs: list) -> None:
    """Keep one rendered batch ordered before terminal section properties."""
    body = doc.element.body
    elements = []
    for paragraph in paragraphs:
        element = paragraph._element
        with contextlib.suppress(ValueError):
            body.remove(element)
        elements.append(element)
    section = body.sectPr
    index = body.index(section) if section is not None else len(body)
    for offset, element in enumerate(elements):
        body.insert(index + offset, element)


def _format_placeholder_value(value: Any, *, list_separator: str) -> str:
    if is_empty_field_value(value):
        return ""
    if isinstance(value, (list, tuple)):
        return list_separator.join(_iter_placeholder_list_items(value))
    return str(value)


def _iter_placeholder_list_items(value: list[Any] | tuple[Any, ...]) -> Iterator[str]:
    for item in value:
        if isinstance(item, (list, tuple)):
            yield from _iter_placeholder_list_items(item)
        elif not is_empty_field_value(item) and not (isinstance(item, str) and item in {"null", "None"}):
            yield str(item)
