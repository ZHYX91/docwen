"""PPTX grouped-shape and YAML frontmatter regressions."""

from __future__ import annotations

from pathlib import Path

import pytest
import yaml

from ._presentation_to_md_support import _run_request, pipeline as pipeline

pytestmark = pytest.mark.integration


def test_grouped_text_and_image_are_exported_and_title_yaml_is_safe(pipeline, tmp_path: Path) -> None:
    from PIL import Image
    from pptx import Presentation
    from pptx.util import Inches

    image_path = tmp_path / "grouped.png"
    Image.new("RGB", (5, 5), (30, 90, 150)).save(image_path)
    expected_image = image_path.read_bytes()

    presentation = Presentation()
    presentation.core_properties.title = "计划: 修订 #1"
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    group = slide.shapes.add_group_shape()
    textbox = group.shapes.add_textbox(Inches(1), Inches(1), Inches(3), Inches(1))
    textbox.text_frame.text = "Grouped text payload"
    group.shapes.add_picture(
        str(image_path),
        Inches(1),
        Inches(2),
        width=Inches(1),
        height=Inches(1),
    )

    source = tmp_path / "grouped.pptx"
    presentation.save(source)
    output = tmp_path / "output"
    output.mkdir()

    _plugin, task_mgr, _workspace = pipeline
    result = _run_request(
        task_mgr,
        source,
        "pptx",
        output,
        image_link_style="markdown_embed",
    )

    assert result.success, result.error
    primary = next(item for item in result.artifacts if item.is_primary)
    markdown = Path(primary.staging_path).read_text(encoding="utf-8")

    assert "Grouped text payload" in markdown
    image_artifacts = [item for item in result.artifacts if item.kind == "image"]
    assert len(image_artifacts) == 1
    assert Path(image_artifacts[0].staging_path).read_bytes() == expected_image

    frontmatter = yaml.safe_load(markdown.split("---", 2)[1])
    assert frontmatter["title"] == "计划: 修订 #1"
    assert frontmatter["aliases"] == ["计划: 修订 #1"]
