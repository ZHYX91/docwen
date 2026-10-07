"""PPTX grouped-content and YAML fidelity regressions."""

from __future__ import annotations

from pathlib import Path

import pytest
import yaml

from ._presentation_to_md_support import _run_request, pipeline as pipeline

pytestmark = pytest.mark.integration


def test_pptx_group_shape_preserves_nested_text_and_image(pipeline, tmp_path: Path) -> None:
    from PIL import Image
    from pptx import Presentation
    from pptx.util import Inches

    image_path = tmp_path / "grouped.png"
    Image.new("RGB", (4, 4), (30, 90, 180)).save(image_path)
    expected_image = image_path.read_bytes()

    presentation = Presentation()
    presentation.core_properties.title = "Grouped content"
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    group = slide.shapes.add_group_shape()
    textbox = group.shapes.add_textbox(Inches(1), Inches(1), Inches(3), Inches(0.5))
    textbox.text_frame.text = "Nested text survives"
    group.shapes.add_picture(str(image_path), Inches(1), Inches(2), width=Inches(1), height=Inches(1))

    source = tmp_path / "grouped.pptx"
    presentation.save(str(source))
    output = tmp_path / "out-grouped"
    output.mkdir()
    _plugin, task_mgr, _ws_mgr = pipeline

    result = _run_request(task_mgr, source, "pptx", output, image_link_style="markdown_embed")

    assert result.success, result.error
    markdown = Path(result.artifacts[0].staging_path).read_text(encoding="utf-8")
    assert "Nested text survives" in markdown
    images = [artifact for artifact in result.artifacts if artifact.kind == "image"]
    assert len(images) == 1
    assert Path(images[0].staging_path).read_bytes() == expected_image
    assert images[0].suggested_name in markdown


def test_pptx_title_is_yaml_serialized_without_changing_value(pipeline, tmp_path: Path) -> None:
    from pptx import Presentation

    title = "计划: 修订 #1"
    presentation = Presentation()
    presentation.core_properties.title = title
    presentation.slides.add_slide(presentation.slide_layouts[6])

    source = tmp_path / "yaml-title.pptx"
    presentation.save(str(source))
    output = tmp_path / "out-yaml-title"
    output.mkdir()
    _plugin, task_mgr, _ws_mgr = pipeline

    result = _run_request(task_mgr, source, "pptx", output)

    assert result.success, result.error
    markdown = Path(result.artifacts[0].staging_path).read_text(encoding="utf-8")
    frontmatter = yaml.safe_load(markdown.split("---", 2)[1])
    assert frontmatter["title"] == title
    assert frontmatter["aliases"] == [title]
