"""Exercise python-docx's real template lookups after PyInstaller freezing."""

from __future__ import annotations

import json
import os
import runpy
import sys
import textwrap
from pathlib import Path, PurePosixPath

import pytest
from tests.support.subprocess_runner import run_subprocess

ROOT = Path(__file__).resolve().parents[2]
HOOKS = ROOT / "scripts" / "build" / "pyinstaller_hooks"
TEMPLATES = ("comments", "footer", "header", "settings", "styles")


@pytest.mark.unit
def test_docx_hook_collects_all_default_xml_templates() -> None:
    hook = runpy.run_path(str(HOOKS / "hook-docx.py"))
    collected = {
        (PurePosixPath(destination.replace("\\", "/")) / Path(source).name).as_posix()
        for source, destination in hook["datas"]
    }

    assert {f"docx/templates/default-{name}.xml" for name in TEMPLATES} <= collected
    assert "docx/templates/default.docx" in collected


@pytest.mark.integration
@pytest.mark.release_gate
@pytest.mark.slow
def test_frozen_docx_loads_default_parts_and_roundtrips_comments(tmp_path: Path) -> None:
    probe = tmp_path / "docx_probe.py"
    probe.write_text(
        textwrap.dedent(
            """
            import io
            import json
            import sys
            from pathlib import Path

            from docx import Document
            from docx.parts.comments import CommentsPart
            from docx.parts.hdrftr import FooterPart, HeaderPart
            from docx.parts.settings import SettingsPart
            from docx.parts.styles import StylesPart

            document = Document()
            package = document.part.package
            parts = (
                CommentsPart.default(package),
                FooterPart.new(package),
                HeaderPart.new(package),
                SettingsPart.default(package),
                StylesPart.default(package),
            )
            roots = [part.element.tag.rsplit('}', 1)[-1] for part in parts]
            document.add_comment(document.add_paragraph('anchor').runs[0], text='review')
            document.sections[0].header.paragraphs[0].text = 'header'
            document.sections[0].footer.paragraphs[0].text = 'footer'
            stream = io.BytesIO()
            document.save(stream)
            stream.seek(0)
            reopened = Document(stream)
            comments = list(reopened.comments)
            assert len(comments) == 1
            assert comments[0].text == 'review'
            assert reopened.sections[0].header.paragraphs[0].text == 'header'
            assert reopened.sections[0].footer.paragraphs[0].text == 'footer'
            print(json.dumps({
                'frozen': bool(getattr(sys, 'frozen', False)),
                'partsDirectoryExists': (Path(sys._MEIPASS) / 'docx' / 'parts').is_dir(),
                'roots': roots,
                'comment': comments[0].text,
            }))
            """
        ),
        encoding="utf-8",
    )
    build = run_subprocess(
        [
            sys.executable,
            "-I",
            "-m",
            "PyInstaller",
            str(probe),
            "--name=docx-probe",
            "--onedir",
            "--clean",
            "--noconfirm",
            f"--additional-hooks-dir={HOOKS}",
            f"--distpath={tmp_path / 'dist'}",
            f"--workpath={tmp_path / 'build'}",
            f"--specpath={tmp_path / 'spec'}",
        ],
        cwd=tmp_path,
        env={**os.environ, "PYINSTALLER_CONFIG_DIR": str(tmp_path / "pyinstaller-cache")},
        timeout=180,
    )
    assert build.returncode == 0, build.stdout + build.stderr
    payload = tmp_path / "dist" / "docx-probe"
    for name in TEMPLATES:
        assert (payload / "_internal" / "docx" / "templates" / f"default-{name}.xml").is_file()
    assert (payload / "_internal" / "docx" / "parts" / "comments.py").is_file()

    executable = payload / ("docx-probe.exe" if sys.platform == "win32" else "docx-probe")
    result = run_subprocess([str(executable)], cwd=tmp_path, timeout=30)
    assert result.returncode == 0, result.stdout + result.stderr
    assert json.loads(result.stdout) == {
        "frozen": True,
        "partsDirectoryExists": True,
        "roots": ["comments", "ftr", "hdr", "settings", "styles"],
        "comment": "review",
    }
