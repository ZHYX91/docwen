"""Repository tests for maintenance template-generation tooling."""

from __future__ import annotations

import subprocess
import sys
from pathlib import Path

import pytest
from docx import Document

pytestmark = pytest.mark.unit

PROJECT_ROOT = Path(__file__).resolve().parents[2]
SCRIPT = PROJECT_ROOT / "scripts" / "maintenance" / "generate_templates.py"


def test_generate_templates_help_uses_current_locale_directory() -> None:
    result = subprocess.run(
        [sys.executable, str(SCRIPT), "--help"],
        cwd=PROJECT_ROOT,
        capture_output=True,
        text=True,
        timeout=30,
    )

    assert result.returncode == 0, result.stderr
    assert "--locale" in result.stdout
    assert "en_US" in result.stdout
    assert "src\\docwen\\i18n\\locales" not in result.stderr


def test_generate_templates_creates_localized_docx_placeholders(tmp_path: Path) -> None:
    result = subprocess.run(
        [
            sys.executable,
            str(SCRIPT),
            "--locale",
            "en_US",
            "--output-dir",
            str(tmp_path),
        ],
        cwd=PROJECT_ROOT,
        capture_output=True,
        text=True,
        timeout=30,
    )

    assert result.returncode == 0, result.stderr

    output = tmp_path / "English General Template.docx"
    assert output.exists()

    doc = Document(str(output))
    paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
    assert paragraphs[:2] == ["{{title}}", "{{body}}"]
