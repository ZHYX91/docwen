"""The packaged field gate rejects semantic fallback, not just missing modules."""

from __future__ import annotations

import json
import subprocess
import tomllib
from pathlib import Path

import pytest
from scripts.release.gongwen_fields_fixture import (
    verify_gongwen_fields_output,
    write_gongwen_fields_cases,
)

pytestmark = pytest.mark.contract


def _write_expected_document(path: Path, *, enabled: bool, empty: bool) -> None:
    """An explicit output golden independent of the conversion implementation."""
    from docx import Document
    from docx.shared import Pt

    def layout(paragraph, index):
        paragraph.paragraph_format.left_indent = Pt(84)
        paragraph.paragraph_format.right_indent = Pt(8)
        paragraph.paragraph_format.first_line_indent = Pt(-72 if index == 0 else -24)
        paragraph.paragraph_format.keep_with_next = True
        paragraph.paragraph_format.tab_stops.add_tab_stop(Pt(84))
        paragraph.runs[0].font.size = Pt(16)
        paragraph.runs[0].bold = True

    document = Document()
    document.add_paragraph("DOCWEN ATTACHMENTS START")
    attachments = (
        ([] if enabled else [""])
        if empty
        else ["附件：1.\t任务分工表", "2.\t工作安排表"]
        if enabled
        else ["任务分工表 / 工作安排表"]
    )
    for index, text in enumerate(attachments):
        paragraph = document.add_paragraph(text)
        if enabled and not empty:
            layout(paragraph, index)
    for text in (
        "DOCWEN ATTACHMENTS END",
        "普通列表：普通甲 / 普通乙",
        "抄送：单位甲、单位乙" if enabled else "抄送：单位甲 / 单位乙",
        "（此件公开发布）" if enabled else "（（此件公开发布））",
        "成文：2026年10月20日" if enabled else "成文：2026-10-20",
        "DOCWEN EMPTY FIELD START",
    ):
        document.add_paragraph(text)
    if not enabled:
        document.add_paragraph("：")
    document.add_paragraph("DOCWEN EMPTY FIELD END")
    cell = document.add_table(rows=1, cols=1).cell(0, 0)
    if attachments:
        cell.paragraphs[0].text = attachments[0]
        for text in attachments[1:]:
            cell.add_paragraph(text)
        if enabled and not empty:
            for index, paragraph in enumerate(cell.paragraphs):
                layout(paragraph, index)
    imprint = document.add_table(rows=1 if enabled and empty else 2, cols=2)
    imprint.cell(0, 0).text = "DOCWEN IMPRINT"
    imprint.cell(0, 1).text = "KEEP THIS ROW"
    if not empty:
        imprint.cell(1, 0).text = "测试机关"
        imprint.cell(1, 1).text = "2026年10月21日印发" if enabled else "2026-10-21"
    document.add_paragraph("DOCWEN PACKAGED GONGWEN FIELDS CONTENT")
    path.parent.mkdir(parents=True, exist_ok=True)
    document.save(path)


def test_gongwen_fields_fixtures_have_independent_config_and_real_templates(tmp_path: Path) -> None:
    from docx import Document

    cases = write_gongwen_fields_cases(tmp_path / "fixtures")
    assert len(cases) == 4
    assert {(case.enabled, case.empty) for case in cases} == {
        (True, True),
        (True, False),
        (False, True),
        (False, False),
    }
    for case in cases:
        config = tomllib.loads((case.work_dir / "config_home" / "template_fill.toml").read_text(encoding="utf-8"))
        assert config == {"gongwen_fields_enabled": case.enabled, "list_separator": " / "}
        state = json.loads((case.work_dir / "data_home" / "template-state.json").read_text(encoding="utf-8"))
        identity = state["user_identities"][case.template_path.name]
        assert state["defaults"]["docx"] == identity["id"]
        assert state["enabled"][identity["id"]] is True
        template = Document(case.template_path)
        assert [paragraph.text for paragraph in template.paragraphs].count("{{附件说明}}") == 1
        assert template.tables[0].cell(0, 0).text == "{{附件说明}}"
        assert template.tables[1].cell(1, 0).text == "{{印发机关}}"
        assert template.tables[1].cell(1, 1).text == "{{印发日期}}"
        output = case.work_dir / "outputs" / "golden.docx"
        _write_expected_document(output, enabled=case.enabled, empty=case.empty)
        verify_gongwen_fields_output(output, case=case)


@pytest.mark.parametrize(
    ("mutation", "case_name", "expected_error"),
    [
        ("body_joined", "on-values", "attachments_mismatch"),
        ("table_joined", "on-values", "attachments_mismatch"),
        ("wrong_number", "on-values", "attachments_mismatch"),
        ("missing_indent", "on-values", "attachment_layout_mismatch"),
        ("wrong_tab", "on-values", "attachment_layout_mismatch"),
        ("lost_font", "on-values", "attachment_layout_mismatch"),
        ("lost_paragraph_format", "on-values", "attachment_layout_mismatch"),
        ("ordinary_separator", "on-values", "scalar_mismatch"),
        ("copy_separator", "on-values", "scalar_mismatch"),
        ("double_parentheses", "on-values", "scalar_mismatch"),
        ("raw_date", "on-values", "scalar_mismatch"),
        ("raw_print_date", "on-values", "imprint_mismatch"),
        ("empty_paragraph", "on-values", "empty_paragraph_mismatch"),
        ("empty_row", "on-empty", "empty_row_mismatch"),
        ("cell_without_paragraph", "on-empty", "attachment_cell_invalid"),
        ("ignored_off", "off-values", "attachments_mismatch"),
        ("off_row_removed", "off-empty", "empty_row_mismatch"),
    ],
)
def test_gongwen_fields_oracle_rejects_semantic_regressions(
    mutation: str,
    case_name: str,
    expected_error: str,
    tmp_path: Path,
) -> None:
    from docx import Document

    case = next(case for case in write_gongwen_fields_cases(tmp_path / "fixtures") if case.name == case_name)
    output = case.work_dir / "outputs" / "incorrect.docx"
    _write_expected_document(output, enabled=case.enabled or mutation == "ignored_off", empty=case.empty)
    document = Document(output)
    paragraphs = {paragraph.text: paragraph for paragraph in document.paragraphs}
    if mutation == "body_joined":
        paragraphs["附件：1.\t任务分工表"].text = "任务分工表、工作安排表"
        second = paragraphs["2.\t工作安排表"]._p
        second.getparent().remove(second)
    elif mutation == "table_joined":
        document.tables[0].cell(0, 0).text = "任务分工表、工作安排表"
    elif mutation == "wrong_number":
        document.tables[0].cell(0, 0).paragraphs[1].text = "1. 工作安排表"
    elif mutation == "missing_indent":
        paragraphs["附件：1.\t任务分工表"].paragraph_format.first_line_indent = None
    elif mutation == "wrong_tab":
        document.tables[0].cell(0, 0).paragraphs[1].paragraph_format.tab_stops.clear_all()
    elif mutation == "lost_font":
        paragraphs["附件：1.\t任务分工表"].runs[0].font.size = None
    elif mutation == "lost_paragraph_format":
        paragraphs["附件：1.\t任务分工表"].paragraph_format.keep_with_next = False
    elif mutation == "ordinary_separator":
        paragraphs["普通列表：普通甲 / 普通乙"].text = "普通列表：普通甲、普通乙"
    elif mutation == "copy_separator":
        paragraphs["抄送：单位甲、单位乙"].text = "抄送：单位甲 / 单位乙"
    elif mutation == "double_parentheses":
        paragraphs["（此件公开发布）"].text = "（（此件公开发布））"
    elif mutation == "raw_date":
        paragraphs["成文：2026年10月20日"].text = "成文：2026-10-20"
    elif mutation == "raw_print_date":
        document.tables[1].cell(1, 1).text = "2026-10-21"
    elif mutation == "empty_paragraph":
        paragraphs["DOCWEN EMPTY FIELD END"].insert_paragraph_before("：")
    elif mutation == "empty_row":
        document.tables[1].add_row()
    elif mutation == "cell_without_paragraph":
        cell = document.tables[0].cell(0, 0)
        cell._tc.remove(cell.paragraphs[0]._p)
    elif mutation == "off_row_removed":
        row = document.tables[1].rows[1]._tr
        row.getparent().remove(row)
    document.save(output)
    with pytest.raises(RuntimeError, match=f"gongwen_fields_{expected_error}"):
        verify_gongwen_fields_output(output, case=case)


def test_packaged_cli_gongwen_fields_gate_runs_four_real_convert_requests(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    from scripts.release import verify_packaged_cli

    observed: list[tuple[Path, bool, bool]] = []

    def fake_run(binary_path: Path, *args: str, cwd: Path) -> subprocess.CompletedProcess[str]:
        assert args[0] == "convert"
        assert args[args.index("--to") + 1] == "docx"
        assert args[-2:] == ("--json", "--quiet")
        template_state = json.loads((cwd / "data_home" / "template-state.json").read_text(encoding="utf-8"))
        assert args[args.index("--template") + 1] == template_state["defaults"]["docx"]
        config = tomllib.loads((cwd / "config_home" / "template_fill.toml").read_text(encoding="utf-8"))
        enabled = config["gongwen_fields_enabled"]
        empty = "附件说明: []" in Path(args[1]).read_text(encoding="utf-8")
        observed.append((cwd, enabled, empty))
        output = Path(args[args.index("--output-dir") + 1]) / "fields.docx"
        _write_expected_document(output, enabled=enabled, empty=empty)
        return subprocess.CompletedProcess(
            [str(binary_path), *args],
            0,
            stdout=json.dumps({"success": True, "command": "convert", "data": {"output": str(output)}}),
            stderr="",
        )

    monkeypatch.setattr(verify_packaged_cli, "_run", fake_run)
    outputs = verify_packaged_cli._run_gongwen_fields_smoke(tmp_path / "DocWenCLI.exe", work_dir=tmp_path / "run")
    assert len(outputs) == 4
    assert len({cwd for cwd, _enabled, _empty in observed}) == 4
    assert {(enabled, empty) for _cwd, enabled, empty in observed} == {
        (True, True),
        (True, False),
        (False, True),
        (False, False),
    }


def test_packaged_gui_gongwen_fields_gate_uses_action_area_and_profile_template(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    from scripts.release import verify_packaged_gui

    binary_dir = tmp_path / "dist"
    binary_dir.mkdir()
    binary = binary_dir / "DocWen.exe"
    binary.write_bytes(b"placeholder")
    monkeypatch.setattr(verify_packaged_gui, "_verify_resource_layout", lambda _path: None)
    monkeypatch.setenv("PYTHONPATH", str(tmp_path / "unrelated-source-tree"))
    monkeypatch.setenv("PYTHONHOME", str(tmp_path / "unrelated-python"))
    observed: list[dict[str, str]] = []

    def fake_run(
        binary_path: Path,
        *args: str,
        cwd: Path,
        env: dict[str, str],
        timeout: int,
    ) -> subprocess.CompletedProcess[str]:
        assert binary_path == binary
        assert not args
        assert timeout >= 90
        assert "PYTHONPATH" not in env and "PYTHONHOME" not in env
        assert env["DOCWEN_GUI_TEST_CONVERSION_SURFACE"] == "action"
        assert env["DOCWEN_GUI_TEST_CONVERSION_TARGET"] == "docx"
        assert env["DOCWEN_GUI_TEST_CONVERSION_ACTION"] == ""
        observed.append(dict(env))
        config_dir, data_dir = Path(env["DOCWEN_CONFIG_DIR"]), Path(env["DOCWEN_DATA_DIR"])
        config = tomllib.loads((config_dir / "template_fill.toml").read_text(encoding="utf-8"))
        template_state = json.loads((data_dir / "template-state.json").read_text(encoding="utf-8"))
        default_id = template_state["defaults"]["docx"]
        template = next(value for value in template_state["user_identities"].values() if value["id"] == default_id)
        assert (data_dir / "templates" / template["filename"]).is_file()
        source = Path(env["DOCWEN_GUI_TEST_CONVERSION_INPUT"])
        assert source.parent == cwd
        output = Path(env["DOCWEN_GUI_TEST_CONVERSION_OUTPUT_DIR"]) / "fields.docx"
        _write_expected_document(
            output,
            enabled=config["gongwen_fields_enabled"],
            empty="附件说明: []" in source.read_text(encoding="utf-8"),
        )
        Path(env["DOCWEN_GUI_TEST_CONVERSION_REPORT"]).write_text(
            json.dumps(
                {
                    "success": True,
                    "status": "completed",
                    "outputPath": str(output),
                    "outputExists": True,
                    "error": None,
                }
            ),
            encoding="utf-8",
        )
        logs = Path(env["DOCWEN_LOG_DIR"]) / "logs"
        logs.mkdir(parents=True, exist_ok=True)
        (logs / "fields.log").write_text("conversion completed", encoding="utf-8")
        return subprocess.CompletedProcess([str(binary_path)], 0, stdout="", stderr="")

    monkeypatch.setattr(verify_packaged_gui, "_run_with_env", fake_run)
    assert (
        verify_packaged_gui.main(
            ["--binary-dir", str(binary_dir), "--binary-name", binary.name, "--gongwen-fields-smoke"]
        )
        == 0
    )
    assert len(observed) == 4
    assert len({env["DOCWEN_CONFIG_DIR"] for env in observed}) == 4
    assert len({env["DOCWEN_DATA_DIR"] for env in observed}) == 4


def test_packaged_cli_child_environment_cannot_import_checkout(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    from scripts.release import verify_packaged_cli

    monkeypatch.setenv("PYTHONPATH", str(tmp_path / "source"))
    monkeypatch.setenv("PYTHONHOME", str(tmp_path / "python"))
    observed: dict[str, str] = {}

    def fake_run(args: list[str], **kwargs: object) -> subprocess.CompletedProcess[str]:
        env = kwargs["env"]
        assert isinstance(env, dict)
        observed.update(env)
        return subprocess.CompletedProcess(args, 0, stdout="", stderr="")

    monkeypatch.setattr(verify_packaged_cli.subprocess, "run", fake_run)
    verify_packaged_cli._run(tmp_path / "DocWenCLI.exe", "--help", cwd=tmp_path)
    assert "PYTHONPATH" not in observed and "PYTHONHOME" not in observed
    assert observed["DOCWEN_CONFIG_DIR"] == str(tmp_path / "config_home")


@pytest.mark.parametrize("other_gate", ["--settings-smoke", "--office-smoke", "--ipc-smoke"])
def test_gongwen_fields_gui_gate_cannot_share_a_live_profile(other_gate: str) -> None:
    from scripts.release import verify_packaged_gui

    with pytest.raises(SystemExit) as exc_info:
        verify_packaged_gui.main(["--binary-dir", "unused", "--gongwen-fields-smoke", other_gate])
    assert exc_info.value.code == 2


def test_release_workflow_runs_field_behavior_on_both_platforms_and_entrypoints() -> None:
    workflow = Path(".github/workflows/release.yml").read_text(encoding="utf-8")
    cli_commands = [line for line in workflow.splitlines() if "scripts/release/verify_packaged_cli.py" in line]
    gui_commands = [
        line
        for line in workflow.splitlines()
        if "scripts/release/verify_packaged_gui.py" in line and "--gongwen-fields-smoke" in line
    ]
    assert len(cli_commands) == 2 and all("--gongwen-fields-smoke" in command for command in cli_commands)
    assert len(gui_commands) == 2
