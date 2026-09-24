"""Small real OOXML fixture with distinct formula-cache states."""

from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile

import openpyxl
from lxml import etree


def write_formula_cache_fixture(path: Path) -> None:
    workbook = openpyxl.Workbook()
    first = workbook.active
    assert first is not None
    first.title = "Calc"
    for coordinate in ("A1", "B1", "C1", "D1"):
        first[coordinate] = '=IF(1=1,"",1)'
    first["E1"] = "ordinary"
    second = workbook.create_sheet("Other")
    for index in range(1, 26):
        second.cell(index, 1, "=1+1")
    workbook.save(path)
    workbook.close()
    with ZipFile(path) as source:
        parts = {name: source.read(name) for name in source.namelist()}
    namespace = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
    sheet = etree.fromstring(parts["xl/worksheets/sheet1.xml"])
    for cell in sheet.iter(f"{namespace}c"):
        value = cell.find(f"{namespace}v")
        if cell.get("r") == "A1":
            cell.set("t", "str")  # explicit empty cached string, valid
        elif cell.get("r") == "B1":
            cell.set("t", "str")  # same type, but missing v is not an empty result
            assert value is not None
            cell.remove(value)
        elif cell.get("r") == "D1":
            assert value is not None
            value.text = "0"  # zero is a usable cached value
    parts["xl/worksheets/sheet1.xml"] = etree.tostring(sheet)
    with ZipFile(path, "w", ZIP_DEFLATED) as target:
        for name, content in parts.items():
            target.writestr(name, content)
