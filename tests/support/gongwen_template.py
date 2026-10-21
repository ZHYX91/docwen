"""Build the formatting conditions of a real Gongwen template without a binary fixture."""

from __future__ import annotations

from typing import Literal

from docx import Document
from docx.document import Document as WordDocument
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt
from docx.styles.style import ParagraphStyle


def attachment_template(
    *,
    inheritance: Literal["style", "ancestor", "defaults", "direct"] = "style",
    location: Literal["body", "cell"] = "body",
) -> WordDocument:
    """Retain conflicting character/length indents and inherited 16 pt font.

    These four indentation declarations reproduce the reported real-template
    failure. Names, content and unrelated formatting are intentionally synthetic;
    the test never needs the original ZIP, DOCX, or a machine-specific path.
    """
    document = Document()
    normal = document.styles["Normal"]
    assert isinstance(normal, ParagraphStyle)
    normal.font.size = Pt(16)
    normal.font.name = "Times New Roman"
    normal.element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "SimSun")
    parent = document.styles.add_style("AttachmentBase", WD_STYLE_TYPE.PARAGRAPH)
    child = document.styles.add_style("Attachment", WD_STYLE_TYPE.PARAGRAPH)
    assert isinstance(parent, ParagraphStyle) and isinstance(child, ParagraphStyle)
    parent.base_style = normal
    child.base_style = parent
    document.add_paragraph("{{正文}}")
    anchor = (
        document.add_paragraph(style=child) if location == "body" else document.add_table(1, 1).cell(0, 0).paragraphs[0]
    )
    anchor.style = child
    anchor.add_run("{{附")
    anchor.add_run("件说明}}")
    if inheritance == "defaults":
        defaults = document.styles.element.find(qn("w:docDefaults"))
        assert defaults is not None
        properties = defaults.find(qn("w:pPrDefault"))
        if properties is None:
            properties = OxmlElement("w:pPrDefault")
            defaults.append(properties)
        paragraph_properties = properties.find(qn("w:pPr"))
        if paragraph_properties is None:
            paragraph_properties = OxmlElement("w:pPr")
            properties.append(paragraph_properties)
    else:
        element = (
            anchor._p if inheritance == "direct" else parent.element if inheritance == "ancestor" else child.element
        )
        paragraph_properties = element.get_or_add_pPr()
    indent = OxmlElement("w:ind")
    for name, value in {"leftChars": "200", "left": "1580", "hangingChars": "300", "hanging": "948"}.items():
        indent.set(qn(f"w:{name}"), value)
    paragraph_properties.append(indent)
    return document
