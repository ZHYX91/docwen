"""Native OOXML facts survive the real document route and Artifact Bundle."""

from __future__ import annotations

import hashlib
from pathlib import Path
from typing import Any
from zipfile import ZIP_DEFLATED, ZipFile

import pytest
from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches
from lxml import etree
from markdown_it import MarkdownIt

from docwen_application.bundle_mapping import build_bundle_draft
from docwen_application.controller import ApplicationController
from docwen_application.conversion_contracts import (
    DOCX_MEDIA_TYPE,
    DOCX_TO_MARKDOWN_CAPABILITY_ID,
    ConversionPlanRequest,
    LocalInputHandle,
    StagingOutputTarget,
)
from docwen_application.conversion_service import ConversionService
from docwen_core.models.file_ref import FileRef
from docwen_core.models.request import ConversionRequest, OutputPolicy
from docwen_core.paths import filesystem_path
from docwen_runtime.output.artifact_bundle import ArtifactBundleCommitter

pytestmark = [pytest.mark.integration, pytest.mark.pr_gate]

_W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
_V = "urn:schemas-microsoft-com:vml"
_WPS = "http://schemas.microsoft.com/office/word/2010/wordprocessingShape"
_MC = "http://schemas.openxmlformats.org/markup-compatibility/2006"
_REL = "http://schemas.openxmlformats.org/package/2006/relationships"
_CT = "http://schemas.openxmlformats.org/package/2006/content-types"


def _publish(
    runtime: Any,
    source: Path,
    target: Path,
    *,
    config_snapshot: dict[str, Any] | None = None,
    **options: Any,
) -> str:
    original = source.read_bytes()
    target.mkdir()
    if config_snapshot is None:
        controller = ApplicationController(runtime_port=runtime)
        controller.start()
        service = ConversionService(controller, ArtifactBundleCommitter())
        request = ConversionPlanRequest(
            capability_id=DOCX_TO_MARKDOWN_CAPABILITY_ID,
            inputs=(
                LocalInputHandle(
                    "native-source",
                    str(source),
                    DOCX_MEDIA_TYPE,
                    len(original),
                    hashlib.sha256(original).hexdigest(),
                    "document",
                    "source",
                    source.name,
                ),
            ),
            output=StagingOutputTarget(str(target)),
            options={
                "remove_numbering": False,
                "markdown_extensions": {"output": {}},
                **options,
            },
        )
        plan = service.plan(request)
        outcome = service.execute_accepted(service.accept(plan.plan_id))
        assert outcome.state == "completed", outcome.error
        assert outcome.bundle is not None
        assert outcome.bundle.layout_schema == "docwen.document_node.v1"
        bundle = outcome.bundle
    else:
        result = runtime.execute(
            ConversionRequest(
                request_id="native-visible-projection",
                input_refs=[FileRef(str(source), "docx", "document", len(original))],
                target_format="md",
                options={
                    "remove_numbering": False,
                    "markdown_extensions": {"output": {}},
                    **options,
                },
                output_policy=OutputPolicy(output_dir=str(target)),
                config_snapshot=config_snapshot,
            )
        )
        assert result.success, result.error
        draft = build_bundle_draft(
            profile="document_with_resources",
            output_media_type="text/markdown",
            artifacts=result.artifacts,
        )
        bundle = ArtifactBundleCommitter().commit(
            task_id="native-visible-projection", staging_root=str(target), draft=draft
        )
        assert bundle.layout_schema == "docwen.document_node.v1"
    preferred = [entry for entry in bundle.entries if entry.preferred]
    assert len(preferred) == 1
    artifact = next(item for item in bundle.artifacts if item.artifact_id == preferred[0].artifact_id)
    assert artifact.media_type == "text/markdown"
    output = filesystem_path(target / artifact.locator)
    payload = output.read_bytes()
    assert artifact.sha256 == hashlib.sha256(payload).hexdigest()
    assert artifact.size_bytes == len(payload)
    assert artifact.logical_path and len(Path(artifact.logical_path).parts) >= 2
    assert source.read_bytes() == original
    markdown = payload.decode("utf-8-sig")
    return markdown.split("\n---\n", 1)[1].lstrip("\n") if markdown.startswith("---\n") else markdown


def _xml(path: Path, name: str = "word/document.xml") -> Any:
    with ZipFile(path) as package:
        return etree.fromstring(package.read(name))


def _value(parent: Any, name: str, value: str) -> Any:
    item = OxmlElement("w:" + name)
    item.set(qn("w:val"), value)
    parent.append(item)
    return item


def _numpr(properties: Any, num_id: str, ilvl: int = 0) -> None:
    node = OxmlElement("w:numPr")
    _value(node, "ilvl", str(ilvl))
    _value(node, "numId", num_id)
    properties.append(node)


def _level(
    parent: Any,
    ilvl: int,
    fmt: str,
    template: str | None,
    *,
    start: int = 1,
    suffix: str = "space",
) -> Any:
    node = OxmlElement("w:lvl")
    node.set(qn("w:ilvl"), str(ilvl))
    _value(node, "start", str(start))
    _value(node, "numFmt", fmt)
    if template is not None:
        _value(node, "lvlText", template)
    _value(node, "suff", suffix)
    parent.append(node)
    return node


def _numbering(document: Any) -> Any:
    root = document.part.numbering_part.element
    for child in list(root):
        root.remove(child)
    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), "0")
    _value(abstract, "multiLevelType", "hybridMultilevel")
    root.append(abstract)
    return abstract


def _instance(document: Any, num_id: str) -> Any:
    instance = OxmlElement("w:num")
    instance.set(qn("w:numId"), num_id)
    _value(instance, "abstractNumId", "0")
    document.part.numbering_part.element.append(instance)
    return instance


def _textbox(parent: Any, kind: str, words: list[str]) -> Any:
    if kind == "dml":
        box = etree.SubElement(parent, f"{{{_WPS}}}txbx")
    else:
        shape = etree.SubElement(parent, f"{{{_V}}}shape")
        box = etree.SubElement(shape, f"{{{_V}}}textbox")
    content = etree.SubElement(box, qn("w:txbxContent"))
    for text in words:
        paragraph = etree.SubElement(content, qn("w:p"))
        run = etree.SubElement(paragraph, qn("w:r"))
        etree.SubElement(run, qn("w:t")).text = text
    return box


@pytest.mark.parametrize(
    "mode",
    [
        "dml",
        "vml",
        "same-anchor",
        "mixed-order",
        "separate-anchors",
        "choice",
        "choices",
        "fallback",
    ],
)
def test_native_textbox_occurrences_keep_complete_document_order(
    round_trip_runtime: Any, tmp_path: Path, mode: str
) -> None:
    document = Document()
    document.add_paragraph("Before.")
    anchor = document.add_paragraph("Anchor.")
    run = anchor.add_run()._r
    expected = ["Before.", "Anchor."]
    raw_boxes: list[list[str]]
    if mode in {"dml", "vml"}:
        _textbox(run, mode, ["A", "B", "A"])
        raw_boxes = [["A", "B", "A"]]
        expected.extend(["A", "B", "A"])
    elif mode == "same-anchor":
        _textbox(run, "vml", ["A", "B", "A"])
        _textbox(run, "dml", ["A", "B", "A"])
        raw_boxes = [["A", "B", "A"], ["A", "B", "A"]]
        expected.extend(["A", "B", "A"] * 2)
    elif mode == "mixed-order":
        for kind, text in [
            ("vml", "VML first."),
            ("dml", "DML middle."),
            ("vml", "VML last."),
        ]:
            _textbox(run, kind, [text])
        raw_boxes = [["VML first."], ["DML middle."], ["VML last."]]
        expected.extend(["VML first.", "DML middle.", "VML last."])
    elif mode == "separate-anchors":
        _textbox(run, "dml", ["A", "B", "A"])
        second = document.add_paragraph("Second anchor.")
        _textbox(second.add_run()._r, "vml", ["A", "B", "A"])
        raw_boxes = [["A", "B", "A"], ["A", "B", "A"]]
        expected.extend(["A", "B", "A", "Second anchor.", "A", "B", "A"])
    else:
        alternate = etree.SubElement(run, f"{{{_MC}}}AlternateContent", nsmap={"mc": _MC, "wps": _WPS})
        choice = etree.SubElement(alternate, f"{{{_MC}}}Choice", Requires="wps")
        choice_words = ["A", "B", "A"] if mode in {"choice", "choices"} else []
        _textbox(choice, "dml", choice_words)
        raw_boxes = [choice_words]
        if mode == "choices":
            other = etree.SubElement(alternate, f"{{{_MC}}}Choice", Requires="wps")
            _textbox(other, "dml", ["Unselected choice."])
            raw_boxes.append(["Unselected choice."])
        fallback = etree.SubElement(alternate, f"{{{_MC}}}Fallback")
        _textbox(fallback, "vml", ["A", "B", "A"])
        raw_boxes.append(["A", "B", "A"])
        expected.extend(["A", "B", "A"])
    document.add_paragraph("After.")
    expected.append("After.")
    source = tmp_path / "textbox.docx"
    document.save(str(source))
    xml = _xml(source)
    assert [
        ["".join(p.itertext()) for p in box.findall(qn("w:p"))] for box in xml.iter(qn("w:txbxContent"))
    ] == raw_boxes
    output = _publish(round_trip_runtime, source, tmp_path / "bundle")
    # markdown-it is independent of the production Mistune/source inventory.
    assert [token.content for token in MarkdownIt("commonmark").parse(output) if token.type == "inline"] == expected


@pytest.mark.parametrize("binding", ["direct", "heading-direct", "style"])
def test_concrete_numbering_overrides_keep_separate_native_counters(
    round_trip_runtime: Any, tmp_path: Path, binding: str
) -> None:
    document = Document()
    abstract = _numbering(document)
    _level(abstract, 0, "decimal", "%1.")
    for num_id, fmt in [("1", "upperLetter"), ("2", "upperRoman")]:
        instance = _instance(document, num_id)
        override = OxmlElement("w:lvlOverride")
        override.set(qn("w:ilvl"), "0")
        _level(override, 0, fmt, "%1.", start=4)
        instance.append(override)
        if binding == "style":
            style = document.styles.add_style("Bound " + num_id, WD_STYLE_TYPE.PARAGRAPH)
            _numpr(etree.SubElement(style.element, qn("w:pPr")), num_id)
    for num_id, text in [
        ("1", "Letter first."),
        ("2", "Roman first."),
        ("1", "Letter next."),
        ("2", "Roman next."),
    ]:
        paragraph = document.add_paragraph(
            text,
            style=("Heading 1" if binding == "heading-direct" else ("Bound " + num_id if binding == "style" else None)),
        )
        if binding != "style":
            _numpr(paragraph._p.get_or_add_pPr(), num_id)
    source = tmp_path / "concrete.docx"
    document.save(str(source))
    numbering = _xml(source, "word/numbering.xml")
    assert [
        e.get(qn("w:val"))
        for e in numbering.findall(f"{qn('w:num')}/{qn('w:lvlOverride')}/{qn('w:lvl')}/{qn('w:numFmt')}")
    ] == ["upperLetter", "upperRoman"]
    assert [
        e.get(qn("w:val"))
        for e in numbering.findall(f"{qn('w:num')}/{qn('w:lvlOverride')}/{qn('w:lvl')}/{qn('w:start')}")
    ] == ["4", "4"]
    output = _publish(round_trip_runtime, source, tmp_path / "bundle")
    assert [token.content for token in MarkdownIt("commonmark").parse(output) if token.type == "inline"] == [
        "D. Letter first.",
        "IV. Roman first.",
        "E. Letter next.",
        "V. Roman next.",
    ]


def test_compound_numbering_uses_the_same_concrete_parent_override(round_trip_runtime: Any, tmp_path: Path) -> None:
    document = Document()
    abstract = _numbering(document)
    _level(abstract, 0, "decimal", "%1.")
    _level(abstract, 1, "decimal", "%1.%2.")
    instance = _instance(document, "1")
    parent_override = OxmlElement("w:lvlOverride")
    parent_override.set(qn("w:ilvl"), "0")
    _level(parent_override, 0, "upperLetter", "%1.", start=2)
    _value(parent_override, "startOverride", "4")
    instance.append(parent_override)
    child_override = OxmlElement("w:lvlOverride")
    child_override.set(qn("w:ilvl"), "1")
    _value(child_override, "startOverride", "3")
    instance.append(child_override)
    for level, text in [(0, "Parent."), (1, "Child."), (1, "Child next.")]:
        paragraph = document.add_paragraph(text)
        _numpr(paragraph._p.get_or_add_pPr(), "1", level)
    source = tmp_path / "compound.docx"
    document.save(str(source))
    assert [e.get(qn("w:val")) for e in _xml(source, "word/numbering.xml").iter(qn("w:startOverride"))] == ["4", "3"]
    output = _publish(round_trip_runtime, source, tmp_path / "bundle")
    assert [line.strip() for line in output.splitlines() if line.strip()] == [
        "D. Parent.",
        "D.3. Child.",
        "D.4. Child next.",
    ]


@pytest.mark.parametrize(
    ("fmt", "template", "expected"),
    [
        ("decimal", "", "Text."),
        ("upperLetter", "", "Text."),
        ("upperLetter", None, "D Text."),
        ("decimal", "Article", "Article Text."),
        ("decimal", "%1.%2", "4. Text."),
    ],
)
def test_numbering_template_empty_missing_literal_and_future_level_have_distinct_meaning(
    round_trip_runtime: Any,
    tmp_path: Path,
    fmt: str,
    template: str | None,
    expected: str,
) -> None:
    document = Document()
    abstract = _numbering(document)
    _level(abstract, 0, fmt, template, start=4)
    _instance(document, "1")
    paragraph = document.add_paragraph("Text.")
    _numpr(paragraph._p.get_or_add_pPr(), "1")
    source = tmp_path / "template.docx"
    document.save(str(source))
    actual = _xml(source, "word/numbering.xml").find(f"{qn('w:abstractNum')}/{qn('w:lvl')}/{qn('w:lvlText')}")
    assert (actual is None) == (template is None)
    if actual is not None:
        assert actual.get(qn("w:val")) == template
    output = _publish(round_trip_runtime, source, tmp_path / "bundle")
    # Raw lines also prove the literal native marker; ordered AST nodes drop it.
    assert [line.strip() for line in output.splitlines() if line.strip()] == [expected]


def test_nine_level_template_uses_only_existing_ancestor_counters(round_trip_runtime: Any, tmp_path: Path) -> None:
    document = Document()
    abstract = _numbering(document)
    for level in range(9):
        _level(
            abstract,
            level,
            "decimal",
            "/".join(f"%{index}" for index in range(1, level + 2)),
            start=level + 1,
            suffix="nothing",
        )
    _instance(document, "1")
    paragraph = document.add_paragraph("Deep.")
    _numpr(paragraph._p.get_or_add_pPr(), "1", 8)
    source = tmp_path / "nine-levels.docx"
    document.save(str(source))
    levels = _xml(source, "word/numbering.xml").findall(f"{qn('w:abstractNum')}/{qn('w:lvl')}")
    assert [level.get(qn("w:ilvl")) for level in levels] == [str(index) for index in range(9)]
    output = _publish(round_trip_runtime, source, tmp_path / "bundle")
    assert output.strip() == "1/2/3/4/5/6/7/8/9Deep."


@pytest.mark.parametrize("preserve_formatting", [False, True])
def test_direct_hidden_body_runs_never_reappear_after_an_empty_visible_projection(
    round_trip_runtime: Any, tmp_path: Path, preserve_formatting: bool
) -> None:
    document = Document()
    hidden = document.add_paragraph().add_run("SECRET-PURE")
    hidden.font.hidden = True
    mixed = document.add_paragraph("Before ")
    mixed.add_run("SECRET-MIXED").font.hidden = True
    mixed.add_run("after.")
    visible = document.add_paragraph().add_run("Visible false.")
    visible.font.hidden = False
    ending = document.add_paragraph().add_run("SECRET-BREAK")
    ending.font.hidden = True
    break_run = document.paragraphs[-1].add_run()
    page = OxmlElement("w:br")
    page.set(qn("w:type"), "page")
    break_run._r.append(page)
    document.add_paragraph("After page.")
    source = tmp_path / "hidden.docx"
    document.save(str(source))
    assert [node.get(qn("w:val")) for node in _xml(source).iter(qn("w:vanish"))] == [
        None,
        None,
        "0",
        None,
    ]
    output = _publish(
        round_trip_runtime,
        source,
        tmp_path / "bundle",
        config_snapshot={"conversion": {"docx_to_md": {"preserve_formatting": preserve_formatting}}},
    )
    assert "SECRET" not in output
    tokens = MarkdownIt("commonmark").parse(output)
    assert [token.content for token in tokens if token.type == "inline"] == [
        "Before after.",
        "Visible false.",
        "After page.",
    ]
    assert len([token for token in tokens if token.type == "hr"]) == 1


@pytest.mark.parametrize("kind", ["footnote", "endnote"])
def test_native_note_hidden_suffix_and_reference_separator_keep_visible_ownership(
    round_trip_runtime: Any, tmp_path: Path, kind: str
) -> None:
    document = Document()
    paragraph = document.add_paragraph("Body")
    reference = OxmlElement("w:" + kind + "Reference")
    reference.set(qn("w:id"), "7")
    paragraph.add_run()._r.append(reference)
    source = tmp_path / "note.docx"
    document.save(str(source))
    with ZipFile(source) as package:
        members = {name: package.read(name) for name in package.namelist()}
    root = etree.Element(qn("w:" + kind + "s"), nsmap={"w": _W})
    note = etree.SubElement(root, qn("w:" + kind), {qn("w:id"): "7"})
    p = etree.SubElement(note, qn("w:p"))
    marker = etree.SubElement(p, qn("w:r"))
    etree.SubElement(marker, qn("w:" + kind + "Ref"))
    separator = etree.SubElement(p, qn("w:r"))
    etree.SubElement(
        separator,
        qn("w:t"),
        {"{http://www.w3.org/XML/1998/namespace}space": "preserve"},
    ).text = " "
    for text, hidden in [("Note visible.", False), ("SECRET-SUFFIX", True)]:
        r = etree.SubElement(p, qn("w:r"))
        props = etree.SubElement(r, qn("w:rPr"))
        vanish = etree.SubElement(props, qn("w:vanish"))
        if not hidden:
            vanish.set(qn("w:val"), "false")
        etree.SubElement(r, qn("w:t")).text = text
    part_name = f"word/{kind}s.xml"
    members[part_name] = etree.tostring(root, xml_declaration=True, encoding="UTF-8")
    types = etree.fromstring(members["[Content_Types].xml"])
    etree.SubElement(
        types,
        f"{{{_CT}}}Override",
        PartName="/" + part_name,
        ContentType=f"application/vnd.openxmlformats-officedocument.wordprocessingml.{kind}s+xml",
    )
    members["[Content_Types].xml"] = etree.tostring(types, xml_declaration=True, encoding="UTF-8")
    rels = etree.fromstring(members["word/_rels/document.xml.rels"])
    etree.SubElement(
        rels,
        f"{{{_REL}}}Relationship",
        Id="rIdNotes",
        Type=f"http://schemas.openxmlformats.org/officeDocument/2006/relationships/{kind}s",
        Target=f"{kind}s.xml",
    )
    members["word/_rels/document.xml.rels"] = etree.tostring(rels, xml_declaration=True, encoding="UTF-8")
    with ZipFile(source, "w", ZIP_DEFLATED) as package:
        for name, data in members.items():
            package.writestr(name, data)
    assert [node.text for node in _xml(source, part_name).iter(qn("w:t"))] == [
        " ",
        "Note visible.",
        "SECRET-SUFFIX",
    ]
    output = _publish(
        round_trip_runtime,
        source,
        tmp_path / "bundle",
        markdown_extensions={"output": {"typed_endnotes": True}},
    )
    marker_id = "1" if kind == "footnote" else "endnote:1"
    assert "SECRET" not in output
    assert [line for line in output.splitlines() if line.strip()] == [
        f"Body[^{marker_id}]",
        f"[^{marker_id}]: Note visible.",
    ]


@pytest.mark.parametrize("ticks", [3, 4, 17])
@pytest.mark.parametrize("context", ["plain", "list", "page", "crlf"])
def test_native_code_body_determines_one_safe_fence_per_physical_block(
    round_trip_runtime: Any, tmp_path: Path, ticks: int, context: str
) -> None:
    document = Document()
    style = document.styles.add_style("Code Block", WD_STYLE_TYPE.PARAGRAPH)
    expected = "  before_code  \n" + "`" * ticks + "\nafter_code\n"
    if context == "list":
        document.add_paragraph("List item.", style="List Bullet")
    if context == "crlf":
        paragraph = document.add_paragraph(style=style)
        run = paragraph.add_run()
        text = OxmlElement("w:t")
        text.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
        text.text = expected.removesuffix("\n").replace("\n", "\r\n")
        run._r.append(text)
        paragraphs = [paragraph]
    else:
        paragraphs = [document.add_paragraph(line, style=style) for line in expected.removesuffix("\n").split("\n")]
    if context == "list":
        for paragraph in paragraphs:
            paragraph.paragraph_format.left_indent = Inches(0.5)
    if context == "page":
        br = OxmlElement("w:br")
        br.set(qn("w:type"), "page")
        paragraphs[-1].add_run()._r.append(br)
    document.add_paragraph("ordinary end")
    source = tmp_path / "code.docx"
    document.save(str(source))
    native = _xml(source)
    assert [node.text for node in native.iter(qn("w:t")) if node.text and "`" in node.text] == (
        ["`" * ticks] if context != "crlf" else [expected.removesuffix("\n").replace("\n", "\r\n")]
    )
    output = _publish(round_trip_runtime, source, tmp_path / "bundle")
    tokens = MarkdownIt("commonmark").parse(output)
    assert [token.content for token in tokens if token.type == "fence"] == [expected]
    assert not [token for token in tokens if token.type == "code_block"]
    assert [token.content for token in tokens if token.type == "inline"] == (
        ["List item.", "ordinary end"] if context == "list" else ["ordinary end"]
    )
    assert len([token for token in tokens if token.type == "hr"]) == int(context == "page")
