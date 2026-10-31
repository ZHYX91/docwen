"""Shipped request cleanup cannot consume copy IDs before metadata parsing."""

from __future__ import annotations

import hashlib
import re
from io import BytesIO
from pathlib import Path
from typing import Any
from urllib.parse import unquote

import pytest
from docx import Document
from markdown_it import MarkdownIt
from PIL import Image

from docwen_bundle.config_port import ConfigPortAdapter
from docwen_bundle.runtime_factory import create_runtime_port
from docwen_runtime.config.loader import ConfigLoader
from tests.integration.test_formula_and_gongwen_source_content import _body, _bundle, _math, _r, _title, _xml

pytestmark = [pytest.mark.integration, pytest.mark.pr_gate]
_W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
_M = "http://schemas.openxmlformats.org/officeDocument/2006/math"
_NS = {"w": _W, "m": _M}


@pytest.fixture
def configured_ports(tmp_path: Path) -> tuple[Any, ConfigPortAdapter]:
    loader = ConfigLoader(
        base_dir=Path(__file__).resolve().parents[2] / "configs",
        user_dir=tmp_path / "user-config",
    )
    port = ConfigPortAdapter(loader)
    rules = port.snapshot()["numbering"]["cleanup"]["rules"]
    assert any(rule["id"] == "arabic_space" and rule["enabled"] for rule in rules)
    return create_runtime_port(config_loader=loader), port


@pytest.mark.parametrize("separator", ["", " ", "\u3000"], ids=["compact", "ascii-space", "ideographic-space"])
@pytest.mark.parametrize("number", ["国发〔2024〕1号", "国发〔2024〕 1号", "国发2024-1号"])
def test_typed_header_owns_copy_id_before_shipped_heading_cleanup(
    configured_ports: tuple[Any, ConfigPortAdapter], tmp_path: Path, separator: str, number: str
) -> None:
    runtime, config_port = configured_ports
    authored = "001" + separator + number
    document = Document()
    document.add_paragraph(authored)
    _title(document)
    document.add_paragraph("完整正文。")
    source = tmp_path / "configured-header.docx"
    document.save(source)
    assert _xml(source).find("w:body/w:p/w:r/w:t", _NS).text == authored
    output, _ = _bundle(
        runtime,
        source,
        tmp_path / "bundle",
        source_format="docx",
        capability="gongwen",
        config_port=config_port,
        options={"remove_numbering": True},
    )
    metadata, body = _body(output.read_text(encoding="utf-8"))
    assert metadata["份号"] == "001" and isinstance(metadata["份号"], str)
    assert metadata["发文字号"] == number
    assert body.splitlines() == ["完整正文。"]


@pytest.mark.parametrize("authored", ["001 国发〔2024〕号", "001 国发2024-号"])
def test_uncertain_header_keeps_complete_source_under_shipped_cleanup(
    configured_ports: tuple[Any, ConfigPortAdapter], tmp_path: Path, authored: str
) -> None:
    runtime, config_port = configured_ports
    document = Document()
    _title(document)
    document.add_paragraph("开始正文。")
    document.add_paragraph(authored)
    document.add_paragraph("尾部正文。")
    source = tmp_path / "configured-uncertain.docx"
    document.save(source)
    output, result = _bundle(
        runtime,
        source,
        tmp_path / "bundle",
        source_format="docx",
        capability="gongwen",
        config_port=config_port,
        options={"remove_numbering": True},
    )
    metadata, body = _body(output.read_text(encoding="utf-8"))
    assert metadata["份号"] == metadata["发文字号"] == ""
    assert [line for line in body.splitlines() if line] == ["开始正文。", authored, "尾部正文。"]
    [diagnostic] = [d for d in result.diagnostics if d.code == "GONGWEN-SOURCE-CONTENT-PRESERVED"]
    assert diagnostic.location == "paragraph:3:incomplete_combined_id"
    assert any(d.code == "GONGWEN-NEEDS-REVIEW" and "发文机关署名" in d.message for d in result.diagnostics)


def test_formula_bearing_header_does_not_lose_copy_prefix_or_resource_with_real_snapshot(
    configured_ports: tuple[Any, ConfigPortAdapter], tmp_path: Path
) -> None:
    runtime, config_port = configured_ports
    document = Document()
    authored = "001 国发〔2024〕 1号"
    paragraph = document.add_paragraph(authored)
    paragraph._p.append(_math(_r("x")))
    picture = BytesIO()
    Image.new("RGB", (3, 2), "blue").save(picture, format="PNG")
    picture_bytes = picture.getvalue()
    paragraph.add_run().add_picture(BytesIO(picture_bytes))
    _title(document)
    document.add_paragraph("完整正文。")
    source = tmp_path / "configured-rich-header.docx"
    document.save(source)
    first = _xml(source).find("w:body/w:p", _NS)
    assert [node.text for node in first.iter() if node.tag in {f"{{{_W}}}t", f"{{{_M}}}t"}] == [authored, "x"]
    output, result = _bundle(
        runtime,
        source,
        tmp_path / "bundle",
        source_format="docx",
        capability="gongwen",
        config_port=config_port,
        options={"remove_numbering": True, "recognize_text": False, "preserve_resources": True},
    )
    metadata, body = _body(output.read_text(encoding="utf-8"))
    assert metadata["份号"] == metadata["发文字号"] == ""
    assert body.splitlines().count(authored + "$x$") == 1
    assert body.splitlines().count("完整正文。") == 1
    tokens = MarkdownIt("commonmark").parse(body)
    assert not any(token.type == "heading_open" for token in tokens)
    [target] = re.findall(r"!\[[^\]]*\]\(([^)]+)\)", body)
    assert (
        hashlib.sha256((output.parent / unquote(target)).read_bytes()).digest()
        == hashlib.sha256(picture_bytes).digest()
    )
    assert any(d.code == "GONGWEN-SOURCE-CONTENT-PRESERVED" for d in result.diagnostics)


def test_ordinary_heading_cleanup_still_uses_the_same_shipped_request_rules(
    configured_ports: tuple[Any, ConfigPortAdapter], tmp_path: Path
) -> None:
    runtime, config_port = configured_ports
    document = Document()
    _title(document)
    document.add_paragraph("开始正文。")
    document.add_paragraph("一、常规标题")
    document.add_paragraph("说明正文。")
    document.add_paragraph("001 工作范围")
    document.add_paragraph("尾部正文。")
    source = tmp_path / "configured-ordinary-heading.docx"
    document.save(source)
    output, result = _bundle(
        runtime,
        source,
        tmp_path / "bundle",
        source_format="docx",
        capability="gongwen",
        config_port=config_port,
        options={"remove_numbering": True},
    )
    _, body = _body(output.read_text(encoding="utf-8"))
    tokens = MarkdownIt("commonmark").parse(body)
    assert [
        (token.tag, tokens[index + 1].content) for index, token in enumerate(tokens) if token.type == "heading_open"
    ] == [
        ("h1", "常规标题"),
        ("h3", "工作范围"),
    ]
    assert not any(d.code == "GONGWEN-SOURCE-CONTENT-PRESERVED" for d in result.diagnostics)
    assert body.index("开始正文。") < body.index("常规标题") < body.index("工作范围") < body.index("尾部正文。")
