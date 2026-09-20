"""Real production composition path from Machine plan to a Gongwen Bundle."""

from __future__ import annotations

import hashlib
from pathlib import Path
from typing import Any

import pytest
from docx import Document

from docwen_application.controller import ApplicationController
from docwen_application.conversion_contracts import ConversionPlanRequest, LocalInputHandle, StagingOutputTarget
from docwen_application.conversion_service import ConversionService
from docwen_core.formats import get_media_type
from docwen_core.paths import filesystem_path
from docwen_runtime.output.artifact_bundle import ArtifactBundleCommitter

pytestmark = pytest.mark.integration


def test_machine_executes_declared_gongwen_capability(round_trip_runtime: Any, tmp_path: Path) -> None:
    source = tmp_path / "关于开展工作的通知.docx"
    document = Document()
    document.add_paragraph("某某市人民政府文件")
    document.add_paragraph("某政发〔2026〕1号")
    document.add_heading("关于开展工作的通知", 0)
    document.add_paragraph("各有关单位：")
    document.add_paragraph("请按照要求开展工作。")
    document.add_paragraph("某某市人民政府")
    document.add_paragraph("2026年9月20日")
    document.save(str(source))
    original = source.read_bytes()
    staging = tmp_path / "output"
    staging.mkdir()
    controller = ApplicationController(runtime_port=round_trip_runtime)
    controller.start()
    service = ConversionService(controller, ArtifactBundleCommitter())
    capability = next(
        item
        for item in service.list_capabilities()
        if item.optimization_id == "gongwen" and item.input_shape.slots[0].media_types == (get_media_type("docx"),)
    )
    request = ConversionPlanRequest(
        capability_id=capability.capability_id,
        inputs=(
            LocalInputHandle(
                "source.1",
                str(source),
                get_media_type("docx"),
                len(original),
                hashlib.sha256(original).hexdigest(),
                "document",
                "source",
                source.name,
            ),
        ),
        output=StagingOutputTarget(str(staging)),
        options={"recognize_text": False, "preserve_resources": True},
    )
    plan = service.plan(request)
    outcome = service.execute_accepted(service.accept(plan.plan_id))
    assert outcome.state == "completed", outcome.error
    assert outcome.bundle is not None
    assert outcome.bundle.layout_schema == "docwen.document_node.v1"
    primary_id = next(entry.artifact_id for entry in outcome.bundle.entries if entry.preferred)
    artifact = next(item for item in outcome.bundle.artifacts if item.artifact_id == primary_id)
    assert len(Path(artifact.logical_path).parts) >= 2
    content = filesystem_path(staging / artifact.locator).read_text(encoding="utf-8")
    assert "关于开展工作的通知" in content and "请按照要求开展工作" in content
    assert source.read_bytes() == original
    assert not list(staging.rglob("docwen-node.json"))
