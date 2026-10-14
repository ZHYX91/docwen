"""Quote conversions preserve useful content across non-identical formats."""

from __future__ import annotations

from pathlib import Path
from zipfile import ZipFile

import pytest
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from docwen_plugin_markdown.document_semantics_v3 import analyze_markdown_semantics_v3
from docwen_plugin_markdown.mistune_extensions import parse_markdown_text
from tests.integration._round_trip_helper import docx_to_md, md_to_docx

pytestmark = pytest.mark.integration


def _body(markdown: str) -> str:
    return markdown.split("\n---\n", 1)[1].lstrip("\n") if markdown.startswith("---\n") else markdown


def _code_bodies(nodes: list) -> list[str]:
    output = []
    for node in nodes:
        if node["type"] == "block_code":
            output.append(node["raw"])
        output.extend(_code_bodies(node.get("children", [])))
    return output


def test_mixed_quote_children_keep_native_styles_and_readable_content(tmp_path: Path, round_trip_runtime) -> None:
    source = Path(__file__).parents[1] / "fixtures/files/quote-content.md"
    original = source.read_bytes()
    docx = md_to_docx(round_trip_runtime, source, tmp_path / "forward")
    document = Document(docx)
    paragraphs = {
        "".join(node.text or "" for node in paragraph._p.iter(qn("w:t"))): paragraph
        for paragraph in document.paragraphs
    }
    assert paragraphs["A quoted heading"].style.style_id == "Heading2"
    assert paragraphs['print("literal **code**")'].style.style_id == "DocWenCodeBlock"
    assert paragraphs["First item."].style.style_id == "DocWenListBlock"
    continuation = paragraphs["Continuation paragraph."]
    assert continuation._p.find(f"{qn('w:pPr')}/{qn('w:numPr')}/{qn('w:numId')}").get(qn("w:val")) == "0"
    properties = paragraphs["First item."]._p.find(qn("w:pPr"))
    assert properties.index(properties.find(qn("w:pStyle"))) < properties.index(properties.find(qn("w:numPr")))
    assert len(document.tables) == 2
    with ZipFile(docx) as package:
        assert b"_DW_C_" not in package.read("word/document.xml")
    returned = _body(docx_to_md(round_trip_runtime, docx, tmp_path / "reverse"))
    expected = [
        "Before quote.",
        "Introductory quote paragraph.",
        "A quoted heading",
        'print("literal **code**")',
        "00123",
        "First item.",
        "Continuation paragraph.",
        "Quote inside item.",
        "List inside nested quote.",
        "Another nested item.",
        "Second item.",
        "Quote after list.",
        "After quote.",
        "Table-only quote",
        "Final outside paragraph.",
    ]
    assert all(returned.count(text) == 1 for text in expected)
    assert [returned.index(text) for text in expected] == sorted(returned.index(text) for text in expected)
    assert _code_bodies(parse_markdown_text(returned)) == ['print("literal **code**")\n']
    assert "preserved" in returned and "42" in returned
    assert source.read_bytes() == original


def test_ordered_native_starts_and_overrides_generate_current_ordinals(tmp_path: Path, round_trip_runtime) -> None:
    source = tmp_path / "starts.md"
    source.write_text("> 7. First.\n> 8. Second.\n\nOutside.\n\n> 3. Third.\n> 4. Fourth.\n", encoding="utf-8")
    docx = md_to_docx(round_trip_runtime, source, tmp_path / "forward")
    returned = _body(docx_to_md(round_trip_runtime, docx, tmp_path / "reverse"))
    assert "7. First." in returned and "8. Second." in returned
    assert "3. Third." in returned and "4. Fourth." in returned
    document = Document(docx)
    first = next(p for p in document.paragraphs if p.text == "First.")
    num_id = first._p.find(f"{qn('w:pPr')}/{qn('w:numPr')}/{qn('w:numId')}").get(qn("w:val"))
    number = next(
        node for node in document.part.numbering_part.element.findall(qn("w:num")) if node.get(qn("w:numId")) == num_id
    )
    override = OxmlElement("w:lvlOverride")
    override.set(qn("w:ilvl"), "0")
    start = OxmlElement("w:startOverride")
    start.set(qn("w:val"), "12")
    override.append(start)
    number.append(override)
    edited = tmp_path / "override.docx"
    document.save(edited)
    overridden = _body(docx_to_md(round_trip_runtime, edited, tmp_path / "override-reverse"))
    assert "12. First." in overridden and "13. Second." in overridden
    assert "3. Third." in overridden


def test_list_item_then_quoted_table_preserves_visible_payload(tmp_path: Path, round_trip_runtime) -> None:
    body = "- Item before table.\n\n  > Quote before table.\n  >\n  > | Unique header |\n  > | --- |\n  > | Unique value |\n"
    source = tmp_path / "source.md"
    source.write_text(body, encoding="utf-8")
    docx = md_to_docx(round_trip_runtime, source, tmp_path / "forward")
    assert len(Document(docx).tables) == 1
    returned = _body(docx_to_md(round_trip_runtime, docx, tmp_path / "reverse"))
    texts = ["Item before table.", "Quote before table.", "Unique header", "Unique value"]
    assert all(returned.count(text) == 1 for text in texts)
    assert [returned.index(text) for text in texts] == sorted(returned.index(text) for text in texts)


@pytest.mark.parametrize("anchored", [False, True])
def test_table_first_list_item_keeps_table_content_and_optional_paragraph_id(
    tmp_path: Path, round_trip_runtime, anchored: bool
) -> None:
    source = tmp_path / "table-first.md"
    item_id = " ^item-id" if anchored else ""
    source.write_text(
        "- | Unique header |\n  | --- |\n  | Unique value |\n\n"
        f"  Continuation after table.{item_id}\n\nOutside after table.\n",
        encoding="utf-8",
    )
    source_analysis = analyze_markdown_semantics_v3(source.read_text(encoding="utf-8"), input_id="table-first.md")
    assert not source_analysis.has_errors
    source_ids = [(item["id"], item["block_kind"]) for item in source_analysis.projection["anchors"]]
    assert source_ids == ([("item-id", "paragraph")] if anchored else [])
    docx = md_to_docx(round_trip_runtime, source, tmp_path / "forward")
    assert len(Document(docx).tables) == 1
    returned = _body(docx_to_md(round_trip_runtime, docx, tmp_path / "reverse"))
    expected = ["Unique header", "Unique value", "Continuation after table.", "Outside after table."]
    assert all(returned.count(text) == 1 for text in expected)
    assert [returned.index(text) for text in expected] == sorted(returned.index(text) for text in expected)
    assert _code_bodies(parse_markdown_text(returned)) == []
    if anchored:
        analysis = analyze_markdown_semantics_v3(returned, input_id="table-first-returned.md")
        assert not analysis.has_errors
        assert [(item["id"], item["block_kind"]) for item in analysis.projection["anchors"]] == source_ids


@pytest.mark.parametrize("body", ["> $$x+1$$\n", "> ```python\n> ```\n", ">\n"])
def test_math_empty_fence_and_empty_quote_are_valid_inputs(tmp_path: Path, round_trip_runtime, body: str) -> None:
    source = tmp_path / "source.md"
    source.write_text(body, encoding="utf-8")
    docx = md_to_docx(round_trip_runtime, source, tmp_path / "forward")
    returned = _body(docx_to_md(round_trip_runtime, docx, tmp_path / "reverse"))
    if "x+1" in body:
        assert "x" in returned and "1" in returned
    if "```" in body:
        assert _code_bodies(parse_markdown_text(returned)) == [""]


def test_new_native_list_after_quote_keeps_body_and_marker(tmp_path: Path, round_trip_runtime) -> None:
    source = tmp_path / "quote.md"
    source.write_text("> Quote body.\n", encoding="utf-8")
    docx = md_to_docx(round_trip_runtime, source, tmp_path / "forward")
    document = Document(docx)
    document.add_paragraph("New native list item.", style="List Bullet")
    edited = tmp_path / "edited.docx"
    document.save(edited)
    returned = _body(docx_to_md(round_trip_runtime, edited, tmp_path / "reverse"))
    assert returned.count("Quote body.") == returned.count("New native list item.") == 1
    assert returned.index("Quote body.") < returned.index("- New native list item.")


@pytest.mark.parametrize(
    ("body", "expected"),
    [
        ("> - First item.\n>\n>   Continuation.\n", ["> - First item.", ">   Continuation."]),
        ("> ## Quoted heading\n", ["> ## Quoted heading"]),
        ("> | Unique header |\n> | --- |\n> | Unique value |\n", ["> | Unique header |", "> | Unique value |"]),
    ],
)
def test_existing_quote_id_keeps_a_legally_attached_visible_block(
    tmp_path: Path, round_trip_runtime, body: str, expected: list[str]
) -> None:
    source = tmp_path / "quote-id.md"
    source.write_text(body + "\n^quote-id\n", encoding="utf-8")
    docx = md_to_docx(round_trip_runtime, source, tmp_path / "forward")
    returned = _body(docx_to_md(round_trip_runtime, docx, tmp_path / "reverse"))
    assert returned.count("^quote-id") == 1
    assert all(returned.count(text) == 1 for text in expected)
    assert [returned.index(text) for text in expected] == sorted(returned.index(text) for text in expected)
    assert returned.index(expected[-1]) < returned.index("^quote-id")
    analysis = analyze_markdown_semantics_v3(returned, input_id="quote-id-returned.md")
    assert not analysis.has_errors
    assert [(item["id"], item["block_kind"]) for item in analysis.projection["anchors"]] == [
        ("quote-id", "block_quote")
    ]
