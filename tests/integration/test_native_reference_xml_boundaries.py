"""Authenticated native REF fields own XML nodes, not their surrounding runs."""

from __future__ import annotations

from pathlib import Path
from typing import Any
from zipfile import ZipFile

import pytest
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from lxml import etree
from markdown_it import MarkdownIt

from docwen_core.docx_bookmarks import build_docx_bookmark_inventory, prove_bookmark_name
from docwen_core.docx_parsing.document_semantics import encode_target_bookmark
from docwen_core.markdown_extensions import MarkdownExtensions
from docwen_plugin_markdown.runtime_semantics_v3 import prepare_runtime_semantics_v3
from tests.integration.test_visible_semantics_and_literal_code import _body

pytestmark = [pytest.mark.integration, pytest.mark.pr_gate]


def _text(run: Any, text: str) -> None:
    node = OxmlElement("w:t")
    node.set(qn("xml:space"), "preserve")
    node.text = text
    run.append(node)


def _field(run: Any, kind: str) -> None:
    node = OxmlElement("w:fldChar")
    node.set(qn("w:fldCharType"), kind)
    run.append(node)


@pytest.mark.parametrize("captions", [False, True], ids=["references-off", "references-on"])
@pytest.mark.parametrize("break_kind", ["none", "br", "cr", "page", "section"])
@pytest.mark.parametrize("single_run", [False, True], ids=["separate-boundary-runs", "one-run"])
def test_complex_reference_preserves_text_outside_its_xml_field_boundaries(
    tmp_path: Path, round_trip_runtime: Any, captions: bool, break_kind: str, single_run: bool
) -> None:
    document = Document()
    target = document.add_paragraph("Target")
    name = encode_target_bookmark("native-target")
    begin = OxmlElement("w:bookmarkStart")
    begin.set(qn("w:id"), "31")
    begin.set(qn("w:name"), name)
    end = OxmlElement("w:bookmarkEnd")
    end.set(qn("w:id"), "31")
    target._p.insert(0, begin)
    target._p.append(end)
    paragraph = document.add_paragraph()
    first = paragraph.add_run()._r
    _text(first, "Before ")
    _field(first, "begin")
    instruction_run = first if single_run else paragraph.add_run()._r
    instruction = OxmlElement("w:instrText")
    instruction.text = f" REF {name} \\h "
    instruction_run.append(instruction)
    _field(instruction_run, "separate")
    _text(instruction_run, "1")
    last = first if single_run else paragraph.add_run()._r
    _field(last, "end")
    _text(last, " after.")
    if break_kind in {"br", "cr", "page"}:
        node = OxmlElement("w:cr" if break_kind == "cr" else "w:br")
        if break_kind == "page":
            node.set(qn("w:type"), "page")
        last.append(node)
        _text(last, "Next.")
    elif break_kind == "section":
        paragraph._p.get_or_add_pPr().append(OxmlElement("w:sectPr"))
    document.add_paragraph("Outside.")
    source = tmp_path / "native-field.docx"
    document.save(source)
    with ZipFile(source) as package:
        root = etree.fromstring(package.read("word/document.xml"))
    [native] = [p for p in root.iter(qn("w:p")) if any(n.text == "Before " for n in p.iter(qn("w:t")))]
    [field_begin] = [node for node in native.iter(qn("w:fldChar")) if node.get(qn("w:fldCharType")) == "begin"]
    [field_end] = [node for node in native.iter(qn("w:fldChar")) if node.get(qn("w:fldCharType")) == "end"]
    assert field_begin.getprevious().tag == qn("w:t") and field_begin.getprevious().text == "Before "
    assert field_end.getnext().tag == qn("w:t") and field_end.getnext().text == " after."
    assert [node.text for node in native.iter(qn("w:t"))] == [
        "Before ",
        "1",
        " after.",
        *(["Next."] if break_kind in {"br", "cr", "page"} else []),
    ]
    fresh = Document(source)
    assert prove_bookmark_name(build_docx_bookmark_inventory(fresh), name).valid
    body = _body(round_trip_runtime, source, tmp_path / "final", captions=captions)
    tokens = MarkdownIt("commonmark").parse(body)
    token = "@native-target" if captions else "1"
    expected = f"Before {token} after."
    if break_kind in {"br", "cr"}:
        expected += "\nNext."
    actual = [item.content for item in tokens if item.type == "inline"]
    assert actual == ["Target", expected, *(["Next."] if break_kind == "page" else []), "Outside."]
    assert len([item for item in tokens if item.type == "hr"]) == int(break_kind in {"page", "section"})
    replan = prepare_runtime_semantics_v3(
        body,
        input_id="native-reader.md",
        extensions=MarkdownExtensions(captions_references=captions),
        consumer_profile="number_suite_direct",
    )
    assert not replan.analysis.has_errors


@pytest.mark.parametrize("captions", [False, True])
def test_non_ref_field_and_literal_code_are_not_semantic_replacements(
    tmp_path: Path, round_trip_runtime: Any, captions: bool
) -> None:
    document = Document()
    paragraph = document.add_paragraph()
    run = paragraph.add_run()._r
    _text(run, "Before ")
    _field(run, "begin")
    instruction = OxmlElement("w:instrText")
    instruction.text = " DATE "
    run.append(instruction)
    _field(run, "separate")
    _text(run, "2024")
    _field(run, "end")
    _text(run, " after.")
    code = document.add_paragraph().add_run("literal <br> @native-target")
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), "D9D9D9")
    code._r.get_or_add_rPr().append(shading)
    source = tmp_path / "literal.docx"
    document.save(source)
    assert Document(source).paragraphs[0].text == "Before 2024 after."
    body = _body(round_trip_runtime, source, tmp_path / "final", captions=captions)
    tokens = MarkdownIt("commonmark").parse(body)
    assert [item.content for item in tokens if item.type == "inline"] == ["Before 2024 after."]
    assert [item.content for item in tokens if item.type == "fence"] == ["literal <br> @native-target\n"]
