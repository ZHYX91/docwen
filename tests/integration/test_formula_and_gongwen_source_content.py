"""Source facts survive real Runtime, finalizer and Artifact Bundle conversion."""

from __future__ import annotations

import hashlib
import re
from io import BytesIO
from pathlib import Path
from typing import Any
from xml.etree import ElementTree as ET
from zipfile import ZipFile

import pytest
import yaml
from docx import Document
from docx.oxml import parse_xml
from docx.shared import Pt
from PIL import Image

from docwen_application.controller import ApplicationController
from docwen_application.conversion_contracts import ConversionPlanRequest, LocalInputHandle, StagingOutputTarget
from docwen_application.conversion_service import ConversionService
from docwen_core.formats import get_media_type
from docwen_core.paths import filesystem_path
from docwen_runtime.output.artifact_bundle import ArtifactBundleCommitter

pytestmark = [pytest.mark.integration, pytest.mark.pr_gate]
_W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
_M = "http://schemas.openxmlformats.org/officeDocument/2006/math"
_NS = {"w": _W, "m": _M}


def _bundle(
    runtime: Any,
    source: Path,
    target: Path,
    *,
    source_format: str,
    capability: str,
    options: dict | None = None,
    config_port: Any | None = None,
):
    authored = source.read_bytes()
    target.mkdir()
    controller = ApplicationController(runtime_port=runtime, config_port=config_port)
    controller.start()
    service = ConversionService(controller, ArtifactBundleCommitter())
    if capability == "gongwen":
        capability = next(
            c.capability_id
            for c in service.list_capabilities()
            if c.optimization_id == "gongwen" and c.input_shape.slots[0].media_types == (get_media_type("docx"),)
        )
    request = ConversionPlanRequest(
        capability_id=capability,
        inputs=(
            LocalInputHandle(
                "source.1",
                str(source),
                get_media_type(source_format),
                len(authored),
                hashlib.sha256(authored).hexdigest(),
                "document",
                "source",
                source.name,
            ),
        ),
        output=StagingOutputTarget(str(target)),
        options=options or {},
    )
    plan = service.plan(request)
    result = service.execute_accepted(service.accept(plan.plan_id))
    assert result.state == "completed", result.error
    assert result.bundle is not None
    entry = next(e for e in result.bundle.entries if e.preferred)
    artifact = next(a for a in result.bundle.artifacts if a.artifact_id == entry.artifact_id)
    output = filesystem_path(target / artifact.locator)
    assert output.exists() and source.read_bytes() == authored
    return output, result


def _xml(path: Path) -> ET.Element:
    with ZipFile(path) as package:
        return ET.fromstring(package.read("word/document.xml"))


def _math(children: str, *, block: bool = False):
    formula = f'<m:oMath xmlns:m="{_M}">{children}</m:oMath>'
    return parse_xml(f'<m:oMathPara xmlns:m="{_M}">{formula}</m:oMathPara>' if block else formula)


def _r(text: str) -> str:
    return f"<m:r><m:t>{text}</m:t></m:r>"


def _body(content: str) -> tuple[dict, str]:
    if content.startswith("---\n"):
        _, metadata, body = content.split("---", 2)
        return yaml.safe_load(metadata), body.strip()
    return {}, content.strip()


def _title(document) -> None:
    p = document.add_paragraph("关于公式和元数据的通知")
    p.runs[0].font.name = "小标宋"
    p.runs[0].font.size = Pt(22)
    document.add_paragraph("各有关单位：")


@pytest.mark.parametrize(
    ("latex", "expected"),
    [
        (r"m=3\text{ kg}", ["m", "=", "3", "\u00a0kg"]),
        (r"x\text{ if }x>0", ["x", "\u00a0if\u00a0", "x", ">", "0"]),
        (r"-\frac{m}{2}\text{ kg}", ["−", "m", "2", "\u00a0kg"]),
    ],
)
def test_markdown_text_math_keeps_exact_native_atoms(
    round_trip_runtime: Any, tmp_path: Path, latex: str, expected: list[str]
) -> None:
    source = tmp_path / "units.md"
    source.write_text(f"Before ${latex}$ after.\n", encoding="utf-8")
    output, _ = _bundle(
        round_trip_runtime,
        source,
        tmp_path / "bundle",
        source_format="markdown",
        capability="convert.markdown_source.to_docx",
    )
    root = _xml(output)
    maths = root.findall(".//m:oMath", _NS)
    assert len(maths) == 1
    assert [n.text for n in maths[0].findall(".//m:t", _NS)] == expected
    assert [n.text for n in root.findall(".//w:t", _NS)] == ["Before ", " after."]
    text_runs = [r for r in maths[0].findall(".//m:r", _NS) if r.find("m:rPr/m:nor", _NS) is not None]
    assert len(text_runs) == 1
    assert text_runs[0].find("m:t", _NS).get("{http://www.w3.org/XML/1998/namespace}space") == "preserve"


@pytest.mark.parametrize(
    ("native", "expected_tokens"),
    [
        (_r("α") + _r("x"), [r"\alpha", "x"]),
        (_r("α") + _r("+") + _r("x"), [r"\alpha", "+", "x"]),
        (
            "<m:sSup><m:e>" + _r("α") + _r("x") + "</m:e><m:sup>" + _r("2") + "</m:sup></m:sSup>",
            [r"\alpha", "x", "^", "2"],
        ),
        (
            "<m:sSub><m:e>" + _r("α") + _r("x") + "</m:e><m:sub>" + _r("1") + "</m:sub></m:sSub>",
            [r"\alpha", "x", "_", "1"],
        ),
        (
            "<m:m><m:mr><m:e>" + _r("α") + _r("x") + "</m:e><m:e>" + _r("y") + "</m:e></m:mr></m:m>",
            [r"\alpha", "x", "&", "y"],
        ),
        (_r("lim") + _r("x"), [r"\lim", "x"]),
    ],
)
def test_native_math_keeps_control_word_and_following_ascii_atoms(
    round_trip_runtime: Any, tmp_path: Path, native: str, expected_tokens: list[str]
) -> None:
    document = Document()
    p = document.add_paragraph("Before ")
    p._p.append(_math(native))
    p.add_run(" after.")
    source = tmp_path / "native.docx"
    document.save(source)
    # The oracle is native source atoms, independent of the conversion chain.
    assert [n.text for n in _xml(source).findall(".//m:t", _NS)] == [
        n.text for n in ET.fromstring(f'<m:oMath xmlns:m="{_M}">{native}</m:oMath>').findall(".//m:t", _NS)
    ]
    output, _ = _bundle(
        round_trip_runtime, source, tmp_path / "bundle", source_format="docx", capability="convert.docx.to_markdown"
    )
    _, body = _body(output.read_text(encoding="utf-8"))
    assert body.startswith("Before ") and body.endswith(" after.")
    assert body.count("Before") == body.count("after.") == 1
    [latex] = re.findall(r"(?<!\$)\$(?!\$)(.+?)(?<!\$)\$(?!\$)", body, re.S)
    latex = re.sub(r"\\(?:begin|end)\{[^}]+\}", "", latex)
    tokens = re.findall(r"\\[A-Za-z]+|[A-Za-z0-9+^_&]", latex)
    assert tokens == expected_tokens


@pytest.mark.parametrize(
    "prefix", ["金额为", "前" * 17, "前" * 18], ids=["short", "19-text-codepoints", "20-text-codepoints"]
)
@pytest.mark.parametrize("two", [False, True], ids=["one-formula", "two-formulas"])
def test_gongwen_keeps_xml_text_math_order_in_the_final_bundle(
    round_trip_runtime: Any, tmp_path: Path, prefix: str, two: bool
) -> None:
    document = Document()
    _title(document)
    document.add_paragraph("开始正文。")
    p = document.add_paragraph(prefix)
    p._p.append(_math(_r("x")))
    p.add_run("元。")
    if two:
        p._p.append(_math(_r("y")))
        p.add_run("结束。")
    document.add_paragraph("尾部正文。")
    source = tmp_path / "gongwen.docx"
    document.save(source)
    paragraph = next(p for p in _xml(source).findall(".//w:p", _NS) if p.find("m:oMath", _NS) is not None)
    assert [n.text for n in paragraph.iter() if n.tag in {f"{{{_W}}}t", f"{{{_M}}}t"}] == [prefix, "x", "元。"] + (
        ["y", "结束。"] if two else []
    )
    output, _ = _bundle(
        round_trip_runtime,
        source,
        tmp_path / "bundle",
        source_format="docx",
        capability="gongwen",
        options={"recognize_text": False, "preserve_resources": True},
    )
    _, body = _body(output.read_text(encoding="utf-8"))
    expected = prefix + "$x$元。" + ("$y$结束。" if two else "")
    assert body.splitlines().count(expected) == 1
    assert body.count("$x$") == 1 and body.count("$y$") == int(two)
    assert body.index("开始正文。") < body.index(expected) < body.index("尾部正文。")


def test_gongwen_standalone_omathpara_owns_nested_formula_once(round_trip_runtime: Any, tmp_path: Path) -> None:
    document = Document()
    _title(document)
    document.add_paragraph("开始正文。")
    document.add_paragraph()._p.append(_math(_r("x"), block=True))
    document.add_paragraph("尾部正文。")
    source = tmp_path / "display.docx"
    document.save(source)
    assert len(_xml(source).findall(".//m:oMath", _NS)) == 1
    output, _ = _bundle(round_trip_runtime, source, tmp_path / "bundle", source_format="docx", capability="gongwen")
    _, body = _body(output.read_text(encoding="utf-8"))
    assert body.splitlines().count("x") == 1
    assert body.count("$$") == 2
    assert body.index("开始正文。") < body.index("$$\nx\n$$") < body.index("尾部正文。")


@pytest.mark.parametrize(
    "number",
    ["国发〔2024〕1号", "国发〔2024〕 1号", "国发2024-1号", "国发[2024] 1号", "国发(2024)1号", "国发（2024）　1号"],
)
@pytest.mark.parametrize("separator", ["", " ", "\u3000"])
def test_combined_copy_and_document_number_preserve_owned_fields(
    round_trip_runtime: Any, tmp_path: Path, number: str, separator: str
) -> None:
    document = Document()
    authored = "001" + separator + number
    document.add_paragraph(authored)
    _title(document)
    document.add_paragraph("完整正文。")
    source = tmp_path / "combined.docx"
    document.save(source)
    assert _xml(source).find("w:body/w:p/w:r/w:t", _NS).text == authored
    output, _ = _bundle(round_trip_runtime, source, tmp_path / "bundle", source_format="docx", capability="gongwen")
    metadata, body = _body(output.read_text(encoding="utf-8"))
    assert metadata["份号"] == "001" and isinstance(metadata["份号"], str)
    assert metadata["发文字号"] == number
    assert body.splitlines().count("完整正文。") == 1
    assert authored not in body


@pytest.mark.parametrize(
    "text", ["001 国发〔2024〕号", "001 国发2024-号", "001 国发〔2024〕1号 附加正文", "001 普通正文。"]
)
def test_uncertain_combined_line_remains_visible(round_trip_runtime: Any, tmp_path: Path, text: str) -> None:
    document = Document()
    _title(document)
    document.add_paragraph("开始正文。")
    document.add_paragraph(text)
    document.add_paragraph("尾部正文。")
    source = tmp_path / "uncertain.docx"
    document.save(source)
    output, _ = _bundle(round_trip_runtime, source, tmp_path / "bundle", source_format="docx", capability="gongwen")
    metadata, body = _body(output.read_text(encoding="utf-8"))
    assert not metadata["份号"] and not metadata["发文字号"]
    assert body.splitlines().count(text) == 1
    assert body.index("开始正文。") < body.index(text) < body.index("尾部正文。")


def test_combined_metadata_cannot_consume_a_formula_and_keeps_the_image(
    round_trip_runtime: Any, tmp_path: Path
) -> None:
    document = Document()
    line = "001 国发〔2024〕1号"
    p = document.add_paragraph(line)
    p._p.append(_math(_r("x")))
    image = BytesIO()
    Image.new("RGB", (2, 2), "blue").save(image, format="PNG")
    image_bytes = image.getvalue()
    p.add_run().add_picture(BytesIO(image_bytes))
    _title(document)
    document.add_paragraph("完整正文。")
    source = tmp_path / "rich-metadata.docx"
    document.save(source)
    output, _ = _bundle(
        round_trip_runtime,
        source,
        tmp_path / "bundle",
        source_format="docx",
        capability="gongwen",
        options={"recognize_text": False, "preserve_resources": True},
    )
    metadata, body = _body(output.read_text(encoding="utf-8"))
    assert not metadata["份号"] and not metadata["发文字号"]
    assert body.splitlines().count(line + "$x$") == 1
    assert body.splitlines().count("完整正文。") == 1
    [target] = re.findall(r"!\[[^\]]*\]\(([^)]+)\)", body)
    assert hashlib.sha256((output.parent / target).read_bytes()).digest() == hashlib.sha256(image_bytes).digest()


@pytest.mark.parametrize("block", [False, True], ids=["inline", "display"])
def test_unknown_math_leaf_preserves_source_formula_and_reports_fallback(
    round_trip_runtime: Any, tmp_path: Path, block: bool
) -> None:
    latex = r"x\hspace{1em}y"
    source = tmp_path / "unsupported.md"
    source.write_text(
        f"Before.\n\n$$\n{latex}\n$$\n\nAfter.\n" if block else f"Before ${latex}$ after.\n",
        encoding="utf-8",
    )
    output, result = _bundle(
        round_trip_runtime,
        source,
        tmp_path / "bundle",
        source_format="markdown",
        capability="convert.markdown_source.to_docx",
    )
    root = _xml(output)
    assert not root.findall(".//m:oMath", _NS)
    paragraphs = ["".join(p.itertext()) for p in root.findall("w:body/w:p", _NS)]
    assert paragraphs == (["Before.", latex, "After."] if block else [f"Before {latex} after."])
    assert sum(d.code == "MD2DOCX-FORMULA-TEXT-FALLBACK" for d in result.diagnostics) == 1


@pytest.mark.parametrize("text", ["001 国发〔2024〕号", "001 国发2024-号"])
def test_incomplete_combined_line_has_precise_preserved_source_diagnostic(
    round_trip_runtime: Any, tmp_path: Path, text: str
) -> None:
    document = Document()
    _title(document)
    document.add_paragraph("开始正文。")
    document.add_paragraph(text)
    document.add_paragraph("尾部正文。")
    source = tmp_path / "incomplete-metadata.docx"
    document.save(source)
    output, result = _bundle(
        round_trip_runtime, source, tmp_path / "bundle", source_format="docx", capability="gongwen"
    )
    metadata, body = _body(output.read_text(encoding="utf-8"))
    assert not metadata["份号"] and not metadata["发文字号"]
    assert body.splitlines().count(text) == 1
    [diagnostic] = [d for d in result.diagnostics if d.code == "GONGWEN-SOURCE-CONTENT-PRESERVED"]
    assert diagnostic.location == "paragraph:3:incomplete_combined_id"
    assert diagnostic.level == "warning"
    # The pre-existing missing-signature warning still applies to this source.
    assert any(d.code == "GONGWEN-NEEDS-REVIEW" and "发文机关署名" in d.message for d in result.diagnostics)


def test_ordinary_year_prose_is_not_an_uncertain_combined_header(round_trip_runtime: Any, tmp_path: Path) -> None:
    text = "001 普通正文讲述2024年的工作。"
    document = Document()
    _title(document)
    document.add_paragraph(text)
    source = tmp_path / "ordinary-year.docx"
    document.save(source)
    output, result = _bundle(
        round_trip_runtime, source, tmp_path / "bundle", source_format="docx", capability="gongwen"
    )
    _, body = _body(output.read_text(encoding="utf-8"))
    assert body.splitlines().count(text) == 1
    assert not any(d.code == "GONGWEN-SOURCE-CONTENT-PRESERVED" for d in result.diagnostics)
