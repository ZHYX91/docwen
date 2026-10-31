"""New synthetic native DOCX cases for REF node ownership in rich runs.

The saved ZIP supplies source facts; CommonMark supplies a separate final-body
oracle. These cases do not alter any frozen semantic writer/proof fixture and
do not claim a real Word save/reopen observation.
"""

from __future__ import annotations

from pathlib import Path
from typing import Any
from zipfile import ZipFile

import pytest
from docx import Document
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from lxml import etree
from markdown_it import MarkdownIt

from docwen_core.docx_bookmarks import build_docx_bookmark_inventory, prove_bookmark_name
from docwen_core.docx_parsing.document_semantics import encode_target_bookmark
from docwen_core.markdown_extensions import MarkdownExtensions
from docwen_plugin_markdown.runtime_semantics_v3 import prepare_runtime_semantics_v3
from tests.integration.test_native_reference_xml_boundaries import _field, _text
from tests.integration.test_visible_semantics_and_literal_code import _body

pytestmark = [pytest.mark.integration, pytest.mark.pr_gate]

_URL = "https://example.invalid/native-ref?view=1#part"


def _bookmark(paragraph: Any, name: str, bookmark_id: str) -> None:
    begin = OxmlElement("w:bookmarkStart")
    begin.set(qn("w:id"), bookmark_id)
    begin.set(qn("w:name"), name)
    end = OxmlElement("w:bookmarkEnd")
    end.set(qn("w:id"), bookmark_id)
    paragraph._p.insert(0, begin)
    paragraph._p.append(end)


def _reference(run: Any, name: str, cached: str) -> None:
    _field(run, "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = f" REF {name} \\h "
    run.append(instruction)
    _field(run, "separate")
    _text(run, cached)
    _field(run, "end")


def _native_facts(path: Path, names: tuple[str, str], break_tag: str) -> None:
    with ZipFile(path) as package:
        root = etree.fromstring(package.read("word/document.xml"))
        relationships = etree.fromstring(package.read("word/_rels/document.xml.rels"))
    [paragraph] = [item for item in root.iter(qn("w:p")) if any(text.text == "Lead " for text in item.iter(qn("w:t")))]
    assert [text.text for text in paragraph.iter(qn("w:t"))] == [
        "Before",
        "Lead ",
        "1",
        " tail",
        " / ",
        "Linked ",
        "2",
        " after.",
    ]
    [native_break] = list(paragraph.iter(qn("w:" + break_tag)))
    assert native_break.getprevious().text == "Before"
    assert native_break.getnext().text == "Lead "
    runs = [item for item in paragraph.iter(qn("w:r")) if list(item.iter(qn("w:fldChar")))]
    assert len(runs) == 2
    assert runs[0].find(f"{qn('w:rPr')}/{qn('w:b')}") is not None
    assert runs[1].find(f"{qn('w:rPr')}/{qn('w:i')}") is not None
    for run, name, before, after in zip(runs, names, ("Lead ", "Linked "), (" tail", " after."), strict=True):
        markers = list(run.iter(qn("w:fldChar")))
        assert [node.get(qn("w:fldCharType")) for node in markers] == ["begin", "separate", "end"]
        assert markers[0].getprevious().text == before
        assert markers[-1].getnext().text == after
        assert [node.text for node in run.iter(qn("w:instrText"))] == [f" REF {name} \\h "]
    [hyperlink] = list(paragraph.iter(qn("w:hyperlink")))
    assert runs[1].getparent() is hyperlink
    relation_id = hyperlink.get(qn("r:id"))
    [relationship] = [node for node in relationships if node.get("Id") == relation_id]
    assert relationship.get("Type") == RT.HYPERLINK
    assert relationship.get("TargetMode") == "External"
    assert relationship.get("Target") == _URL
    fresh = Document(path)
    inventory = build_docx_bookmark_inventory(fresh)
    assert all(prove_bookmark_name(inventory, name).valid for name in names)


def _inline_characters(inline: Any) -> list[tuple[str, bool, bool, str | None]]:
    """Read visible text and actual CommonMark formatting/link ownership."""
    characters: list[tuple[str, bool, bool, str | None]] = []
    strong = italic = 0
    href: str | None = None
    for token in inline.children or []:
        if token.type == "strong_open":
            strong += 1
        elif token.type == "strong_close":
            strong -= 1
        elif token.type == "em_open":
            italic += 1
        elif token.type == "em_close":
            italic -= 1
        elif token.type == "link_open":
            assert href is None
            href = token.attrGet("href")
        elif token.type == "link_close":
            href = None
        elif token.type == "text":
            characters.extend((char, bool(strong), bool(italic), href) for char in token.content)
        elif token.type in {"softbreak", "hardbreak"}:
            characters.append(("\n", bool(strong), bool(italic), href))
        else:
            raise AssertionError(f"Unexpected final body token: {token.type}")
    assert strong == italic == 0 and href is None
    return characters


@pytest.mark.parametrize("captions", [False, True], ids=["references-off", "references-on"])
@pytest.mark.parametrize("break_tag", ["br", "cr"])
def test_two_native_references_preserve_rich_partial_runs_links_and_preceding_break(
    tmp_path: Path, round_trip_runtime: Any, captions: bool, break_tag: str
) -> None:
    document = Document()
    names = (encode_target_bookmark("native-one"), encode_target_bookmark("native-two"))
    _bookmark(document.add_paragraph("First target"), names[0], "41")
    _bookmark(document.add_paragraph("Second target"), names[1], "42")
    paragraph = document.add_paragraph()
    first = paragraph.add_run()
    first.bold = True
    _text(first._r, "Before")
    first._r.append(OxmlElement("w:" + break_tag))
    _text(first._r, "Lead ")
    _reference(first._r, names[0], "1")
    _text(first._r, " tail")
    paragraph.add_run(" / ")
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), paragraph.part.relate_to(_URL, RT.HYPERLINK, is_external=True))
    linked_run = OxmlElement("w:r")
    properties = OxmlElement("w:rPr")
    properties.append(OxmlElement("w:i"))
    linked_run.append(properties)
    _text(linked_run, "Linked ")
    _reference(linked_run, names[1], "2")
    _text(linked_run, " after.")
    hyperlink.append(linked_run)
    paragraph._p.append(hyperlink)
    document.add_paragraph("Outside.")
    source = tmp_path / "rich-native-fields.docx"
    document.save(source)
    _native_facts(source, names, break_tag)

    body = _body(round_trip_runtime, source, tmp_path / "final", captions=captions)
    inlines = [token for token in MarkdownIt("commonmark").parse(body) if token.type == "inline"]
    assert len(inlines) == 4
    assert [inlines[index].content for index in (0, 1, 3)] == ["First target", "Second target", "Outside."]
    characters = _inline_characters(inlines[2])
    one, two = ("@native-one", "@native-two") if captions else ("1", "2")
    expected = f"Before\nLead {one} tail / Linked {two} after."
    assert "".join(char for char, *_ in characters) == expected
    for text in ("Before", "Lead", "tail"):
        start = expected.index(text)
        assert all(char[1] for char in characters[start : start + len(text)])
    for text in ("Linked", "after."):
        start = expected.index(text)
        assert all(char[2] for char in characters[start : start + len(text)])
    linked_start = expected.index("Linked")
    assert all(char[3] is None for char in characters[:linked_start])
    assert all(char[3] == _URL for char in characters[linked_start:])
    replan = prepare_runtime_semantics_v3(
        body,
        input_id="rich-native-reader.md",
        extensions=MarkdownExtensions(captions_references=captions),
        consumer_profile="number_suite_direct",
    )
    assert not replan.analysis.has_errors
