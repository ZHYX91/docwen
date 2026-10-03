"""Declared Markdown keeps its output layout after content admission."""

from __future__ import annotations

import hashlib
from pathlib import Path
from typing import Any

import pytest
from docx import Document

from docwen_application.controller import ApplicationController
from docwen_application.conversion_contracts import ConversionPlanRequest, LocalInputHandle, StagingOutputTarget
from docwen_application.conversion_service import ConversionService
from docwen_core.models import ConversionRequest, FileRef, OutputPolicy
from docwen_core.paths import filesystem_path
from docwen_runtime.output.artifact_bundle import ArtifactBundleCommitter

pytestmark = [pytest.mark.integration, pytest.mark.pr_gate]


@pytest.mark.parametrize("authored", ["####### Level seven\n", "Plain paragraph.\n"])
def test_declared_markdown_source_retains_document_node_layout(
    round_trip_runtime: Any, tmp_path: Path, authored: str
) -> None:
    source = tmp_path / "authored.md"
    original = authored.encode("utf-8")
    source.write_bytes(original)
    staging = tmp_path / "output"
    staging.mkdir()
    controller = ApplicationController(runtime_port=round_trip_runtime)
    controller.start()
    service = ConversionService(controller, ArtifactBundleCommitter())
    request = ConversionPlanRequest(
        capability_id="convert.markdown_source.to_docx",
        inputs=(
            LocalInputHandle(
                "source.1",
                str(source),
                "text/markdown",
                len(original),
                hashlib.sha256(original).hexdigest(),
                "document",
                "source",
                "notes/authored.md",
            ),
        ),
        output=StagingOutputTarget(str(staging)),
        options={"markdown_extensions": {"input": {"extended_headings": True}}},
    )
    plan = service.plan(request)
    outcome = service.execute_accepted(service.accept(plan.plan_id))
    assert outcome.state == "completed", outcome.error
    assert outcome.bundle is not None
    assert outcome.bundle.layout_schema == "docwen.document_node.v1"
    [artifact] = outcome.bundle.artifacts
    output = filesystem_path(staging / artifact.locator)
    assert output.parent.name == output.stem
    document = Document(str(output))
    assert any(p.text == authored.strip().lstrip("# ") for p in document.paragraphs)
    if authored.startswith("#######"):
        assert any(p.style is not None and p.style.style_id == "Heading7" for p in document.paragraphs)
    assert source.read_bytes() == original

    direct = controller.execute_single(
        ConversionRequest(
            request_id="direct-markdown-layout",
            input_refs=[FileRef(str(source), "markdown", "markdown")],
            target_format="docx",
            output_policy=OutputPolicy(output_dir=str(tmp_path / "direct")),
            options=request.options,
        )
    )
    assert direct.success, direct.error
    [direct_artifact] = direct.artifacts
    direct_output = filesystem_path(direct_artifact.staging_path)
    assert direct_output.parent.name == direct_output.stem
    assert direct_artifact.metadata["document_node_schema"] == "docwen.document_node.v1"
    direct_document = Document(str(direct_output))
    assert [p.text for p in direct_document.paragraphs] == [p.text for p in document.paragraphs]
    assert source.read_bytes() == original

    exact = tmp_path / "exact.docx"
    refused = controller.execute_single(
        ConversionRequest(
            request_id="direct-markdown-exact-path",
            input_refs=[FileRef(str(source), "markdown", "markdown")],
            target_format="docx",
            output_policy=OutputPolicy(output_path=str(exact)),
            options=request.options,
        )
    )
    assert not refused.success
    assert not exact.exists()
