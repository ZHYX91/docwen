"""Image occurrences keep actual XML order across DrawingML and VML."""

from __future__ import annotations

import hashlib
import posixpath
from pathlib import Path
from typing import Any
from zipfile import ZipFile

import pytest
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from lxml import etree

from tests.integration.test_resource_identity_pipeline import _bundle, _image_targets, _png, _target_artifact

pytestmark = [pytest.mark.integration, pytest.mark.pr_gate]

_A = "http://schemas.openxmlformats.org/drawingml/2006/main"
_R = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
_V = "urn:schemas-microsoft-com:vml"


@pytest.mark.parametrize("container", ["paragraph", "table"])
@pytest.mark.parametrize("vml_first", [False, True])
def test_mixed_drawingml_vml_occurrences_keep_native_xml_order(
    tmp_path: Path, round_trip_runtime: Any, container: str, vml_first: bool
) -> None:
    document = Document()
    paragraph = (
        document.add_paragraph()
        if container == "paragraph"
        else document.add_table(rows=1, cols=1).cell(0, 0).paragraphs[0]
    )
    paragraph.add_run("Native image sequence.")
    originals = [_png((index * 45, 20, 230 - index * 40)) for index in range(5)]
    for index, content in enumerate(originals):
        image = tmp_path / f"source-{index}.png"
        image.write_bytes(content)
        shape = paragraph.add_run().add_picture(str(image))
        shape._inline.docPr.set("title", "photo.png")
        shape._inline.docPr.set("descr", f"Occurrence {index}")
        if (index % 2 == 0) == vml_first:
            rid = next(shape._inline.iter(f"{{{_A}}}blip")).get(qn("r:embed"))
            assert isinstance(rid, str)
            drawing = shape._inline.getparent()
            assert drawing is not None and drawing.getparent() is not None
            pict = OxmlElement("w:pict")
            node = etree.SubElement(pict, f"{{{_V}}}shape")
            etree.SubElement(node, f"{{{_V}}}imagedata").set(qn("r:id"), rid)
            drawing.getparent().replace(drawing, pict)
    source = tmp_path / "mixed-images.docx"
    document.save(str(source))
    with ZipFile(source) as package:
        root = etree.fromstring(package.read("word/document.xml"))
        relations = etree.fromstring(package.read("word/_rels/document.xml.rels"))
        targets = {item.get("Id"): item.get("Target") for item in relations}
        observed: list[bytes] = []
        for element in root.iter():
            if element.tag == f"{{{_A}}}blip":
                rid = element.get(f"{{{_R}}}embed")
            elif element.tag == f"{{{_V}}}imagedata":
                rid = element.get(f"{{{_R}}}id")
            else:
                continue
            observed.append(package.read(posixpath.normpath(posixpath.join("word", targets[rid]))))
        assert observed == originals
    bundle, markdown, primary = _bundle(
        round_trip_runtime, source, tmp_path / "output", image_link_style="markdown_embed"
    )
    assert markdown.count("Native image sequence.") == 1
    links = _image_targets(markdown)
    assert [_target_artifact(bundle, primary, target).sha256 for _, target in links] == [
        hashlib.sha256(data).hexdigest() for data in originals
    ]
