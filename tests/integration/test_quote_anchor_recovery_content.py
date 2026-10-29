"""Authenticated quote IDs survive content rendering and real break boundaries."""

from __future__ import annotations

import hashlib
from copy import deepcopy
from pathlib import Path
from typing import Any

import pytest
from docx import Document
from docx.enum.text import WD_BREAK

from docwen_core.docx_semantics_v3 import DocxSemanticsV3Session
from docwen_core.markdown_extensions import MarkdownExtensions
from docwen_plugin_markdown.runtime_semantics_v3 import prepare_runtime_semantics_v3
from tests.integration._round_trip_helper import docx_to_md
from tests.integration.test_markdown_visible_content import _NS, _text, _xml
from tests.integration.test_nested_quote_anchor_packages import _MAP_NS, _identity, _map_roots, _read, _write

pytestmark = [pytest.mark.integration, pytest.mark.pr_gate]


def _body(markdown: str) -> str:
    return markdown.split("\n---\n", 1)[1] if markdown.startswith("---\n") else markdown


@pytest.mark.parametrize("anchored", [False, True], ids=["anonymous", "paragraph-id"])
def test_explicit_quote_break_preserves_literal_code_and_one_paragraph_identity(
    tmp_path: Path, round_trip_runtime: Any, anchored: bool
) -> None:
    marker = " ^id" if anchored else ""
    source = f"> Before `literal <br>` and First<br>Second{marker}\n\nOutside.\n"
    path = _write(round_trip_runtime, tmp_path, source)
    root = _xml(path)
    assert [_text(p) for p in root.findall(".//w:p", _NS) if _text(p)] == [
        "Before literal <br> and FirstSecond",
        "Outside.",
    ]
    [line_break] = root.findall(".//w:br", _NS)
    assert line_break.get(f"{{{_NS['w']}}}type") in {None, "textWrapping"}
    assert root.find(".//w:bookmarkStart", _NS) is None
    assert root.find(".//w:instrText", _NS) is None
    maps = _map_roots(path)
    assert [record.attrib for item in maps for record in item.findall(f"{{{_MAP_NS}}}anchors/{{{_MAP_NS}}}anchor")] == (
        [_identity("id", "paragraph")] if anchored else []
    )

    returned = _read(round_trip_runtime, path, tmp_path)
    body = _body(returned)
    expected = (
        "> Before `literal <br>` and First<br>Second ^id" if anchored else "> Before `literal <br>` and First\n> Second"
    )
    assert expected in body
    assert body.count("`literal <br>`") == body.count("First") == body.count("Second") == 1
    assert body.count("^id") == int(anchored)
    assert body.index("Second") < body.index("Outside.")
    recovered_plan = prepare_runtime_semantics_v3(returned, input_id="recovered.md", extensions=MarkdownExtensions())
    assert not recovered_plan.analysis.has_errors
    assert [(anchor["id"], anchor["block_kind"]) for anchor in recovered_plan.analysis.projection["anchors"]] == (
        [("id", "paragraph")] if anchored else []
    )


def test_unanchored_outer_quote_cannot_take_its_already_bound_inner_quote(
    tmp_path: Path, round_trip_runtime: Any
) -> None:
    source = "> Before.\n>\n> > Inner.\n>\n> ^inner\n>\n> After.\n\nOutside.\n"
    path = _write(round_trip_runtime, tmp_path, source)
    root = _xml(path)
    assert [_text(p) for p in root.findall(".//w:p", _NS) if _text(p)] == ["Before.", "Inner.", "After.", "Outside."]
    expected = _identity("inner", "block_quote")
    wrappers = [
        wrapper
        for wrapper in root.findall(".//w:sdt", _NS)
        if (tag := wrapper.find("w:sdtPr/w:tag", _NS)) is not None and tag.get(f"{{{_NS['w']}}}val") == expected["tag"]
    ]
    assert len(wrappers) == 1
    content = wrappers[0].find("w:sdtContent", _NS)
    assert content is not None
    assert [node.tag for node in content] == [f"{{{_NS['w']}}}p"]
    assert _text(content) == "Inner."
    assert [
        record.attrib
        for item in _map_roots(path)
        for record in item.findall(f"{{{_MAP_NS}}}anchors/{{{_MAP_NS}}}anchor")
    ] == [expected]
    assert root.find(".//w:bookmarkStart", _NS) is None
    returned = _body(_read(round_trip_runtime, path, tmp_path))
    expected_texts = ["Before.", "Inner.", "^inner", "After.", "Outside."]
    assert all(returned.count(text) == 1 for text in expected_texts)
    assert [returned.index(text) for text in expected_texts] == sorted(returned.index(text) for text in expected_texts)


@pytest.mark.parametrize("ending", ["page", "section"])
@pytest.mark.parametrize("multiline", [False, True], ids=["one-line", "line-break"])
def test_native_quote_inline_id_precedes_its_trailing_page_or_section_separator(
    tmp_path: Path, round_trip_runtime: Any, ending: str, multiline: bool
) -> None:
    document = Document()
    paragraph = document.add_paragraph("First line", style="Quote")
    if multiline:
        paragraph.add_run().add_break(WD_BREAK.LINE)
        paragraph.add_run("Second line")
    if ending == "page":
        paragraph.add_run().add_break(WD_BREAK.PAGE)
    else:
        section = document.element.body.sectPr
        assert section is not None
        paragraph._p.get_or_add_pPr().append(deepcopy(section))
    document.add_paragraph("Outside.")
    session = DocxSemanticsV3Session(document, source_sha256=hashlib.sha256(b"native quote fixture").hexdigest())
    session.bind_paragraph_anchor(paragraph, {"block_kind": "paragraph", "id": "native"})
    session.finalize_document()
    path = tmp_path / "native.docx"
    document.save(str(path))
    session.write_package(path)

    root = _xml(path)
    expected = _identity("native", "paragraph")
    [wrapper] = root.findall(".//w:sdt", _NS)
    tag = wrapper.find("w:sdtPr/w:tag", _NS)
    assert tag is not None and tag.get(f"{{{_NS['w']}}}val") == expected["tag"]
    content = wrapper.find("w:sdtContent", _NS)
    assert content is not None
    assert [node.tag for node in content] == [f"{{{_NS['w']}}}p"]
    assert _text(content) == "First line" + ("Second line" if multiline else "")
    assert len(content.findall(".//w:br", _NS)) == int(multiline) + int(ending == "page")
    assert (content.find("w:p/w:pPr/w:sectPr", _NS) is not None) == (ending == "section")
    assert "Outside." not in _text(content)
    assert [
        record.attrib
        for item in _map_roots(path)
        for record in item.findall(f"{{{_MAP_NS}}}anchors/{{{_MAP_NS}}}anchor")
    ] == [expected]
    returned = _body(
        docx_to_md(
            round_trip_runtime,
            path,
            tmp_path / "reverse",
            options={
                "markdown_extensions": {"output": MarkdownExtensions().to_dict()},
                "page_break_separator": "---",
                "section_break_separator": "***",
            },
        )
    )
    expected_line = "> First line" + ("<br>Second line" if multiline else "") + " ^native"
    separator = "---" if ending == "page" else "***"
    assert expected_line in returned
    assert returned.count("^native") == 1
    assert f"{separator} ^native" not in returned
    assert returned.index("^native") < returned.index(f"\n{separator}\n") < returned.index("Outside.")
