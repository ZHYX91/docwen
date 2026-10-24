"""DOCX metadata survives the Markdown consumer's real YAML parser."""

from __future__ import annotations

from pathlib import Path

import pytest
from docx import Document

from docwen_core.cancellation import CancellationToken
from docwen_core.models.file_ref import FileRef
from docwen_core.models.request import ConversionRequest
from docwen_plugin_document.to_markdown.converter import DocxToMarkdownConverter
from docwen_plugin_markdown.yaml_processor import extract_yaml_front_matter
from tests.support.config import FakeConfigView
from tests.support.execution import FakeExecutionContext
from tests.support.logging import FakePluginLogger
from tests.support.progress import FakeProgressSink
from tests.support.workspace import FakeWorkspaceHandle

pytestmark = pytest.mark.integration


@pytest.mark.parametrize(
    ("title", "subtitle"),
    [
        (r"Report: C:\new\docs", r"C:\Users\name"),
        ("2026-10-24", "-1"),
        ("1.2", "00123456"),
        ('it\'s "ok"\\here', "true"),
    ],
)
def test_docx_titles_roundtrip_as_strings(tmp_path: Path, title: str, subtitle: str) -> None:
    source = tmp_path / "source.docx"
    document = Document()
    document.add_paragraph(title, style="Title")
    document.add_paragraph(subtitle, style="Subtitle")
    document.add_paragraph("Body remains here.")
    document.save(str(source))
    original = source.read_bytes()
    staging = tmp_path / "staging"
    staging.mkdir()
    context = FakeExecutionContext(
        request=ConversionRequest(
            request_id="metadata-yaml-roundtrip",
            input_refs=[FileRef(path=str(source), format="docx", category="document")],
            target_format="md",
            options={"locale": "en_US", "remove_numbering": False, "add_numbering": False},
        ),
        workspace=FakeWorkspaceHandle(str(source), str(staging)),
        config=FakeConfigView(),
        progress=FakeProgressSink(),
        cancellation=CancellationToken().view(),
        logger=FakePluginLogger(),
    )
    result = DocxToMarkdownConverter().convert(context)
    assert result.success, result.error
    content = Path(result.artifacts[0].staging_path).read_text(encoding="utf-8")
    metadata, body = extract_yaml_front_matter(content)
    assert metadata == {"aliases": [title], "title": title, "subtitle": subtitle}
    assert type(metadata["title"]) is str
    assert type(metadata["subtitle"]) is str
    assert body.strip() == "Body remains here."
    assert source.read_bytes() == original
