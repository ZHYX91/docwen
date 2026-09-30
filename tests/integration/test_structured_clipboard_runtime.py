"""Recursive clipboard documents through Application, Runtime, converters, and Finalizer."""

from __future__ import annotations

from pathlib import Path
from zipfile import ZipFile

import pytest
from docx import Document
from openpyxl import load_workbook

from docwen_application.controller import ApplicationController
from docwen_core.detection import inspect_file, inspect_structured_clipboard_snapshot
from docwen_core.models import FILE_INSPECTION_METADATA_KEY
from docwen_core.models.clipboard_document import CLIPBOARD_DOCUMENT_MEDIA_TYPE, clipboard_document_to_bytes
from docwen_core.models.file_ref import (
    MANAGED_INPUT_SHA256_METADATA_KEY,
    MANAGED_INPUT_SIZE_BYTES_METADATA_KEY,
    FileRef,
)
from docwen_core.models.request import ConversionRequest, OutputPolicy
from docwen_gui.clipboard_structured import project_structured_clipboard_html
from docwen_plugin_markdown.plugin import MarkdownPlugin
from docwen_runtime.adapters import RuntimePortAdapter
from docwen_runtime.engine.route_resolver import RouteResolver
from docwen_runtime.engine.task_manager import TaskManager
from docwen_runtime.output.finalizer import OutputFinalizer
from docwen_runtime.plugin_registry.registry import PluginRegistry
from docwen_runtime.templates import TemplateRegistry
from docwen_runtime.workspace.manager import WorkspaceManager

pytestmark = [pytest.mark.integration, pytest.mark.pr_gate]


HTML = b"""
<p>Intro 00123</p>
<table>
<thead>
<tr><th rowspan="2">Group</th><th colspan="2">Scores</th></tr>
<tr><th>A</th><th>B</th></tr>
</thead>
<tbody>
<tr><td rowspan="2">R1</td><td>10</td><td><p>before</p><table><tr><td>nested</td></tr></table><p>after</p></td></tr>
<tr><td>20</td><td>30</td></tr>
</tbody>
</table>
<p>Outro</p>
"""


def _template_id(kind: str, filename: str) -> str:
    matches = [item.id for item in TemplateRegistry.default().list_templates(kind) if item.path.name == filename]
    assert len(matches) == 1
    return matches[0]


def _controller(tmp_path: Path) -> ApplicationController:
    registry = PluginRegistry()
    registry.register(MarkdownPlugin())
    runtime = RuntimePortAdapter(
        TaskManager(
            registry,
            RouteResolver(registry),
            WorkspaceManager(root_dir=str(tmp_path / "runtime-workspaces")),
            OutputFinalizer(),
        )
    )
    return ApplicationController(runtime_port=runtime)


def _request(tmp_path: Path, target: str, *, options: dict | None = None, html: bytes = HTML) -> ConversionRequest:
    model = project_structured_clipboard_html(html)
    payload = clipboard_document_to_bytes(model)
    source = tmp_path / "managed.dwclip"
    source.write_bytes(payload)
    inspection = inspect_structured_clipboard_snapshot(str(source))
    return ConversionRequest(
        request_id=f"structured-{target}",
        input_refs=[
            FileRef(
                path=str(source),
                format="clipboard_document",
                category="markdown",
                size_bytes=len(payload),
                input_kind="document",
                input_role="source",
                logical_path="document.dwclip",
                media_type=CLIPBOARD_DOCUMENT_MEDIA_TYPE,
                metadata={
                    FILE_INSPECTION_METADATA_KEY: inspection.to_dict(),
                    MANAGED_INPUT_SHA256_METADATA_KEY: inspection.content_sha256,
                    MANAGED_INPUT_SIZE_BYTES_METADATA_KEY: inspection.size_bytes,
                },
            )
        ],
        target_format=target,
        options=options or {},
        output_policy=OutputPolicy(output_dir=str(tmp_path / f"out-{target}")),
    )


def _word_cell_block_tags(cell) -> list[str]:
    return [child.tag.rsplit("}", 1)[-1] for child in cell._tc if child.tag.rsplit("}", 1)[-1] != "tcPr"]


def test_docx_route_preserves_merges_header_and_nested_cell_order(tmp_path: Path) -> None:
    request = _request(
        tmp_path,
        "docx",
        options={"template_name": _template_id("docx", "English General Template.docx")},
    )
    result = _controller(tmp_path).execute_single(request)
    assert result.success, result.error
    output = Path(next(item for item in result.artifacts if item.is_primary).staging_path)
    document = Document(output)

    body_text = [paragraph.text for paragraph in document.paragraphs]
    assert "Intro 00123" in body_text
    assert "Outro" in body_text
    assert "{{body}}" not in "\n".join(body_text)
    assert "{{title}}" not in "\n".join(body_text)

    outer = next(
        table
        for table in document.tables
        if any(cell.text.startswith("Group") for row in table.rows for cell in row.cells)
    )
    assert outer.cell(0, 0)._tc is outer.cell(1, 0)._tc
    assert outer.cell(0, 1)._tc is outer.cell(0, 2)._tc
    header_xml = outer.rows[0]._tr.xml + outer.rows[1]._tr.xml
    assert header_xml.count("tblHeader") >= 2

    nested_anchor = outer.cell(2, 2)
    tags = _word_cell_block_tags(nested_anchor)
    table_index = tags.index("tbl")
    assert "p" in tags[:table_index]
    assert "p" in tags[table_index + 1 :]
    assert "before" in nested_anchor.paragraphs[0].text
    assert "after" in nested_anchor.paragraphs[-1].text
    assert nested_anchor.tables[0].cell(0, 0).text.startswith("nested")

    with ZipFile(output) as package:
        xml = package.read("word/document.xml").decode("utf-8")
    assert xml.index("before") < xml.index("nested") < xml.index("after")


def test_xlsx_route_preserves_string_values_merges_and_document_order(tmp_path: Path) -> None:
    request = _request(
        tmp_path,
        "xlsx",
        options={"template_name": _template_id("xlsx", "English Sample Sheet Template.xlsx")},
    )
    result = _controller(tmp_path).execute_single(request)
    assert result.success, result.error
    output = Path(next(item for item in result.artifacts if item.is_primary).staging_path)
    workbook = load_workbook(output)
    assert workbook.sheetnames[:3] == ["Document Order", "Table 1", "Table 2"]

    outer = workbook["Table 1"]
    assert {"A1:A2", "B1:C1", "A3:A4"} <= {str(item) for item in outer.merged_cells.ranges}
    assert outer["A1"].value == "Group"
    assert outer["B3"].value == "10"
    assert outer["B4"].value == "20"
    assert outer["B3"].data_type == "s"
    assert outer["B3"].number_format == "@"

    order = workbook["Document Order"]
    rows = [tuple(str(cell.value or "") for cell in row[:7]) for row in order.iter_rows(min_row=2)]
    assert any(row[1] == "paragraph" and row[2] == "Intro 00123" for row in rows)
    assert any(row[1] == "paragraph" and row[2] == "Outro" for row in rows)
    assert any(row[1] == "cell_paragraph" and row[2] == "before" and row[4] == "T1" for row in rows)
    assert any(row[1] == "table_ref" and row[4] == "T1" and row[5] == "R3C3" and row[6] == "T2" for row in rows)
    assert any(row[1] == "cell_paragraph" and row[2] == "after" and row[4] == "T1" for row in rows)


@pytest.mark.parametrize("target", ["md", "csv"])
def test_lossy_text_targets_publish_visible_projection_and_diagnostic(tmp_path: Path, target: str) -> None:
    result = _controller(tmp_path).execute_single(_request(tmp_path, target))
    assert result.success, result.error
    codes = {item.code for item in result.diagnostics}
    expected = "CLIPBOARD-MARKDOWN-STRUCTURE-PROJECTION" if target == "md" else "CLIPBOARD-CSV-STRUCTURE-PROJECTION"
    assert expected in codes
    paths = [Path(item.staging_path) for item in result.artifacts]
    if target == "md":
        text = paths[0].read_text(encoding="utf-8")
        assert "Intro 00123" in text and "before" in text and "nested" in text and "after" in text
        assert "<table" not in text.lower()
    else:
        assert len(paths) == 3
        combined = "\n".join(path.read_text(encoding="utf-8-sig") for path in paths)
        assert "Intro 00123" in combined and "before" in combined and "nested" in combined and "after" in combined


def test_ordinary_json_never_acquires_structured_clipboard_admission(tmp_path: Path) -> None:
    source = tmp_path / "ordinary.json"
    source.write_bytes(clipboard_document_to_bytes(project_structured_clipboard_html(HTML)))
    inspection = inspect_file(str(source))
    assert inspection.detected_format != "clipboard_document"
    assert inspection.decision.value != "allow"


def test_xlsx_authored_empty_strings_remain_distinct_from_covered_blank_cells(tmp_path: Path) -> None:
    from xml.etree import ElementTree

    html = b"""<table>
    <tr><th>empty</th><th>space</th><th>nbsp</th><th>zero</th><th>formula text</th>
    <th>pipe</th><th>less</th><th>caret</th><th>break</th></tr>
    <tr><td></td><td> </td><td>&nbsp;</td><td>00123</td><td>=1+1</td>
    <td>|</td><td>&lt;</td><td>^</td><td>before<br>after</td></tr>
    <tr><td colspan="9">merged anchor</td></tr></table>"""
    request = _request(
        tmp_path,
        "xlsx",
        options={"template_name": _template_id("xlsx", "English Sample Sheet Template.xlsx")},
        html=html,
    )
    result = _controller(tmp_path).execute_single(request)
    assert result.success, result.error
    output = Path(next(item for item in result.artifacts if item.is_primary).staging_path)
    workbook = load_workbook(output, data_only=False)
    table = workbook["Table 1"]
    expected = ("", " ", "\u00a0", "00123", "=1+1", "|", "<", "^", "before\nafter")
    assert tuple(table.cell(2, column).value for column in range(1, 10)) == expected
    assert all(table.cell(2, column).data_type == "s" for column in range(1, 10))
    assert table["A3"].value == "merged anchor"
    assert table["B3"].value is None

    namespace = {"m": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
    with ZipFile(output) as package:
        worksheets = [
            ElementTree.fromstring(package.read(name))
            for name in package.namelist()
            if name.startswith("xl/worksheets/") and name.endswith(".xml")
        ]
    table_xml = next(
        root
        for root in worksheets
        if any(node.attrib.get("ref") == "A3:I3" for node in root.findall("./m:mergeCells/m:mergeCell", namespace))
    )
    cells = {cell.attrib["r"]: cell for cell in table_xml.findall("./m:sheetData/m:row/m:c", namespace)}
    assert cells["A2"].attrib["t"] == "inlineStr"
    empty_text = cells["A2"].find("m:is/m:t", namespace)
    assert empty_text is not None and (empty_text.text or "") == ""
    assert cells["E2"].find("m:f", namespace) is None
    assert "B3" not in cells or cells["B3"].find("m:is", namespace) is None
