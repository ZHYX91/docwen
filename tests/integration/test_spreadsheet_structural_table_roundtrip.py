"""Spreadsheet output ↔ Structural Tables parser integration regressions."""

from __future__ import annotations

from pathlib import Path

import openpyxl
import pytest

from docwen_core.cancellation import CancellationToken
from docwen_core.markdown_extensions import MarkdownExtensions
from docwen_core.models.file_ref import FileRef
from docwen_core.models.request import ConversionRequest, OutputPolicy
from docwen_plugin_markdown.document_semantics import analyze_document_semantics
from docwen_plugin_markdown.mistune_extensions import parse_markdown_text
from docwen_plugin_spreadsheet.to_markdown.converter import SpreadsheetToMarkdownConverter
from tests.support.config import FakeConfigView
from tests.support.execution import FakeExecutionContext
from tests.support.logging import FakePluginLogger
from tests.support.progress import FakeProgressSink
from tests.support.workspace import FakeWorkspaceHandle

pytestmark = pytest.mark.integration


def _convert_structural(tmp_path: Path, workbook: openpyxl.Workbook) -> str:
    source = tmp_path / "source.xlsx"
    workbook.save(source)
    workbook.close()
    staging = tmp_path / "staging"
    staging.mkdir()
    ref = FileRef(path=str(source), format="xlsx", category="spreadsheet")
    request = ConversionRequest(
        request_id="spreadsheet-structural-roundtrip",
        input_refs=[ref],
        target_format="md",
        options={"markdown_extensions": {"output": {"structural_tables": True}}},
        output_policy=OutputPolicy(),
    )
    context = FakeExecutionContext(
        request=request,
        workspace=FakeWorkspaceHandle(str(source), str(staging)),
        config=FakeConfigView(),
        progress=FakeProgressSink(),
        cancellation=CancellationToken(),
        logger=FakePluginLogger(),
    )
    result = SpreadsheetToMarkdownConverter().convert(context)
    assert result.success, result.error
    return Path(result.artifacts[0].staging_path).read_text(encoding="utf-8")


def _table(markdown: str) -> str:
    return "\n".join(line for line in markdown.splitlines() if line.startswith("|"))


def _analyze_structural(table: str):
    ast = parse_markdown_text(
        table,
        extensions=MarkdownExtensions(structural_tables=True),
    )
    return analyze_document_semantics(ast, current_v3=True)


def test_literal_structural_markers_round_trip_as_literal_cells(tmp_path: Path) -> None:
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet["A1"] = "Kind"
    sheet["B1"] = "Value"
    sheet["A2"] = "<"
    sheet["B2"] = "^"

    analysis = _analyze_structural(_table(_convert_structural(tmp_path, workbook)))

    assert not analysis.has_errors
    metadata = analysis.ast[0]["_document_semantics_table"]
    assert all(anchor["row_span"] == anchor["column_span"] == 1 for anchor in metadata["anchors"])
    values = [
        child["raw"] for anchor in metadata["anchors"] for child in anchor["children"] if child.get("type") == "text"
    ]
    assert "<" in values
    assert "^" in values


def test_cross_header_merge_round_trips_as_zero_header_geometry(tmp_path: Path) -> None:
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet["A1"] = "Merged"
    sheet["B1"] = "X"
    sheet["B2"] = "Y"
    sheet.merge_cells("A1:A2")

    analysis = _analyze_structural(_table(_convert_structural(tmp_path, workbook)))

    assert not analysis.has_errors
    metadata = analysis.ast[0]["_document_semantics_table"]
    assert metadata["header_rows"] == 0
    assert any(
        anchor["row"] == 0 and anchor["column"] == 0 and anchor["row_span"] == 2 and anchor["column_span"] == 1
        for anchor in metadata["anchors"]
    )
