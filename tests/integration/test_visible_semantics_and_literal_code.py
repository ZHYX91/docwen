"""Visible semantic references and literal code use the same native run facts."""

from __future__ import annotations

import hashlib
from io import BytesIO
from pathlib import Path
from typing import Any
from zipfile import ZIP_DEFLATED, ZipFile

import pytest
from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from lxml import etree
from markdown_it import MarkdownIt

from docwen_application.bundle_mapping import build_bundle_draft
from docwen_core.docx_bookmarks import build_docx_bookmark_inventory, prove_bookmark_name
from docwen_core.docx_parsing.document_semantics import encode_target_bookmark
from docwen_core.docx_semantics_v3 import DocxSemanticsV3Recovery, DocxSemanticsV3Session
from docwen_core.markdown_extensions import MarkdownExtensions
from docwen_plugin_markdown.mistune_extensions import parse_markdown_text
from docwen_plugin_markdown.renderer import MdToDocxRenderer
from docwen_plugin_markdown.runtime_semantics_v3 import apply_runtime_semantics_v3, prepare_runtime_semantics_v3
from docwen_runtime.output.artifact_bundle import ArtifactBundleCommitter
from tests.integration._round_trip_helper import _run

pytestmark = [pytest.mark.integration, pytest.mark.pr_gate]


def _body(runtime: Any, source: Path, output: Path, *, captions: bool = True) -> str:
    original = source.read_bytes()
    result = _run(
        runtime,
        "visible-native-content",
        source,
        source_format="docx",
        target_format="md",
        output_dir=output,
        options={
            "remove_numbering": False,
            "markdown_extensions": {"output": MarkdownExtensions(captions_references=captions).to_dict()},
        },
    )
    assert result.success, result.error
    draft = build_bundle_draft(
        profile="document_with_resources", output_media_type="text/markdown", artifacts=result.artifacts
    )
    bundle = ArtifactBundleCommitter().commit(task_id="visible-native-content", staging_root=str(output), draft=draft)
    [preferred] = [entry.artifact_id for entry in bundle.entries if entry.preferred]
    [artifact] = [item for item in bundle.artifacts if item.artifact_id == preferred]
    data = (output / artifact.locator).read_bytes()
    assert artifact.sha256 == hashlib.sha256(data).hexdigest()
    assert source.read_bytes() == original
    text = data.decode("utf-8-sig")
    return text.split("\n---\n", 1)[1].lstrip("\n") if text.startswith("---\n") else text


def _legacy_source(path: Path, *, simple: bool) -> str:
    document = Document()
    target = document.add_paragraph("Target")
    bookmark_name = encode_target_bookmark("fig-target")
    start = OxmlElement("w:bookmarkStart")
    start.set(qn("w:name"), bookmark_name)
    start.set(qn("w:id"), "91")
    end = OxmlElement("w:bookmarkEnd")
    end.set(qn("w:id"), "91")
    target._p.insert(0, start)
    target._p.append(end)
    paragraph = document.add_paragraph("Before ")
    instruction = f" REF {bookmark_name} \\h "
    if simple:
        field = OxmlElement("w:fldSimple")
        field.set(qn("w:instr"), instruction)
        run = OxmlElement("w:r")
        text = OxmlElement("w:t")
        text.text = "1"
        run.append(text)
        field.append(run)
        paragraph._p.append(field)
    else:
        for kind in ("begin", "instruction", "separate", "result", "end"):
            run = paragraph.add_run()._r
            if kind in {"begin", "separate", "end"}:
                element = OxmlElement("w:fldChar")
                element.set(qn("w:fldCharType"), kind)
            else:
                element = OxmlElement("w:instrText" if kind == "instruction" else "w:t")
                element.text = instruction if kind == "instruction" else "1"
            run.append(element)
    paragraph.add_run(" after.")
    document.save(path)
    return "@fig-target"


def _v3_source(path: Path, *, soft: bool) -> str:
    token = "@[[#Target]]" if soft else "@[[#^target]]"
    source = f"# Target{'' if soft else ' ^target'}\n\nBefore {token} after.\n"
    dialect = MarkdownExtensions(captions_references=True)
    plan = prepare_runtime_semantics_v3(
        source, input_id="source.md", extensions=dialect, consumer_profile="number_suite_direct"
    )
    assert not plan.analysis.has_errors
    document = Document()
    session = DocxSemanticsV3Session(document, source_sha256=plan.source_sha256, standalone_caption_authority=True)
    parsed = parse_markdown_text(plan.shielded_source, extensions=dialect, auto_link_bare_url=False)
    assert isinstance(parsed, list)
    MdToDocxRenderer(doc=document, semantic_v3_session=session).render(apply_runtime_semantics_v3(parsed, plan))
    session.finalize_document()
    stream = BytesIO()
    document.save(stream)
    path.write_bytes(stream.getvalue())
    session.write_package(path)
    return token


def _suffix_variant(original: Path, candidate: Path, *, hidden: bool, page: bool) -> None:
    with ZipFile(original) as source, ZipFile(candidate, "w", ZIP_DEFLATED) as output:
        root = etree.fromstring(source.read("word/document.xml"))
        [paragraph] = [
            item
            for item in root.iter(qn("w:p"))
            if "Before " in "".join(node.text or "" for node in item.iter(qn("w:t")))
        ]
        run = etree.SubElement(paragraph, qn("w:r"))
        properties = etree.SubElement(run, qn("w:rPr"))
        vanish = etree.SubElement(properties, qn("w:vanish"))
        if not hidden:
            vanish.set(qn("w:val"), "0")
        etree.SubElement(run, qn("w:t")).text = "SUFFIX"
        if page:
            page_run = etree.SubElement(paragraph, qn("w:r"))
            etree.SubElement(page_run, qn("w:br")).set(qn("w:type"), "page")
        for info in source.infolist():
            output.writestr(
                info,
                etree.tostring(root, encoding="UTF-8", xml_declaration=True, standalone=True)
                if info.filename == "word/document.xml"
                else source.read(info.filename),
            )
    with ZipFile(original) as source, ZipFile(candidate) as output:
        assert source.namelist() == output.namelist()
        assert all(source.read(name) == output.read(name) for name in source.namelist() if name != "word/document.xml")


_CASES = [
    (kind, captions, hidden, False)
    for kind in ("legacy-complex", "legacy-simple", "v3-soft", "v3-ref")
    for captions in (False, True)
    for hidden in (False, True)
]
_CASES += [(kind, True, True, True) for kind in ("legacy-complex", "legacy-simple", "v3-soft", "v3-ref")]


@pytest.mark.parametrize("kind,captions,hidden,page", _CASES)
def test_direct_visibility_does_not_get_bypassed_by_a_proven_reference(
    tmp_path: Path, round_trip_runtime: Any, kind: str, captions: bool, hidden: bool, page: bool
) -> None:
    original = tmp_path / "writer.docx"
    token = (
        _legacy_source(original, simple=kind == "legacy-simple")
        if kind.startswith("legacy")
        else _v3_source(original, soft=kind == "v3-soft")
    )
    original_bytes = original.read_bytes()
    native = tmp_path / "native.docx"
    _suffix_variant(original, native, hidden=hidden, page=page)
    assert original.read_bytes() == original_bytes
    with ZipFile(native) as package:
        root = etree.fromstring(package.read("word/document.xml"))
    [suffix] = [node for node in root.iter(qn("w:t")) if node.text == "SUFFIX"]
    vanish = suffix.getparent().find(f"{qn('w:rPr')}/{qn('w:vanish')}")
    assert vanish is not None and vanish.get(qn("w:val")) == (None if hidden else "0")
    fresh = Document(BytesIO(native.read_bytes()))
    recovery = DocxSemanticsV3Recovery.load(native, fresh)
    if kind.startswith("legacy"):
        assert prove_bookmark_name(build_docx_bookmark_inventory(fresh), encode_target_bookmark("fig-target")).valid
        assert not recovery.soft_reference_identities and not recovery.reference_occurrence_identities
    else:
        assert len(recovery.soft_reference_identities) == int(kind == "v3-soft")
        assert len(recovery.reference_occurrence_identities) == int(kind == "v3-ref")
    body = _body(round_trip_runtime, native, tmp_path / "output", captions=captions)
    expected = f"Before {token if captions else '1'} after.{'' if hidden else 'SUFFIX'}"
    tokens = MarkdownIt("commonmark").parse(body)
    assert [item.content for item in tokens if item.type == "inline" and item.content.startswith("Before ")] == [
        expected
    ]
    assert body.count("SUFFIX") == int(not hidden)
    assert body.count("Before") == body.count("after.") == 1
    assert len([item for item in tokens if item.type == "hr"]) == int(page)


@pytest.mark.parametrize("page", [False, True])
def test_native_code_literal_projection_keeps_glyphs_links_ticks_and_body_default(
    tmp_path: Path, round_trip_runtime: Any, page: bool
) -> None:
    document = Document()
    code_style = document.styles.add_style("Code Block", WD_STYLE_TYPE.PARAGRAPH)
    lines = [
        "  before [literal](target) <br> `text`  ",
        "☐",
        "☑",
        "☒",
        "¨",
        "°",
        "×",
        "literal linked text",
        "```",
        "~~~~",
        "after",
    ]
    for index, text in enumerate(lines):
        paragraph = document.add_paragraph(style=code_style)
        if index == 7:
            link = OxmlElement("w:hyperlink")
            link.set(qn("r:id"), paragraph.part.relate_to("https://example.test/code", RT.HYPERLINK, is_external=True))
            run = OxmlElement("w:r")
            node = OxmlElement("w:t")
            node.text = text
            run.append(node)
            link.append(run)
            paragraph._p.append(link)
        else:
            run = paragraph.add_run(text)
            if 4 <= index <= 6:
                run.font.name = "Wingdings"
        if index == 0:
            paragraph.add_run("SECRET").font.hidden = True
        if index == len(lines) - 1 and page:
            run = paragraph.add_run()._r
            br = OxmlElement("w:br")
            br.set(qn("w:type"), "page")
            run.append(br)
    document.add_paragraph("ordinary end")
    document.add_paragraph("☐")
    source = tmp_path / "literal-code.docx"
    document.save(source)
    with ZipFile(source) as package:
        root = etree.fromstring(package.read("word/document.xml"))
    assert [node.text for node in root.iter(qn("w:t"))] == [lines[0], "SECRET", *lines[1:], "ordinary end", "☐"]
    body = _body(round_trip_runtime, source, tmp_path / "output")
    tokens = MarkdownIt("commonmark").parse(body)
    assert [item.content for item in tokens if item.type == "fence"] == ["\n".join(lines) + "\n"]
    assert not [item for item in tokens if item.type == "code_block"]
    assert [item.content for item in tokens if item.type == "inline"] == ["ordinary end", "[ ]"]
    assert len([item for item in tokens if item.type == "hr"]) == int(page)
    assert "SECRET" not in body
