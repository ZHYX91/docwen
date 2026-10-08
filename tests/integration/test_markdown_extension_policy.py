"""Exercise dialect switches on real DOCX bytes, independently in both directions."""

import hashlib
import json
from pathlib import Path
from typing import Any
from zipfile import ZipFile

import pytest
from docx import Document
from docx.oxml.ns import qn

from docwen_core.docx_parsing.document_semantics import extract_semantic_table_metadata
from docwen_core.docx_parsing.table_extraction import build_docx_table_semantic_grid
from docwen_core.markdown_extensions import EXTENSION_NAMES, MarkdownExtensions, resolve_markdown_extensions
from docwen_core.models.file_ref import FileRef
from docwen_core.models.request import ConversionRequest, OutputPolicy
from docwen_core.models.resolved_numbering import canonicalize_numbering_plan
from docwen_plugin_document.to_markdown.converter import DocxToMarkdownConverter
from docwen_plugin_markdown.to_docx.converter import MdToDocxConverter
from docwen_runtime.config.document_styles import build_document_style_catalog
from tests.support.cancellation import FakeCancellationTokenView
from tests.support.config import FakeConfigView
from tests.support.execution import FakeExecutionContext
from tests.support.logging import FakePluginLogger
from tests.support.progress import FakeProgressSink
from tests.support.workspace import FakeWorkspaceHandle

pytestmark = pytest.mark.integration
ROOT = Path(__file__).resolve().parents[2]
DIALECTS = [
    MarkdownExtensions(),
    MarkdownExtensions.obsidian(),
    *(MarkdownExtensions(**{name: True}) for name in EXTENSION_NAMES),
]
SOURCE = """# Example

**Strong** and *emphasis*.

Table: Metrics ^metrics

| Key | Value |
| --- | --- |
| A | 1 |
| B | ^ |

See @[[#Table: Metrics]].

####### Deep

Note[^endnote:n].

[^endnote:n]: Endnote body.
"""


def _context(tmp_path: Path, source: Path, target: str, extensions: dict) -> FakeExecutionContext:
    staging = tmp_path / f"staging-{target}"
    staging.mkdir()
    ref = FileRef(path=str(source), format="markdown" if target == "docx" else "docx", category="document")
    request = ConversionRequest(
        request_id=f"dialect-{target}",
        input_refs=[ref],
        target_format=target,
        options={"markdown_extensions": extensions, "heading_merge_mode": "never", "remove_numbering": False},
        output_policy=OutputPolicy(),
    )
    return FakeExecutionContext(
        request,
        FakeWorkspaceHandle(str(source), str(staging), (ref,)),
        FakeConfigView(),
        FakeProgressSink(),
        FakeCancellationTokenView(),
        FakePluginLogger(),
        document_style_catalog=build_document_style_catalog(
            {"gui": {"language": {"locale": "zh_CN"}}},
            locales_dir=ROOT / "i18n" / "locales",
        ),
    )


def _structural_config() -> FakeConfigView:
    return FakeConfigView({"conversion": {"markdown_extensions": {"input": {"structural_tables": True}}}})


def _structural_direct_context(tmp_path: Path, source: Path) -> FakeExecutionContext:
    staging = tmp_path / "staging-direct"
    staging.mkdir()
    ref = FileRef(path=str(source), format="markdown", category="document")
    request = ConversionRequest(
        request_id="structural-direct",
        input_refs=[ref],
        target_format="docx",
        options={"heading_merge_mode": "never"},
        output_policy=OutputPolicy(),
    )
    return FakeExecutionContext(
        request,
        FakeWorkspaceHandle(str(source), str(staging), (ref,)),
        _structural_config(),
        FakeProgressSink(),
        FakeCancellationTokenView(),
        FakePluginLogger(),
        document_style_catalog=build_document_style_catalog(
            {"gui": {"language": {"locale": "zh_CN"}}},
            locales_dir=ROOT / "i18n" / "locales",
        ),
    )


def _structural_resolved_context(tmp_path: Path, authored_markdown: str) -> FakeExecutionContext:
    plan_value = {"heading_definitions": [], "heading_instances": [], "targets": []}
    source_sha256 = hashlib.sha256(authored_markdown.encode()).hexdigest()
    plan_sha256 = hashlib.sha256(canonicalize_numbering_plan(plan_value)).hexdigest()
    document_value = {
        "$schema": "urn:docwen:schema:resolved-document:v1",
        "schema": "docwen.resolved_document.v1",
        "input_id": "structural-parity",
        "source_sha256": source_sha256,
        "plan_sha256": plan_sha256,
        "document": {
            "authored_markdown": authored_markdown,
            "targets": [],
            "references": [],
            "resource_occurrences": [],
            "citations": [],
            "resources": [],
        },
    }
    plan_envelope = {
        "$schema": "urn:docwen:schema:numbering-export-plan:v1",
        "schema": "docwen.numbering_export_plan.v1",
        "input_id": "structural-parity",
        "source_sha256": source_sha256,
        "plan_sha256": plan_sha256,
        "plan": plan_value,
    }
    neutral = tmp_path / "resolved-document.json"
    plan = tmp_path / "numbering-export-plan.json"
    neutral.write_text(json.dumps(document_value, separators=(",", ":")), encoding="utf-8")
    plan.write_text(json.dumps(plan_envelope, separators=(",", ":")), encoding="utf-8")
    refs = (
        FileRef(
            path=str(neutral),
            format="markdown",
            category="document",
            input_kind="document",
            input_role="neutral_document",
            logical_path="tables/structural.md",
            media_type="application/vnd.docwen.resolved-document+json",
        ),
        FileRef(
            path=str(plan),
            format="json",
            category="resource",
            input_kind="resource",
            input_role="numbering_export_plan",
            logical_path="numbering-export-plan.json",
            media_type="application/vnd.docwen.numbering-export-plan+json",
        ),
    )
    staging = tmp_path / "staging-resolved"
    staging.mkdir()
    request = ConversionRequest(
        request_id="structural-resolved",
        input_refs=list(refs),
        target_format="docx",
        options={"heading_merge_mode": "never"},
        output_policy=OutputPolicy(),
    )
    return FakeExecutionContext(
        request,
        FakeWorkspaceHandle(str(neutral), str(staging), refs),
        _structural_config(),
        FakeProgressSink(),
        FakeCancellationTokenView(),
        FakePluginLogger(),
        document_style_catalog=build_document_style_catalog(
            {"gui": {"language": {"locale": "zh_CN"}}},
            locales_dir=ROOT / "i18n" / "locales",
        ),
    )


def _table_signatures(path: Path) -> list[tuple[Any, ...]]:
    document = Document(path)

    def cell_text(cell, _row: int, _column: int) -> str:
        return "".join(node.text or "" for node in cell.iter(qn("w:t")))

    signatures: list[tuple[Any, ...]] = []
    for table in document.tables:
        metadata = extract_semantic_table_metadata(table._tbl)
        grid = build_docx_table_semantic_grid(table._tbl, cell_text_resolver=cell_text)
        signatures.append(
            (
                metadata.header_rows,
                metadata.header_columns,
                metadata.repeat_header,
                tuple(
                    tuple(
                        (
                            cell.anchor_text,
                            cell.anchor_row,
                            cell.anchor_col,
                            cell.rowspan,
                            cell.colspan,
                            cell.is_covered,
                        )
                        for cell in row
                    )
                    for row in grid
                ),
            )
        )
    return signatures


@pytest.mark.parametrize("dialect", DIALECTS)
def test_input_switches_control_native_structures_and_preserve_disabled_syntax(
    tmp_path: Path, dialect: MarkdownExtensions
) -> None:
    source = tmp_path / "source.md"
    source.write_text(SOURCE, encoding="utf-8")
    context = _context(tmp_path, source, "docx", {"input": dialect.to_dict()})
    result = MdToDocxConverter().convert(context)
    assert result.success, result.error
    assert len(result.artifacts) == 1
    output = Path(result.artifacts[0].staging_path)
    with ZipFile(output) as package:
        xml = package.read("word/document.xml").decode()
        assert ("SEQ Table" in xml) is dialect.captions_references
        assert ("<w:vMerge" in xml) is dialect.structural_tables
        assert ("<w:endnoteReference" in xml) is dialect.typed_endnotes
        assert ("<w:footnoteReference" in xml) is not dialect.typed_endnotes
        if not dialect.captions_references:
            assert "Table: Metrics ^metrics" in xml
            assert "@[[#Table: Metrics]]" in xml
        assert ("####### Deep" in xml) is not dialect.extended_headings
    assert source.read_text(encoding="utf-8") == SOURCE


@pytest.mark.parametrize("dialect", DIALECTS)
def test_output_switches_are_independent_and_need_only_docx(tmp_path: Path, dialect: MarkdownExtensions) -> None:
    source = tmp_path / "source.md"
    source.write_text(SOURCE, encoding="utf-8")
    forward = MdToDocxConverter().convert(
        _context(tmp_path, source, "docx", {"input": MarkdownExtensions.obsidian().to_dict()})
    )
    assert forward.success, forward.error
    isolated = tmp_path / "input.docx"
    isolated.write_bytes(Path(forward.artifacts[0].staging_path).read_bytes())
    source.unlink()
    context = _context(tmp_path, isolated, "md", {"output": dialect.to_dict()})
    result = DocxToMarkdownConverter().convert(context)
    assert result.success, result.error
    markdown = Path(result.artifacts[0].staging_path).read_text(encoding="utf-8")
    assert ("@[[#Table: Metrics]]" in markdown) is dialect.captions_references
    assert ("[^endnote:1]" in markdown) is dialect.typed_endnotes
    assert ("####### Deep" in markdown) is dialect.extended_headings
    assert "Endnote body." in markdown
    assert "| A | 1 |" in markdown
    assert "**Strong** and *emphasis*." in markdown
    assert ("| B | ^ |" in markdown) is dialect.structural_tables
    if not dialect.extended_headings:
        assert "###### Deep" in markdown
    if not dialect.typed_endnotes:
        assert "[^endnote-1]" in markdown
    if not dialect.captions_references:
        assert "^metrics" not in markdown
    if dialect != MarkdownExtensions.obsidian():
        assert any("flattened" in (item.code or "") for item in result.diagnostics)


def test_number_suite_note_identity_and_multiline_round_trip_from_isolated_docx(tmp_path: Path) -> None:
    source = tmp_path / "notes.md"
    source.write_text(
        "One[^endnote-topic], two[^Straße], three[^Strasse], rich[^rich].\n\n"
        "[^endnote-topic]: Ordinary footnote.\n"
        "[^Straße]: Sharp-s identity.\n"
        "[^Strasse]: Latin ss identity.\n"
        "[^rich]: First line\n"
        "  second line\n"
        "  **third line**\n",
        encoding="utf-8",
    )
    obsidian = MarkdownExtensions.obsidian().to_dict()

    forward_root = tmp_path / "notes-forward"
    forward_root.mkdir()
    forward = MdToDocxConverter().convert(_context(forward_root, source, "docx", {"input": obsidian}))
    assert forward.success, forward.error
    generated = Path(forward.artifacts[0].staging_path)
    with ZipFile(generated) as package:
        document_xml = package.read("word/document.xml").decode()
        assert document_xml.count("<w:footnoteReference") == 4
        assert "<w:endnoteReference" not in document_xml

    isolated = tmp_path / "notes-isolated.docx"
    isolated.write_bytes(generated.read_bytes())
    source.unlink()

    reverse_root = tmp_path / "notes-reverse"
    reverse_root.mkdir()
    reverse = DocxToMarkdownConverter().convert(_context(reverse_root, isolated, "md", {"output": obsidian}))
    assert reverse.success, reverse.error
    markdown = Path(reverse.artifacts[0].staging_path).read_text(encoding="utf-8")
    assert "[^endnote:" not in markdown
    assert "[^1]: Ordinary footnote." in markdown
    assert "[^2]: Sharp-s identity." in markdown
    assert "[^3]: Latin ss identity." in markdown
    assert "[^4]: First line\n    second line\n    **third line**" in markdown


@pytest.mark.parametrize("label,expected", [("n", "[^1]"), ("endnote:n", "[^endnote:1]")])
@pytest.mark.parametrize("complex_fields", [False, True])
def test_repeated_note_fields_survive_full_docx_only_conversion(tmp_path, label, expected, complex_fields):
    from copy import deepcopy

    from lxml import etree

    source = tmp_path / "repeated.md"
    source.write_text(
        f"First[^{label}].\n\n[^{label}]\n\nAgain[^{label}].\n\n[^{label}]: Kept note.\n", encoding="utf-8"
    )
    policy = MarkdownExtensions.obsidian().to_dict()
    forward_root = tmp_path / "forward"
    forward_root.mkdir()
    forward = MdToDocxConverter().convert(_context(forward_root, source, "docx", {"input": policy}))
    assert forward.success, forward.error
    isolated = tmp_path / "isolated.docx"
    isolated.write_bytes(Path(forward.artifacts[0].staging_path).read_bytes())
    source.unlink()
    if not complex_fields:
        doc = Document(isolated)
        instructions = [n for n in doc.element.iter(qn("w:instrText")) if (n.text or "").strip().startswith("NOTEREF ")]
        assert len(instructions) == 2
        for instruction in instructions:
            begin = instruction.getparent().getprevious()
            runs = [begin]
            for _ in range(4):
                runs.append(runs[-1].getnext())
            field = etree.Element(qn("w:fldSimple"), {qn("w:instr"): instruction.text})
            field.append(deepcopy(runs[3]))
            begin.addprevious(field)
            for run in runs:
                run.getparent().remove(run)
        doc.save(isolated)
    reverse_root = tmp_path / "reverse"
    reverse_root.mkdir()
    reverse = DocxToMarkdownConverter().convert(_context(reverse_root, isolated, "md", {"output": policy}))
    assert reverse.success, reverse.error
    markdown = Path(reverse.artifacts[0].staging_path).read_text(encoding="utf-8")
    assert f"First{expected}." in markdown
    assert f"\n\n{expected}\n\n" in markdown
    assert f"Again{expected}." in markdown
    assert f"{expected}: Kept note." in markdown


def test_structural_tables_direct_and_resolved_routes_share_docx_semantics(tmp_path: Path) -> None:
    authored = """| Region | Sales | < |
| Quarter | Q1 | Q2 |
| --- || --- | --- |
| North | 10 | 12 |
| ^ | 8 | 11 |

| --- | --- | --- |
| Block | < | Tail |
| ^ | ^ | Done |

| Code left | Strong up |
| --- | --- |
| `<` | **^** |
"""
    source = tmp_path / "structural.md"
    source.write_text(authored, encoding="utf-8")

    direct_root = tmp_path / "direct"
    direct_root.mkdir()
    direct = MdToDocxConverter().convert(_structural_direct_context(direct_root, source))
    assert direct.success, direct.error

    resolved_root = tmp_path / "resolved"
    resolved_root.mkdir()
    resolved = MdToDocxConverter().convert(_structural_resolved_context(resolved_root, authored))
    assert resolved.success, resolved.error

    direct_signatures = _table_signatures(Path(direct.artifacts[0].staging_path))
    resolved_signatures = _table_signatures(Path(resolved.artifacts[0].staging_path))
    assert resolved_signatures == direct_signatures
    assert len(direct_signatures) == 3
    assert direct_signatures[0][:2] == (2, 1)
    assert direct_signatures[1][:2] == (0, 0)
    assert any(cell[3:5] == (2, 2) for row in direct_signatures[1][3] for cell in row if not cell[5])


def test_structural_tables_in_quote_callout_and_list_render_as_native_docx_tables(tmp_path: Path) -> None:
    source = tmp_path / "container-tables.md"
    source.write_text(
        "> | - | - |\n"
        "> | Quote A | Quote B |\n"
        "> | Quote C | Quote D |\n\n"
        "> [!note]\n"
        ">\n"
        "> | - | - |\n"
        "> | Callout A | Callout B |\n"
        "> | Callout C | Callout D |\n\n"
        "- Item\n\n"
        "  | - | - |\n"
        "  | List A | List B |\n"
        "  | List C | List D |\n",
        encoding="utf-8",
    )
    root = tmp_path / "container-tables-out"
    root.mkdir()

    result = MdToDocxConverter().convert(_structural_direct_context(root, source))

    assert result.success, result.error
    document = Document(str(result.artifacts[0].staging_path))
    assert len(document.tables) == 3
    values = [[[cell.text for cell in row.cells] for row in table.rows] for table in document.tables]
    assert values[0] == [["Quote A", "Quote B"], ["Quote C", "Quote D"]]
    assert values[1] == [["Callout A", "Callout B"], ["Callout C", "Callout D"]]
    assert values[2] == [["List A", "List B"], ["List C", "List D"]]


def test_no_header_structural_table_round_trips_from_isolated_docx(tmp_path: Path) -> None:
    source = tmp_path / "no-header.md"
    source.write_text("| --- | --- |\n| Alice | 10 |\n| Bob | 20 |\n", encoding="utf-8")
    structural = MarkdownExtensions(structural_tables=True).to_dict()

    forward_root = tmp_path / "forward"
    forward_root.mkdir()
    forward_context = _context(forward_root, source, "docx", {"input": structural})
    forward = MdToDocxConverter().convert(forward_context)
    assert forward.success, forward.error
    generated = Path(forward.artifacts[0].staging_path)
    with ZipFile(generated) as package:
        xml = package.read("word/document.xml").decode()
        assert 'w:firstRow="0"' in xml

    isolated = tmp_path / "isolated.docx"
    isolated.write_bytes(generated.read_bytes())
    source.unlink()

    enabled_root = tmp_path / "enabled"
    enabled_root.mkdir()
    enabled_context = _context(enabled_root, isolated, "md", {"output": structural})
    enabled = DocxToMarkdownConverter().convert(enabled_context)
    assert enabled.success, enabled.error
    markdown = Path(enabled.artifacts[0].staging_path).read_text(encoding="utf-8")
    assert "| --- | --- |\n| Alice | 10 |\n| Bob | 20 |" in markdown

    disabled_root = tmp_path / "disabled"
    disabled_root.mkdir()
    disabled_context = _context(disabled_root, isolated, "md", {"output": MarkdownExtensions().to_dict()})
    disabled = DocxToMarkdownConverter().convert(disabled_context)
    assert disabled.success, disabled.error
    codes = {item.code for item in disabled.diagnostics}
    assert "docwen.conversion.markdown_extension.structural_tables.flattened" in codes


def test_request_override_does_not_change_other_direction_or_config() -> None:
    config = FakeConfigView({"conversion": {"markdown_extensions": {"input": MarkdownExtensions.obsidian().to_dict()}}})
    options = {"markdown_extensions": {"input": {"structural_tables": False}}}
    assert not resolve_markdown_extensions(options, config, direction="input").structural_tables
    assert resolve_markdown_extensions(options, config, direction="input").captions_references
    assert resolve_markdown_extensions(options, config, direction="output") == MarkdownExtensions()
    assert resolve_markdown_extensions({}, config, direction="input") == MarkdownExtensions.obsidian()


@pytest.mark.parametrize("dialect", [MarkdownExtensions(), MarkdownExtensions.obsidian()])
def test_real_word_saved_docx_uses_current_body_and_semantics(tmp_path: Path, dialect: MarkdownExtensions) -> None:
    isolated = tmp_path / "edited.docx"
    isolated.write_bytes((ROOT / "tests/fixtures/word-save/edited-markdown-extensions.docx").read_bytes())
    result = DocxToMarkdownConverter().convert(_context(tmp_path, isolated, "md", {"output": dialect.to_dict()}))
    assert result.success, result.error
    markdown = Path(result.artifacts[0].staging_path).read_text(encoding="utf-8")
    assert "Word 实测追加：从独立 DOCX 的当前内容回转。" in markdown
    assert "**粗体正文**与*斜体正文*应保留。" in markdown
    assert "| 甲 | 12 |" in markdown
    assert ("| 乙 | ^ |" in markdown) is dialect.structural_tables
    assert ("@[[#Table: 指标]]" in markdown) is dialect.captions_references
    assert ("####### 深层标题" in markdown) is dialect.extended_headings
    assert ("[^endnote:1]" in markdown) is dialect.typed_endnotes
    assert "这是尾注正文。" in markdown


def test_word_edit_that_adds_a_block_inside_caption_control_reports_invalid_structure(tmp_path: Path) -> None:
    isolated = tmp_path / "expanded-caption.docx"
    isolated.write_bytes((ROOT / "tests/fixtures/word-save/expanded-caption-control.docx").read_bytes())
    result = DocxToMarkdownConverter().convert(
        _context(tmp_path, isolated, "md", {"output": MarkdownExtensions.obsidian().to_dict()})
    )
    assert not result.success
    assert result.error is not None
    assert "one caption and one logical object" in result.error.message
    assert not result.artifacts
