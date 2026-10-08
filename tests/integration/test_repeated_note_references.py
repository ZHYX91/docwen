"""Repeated notes must open in Word without duplicating native note references."""

from copy import deepcopy
from io import BytesIO
from zipfile import ZipFile

import pytest
from docx import Document
from lxml import etree

from docwen_core.docx_parsing.format_features import DocxMarkdownSyntaxConfig
from docwen_plugin_document.shared.markdown_runs import render_paragraph_runs
from docwen_plugin_document.shared.note_extraction import NoteExtractor
from docwen_plugin_markdown.to_docx.converter import MdToDocxConverter
from docwen_plugin_markdown.to_docx.note_references import project_repeated_note_references
from tests.integration.test_markdown_extension_policy import _context

pytestmark = pytest.mark.integration
NS = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}


def _note_instructions(root):
    return [node for node in root.findall(".//w:instrText", NS) if (node.text or "").strip().startswith("NOTEREF ")]


def _cached_marks(root):
    return [node.getparent().getnext().getnext().find("w:t", NS).text for node in _note_instructions(root)]


@pytest.mark.parametrize(
    "restart,number_format,start",
    [
        ("eachPage", "decimal", 1),
        ("continuous", "chicago", 1),
        ("continuous", "upperRoman", 4000),
        ("continuous", "chineseCounting", 100),
        ("continuous", "decimalEnclosedCircle", 51),
    ],
)
def test_uncomputed_note_number_is_explicit_and_reported(tmp_path, restart, number_format, start):
    q = f"{{{NS['w']}}}"
    template = tmp_path / "template.docx"
    doc = Document()
    doc.add_paragraph("{{body}}")
    props = etree.SubElement(doc.sections[0]._sectPr, q + "footnotePr")
    etree.SubElement(props, q + "numFmt", {q + "val": number_format})
    etree.SubElement(props, q + "numRestart", {q + "val": restart})
    etree.SubElement(props, q + "numStart", {q + "val": str(start)})
    doc.save(str(template))
    source = tmp_path / "source.md"
    source.write_text("First[^a], repeated[^a].\n\n[^a]: Body.\n", encoding="utf-8")
    context = _context(tmp_path, source, "docx", {})
    context.request.options["template_name"] = str(template)
    result = MdToDocxConverter().convert(context)
    assert result.success, result.error
    warnings = [d for d in result.diagnostics if d.code == "MD2DOCX-NOTE-FIELD-UPDATE-REQUIRED"]
    assert len(warnings) == 1 and warnings[0].level == "warning"
    output = result.artifacts[0].staging_path
    with ZipFile(output) as archive:
        root = etree.fromstring(archive.read("word/document.xml"))
    assert _cached_marks(root) == ["?"]
    begin = _note_instructions(root)[0].getparent().getprevious()
    assert begin.find("w:fldChar", NS).get(q + "dirty") == "true"
    loaded = Document(output)
    extractor = NoteExtractor(loaded)
    text = "\n".join(
        render_paragraph_runs(p, extractor, syntax_config=DocxMarkdownSyntaxConfig()) for p in loaded.paragraphs
    )
    assert text.count("[^1]") == 2


@pytest.mark.parametrize(
    "kind,number_format,start,expected",
    [
        ("footnote", None, None, ["1", "2"]),
        ("endnote", None, None, ["i", "ii"]),
        ("footnote", "upperRoman", "4", ["IV", "V"]),
        ("endnote", "decimal", "7", ["7", "8"]),
        ("footnote", "lowerLetter", "3", ["c", "d"]),
        ("footnote", "upperLetter", "27", ["AA", "BB"]),
        ("endnote", "lowerLetter", "52", ["zz", "aaa"]),
    ],
)
@pytest.mark.parametrize("section_override", [False, True])
def test_repeated_note_cache_obeys_template_numbering(kind, number_format, start, expected, section_override):
    q = f"{{{NS['w']}}}"
    doc = Document()
    parent = doc.sections[0]._sectPr if section_override else doc.settings.element
    props = etree.SubElement(parent, q + kind + "Pr")
    if number_format:
        etree.SubElement(props, q + "numFmt", {q + "val": number_format})
    if start:
        etree.SubElement(props, q + "numStart", {q + "val": start})
    for note_id in (18, 4, 18, 4):
        run = doc.add_paragraph().add_run()._r
        etree.SubElement(run, q + kind + "Reference", {q + "id": str(note_id)})
    stream = BytesIO()
    doc.save(stream)
    xml = project_repeated_note_references(
        stream.getvalue(), {4, 18} if kind == "footnote" else set(), {4, 18} if kind == "endnote" else set()
    )
    root = etree.fromstring(xml.document_xml)
    assert _cached_marks(root) == expected


@pytest.mark.parametrize("restart,expected", [("eachSect", ["V", "i"]), ("continuous", ["V", "f"])])
def test_repeated_note_cache_keeps_original_section_number(restart, expected):
    q = f"{{{NS['w']}}}"
    doc = Document()
    defaults = etree.SubElement(doc.settings.element, q + "footnotePr")
    etree.SubElement(defaults, q + "numFmt", {q + "val": "lowerLetter"})
    etree.SubElement(defaults, q + "numStart", {q + "val": "3"})
    first = etree.SubElement(doc.sections[0]._sectPr, q + "footnotePr")
    etree.SubElement(first, q + "numFmt", {q + "val": "upperRoman"})
    etree.SubElement(first, q + "numStart", {q + "val": "5"})
    etree.SubElement(doc.add_paragraph().add_run()._r, q + "footnoteReference", {q + "id": "18"})
    second_section = doc.add_section()._sectPr
    second_section.remove(second_section.find(q + "footnotePr"))
    second = etree.SubElement(second_section, q + "footnotePr")
    etree.SubElement(second, q + "numStart", {q + "val": "9"})
    etree.SubElement(second, q + "numRestart", {q + "val": restart})
    for note_id in (4, 18, 4):
        etree.SubElement(doc.add_paragraph().add_run()._r, q + "footnoteReference", {q + "id": str(note_id)})
    stream = BytesIO()
    doc.save(stream)
    root = etree.fromstring(project_repeated_note_references(stream.getvalue(), {4, 18}, set()).document_xml)
    assert _cached_marks(root) == expected


@pytest.mark.parametrize("kind", ["footnote", "endnote"])
@pytest.mark.parametrize("custom", ["true", "1", "on", "false", "0", "off"])
def test_template_custom_marks_do_not_consume_automatic_numbers(kind, custom):
    q = f"{{{NS['w']}}}"
    doc = Document()
    props = etree.SubElement(doc.settings.element, q + kind + "Pr")
    etree.SubElement(props, q + "numFmt", {q + "val": "upperRoman"})
    for note_id in (10, 11, 18, 18):
        run = doc.add_paragraph().add_run()._r
        attributes = {q + "id": str(note_id)}
        if note_id == 11:
            attributes[q + "customMarkFollows"] = custom
        etree.SubElement(run, q + kind + "Reference", attributes)
        if note_id == 11:
            etree.SubElement(run, q + "t").text = "*"
    stream = BytesIO()
    doc.save(stream)
    root = etree.fromstring(
        project_repeated_note_references(
            stream.getvalue(), {18} if kind == "footnote" else set(), {18} if kind == "endnote" else set()
        ).document_xml
    )
    assert _cached_marks(root) == (["II"] if custom in {"true", "1", "on"} else ["III"])
    original = root.find(f".//w:{kind}Reference[@w:id='11']", NS)
    assert original.get(q + "customMarkFollows") == custom
    assert original.getnext().text == "*"


def test_repeated_note_bookmark_avoids_header_and_body_collisions():
    q = f"{{{NS['w']}}}"
    doc = Document()
    for paragraph, bookmark_id, name in (
        (doc.sections[0].header.paragraphs[0], "00", "_NOTE_FOOTNOTE_18"),
        (doc.add_paragraph(), "1", "_Note_footnote_18_1"),
    ):
        etree.SubElement(paragraph._p, q + "bookmarkStart", {q + "id": bookmark_id, q + "name": name})
        paragraph.add_run("Existing target")
        etree.SubElement(paragraph._p, q + "bookmarkEnd", {q + "id": bookmark_id})
    for _ in range(2):
        etree.SubElement(doc.add_paragraph().add_run()._r, q + "footnoteReference", {q + "id": "18"})
    stream = BytesIO()
    doc.save(stream)
    root = etree.fromstring(project_repeated_note_references(stream.getvalue(), {18}, set()).document_xml)
    target = root.find(".//w:bookmarkStart[@w:name='_Note_footnote_18_2']", NS)
    assert target is not None and target.get(q + "id") == "2"
    assert root.find(".//w:bookmarkEnd[@w:id='2']", NS) is not None
    assert _note_instructions(root)[0].text.split()[1] == "_Note_footnote_18_2"


@pytest.mark.parametrize("kind,label", [("footnote", "n"), ("endnote", "endnote:n")])
@pytest.mark.parametrize("in_table", [False, True])
@pytest.mark.parametrize("complex_fields", [False, True])
@pytest.mark.parametrize("proofing", [False, True])
def test_repeated_notes_use_one_native_reference_and_linked_fields(
    tmp_path, kind, label, in_table, complex_fields, proofing
):
    repeated = f"Again[^{label}], third[^{label}]."
    if in_table:
        repeated = f"| Column |\n| --- |\n| {repeated} |"
    source = tmp_path / "repeated.md"
    source.write_text(f"First[^{label}].\n\n{repeated}\n\n[^{label}]: **One** note body.\n", encoding="utf-8")
    context = _context(tmp_path, source, "docx", {"input": {"typed_endnotes": True}})
    result = MdToDocxConverter().convert(context)
    assert result.success, result.error
    with ZipFile(result.artifacts[0].staging_path) as archive:
        document = etree.fromstring(archive.read("word/document.xml"))
        references = document.findall(f".//w:{kind}Reference", NS)
        assert len(references) == 1, "Word repairs duplicate native references to one note ID"
        fields = _note_instructions(document)
        instructions = [field.text for field in fields]
        assert len(instructions) == 2 and all(value.strip().startswith("NOTEREF ") for value in instructions)
        names = {value.split()[1] for value in instructions}
        assert len(names) == 1
        starts = document.findall(".//w:bookmarkStart", NS)
        assert sum(node.get(f"{{{NS['w']}}}name") in names for node in starts) == 1
        for field in fields:
            run = field.getparent().getprevious()
            for index in range(5):
                assert run.find("w:rPr/w:vertAlign", NS).get(f"{{{NS['w']}}}val") == "superscript"
                if index == 3:
                    assert run.find("w:t", NS).text, "Keep a visible cached number for non-updating readers"
                run = run.getnext()
        notes = etree.fromstring(archive.read(f"word/{kind}s.xml"))
        assert len([node for node in notes if int(node.get(f"{{{NS['w']}}}id", "0")) > 0]) == 1
    path = str(result.artifacts[0].staging_path)
    doc = Document(path)
    if proofing and complex_fields:
        q = f"{{{NS['w']}}}"
        for instruction in _note_instructions(doc.element):
            cache = instruction.getparent().getnext().getnext()
            cache.addprevious(etree.Element(q + "proofErr", {q + "type": "spellStart"}))
            cache.addnext(etree.Element(q + "proofErr", {q + "type": "spellEnd"}))
    if not complex_fields:
        # External producers may use simple fields. Preserve inverse coverage
        # even though our writer uses complex fields to retain Word formatting.
        q = f"{{{NS['w']}}}"
        for instruction in _note_instructions(doc.element):
            begin = instruction.getparent().getprevious()
            parts = [begin]
            for _ in range(4):
                parts.append(parts[-1].getnext())
            field = etree.Element(q + "fldSimple", {q + "instr": instruction.text})
            field.append(deepcopy(parts[3]))
            begin.addprevious(field)
            for run in parts:
                run.getparent().remove(run)
        path = str(tmp_path / "simple-notes.docx")
        doc.save(path)
        doc = Document(path)
    extractor = NoteExtractor(doc, path)
    paragraphs = list(doc.paragraphs)
    paragraphs.extend(p for table in doc.tables for row in table.rows for cell in row.cells for p in cell.paragraphs)
    text = "\n".join(render_paragraph_runs(p, extractor, syntax_config=DocxMarkdownSyntaxConfig()) for p in paragraphs)
    expected = "[^1]" if kind == "footnote" else "[^endnote:1]"
    assert text.count(expected) == 3, "DOCX-only recovery must preserve every reference, including fields"
    assert extractor.build_definitions_block().count("One") == 1


@pytest.mark.parametrize("case", ["missing", "duplicate", "unbalanced", "text_range", "position_switch"])
def test_unproven_noteref_keeps_cached_text(case):
    q = f"{{{NS['w']}}}"
    doc = Document()
    anchor = doc.add_paragraph()
    start = etree.SubElement(anchor._p, q + "bookmarkStart", {q + "id": "1", q + "name": "NoteTarget"})
    run = etree.SubElement(anchor._p, q + "r")
    etree.SubElement(run, q + "footnoteReference", {q + "id": "1"})
    end = etree.SubElement(anchor._p, q + "bookmarkEnd", {q + "id": "1"})
    if case == "missing":
        start.set(q + "name", "AnotherTarget")
    elif case == "duplicate":
        anchor._p.append(deepcopy(start))
        anchor._p.append(deepcopy(end))
    elif case == "unbalanced":
        anchor._p.remove(end)
    elif case == "text_range":
        etree.SubElement(run, q + "t").text = "Unrelated text"
    paragraph = doc.add_paragraph("Before ")
    field = etree.SubElement(paragraph._p, q + "fldSimple")
    instruction = " NOTEREF NoteTarget \\h" + (" \\p" if case == "position_switch" else "")
    field.set(q + "instr", instruction)
    cached = etree.SubElement(field, q + "r")
    etree.SubElement(cached, q + "t").text = "Cached result"
    paragraph.add_run(" after.")
    extractor = NoteExtractor(doc)
    assert extractor.get_noteref_text(instruction) is None
    assert render_paragraph_runs(paragraph, extractor, syntax_config=DocxMarkdownSyntaxConfig()) == (
        "Before Cached result after."
    )


@pytest.mark.parametrize("simple", [False, True])
@pytest.mark.parametrize("hidden", [False, True])
def test_note_field_visibility_matches_ordinary_runs(simple, hidden):
    q = f"{{{NS['w']}}}"
    doc = Document()
    anchor = doc.add_paragraph()
    etree.SubElement(anchor._p, q + "bookmarkStart", {q + "id": "1", q + "name": "NoteTarget"})
    etree.SubElement(anchor.add_run()._r, q + "footnoteReference", {q + "id": "1"})
    etree.SubElement(anchor._p, q + "bookmarkEnd", {q + "id": "1"})
    paragraph = doc.add_paragraph("Before ")
    instruction = " NOTEREF NoteTarget \\h \\f "
    parent = etree.SubElement(paragraph._p, q + "fldSimple", {q + "instr": instruction}) if simple else paragraph._p
    values = (
        [("t", "Cached result")]
        if simple
        else [
            ("fldChar", "begin"),
            ("instrText", instruction),
            ("fldChar", "separate"),
            ("t", "Cached result"),
            ("fldChar", "end"),
        ]
    )
    for tag, value in values:
        run = etree.SubElement(parent, q + "r")
        props = etree.SubElement(run, q + "rPr")
        etree.SubElement(props, q + "vanish", {q + "val": "true" if hidden else "false"})
        payload = etree.SubElement(run, q + tag)
        if tag == "fldChar":
            payload.set(q + "fldCharType", value)
        else:
            payload.text = value
    paragraph.add_run(" after.")
    extractor = NoteExtractor(doc)
    expected = "Before  after." if hidden else "Before [^1] after."
    assert render_paragraph_runs(paragraph, extractor, syntax_config=DocxMarkdownSyntaxConfig()) == expected


@pytest.mark.parametrize(
    "case", ["missing_end", "position_switch", "mixed_end", "nested", "proof_type", "proof_payload", "proof_attribute"]
)
def test_malformed_complex_note_fields_preserve_visible_cached_text(case):
    q = f"{{{NS['w']}}}"
    doc = Document()
    anchor = doc.add_paragraph()
    etree.SubElement(anchor._p, q + "bookmarkStart", {q + "id": "1", q + "name": "NoteTarget"})
    etree.SubElement(anchor.add_run()._r, q + "footnoteReference", {q + "id": "1"})
    etree.SubElement(anchor._p, q + "bookmarkEnd", {q + "id": "1"})
    paragraph = doc.add_paragraph("Before ")
    instruction = " NOTEREF NoteTarget \\h" + (" \\p" if case == "position_switch" else "")
    values = [("fldChar", "begin"), ("instrText", instruction), ("fldChar", "separate"), ("t", "Cached result")]
    if case == "nested":
        values.insert(2, ("fldChar", "begin"))
    if case != "missing_end":
        values.append(("fldChar", "end"))
    for tag, value in values:
        run = etree.SubElement(paragraph._p, q + "r")
        payload = etree.SubElement(run, q + tag)
        if tag == "fldChar":
            payload.set(q + "fldCharType", value)
        else:
            payload.text = value
    if case == "mixed_end":
        etree.SubElement(run, q + "t").text = " tail"
    if case.startswith("proof_"):
        marker = etree.Element(q + "proofErr", {q + "type": "spellEnd"})
        if case == "proof_type":
            marker.set(q + "type", "unknown")
        elif case == "proof_payload":
            etree.SubElement(marker, q + "t").text = "invalid child"
        else:
            marker.set(q + "unexpected", "1")
        run.addprevious(marker)
    paragraph.add_run(" after.")
    extractor = NoteExtractor(doc)
    expected = "Before Cached result" + (" tail" if case == "mixed_end" else "") + " after."
    assert render_paragraph_runs(paragraph, extractor, syntax_config=DocxMarkdownSyntaxConfig()) == expected
