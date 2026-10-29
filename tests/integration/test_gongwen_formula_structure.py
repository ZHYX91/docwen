"""Real Gongwen formula paragraphs retain their source structure and visibility."""

from __future__ import annotations

import hashlib
import re
from io import BytesIO
from pathlib import Path
from typing import Any
from urllib.parse import unquote
from xml.etree import ElementTree as ET
from zipfile import ZipFile

import pytest
from docx import Document
from docx.oxml import parse_xml
from markdown_it import MarkdownIt
from PIL import Image

from docwen_bundle.config_port import ConfigPortAdapter
from docwen_bundle.runtime_factory import create_runtime_port
from docwen_core.paths import filesystem_path
from docwen_runtime.config.loader import ConfigLoader
from tests.integration.test_formula_and_gongwen_source_content import _body, _bundle, _math, _r, _title, _xml

pytestmark = [pytest.mark.integration, pytest.mark.pr_gate]
_W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
_M = "http://schemas.openxmlformats.org/officeDocument/2006/math"
_NS = {"w": _W, "m": _M}


@pytest.fixture
def configured_gongwen_runtime(tmp_path: Path) -> tuple[Any, ConfigPortAdapter]:
    """Admit the shipped snapshot through the same port as GUI/CLI/Machine."""
    loader = ConfigLoader(
        base_dir=Path(__file__).resolve().parents[2] / "configs",
        user_dir=tmp_path / "user-config",
    )
    config_port = ConfigPortAdapter(loader)
    assert config_port.snapshot()["numbering"]["cleanup"]["rules"]
    return create_runtime_port(config_loader=loader), config_port


@pytest.fixture
def no_cleanup_gongwen_runtime(tmp_path: Path) -> tuple[Any, ConfigPortAdapter]:
    """An explicit empty request rule set cannot infer authored prefixes."""
    loader = ConfigLoader(
        base_dir=Path(__file__).resolve().parents[2] / "configs",
        user_dir=tmp_path / "no-cleanup-user-config",
        runtime_overrides={"numbering": {"cleanup": {"rules": [], "settings": {"order": []}}}},
    )
    config_port = ConfigPortAdapter(loader)
    assert config_port.snapshot()["numbering"]["cleanup"]["rules"] == []
    return create_runtime_port(config_loader=loader), config_port


def _numbering(document, *, kind: str = "decimal", start: int = 1, heading_style: str = "") -> str:
    abstract = parse_xml(
        f'<w:abstractNum xmlns:w="{_W}" w:abstractNumId="99100">'
        '<w:multiLevelType w:val="multilevel"/>'
        + "".join(
            f'<w:lvl w:ilvl="{level}"><w:start w:val="{start if level == 0 else 1}"/>'
            f'<w:numFmt w:val="{kind}"/>'
            + (f'<w:pStyle w:val="{heading_style}"/>' if heading_style and level == 0 else "")
            + '<w:suff w:val="space"/>'
            + f'<w:lvlText w:val="{("•" if kind == "bullet" else "%" + str(level + 1) + ".")}"/></w:lvl>'
            for level in (0, 1)
        )
        + "</w:abstractNum>"
    )
    root = document.part.numbering_part.element
    root.insert(0, abstract)
    root.append(parse_xml(f'<w:num xmlns:w="{_W}" w:numId="99101"><w:abstractNumId w:val="99100"/></w:num>'))
    return "99101"


def _bind(paragraph, num_id: str, level: int = 0) -> None:
    paragraph._p.get_or_add_pPr().append(
        parse_xml(f'<w:numPr xmlns:w="{_W}"><w:ilvl w:val="{level}"/><w:numId w:val="{num_id}"/></w:numPr>')
    )


def _source_math_is_one_x(path: Path) -> None:
    assert [n.text for n in _xml(path).findall(".//m:t", _NS)] == ["x"]


@pytest.mark.parametrize("remove_numbering", [False, True])
@pytest.mark.parametrize("native", [False, True], ids=["authored-prefix", "native-counter"])
def test_formula_heading_uses_the_same_prefix_policy_and_heading_structure(
    configured_gongwen_runtime: Any, tmp_path: Path, remove_numbering: bool, native: bool
) -> None:
    document = Document()
    _title(document)
    document.add_paragraph("开始正文。")
    paragraph = document.add_paragraph("金额为" if native else "一、金额为", style="Heading 1")
    if native:
        _bind(paragraph, _numbering(document, start=4, heading_style="Heading1"))
    paragraph._p.append(_math(_r("x")))
    paragraph.add_run("元")
    document.add_paragraph("结束正文。")
    source = tmp_path / "formula-heading.docx"
    document.save(source)
    _source_math_is_one_x(source)
    if native:
        with ZipFile(source) as package:
            numbering = ET.fromstring(package.read("word/numbering.xml"))
        level = numbering.find("w:abstractNum[@w:abstractNumId='99100']/w:lvl", _NS)
        assert level.find("w:start", _NS).get(f"{{{_W}}}val") == "4"
        assert level.find("w:pStyle", _NS).get(f"{{{_W}}}val") == "Heading1"
    output, _ = _bundle(
        configured_gongwen_runtime[0],
        source,
        tmp_path / "bundle",
        source_format="docx",
        capability="gongwen",
        config_port=configured_gongwen_runtime[1],
        options={"remove_numbering": remove_numbering},
    )
    _, body = _body(output.read_text(encoding="utf-8"))
    prefix = "" if remove_numbering else ("4. " if native else "一、")
    expected = "# " + prefix + "金额为$x$元"
    assert body.splitlines().count(expected) == 1
    tokens = MarkdownIt("commonmark").parse(body)
    headings = [tokens[index + 1].content for index, token in enumerate(tokens) if token.type == "heading_open"]
    assert headings == [prefix + "金额为$x$元"]
    assert body.count("$x$") == 1
    assert body.index("开始正文。") < body.index(expected) < body.index("结束正文。")


@pytest.mark.parametrize("remove_numbering", [False, True])
def test_mixed_formula_heading_projects_the_plain_text_boundary_once(
    configured_gongwen_runtime: Any, tmp_path: Path, remove_numbering: bool
) -> None:
    document = Document()
    _title(document)
    paragraph = document.add_paragraph()
    paragraph.add_run("一、金额为").bold = True
    paragraph._p.append(_math(_r("x")))
    paragraph.add_run("元。").bold = True
    paragraph.add_run("请核对正文。").bold = False
    document.add_paragraph("后续段落。")
    source = tmp_path / "mixed-formula.docx"
    document.save(source)
    _source_math_is_one_x(source)
    output, _ = _bundle(
        configured_gongwen_runtime[0],
        source,
        tmp_path / "bundle",
        source_format="docx",
        capability="gongwen",
        config_port=configured_gongwen_runtime[1],
        options={"remove_numbering": remove_numbering},
    )
    _, body = _body(output.read_text(encoding="utf-8"))
    expected = "# " + ("" if remove_numbering else "一、") + "金额为$x$元。"
    assert body.splitlines().count(expected) == 1
    assert body.splitlines().count("请核对正文。") == 1
    assert body.index(expected) < body.index("请核对正文。") < body.index("后续段落。")
    tokens = MarkdownIt("commonmark").parse(body)
    assert sum(t.type == "heading_open" for t in tokens) == 1
    assert body.count("$x$") == 1


@pytest.mark.parametrize("remove_numbering", [False, True])
@pytest.mark.parametrize("kind", ["decimal", "bullet"])
@pytest.mark.parametrize("level", [0, 1], ids=["top-level", "nested"])
def test_formula_list_item_keeps_markers_nesting_and_concrete_counter(
    round_trip_runtime: Any, tmp_path: Path, remove_numbering: bool, kind: str, level: int
) -> None:
    document = Document()
    _title(document)
    num_id = _numbering(document, kind=kind)
    if level:
        _bind(document.add_paragraph("父项。"), num_id)
    paragraph = document.add_paragraph("金额为")
    _bind(paragraph, num_id, level)
    paragraph._p.append(_math(_r("x")))
    paragraph.add_run("元。")
    _bind(document.add_paragraph("后项。"), num_id, level)
    source = tmp_path / "formula-list.docx"
    document.save(source)
    _source_math_is_one_x(source)
    output, _ = _bundle(
        round_trip_runtime,
        source,
        tmp_path / "bundle",
        source_format="docx",
        capability="gongwen",
        options={"remove_numbering": remove_numbering},
    )
    _, body = _body(output.read_text(encoding="utf-8"))
    indent = "    " * level
    first = "-" if kind == "bullet" else "1."
    second = "-" if kind == "bullet" else "2."
    assert body.splitlines().count(indent + first + " 金额为$x$元。") == 1
    assert body.splitlines().count(indent + second + " 后项。") == 1
    tokens = MarkdownIt("commonmark").parse(body)
    assert sum(t.type == "list_item_open" for t in tokens) == 2 + bool(level)
    assert not any(t.type == "code_block" for t in tokens)


@pytest.mark.parametrize("remove_numbering", [False, True])
def test_attachment_formula_obeys_prefix_policy_in_the_related_final_artifact(
    configured_gongwen_runtime: Any, tmp_path: Path, remove_numbering: bool
) -> None:
    document = Document()
    _title(document)
    document.add_paragraph("完整主文。")
    document.add_paragraph("附件：计算规则")
    document.add_paragraph("测试市人民政府")
    document.add_paragraph("2024年1月15日")
    paragraph = document.add_paragraph("一、金额为", style="Heading 1")
    paragraph._p.append(_math(_r("x")))
    paragraph.add_run("元。")
    document.add_paragraph("附件结束。")
    source = tmp_path / "formula-annex.docx"
    document.save(source)
    _source_math_is_one_x(source)
    target = tmp_path / "bundle"
    output, result = _bundle(
        configured_gongwen_runtime[0],
        source,
        target,
        source_format="docx",
        capability="gongwen",
        config_port=configured_gongwen_runtime[1],
        options={"remove_numbering": remove_numbering},
    )
    [relation] = [r for r in result.bundle.relations if r.type == "attachment_of"]
    artifact = next(a for a in result.bundle.artifacts if a.artifact_id == relation.source_artifact_id)
    annex = filesystem_path(target / artifact.locator).read_text(encoding="utf-8")
    expected = ("" if remove_numbering else "一、") + "金额为$x$元。"
    assert annex.splitlines().count(expected) == 1
    assert annex.splitlines().count("附件结束。") == 1
    assert annex.index(expected) < annex.index("附件结束。")
    assert "$x$" not in output.read_text(encoding="utf-8")


@pytest.mark.parametrize("hidden", [False, True])
def test_formula_walker_respects_direct_visibility_without_reviving_moved_text(
    round_trip_runtime: Any, tmp_path: Path, hidden: bool
) -> None:
    document = Document()
    _title(document)
    paragraph = document.add_paragraph("金额为")
    paragraph.add_run("SECRET").font.hidden = hidden
    paragraph._p.append(parse_xml(f'<w:moveFrom xmlns:w="{_W}" w:id="7"><w:r><w:t>MOVED</w:t></w:r></w:moveFrom>'))
    paragraph._p.append(_math(_r("x")))
    paragraph.add_run("元。")
    paragraph.add_run("KEEP").font.hidden = False
    source = tmp_path / "formula-visible.docx"
    document.save(source)
    _source_math_is_one_x(source)
    output, _ = _bundle(round_trip_runtime, source, tmp_path / "bundle", source_format="docx", capability="gongwen")
    _, body = _body(output.read_text(encoding="utf-8"))
    expected = "金额为" + ("" if hidden else "SECRET") + "$x$元。KEEP"
    assert body.splitlines().count(expected) == 1
    assert "MOVED" not in body
    assert body.count("SECRET") == int(not hidden)


def test_formula_paragraph_keeps_the_declared_image_occurrence_order_and_hashes(
    round_trip_runtime: Any, tmp_path: Path
) -> None:
    document = Document()
    _title(document)
    paragraph = document.add_paragraph("金额为")
    hashes = []
    for color in ("red", "blue"):
        image = BytesIO()
        Image.new("RGB", (3, 2), color).save(image, format="PNG")
        hashes.append(hashlib.sha256(image.getvalue()).hexdigest())
        paragraph.add_run().add_picture(BytesIO(image.getvalue()))
        if color == "red":
            paragraph._p.append(_math(_r("x")))
    paragraph.add_run("元。")
    source = tmp_path / "formula-images.docx"
    document.save(source)
    _source_math_is_one_x(source)
    output, _ = _bundle(
        round_trip_runtime,
        source,
        tmp_path / "bundle",
        source_format="docx",
        capability="gongwen",
        options={"recognize_text": False, "preserve_resources": True},
    )
    _, body = _body(output.read_text(encoding="utf-8"))
    assert body.splitlines().count("金额为$x$元。") == 1
    targets = re.findall(r"!\[[^\]]*\]\(([^)]+)\)", body)
    assert [hashlib.sha256((output.parent / unquote(target)).read_bytes()).hexdigest() for target in targets] == hashes


@pytest.mark.parametrize("styled_heading", [False, True], ids=["body", "heading"])
def test_formula_prefix_without_request_cleanup_remains_authored_content(
    no_cleanup_gongwen_runtime: Any, tmp_path: Path, styled_heading: bool
) -> None:
    document = Document()
    _title(document)
    document.add_paragraph("开始正文。")
    paragraph = document.add_paragraph("一、金额为", style="Heading 1" if styled_heading else None)
    paragraph._p.append(_math(_r("x")))
    paragraph.add_run("元。")
    document.add_paragraph("结束正文。")
    source = tmp_path / "unconfigured-formula-prefix.docx"
    document.save(source)
    _source_math_is_one_x(source)
    output, _ = _bundle(
        no_cleanup_gongwen_runtime[0],
        source,
        tmp_path / "bundle",
        source_format="docx",
        capability="gongwen",
        config_port=no_cleanup_gongwen_runtime[1],
        options={"remove_numbering": True},
    )
    _, body = _body(output.read_text(encoding="utf-8"))
    expected = ("# " if styled_heading else "") + "一、金额为$x$元。"
    assert body.splitlines().count(expected) == 1
    tokens = MarkdownIt("commonmark").parse(body)
    assert sum(token.type == "heading_open" for token in tokens) == int(styled_heading)
    assert body.index("开始正文。") < body.index(expected) < body.index("结束正文。")
    assert body.count("$x$") == 1
