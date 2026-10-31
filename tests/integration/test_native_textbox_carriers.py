"""Complete native textbox envelopes, without claiming installed Word acceptance."""

from __future__ import annotations

from pathlib import Path
from typing import Any
from zipfile import ZipFile

import pytest
from docx import Document
from docx.oxml.ns import qn
from lxml import etree
from markdown_it import MarkdownIt

from tests.integration.test_native_docx_reading_contract import _publish

pytestmark = [pytest.mark.integration, pytest.mark.pr_gate]

_NS = {
    "w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main",
    "wp": "http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing",
    "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
    "wps": "http://schemas.microsoft.com/office/word/2010/wordprocessingShape",
    "v": "urn:schemas-microsoft-com:vml",
    "o": "urn:schemas-microsoft-com:office:office",
    "mc": "http://schemas.openxmlformats.org/markup-compatibility/2006",
    "future": "urn:docwen:test:unsupported-drawing-version",
}


def _element(parent: Any, tag_name: str, **attributes: str) -> Any:
    prefix, local = tag_name.split(":")
    return etree.SubElement(parent, f"{{{_NS[prefix]}}}{local}", attributes)


def _content(parent: Any, words: list[str]) -> None:
    content = _element(parent, "w:txbxContent")
    for word in words:
        paragraph = _element(content, "w:p")
        run = _element(paragraph, "w:r")
        _element(run, "w:t").text = word


def _drawingml(parent: Any, shape_id: int, words: list[str]) -> None:
    # CT_Inline and CT_WordprocessingShape child ordering follows the public
    # Open XML schemas; the test below independently inspects the envelope.
    drawing = _element(parent, "w:drawing")
    inline = _element(drawing, "wp:inline", distT="0", distB="0", distL="0", distR="0")
    _element(inline, "wp:extent", cx="1828800", cy="914400")
    _element(inline, "wp:docPr", id=str(shape_id), name=f"Text Box {shape_id}")
    _element(inline, "wp:cNvGraphicFramePr")
    graphic = _element(inline, "a:graphic")
    data = _element(graphic, "a:graphicData", uri=_NS["wps"])
    shape = _element(data, "wps:wsp")
    _element(shape, "wps:cNvSpPr", txBox="1")
    properties = _element(shape, "wps:spPr")
    transform = _element(properties, "a:xfrm")
    _element(transform, "a:off", x="0", y="0")
    _element(transform, "a:ext", cx="1828800", cy="914400")
    geometry = _element(properties, "a:prstGeom", prst="rect")
    _element(geometry, "a:avLst")
    _element(properties, "a:noFill")
    line = _element(properties, "a:ln")
    _element(line, "a:noFill")
    _content(_element(shape, "wps:txbx"), words)
    body_properties = _element(shape, "wps:bodyPr", wrap="square", anchor="t")
    _element(body_properties, "a:noAutofit")


def _vml(parent: Any, shape_id: int, words: list[str]) -> None:
    pict = _element(parent, "w:pict")
    shape_type = _element(
        pict, "v:shapetype", id=f"textbox-type-{shape_id}", coordsize="21600,21600", path="m,l,21600r21600,l21600,xe"
    )
    shape_type.set(f"{{{_NS['o']}}}spt", "202")
    _element(shape_type, "v:stroke", joinstyle="miter")
    path = _element(shape_type, "v:path", gradientshapeok="t")
    path.set(f"{{{_NS['o']}}}connecttype", "rect")
    shape = _element(
        pict,
        "v:shape",
        id=f"textbox-{shape_id}",
        type=f"#textbox-type-{shape_id}",
        style="width:144pt;height:72pt",
    )
    shape.set(f"{{{_NS['o']}}}spid", f"_x0000_s{1024 + shape_id}")
    _content(_element(shape, "v:textbox", inset="0,0,0,0"), words)


@pytest.mark.parametrize("mode", ["drawingml", "vml", "same-anchor", "separate-anchors", "mixed", "choice", "fallback"])
def test_complete_native_textbox_carriers_preserve_occurrence_and_document_order(
    tmp_path: Path, round_trip_runtime: Any, mode: str
) -> None:
    document = Document()
    document.add_paragraph("Before.")
    anchor = document.add_paragraph("Anchor.")
    run = anchor.add_run()._r
    expected = ["Before.", "Anchor."]
    if mode in {"drawingml", "vml"}:
        (_drawingml if mode == "drawingml" else _vml)(run, 1, ["A", "B", "A"])
        expected += ["A", "B", "A"]
    elif mode == "same-anchor":
        _drawingml(run, 1, ["A", "B", "A"])
        _vml(run, 2, ["A", "B", "A"])
        expected += ["A", "B", "A"] * 2
    elif mode == "separate-anchors":
        _drawingml(run, 1, ["A", "B", "A"])
        second = document.add_paragraph("Second anchor.")
        _vml(second.add_run()._r, 2, ["A", "B", "A"])
        expected += ["A", "B", "A", "Second anchor.", "A", "B", "A"]
    elif mode == "mixed":
        _vml(run, 1, ["First."])
        _drawingml(run, 2, ["Middle."])
        _vml(run, 3, ["Last."])
        expected += ["First.", "Middle.", "Last."]
    else:
        alternate = etree.SubElement(run, f"{{{_NS['mc']}}}AlternateContent", nsmap=_NS)
        choice = _element(alternate, "mc:Choice", Requires="wps" if mode == "choice" else "future")
        if mode == "choice":
            _drawingml(choice, 1, ["A", "B", "A"])
        else:
            _element(choice, "future:drawing")
        fallback = _element(alternate, "mc:Fallback")
        _vml(fallback, 2, ["A", "B", "A"])
        expected += ["A", "B", "A"]
    document.add_paragraph("After.")
    expected.append("After.")
    source = tmp_path / "complete-textboxes.docx"
    document.save(source)
    fresh = Document(source)
    assert [paragraph.text for paragraph in fresh.paragraphs] == [
        "Before.",
        "Anchor.",
        *(["Second anchor."] if mode == "separate-anchors" else []),
        "After.",
    ]
    with ZipFile(source) as package:
        root = etree.fromstring(package.read("word/document.xml"))
    shapes = root.findall(".//w:drawing/wp:inline/a:graphic/a:graphicData/wps:wsp", _NS)
    for shape in shapes:
        assert [etree.QName(child).localname for child in shape] == ["cNvSpPr", "spPr", "txbx", "bodyPr"]
        assert shape.find("wps:cNvSpPr", _NS).get("txBox") == "1"
        assert shape.find("wps:spPr/a:xfrm/a:ext", _NS).get("cx") == "1828800"
        assert shape.find("wps:spPr/a:prstGeom", _NS).get("prst") == "rect"
        assert shape.getparent().get("uri") == _NS["wps"]
        inline = shape.getparent().getparent().getparent()
        assert inline.find("wp:extent", _NS).get("cy") == "914400"
        assert inline.find("wp:docPr", _NS).get("id")
    vml_shapes = root.findall(".//w:pict/v:shape", _NS)
    for shape in vml_shapes:
        shape_type = shape.getparent().find("v:shapetype", _NS)
        assert shape_type is not None
        assert shape.get("type") == "#" + shape_type.get("id")
        assert shape_type.get(f"{{{_NS['o']}}}spt") == "202"
        assert shape.find("v:textbox/w:txbxContent", _NS) is not None
        assert shape.get("style") == "width:144pt;height:72pt"
    for box in root.iter(qn("w:txbxContent")):
        parent = box.getparent()
        assert parent.tag in {f"{{{_NS['wps']}}}txbx", f"{{{_NS['v']}}}textbox"}
        assert all(paragraph.tag == qn("w:p") for paragraph in box)
    assert len(shapes) + len(vml_shapes) == (
        2 if mode in {"same-anchor", "separate-anchors", "choice"} else 3 if mode == "mixed" else 1
    )
    body = _publish(round_trip_runtime, source, tmp_path / "published")
    tokens = MarkdownIt("commonmark").parse(body)
    assert [token.content for token in tokens if token.type == "inline"] == expected
