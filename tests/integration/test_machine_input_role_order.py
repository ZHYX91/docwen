"""Typed Machine inputs route by role all the way to published DOCX bytes."""

from __future__ import annotations

import hashlib
import json
from dataclasses import replace
from pathlib import Path
from typing import Any, Literal
from zipfile import ZipFile

import pytest
from docx import Document
from PIL import Image

from docwen_application.controller import ApplicationController
from docwen_application.conversion_contracts import (
    MARKDOWN_SOURCE_TO_DOCX_CAPABILITY_ID,
    MARKDOWN_TO_DOCX_CAPABILITY_ID,
    ConversionPlanRequest,
    ConversionServiceError,
    LocalInputHandle,
    StagingOutputTarget,
)
from docwen_application.conversion_service import ConversionService
from docwen_core.models import (
    NUMBERING_EXPORT_PLAN_MEDIA_TYPE,
    RESOLVED_DOCUMENT_MEDIA_TYPE,
    canonicalize_numbering_plan,
)
from docwen_runtime.output.artifact_bundle import ArtifactBundleCommitter

pytestmark = [pytest.mark.integration, pytest.mark.pr_gate]


def _handle(
    path: Path, *, role: str, media_type: str, kind: Literal["document", "resource"], logical_path: str
) -> LocalInputHandle:
    content = path.read_bytes()
    return LocalInputHandle(
        input_id=path.name,
        path=str(path),
        media_type=media_type,
        size_bytes=len(content),
        sha256=hashlib.sha256(content).hexdigest(),
        kind=kind,
        role=role,
        logical_path=logical_path,
    )


def _service(runtime: Any) -> ConversionService:
    controller = ApplicationController(runtime_port=runtime)
    controller.start()
    return ConversionService(controller, ArtifactBundleCommitter())


def _published_docx(service: ConversionService, request: ConversionPlanRequest) -> Path:
    plan = service.plan(request)
    outcome = service.execute_accepted(service.accept(plan.plan_id))
    assert outcome.state == "completed", outcome.error
    assert outcome.bundle is not None
    [artifact] = outcome.bundle.artifacts
    assert artifact.kind == "document"
    assert artifact.media_type == "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    [entry] = outcome.bundle.entries
    assert entry.artifact_id == artifact.artifact_id
    assert entry.role == "primary" and entry.preferred
    output = Path(request.output.staging_root) / artifact.locator
    content = output.read_bytes()
    assert len(content) == artifact.size_bytes
    assert hashlib.sha256(content).hexdigest() == artifact.sha256
    return output


def _source_inputs(tmp_path: Path) -> tuple[LocalInputHandle, ...]:
    source = tmp_path / "snapshot.md"
    source.write_text(
        "# Authored heading\n\nVisible body.\n\n![red](assets/red.png)\n\n![blue](assets/blue.png)\n",
        encoding="utf-8",
    )
    handles = [
        _handle(source, role="source", media_type="text/markdown", kind="document", logical_path="notes/authored.md")
    ]
    for color in ("red", "blue"):
        image = tmp_path / f"{color}.png"
        Image.new("RGB", (3, 2), color).save(image)
        handles.append(
            _handle(
                image,
                role="linked_resource",
                media_type="image/png",
                kind="resource",
                logical_path=f"notes/assets/{color}.png",
            )
        )
    return tuple(handles)


def _resolved_inputs(tmp_path: Path) -> tuple[LocalInputHandle, ...]:
    authored = "Resolved body.\n"
    plan = {"heading_definitions": [], "heading_instances": [], "targets": []}
    pointers = {
        "input_id": "opaque-input-id",
        "source_sha256": hashlib.sha256(authored.encode()).hexdigest(),
        "plan_sha256": hashlib.sha256(canonicalize_numbering_plan(plan)).hexdigest(),
    }
    document_path = tmp_path / "physical-snapshot.json"
    document_path.write_text(
        json.dumps(
            {
                "$schema": "urn:docwen:schema:resolved-document:v1",
                "schema": "docwen.resolved_document.v1",
                **pointers,
                "document": {
                    "authored_markdown": authored,
                    "targets": [],
                    "references": [],
                    "resource_occurrences": [],
                    "citations": [],
                    "resources": [],
                },
            }
        ),
        encoding="utf-8",
    )
    plan_path = tmp_path / "physical-numbering.json"
    plan_path.write_text(
        json.dumps(
            {
                "$schema": "urn:docwen:schema:numbering-export-plan:v1",
                "schema": "docwen.numbering_export_plan.v1",
                **pointers,
                "plan": plan,
            }
        ),
        encoding="utf-8",
    )
    return (
        _handle(
            document_path,
            role="neutral_document",
            media_type=RESOLVED_DOCUMENT_MEDIA_TYPE,
            kind="document",
            logical_path="notes/visible-name.md",
        ),
        _handle(
            plan_path,
            role="numbering_export_plan",
            media_type=NUMBERING_EXPORT_PLAN_MEDIA_TYPE,
            kind="resource",
            logical_path="numbering.json",
        ),
    )


@pytest.mark.parametrize("source_index", [0, 1, 2])
def test_source_native_input_can_follow_resources(round_trip_runtime: Any, tmp_path: Path, source_index: int) -> None:
    source, *resources = _source_inputs(tmp_path)
    handles = list(resources)
    handles.insert(source_index, source)
    originals = {handle.path: Path(handle.path).read_bytes() for handle in handles}
    staging = tmp_path / "published"
    staging.mkdir()
    output = _published_docx(
        _service(round_trip_runtime),
        ConversionPlanRequest(
            capability_id=MARKDOWN_SOURCE_TO_DOCX_CAPABILITY_ID,
            inputs=tuple(handles),
            output=StagingOutputTarget(str(staging)),
            options={"remove_numbering": False, "add_numbering": False},
        ),
    )
    document = Document(str(output))
    assert "Authored heading" in [paragraph.text for paragraph in document.paragraphs]
    assert "Visible body." in [paragraph.text for paragraph in document.paragraphs]
    assert len(document.inline_shapes) == 2
    # Package bytes are compared with independent input images, not renderer
    # counters or a second conversion using the same parsing assumptions.
    with ZipFile(output) as package:
        actual_images = [package.read(name) for name in package.namelist() if name.startswith("word/media/")]
    assert sorted(actual_images) == sorted(originals[resource.path] for resource in resources)
    assert all(Path(path).read_bytes() == content for path, content in originals.items())


@pytest.mark.parametrize("neutral_first", [True, False])
def test_resolved_document_can_follow_numbering_plan(
    round_trip_runtime: Any, tmp_path: Path, neutral_first: bool
) -> None:
    inputs = _resolved_inputs(tmp_path)
    originals = {handle.path: Path(handle.path).read_bytes() for handle in inputs}
    staging = tmp_path / "published"
    staging.mkdir()
    output = _published_docx(
        _service(round_trip_runtime),
        ConversionPlanRequest(
            capability_id=MARKDOWN_TO_DOCX_CAPABILITY_ID,
            inputs=inputs if neutral_first else tuple(reversed(inputs)),
            output=StagingOutputTarget(str(staging)),
            options={},
        ),
    )
    document = Document(str(output))
    assert "Resolved body." in [paragraph.text for paragraph in document.paragraphs]
    assert "opaque-input-id" not in "\n".join(paragraph.text for paragraph in document.paragraphs)
    assert output.name.startswith("visible-name")
    assert all(Path(path).read_bytes() == content for path, content in originals.items())


@pytest.mark.parametrize(
    ("invalid", "expected_code"),
    [
        ("missing_source", "input_slot_cardinality_mismatch"),
        ("duplicate_source", "input_slot_cardinality_mismatch"),
        ("unknown_role", "undeclared_input_role"),
        ("wrong_kind", "input_slot_kind_mismatch"),
    ],
)
def test_role_dispatch_does_not_relax_source_admission(
    round_trip_runtime: Any, tmp_path: Path, invalid: str, expected_code: str
) -> None:
    source, *resources = _source_inputs(tmp_path)
    inputs = [*resources, source]
    if invalid == "missing_source":
        inputs = resources
    elif invalid == "duplicate_source":
        inputs.append(replace(source, input_id="second-source", logical_path="notes/second.md"))
    elif invalid == "unknown_role":
        inputs[0] = replace(inputs[0], role="unknown")
    else:
        inputs[-1] = replace(source, kind="resource")
    staging = tmp_path / "published"
    staging.mkdir()
    request = ConversionPlanRequest(
        capability_id=MARKDOWN_SOURCE_TO_DOCX_CAPABILITY_ID,
        inputs=tuple(inputs),
        output=StagingOutputTarget(str(staging)),
        options={},
    )
    with pytest.raises(ConversionServiceError) as refused:
        _service(round_trip_runtime).plan(request)
    assert refused.value.code == expected_code
    assert list(staging.iterdir()) == []


@pytest.mark.parametrize(
    ("invalid", "expected_code"),
    [
        ("missing_document", "docwen.resolved_document.missing"),
        ("duplicate_document", "docwen.resolved_document.invalid"),
        ("missing_plan", "docwen.numbering_export_plan.missing"),
    ],
)
def test_role_dispatch_keeps_resolved_pair_cardinality(
    round_trip_runtime: Any, tmp_path: Path, invalid: str, expected_code: str
) -> None:
    document, plan = _resolved_inputs(tmp_path)
    inputs = [plan, document]
    if invalid == "missing_document":
        inputs = [plan]
    elif invalid == "missing_plan":
        inputs = [document]
    else:
        inputs.append(replace(document, input_id="second-document", logical_path="notes/second.md"))
    staging = tmp_path / "published"
    staging.mkdir()
    request = ConversionPlanRequest(
        capability_id=MARKDOWN_TO_DOCX_CAPABILITY_ID,
        inputs=tuple(inputs),
        output=StagingOutputTarget(str(staging)),
        options={},
    )
    with pytest.raises(ConversionServiceError) as refused:
        _service(round_trip_runtime).plan(request)
    assert refused.value.code == expected_code
    assert list(staging.iterdir()) == []
