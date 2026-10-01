"""Derived Office PNGs in an explicit model through actual final converters.

These are Git regression fixtures, not native rich binding or host acceptance.
"""

import csv
import hashlib
import json
import zipfile
from collections import Counter
from dataclasses import replace
from io import BytesIO
from pathlib import Path
from xml.etree import ElementTree as ET

import pytest
from openpyxl import load_workbook
from PIL import Image

from docwen_core.models.clipboard_document import (
    ClipboardDocument,
    ClipboardImageRef,
    ClipboardParagraph,
    ClipboardTable,
    ClipboardTableCell,
    ClipboardText,
    clipboard_document_to_bytes,
)
from docwen_core.models.request import OutputPolicy
from docwen_gui.clipboard_inputs import ClipboardInputStore
from docwen_gui.clipboard_office_provider import parse_word_embed_source, parse_wps_writer
from docwen_plugin_markdown.plugin import MarkdownPlugin
from tests.integration.test_structured_clipboard_runtime import _template_id
from tests.integration.test_structured_resource_group_runtime import _builder, _controller, _source_ref

pytestmark = [pytest.mark.integration, pytest.mark.pr_gate, pytest.mark.release_gate]
_SAMPLES = Path(__file__).resolve().parents[1] / "fixtures/files/clipboard-office"
EXPECTED_SHAS = {
    "66d68226db981e8856cd373abedf046917cd5d801498754efdf1b4ecf70d6633",
    "020d53090eb9e490942c26afbfee75b8b568059ad5955aad1f7616dfaf28e8c3",
}
EXPECTED_EXTENTS = Counter({(914400, 609600): 2, (457200, 304800): 1})


def _controlled_model(provider):
    """Hand-authored positions; identities and extents come from the derived provider parser."""
    identities = dict(provider.identifier_to_resource_id)
    paragraphs = []
    for index, occurrence in enumerate(provider.occurrences):
        paragraphs.append(
            ClipboardParagraph(
                (
                    ClipboardText(f"probe-{index}-before"),
                    ClipboardImageRef(
                        identities[occurrence.identifier],
                        f"probe-{index}-image",
                        "",
                        occurrence.extent_cx_emu,
                        occurrence.extent_cy_emu,
                    ),
                    ClipboardText(f"probe-{index}-after"),
                )
            )
        )
    nested = ClipboardTable(1, 1, (ClipboardTableCell(0, 0, 1, 1, (paragraphs[1],)),))
    outer = ClipboardTable(1, 1, (ClipboardTableCell(0, 0, 1, 1, (nested,)),))
    return ClipboardDocument((paragraphs[0], outer, paragraphs[2]), provider.resources)


def _assert_pngs(payloads):
    assert {hashlib.sha256(payload).hexdigest() for payload in payloads} == EXPECTED_SHAS
    for payload in payloads:
        with Image.open(BytesIO(payload)) as image:
            image.load()
            assert image.size == (96, 64)
            assert image.format == "PNG"


@pytest.mark.parametrize("producer", ["word", "wps"])
@pytest.mark.parametrize("target", ["md", "docx", "xlsx", "csv"])
def test_derived_provider_resources_through_final_artifacts(tmp_path, producer, target):
    def sample(name):
        data = (_SAMPLES / name).read_bytes()
        facts = json.loads((_SAMPLES / "provenance.json").read_text(encoding="utf-8"))["files"][name]
        assert len(data) == facts["bytes"]
        assert hashlib.sha256(data).hexdigest() == facts["sha256"]
        return data

    if producer == "word":
        provider = parse_word_embed_source(sample("word-derived.ole"))
    else:
        provider = parse_wps_writer(sample("wps-writer-derived.zip"), sample("wps-images.bin"))
    model = _controlled_model(provider)
    store = ClipboardInputStore(tmp_path / "managed")
    try:
        bundle = store.create_bundle(
            clipboard_document_to_bytes(model),
            display_name_template="Controlled output {index}.dwclip",
            preview="controlled",
            resources=provider.resource_bytes,
        )
        store.sync_visible([bundle.main.path])
        request, _ = _builder(store, [_source_ref(bundle)]).single(
            file_path=bundle.main.path,
            target_format=target,
            action_name="",
            options={},
            output_policy=OutputPolicy(output_dir=str(tmp_path / "published")),
        )
        options = {}
        if target == "md":
            options = {
                "image_mode": "file",
                "image_link_style": "markdown_embed",
                "markdown_extensions": {"output": {"structural_tables": True}},
            }
        if target == "docx":
            options["template_name"] = _template_id("docx", "English General Template.docx")
        elif target == "xlsx":
            options["template_name"] = _template_id("xlsx", "English Sample Sheet Template.xlsx")
        result = _controller(tmp_path, MarkdownPlugin()).execute_single(replace(request, options=options))
        assert result.success, result.error
        assert "CLIPBOARD-IMAGE-RESOURCE-UNAVAILABLE" not in {item.code for item in result.diagnostics}
        primary = Path(next(item for item in result.artifacts if item.is_primary).staging_path)
        assert primary.is_file()
        if target in {"docx", "xlsx"}:
            prefix = "word/media/" if target == "docx" else "xl/media/"
            with zipfile.ZipFile(primary) as package:
                media_names = [
                    name
                    for name in package.namelist()
                    if name.startswith(prefix) and package.read(name).startswith(b"\x89PNG\r\n\x1a\n")
                ]
                media = [package.read(name) for name in media_names]
                _assert_pngs(media)
                assert len(media) == 2, "Repeated occurrences must reuse each unique embedded PNG"
                assert all(name.endswith(".png") for name in media_names), media_names
                if target == "docx":
                    root = ET.fromstring(package.read("word/document.xml"))
                    ns = {
                        "wp": "http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing",
                        "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
                        "w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main",
                    }
                    extents = [(int(node.get("cx")), int(node.get("cy"))) for node in root.findall(".//wp:extent", ns)]
                    assert Counter(extents) == EXPECTED_EXTENTS
                    assert len(root.findall(".//a:blip", ns)) == 3
                    assert len(root.findall(".//w:tbl/w:tr/w:tc/w:tbl//a:blip", ns)) == 1
                    xml = package.read("word/document.xml").decode()
                    assert xml.index("probe-0-before") < xml.index("probe-1-before") < xml.index("probe-2-before")
                else:
                    ns = {"xdr": "http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing"}
                    extents = []
                    for name in package.namelist():
                        if name.startswith("xl/drawings/drawing") and name.endswith(".xml"):
                            root = ET.fromstring(package.read(name))
                            extents.extend(
                                (int(node.get("cx")), int(node.get("cy"))) for node in root.findall(".//xdr:ext", ns)
                            )
                    assert Counter(extents) == EXPECTED_EXTENTS
            if target == "xlsx":
                workbook = load_workbook(primary)
                try:
                    rows = list(workbook["Image Semantics"].values)
                    assert len(rows) == 4
                    assert [row[1] for row in rows[1:]] == ["bound"] * 3
                    assert rows[1][2] == rows[3][2] != rows[2][2]
                finally:
                    workbook.close()
        else:
            images = [
                Path(item.staging_path).read_bytes() for item in result.artifacts if item.media_type == "image/png"
            ]
            _assert_pngs(images)
            assert len(images) == 2
            text = primary.read_text(encoding="utf-8-sig")
            if target == "md":
                assert text.count("![probe-") == 3
                assert text.index("probe-0-before") < text.index("probe-1-before") < text.index("probe-2-before")
            else:
                assert "CLIPBOARD-CSV-IMAGE-PROJECTION" in {item.code for item in result.diagnostics}
                semantics = next(
                    Path(item.staging_path)
                    for item in result.artifacts
                    if "image-semantics" in item.suggested_name.lower()
                )
                rows = list(csv.reader(semantics.read_text(encoding="utf-8-sig").splitlines()))
                assert len(rows) == 4 and [row[1] for row in rows[1:]] == ["bound"] * 3
                assert Counter((int(row[-2]), int(row[-1])) for row in rows[1:]) == EXPECTED_EXTENTS
    finally:
        store.close()
