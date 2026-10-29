"""Native cell literals survive both published Markdown and a fresh DOCX package."""

from __future__ import annotations

import hashlib
from html.parser import HTMLParser
from pathlib import Path
from typing import Any
from xml.etree import ElementTree as ET
from zipfile import ZipFile

import mistune
import openpyxl
import pytest
from docx import Document

from docwen_application.bundle_mapping import build_bundle_draft
from docwen_core.models.file_ref import FileRef
from docwen_core.models.request import ConversionRequest, OutputPolicy
from docwen_runtime.output.artifact_bundle import ArtifactBundleCommitter

pytestmark = [pytest.mark.integration, pytest.mark.pr_gate]

_W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
_NS = {"w": _W}
_MEDIA = {
    "md": "text/markdown",
    "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
}


class _HtmlTableReader(HTMLParser):
    """Consume independently rendered HTML once, including literal entities."""

    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self.tables: list[list[list[str]]] = []
        self.row: list[str] | None = None
        self.cell: str | None = None

    def handle_starttag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        if tag == "table":
            self.tables.append([])
        elif tag == "tr":
            self.row = []
        elif tag in {"td", "th"}:
            self.cell = ""
        elif tag == "br" and self.cell is not None:
            self.cell += "\n"

    def handle_endtag(self, tag: str) -> None:
        if tag in {"td", "th"}:
            assert self.row is not None and self.cell is not None
            self.row.append(self.cell)
            self.cell = None
        elif tag == "tr":
            assert self.row is not None
            self.tables[-1].append(self.row)
            self.row = None

    def handle_data(self, data: str) -> None:
        if self.cell is not None:
            self.cell += data


def _publish(runtime: Any, source: Path, output: Path, target: str, *, structural: bool) -> Path:
    authored = source.read_bytes()
    is_sheet = source.suffix == ".xlsx"
    result = runtime.execute(
        ConversionRequest(
            request_id="literal-table-" + target,
            input_refs=[
                FileRef(
                    path=str(source),
                    format="xlsx" if is_sheet else "markdown",
                    category="spreadsheet" if is_sheet else "markdown",
                    input_role="source",
                    size_bytes=len(authored),
                )
            ],
            target_format=target,
            output_policy=OutputPolicy(output_dir=str(output)),
            options={
                **(
                    {"to_md_keep_images": False, "to_md_enable_ocr": False} if is_sheet else {"remove_numbering": False}
                ),
                "markdown_extensions": {"output" if is_sheet else "input": {"structural_tables": structural}},
            },
        )
    )
    assert result.success, result.error
    assert source.read_bytes() == authored
    draft = build_bundle_draft(profile="single_document", output_media_type=_MEDIA[target], artifacts=result.artifacts)
    bundle = ArtifactBundleCommitter().commit(task_id="literal-table-" + target, staging_root=str(output), draft=draft)
    for artifact in bundle.artifacts:
        content = (output / artifact.locator).read_bytes()
        assert artifact.size_bytes == len(content)
        assert artifact.sha256 == hashlib.sha256(content).hexdigest()
    preferred = [entry for entry in bundle.entries if entry.preferred]
    assert len(preferred) == 1
    artifact = next(item for item in bundle.artifacts if item.artifact_id == preferred[0].artifact_id)
    return output / artifact.locator


def _xml_table(path: Path) -> ET.Element:
    with ZipFile(path) as package:
        root = ET.fromstring(package.read("word/document.xml"))
    tables = root.findall(".//w:tbl", _NS)
    assert len(tables) == 1
    return tables[0]


def _cell_text(cell: ET.Element) -> str:
    parts: list[str] = []
    for node in cell.iter():
        if node.tag == f"{{{_W}}}t":
            parts.append(node.text or "")
        elif node.tag in {f"{{{_W}}}br", f"{{{_W}}}cr"}:
            parts.append("\n")
    return "".join(parts)


@pytest.mark.parametrize("structural", [False, True], ids=["ordinary", "structural-enabled"])
def test_native_literals_reach_final_docx_as_authored_characters(
    tmp_path: Path, round_trip_runtime: Any, structural: bool
) -> None:
    expected = [
        ["Kind", "Value"],
        ["lt", "<"],
        ["gt", ">"],
        ["amp", "&"],
        ["entity", "&lt;"],
        ["caret", "^"],
        ["pipe", r"a\|b"],
        ["line", "top\nbottom"],
        ["html", "<br>"],
        ["syntax", "*bold* [x](y)"],
    ]
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    assert sheet is not None
    for row in expected:
        sheet.append(row)
    source = tmp_path / "literals.xlsx"
    workbook.save(source)
    workbook.close()
    reopened = openpyxl.load_workbook(source)
    try:
        assert list(reopened.worksheets[0].values) == [tuple(row) for row in expected]
        assert all(cell.data_type == "s" for row in reopened.worksheets[0] for cell in row)
    finally:
        reopened.close()

    markdown = _publish(round_trip_runtime, source, tmp_path / "markdown", "md", structural=structural)
    consumer = _HtmlTableReader()
    html = mistune.create_markdown(renderer="html", escape=False, plugins=["table"])(
        markdown.read_text(encoding="utf-8-sig")
    )
    assert isinstance(html, str)
    consumer.feed(html)
    consumer.close()
    assert consumer.tables == [expected]

    document = _publish(round_trip_runtime, markdown, tmp_path / "document", "docx", structural=structural)
    table = _xml_table(document)
    actual = [[_cell_text(cell) for cell in row.findall("w:tc", _NS)] for row in table.findall("w:tr", _NS)]
    assert actual == expected
    assert table.find(".//w:gridSpan", _NS) is None
    assert table.find(".//w:vMerge", _NS) is None
    fresh = Document(document)
    assert [[cell.text for cell in row.cells] for row in fresh.tables[0].rows] == expected


@pytest.mark.parametrize("merge", ["A1:B1", "A1:A2"])
def test_literal_encoding_keeps_real_structural_merge_geometry(
    tmp_path: Path, round_trip_runtime: Any, merge: str
) -> None:
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet.append(["Merged", "X"])
    sheet.append(["Y", "&lt;"])
    sheet.merge_cells(merge)
    source = tmp_path / "merged.xlsx"
    workbook.save(source)
    workbook.close()
    reopened = openpyxl.load_workbook(source)
    try:
        assert list(map(str, reopened.worksheets[0].merged_cells.ranges)) == [merge]
        assert reopened.worksheets[0]["B2"].value == "&lt;"
    finally:
        reopened.close()
    markdown = _publish(round_trip_runtime, source, tmp_path / "markdown", "md", structural=True)
    document = _publish(round_trip_runtime, markdown, tmp_path / "document", "docx", structural=True)
    table = _xml_table(document)
    rows = table.findall("w:tr", _NS)
    assert len(rows) == 2
    if merge == "A1:B1":
        assert [[_cell_text(cell) for cell in row.findall("w:tc", _NS)] for row in rows] == [["Merged"], ["Y", "&lt;"]]
        span = rows[0].find("w:tc/w:tcPr/w:gridSpan", _NS)
        assert span is not None and span.get(f"{{{_W}}}val") == "2"
        assert table.find(".//w:vMerge", _NS) is None
    else:
        assert [[_cell_text(cell) for cell in row.findall("w:tc", _NS)] for row in rows] == [
            ["Merged", "X"],
            ["", "&lt;"],
        ]
        first = rows[0].find("w:tc/w:tcPr/w:vMerge", _NS)
        second = rows[1].find("w:tc/w:tcPr/w:vMerge", _NS)
        assert first is not None and first.get(f"{{{_W}}}val") == "restart"
        assert second is not None and second.get(f"{{{_W}}}val") in {None, "continue"}
        assert table.find(".//w:gridSpan", _NS) is None
