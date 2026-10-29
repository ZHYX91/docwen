"""Authenticated references and native breaks keep one source paragraph owner."""

from __future__ import annotations

import hashlib
from copy import deepcopy
from io import BytesIO
from pathlib import Path
from typing import Any
from zipfile import ZIP_DEFLATED, ZipFile

import pytest
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from lxml import etree

from docwen_application.bundle_mapping import build_bundle_draft
from docwen_core.docx_parsing.format_features import DocxMarkdownSyntaxConfig
from docwen_core.docx_semantics_v3 import SOFT_REFERENCE_TAG_PREFIX, DocxSemanticsV3Recovery, DocxSemanticsV3Session
from docwen_core.markdown_extensions import MarkdownExtensions
from docwen_plugin_document.shared.markdown_runs import render_paragraph_runs_split_on_page_breaks
from docwen_plugin_markdown.mistune_extensions import parse_markdown_text
from docwen_plugin_markdown.renderer import MdToDocxRenderer
from docwen_plugin_markdown.runtime_semantics_v3 import apply_runtime_semantics_v3, prepare_runtime_semantics_v3
from docwen_runtime.output.artifact_bundle import ArtifactBundleCommitter
from tests.integration._round_trip_helper import _run
from tests.integration.test_nested_quote_anchor_packages import _identity

pytestmark = [pytest.mark.integration, pytest.mark.pr_gate]


def _variant(original: Path, target: Path, *, break_tag: str, ending: str) -> None:
    """Keep every writer-produced package member except the native break variant."""
    with ZipFile(original) as source, ZipFile(target, "w", ZIP_DEFLATED) as output:
        root = etree.fromstring(source.read("word/document.xml"))
        [paragraph] = [item for item in root.iter(qn("w:p")) if "Before " in "".join(item.itertext())]
        [native_break] = list(paragraph.iter(qn("w:br")))
        native_break.tag = qn("w:" + break_tag)
        if ending == "page":
            run = OxmlElement("w:r")
            page_break = OxmlElement("w:br")
            page_break.set(qn("w:type"), "page")
            run.append(page_break)
            paragraph.append(run)
        elif ending == "section":
            section = root.find(f"{qn('w:body')}/{qn('w:sectPr')}")
            assert section is not None
            properties = paragraph.find(qn("w:pPr"))
            assert properties is not None
            properties.append(deepcopy(section))
        for item in source.infolist():
            output.writestr(
                item,
                etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
                if item.filename == "word/document.xml"
                else source.read(item.filename),
            )
    with ZipFile(original) as source, ZipFile(target) as output:
        assert source.namelist() == output.namelist()
        assert all(source.read(name) == output.read(name) for name in source.namelist() if name != "word/document.xml")


@pytest.mark.parametrize("captions", [False, True], ids=["captions-off", "captions-on"])
@pytest.mark.parametrize("break_tag", ["br", "cr"])
@pytest.mark.parametrize("ending", ["none", "page", "section"])
@pytest.mark.parametrize("literal_code", [False, True], ids=["original", "literal-code"])
def test_semantic_quote_break_retains_reference_content_and_valid_owner(
    tmp_path: Path, round_trip_runtime: Any, captions: bool, break_tag: str, ending: str, literal_code: bool
) -> None:
    literal = " and `literal <br>`" if literal_code else ""
    source = f"# Target\n\n> Before @[[#Target]]<br>After{literal} ^quote-id\n\n^outer\n"
    dialect = MarkdownExtensions(captions_references=True)
    plan = prepare_runtime_semantics_v3(
        source, input_id="source.md", extensions=dialect, consumer_profile="number_suite_direct"
    )
    assert not plan.analysis.has_errors
    source_path = tmp_path / "source.md"
    source_path.write_bytes(source.encode())
    document = Document()
    session = DocxSemanticsV3Session(document, source_sha256=plan.source_sha256, standalone_caption_authority=True)
    parsed = parse_markdown_text(plan.shielded_source, extensions=dialect, auto_link_bare_url=False)
    assert isinstance(parsed, list)
    MdToDocxRenderer(doc=document, semantic_v3_session=session).render(apply_runtime_semantics_v3(parsed, plan))
    session.finalize_document()
    buffer = BytesIO()
    document.save(buffer)
    original = tmp_path / "writer.docx"
    original.write_bytes(buffer.getvalue())
    session.write_package(original)
    original_digest = hashlib.sha256(original.read_bytes()).hexdigest()
    candidate = tmp_path / "native-break.docx"
    _variant(original, candidate, break_tag=break_tag, ending=ending)
    assert hashlib.sha256(original.read_bytes()).hexdigest() == original_digest
    assert source_path.read_bytes() == source.encode()

    with ZipFile(candidate) as package:
        root = etree.fromstring(package.read("word/document.xml"))
    tags = {
        tag.get(qn("w:val")): element
        for element in root.iter(qn("w:sdt"))
        if (tag := element.find(f"{qn('w:sdtPr')}/{qn('w:tag')}")) is not None
    }
    [soft] = [element for tag, element in tags.items() if tag and tag.startswith(SOFT_REFERENCE_TAG_PREFIX)]
    assert [node.text for node in soft.iter(qn("w:t"))] == ["1"]
    assert not list(root.iter(qn("w:bookmarkStart")))
    assert not list(root.iter(qn("w:instrText")))
    inner = tags[_identity("quote-id", "paragraph")["tag"]]
    outer = tags[_identity("outer", "block_quote")["tag"]]
    assert inner in outer.iter()
    [paragraph] = inner.findall(f"{qn('w:sdtContent')}/{qn('w:p')}")
    assert "".join(node.text or "" for node in paragraph.iter(qn("w:t"))) == (
        "Before 1After" + (" and literal <br>" if literal_code else "")
    )
    assert len(list(paragraph.iter(qn("w:" + break_tag)))) == 1 + int(ending == "page" and break_tag == "br")
    assert len([node for node in paragraph.iter(qn("w:br")) if node.get(qn("w:type")) == "page"]) == int(
        ending == "page"
    )
    assert (paragraph.find(f"{qn('w:pPr')}/{qn('w:sectPr')}") is not None) == (ending == "section")

    fresh = Document(BytesIO(candidate.read_bytes()))
    authenticated = DocxSemanticsV3Recovery.load(candidate, fresh)
    assert len(authenticated.soft_reference_identities) == 1
    assert len(authenticated.reference_occurrence_identities) == 0
    result = _run(
        round_trip_runtime,
        "semantic-quote-reader",
        candidate,
        source_format="docx",
        target_format="md",
        output_dir=tmp_path / "reverse",
        options={
            "remove_numbering": False,
            "markdown_extensions": {"output": MarkdownExtensions(captions_references=captions).to_dict()},
            "page_break_separator": "---",
            "section_break_separator": "***",
        },
    )
    assert result.success, result.error
    draft = build_bundle_draft(
        profile="document_with_resources", output_media_type="text/markdown", artifacts=result.artifacts
    )
    bundle = ArtifactBundleCommitter().commit(
        task_id="quote-reader", staging_root=str(tmp_path / "reverse"), draft=draft
    )
    [preferred] = [entry.artifact_id for entry in bundle.entries if entry.preferred]
    [artifact] = [item for item in bundle.artifacts if item.artifact_id == preferred]
    data = (tmp_path / "reverse" / artifact.locator).read_bytes()
    assert artifact.sha256 == hashlib.sha256(data).hexdigest()
    returned = data.decode("utf-8-sig")
    reference = "@[[#Target]]" if captions else "1"
    expected = f"> Before {reference}<br>After{literal} ^quote-id"
    assert expected in returned
    assert (
        returned.count("Before")
        == returned.count("After")
        == returned.count("^quote-id")
        == returned.count("^outer")
        == 1
    )
    assert returned.count("@[[#Target]]") == int(captions)
    assert returned.count("`literal <br>`") == int(literal_code)
    if ending != "none":
        separator = "---" if ending == "page" else "***"
        body = returned.split("\n---\n", 1)[1] if returned.startswith("---\n") else returned
        assert body.index("^quote-id") < body.index(f"\n{separator}\n")
    replanned = prepare_runtime_semantics_v3(
        returned, input_id="recovered.md", extensions=dialect, consumer_profile="number_suite_direct"
    )
    assert not replanned.analysis.has_errors
    assert [(item["id"], item["block_kind"]) for item in replanned.analysis.projection["anchors"]] == [
        ("quote-id", "paragraph"),
        ("outer", "block_quote"),
    ]
    assert all(target.get("id") is None for target in replanned.analysis.projection["targets"])


@pytest.mark.parametrize("break_tag", ["br", "cr"])
@pytest.mark.parametrize("projection", ["\n", "<br>"])
def test_shared_break_projection_does_not_rewrite_literal_text_or_code(break_tag: str, projection: str) -> None:
    document = Document()
    paragraph = document.add_paragraph()
    run = paragraph.add_run()
    text = OxmlElement("w:t")
    text.text = "literal <br>\nuntouched"
    run._r.append(text)
    run._r.append(OxmlElement("w:" + break_tag))
    code_run = paragraph.add_run("code <br>")
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), "D9D9D9")
    code_run._r.get_or_add_rPr().append(shading)
    assert render_paragraph_runs_split_on_page_breaks(
        paragraph, syntax_config=DocxMarkdownSyntaxConfig(), line_break_text=projection
    ) == [f"literal <br>\nuntouched{projection}`code <br>`"]
