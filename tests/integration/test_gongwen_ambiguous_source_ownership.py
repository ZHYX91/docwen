"""Generated native source cannot lose unrepresented numeric header prefixes."""

from __future__ import annotations

import hashlib
from io import BytesIO
from pathlib import Path
from typing import Any
from urllib.parse import unquote, urlsplit

import pytest
from docx import Document
from markdown_it import MarkdownIt
from PIL import Image

from docwen_bundle.config_port import ConfigPortAdapter
from docwen_core.paths import filesystem_path
from tests.integration.test_formula_and_gongwen_source_content import _body, _bundle, _math, _r, _title, _xml
from tests.integration.test_gongwen_request_source_ownership import configured_ports as configured_ports

pytestmark = [pytest.mark.integration, pytest.mark.pr_gate]
_W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
_M = "http://schemas.openxmlformats.org/officeDocument/2006/math"
_NS = {"w": _W, "m": _M}


@pytest.mark.parametrize(
    "authored",
    [
        "001 002 国发〔2024〕1号",
        "001\u3000002\u3000国发〔2024〕1号",
        "001\t002 国发〔2024〕1号",
        "001 002 003 国发〔2024〕1号",
        "001 002 国发〔2024〕号",
        "001 002 国发2024-1号",
        "001 002国发〔2024〕1号",
    ],
)
@pytest.mark.parametrize("rich", [False, True], ids=["plain", "formula-image"])
def test_ambiguous_numeric_header_keeps_full_native_source_in_final_bundle(
    configured_ports: tuple[Any, ConfigPortAdapter], tmp_path: Path, authored: str, rich: bool
) -> None:
    runtime, config_port = configured_ports
    rules = config_port.snapshot()["numbering"]["cleanup"]["rules"]
    arabic_space = next(rule for rule in rules if rule["id"] == "arabic_space")
    assert arabic_space["enabled"] is True and arabic_space["level"] == 3
    document = Document()
    _title(document)
    document.add_paragraph("开始正文。")
    paragraph = document.add_paragraph(authored)
    picture_bytes = b""
    if rich:
        paragraph._p.append(_math(_r("x")))
        paragraph.add_run("附带说明。")
        picture = BytesIO()
        Image.new("RGB", (3, 2), "blue").save(picture, format="PNG")
        picture_bytes = picture.getvalue()
        paragraph.add_run().add_picture(BytesIO(picture_bytes))
    document.add_paragraph("尾部正文。")
    source = tmp_path / "ambiguous-header.docx"
    document.save(source)
    reopened = Document(source)
    assert reopened.paragraphs[3].text == authored + ("附带说明。" if rich else "")
    native = _xml(source).findall("w:body/w:p", _NS)[3]
    source_atoms = [
        "\t" if node.tag == f"{{{_W}}}tab" else (node.text or "")
        for node in native.iter()
        if node.tag in {f"{{{_W}}}t", f"{{{_M}}}t", f"{{{_W}}}tab"}
    ]
    assert "".join(source_atoms) == authored + ("x附带说明。" if rich else "")
    assert [node.text for node in native.findall(".//m:t", _NS)] == (["x"] if rich else [])
    target = tmp_path / "bundle"
    output, result = _bundle(
        runtime,
        source,
        target,
        source_format="docx",
        capability="gongwen",
        config_port=config_port,
        options={"remove_numbering": True, "recognize_text": False, "preserve_resources": True},
    )
    metadata, body = _body(output.read_text(encoding="utf-8"))
    assert metadata["份号"] == metadata["发文字号"] == ""
    expected = authored + ("$x$附带说明。" if rich else "")
    tokens = MarkdownIt("commonmark").parse(body)
    assert not any(token.type == "heading_open" for token in tokens)
    children = [child for token in tokens if token.type == "inline" for child in (token.children or [])]
    assert all(child.type in {"text", "softbreak", "image"} for child in children)
    # A renderer may keep an image after a soft break in the same paragraph.
    # Retain that paragraph's text in the independent consumer projection.
    prose = [child.content for child in children if child.type == "text"]
    assert prose == ["开始正文。", expected, "尾部正文。"]
    occurrence_roles = [child.type for child in children if child.type != "softbreak"]
    assert occurrence_roles == (["text", "text", "image", "text"] if rich else ["text"] * 3)
    [diagnostic] = [d for d in result.diagnostics if d.code == "GONGWEN-SOURCE-CONTENT-PRESERVED"]
    assert diagnostic.location == "paragraph:3:ambiguous_combined_id"
    assert any(d.code == "GONGWEN-NEEDS-REVIEW" and "发文机关署名" in d.message for d in result.diagnostics)
    images = [child for token in tokens for child in (token.children or []) if child.type == "image"]
    assert len(images) == int(rich)
    if rich:
        reference = urlsplit(images[0].attrGet("src") or "")
        assert not reference.scheme and not reference.netloc and not reference.query and not reference.fragment
        resource = output.parent / unquote(reference.path)
        assert resource.read_bytes() == picture_bytes
        assert result.bundle is not None
        [artifact] = [
            artifact
            for artifact in result.bundle.artifacts
            if filesystem_path(target / artifact.locator).resolve() == resource.resolve()
        ]
        assert artifact.sha256 == hashlib.sha256(picture_bytes).hexdigest()


@pytest.mark.parametrize("separator", [" ", "\u3000"])
def test_ambiguous_header_does_not_replace_a_fully_owned_single_header(
    configured_ports: tuple[Any, ConfigPortAdapter], tmp_path: Path, separator: str
) -> None:
    runtime, config_port = configured_ports
    authored = f"001{separator}002{separator}国发〔2024〕1号"
    document = Document()
    document.add_paragraph("009 国办发〔2024〕8号")
    _title(document)
    document.add_paragraph("开始正文。")
    document.add_paragraph(authored)
    document.add_paragraph("尾部正文。")
    source = tmp_path / "conflicting-header.docx"
    document.save(source)
    assert Document(source).paragraphs[4].text == authored
    output, result = _bundle(
        runtime,
        source,
        tmp_path / "bundle",
        source_format="docx",
        capability="gongwen",
        config_port=config_port,
        options={"remove_numbering": True},
    )
    metadata, body = _body(output.read_text(encoding="utf-8"))
    assert metadata["份号"] == "009" and metadata["发文字号"] == "国办发〔2024〕8号"
    assert [line for line in body.splitlines() if line] == ["开始正文。", authored, "尾部正文。"]
    [diagnostic] = [d for d in result.diagnostics if d.code == "GONGWEN-SOURCE-CONTENT-PRESERVED"]
    assert diagnostic.location == "paragraph:4:ambiguous_combined_id"


@pytest.mark.parametrize("authored,expected", [("001 工作范围", "工作范围"), ("001 002 工作范围", "002 工作范围")])
def test_numeric_heading_without_document_number_keeps_shipped_cleanup(
    configured_ports: tuple[Any, ConfigPortAdapter], tmp_path: Path, authored: str, expected: str
) -> None:
    runtime, config_port = configured_ports
    document = Document()
    _title(document)
    document.add_paragraph("开始正文。")
    document.add_paragraph(authored)
    document.add_paragraph("尾部正文。")
    source = tmp_path / "ordinary-numeric-heading.docx"
    document.save(source)
    assert Document(source).paragraphs[3].text == authored
    output, result = _bundle(
        runtime,
        source,
        tmp_path / "bundle",
        source_format="docx",
        capability="gongwen",
        config_port=config_port,
        options={"remove_numbering": True},
    )
    _, body = _body(output.read_text(encoding="utf-8"))
    tokens = MarkdownIt("commonmark").parse(body)
    assert [
        (token.tag, tokens[index + 1].content) for index, token in enumerate(tokens) if token.type == "heading_open"
    ] == [("h3", expected)]
    assert [token.content for token in tokens if token.type == "inline"] == ["开始正文。", expected, "尾部正文。"]
    assert not any(d.code == "GONGWEN-SOURCE-CONTENT-PRESERVED" for d in result.diagnostics)
