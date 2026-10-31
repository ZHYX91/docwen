"""Sparse native geometry, resources, budgets and cancellation at final publication."""

from __future__ import annotations

import csv
import hashlib
import posixpath
from io import BytesIO
from pathlib import Path, PurePosixPath
from typing import Any
from urllib.parse import unquote
from zipfile import ZIP_DEFLATED, ZipFile

import mistune
import openpyxl
import pytest
from lxml import etree
from openpyxl.drawing.image import Image as SheetImage
from PIL import Image

from tests.integration.test_spreadsheet_final_content import _primary, _publish, _request, _tables

pytestmark = [pytest.mark.integration, pytest.mark.pr_gate]

_S = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"


@pytest.mark.parametrize("target", ["csv", "tsv"])
def test_sparse_delimited_grid_retains_coordinates_without_creating_blank_cells(
    tmp_path: Path, round_trip_runtime: Any, monkeypatch: pytest.MonkeyPatch, target: str
) -> None:
    source = tmp_path / "sparse-grid.xlsx"
    workbook = openpyxl.Workbook()
    worksheet = workbook.active
    assert worksheet is not None
    worksheet.cell(250, 250, "ONLY")
    workbook.save(source)
    workbook.close()
    reopened = openpyxl.load_workbook(source)
    try:
        assert set(reopened.worksheets[0]._cells) == {(250, 250)}
        assert reopened.worksheets[0]["IP250"].value == "ONLY"
    finally:
        reopened.close()
    loaded: list[Any] = []
    load = openpyxl.load_workbook

    def capture(*args: Any, **kwargs: Any) -> Any:
        view = load(*args, **kwargs)
        loaded.append(view)
        return view

    monkeypatch.setattr(openpyxl, "load_workbook", capture)
    output = tmp_path / "published"
    _, bundle = _publish(round_trip_runtime, source, output, target)
    with _primary(bundle, output).open(encoding="utf-8-sig", newline="") as handle:
        actual = list(csv.reader(handle, delimiter="\t" if target == "tsv" else ","))
    expected = [[""] * 250 for _ in range(250)]
    expected[249][249] = "ONLY"
    assert actual == expected
    assert len(loaded) == 2
    assert all(set(view.worksheets[0]._cells) == {(250, 250)} for view in loaded)


def _png() -> bytes:
    with BytesIO() as stream, Image.new("RGB", (3, 2), (20, 100, 200)) as image:
        image.save(stream, format="PNG")
        return stream.getvalue()


def test_distant_merge_and_image_preserve_cells_and_real_resource_identity(
    tmp_path: Path, round_trip_runtime: Any, monkeypatch: pytest.MonkeyPatch
) -> None:
    image = tmp_path / "source.png"
    pixels = _png()
    image.write_bytes(pixels)
    source = tmp_path / "merge-and-image.xlsx"
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet.title = "Data"
    sheet["IP250"] = 0
    sheet.merge_cells("IP250:IQ251")
    sheet["IP260"] = r"a\|b"
    sheet.add_image(SheetImage(image), "IP260")
    workbook.save(source)
    workbook.close()
    reopened = openpyxl.load_workbook(source)
    try:
        assert reopened["Data"]["IP250"].value == 0
        assert str(next(iter(reopened["Data"].merged_cells.ranges))) == "IP250:IQ251"
        assert reopened["Data"]["IP260"].value == r"a\|b"
        assert len(reopened["Data"]._images) == 1
    finally:
        reopened.close()
    with ZipFile(source) as package:
        assert package.read("xl/media/image1.png") == pixels
        drawing = etree.fromstring(package.read("xl/drawings/drawing1.xml"))
        namespace = {"xdr": "http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing"}
        assert drawing.findtext("xdr:oneCellAnchor/xdr:from/xdr:row", namespaces=namespace) == "259"
        assert drawing.findtext("xdr:oneCellAnchor/xdr:from/xdr:col", namespaces=namespace) == "249"
    loaded: list[Any] = []
    load = openpyxl.load_workbook

    def capture(*args: Any, **kwargs: Any) -> Any:
        view = load(*args, **kwargs)
        loaded.append(view)
        return view

    monkeypatch.setattr(openpyxl, "load_workbook", capture)
    output = tmp_path / "published"
    _, bundle = _publish(
        round_trip_runtime,
        source,
        output,
        "md",
        to_md_keep_images=True,
        image_link_style="markdown_embed",
        table_merge_strategy="fill",
    )
    primary = _primary(bundle, output)
    markdown = primary.read_text(encoding="utf-8-sig")
    assert _tables(markdown) == [[["0", "0"], ["0", "0"]], [["a\\|b\nData_image1.png"]]]
    ast = mistune.create_markdown(renderer="ast", plugins=["table"])(markdown)
    assert isinstance(ast, list)
    images: list[dict[str, Any]] = []

    def walk(nodes: list[dict[str, Any]]) -> None:
        for node in nodes:
            if node["type"] == "image":
                images.append(node)
            walk(node.get("children", []))

    walk(ast)
    assert len(images) == 1
    root_artifact = next(item for item in bundle.artifacts if output / item.locator == primary)
    target = unquote(images[0]["attrs"]["url"])
    logical = posixpath.normpath(posixpath.join(str(PurePosixPath(root_artifact.logical_path).parent), target))
    resources = [item for item in bundle.artifacts if item.logical_path == logical]
    assert len(resources) == 1
    assert resources[0].sha256 == hashlib.sha256(pixels).hexdigest()
    assert (output / resources[0].locator).read_bytes() == pixels
    assert loaded and all(len(view["Data"]._cells) == 5 for view in loaded)


@pytest.mark.parametrize("kind", ["merge", "cumulative-components"])
def test_projection_budget_rejects_before_expansion_or_final_publication(
    tmp_path: Path, round_trip_runtime: Any, monkeypatch: pytest.MonkeyPatch, kind: str
) -> None:
    from docwen_plugin_spreadsheet.to_markdown import converter

    source = tmp_path / "bounded-projection.xlsx"
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    assert sheet is not None
    if kind == "merge":
        sheet["A1"] = "Merged"
        sheet.merge_cells("A1:H8")
    else:
        for current in (sheet, workbook.create_sheet("Second")):
            for index in range(1, 7):
                current.cell(1, index, f"top-{index}")
                current.cell(index, 1, f"left-{index}")
    workbook.save(source)
    workbook.close()
    reopened = openpyxl.load_workbook(source)
    try:
        if kind == "merge":
            assert str(next(iter(reopened.worksheets[0].merged_cells.ranges))) == "A1:H8"
        else:
            assert [[cell.value for cell in current[1]] for current in reopened.worksheets] == [
                ["left-1", "top-2", "top-3", "top-4", "top-5", "top-6"]
            ] * 2
    finally:
        reopened.close()
    # Scale the finite policy to keep this negative bounded on the old code.
    # The source files and runtime/finalizer are otherwise real.
    monkeypatch.setattr(converter, "MAX_SPREADSHEET_PROJECTION_CELLS", 40, raising=False)
    loaded: list[Any] = []
    closed: list[Any] = []
    load, close = openpyxl.load_workbook, openpyxl.Workbook.close

    def capture(*args: Any, **kwargs: Any) -> Any:
        view = load(*args, **kwargs)
        loaded.append(view)
        return view

    def close_view(view: Any) -> None:
        closed.append(view)
        close(view)

    monkeypatch.setattr(openpyxl, "load_workbook", capture)
    monkeypatch.setattr(openpyxl.Workbook, "close", close_view)
    output = tmp_path / "published"
    result = _request(round_trip_runtime, source, output, "md")
    assert not result.success and result.error is not None
    assert result.error.diagnostic_code == "SPREADSHEET-PROJECTION-LIMIT"
    assert result.artifacts == []
    assert not output.exists() or not list(output.rglob("*"))
    if kind == "merge":
        assert loaded == []
    else:
        assert len(loaded) == 2
        assert all(view in closed for view in loaded)


@pytest.mark.parametrize("phase", ["formula-scan", "grid"])
def test_real_runtime_cancellation_closes_both_views_without_partial_output(
    tmp_path: Path, round_trip_runtime: Any, monkeypatch: pytest.MonkeyPatch, phase: str
) -> None:
    import docwen_core.export_semantics as semantics

    source = tmp_path / "cancel.xlsx"
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    assert sheet is not None
    for row in range(1, 65):
        for column in range(1, 65):
            sheet.cell(row, column, f"cell-{row}-{column}")
    workbook.save(source)
    workbook.close()
    reopened = openpyxl.load_workbook(source)
    try:
        assert reopened.worksheets[0]["A1"].value == "cell-1-1"
        assert reopened.worksheets[0]["BL64"].value == "cell-64-64"
    finally:
        reopened.close()
    loaded: list[Any] = []
    closed: list[Any] = []
    load, close = openpyxl.load_workbook, openpyxl.Workbook.close

    def capture(*args: Any, **kwargs: Any) -> Any:
        view = load(*args, **kwargs)
        loaded.append(view)
        if phase == "formula-scan" and len(loaded) == 2:
            round_trip_runtime.cancel("spreadsheet-final-content")
        return view

    def close_view(view: Any) -> None:
        closed.append(view)
        close(view)

    built = 0
    original_cell = semantics.TableSemanticCell

    def cell(*args: Any, **kwargs: Any) -> Any:
        nonlocal built
        built += 1
        if phase == "grid" and built == 32:
            round_trip_runtime.cancel("spreadsheet-final-content")
        return original_cell(*args, **kwargs)

    monkeypatch.setattr(openpyxl, "load_workbook", capture)
    monkeypatch.setattr(openpyxl.Workbook, "close", close_view)
    monkeypatch.setattr(semantics, "TableSemanticCell", cell)
    output = tmp_path / "published"
    result = _request(round_trip_runtime, source, output, "md")
    assert not result.success and result.error is not None
    assert result.error.error_type == "cancelled"
    assert result.artifacts == []
    assert not output.exists() or not list(output.rglob("*"))
    assert len(loaded) == 2 and all(view in closed for view in loaded)
    if phase == "grid":
        assert 32 <= built <= 1032


def test_default_merge_budget_rejects_small_package_before_native_expansion(
    tmp_path: Path, round_trip_runtime: Any, monkeypatch: pytest.MonkeyPatch
) -> None:
    # Run this new policy boundary only after the fix: the original implementation
    # would ask openpyxl to instantiate over a million covered cells.
    source = tmp_path / "large-merge.xlsx"
    workbook = openpyxl.Workbook()
    worksheet = workbook.active
    assert worksheet is not None
    worksheet["A1"] = "Anchor"
    workbook.save(source)
    workbook.close()
    with ZipFile(source) as package:
        members = {name: package.read(name) for name in package.namelist()}
    root = etree.fromstring(members["xl/worksheets/sheet1.xml"])
    merges = etree.SubElement(root, f"{{{_S}}}mergeCells", count="1")
    etree.SubElement(merges, f"{{{_S}}}mergeCell", ref="A1:A1000001")
    members["xl/worksheets/sheet1.xml"] = etree.tostring(root)
    with ZipFile(source, "w", ZIP_DEFLATED) as package:
        for name, payload in members.items():
            package.writestr(name, payload)
    with ZipFile(source) as package:
        xml = etree.fromstring(package.read("xl/worksheets/sheet1.xml"))
        assert [item.get("r") for item in xml.iter(f"{{{_S}}}c")] == ["A1"]
        assert [item.get("ref") for item in xml.iter(f"{{{_S}}}mergeCell")] == ["A1:A1000001"]

    def no_expansion(*_args: Any, **_kwargs: Any) -> Any:
        pytest.fail("native loader must not expand an over-budget merge")

    monkeypatch.setattr(openpyxl, "load_workbook", no_expansion)
    output = tmp_path / "published"
    result = _request(round_trip_runtime, source, output, "md")
    assert not result.success and result.error is not None
    assert result.error.diagnostic_code == "SPREADSHEET-PROJECTION-LIMIT"
    assert result.artifacts == []
    assert not output.exists() or not list(output.rglob("*"))
