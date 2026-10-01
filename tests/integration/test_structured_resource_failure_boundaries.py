"""Final Runtime failure privacy, image states and group staging lifetime."""

from __future__ import annotations

import json
import logging
import tempfile
from dataclasses import replace
from pathlib import Path

import pytest

from docwen_application.controller import ApplicationController, _ManagedPreconversion
from docwen_core.events.task_events import TASK_FAILED
from docwen_plugin_markdown.plugin import MarkdownPlugin
from docwen_runtime._execution_context import _RuntimePluginLogger
from docwen_runtime.adapters import RuntimePortAdapter
from docwen_runtime.engine.route_resolver import RouteResolver
from docwen_runtime.engine.task_manager import TaskManager
from docwen_runtime.output.finalizer import OutputFinalizer
from docwen_runtime.plugin_registry.registry import PluginRegistry
from docwen_runtime.workspace.manager import WorkspaceManager
from tests.integration.test_structured_clipboard_converter_semantics import _template_id
from tests.integration.test_structured_resource_group_runtime import (
    _build_grouped_request,
    _builder,
    _controller,
    _source_ref,
)

pytestmark = [pytest.mark.integration, pytest.mark.pr_gate, pytest.mark.release_gate]


@pytest.mark.parametrize("channel", ["default-error", "task-event"])
def test_actual_runtime_plugin_exception_does_not_disclose_raw_content(tmp_path, monkeypatch, caplog, channel):
    sentinel = "RUNTIME_PRIVATE_EXCEPTION_392b"
    plugin = MarkdownPlugin()

    def explode(_context):
        raise ValueError(sentinel + " authored content")

    monkeypatch.setattr(plugin, "convert", explode)
    store, _bundles, parent, groups = _build_grouped_request(tmp_path)
    events = []
    registry = PluginRegistry()
    registry.register(plugin)
    runtime = RuntimePortAdapter(
        TaskManager(
            registry, RouteResolver(registry), WorkspaceManager(root_dir=str(tmp_path / "runtime")), OutputFinalizer()
        ),
        event_callback=events.append,
    )
    controller = ApplicationController(runtime_port=runtime)
    loggers = []
    original_logger_init = _RuntimePluginLogger.__init__

    def capture_logger(self, task_id):
        original_logger_init(self, task_id)
        loggers.append(self)

    monkeypatch.setattr(_RuntimePluginLogger, "__init__", capture_logger)
    caplog.set_level(logging.DEBUG)
    try:
        results = controller.execute_document_group_batch(parent, groups)
        assert len(results) == 2 and all(not result.success for result in results)
        failed_events = [event for event in events if event.event_type == TASK_FAILED]
        assert len(failed_events) == 2
        assert loggers
        assert sentinel not in caplog.text
        assert sentinel not in repr([logger.messages for logger in loggers])
        if channel == "default-error":
            assert sentinel not in repr([result.error for result in results])
        else:
            assert sentinel not in repr([event.payload for event in failed_events])
    finally:
        store.close()


@pytest.mark.parametrize("target", ["md", "docx", "xlsx", "csv"])
@pytest.mark.parametrize("bound", [True, False])
def test_final_artifact_distinguishes_bound_but_not_rendered_from_missing(tmp_path, target, bound):
    store, bundles, _parent, groups = _build_grouped_request(tmp_path)
    controller = _controller(tmp_path, MarkdownPlugin())
    try:
        request = groups[0]
        if not bound:
            payload = json.loads(Path(bundles[0].main.path).read_bytes())
            payload["resources"] = []
            image = payload["blocks"][0]["inlines"][1]
            image["resourceId"] = None
            image["missingReason"] = "unverified_source"
            missing = store.create_bundle(
                json.dumps(payload).encode(), display_name_template="Missing {index}.dwclip", preview="missing"
            )
            request, _context = _builder(store, [_source_ref(missing)]).single(
                file_path=missing.main.path, target_format="md", action_name="", options={}
            )
        options = {}
        if target == "docx":
            options["template_name"] = _template_id("docx", "English General Template.docx")
        elif target == "xlsx":
            options["template_name"] = _template_id("xlsx", "English Sample Sheet Template.xlsx")
        result = controller.execute_single(
            replace(request, target_format=target, options=options, output_policy=groups[0].output_policy)
        )
        assert result.success, result.error
        code = "NOT-RENDERED" if bound else "UNAVAILABLE"
        assert f"CLIPBOARD-IMAGE-RESOURCE-{code}" in {item.code for item in result.diagnostics}
        primary = next(item for item in result.artifacts if item.is_primary)
        output = Path(primary.staging_path)
        if target == "docx":
            from docx import Document

            text = "\n".join(paragraph.text for paragraph in Document(output).paragraphs)
        elif target == "xlsx":
            from openpyxl import load_workbook

            workbook = load_workbook(output)
            text = repr([list(sheet.values) for sheet in workbook])
            workbook.close()
        else:
            text = output.read_text(encoding="utf-8-sig")
        expected, absent = ("not rendered", "unavailable") if bound else ("unavailable", "not rendered")
        assert f"Image {expected}" in text
        assert f"Image {absent}" not in text
    finally:
        store.close()


@pytest.mark.parametrize("failure", [False, True])
def test_group_managed_staging_lives_until_runtime_then_is_retired(tmp_path, monkeypatch, failure):
    store, bundles, parent, groups = _build_grouped_request(tmp_path)
    controller = _controller(tmp_path, MarkdownPlugin())
    staging: dict[str, Path] = {}
    observed: list[str] = []
    original_execute = controller._execute_runtime_request

    def managed_preconversion(request, **_kwargs):
        owner = tempfile.TemporaryDirectory(prefix="group-pre-", dir=tmp_path)
        stage = Path(owner.name)
        (stage / "ready").write_text(request.request_id, encoding="utf-8")
        staging[request.request_id] = stage
        return _ManagedPreconversion(payload=request, temp_owner=owner, manifest_request=request)

    def observe_runtime(request, scope, task_id):
        stage = staging[request.request_id]
        assert (stage / "ready").read_text(encoding="utf-8") == task_id
        observed.append(task_id)
        if failure and task_id == f"{parent.request_id}-0":
            raise ValueError("controlled first-group runtime failure")
        return original_execute(request, scope, task_id)

    monkeypatch.setattr(controller, "_maybe_preconvert", managed_preconversion)
    monkeypatch.setattr(controller, "_execute_runtime_request", observe_runtime)
    try:
        results = controller.execute_document_group_batch(parent, groups)
        assert observed == [f"{parent.request_id}-0", f"{parent.request_id}-1"]
        assert [result.task_id for result in results] == observed
        assert [result.success for result in results] == [not failure, True]
        assert len(set(staging.values())) == 2
        assert all(not stage.exists() for stage in staging.values())
        assert all(Path(bundle.main.path).is_file() for bundle in bundles)
        assert [Path(bundle.resources[0].path).read_bytes() for bundle in bundles] == [
            b"first-resource",
            b"second-resource",
        ]
    finally:
        store.close()
