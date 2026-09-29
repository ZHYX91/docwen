"""Managed clipboard Markdown through Application, Runtime, plugin, and final publication."""

from __future__ import annotations

import logging
from pathlib import Path
from typing import Any

import pytest
from docx import Document
from PIL import Image

from docwen_application.controller import ApplicationController
from docwen_core.detection import inspect_file, inspect_utf8_markdown_snapshot
from docwen_core.models import FILE_INSPECTION_METADATA_KEY
from docwen_core.models.file_ref import FileRef
from docwen_core.models.request import ConversionRequest, OutputPolicy
from docwen_gui.clipboard_inputs import ClipboardInputStore
from docwen_plugin_markdown.plugin import MarkdownPlugin
from docwen_runtime.adapters import RuntimePortAdapter
from docwen_runtime.engine.route_resolver import RouteResolver
from docwen_runtime.engine.task_manager import TaskManager
from docwen_runtime.output.finalizer import OutputFinalizer
from docwen_runtime.plugin_registry.registry import PluginRegistry
from docwen_runtime.workspace.manager import WorkspaceManager

pytestmark = [pytest.mark.integration, pytest.mark.pr_gate, pytest.mark.release_gate]


def test_clipboard_markdown_snapshot_runs_through_existing_runtime_pipeline(tmp_path: Path) -> None:
    text = "| Name | Value |\n| --- | --- |\n| A | 00123 |\n"
    store = ClipboardInputStore(tmp_path / "managed")
    snapshot = store.create(text, display_name_template="Clipboard Markdown {index}.md")
    source = Path(snapshot.path)
    inspection = inspect_utf8_markdown_snapshot(source)
    assert inspection.workflow_category == "markdown"

    registry = PluginRegistry()
    registry.register(MarkdownPlugin())
    workspaces = WorkspaceManager(root_dir=str(tmp_path / "workspaces"))
    runtime = RuntimePortAdapter(
        TaskManager(
            registry,
            RouteResolver(registry),
            workspaces,
            OutputFinalizer(),
        )
    )
    controller = ApplicationController(runtime_port=runtime)
    request = ConversionRequest(
        request_id="clipboard-markdown-runtime",
        input_refs=[
            FileRef(
                path=str(source),
                format=inspection.detected_format,
                category=inspection.workflow_category,
                size_bytes=source.stat().st_size,
                metadata={FILE_INSPECTION_METADATA_KEY: inspection.to_dict()},
            )
        ],
        target_format="csv",
        output_policy=OutputPolicy(output_dir=str(tmp_path / "published")),
    )

    result = controller.execute_single(request)

    assert result.success is True, result.error
    assert source.read_bytes() == text.encode("utf-8")
    assert result.artifacts
    output = Path(result.artifacts[0].staging_path)
    assert output.is_file()
    delivered = output.read_text(encoding="utf-8-sig")
    assert "Name" in delivered
    assert "00123" in delivered
    assert len(workspaces) == 0

    store.close()


@pytest.mark.parametrize("image_location", ["snapshot-neighbor", "cwd"])
@pytest.mark.parametrize("image_target", ["nearby.png", "file:nearby.png", "file:./nearby.png"])
@pytest.mark.parametrize("syntax", ["inline", "reference"])
def test_clipboard_relative_image_has_no_implicit_source_directory(
    tmp_path: Path,
    round_trip_runtime: Any,
    monkeypatch: pytest.MonkeyPatch,
    image_location: str,
    image_target: str,
    syntax: str,
) -> None:
    store = ClipboardInputStore(tmp_path / "managed")
    text = (
        f"# Clipboard image\n\n![Local][r]\n\n[r]: {image_target}\n"
        if syntax == "reference"
        else f"# Clipboard image\n\n![Local]({image_target})\n"
    )
    snapshot = store.create(text, display_name_template="Clipboard Markdown {index}.md")
    source = Path(snapshot.path)
    image_root = source.parent if image_location == "snapshot-neighbor" else tmp_path
    image_path = image_root / "nearby.png"
    Image.new("RGB", (2, 2), "red").save(image_path)
    image_bytes = image_path.read_bytes()
    monkeypatch.chdir(tmp_path)
    inspection = inspect_utf8_markdown_snapshot(source)
    ref = FileRef(
        path=str(source),
        format="markdown",
        category="markdown",
        metadata={FILE_INSPECTION_METADATA_KEY: inspection.to_dict()},
    )
    request = ConversionRequest(
        request_id="clipboard-resource-boundary",
        input_refs=[ref],
        target_format="docx",
        output_policy=OutputPolicy(output_dir=str(tmp_path / "synthetic-output")),
    )
    controller = ApplicationController(runtime_port=round_trip_runtime)
    result = controller.execute_single(request)
    assert result.success, result.error
    primary = next(artifact for artifact in result.artifacts if artifact.kind == "primary")
    document = Document(primary.staging_path)
    assert not document.inline_shapes
    assert "nearby.png" in "\n".join(paragraph.text for paragraph in document.paragraphs)
    assert source.read_bytes() == text.encode("utf-8")
    assert image_path.read_bytes() == image_bytes

    # An ordinary file genuinely located beside that image retains its normal
    # source-relative behavior through exactly the same conversion pipeline.
    regular = image_root / "regular.md"
    regular.write_text(text, encoding="utf-8")
    ref.path = str(regular)
    ref.metadata = {FILE_INSPECTION_METADATA_KEY: inspect_file(regular).to_dict()}
    request.request_id = "ordinary-resource-control"
    request.output_policy = OutputPolicy(output_dir=str(tmp_path / "ordinary-output"))
    control = controller.execute_single(request)
    assert control.success, control.error
    primary = next(artifact for artifact in control.artifacts if artifact.kind == "primary")
    assert len(Document(primary.staging_path).inline_shapes) == 1
    assert image_path.read_bytes() == image_bytes
    store.close()


@pytest.mark.parametrize("text", ["Ordinary unmarked clipboard text 00123", "https://example.invalid/no-request"])
def test_plain_clipboard_text_exports_real_docx_without_reclassification(
    tmp_path: Path, round_trip_runtime: Any, text: str
) -> None:
    store = ClipboardInputStore(tmp_path / "managed")
    snapshot = store.create(text, display_name_template="Clipboard Markdown {index}.md")
    source = Path(snapshot.path)
    inspection = inspect_utf8_markdown_snapshot(source)
    request = ConversionRequest(
        request_id="clipboard-plain-docx",
        input_refs=[
            FileRef(
                path=str(source),
                format="markdown",
                category="markdown",
                metadata={FILE_INSPECTION_METADATA_KEY: inspection.to_dict()},
            )
        ],
        target_format="docx",
        output_policy=OutputPolicy(output_dir=str(tmp_path / "published")),
    )
    result = ApplicationController(runtime_port=round_trip_runtime).execute_single(request)
    assert result.success, result.error
    primary = next(artifact for artifact in result.artifacts if artifact.kind == "primary")
    document = Document(primary.staging_path)
    assert text in "\n".join(paragraph.text for paragraph in document.paragraphs)
    assert source.read_bytes() == text.encode("utf-8")
    store.close()


@pytest.mark.parametrize(
    ("synthetic", "text"),
    [
        (False, "# Ordinary Markdown\n\nExecutionThread control path.\n"),
        (True, "<html><body>literal clipboard markup</body></html>\n\n# Synthetic Markdown\n"),
    ],
)
def test_execution_thread_revalidates_matching_inspection_through_real_runtime(
    tmp_path: Path,
    round_trip_runtime: Any,
    qtbot,
    synthetic: bool,
    text: str,
) -> None:
    from docwen_gui.qt_bridge.execution import ExecutionThread

    source = tmp_path / ("clipboard.md" if synthetic else "ordinary.md")
    source.write_text(text, encoding="utf-8")
    inspection = inspect_utf8_markdown_snapshot(source) if synthetic else inspect_file(source)
    request = ConversionRequest(
        request_id=f"execution-thread-{'synthetic' if synthetic else 'ordinary'}",
        input_refs=[
            FileRef(
                path=str(source),
                format=inspection.detected_format,
                category=inspection.workflow_category,
                size_bytes=inspection.size_bytes,
                metadata={FILE_INSPECTION_METADATA_KEY: inspection.to_dict()},
            )
        ],
        target_format="docx",
        output_policy=OutputPolicy(output_dir=str(tmp_path / "thread-output")),
    )
    controller = ApplicationController(runtime_port=round_trip_runtime)
    context = {
        "request_id": request.request_id,
        "file_path": str(source),
        "display_name": source.name,
    }
    results: list[Any] = []
    errors: list[str] = []
    thread = ExecutionThread(
        controller=controller,
        request=request,
        context=context,
    )
    thread.result_signal.connect(lambda result, _context: results.append(result))
    thread.error_signal.connect(lambda message, _context: errors.append(message))
    thread.start()
    qtbot.waitUntil(lambda: bool(results or errors), timeout=15000)
    thread.wait(15000)

    assert errors == []
    assert len(results) == 1
    result = results[0]
    assert result.success, result.error
    primary = next(artifact for artifact in result.artifacts if artifact.kind == "primary")
    assert Path(primary.staging_path).is_file()


def test_synthetic_authored_link_targets_do_not_enter_conversion_logs(
    tmp_path: Path,
    round_trip_runtime: Any,
    caplog: pytest.LogCaptureFixture,
) -> None:
    sentinel = "CLIPBOARD_SECRET_QUERY_TOKEN_9f6d"
    text = (
        "# Sensitive links\n\n"
        f"![Remote](https://example.invalid/private.png?token={sentinel})\n\n"
        f"![Missing]({sentinel}.png)\n"
    )
    store = ClipboardInputStore(tmp_path / "managed")
    snapshot = store.create(text, display_name_template="Clipboard Markdown {index}.md")
    source = Path(snapshot.path)
    inspection = inspect_utf8_markdown_snapshot(source)
    request = ConversionRequest(
        request_id="clipboard-log-privacy",
        input_refs=[
            FileRef(
                path=str(source),
                format="markdown",
                category="markdown",
                metadata={FILE_INSPECTION_METADATA_KEY: inspection.to_dict()},
            )
        ],
        target_format="docx",
        output_policy=OutputPolicy(output_dir=str(tmp_path / "published")),
    )

    caplog.set_level(logging.DEBUG)
    result = ApplicationController(runtime_port=round_trip_runtime).execute_single(request)

    assert result.success, result.error
    assert sentinel not in caplog.text
    assert "private.png?token=" not in caplog.text
    assert source.read_bytes() == text.encode("utf-8")
    store.close()


@pytest.mark.parametrize(
    ("yaml_title", "expected_title"),
    [
        (None, "剪贴板 Markdown 1"),
        ("显式宠物标题", "显式宠物标题"),
    ],
)
def test_clipboard_docx_uses_logical_name_for_publication_and_only_as_title_fallback(
    tmp_path: Path, round_trip_runtime: Any, yaml_title: str | None, expected_title: str
) -> None:
    yaml_lines = [
        "---",
        "文字测试: 保留普通字段",
        "列表测试:",
        "  - 第一项",
        "  - 第二项",
    ]
    if yaml_title is not None:
        yaml_lines.append(f"title: {yaml_title}")
    yaml_lines.extend(
        [
            "---",
            "",
            "# 宠物档案",
            "",
            "## 小猫",
            "",
            "| 名称 | 年龄 |",
            "| --- | --- |",
            "| 花花 | 2 |",
            "",
            "## 小狗",
            "",
            "| 名称 | 年龄 |",
            "| --- | --- |",
            "| 旺财 | 3 |",
            "",
        ]
    )
    text = "\n".join(yaml_lines)
    store = ClipboardInputStore(tmp_path / "managed")
    snapshot = store.create(text, display_name_template="剪贴板 Markdown {index}.md")
    source = Path(snapshot.path)
    physical_stem = source.stem
    inspection = inspect_utf8_markdown_snapshot(source)
    logical_name = snapshot.display_name
    logical_stem = Path(logical_name).stem
    request = ConversionRequest(
        request_id=f"clipboard-logical-docx-{'explicit' if yaml_title else 'fallback'}",
        input_refs=[
            FileRef(
                path=str(source),
                format="markdown",
                category="markdown",
                logical_path=logical_name,
                metadata={FILE_INSPECTION_METADATA_KEY: inspection.to_dict()},
            )
        ],
        target_format="docx",
        output_policy=OutputPolicy(output_dir=str(tmp_path / "published")),
    )

    result = ApplicationController(runtime_port=round_trip_runtime).execute_single(request)

    assert result.success, result.error
    primary = next(artifact for artifact in result.artifacts if artifact.kind == "primary")
    output = Path(primary.staging_path)
    assert primary.suggested_name == f"{logical_stem}.docx"
    assert logical_stem in output.name
    assert logical_stem in output.parent.name
    assert physical_stem not in output.name
    assert physical_stem not in output.parent.name

    document = Document(output)
    paragraphs = [paragraph.text for paragraph in document.paragraphs if paragraph.text]
    assert paragraphs[:4] == [expected_title, "宠物档案", "小猫", "小狗"]
    assert physical_stem not in paragraphs
    assert len(document.tables) == 2
    assert source.read_bytes() == text.encode("utf-8")
    store.close()

