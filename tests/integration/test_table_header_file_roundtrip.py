"""Direct Markdown/DOCX files keep roles separate from style and pagination."""

from __future__ import annotations

from functools import cache
from pathlib import Path
from typing import Any
from zipfile import ZipFile

import lxml.etree as etree
import pytest
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from docwen_core.models.file_ref import FileRef
from docwen_core.models.request import ConversionRequest
from docwen_plugin_document.to_markdown.converter import DocxToMarkdownConverter
from docwen_plugin_markdown.document_semantics import analyze_document_semantics
from docwen_plugin_markdown.mistune_extensions import parse_markdown_text
from docwen_plugin_markdown.to_docx.converter import MdToDocxConverter
from docwen_runtime.config.document_styles import build_document_style_catalog
from tests.support.cancellation import FakeCancellationTokenView
from tests.support.config import FakeConfigView
from tests.support.execution import FakeExecutionContext
from tests.support.logging import FakePluginLogger
from tests.support.progress import FakeProgressSink
from tests.support.workspace import FakeWorkspaceHandle

pytestmark = [pytest.mark.integration, pytest.mark.pr_gate]

_SOURCE = """| Region | Sales | < |
| Quarter | Q1 | Q2 |
| --- || --- | --- |
| North | 10 | 12 |
| ^ | 8 | 11 |
"""


@cache
def _style_catalog():
    return build_document_style_catalog(
        {"gui": {"language": {"locale": "en_US"}}},
        locales_dir=Path(__file__).resolve().parents[2] / "i18n" / "locales",
    )


def _context(path: Path, staging: Path, *, target: str, structural: bool = True) -> FakeExecutionContext:
    staging.mkdir()
    request = ConversionRequest(
        request_id=f"table-{target}",
        input_refs=[FileRef(path=str(path), format="markdown" if target == "docx" else "docx", category="document")],
        target_format=target,
        options={
            "locale": "en_US",
            "heading_merge_mode": "never",
            "markdown_extensions": {
                "input": {"structural_tables": structural},
                "output": {"structural_tables": structural},
            },
        },
    )
    return FakeExecutionContext(
        request,
        FakeWorkspaceHandle(str(path), str(staging)),
        FakeConfigView(),
        FakeProgressSink(),
        FakeCancellationTokenView(),
        FakePluginLogger(),
        document_style_catalog=_style_catalog(),
    )


def _export(tmp_path: Path, source: str) -> Path:
    path = tmp_path / "source.md"
    path.write_text(source, encoding="utf-8")
    result = MdToDocxConverter().convert(_context(path, tmp_path / "docx-output", target="docx"))
    assert result.success, result.error
    # Reverse conversion must depend only on the isolated DOCX, not its source.
    path.unlink()
    return Path(result.artifacts[0].staging_path)


def _import(tmp_path: Path, path: Path, *, structural: bool = True):
    result = DocxToMarkdownConverter().convert(
        _context(path, tmp_path / "md-output", target="md", structural=structural)
    )
    return result, Path(result.artifacts[0].staging_path).read_text(encoding="utf-8") if result.success else ""


def _metadata(markdown: str) -> list[dict[str, Any]]:
    analysis = analyze_document_semantics(parse_markdown_text(markdown), current_v3=True)
    assert not analysis.has_errors, analysis.diagnostics
    return [node["_document_semantics_table"] for node in analysis.ast if node["type"] == "table"]


def _disable_styles(document: Any, *, value: str = "0") -> None:
    for table in document.tables:
        for marker in list(table._tbl.iter(qn("w:cnfStyle"))):
            marker.getparent().remove(marker)
        look = table._tbl.tblPr.find(qn("w:tblLook"))
        look.set(qn("w:firstRow"), value)
        look.set(qn("w:firstColumn"), value)


@pytest.mark.parametrize("repeat, expected", [("true", "always"), ("false", "never"), (None, "inherit")])
@pytest.mark.parametrize("disabled", ["0", "false", "off"])
def test_direct_file_roles_merges_and_repeat_policy_survive_style_edits(
    tmp_path: Path, repeat: str | None, expected: str, disabled: str
) -> None:
    path = _export(tmp_path, _SOURCE + (f"{{repeat-header={repeat}}}\n" if repeat is not None else ""))
    with ZipFile(path) as package:
        maps = [
            etree.fromstring(package.read(name))
            for name in package.namelist()
            if name.startswith("customXml/item") and name.endswith(".xml") and "/" not in name[len("customXml/") :]
        ]
        [role_map] = [root for root in maps if root.tag == "{urn:docwen:table-roles:v1}tableRoles"]
        assert set(role_map[0].attrib) == {"bookmark", "rows", "columns", "shape"}
        assert (role_map[0].get("rows"), role_map[0].get("columns")) == ("2", "1")
        assert b"w:gridSpan" in package.read("word/document.xml")
        assert b"w:vMerge" in package.read("word/document.xml")

    document: Any = Document(str(path))
    table = document.tables[0]
    markers = list(table._tbl.iter(qn("w:tblHeader")))
    assert len(markers) == (2 if repeat is not None else 0)
    assert all(marker.get(qn("w:val")) == ("1" if repeat == "true" else "0") for marker in markers)
    _disable_styles(document, value=disabled)
    document.save(str(path))
    edited_bytes = path.read_bytes()

    result, markdown = _import(tmp_path, path)

    assert result.success, result.error
    assert not result.diagnostics
    [metadata] = _metadata(markdown)
    assert (metadata["header_rows"], metadata["header_columns"], metadata["repeat_header"]) == (2, 1, expected)
    assert [
        (anchor["row"], anchor["column"], anchor["row_span"], anchor["column_span"]) for anchor in metadata["anchors"]
    ] == [
        (0, 0, 1, 1),
        (0, 1, 1, 2),
        (1, 0, 1, 1),
        (1, 1, 1, 1),
        (1, 2, 1, 1),
        (2, 0, 2, 1),
        (2, 1, 1, 1),
        (2, 2, 1, 1),
        (3, 1, 1, 1),
        (3, 2, 1, 1),
    ]
    assert path.read_bytes() == edited_bytes


def test_direct_file_text_edits_and_table_moves_keep_distinct_roles(tmp_path: Path) -> None:
    second = "| Second | B | C |\n| --- | --- || --- |\n| A | D | E |\n| F | G | H |\n| I | J | K |\n"
    path = _export(tmp_path, _SOURCE + "\nBetween tables.\n\n" + second)
    document: Any = Document(str(path))
    _disable_styles(document)
    document.tables[0].cell(2, 1).text = "Current edited value"
    moved = document.tables[1]._tbl
    document.element.body.insert(0, moved)
    document.save(str(path))

    result, markdown = _import(tmp_path, path)

    assert result.success, result.error
    assert not result.diagnostics
    assert [(item["header_rows"], item["header_columns"]) for item in _metadata(markdown)] == [(1, 2), (2, 1)]
    assert markdown.index("Second") < markdown.index("Region")
    assert "Current edited value" in markdown
    assert "| North | 10 |" not in markdown


@pytest.mark.parametrize("edit", ["add_row", "merge", "delete_binding"])
def test_direct_file_stale_roles_warn_and_use_current_geometry(tmp_path: Path, edit: str) -> None:
    path = _export(tmp_path, _SOURCE)
    document: Any = Document(str(path))
    _disable_styles(document)
    table = document.tables[0]
    if edit == "add_row":
        table.add_row()
    elif edit == "merge":
        table.cell(2, 1).merge(table.cell(2, 2))
    else:
        for marker in list(table._tbl.iter(qn("w:bookmarkStart"))) + list(table._tbl.iter(qn("w:bookmarkEnd"))):
            marker.getparent().remove(marker)
    document.save(str(path))

    result, markdown = _import(tmp_path, path)

    assert result.success, result.error
    assert [item.code for item in result.diagnostics] == ["DOCX2MD-TABLE-ROLES-STALE"]
    [metadata] = _metadata(markdown)
    assert (metadata["header_rows"], metadata["header_columns"]) == (0, 0)
    assert metadata["row_count"] == (5 if edit == "add_row" else 4)
    if edit == "merge":
        assert any(
            anchor["row"] == 2 and anchor["column"] == 1 and anchor["column_span"] == 2
            for anchor in metadata["anchors"]
        )


def test_direct_file_native_pagination_does_not_expand_verified_header_roles(tmp_path: Path) -> None:
    path = _export(tmp_path, _SOURCE)
    document: Any = Document(str(path))
    _disable_styles(document)
    for row in document.tables[0].rows[:3]:
        repeat = OxmlElement("w:tblHeader")
        repeat.set(qn("w:val"), "1")
        row._tr.get_or_add_trPr().append(repeat)
    document.save(str(path))

    result, markdown = _import(tmp_path, path)

    assert result.success, result.error
    [metadata] = _metadata(markdown)
    assert (metadata["header_rows"], metadata["header_columns"], metadata["repeat_header"]) == (2, 1, "always")


def test_direct_file_disabled_output_dialect_keeps_explicit_flattening_warning(tmp_path: Path) -> None:
    path = _export(tmp_path, _SOURCE + "{repeat-header=true}\n")
    result, markdown = _import(tmp_path, path, structural=False)

    assert result.success, result.error
    assert "docwen.conversion.markdown_extension.structural_tables.flattened" in [
        item.code for item in result.diagnostics
    ]
    assert "repeat-header" not in markdown
    assert "||" not in markdown


@pytest.mark.parametrize("repeat, expected", [("true", "always"), ("false", "never"), (None, "inherit")])
@pytest.mark.parametrize("structural", [False, True])
def test_direct_file_ordinary_table_repeat_only_output_policy(
    tmp_path: Path, repeat: str | None, expected: str, structural: bool
) -> None:
    source = "| Name | Value |\n| --- | --- |\n| A | 1 |\n"
    if repeat is not None:
        source += f"{{repeat-header={repeat}}}\n"
    path = _export(tmp_path, source)

    result, markdown = _import(tmp_path, path, structural=structural)

    assert result.success, result.error
    [metadata] = _metadata(markdown)
    assert (metadata["header_rows"], metadata["header_columns"]) == (1, 0)
    assert all(anchor["row_span"] == anchor["column_span"] == 1 for anchor in metadata["anchors"])
    assert metadata["repeat_header"] == (expected if structural else "inherit")
    assert [item.code for item in result.diagnostics] == (
        ["docwen.conversion.markdown_extension.structural_tables.flattened"]
        if not structural and repeat == "true"
        else []
    )
