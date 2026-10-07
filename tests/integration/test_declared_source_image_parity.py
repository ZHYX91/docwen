"""Raw and declared source inputs produce the same image and table semantics."""

from __future__ import annotations

import hashlib
from pathlib import Path
from typing import Any
from zipfile import ZipFile

import pytest
from docx import Document
from docx.oxml.ns import qn
from PIL import Image

from docwen_core.models.file_ref import FileRef
from docwen_core.models.request import ConversionRequest, OutputPolicy
from tests.integration._round_trip_helper import md_to_docx

pytestmark = [pytest.mark.integration, pytest.mark.pr_gate]


def test_declared_source_images_match_raw_source_with_tables_and_captions(
    round_trip_runtime: Any,
    tmp_path: Path,
) -> None:
    raw = tmp_path / "raw"
    raw.mkdir()
    for directory in ("assets", "other"):
        (raw / directory).mkdir()
    images = {"assets/chart.png": "red", "other/chart.png": "blue", "assets/中文 图.png": "green"}
    for relative, color in images.items():
        Image.new("RGB", (6, 4), color).save(raw / relative)
    tokens = ["![[chart.png|120x80]]", "![Second](other/chart.png)", "![Third](<assets/中文 图.png>)"]
    source = (
        "# Report\n\nFigure: Chart ^chart\n\n"
        + "\n\n".join(tokens)
        + "\n\n| Group | A | < |\n| Detail | B | C |\n| --- || --- | --- |\n| East | 1 | 2 |\n| ^ | 3 | 4 |\n\n`![[missing.png]]`\n"
    )
    raw_source = raw / "note.md"
    raw_source.write_text(source, encoding="utf-8", newline="")
    options = {
        "markdown_extensions": {"input": {"structural_tables": True, "captions_references": True}},
        "remove_numbering": False,
    }
    direct = md_to_docx(round_trip_runtime, raw_source, tmp_path / "direct", options=options)

    isolated = tmp_path / "isolated"
    isolated.mkdir()
    snapshot = isolated / "source.md"
    snapshot.write_bytes(raw_source.read_bytes())
    refs = [
        FileRef(
            path=str(snapshot),
            format="markdown",
            category="document",
            input_kind="document",
            input_role="source",
            logical_path="notes/note.md",
            media_type="text/markdown",
        )
    ]
    for index, relative in enumerate(images):
        resource = isolated / f"linked-{index}.png"
        resource.write_bytes((raw / relative).read_bytes())
        refs.append(
            FileRef(
                path=str(resource),
                format="png",
                category="image",
                input_kind="resource",
                input_role="linked_resource",
                logical_path=relative,
                media_type="image/png",
            )
        )
    result = round_trip_runtime.execute(
        ConversionRequest(
            request_id="declared-source-parity",
            input_refs=refs,
            target_format="docx",
            output_policy=OutputPolicy(output_dir=str(tmp_path / "declared")),
            options={
                **options,
                "markdown_resource_bindings": {
                    "authored_sha256": hashlib.sha256(source.encode()).hexdigest(),
                    "images": [
                        {"authored_token": token, "logical_path": relative}
                        for token, relative in zip(tokens, images, strict=True)
                    ],
                },
            },
        )
    )
    assert result.success, result.error
    declared = Path(next(artifact.staging_path for artifact in result.artifacts if artifact.kind == "primary"))

    def semantics(path: Path) -> tuple:
        doc = Document(str(path))
        with ZipFile(path) as archive:
            media = sorted(
                hashlib.sha256(archive.read(name)).hexdigest()
                for name in archive.namelist()
                if name.startswith("word/media/")
            )
        return (
            [element.text for element in doc.element.body.iter(qn("w:t"))],
            [(shape.width, shape.height) for shape in doc.inline_shapes],
            [table._tbl.xml for table in doc.tables],
            [element.text for element in doc.element.body.iter(qn("w:instrText"))],
            media,
        )

    direct_semantics = semantics(direct)
    assert len(direct_semantics[1]) == 3
    assert len(direct_semantics[2]) == 1
    assert any("SEQ Figure" in (instruction or "") for instruction in direct_semantics[3])
    assert semantics(declared) == direct_semantics
    assert snapshot.read_bytes() == raw_source.read_bytes() == source.encode()


def test_declared_source_wikilink_uses_authenticated_navigation_uri(
    round_trip_runtime: Any,
    tmp_path: Path,
) -> None:
    source = "See [[Other#Section|Other note]].\n"
    snapshot = tmp_path / "source.md"
    snapshot.write_text(source, encoding="utf-8")
    href = "obsidian://open?vault=Knowledge&file=Notes%2FOther.md%23Section"
    source_ref = FileRef(
        path=str(snapshot),
        format="markdown",
        category="document",
        input_kind="document",
        input_role="source",
        logical_path="Notes/Current.md",
        media_type="text/markdown",
    )

    result = round_trip_runtime.execute(
        ConversionRequest(
            request_id="declared-wiki-navigation",
            input_refs=[source_ref],
            target_format="docx",
            output_policy=OutputPolicy(output_dir=str(tmp_path / "declared-link")),
            options={
                "markdown_resource_bindings": {
                    "authored_sha256": hashlib.sha256(source.encode()).hexdigest(),
                    "images": [],
                    "wiki_links": [
                        {
                            "authored_token": "[[Other#Section|Other note]]",
                            "href": href,
                        }
                    ],
                },
            },
        )
    )

    assert result.success, result.error
    output = Path(next(artifact.staging_path for artifact in result.artifacts if artifact.kind == "primary"))
    document = Document(str(output))
    hyperlink_targets = {
        relationship.target_ref
        for relationship in document.part.rels.values()
        if relationship.reltype.endswith("/hyperlink")
    }
    assert href in hyperlink_targets
    assert "Other note" in "\n".join(paragraph.text for paragraph in document.paragraphs)
    assert snapshot.read_text(encoding="utf-8") == source
