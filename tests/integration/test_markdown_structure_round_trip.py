"""Real conversion preserves Markdown block boundaries and literal examples."""

from __future__ import annotations

from pathlib import Path

import pytest
from docx import Document

from tests.integration._round_trip_helper import docx_to_md, md_to_docx

pytestmark = pytest.mark.integration


@pytest.mark.parametrize("newline", ["\n", "\r\n"])
def test_setext_and_html_breaks_survive_production_conversion(tmp_path: Path, round_trip_runtime, newline):
    fenced = "```markdown\n---\nReportName: literal<br>\nUnit: office\n---\n```"
    source = (
        fenced
        + "\n\nFirst line\nSecond `code` tail\n---\nBody `<br>` and $x+y$.\n\n"
        + "# First<br>Second\n\n| Head |\n| --- |\n| Left<br>Right |\n\n"
        + "- Kept item\n---\nAfter rule\n"
    ).replace("\n", newline)
    input_path = tmp_path / "structure.md"
    original = source.encode("utf-8")
    input_path.write_bytes(original)
    output = md_to_docx(
        round_trip_runtime,
        input_path,
        tmp_path / "forward",
        options={"heading_merge_mode": "never"},
    )
    document = Document(str(output))
    paragraphs = document.paragraphs
    assert any(
        "First line" in paragraph.text and "Second" in paragraph.text and "tail" in paragraph.text
        for paragraph in paragraphs
    )
    assert any(paragraph.text == "First\nSecond" for paragraph in paragraphs)
    assert any(paragraph.text == "Kept item" for paragraph in paragraphs)
    # Authenticated fenced blocks live in a native content control; the
    # python-docx top-level paragraphs collection excludes those paragraphs.
    assert document.element.xpath('.//w:p[.//w:t[contains(., "ReportName")]]//w:t/text()') == [
        "---",
        "ReportName: literal<br>",
        "Unit: office",
        "---",
    ]
    assert document.tables[0].cell(1, 0).text == "Left\nRight"
    assert not any("# First" in paragraph.text or "## - Kept item" in paragraph.text for paragraph in paragraphs)
    returned = docx_to_md(
        round_trip_runtime,
        output,
        tmp_path / "reverse",
        options={"to_md_keep_images": False, "to_md_enable_ocr": False},
    )
    assert fenced in returned.replace("\r\n", "\n")
    assert "`<br>`" in returned
    assert "# First<br>Second" in returned
    assert "x+y" in returned
    assert "Kept item" in returned and "After rule" in returned
    assert input_path.read_bytes() == original
    for internal_marker in ("DOCWEN_SEMANTIC", "{{IMAGE@", "docwen-runtime-marker"):
        assert internal_marker not in returned

    returned_path = tmp_path / "returned.md"
    returned_path.write_text(returned, encoding="utf-8")
    second_output = md_to_docx(
        round_trip_runtime, returned_path, tmp_path / "second", options={"heading_merge_mode": "never"}
    )
    second_document = Document(str(second_output))
    assert any(paragraph.text == "First\nSecond" for paragraph in second_document.paragraphs)
    assert second_document.tables[0].cell(1, 0).text == "Left\nRight"
