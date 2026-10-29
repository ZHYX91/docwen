"""Native resources retain occurrence identity through Runtime and Bundle commit."""

from __future__ import annotations

import base64
import hashlib
import html
import posixpath
import re
from io import BytesIO
from pathlib import Path, PurePosixPath
from typing import Any
from urllib.parse import unquote
from zipfile import ZipFile

import mistune
import pytest
from docx import Document
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from lxml import etree
from PIL import Image

from docwen_application.bundle_mapping import build_bundle_draft
from docwen_core.markdown_extensions import MarkdownExtensions
from docwen_core.models.file_ref import FileRef
from docwen_core.models.request import ConversionRequest, OutputPolicy
from docwen_runtime.output.artifact_bundle import ArtifactBundleCommitter
from tests.integration._round_trip_helper import _primary_path, _run

pytestmark = [pytest.mark.integration, pytest.mark.pr_gate]

_W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
_R = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
_A = "http://schemas.openxmlformats.org/drawingml/2006/main"
_V = "urn:schemas-microsoft-com:vml"


def _png(color: tuple[int, int, int]) -> bytes:
    with BytesIO() as buffer, Image.new("RGB", (3, 2), color) as image:
        image.save(buffer, format="PNG")
        return buffer.getvalue()


def _bundle(
    runtime: Any, source: Path, output: Path, *, config_snapshot: dict[str, Any] | None = None, **options: Any
) -> tuple[Any, str, Any]:
    output.mkdir()
    original = source.read_bytes()
    format_name = source.suffix.lstrip(".")
    category = {"docx": "document", "enex": "markup", "pptx": "presentation"}[format_name]
    result = runtime.execute(
        ConversionRequest(
            request_id="native-resource-identity",
            input_refs=[FileRef(str(source), format_name, category, len(original))],
            target_format="md",
            config_snapshot=config_snapshot or {},
            output_policy=OutputPolicy(output_dir=str(output)),
            options={
                "to_md_keep_images": True,
                "to_md_enable_ocr": False,
                **({"remove_numbering": False} if format_name == "docx" else {}),
                **options,
            },
        )
    )
    assert result.success, result.error
    assert source.read_bytes() == original
    draft = build_bundle_draft(
        profile="document_with_resources", output_media_type="text/markdown", artifacts=result.artifacts
    )
    bundle = ArtifactBundleCommitter().commit(task_id="native-resource-identity", staging_root=str(output), draft=draft)
    preferred = next(entry.artifact_id for entry in bundle.entries if entry.preferred)
    primary = next(item for item in bundle.artifacts if item.artifact_id == preferred)
    for artifact in bundle.artifacts:
        content = (output / artifact.locator).read_bytes()
        assert artifact.logical_path == artifact.locator
        assert artifact.size_bytes == len(content)
        assert artifact.sha256 == hashlib.sha256(content).hexdigest()
    return bundle, (output / primary.locator).read_text(encoding="utf-8-sig"), primary


def _target_artifact(bundle: Any, primary: Any, target: str) -> Any:
    # Independent filesystem/Bundle interpretation: this does not call any
    # producer/finalizer link parser or filename-replacement helper.
    path = unquote(target.split("#", 1)[0].split("?", 1)[0])
    logical = posixpath.normpath(posixpath.join(str(PurePosixPath(primary.logical_path).parent), path))
    matches = [item for item in bundle.artifacts if item.logical_path == logical]
    assert len(matches) == 1, (target, logical, [item.logical_path for item in bundle.artifacts])
    return matches[0]


def _image_targets(markdown: str) -> list[tuple[str, str]]:
    ast = mistune.create_markdown(
        renderer="ast", plugins=["table", "mark", "superscript", "subscript", "strikethrough", "math"]
    )(markdown)
    result: list[tuple[str, str]] = []

    def walk(nodes: list[dict[str, Any]]) -> None:
        for node in nodes:
            if node["type"] == "image":
                label = "".join(child.get("raw", "") for child in node.get("children", []))
                result.append((label, node["attrs"]["url"]))
            walk(node.get("children", []))

    walk(ast)
    return result


@pytest.mark.parametrize("order", [(0, 1, 2), (2, 0, 1)])
@pytest.mark.parametrize("preoccupied", [False, True])
def test_caption_image_occurrences_keep_original_bytes_after_repeated_source_names(
    tmp_path: Path, round_trip_runtime: Any, order: tuple[int, ...], preoccupied: bool
) -> None:
    images = [_png(color) for color in ((230, 0, 0), (0, 230, 0), (0, 0, 230))]
    expected: list[tuple[str, bytes]] = []
    source_lines: list[str] = []
    selections = [(-1, "photo_1.png")] if preoccupied else []
    selections += [(index, "photo.png") for index in order]
    for occurrence, (index, name) in enumerate(selections):
        folder = tmp_path / f"source-{occurrence}"
        folder.mkdir()
        content = _png((80, 80, 80)) if index < 0 else images[index]
        (folder / name).write_bytes(content)
        label = f"Occurrence {occurrence}"
        expected.append((label, content))
        source_lines.append(
            f"Figure: Caption {occurrence} ^figure-{occurrence}\n\n![{label}](source-{occurrence}/{name})\n\nEnd of occurrence {occurrence}.\n"
        )
    source = tmp_path / "captions.md"
    source.write_text("\n".join(source_lines), encoding="utf-8")
    forward = _run(
        round_trip_runtime,
        "resource-caption-writer",
        source,
        source_format="markdown",
        target_format="docx",
        output_dir=tmp_path / "forward",
    )
    docx = _primary_path(forward)
    with ZipFile(docx) as package:
        root = etree.fromstring(package.read("word/document.xml"))
        relationships = etree.fromstring(package.read("word/_rels/document.xml.rels"))
        targets = {item.get("Id"): item.get("Target") for item in relationships}
        observed: list[bytes] = []
        for blip in root.iter(f"{{{_A}}}blip"):
            target = targets[blip.get(f"{{{_R}}}embed")]
            observed.append(package.read(posixpath.normpath(posixpath.join("word", target))))
        assert observed == [content for _, content in expected]
        titles = [
            item.get("title")
            for item in root.iter("{http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing}docPr")
        ]
        assert titles == [name for _, name in selections]
    bundle, markdown, primary = _bundle(
        round_trip_runtime,
        docx,
        tmp_path / "reverse",
        markdown_extensions={"output": MarkdownExtensions(captions_references=True).to_dict()},
        image_link_style="markdown_embed",
    )
    links = _image_targets(markdown)
    assert [label for label, _ in links] == [label for label, _ in expected]
    assert [_target_artifact(bundle, primary, target).sha256 for _, target in links] == [
        hashlib.sha256(content).hexdigest() for _, content in expected
    ]
    assert [markdown.index(f"Caption {index}") for index in range(len(expected))] == sorted(
        markdown.index(f"Caption {index}") for index in range(len(expected))
    )


@pytest.mark.parametrize("vml", [False, True], ids=["drawingml", "vml"])
def test_repeated_native_image_occurrences_are_never_overwritten(
    tmp_path: Path, round_trip_runtime: Any, vml: bool
) -> None:
    document = Document()
    values = [_png(color) for color in ((240, 0, 0), (0, 0, 240), (240, 0, 0), (0, 240, 0))]
    for index, value in enumerate(values):
        path = tmp_path / f"image-{index}.png"
        path.write_bytes(value)
        paragraph = document.add_paragraph(f"Image {index}")
        shape = paragraph.add_run().add_picture(str(path))
        if vml:
            rid = next(shape._inline.iter(f"{{{_A}}}blip")).get(qn("r:embed"))
            drawing = shape._inline.getparent()
            pict = OxmlElement("w:pict")
            node = etree.SubElement(pict, f"{{{_V}}}shape")
            etree.SubElement(node, f"{{{_V}}}imagedata").set(qn("r:id"), rid)
            drawing.getparent().replace(drawing, pict)
    source = tmp_path / "native-images.docx"
    document.save(str(source))
    bundle, markdown, primary = _bundle(
        round_trip_runtime, source, tmp_path / "published", image_link_style="markdown_embed"
    )
    links = _image_targets(markdown)
    assert len(links) == len(values)
    assert [_target_artifact(bundle, primary, target).sha256 for _, target in links] == [
        hashlib.sha256(value).hexdigest() for value in values
    ]


@pytest.mark.parametrize(
    "target",
    [
        "../unrelated/docx-image_1_rId9.png",
        "C:/unrelated/docx-image_1_rId9.png",
        "../my%20folder/docx-image_1_rId9.png",
        "../literal%2520/docx-image_1_rId9.png",
        "../unrelated/docx-image_1_rId9.png?download=1#authored%29part",
        "https://example.invalid/docx-image_1_rId9.png",
        r"\\server\docx-image_1_rId9.png",
    ],
)
def test_unrelated_native_hyperlink_keeps_qualified_identity(
    tmp_path: Path, round_trip_runtime: Any, target: str
) -> None:
    document = Document()
    image = tmp_path / "red.png"
    image.write_bytes(_png((230, 0, 0)))
    document.add_paragraph().add_run().add_picture(str(image))
    # python-docx's fresh package image relationship is an independent OOXML
    # fact, not a mocked finalizer filename.
    assert document.part.rels["rId9"].reltype == RT.IMAGE
    paragraph = document.add_paragraph()
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), document.part.relate_to(target, RT.HYPERLINK, is_external=True))
    run = OxmlElement("w:r")
    text = OxmlElement("w:t")
    text.text = "Unrelated native hyperlink"
    run.append(text)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)
    source = tmp_path / "qualified-link.docx"
    document.save(str(source))
    with ZipFile(source) as package:
        rels = etree.fromstring(package.read("word/_rels/document.xml.rels"))
        assert [item.get("Target") for item in rels if item.get("Type") == RT.HYPERLINK] == [target]
    bundle, markdown, primary = _bundle(
        round_trip_runtime, source, tmp_path / "output", image_link_style="markdown_embed"
    )
    assert f"[Unrelated native hyperlink]({target})" in markdown
    [(_, actual_image)] = _image_targets(markdown)
    assert _target_artifact(bundle, primary, actual_image).sha256 == hashlib.sha256(image.read_bytes()).hexdigest()


@pytest.mark.parametrize("description", ["portrait", "", "a|b [close] & literal"])
def test_native_wiki_image_link_target_is_left_and_alias_is_its_occurrence(
    tmp_path: Path, round_trip_runtime: Any, description: str
) -> None:
    document = Document()
    image = tmp_path / "image.png"
    image.write_bytes(_png((18, 99, 202)))
    shape = document.add_paragraph().add_run().add_picture(str(image))
    shape._inline.docPr.set("descr", description)
    source = tmp_path / "wiki-image.docx"
    document.save(str(source))
    bundle, markdown, primary = _bundle(round_trip_runtime, source, tmp_path / "output", image_link_style="wiki_link")
    # This oracle deliberately does not import DocWen's Wiki regex/parser.
    tokens = re.findall(r"(?<!!)\[\[([^\]\r\n]+)\]\]", markdown)
    assert len(tokens) == 1
    target, separator, alias = tokens[0].partition("|")
    assert separator == "|"
    assert html.unescape(alias) == description
    assert _target_artifact(bundle, primary, target).sha256 == hashlib.sha256(image.read_bytes()).hexdigest()


def _enex_note(title: str, payloads: list[tuple[bytes, str, str]]) -> str:
    body: list[str] = []
    resources: list[str] = []
    for index, (content, name, mime) in enumerate(payloads):
        digest = hashlib.md5(content).hexdigest()
        body.append(f'<p>Occurrence {index}</p><en-media hash="{digest}" type="{mime}"/>')
        resources.append(
            f'<resource><data encoding="base64">{base64.b64encode(content).decode()}</data><mime>{mime}</mime><resource-attributes><file-name>{name}</file-name></resource-attributes></resource>'
        )
    return f"<note><title>{title}</title><content><![CDATA[<en-note>{''.join(body)}</en-note>]]></content>{''.join(resources)}</note>"


@pytest.mark.parametrize("same_note", [False, True])
@pytest.mark.parametrize("same_bytes", [False, True])
def test_enex_resource_links_keep_each_note_and_occurrence_bytes(
    tmp_path: Path, round_trip_runtime: Any, same_note: bool, same_bytes: bool
) -> None:
    red = _png((230, 0, 0))
    blue = red if same_bytes else _png((0, 0, 230))
    first = (red, "photo.png", "image/png")
    second = (blue, "photo.png", "image/png")
    notes = (
        [_enex_note("First", [first, second])]
        if same_note
        else [_enex_note("First", [first]), _enex_note("Second", [second])]
    )
    source = tmp_path / "notes.enex"
    source.write_text("<en-export>" + "".join(notes) + "</en-export>", encoding="utf-8")
    bundle, markdown, primary = _bundle(
        round_trip_runtime, source, tmp_path / "output", image_link_style="markdown_embed"
    )
    links = _image_targets(markdown)
    assert len(links) == 2
    assert [_target_artifact(bundle, primary, target).sha256 for _, target in links] == [
        hashlib.sha256(red).hexdigest(),
        hashlib.sha256(blue).hexdigest(),
    ]
    if not same_note:
        first_text, second_text = markdown.split("# First", 1)[1].split("# Second", 1)
        assert len(_image_targets(first_text)) == len(_image_targets(second_text)) == 1


def test_enex_same_named_non_image_attachments_keep_note_ownership(tmp_path: Path, round_trip_runtime: Any) -> None:
    values: list[bytes] = []
    for text in ("First embedded document", "Second embedded document"):
        attachment = Document()
        attachment.add_paragraph(text)
        with BytesIO() as buffer:
            attachment.save(buffer)
            values.append(buffer.getvalue())
    source = tmp_path / "attachments.enex"
    source.write_text(
        "<en-export>"
        + "".join(
            _enex_note(
                title,
                [(content, "photo.docx", "application/vnd.openxmlformats-officedocument.wordprocessingml.document")],
            )
            for title, content in zip(("First", "Second"), values, strict=True)
        )
        + "</en-export>",
        encoding="utf-8",
    )
    bundle, markdown, primary = _bundle(round_trip_runtime, source, tmp_path / "output")
    targets = re.findall(r"\[photo[^\]]*\]\(([^)]+)\)", markdown)
    assert [_target_artifact(bundle, primary, target).sha256 for target in targets] == [
        hashlib.sha256(value).hexdigest() for value in values
    ]


@pytest.mark.parametrize("field", ["category", "series"])
@pytest.mark.parametrize(
    "category",
    [
        "A|B",
        r"A\|B",
        "line one\nline two",
        "[label](target) *value* `tick` <tag>",
        "==mark== ^sup^ ~sub~ $m$ {#id} &amp;",
    ],
)
def test_pptx_chart_cells_are_literal_and_workbook_bytes_remain_exact(
    tmp_path: Path, round_trip_runtime: Any, category: str, field: str
) -> None:
    from pptx import Presentation
    from pptx.chart.data import CategoryChartData
    from pptx.enum.chart import XL_CHART_TYPE
    from pptx.util import Inches

    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    label, series_name = (category, "Revenue") if field == "category" else ("A", category)
    data = CategoryChartData()
    data.categories = [label]
    data.add_series(series_name, (125,))
    frame = slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(1), Inches(1), Inches(5), Inches(3), data)
    expected_workbook = bytes(
        next(rel.target_part.blob for rel in frame.chart.part.rels.values() if rel.reltype.endswith("/package"))
    )
    source = tmp_path / "chart.pptx"
    presentation.save(str(source))
    with ZipFile(source) as package:
        assert [package.read(name) for name in package.namelist() if name.startswith("ppt/embeddings/")] == [
            expected_workbook
        ]
    bundle, markdown, primary = _bundle(round_trip_runtime, source, tmp_path / "output")
    ast = mistune.create_markdown(
        renderer="ast", plugins=["table", "mark", "superscript", "subscript", "strikethrough", "math"]
    )(markdown)
    tables = [node for node in ast if node["type"] == "table"]
    assert len(tables) == 1
    [header, body] = tables[0]["children"]
    assert len(header["children"]) == 2
    [row] = body["children"]
    assert len(row["children"]) == 2

    def text_of(node: dict[str, Any]) -> str:
        if node["type"] == "inline_html":
            assert node["raw"] == "<br>"
            return "\n"
        assert node["type"] in {"table_cell", "text"}, node
        return html.unescape(node.get("raw", "")) + "".join(text_of(child) for child in node.get("children", []))

    assert [text_of(cell) for cell in header["children"]] == ["Category", series_name]
    assert text_of(row["children"][0]) == label
    assert float(text_of(row["children"][1])) == 125
    target = re.search(r"\[Embedded chart workbook\]\(([^)]+)\)", markdown)
    assert target is not None
    workbook = _target_artifact(bundle, primary, target.group(1))
    assert workbook.sha256 == hashlib.sha256(expected_workbook).hexdigest()


def _navigation_targets(markdown: str) -> list[str]:
    ast = mistune.create_markdown(
        renderer="ast", plugins=["table", "mark", "superscript", "subscript", "strikethrough", "math"]
    )(markdown)
    result: list[str] = []

    def walk(nodes: list[dict[str, Any]]) -> None:
        for node in nodes:
            if node["type"] == "link":
                result.append(node["attrs"]["url"])
            walk(node.get("children", []))

    walk(ast)
    return result


def test_enex_ocr_sidecars_and_authored_markdown_share_one_request_name_inventory(
    tmp_path: Path, round_trip_runtime: Any, monkeypatch: pytest.MonkeyPatch
) -> None:
    from docwen_core.text.ocr import OcrOutcome, OcrStatus

    red, blue = _png((210, 10, 10)), _png((10, 10, 210))
    recognized = {
        hashlib.sha256(red).hexdigest(): "OCR for first note",
        hashlib.sha256(blue).hexdigest(): "OCR for second note",
    }

    def recognize(path: str, **_kwargs: Any) -> OcrOutcome:
        # Only the OCR engine is doubled. Every resource write, sidecar link,
        # output relocation and Bundle byte/hash check remains production.
        return OcrOutcome(OcrStatus.SUCCESS, text=recognized[hashlib.sha256(Path(path).read_bytes()).hexdigest()])

    monkeypatch.setattr("docwen_plugin_markup.markdown_resources.run_ocr_outcome", recognize)
    authored = b"# Authored attachment\n\nThis is not OCR output.\n"
    first = _enex_note("First", [(red, "photo.png", "image/png"), (authored, "photo_ocr.md", "text/markdown")])
    second = _enex_note("Second", [(blue, "photo.png", "image/png")])
    source = tmp_path / "ocr-notes.enex"
    source.write_text(f"<en-export>{first}{second}</en-export>", encoding="utf-8")
    output = tmp_path / "output"
    bundle, markdown, primary = _bundle(
        round_trip_runtime,
        source,
        output,
        image_link_style="markdown_embed",
        to_md_enable_ocr=True,
        ocr_placement="image_md",
        config_snapshot={"link": {"format": {"md_file_link_style": "markdown_link"}}},
    )
    first_note, second_note = markdown.split("# First", 1)[1].split("# Second", 1)
    [first_ocr, attachment] = _navigation_targets(first_note)
    [second_ocr] = _navigation_targets(second_note)
    attached = _target_artifact(bundle, primary, attachment)
    assert (output / attached.locator).read_bytes() == authored
    for target, image_bytes, expected_text, other_text in (
        (first_ocr, red, "OCR for first note", "OCR for second note"),
        (second_ocr, blue, "OCR for second note", "OCR for first note"),
    ):
        sidecar = _target_artifact(bundle, primary, target)
        body = (output / sidecar.locator).read_text(encoding="utf-8-sig")
        [(_, image_target)] = _image_targets(body)
        assert _target_artifact(bundle, sidecar, image_target).sha256 == hashlib.sha256(image_bytes).hexdigest()
        assert body.count(expected_text) == 1
        assert other_text not in body
    assert len({first_ocr, second_ocr, attachment}) == 3


@pytest.mark.parametrize(
    "native_cell",
    [
        "A|B",
        r"A\|B",
        "line one\nline two",
        "[label](target) *value* `tick` <tag>",
        "==mark== ^sup^ ~sub~ $m$ {#id} &amp;",
    ],
)
def test_native_pptx_table_header_and_body_keep_literal_cells(
    tmp_path: Path, round_trip_runtime: Any, native_cell: str
) -> None:
    from pptx import Presentation
    from pptx.util import Inches

    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    frame = slide.shapes.add_table(2, 2, Inches(1), Inches(1), Inches(6), Inches(2))
    table = frame.table
    expected = [[native_cell, "**Header**"], [native_cell, "125"]]
    for row_index, row in enumerate(expected):
        for column_index, value in enumerate(row):
            table.cell(row_index, column_index).text = value
    source = tmp_path / "table.pptx"
    presentation.save(str(source))
    bundle, markdown, _ = _bundle(round_trip_runtime, source, tmp_path / "output")
    assert bundle.artifacts
    [table_ast] = [
        node
        for node in mistune.create_markdown(
            renderer="ast", plugins=["table", "mark", "superscript", "subscript", "strikethrough", "math"]
        )(markdown)
        if node["type"] == "table"
    ]
    header, body = table_ast["children"]
    [row] = body["children"]

    def text_of(node: dict[str, Any]) -> str:
        if node["type"] == "inline_html":
            assert node["raw"] == "<br>"
            return "\n"
        assert node["type"] in {"table_cell", "text"}, node
        return html.unescape(node.get("raw", "")) + "".join(text_of(child) for child in node.get("children", []))

    assert [[text_of(cell) for cell in cells["children"]] for cells in (header, row)] == expected
