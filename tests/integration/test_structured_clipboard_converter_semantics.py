"""Final-artifact semantics for structured clipboard converter projections."""

from __future__ import annotations

import csv
from pathlib import Path

import pytest
from docx import Document
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill

from docwen_application.controller import ApplicationController
from docwen_core.clipboard_table_associations import read_clipboard_table_associations
from docwen_core.detection import inspect_structured_clipboard_snapshot
from docwen_core.docx_parsing.document_semantics import extract_semantic_table_metadata
from docwen_core.markdown_extensions import MarkdownExtensions
from docwen_core.models import FILE_INSPECTION_METADATA_KEY
from docwen_core.models.clipboard_document import CLIPBOARD_DOCUMENT_MEDIA_TYPE, clipboard_document_to_bytes
from docwen_core.models.file_ref import (
    MANAGED_INPUT_SHA256_METADATA_KEY,
    MANAGED_INPUT_SIZE_BYTES_METADATA_KEY,
    FileRef,
)
from docwen_core.models.request import ConversionRequest, OutputPolicy
from docwen_gui.clipboard_structured import project_structured_clipboard_html
from docwen_plugin_markdown.mistune_extensions import parse_markdown_text
from docwen_plugin_markdown.plugin import MarkdownPlugin
from docwen_runtime.adapters import RuntimePortAdapter
from docwen_runtime.engine.route_resolver import RouteResolver
from docwen_runtime.engine.task_manager import TaskManager
from docwen_runtime.output.finalizer import OutputFinalizer
from docwen_runtime.plugin_registry.registry import PluginRegistry
from docwen_runtime.templates import TemplateRegistry
from docwen_runtime.templates.state import user_templates_dir
from docwen_runtime.workspace.manager import WorkspaceManager

pytestmark = [pytest.mark.integration, pytest.mark.pr_gate, pytest.mark.release_gate]


HEADER_HTML = b"""
<table>
<tr><th id="a" scope="row">A</th><th id="b" scope="row">B</th></tr>
<tr><td headers="a">1</td><td headers="b">2</td></tr>
</table>
<table>
<tr><th id="corner">Region</th><th id="q1" scope="col">Q1</th></tr>
<tr><th id="north" scope="row">North</th><td headers="north q1">10</td></tr>
</table>
<table><tr><td>plain-a</td><td>plain-b</td></tr><tr><td>3</td><td>4</td></tr></table>
"""

STRUCTURAL_HTML = b"""
<table>
<tr><th id="corner">Region</th><th id="q1" scope="col">Q1</th><th id="q2" scope="col">Q2</th></tr>
<tr><th id="north" scope="row">North</th><td colspan="2">**literal** | &lt; ^ \\ path</td></tr>
</table>
"""


def _controller(tmp_path: Path) -> ApplicationController:
    registry = PluginRegistry()
    registry.register(MarkdownPlugin())
    return ApplicationController(
        runtime_port=RuntimePortAdapter(
            TaskManager(
                registry,
                RouteResolver(registry),
                WorkspaceManager(root_dir=str(tmp_path / "runtime")),
                OutputFinalizer(),
            )
        )
    )


def _request(
    tmp_path: Path,
    *,
    request_id: str,
    target: str,
    html: bytes,
    options: dict | None = None,
) -> ConversionRequest:
    document = project_structured_clipboard_html(html)
    payload = clipboard_document_to_bytes(document)
    source = tmp_path / f"{request_id}.dwclip"
    source.write_bytes(payload)
    inspection = inspect_structured_clipboard_snapshot(str(source))
    return ConversionRequest(
        request_id=request_id,
        input_refs=[
            FileRef(
                path=str(source),
                format="clipboard_document",
                category="markdown",
                size_bytes=len(payload),
                input_kind="document",
                input_role="source",
                logical_path="document.dwclip",
                media_type=CLIPBOARD_DOCUMENT_MEDIA_TYPE,
                metadata={
                    FILE_INSPECTION_METADATA_KEY: inspection.to_dict(),
                    MANAGED_INPUT_SHA256_METADATA_KEY: inspection.content_sha256,
                    MANAGED_INPUT_SIZE_BYTES_METADATA_KEY: inspection.size_bytes,
                },
            )
        ],
        target_format=target,
        options=options or {},
        output_policy=OutputPolicy(output_dir=str(tmp_path / f"out-{request_id}")),
    )


def _template_id(kind: str, filename: str) -> str:
    matches = [item.id for item in TemplateRegistry.default().list_templates(kind) if item.path.name == filename]
    assert len(matches) == 1
    return matches[0]


def _primary(result) -> Path:
    assert result.success, result.error
    return Path(next(item for item in result.artifacts if item.is_primary).staging_path)


def test_docx_preserves_native_header_shape_and_reversible_html_associations(tmp_path: Path) -> None:
    result = _controller(tmp_path).execute_single(
        _request(
            tmp_path,
            request_id="header-docx",
            target="docx",
            html=HEADER_HTML,
            options={"template_name": _template_id("docx", "English General Template.docx")},
        )
    )
    output = _primary(result)
    document = Document(output)
    assert len(document.tables) == 3

    first = extract_semantic_table_metadata(document.tables[0]._tbl, default_first_row=False)
    second = extract_semantic_table_metadata(document.tables[1]._tbl, default_first_row=False)
    third = extract_semantic_table_metadata(document.tables[2]._tbl, default_first_row=False)
    assert (first.header_rows, first.header_columns) == (0, 0)
    assert (second.header_rows, second.header_columns) == (1, 1)
    assert (third.header_rows, third.header_columns) == (0, 0)
    assert "tblHeader" not in document.tables[0]._tbl.xml
    assert "tblHeader" in document.tables[1]._tbl.xml
    assert "tblHeader" not in document.tables[2]._tbl.xml
    assert 'w:firstRow="0"' in document.tables[0]._tbl.xml
    assert 'w:firstRow="0"' in document.tables[2]._tbl.xml

    associations = read_clipboard_table_associations(output)
    assert [(item.header_rows, item.header_columns) for item in associations] == [(0, 0), (1, 1), (0, 0)]
    first_cells = associations[0].cells
    assert (first_cells[0].header, first_cells[0].scope, first_cells[0].html_id) == (True, "row", "a")
    assert (first_cells[1].header, first_cells[1].scope, first_cells[1].html_id) == (True, "row", "b")
    assert first_cells[2].headers == ("a",)
    assert first_cells[3].headers == ("b",)
    second_cells = associations[1].cells
    assert second_cells[0].html_id == "corner"
    assert second_cells[1].scope == "col"
    assert second_cells[2].scope == "row"
    assert second_cells[3].headers == ("north", "q1")
    assert all(not cell.header and not cell.scope and not cell.headers for cell in associations[2].cells)


def test_xlsx_and_csv_publish_visible_header_association_indices(tmp_path: Path) -> None:
    controller = _controller(tmp_path)
    xlsx = controller.execute_single(
        _request(
            tmp_path,
            request_id="header-xlsx",
            target="xlsx",
            html=HEADER_HTML,
            options={"template_name": _template_id("xlsx", "English Sample Sheet Template.xlsx")},
        )
    )
    workbook = load_workbook(_primary(xlsx))
    assert "Table Semantics" in workbook.sheetnames
    semantics = workbook["Table Semantics"]
    headers = [cell.value for cell in semantics[1]]
    rows = [dict(zip(headers, [cell.value for cell in row], strict=True)) for row in semantics.iter_rows(min_row=2)]
    assert any(
        row["Table"] == "T1"
        and row["Anchor"] == "R1C1"
        and row["Role"] == "associated_header"
        and row["Scope"] == "row"
        and row["HTML ID"] == "a"
        for row in rows
    )
    assert any(row["Table"] == "T1" and row["Anchor"] == "R2C1" and row["Headers"] == "a" for row in rows)
    assert any(row["Table"] == "T2" and row["Anchor"] == "R1C1" and row["Role"] == "corner_header" for row in rows)
    assert all(row["Role"] == "data" for row in rows if row["Table"] == "T3")
    assert "CLIPBOARD-XLSX-HEADER-ASSOCIATIONS-PROJECTED" in {item.code for item in xlsx.diagnostics}

    csv_result = controller.execute_single(_request(tmp_path, request_id="header-csv", target="csv", html=HEADER_HTML))
    assert csv_result.success, csv_result.error
    semantics_artifact = next(
        item for item in csv_result.artifacts if item.suggested_name.endswith("-table-semantics.csv")
    )
    with Path(semantics_artifact.staging_path).open(encoding="utf-8-sig", newline="") as stream:
        csv_rows = list(csv.DictReader(stream))
    assert any(row["Table"] == "T1" and row["Scope"] == "row" and row["HTML ID"] == "a" for row in csv_rows)
    assert any(row["Table"] == "T2" and row["Headers"] == "north q1" for row in csv_rows)


def _install_xlsx_template(
    name: str,
    *,
    bold: bool,
    width: float,
    height: float,
    orientation: str,
    header_text: str,
    show_gridlines: bool,
) -> str:
    directory = user_templates_dir()
    directory.mkdir(parents=True, exist_ok=True)
    path = directory / f"{name}.xlsx"
    workbook = Workbook()
    prototype = workbook.active
    assert prototype is not None
    prototype.title = "Prototype"
    prototype["A1"] = "template-value"
    prototype["A1"].font = Font(bold=bold, italic=not bold)
    prototype["A1"].fill = PatternFill(fill_type="solid", fgColor="00FF00" if bold else "FFFF00")
    prototype.column_dimensions["A"].width = width
    prototype.row_dimensions[1].height = height
    prototype.page_setup.orientation = orientation
    prototype.print_title_rows = "1:1"
    prototype.print_title_cols = "A:A"
    prototype.print_area = "A1:C8"
    prototype.oddHeader.center.text = header_text
    prototype.freeze_panes = "B2"
    prototype.sheet_view.showGridLines = show_gridlines
    prototype.sheet_view.zoomScale = 120 if bold else 90
    lookup = workbook.create_sheet("Lookup")
    lookup["A1"] = f"lookup-{name}"
    reserved = workbook.create_sheet("Table 1")
    reserved["A1"] = f"reserved-{name}"
    workbook.save(path)
    return _template_id("xlsx", path.name)


@pytest.mark.parametrize(
    ("name", "bold", "width", "height", "orientation", "header_text", "show_gridlines"),
    [
        ("structured-prototype-a", True, 27.0, 33.0, "landscape", "HEADER-A", False),
        ("structured-prototype-b", False, 41.0, 47.0, "portrait", "HEADER-B", True),
    ],
)
def test_xlsx_table_sheets_really_derive_selected_template_prototype(
    tmp_path: Path,
    name: str,
    bold: bool,
    width: float,
    height: float,
    orientation: str,
    header_text: str,
    show_gridlines: bool,
) -> None:
    template_id = _install_xlsx_template(
        name,
        bold=bold,
        width=width,
        height=height,
        orientation=orientation,
        header_text=header_text,
        show_gridlines=show_gridlines,
    )
    html = b"<table><tr><th>A</th><th>B</th></tr><tr><td>1</td><td>2</td></tr></table>"
    result = _controller(tmp_path).execute_single(
        _request(
            tmp_path,
            request_id=f"template-{name}",
            target="xlsx",
            html=html,
            options={"template_name": template_id},
        )
    )
    workbook = load_workbook(_primary(result))
    generated = next(title for title in workbook.sheetnames if title.startswith("Table 1 ") and title != "Table 1")
    sheet = workbook[generated]
    assert sheet["A1"].value == "A"
    assert sheet["A1"].font.bold is bold
    assert sheet["A1"].font.italic is (not bold)
    assert sheet.column_dimensions["A"].width == width
    assert sheet.row_dimensions[1].height == height
    assert sheet.page_setup.orientation == orientation
    assert sheet.print_title_rows == "$1:$1"
    assert sheet.print_title_cols == "$A:$A"
    assert f"'{generated}'!$A$1:$C$8" == sheet.print_area
    assert sheet.oddHeader.center.text == header_text
    assert sheet.freeze_panes == "B2"
    assert sheet.sheet_view.showGridLines is show_gridlines
    assert sheet.sheet_view.zoomScale == (120 if bold else 90)
    assert workbook["Lookup"]["A1"].value == f"lookup-{name}"
    assert workbook["Table 1"]["A1"].value == f"reserved-{name}"
    assert workbook["Prototype"]["A1"].value == "template-value"

    order = workbook["Document Order"]
    order_headers = [cell.value for cell in order[1]]
    order_rows = [
        dict(zip(order_headers, [cell.value for cell in row], strict=True)) for row in order.iter_rows(min_row=2)
    ]
    assert any(row["Kind"] == "table_ref" and row["Sheet"] == generated for row in order_rows)


def test_markdown_output_extension_changes_real_projection_without_reinterpreting_values(tmp_path: Path) -> None:
    controller = _controller(tmp_path)
    enabled = controller.execute_single(
        _request(
            tmp_path,
            request_id="md-structural-on",
            target="md",
            html=STRUCTURAL_HTML,
            options={"markdown_extensions": {"output": {"structural_tables": True}}},
        )
    )
    disabled = controller.execute_single(
        _request(
            tmp_path,
            request_id="md-structural-off",
            target="md",
            html=STRUCTURAL_HTML,
            options={"markdown_extensions": {"output": {"structural_tables": False}}},
        )
    )
    enabled_text = _primary(enabled).read_text(encoding="utf-8")
    disabled_text = _primary(disabled).read_text(encoding="utf-8")
    assert enabled_text != disabled_text
    assert "| --- || --- | --- |" in enabled_text
    assert r"\*\*literal\*\*" in enabled_text
    assert r"\|" in enabled_text
    assert r"\<" in enabled_text
    assert r"\^" in enabled_text
    assert "role=row_header" in enabled_text
    assert "scope=row" in enabled_text
    assert "id=north" in enabled_text
    assert "headers=-" in enabled_text
    assert "CLIPBOARD-MARKDOWN-STRUCTURAL-FALLBACK" not in {item.code for item in enabled.diagnostics}
    assert "CLIPBOARD-MARKDOWN-HEADER-ASSOCIATIONS-PROJECTED" in {item.code for item in enabled.diagnostics}
    assert "CLIPBOARD-MARKDOWN-STRUCTURE-PROJECTION" in {item.code for item in disabled.diagnostics}
    assert "| --- || --- | --- |" not in disabled_text
    assert "**literal** | < ^ \\ path" in disabled_text


def test_structural_markdown_falls_back_for_nested_or_headerless_tables(tmp_path: Path) -> None:
    html = (
        b"<table><tr><td>no-head</td><td>value</td></tr></table>"
        b"<table><tr><th>A</th></tr><tr><td><p>before</p>"
        b"<table><tr><td>nested</td></tr></table><p>after</p></td></tr></table>"
    )
    result = _controller(tmp_path).execute_single(
        _request(
            tmp_path,
            request_id="md-structural-fallback",
            target="md",
            html=html,
            options={"markdown_extensions": {"output": {"structural_tables": True}}},
        )
    )
    text = _primary(result).read_text(encoding="utf-8")
    assert "no-head" in text
    assert "before" in text and "nested" in text and "after" in text
    assert "header_rows=0" in text
    assert "CLIPBOARD-MARKDOWN-STRUCTURAL-FALLBACK" in {item.code for item in result.diagnostics}


def _ast_descendants(nodes):
    for node in nodes:
        yield node
        yield from _ast_descendants(node.get("children", []))


@pytest.mark.parametrize("header_rows", [1, 2])
def test_structural_markdown_final_file_reparses_authored_punctuation_as_text(tmp_path: Path, header_rows: int) -> None:
    literals = (
        "$A^2$",
        "$P(A|B)$",
        "==mark==",
        "`code`",
        "**bold**",
        "<b>raw</b>",
        r"pipe|slash\literal",
        "<",
        "^",
    )
    header = "".join(f"<th>H{index}</th>" for index in range(len(literals)))
    body = "".join(
        f"<td>{value.replace('&', '&amp;').replace('<', '&lt;').replace('>', '&gt;')}</td>" for value in literals
    )
    html = ("<table>" + f"<tr>{header}</tr>" * header_rows + f"<tr>{body}</tr></table>").encode()
    result = _controller(tmp_path).execute_single(
        _request(
            tmp_path,
            request_id="md-literal-reparse",
            target="md",
            html=html,
            options={"markdown_extensions": {"output": {"structural_tables": True}}},
        )
    )
    output = _primary(result)
    markdown = output.read_text(encoding="utf-8")
    ast = parse_markdown_text(markdown, extensions=MarkdownExtensions(structural_tables=True))
    table = next(node for node in ast if node.get("type") == "table")
    descendants = tuple(_ast_descendants([table]))
    forbidden = {"inline_math", "mark", "highlight", "codespan", "emphasis", "strong", "inline_html"}
    assert not any(node.get("type") in forbidden for node in descendants)
    visible_text = "".join(str(node.get("raw", "")) for node in descendants if node.get("type") == "text")
    for literal in literals:
        assert literal in visible_text
    assert ("_structural_table" in table) is (header_rows > 1)
    authored_cells = table["children"][-1]["children"][-1]["children"]
    assert authored_cells[-2]["attrs"]["docwen_literal_merge_marker"] is True
    assert authored_cells[-1]["attrs"]["docwen_literal_merge_marker"] is True
