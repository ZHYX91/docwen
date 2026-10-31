"""Generated native spreadsheet files retain final cell and worksheet identity."""

from __future__ import annotations

import csv
import hashlib
import html
from pathlib import Path
from typing import Any
from zipfile import ZipFile

import mistune
import openpyxl
import pytest
from lxml import etree
from openpyxl.chart import BarChart
from openpyxl.chart.data_source import NumData, NumDataSource, NumVal
from openpyxl.chart.series import Series

from docwen_application.bundle_mapping import BundleProfile, build_bundle_draft
from docwen_core.detection import inspect_file
from docwen_core.models.file_ref import FileRef
from docwen_core.models.request import ConversionRequest, OutputPolicy
from docwen_runtime.output.artifact_bundle import ArtifactBundleCommitter

pytestmark = [pytest.mark.integration, pytest.mark.pr_gate]

_S = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
_R = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
_MEDIA = {
    "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    "md": "text/markdown",
    "csv": "text/csv",
    "tsv": "text/tab-separated-values",
}


def _request(runtime: Any, source: Path, output: Path, target: str, **options: Any) -> Any:
    inspection = inspect_file(str(source))
    assert inspection.detected_format == source.suffix.removeprefix(".")
    original = source.read_bytes()
    result = runtime.execute(
        ConversionRequest(
            request_id="spreadsheet-final-content",
            input_refs=[FileRef(str(source), inspection.detected_format, "spreadsheet", len(original))],
            target_format=target,
            output_policy=OutputPolicy(output_dir=str(output)),
            options={"to_md_keep_images": False, "to_md_enable_ocr": False, **options} if target == "md" else options,
        )
    )
    assert source.read_bytes() == original
    return result


def _publish(runtime: Any, source: Path, output: Path, target: str, **options: Any) -> tuple[Any, Any]:
    result = _request(runtime, source, output, target, **options)
    assert result.success, result.error
    profile: BundleProfile = "worksheet_resources" if target in {"csv", "tsv"} else "single_document"
    if target == "md" and options.get("to_md_keep_images"):
        profile = "document_with_resources"
    draft = build_bundle_draft(profile=profile, output_media_type=_MEDIA[target], artifacts=result.artifacts)
    bundle = ArtifactBundleCommitter().commit(
        task_id="spreadsheet-final-content", staging_root=str(output), draft=draft
    )
    for artifact in bundle.artifacts:
        content = (output / artifact.locator).read_bytes()
        assert artifact.size_bytes == len(content)
        assert artifact.sha256 == hashlib.sha256(content).hexdigest()
    return result, bundle


def _primary(bundle: Any, output: Path) -> Path:
    preferred = [entry for entry in bundle.entries if entry.preferred]
    assert len(preferred) == 1
    artifact = next(item for item in bundle.artifacts if item.artifact_id == preferred[0].artifact_id)
    return output / artifact.locator


def _plain(nodes: list[dict[str, Any]]) -> str:
    parts: list[str] = []
    for node in nodes:
        if node["type"] in {"text", "codespan"}:
            parts.append(html.unescape(node.get("raw", "")))
        elif node["type"] == "inline_html" and node.get("raw", "").lower() in {"<br>", "<br/>", "<br />"}:
            parts.append("\n")
        elif "children" in node:
            parts.append(_plain(node["children"]))
    return "".join(parts)


def _tables(markdown: str) -> list[list[list[str]]]:
    # Independent Markdown consumer: no DocWen parser or serializer is used.
    ast = mistune.create_markdown(renderer="ast", plugins=["table"])(markdown)
    assert isinstance(ast, list)
    result: list[list[list[str]]] = []
    for node in ast:
        if node["type"] != "table":
            continue
        rows: list[list[str]] = []
        for section in node["children"]:
            if section["type"] == "table_head":
                rows.append([_plain(cell["children"]) for cell in section["children"]])
            else:
                for row in section["children"]:
                    rows.append([_plain(cell["children"]) for cell in row["children"]])
        result.append(rows)
    return result


def _saved_values(path: Path) -> list[list[Any]]:
    workbook = openpyxl.load_workbook(path)
    try:
        assert len(workbook.worksheets) == 1
        worksheet = workbook.worksheets[0]
        rows = list(worksheet.iter_rows())
        assert all(cell.data_type == "s" for row in rows for cell in row)
        return [[cell.value for cell in row] for row in rows]
    finally:
        workbook.close()


@pytest.mark.parametrize("encoding", ["utf-8", "utf-8-sig", "utf-16", "gbk"])
@pytest.mark.parametrize("target", ["xlsx", "md"])
def test_semicolon_content_uses_same_two_column_fact_at_final_output(
    tmp_path: Path, round_trip_runtime: Any, encoding: str, target: str
) -> None:
    source = tmp_path / "semicolon.csv"
    source.write_bytes("Name;Amount\nAlice;7\nBob;9\n".encode(encoding))
    with source.open(encoding=encoding, newline="") as handle:
        assert list(csv.reader(handle, delimiter=";")) == [["Name", "Amount"], ["Alice", "7"], ["Bob", "9"]]
    output = tmp_path / "published"
    _, bundle = _publish(round_trip_runtime, source, output, target)
    primary = _primary(bundle, output)
    expected = [["Name", "Amount"], ["Alice", "7"], ["Bob", "9"]]
    assert (
        _saved_values(primary) if target == "xlsx" else _tables(primary.read_text(encoding="utf-8-sig"))[0]
    ) == expected


@pytest.mark.parametrize("target", ["xlsx", "md"])
def test_delimiter_fact_survives_full_file_utf8_to_gbk_retry(
    tmp_path: Path, round_trip_runtime: Any, target: str
) -> None:
    expected = [["Name", "Amount"], *[[f"record-{index:05d}", "001"] for index in range(5000)], ["末行", "0007"]]
    source = tmp_path / "late-gbk.csv"
    source.write_bytes(("\n".join(";".join(row) for row in expected) + "\n").encode("gbk"))
    assert len(source.read_bytes().split("末行".encode("gbk"))[0]) > 65536
    output = tmp_path / "published"
    _, bundle = _publish(round_trip_runtime, source, output, target)
    primary = _primary(bundle, output)
    actual = _saved_values(primary) if target == "xlsx" else _tables(primary.read_text(encoding="utf-8-sig"))[0]
    assert len(actual) == len(expected)
    for row_index, (actual_row, expected_row) in enumerate(zip(actual, expected, strict=True)):
        assert len(actual_row) == len(expected_row), (row_index, len(actual_row), len(expected_row))
        for column, (actual_cell, expected_cell) in enumerate(zip(actual_row, expected_row, strict=True)):
            matches = actual_cell == expected_cell
            assert matches, (
                row_index,
                column,
                len(actual_cell),
                len(expected_cell),
                hashlib.sha256(actual_cell.encode()).hexdigest(),
                hashlib.sha256(expected_cell.encode()).hexdigest(),
            )


def _chartsheet_workbook(path: Path, position: int, *, only_chart: bool = False) -> None:
    workbook = openpyxl.Workbook()
    data = workbook.active
    assert data is not None
    data.title = "Data"
    data.append(["Name", "Value"])
    data.append(["zero", 0])
    data.append(["false", False])
    if not only_chart:
        second = workbook.create_sheet("Later")
        second.append(["Tail", "Number"])
        second.append(["last", 3])
    chart = BarChart()
    chart.series = [Series(idx=0, order=0, val=NumDataSource(numLit=NumData(ptCount=1, pt=[NumVal(idx=0, v=1)])))]
    chart_sheet = workbook.create_chartsheet("Chart", position)
    chart_sheet.add_chart(chart)
    if only_chart:
        workbook.remove(data)
    workbook.save(path)
    workbook.close()
    reopened = openpyxl.load_workbook(path)
    try:
        assert [sheet.title for sheet in reopened.worksheets] == ([] if only_chart else ["Data", "Later"])
        assert [sheet.title for sheet in reopened.chartsheets] == ["Chart"]
    finally:
        reopened.close()
    with ZipFile(path) as package:
        relations = etree.fromstring(package.read("xl/_rels/workbook.xml.rels"))
        assert sum(item.get("Type", "").endswith("/chartsheet") for item in relations) == 1
        assert sum(item.get("Type", "").endswith("/worksheet") for item in relations) == (0 if only_chart else 2)


@pytest.mark.parametrize("position", [0, 1, 2], ids=["first", "middle", "last"])
@pytest.mark.parametrize("target", ["md", "csv", "tsv"])
def test_typed_worksheets_own_order_primary_and_content_across_chart_sheets(
    tmp_path: Path, round_trip_runtime: Any, position: int, target: str
) -> None:
    source = tmp_path / "chart-position.xlsx"
    _chartsheet_workbook(source, position)
    output = tmp_path / "published"
    result, bundle = _publish(round_trip_runtime, source, output, target)
    expected = [[["Name", "Value"], ["zero", "0"], ["false", "False"]], [["Tail", "Number"], ["last", "3"]]]
    if target == "md":
        markdown = _primary(bundle, output).read_text(encoding="utf-8-sig")
        assert _tables(markdown) == expected
        ast = mistune.create_markdown(renderer="ast")(markdown)
        assert isinstance(ast, list)
        assert [_plain(node["children"]) for node in ast if node["type"] == "heading"] == ["Data", "Later"]
    else:
        ordered = sorted(bundle.entries, key=lambda entry: entry.ordinal)
        assert [entry.ordinal for entry in ordered] == [0, 1]
        assert [entry.preferred for entry in ordered] == [True, False]
        actual = []
        for entry in ordered:
            artifact = next(item for item in bundle.artifacts if item.artifact_id == entry.artifact_id)
            with (output / artifact.locator).open(encoding="utf-8-sig", newline="") as handle:
                actual.append(list(csv.reader(handle, delimiter="\t" if target == "tsv" else ",")))
        assert actual == expected
        assert [artifact.metadata["sheet_index"] for artifact in result.artifacts] == [0, 1]
        assert [artifact.metadata["sheet_name"] for artifact in result.artifacts] == ["Data", "Later"]


@pytest.mark.parametrize("target", ["md", "csv", "tsv"])
def test_only_chart_workbook_is_an_explicit_failure_without_final_output(
    tmp_path: Path, round_trip_runtime: Any, target: str
) -> None:
    source = tmp_path / "only-chart.xlsx"
    _chartsheet_workbook(source, 0, only_chart=True)
    output = tmp_path / "published"
    result = _request(round_trip_runtime, source, output, target)
    assert not result.success
    assert result.error is not None
    assert result.error.diagnostic_code == "SPREADSHEET-NO-WORKSHEETS"
    assert result.artifacts == []
    assert not output.exists() or not list(output.rglob("*"))


@pytest.mark.parametrize("source_format", ["csv", "xlsx"])
def test_native_cell_punctuation_is_literal_to_an_independent_table_consumer(
    tmp_path: Path, round_trip_runtime: Any, source_format: str
) -> None:
    expected = [
        ["Author|Text", "Other"],
        [r"a\|b", "second"],
        [r"**literal** `code` [x](y) & <tag>", "third"],
        ["line\nbreak", "fourth"],
    ]
    if source_format == "csv":
        # The content-first ingress deliberately treats a Markdown link marker
        # as Markdown even under .csv. Exercise that raw literal in native XLSX,
        # and keep this CSV fixture within the documented delimited admission.
        expected[2][0] = r"**literal** `code` & <tag>"
    source = tmp_path / f"literal.{source_format}"
    if source_format == "csv":
        with source.open("w", encoding="utf-8", newline="") as handle:
            csv.writer(handle).writerows(expected)
    else:
        workbook = openpyxl.Workbook()
        worksheet = workbook.active
        assert worksheet is not None
        for row in expected:
            worksheet.append(row)
        workbook.save(source)
        workbook.close()
        assert _saved_values(source) == expected
    output = tmp_path / "published"
    _, bundle = _publish(round_trip_runtime, source, output, "md")
    assert _tables(_primary(bundle, output).read_text(encoding="utf-8-sig")) == [expected]


@pytest.mark.parametrize("shape", ["distant", "ring-and-island"])
def test_sparse_workbook_emits_each_native_value_once_in_row_major_blocks(
    tmp_path: Path, round_trip_runtime: Any, monkeypatch: pytest.MonkeyPatch, shape: str
) -> None:
    source = tmp_path / "sparse.xlsx"
    workbook = openpyxl.Workbook()
    worksheet = workbook.active
    assert worksheet is not None
    expected: list[list[list[str]]]
    if shape == "distant":
        worksheet.cell(250, 250, "ONLY")
        expected = [[["ONLY"]]]
    else:
        for row, values in enumerate(
            [
                ["A", "B", "C", "D", "E"],
                ["F", None, None, None, "G"],
                ["H", None, "ISLAND", None, "I"],
                ["J", None, None, None, "K"],
                ["L", "M", "N", "O", "P"],
            ],
            start=201,
        ):
            for column, value in enumerate(values, start=201):
                if value is not None:
                    worksheet.cell(row, column, value)
        worksheet.cell(250, 250, 0)
        worksheet.cell(252, 252, False)
        expected = [
            [
                ["A", "B", "C", "D", "E"],
                ["F", "", "", "", "G"],
                ["H", "", "", "", "I"],
                ["J", "", "", "", "K"],
                ["L", "M", "N", "O", "P"],
            ],
            [["ISLAND"]],
            [["0"]],
            [["False"]],
        ]
    workbook.save(source)
    workbook.close()
    with ZipFile(source) as package:
        xml = etree.fromstring(package.read("xl/worksheets/sheet1.xml"))
        source_positions = [cell.get("r") for cell in xml.iter(f"{{{_S}}}c")]
        if shape == "distant":
            assert source_positions == ["IP250"]
    loaded: list[Any] = []
    original_load = openpyxl.load_workbook

    def capture(*args: Any, **kwargs: Any) -> Any:
        result = original_load(*args, **kwargs)
        loaded.append(result)
        return result

    monkeypatch.setattr(openpyxl, "load_workbook", capture)
    output = tmp_path / "published"
    _, bundle = _publish(round_trip_runtime, source, output, "md")
    assert _tables(_primary(bundle, output).read_text(encoding="utf-8-sig")) == expected
    # This supplements the final-content oracle with the actual loaded object:
    # distant blank coordinates must not be materialized by traversal.
    assert loaded
    for view in loaded:
        assert len(view.worksheets[0]._cells) == len(source_positions)
