"""converter 单元测试。"""

from __future__ import annotations

from pathlib import Path

import pytest
from docx import Document

from docwen.config.config_manager import config_manager
from docwen.converter.docx2md.shared.table_processor import convert_table_to_md_with_images

pytestmark = pytest.mark.unit


def test_convert_table_to_md_with_images_horizontal_merged_cell_outputs_blank_for_repeat(
    tmp_path: Path,
) -> None:
    doc = Document()
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "H"
    table.cell(0, 1).text = ""
    table.cell(0, 0).merge(table.cell(0, 1))

    table.cell(1, 0).text = "A"
    table.cell(1, 1).text = "B"

    md = convert_table_to_md_with_images(
        table=table,
        table_index=0,
        images_info=[],
        output_folder=str(tmp_path),
        parent_file_name="in.docx",
        config_manager=config_manager,
        options={"keep_images": False, "enable_ocr": False},
    )

    lines = md.splitlines()
    assert lines[0] == "| H |  |"
    assert lines[1] == "| --- | --- |"
    assert lines[2] == "| A | B |"
