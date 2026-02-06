from __future__ import annotations

import openpyxl
from openpyxl.drawing.image import Image as XlsxImage
from pathlib import Path

import pytest
from PIL import Image

from docwen.converter.formats.image.core import convert_image
from docwen.converter.md2docx.handlers import formula_handler as fh
from docwen.converter.md2docx.processors.md_processor import process_md_body
from docwen.converter.md2docx.processors.xml_processor import process_xml_file
from docwen.converter.xlsx2md.image_processor import extract_images_from_xlsx


pytestmark = pytest.mark.unit


@pytest.mark.slow
def test_convert_image_rgba_to_jpeg(tmp_path: Path) -> None:
    src = tmp_path / "in.png"
    Image.new("RGBA", (2, 2), (255, 0, 0, 128)).save(src)

    out = tmp_path / "out.jpg"
    convert_image(str(src), "jpeg", str(out), {"compress_mode": "lossless"})
    assert out.exists() is True

    with Image.open(out) as img:
        assert img.mode == "RGB"


@pytest.mark.slow
def test_extract_images_from_xlsx_extracts_one(tmp_path: Path) -> None:
    png_path = tmp_path / "img.png"
    Image.new("RGB", (1, 1), (0, 255, 0)).save(png_path)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "S1"
    ws.add_image(XlsxImage(str(png_path)), "A1")

    xlsx_path = tmp_path / "book.xlsx"
    wb.save(xlsx_path)

    loaded = openpyxl.load_workbook(xlsx_path, data_only=True)
    out_dir = tmp_path / "images"
    images = extract_images_from_xlsx(loaded, str(out_dir), str(xlsx_path))
    assert len(images) == 1
    assert Path(images[0]["image_path"]).exists() is True


def test_formula_handler_placeholders_roundtrip() -> None:
    text = "a $E=mc^2$ b\n\n$$\nx\n$$"
    assert fh.has_latex_formulas(text) is True

    replaced, formulas = fh.replace_formulas_with_placeholders(text)
    assert "{{FORMULA_0}}" in replaced
    restored = fh.restore_formulas_from_placeholders(replaced, formulas)
    assert "E=mc^2" in restored


def test_formula_handler_convert_block_formulas_to_paragraphs() -> None:
    lines = ["x", "$$", "a", "$$", "y"]
    out = fh.convert_block_formulas_to_paragraphs(None, lines)
    assert out == ["x", "$$\na\n$$", "y"]


def test_md_processor_process_md_body_normalizes_crlf() -> None:
    md = "# 标题\r\n正文\r\n- 列表1\r\n"
    items = process_md_body(md)
    assert items
    assert all("text" in it and "type" in it and "level" in it for it in items)


def test_xml_processor_process_xml_file_minimal(tmp_path: Path) -> None:
    xml_path = tmp_path / "doc.xml"
    xml_path.write_text(
        '<?xml version="1.0" encoding="UTF-8"?>'
        '<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
        "<w:body><w:p><w:r><w:t>hello</w:t></w:r></w:p></w:body>"
        "</w:document>",
        encoding="utf-8",
    )

    process_xml_file(str(xml_path), {})
    assert "hello" in xml_path.read_text(encoding="utf-8")
