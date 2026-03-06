"""converter 单元测试。"""

from __future__ import annotations

import re
from pathlib import Path

import openpyxl
import pytest

from docwen.converter.md2xlsx.placeholder_processor import process_image_placeholders
from docwen.utils.link_processing import process_markdown_links

pytestmark = pytest.mark.unit


_BASE64_PNG_1X1 = "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVR42mP8/5+hHgAHggJ/PchI7wAAAABJRU5ErkJggg=="


def _extract_image_path(result: str) -> str:
    m = re.search(r"\{\{IMAGE:([^}]+)\}\}", result)
    assert m, result
    payload = m.group(1)
    return payload.split(r"\|", 1)[0]


def test_xlsx_inserts_image_from_data_uri_placeholder(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr("docwen.utils.link_processing.config_manager.get_max_embed_depth", lambda: 10, raising=True)
    monkeypatch.setattr(
        "docwen.utils.link_processing.config_manager.get_markdown_embed_image_mode",
        lambda: "embed",
        raising=True,
    )
    monkeypatch.setattr("docwen.utils.link_processing.config_manager.get_wiki_link_mode", lambda: "keep", raising=True)
    monkeypatch.setattr(
        "docwen.utils.link_processing.config_manager.get_markdown_link_mode", lambda: "keep", raising=True
    )

    md = f"![img](data:image/png;base64,{_BASE64_PNG_1X1} =300x200)"
    source = tmp_path / "src.md"
    source.write_text(md, encoding="utf-8")

    processed = process_markdown_links(md, str(source), temp_dir=str(tmp_path))
    image_path = _extract_image_path(processed)
    assert Path(image_path).exists() is True

    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"] = processed

    inserted = process_image_placeholders(wb)
    assert inserted == 1
    assert len(ws._images) == 1
