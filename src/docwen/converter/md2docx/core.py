"""
MD转DOCX转换器核心模块

统一的临时目录管理和转换流程
"""

import os
import re
import datetime
import logging
import shutil
import tempfile
import threading
from typing import Callable, Optional

from docx import Document as open_docx
from docx.document import Document as DocxDocument

from .processors import md_processor
from .processors import docx_processor
from .style.injector import ensure_styles
from .handlers.notes_part_handler import ensure_notes_parts
from .handlers.numbering_handler import ensure_numbering_part
from docwen.template.loader import TemplateLoader
from docwen.utils.validation_utils import is_value_empty
from docwen.utils.heading_utils import convert_to_halfwidth
from docwen.utils.workspace_manager import prepare_input_file, should_save_intermediate_files
from docwen.utils.path_utils import generate_output_path
from docwen.utils import yaml_utils
from docwen.i18n import t, t_all_locales

logger = logging.getLogger(__name__)

# 附件序号清理正则
ATTACH_NUM_PATTERN = re.compile(
    r'^[一二三四五六七八九十㈠㈡㈢㈣㈤㈥㈦㈧㈨㈩]+、|'  # 中文序号（含带括号中文数字）
    r'^（[一二三四五六七八九十]+）|'  # 带括号中文序号
    r'^\d+[\.．]\s*|'  # 半角数字 + 点
    r'^[０１２３４５６７８９]+[\.．]\s*|'  # 全角数字 + 点
    r'^（\d+）\s*|'  # 带括号半角数字
    r'^（[０１２３４５６７８９]+）\s*|'  # 带括号全角数字
    r'^[①②③④⑤⑥⑦⑧⑨⑩⑪⑫⑬⑭⑮⑯⑰⑱⑲⑳㉑㉒㉓㉔㉕㉖㉗㉘㉙㉚㉛㉜㉝㉞㉟㊱㊲㊳㊴㊵㊶㊷㊸㊹㊺㊻㊼㊽㊾㊿⓪⓵⓶⓷⓸⓹⓺⓻⓼⓽⓾⓿❶❷❸❹❺❻❼❽❾❿⓫⓬⓭⓮⓯⓰⓱⓲⓳⓴]\s*'  # 带圈数字
)


def convert(
    md_path: str,
    output_path: str,
    *,
    template_name: str,
    progress_callback: Optional[Callable[[str], None]] = None,
    cancel_event: Optional[threading.Event] = None,
    original_source_path: Optional[str] = None,
    options: Optional[dict] = None
) -> Optional[str]:
    """
    将Markdown文件转换为DOCX文档
    
    统一的临时目录管理流程：
    1. 创建临时目录
    2. 准备input.md副本
    3. 对副本进行转换
    4. 保存到指定输出路径
    5. 自动清理临时目录
    
    参数:
        md_path: Markdown文件路径
        output_path: 输出文件完整路径
        template_name: 模板名称（必需）
        progress_callback: 进度回调函数
        cancel_event: 取消事件
        original_source_path: 原始源文件路径（用于嵌入功能的路径解析，可选）
        options: 转换选项字典，包含序号配置等
    
    返回:
        str: 输出文件完整路径，失败时返回None
    """
    try:
        logger.info("=" * 60)
        logger.info(f"开始MD转DOCX: {os.path.basename(md_path)}")
        logger.info(f"使用模板: {template_name}")
        logger.info("=" * 60)
        
        with tempfile.TemporaryDirectory() as temp_dir:
            # 1. 创建input.md副本
            temp_input = prepare_input_file(md_path, temp_dir, 'md')
            logger.debug(f"已创建输入副本: {os.path.basename(temp_input)}")
            
            # 2. 读取和解析MD文件（传递原始路径用于路径解析）
            if progress_callback:
                progress_callback(t('conversion.progress.parsing_markdown'))
            
            # 如果提供了original_source_path，使用它；否则使用md_path
            path_for_resolve = original_source_path if original_source_path else md_path
            yaml_data, md_body = _read_and_parse_md(temp_input, path_for_resolve)
            
            if cancel_event and cancel_event.is_set():
                logger.info("操作已取消")
                return None
            
            # 3. 在临时目录生成临时输出文件
            temp_output = os.path.join(temp_dir, "temp_output.docx")
            
            # 4. 在临时目录转换（传递原始文件路径用于标题提取和options）
            success = _convert_internal(
                yaml_data, md_body, temp_output,
                template_name,
                progress_callback, cancel_event,
                original_source_path if original_source_path else md_path,
                options
            )
            
            if not success:
                logger.error("内部转换失败")
                return None
            
            if cancel_event and cancel_event.is_set():
                logger.info("操作已取消")
                return None
            
            # 5. 确保输出目录存在
            output_dir = os.path.dirname(output_path)
            if output_dir:
                os.makedirs(output_dir, exist_ok=True)
            
            # 6. 移动到最终输出路径
            shutil.move(temp_output, output_path)
            
            # 8. 保留中间文件（如需要）
            if should_save_intermediate_files():
                logger.debug("保留中间文件到输出目录")
                for filename in os.listdir(temp_dir):
                    src = os.path.join(temp_dir, filename)
                    if os.path.isfile(src):
                        dst = os.path.join(output_dir, filename)
                        shutil.copy(src, dst)
                        logger.debug(f"保留中间文件: {filename}")
            
            logger.info(f"MD转DOCX成功: {output_path}")
            return output_path
        
    except ValueError:
        raise
    except Exception as e:
        logger.error(f"MD转DOCX失败: {e}", exc_info=True)
        return None


def _read_and_parse_md(temp_md_path: str, original_md_path: Optional[str] = None) -> tuple:
    """
    读取并解析Markdown文件，返回YAML数据和YAML后的Markdown内容
    
    新增功能：
    - 在返回前展开Markdown内容中的嵌入链接（如果启用）
    
    参数:
        temp_md_path: 临时Markdown文件路径（用于读取内容）
        original_md_path: 原始Markdown文件路径（用于路径解析，可选）
    
    返回:
        tuple: (yaml_data字典, md_body字符串)
    """
    logger.debug(f"读取Markdown文件: {temp_md_path}")
    with open(temp_md_path, 'r', encoding='utf-8') as f:
        content = f.read()
    
    # 匹配YAML格式部分
    yaml_match = re.search(r'---\n(.*?)\n---', content, re.DOTALL)
    
    if yaml_match:
        yaml_content = yaml_match.group(1).replace('\t', '  ')
        md_body = content[yaml_match.end():].strip()
        logger.debug(f"提取YAML内容: {len(yaml_content)} 字符")
        logger.debug(f"提取YAML后的内容: {len(md_body)} 字符")
    else:
        logger.warning("未找到YAML头部，使用空YAML")
        yaml_content = ""
        md_body = content.strip()
    
    yaml_data = yaml_utils.parse_yaml(yaml_content)
    
    # 处理所有Markdown链接（嵌入和非嵌入）
    try:
        from docwen.utils.link_processing import process_markdown_links
        
        # 使用原始文件路径进行路径解析（如果提供），否则使用临时路径
        source_path_for_resolve = original_md_path if original_md_path else temp_md_path
        
        # 1. 处理YAML字段值中的链接
        logger.info("开始处理YAML字段值中的链接...")
        
        def process_yaml_links(data, source_path):
            """递归处理YAML数据结构中的所有链接"""
            if isinstance(data, dict):
                return {k: process_yaml_links(v, source_path) for k, v in data.items()}
            elif isinstance(data, list):
                return [process_yaml_links(item, source_path) for item in data]
            elif isinstance(data, str):
                # 对字符串调用链接处理（depth=0，允许嵌套展开）
                return process_markdown_links(data, source_path, set(), depth=0)
            else:
                return data
        
        yaml_data = process_yaml_links(yaml_data, source_path_for_resolve)
        logger.info("YAML字段链接处理完成")
        
        # 2. 处理Markdown内容中的链接
        logger.info("开始处理Markdown内容中的链接...")
        original_length = len(md_body)
        md_body = process_markdown_links(md_body, source_path_for_resolve)
        new_length = len(md_body)
        logger.info(f"Markdown内容链接处理完成 | 长度变化: {original_length} → {new_length}")
    except Exception as e:
        logger.error(f"处理Markdown链接失败: {e}", exc_info=True)
        logger.warning("将继续使用原始MD内容")
    
    return yaml_data, md_body


def _convert_internal(
    yaml_data: dict,
    md_body: str,
    output_path: str,
    template_name: str,
    progress_callback,
    cancel_event,
    md_path: Optional[str] = None,
    options: Optional[dict] = None
) -> bool:
    """
    内部转换实现
    
    复用现有的md2docx模块代码
    
    参数:
        yaml_data: YAML数据字典
        md_body: Markdown正文
        output_path: 输出文件路径
        template_name: 模板名称
        progress_callback: 进度回调
        cancel_event: 取消事件
        md_path: 原始MD文件路径（用于从文件名提取标题）
        options: 转换选项字典，包含序号配置等
    
    返回:
        bool: 转换是否成功
    """
    try:
        # 1. 预处理DOCX相关的YAML数据
        logger.debug("预处理DOCX相关的YAML数据...")
        process_docx_specific_yaml(yaml_data)
        
        # 2. 为所有语言版本的标题键设置回退值
        _ensure_title_fallbacks(yaml_data, md_path)
        logger.debug(f"标题回退值设置完成")
        
        # 3. 处理Markdown正文（支持序号配置）
        if progress_callback:
            progress_callback(t('conversion.progress.processing_markdown_body'))
        logger.debug("处理Markdown正文...")
        
        # 提取序号配置参数
        if options is None:
            options = {}
        
        remove_numbering = options.get('md_remove_numbering', False)
        add_numbering = options.get('md_add_numbering', False)
        scheme_name = options.get('md_numbering_scheme', 'gongwen_standard')
        
        # 创建序号格式化器（如果需要添加序号）
        formatter = None
        if add_numbering:
            from docwen.utils.heading_numbering import get_formatter_from_config
            from docwen.config.config_manager import config_manager
            formatter = get_formatter_from_config(config_manager, scheme_name)
            if not formatter:
                logger.warning(f"无法创建序号格式化器（方案：{scheme_name}），将不添加序号")
                add_numbering = False
        
        logger.info(f"处理Markdown正文：清除序号={remove_numbering}, 添加序号={add_numbering}")
        
        # 统一调用 process_md_body_with_notes（处理正文并提取脚注/尾注）
        processed_body, footnotes, endnotes = md_processor.process_md_body_with_notes(
            md_body, 
            remove_numbering=remove_numbering,
            add_numbering=add_numbering,
            formatter=formatter
        )
        
        logger.info(f"正文处理完成, 生成 {len(processed_body)} 个段落, 脚注: {len(footnotes)} 个, 尾注: {len(endnotes)} 个")
        
        # 4. 加载模板（先确保标题样式存在）
        if progress_callback:
            progress_callback(t('conversion.progress.loading_word_template'))
        logger.debug(f"加载Word模板: {template_name}")
        
        # 获取原始模板路径，检查并注入缺失的标题样式
        # 使用统一的临时目录清理回调
        temp_dirs_to_cleanup = []
        def cleanup_temp_dir(temp_dir):
            temp_dirs_to_cleanup.append(temp_dir)
        
        template_loader = TemplateLoader()
        original_template_path = template_loader.get_template_path("docx", template_name)
        processed_template_path = ensure_styles(original_template_path, progress_callback, cleanup_callback=cleanup_temp_dir)
        
        # 如果有脚注或尾注，确保模板包含脚注/尾注Part
        if footnotes or endnotes:
            processed_template_path = ensure_notes_parts(processed_template_path, cleanup_callback=cleanup_temp_dir)
            logger.info("已确保模板包含脚注/尾注Part")
        
        # 如果有列表项，确保模板包含numbering.xml
        has_list_items = any(item.get('type') == 'list_item' for item in processed_body)
        if has_list_items:
            processed_template_path = ensure_numbering_part(processed_template_path, cleanup_callback=cleanup_temp_dir)
            logger.info("已确保模板包含numbering.xml")
        
        # 如果返回了临时路径（说明注入了样式或添加了Part），使用临时路径加载
        if processed_template_path != original_template_path:
            doc = open_docx(processed_template_path)
            logger.info(f"使用增强后的临时模板: {processed_template_path}")
        else:
            doc = template_loader.load_docx_template(template_name)
        
        logger.info("模板加载成功")
        
        # 5. 替换占位符（传递脚注/尾注数据）
        if progress_callback:
            progress_callback(t('conversion.progress.filling_template'))
        logger.debug("处理模板占位符...")
        success = _process_docx_template(doc, output_path, yaml_data, processed_body, template_name, footnotes, endnotes)
        
        if not success:
            logger.error("在模板处理过程中发生错误")
            return False
            
        if cancel_event and cancel_event.is_set():
            logger.info("操作被用户取消")
            return False

        logger.info(f"转换成功! DOCX文件已保存到: {output_path}")
        return True
        
    except ValueError:
        raise
    except Exception as e:
        logger.error(f"内部转换失败: {e}", exc_info=True)
        return False
    finally:
        # 清理 ensure_notes_parts 创建的临时目录
        for temp_dir in temp_dirs_to_cleanup:
            try:
                if os.path.isdir(temp_dir):
                    shutil.rmtree(temp_dir)
                    logger.debug(f"已清理临时目录: {temp_dir}")
            except Exception as e:
                logger.warning(f"清理临时目录失败: {temp_dir}, 错误: {e}")


def _get_fallback_title(yaml_data: dict, md_path: Optional[str] = None) -> str:
    """
    获取标题的回退值
    
    优先级：aliases → 文件名 → "标题"
    
    参数:
        yaml_data: YAML数据字典
        md_path: 原始MD文件路径（用于提取文件名）
    
    返回:
        str: 回退标题文本
    """
    # 优先级1：aliases
    if 'aliases' in yaml_data:
        aliases = yaml_data['aliases']
        if isinstance(aliases, str) and aliases.strip():
            logger.debug("使用YAML中的'aliases'键（字符串）作为回退标题")
            return aliases.strip()
        elif isinstance(aliases, list) and aliases:
            logger.debug("使用YAML中的'aliases'键（列表第一项）作为回退标题")
            return str(aliases[0]).strip()
    
    # 优先级2：文件名
    if md_path:
        filename = os.path.splitext(os.path.basename(md_path))[0]
        if filename:
            logger.debug(f"从文件名提取回退标题: {filename}")
            return filename
    
    # 最终默认值
    logger.debug("使用默认回退标题")
    return "标题"


def _ensure_title_fallbacks(yaml_data: dict, md_path: Optional[str] = None):
    """
    为所有语言版本的标题键设置回退值
    
    逻辑：对于每个标题键（如 'title', '标题', 'Titel'...）：
    1. 如果 YAML 中已有该键且非空 → 保持原值
    2. 如果 YAML 中没有该键或为空 → 设置回退值（aliases → 文件名）
    
    参数:
        yaml_data: YAML数据字典（会被原地修改）
        md_path: 原始MD文件路径（用于提取文件名）
    """
    # 获取所有语言版本的标题键名（从 placeholders.title）
    title_variants = set(t_all_locales('placeholders.title').values())
    logger.debug(f"标题键变体: {title_variants}")
    
    # 计算回退值
    fallback_value = _get_fallback_title(yaml_data, md_path)
    
    # 为每个变体设置回退值
    for title_key in title_variants:
        if title_key not in yaml_data or not yaml_data[title_key]:
            yaml_data[title_key] = fallback_value
            logger.debug(f"为 '{title_key}' 设置回退值: {fallback_value}")


def _process_docx_template(
    doc: DocxDocument,
    output_path: str,
    yaml_data: dict,
    body_data: list,
    template_name: Optional[str] = None,
    footnotes: Optional[dict] = None,
    endnotes: Optional[dict] = None,
) -> bool:
    """
    处理DOCX模板并保存结果
    
    参数:
        doc: Document对象
        output_path: 输出路径
        yaml_data: YAML数据
        body_data: 正文数据
        template_name: 模板名称
        footnotes: 脚注字典 {id: content}
        endnotes: 尾注字典 {id: content}
    
    返回:
        bool: 是否成功
    """
    try:
        logger.debug("替换模板中的占位符...")
        temp_path, warnings = docx_processor.replace_placeholders(doc, yaml_data, body_data, template_name, 
                                                        footnotes=footnotes, endnotes=endnotes)
        
        # 记录所有警告
        for warning_msg in warnings:
            logger.warning(warning_msg)
        
        output_dir = os.path.dirname(output_path)
        os.makedirs(output_dir, exist_ok=True)
        
        shutil.move(temp_path, output_path)
        
        if not os.path.exists(output_path):
            logger.error(f"移动文件到最终位置失败: {output_path}")
            return False
            
        return True
    except ValueError:
        # 精确放行：如果是我们预期的业务错误（如缺少样式），直接向上抛出
        # 这样上层策略就能捕获到具体的错误信息并显示给用户
        raise
    except Exception as e:
        # 其他未知错误，记录日志并返回False（显示通用错误信息）
        logger.error(f"处理DOCX模板时发生错误: {str(e)}", exc_info=True)
        return False


# --- DOCX相关的YAML处理函数 ---

def process_docx_specific_yaml(data: dict):
    """处理DOCX输出特有的YAML字段"""
    logger.info("处理DOCX特定的YAML字段...")
    process_attachment_description(data)
    process_cc_orgs(data)
    process_special_fields(data)


def process_attachment_description(data: dict):
    """
    处理附件说明字段
    第2行起添加全角空格对齐，配合悬挂缩进使用
    """
    if '附件说明' not in data:
        return
    
    attachments = data['附件说明']
    if is_value_empty(attachments):
        data['附件说明'] = []
        return
    
    if not isinstance(attachments, list):
        attachments = [attachments]
    
    # 清理序号
    cleaned_attachments = []
    for item in attachments:
        content = str(item).strip() if item is not None else ""
        normalized_content = convert_to_halfwidth(content)
        cleaned = ATTACH_NUM_PATTERN.sub('', normalized_content).strip()
        if cleaned == "" and content != "":
            cleaned = content
        cleaned_attachments.append(cleaned)
    
    # 构建附件说明列表
    formatted = []
    for i, content in enumerate(cleaned_attachments, 1):
        if len(cleaned_attachments) == 1:
            # 单附件：附件：<内容>
            formatted.append(f"附件：{content}")
        else:
            # 多附件
            if i == 1:
                formatted.append(f"附件：{i}. {content}")
            else:
                # 第2行起：加3个全角空格（序号10开始，加两个全角空格和一个半角空格）
                indent = "\u3000\u3000\u3000" if i < 10 else "\u3000\u3000 "
                formatted.append(f"{indent}{i}. {content}")
    
    data['附件说明'] = formatted
    logger.debug(f"处理附件说明完成，共 {len(formatted)} 项")


def process_cc_orgs(data: dict):
    """处理抄送机关字段"""
    if '抄送机关' in data:
        cc_orgs = data['抄送机关']
        if is_value_empty(cc_orgs):
            data['抄送机关'] = ""
        else:
            if not isinstance(cc_orgs, list):
                cc_orgs = [cc_orgs]
            
            # 使用format_display_value处理可能的嵌套列表
            from docwen.utils.text_utils import format_display_value
            valid_orgs = [format_display_value(org).strip() for org in cc_orgs if not is_value_empty(org)]
            data['抄送机关'] = "，".join(valid_orgs)


def process_special_fields(data: dict):
    """处理特殊字段"""
    if '附注' in data:
        data['附注'] = process_notes(data['附注'])
    if '印发日期' in data:
        data['印发日期'] = format_date(data['印发日期'], suffix="印发")
    if '成文日期' in data:
        data['成文日期'] = format_date(data['成文日期'])


def process_notes(notes: str) -> str:
    """处理附注字段"""
    if not notes:
        return ""
    if re.match(r'^[（(].*?[)）]$', notes):
        return notes[1:-1]
    return notes


def format_date(date_str, suffix: str = "") -> str:
    """格式化日期字段"""
    if not date_str:
        return ""
    if isinstance(date_str, (datetime.date, datetime.datetime)):
        date_obj = date_str
    else:
        date_formats = ["%Y-%m-%d", "%Y/%m/%d", "%Y年%m月%d日", "%Y.%m.%d", "%Y年%m月%d号"]
        date_obj = None
        for fmt in date_formats:
            try:
                date_obj = datetime.datetime.strptime(str(date_str), fmt)
                break
            except ValueError:
                continue
    if date_obj:
        formatted = f"{date_obj.year}年{date_obj.month}月{date_obj.day}日"
        if suffix:
            formatted += suffix
        return formatted
    return str(date_str)
