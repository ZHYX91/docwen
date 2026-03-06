"""
模板加载器模块

本模块负责模板文件的加载、缓存和管理，提供统一的模板访问接口。
主要功能包括：
- 支持DOCX和XLSX两种格式的模板加载
- 智能模板缓存机制，提高加载性能
- 模板文件的模糊匹配和智能查找
- 自动检测模板文件变更并更新缓存
- 支持开发和生产环境的多路径适配

"""

import copy
import logging
import sys
from pathlib import Path

from docx import Document
from openpyxl import load_workbook

from docwen.utils.path_utils import ensure_dir_exists_or_raise, get_project_root

# 配置日志
logger = logging.getLogger(__name__)

# 全局模板缓存
_TEMPLATE_CACHE = {}
_TEMPLATE_MTIMES = {}  # 记录模板文件的最后修改时间


class TemplateLoader:
    """模板加载器核心类"""

    def __init__(self, template_dir=None):
        """
        初始化模板加载器
        :param template_dir: 自定义模板目录（可选）
        """
        self.template_dir = template_dir or self.get_default_template_dir()
        logger.info(f"初始化模板加载器 | 模板目录: {self.template_dir}")
        self._available_templates_cache: dict[tuple[str, bool], tuple[float, list[str]]] = {}
        self._validate_template_dir()
        self._refresh_cache()

    def get_default_template_dir(self) -> str:
        """
        获取默认模板目录路径
        """
        # 获取项目根目录
        project_root = get_project_root()
        logger.debug(f"项目根目录: {project_root}")

        # 开发环境：项目根目录的templates子目录
        templates_dir = Path(project_root) / "templates"

        # 如果模板目录存在，直接使用
        if templates_dir.exists():
            logger.debug(f"使用开发环境模板目录: {templates_dir}")
            return str(templates_dir)

        # 打包环境：可执行文件同级的templates目录
        if getattr(sys, "frozen", False):
            exe_dir = Path(sys.executable).resolve().parent
            templates_dir = exe_dir / "templates"
            if templates_dir.exists():
                logger.debug(f"使用生产环境模板目录: {templates_dir}")
                return str(templates_dir)

        # 回退到用户目录
        templates_dir = Path("~").expanduser() / ".docwen" / "templates"
        ensure_dir_exists_or_raise(str(templates_dir))
        logger.warning(f"使用用户模板目录: {templates_dir}")
        return str(templates_dir)

    def _validate_template_dir(self):
        """验证模板目录是否存在"""
        template_dir = Path(self.template_dir)
        if not template_dir.exists():
            logger.error(f"模板目录不存在: {self.template_dir}")
            raise FileNotFoundError(f"模板目录不存在: {self.template_dir}")

        if not template_dir.is_dir():
            logger.error(f"模板路径不是目录: {self.template_dir}")
            raise NotADirectoryError(f"模板路径不是目录: {self.template_dir}")

        logger.info(f"模板目录验证通过: {self.template_dir}")

    def _normalize_template_name(self, template_type: str, template_name: str) -> str:
        if not template_name:
            raise ValueError("模板名称是必需的，请指定要使用的模板")
        normalized = template_name.strip()
        if not normalized:
            raise ValueError("模板名称是必需的，请指定要使用的模板")
        if not normalized.lower().endswith(f".{template_type}"):
            normalized = f"{normalized}.{template_type}"
        return normalized

    def get_template_path(self, template_type: str, template_name: str) -> str:
        """
        获取模板完整路径
        支持自动添加文件扩展名

        :param template_type: 模板类型 ('docx' 或 'xlsx')
        :param template_name: 模板名 (必需)
        :return: 模板完整路径
        """
        template_name = self._normalize_template_name(template_type, template_name)

        # 构建完整路径
        template_path = str(Path(self.template_dir) / template_name)

        # 检查模板是否存在
        if not Path(template_path).exists():
            logger.error(f"模板文件不存在: {template_path}")
            raise FileNotFoundError(f"模板文件不存在: {template_path}")

        logger.info(f"获取模板路径: {template_path}")
        return template_path

    def load_docx_template(self, template_name: str):
        """
        加载DOCX模板

        参数:
            template_name: 模板名 (必需)

        返回:
            Document对象（深拷贝）
        """
        normalized_name = self._normalize_template_name("docx", template_name)
        template_path = None
        try:
            # 获取模板路径
            template_path = self.get_template_path("docx", normalized_name)
            logger.debug(f"DOCX模板路径: {template_path}")

            # 生成缓存键
            cache_key = self._generate_cache_key("docx", normalized_name)
            logger.debug(f"缓存键: {cache_key}")

            # 检查文件是否修改
            current_mtime = Path(template_path).stat().st_mtime
            last_mtime = _TEMPLATE_MTIMES.get(cache_key, 0)
            logger.debug(f"文件修改时间检查 | 当前: {current_mtime} | 上次: {last_mtime}")

            # 文件已修改或缓存不存在时重新加载
            if current_mtime > last_mtime or cache_key not in _TEMPLATE_CACHE:
                if cache_key in _TEMPLATE_CACHE:
                    logger.info(f"模板已修改或缓存过期，重新加载: {template_path}")
                else:
                    logger.info(f"模板未缓存，首次加载: {template_path}")

                # 清除旧缓存
                if cache_key in _TEMPLATE_CACHE:
                    del _TEMPLATE_CACHE[cache_key]
                    logger.debug(f"清除旧缓存: {cache_key}")

                # 加载新模板
                logger.debug(f"加载模板文件: {template_path}")
                doc = Document(template_path)

                # 更新缓存和修改时间
                _TEMPLATE_CACHE[cache_key] = doc
                _TEMPLATE_MTIMES[cache_key] = current_mtime
                logger.info(f"缓存更新 | 键: {cache_key} | 大小: {Path(template_path).stat().st_size}字节")
            else:
                logger.debug(f"使用缓存DOCX模板: {cache_key}")
                doc = _TEMPLATE_CACHE[cache_key]

            # 返回模板的深拷贝，避免修改缓存中的原始模板
            logger.debug("创建模板的深拷贝...")
            try:
                # 尝试深拷贝
                doc_copy = copy.deepcopy(doc)
                logger.debug("深拷贝创建成功")
                return doc_copy
            except Exception as copy_error:
                # 深拷贝失败时回退到浅拷贝
                logger.warning(f"深拷贝失败: {copy_error!s}，使用浅拷贝")
                return copy.copy(doc)

        except FileNotFoundError as e:
            missing = Path(template_path).name if template_path else normalized_name
            logger.error(f"模板文件不存在: {missing}")
            raise RuntimeError(f"模板文件不存在: {Path(missing).name}") from e
        except PermissionError as e:
            logger.error(f"无权限访问模板文件: {template_path or normalized_name}")
            raise RuntimeError("无权限访问模板文件") from e
        except Exception as e:
            error_str = str(e).lower()
            # 检测是否为无效的 DOCX 文件（不是有效的 ZIP 压缩包）
            if "not a zip file" in error_str or "package not found" in error_str:
                from docwen.translation import t

                template_basename = Path(template_path).name if template_path else Path(normalized_name).name
                friendly_msg = t("messages.errors.invalid_docx_template", template_name=template_basename)
                logger.error(f"模板文件无效或损坏: {template_path or normalized_name}")
                raise RuntimeError(friendly_msg) from e
            logger.error(f"加载DOCX模板失败: {e!s}", exc_info=True)
            raise RuntimeError(f"加载模板失败: {e!s}") from e

    def load_xlsx_template(self, template_name: str):
        """
        加载Excel模板

        :param template_name: 模板名 (必需)
        :return: Workbook对象
        """
        normalized_name = self._normalize_template_name("xlsx", template_name)
        try:
            # 获取模板路径
            template_path = self.get_template_path("xlsx", normalized_name)
            logger.debug(f"Excel模板路径: {template_path}")

            # 生成缓存键
            cache_key = self._generate_cache_key("xlsx", normalized_name)

            # 检查文件是否修改
            current_mtime = Path(template_path).stat().st_mtime
            last_mtime = _TEMPLATE_MTIMES.get(cache_key, 0)

            # 文件已修改或缓存不存在时重新加载
            if current_mtime > last_mtime or cache_key not in _TEMPLATE_CACHE:
                logger.info(f"模板已修改或缓存过期，重新加载: {template_path}")

                # 清除旧缓存
                _TEMPLATE_CACHE.pop(cache_key, None)

                # 加载新模板
                wb = load_workbook(template_path)

                # 更新缓存和修改时间
                _TEMPLATE_CACHE[cache_key] = wb
                _TEMPLATE_MTIMES[cache_key] = current_mtime
                logger.debug(f"缓存更新 | 键: {cache_key} | 大小: {Path(template_path).stat().st_size}字节")
            else:
                logger.debug(f"使用缓存Excel模板: {cache_key}")
                wb = _TEMPLATE_CACHE[cache_key]

            return wb
        except Exception as e:
            logger.error(f"加载Excel模板失败: {e!s}")
            raise

    def _generate_cache_key(self, template_type: str, template_name: str) -> str:
        """生成模板缓存键（唯一标识）"""
        normalized_name = self._normalize_template_name(template_type, template_name)
        return f"{template_type}:{normalized_name}".lower()

    def get_available_templates(self, template_type: str, name_only: bool = True) -> list[str]:
        """
        获取指定类型的所有可用模板
        支持只返回名称（不带扩展名）

        :param template_type: 模板类型 ('docx' 或 'xlsx')
        :param name_only: 是否只返回名称（不带扩展名）
        :return: 模板名称列表
        """
        logger.debug(f"获取可用{template_type.upper()}模板...")

        # 检查模板目录是否存在
        template_dir = Path(self.template_dir)
        if not template_dir.exists():
            logger.warning(f"模板目录不存在: {self.template_dir}")
            return []

        try:
            dir_mtime = template_dir.stat().st_mtime
        except Exception:
            dir_mtime = -1.0

        cache_key = (template_type, bool(name_only))
        cached = self._available_templates_cache.get(cache_key)
        if cached and cached[0] == dir_mtime:
            return cached[1].copy()

        # 获取指定类型的模板文件
        extension = ".docx" if template_type == "docx" else ".xlsx"
        templates = []

        try:
            for entry in template_dir.iterdir():
                if not entry.is_file():
                    continue
                f = entry.name
                if not f.lower().endswith(extension):
                    continue
                name = entry.stem if name_only else f
                templates.append(name)
                logger.debug(f"找到模板: {name}")
        except Exception:
            templates = []

        # 按字母排序
        templates.sort()
        logger.debug(f"找到 {len(templates)} 个{template_type.upper()}模板")
        self._available_templates_cache[cache_key] = (dir_mtime, templates)
        return templates

    def get_template_list(self, template_type: str) -> list[str]:
        """
        获取指定类型的模板列表（GUI专用）

        :param template_type: 模板类型 ('docx' 或 'xlsx')
        :return: 模板名称列表
        """
        return self.get_available_templates(template_type)

    def find_template(self, template_type: str, search_term: str) -> str:
        """
        模糊查找模板

        :param template_type: 模板类型 ('docx' 或 'xlsx')
        :param search_term: 搜索关键词
        :return: 匹配的模板名称（找不到时返回空字符串）
        """
        # 获取所有模板
        templates = self.get_available_templates(template_type)

        # 完全匹配
        if search_term in templates:
            logger.debug(f"完全匹配模板: {search_term}")
            return search_term

        # 模糊匹配
        matches = []
        for template in templates:
            if search_term.lower() in template.lower():
                matches.append(template)

        # 单个匹配
        if len(matches) == 1:
            logger.debug(f"模糊匹配模板: {matches[0]}")
            return matches[0]

        # 多个匹配
        if matches:
            logger.debug(f"找到多个匹配模板: {matches}")
            return matches[0]  # 返回第一个匹配项

        # 无匹配
        logger.warning(f"未找到匹配模板: {search_term}")
        return ""

    def clear_cache(self, template_type: str | None = None, template_name: str | None = None):
        """
        清空模板缓存（可指定类型或特定模板）

        :param template_type: 模板类型 ('docx' 或 'xlsx')，可选
        :param template_name: 模板名称，可选
        """
        global _TEMPLATE_CACHE, _TEMPLATE_MTIMES

        if template_type and template_name:
            # 清除特定模板缓存
            cache_key = self._generate_cache_key(template_type, template_name)
            _TEMPLATE_CACHE.pop(cache_key, None)
            _TEMPLATE_MTIMES.pop(cache_key, None)
            logger.info(f"清除特定模板缓存: {template_type}/{template_name}")
        elif template_type:
            # 清除特定类型的所有缓存
            normalized_type = template_type.strip().lower()
            prefix = f"{normalized_type}:"
            keys_to_remove = [k for k in _TEMPLATE_CACHE if k.startswith(prefix)]
            for key in keys_to_remove:
                del _TEMPLATE_CACHE[key]
                _TEMPLATE_MTIMES.pop(key, None)
            logger.info(f"清除{normalized_type.upper()}类型所有模板缓存")
        else:
            # 清空所有缓存
            cache_size = len(_TEMPLATE_CACHE)
            _TEMPLATE_CACHE = {}
            _TEMPLATE_MTIMES = {}
            logger.info(f"清空所有模板缓存 | 已移除 {cache_size} 个模板")

    def refresh_cache(self):
        """刷新缓存（检查所有模板的修改时间）"""
        logger.info("开始刷新模板缓存...")

        # 获取所有模板路径
        docx_templates = self.get_available_templates("docx", name_only=False)
        xlsx_templates = self.get_available_templates("xlsx", name_only=False)
        all_templates = [(name, "docx") for name in docx_templates] + [(name, "xlsx") for name in xlsx_templates]

        refreshed_count = 0

        for name, t_type in all_templates:
            try:
                # 获取模板路径
                template_path = self.get_template_path(t_type, name)

                # 获取当前修改时间
                current_mtime = Path(template_path).stat().st_mtime

                # 生成缓存键
                cache_key = self._generate_cache_key(t_type, name)

                # 检查是否需要刷新
                if cache_key in _TEMPLATE_MTIMES and current_mtime > _TEMPLATE_MTIMES[cache_key]:
                    logger.info(f"模板已修改: {name}.{t_type}")

                    # 清除缓存
                    _TEMPLATE_CACHE.pop(cache_key, None)

                    # 更新修改时间
                    _TEMPLATE_MTIMES[cache_key] = current_mtime
                    refreshed_count += 1
            except Exception as e:
                logger.error(f"刷新模板缓存失败: {name}.{t_type} - {e!s}")

        logger.info(f"缓存刷新完成 | 更新了 {refreshed_count}/{len(all_templates)} 个模板")

    def _refresh_cache(self):
        """内部方法：初始加载时刷新缓存"""
        try:
            self.refresh_cache()
        except Exception as e:
            logger.error(f"初始缓存刷新失败: {e!s}")
