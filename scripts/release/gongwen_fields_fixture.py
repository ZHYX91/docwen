"""Shared input and DOCX result oracle for the frozen Gongwen field gate.

The verifier generates documents but never imports the field implementation.
Both packaged entry points must produce the same observable document content.
"""

from __future__ import annotations

import hashlib
import json
import zipfile
from dataclasses import dataclass
from pathlib import Path
from xml.etree import ElementTree

_WORD = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
_BODY_START = "DOCWEN ATTACHMENTS START"
_BODY_END = "DOCWEN ATTACHMENTS END"
_EMPTY_START = "DOCWEN EMPTY FIELD START"
_EMPTY_END = "DOCWEN EMPTY FIELD END"
_CONTENT = "DOCWEN PACKAGED GONGWEN FIELDS CONTENT"
_TEMPLATE_ID = f"template.docx.{hashlib.sha256(b'docwen-gongwen-fields-release-fixture').hexdigest()}"


@dataclass(frozen=True)
class GongwenFieldsCase:
    name: str
    work_dir: Path
    source_path: Path
    template_path: Path
    template_id: str
    enabled: bool
    empty: bool


def write_gongwen_fields_cases(root: Path) -> tuple[GongwenFieldsCase, ...]:
    """Create four independent profiles, including the GUI's normal template state."""
    from docx import Document
    from docx.shared import Pt

    cases = []
    for enabled in (True, False):
        for empty in (False, True):
            name = f"{'on' if enabled else 'off'}-{'empty' if empty else 'values'}"
            work_dir = root / name
            config_dir = work_dir / "config_home"
            data_dir = work_dir / "data_home"
            template_path = data_dir / "templates" / "fields.docx"
            config_dir.mkdir(parents=True)
            template_path.parent.mkdir(parents=True)
            (work_dir / "outputs").mkdir()
            (config_dir / "template_fill.toml").write_text(
                f"gongwen_fields_enabled = {'true' if enabled else 'false'}\nlist_separator = ' / '\n",
                encoding="utf-8",
            )
            # An explicitly enabled Chinese field must also work with English
            # output styles and interface text. These are independent choices.
            (config_dir / "conversion.toml").write_text('output_locale = "en_US"\n', encoding="utf-8")
            (config_dir / "gui.toml").write_text('[language]\nlocale = "en_US"\n', encoding="utf-8")

            document = Document()
            for text in (
                _BODY_START,
                "{{附件说明}}",
                _BODY_END,
                "普通列表：{{普通列表}}",
                "抄送：{{抄送机关}}",
                "（{{附注}}）",
                "成文：{{成文日期}}",
                _EMPTY_START,
                "{{主送机关}}：",
                _EMPTY_END,
            ):
                document.add_paragraph(text)
            attachments = document.add_table(rows=1, cols=1)
            attachments.cell(0, 0).text = "{{附件说明}}"
            anchors = [document.paragraphs[1], attachments.cell(0, 0).paragraphs[0]]
            for anchor in anchors:
                anchor.paragraph_format.left_indent = Pt(12)
                anchor.paragraph_format.right_indent = Pt(8)
                anchor.paragraph_format.keep_with_next = True
                anchor.runs[0].font.size = Pt(16)
                anchor.runs[0].bold = True
            imprint = document.add_table(rows=2, cols=2)
            imprint.cell(0, 0).text = "DOCWEN IMPRINT"
            imprint.cell(0, 1).text = "KEEP THIS ROW"
            imprint.cell(1, 0).text = "{{印发机关}}"
            imprint.cell(1, 1).text = "{{印发日期}}"
            document.add_paragraph("{{正文}}")
            document.save(str(template_path))

            # Seed an ordinary user template and its persisted default. The
            # frozen GUI still discovers/selects it through its normal catalog;
            # no plugin import or special conversion hook is added for this gate.
            state = {
                "version": 1,
                "defaults_initialized": True,
                "user_identities": {
                    template_path.name: {
                        "filename": template_path.name,
                        "id": _TEMPLATE_ID,
                        "target": "docx",
                    }
                },
                "enabled": {_TEMPLATE_ID: True},
                "order": {"docx": [_TEMPLATE_ID], "xlsx": []},
                "defaults": {"docx": _TEMPLATE_ID, "xlsx": None},
            }
            (data_dir / "template-state.json").write_text(json.dumps(state), encoding="utf-8")
            source_path = work_dir / "fields.md"
            attachments_yaml = "[]" if empty else '["任务分工表", "工作安排表"]'
            issuer = '""' if empty else '"测试机关"'
            print_date = '""' if empty else '"2026-10-21"'
            source_path.write_text(
                "---\n"
                f"附件说明: {attachments_yaml}\n"
                '普通列表: ["普通甲", "普通乙"]\n'
                '抄送机关: ["单位甲", "单位乙"]\n'
                '附注: "（此件公开发布）"\n'
                '成文日期: "2026-10-20"\n'
                f"印发机关: {issuer}\n"
                f"印发日期: {print_date}\n"
                '主送机关: ""\n'
                "---\n\n"
                f"{_CONTENT}\n",
                encoding="utf-8",
            )
            cases.append(GongwenFieldsCase(name, work_dir, source_path, template_path, _TEMPLATE_ID, enabled, empty))
    return tuple(cases)


def _paragraph_text(element: ElementTree.Element) -> str:
    return "".join(
        node.text or "" if node.tag == f"{_WORD}t" else "\t" if node.tag == f"{_WORD}tab" else "\n"
        for run in element.iter(f"{_WORD}r")
        for node in run
        if node.tag in {f"{_WORD}t", f"{_WORD}tab", f"{_WORD}br"}
    )


def _between(paragraphs: list[str], start: str, end: str) -> list[str]:
    if paragraphs.count(start) != 1 or paragraphs.count(end) != 1:
        raise RuntimeError(f"gongwen_fields_fixture_markers_missing:{start}")
    first, last = paragraphs.index(start), paragraphs.index(end)
    if first >= last:
        raise RuntimeError(f"gongwen_fields_fixture_markers_reordered:{start}")
    return paragraphs[first + 1 : last]


def verify_gongwen_fields_output(output_path: Path, *, case: GongwenFieldsCase) -> None:
    """Require specialized paragraphs and generic values to survive real export."""
    if output_path.suffix.lower() != ".docx" or not output_path.is_file():
        raise RuntimeError(f"gongwen_fields_output_missing:{case.name}:{output_path}")
    try:
        with zipfile.ZipFile(output_path) as package:
            root = ElementTree.fromstring(package.read("word/document.xml"))
    except (OSError, KeyError, zipfile.BadZipFile, ElementTree.ParseError) as exc:
        raise RuntimeError(f"gongwen_fields_output_invalid:{case.name}") from exc
    body = root.find(f"{_WORD}body")
    if body is None:
        raise RuntimeError(f"gongwen_fields_body_missing:{case.name}")
    paragraphs = [_paragraph_text(paragraph) for paragraph in body.findall(f"{_WORD}p")]
    all_text = "\n".join(_paragraph_text(paragraph) for paragraph in root.iter(f"{_WORD}p"))
    if "{{" in all_text or _CONTENT not in paragraphs:
        raise RuntimeError(f"gongwen_fields_template_not_filled:{case.name}")

    body_attachments = _between(paragraphs, _BODY_START, _BODY_END)
    tables = body.findall(f"{_WORD}tbl")
    if len(tables) != 2:
        raise RuntimeError(f"gongwen_fields_tables_changed:{case.name}:{len(tables)}")
    attachment_cells = tables[0].findall(f"{_WORD}tr/{_WORD}tc")
    if len(attachment_cells) != 1:
        raise RuntimeError(f"gongwen_fields_attachment_cell_changed:{case.name}")
    cell = attachment_cells[0]
    if not len(cell) or cell[-1].tag != f"{_WORD}p":
        raise RuntimeError(f"gongwen_fields_attachment_cell_invalid:{case.name}")
    cell_attachments = [_paragraph_text(paragraph) for paragraph in cell.findall(f"{_WORD}p")]

    if case.enabled and not case.empty:
        # Space-padded literal numbering must remain editable ordinary text.
        for location, actual in (("body", body_attachments), ("table", cell_attachments)):
            if len(actual) != 2 or not all(
                expected == text
                for expected, text in zip(
                    ("附件：1. 任务分工表", "\u3000\u3000\u30002. 工作安排表"), actual, strict=True
                )
            ):
                raise RuntimeError(f"gongwen_fields_attachments_mismatch:{case.name}:{location}:{actual!r}")
        first, last = paragraphs.index(_BODY_START), paragraphs.index(_BODY_END)
        body_nodes = body.findall(f"{_WORD}p")[first + 1 : last]
        for location, nodes in (("body", body_nodes), ("table", cell.findall(f"{_WORD}p"))):
            for index, paragraph in enumerate(nodes):
                indent = paragraph.find(f"{_WORD}pPr/{_WORD}ind")
                size = paragraph.find(f"{_WORD}r/{_WORD}rPr/{_WORD}sz")
                bold = paragraph.find(f"{_WORD}r/{_WORD}rPr/{_WORD}b")
                keep_next = paragraph.find(f"{_WORD}pPr/{_WORD}keepNext")
                # This fixture authors 16 pt characters and a 12 pt base indent:
                # the 48 pt label and 16 pt number/period/space place names at 76 pt.
                if (
                    indent is None
                    or indent.get(f"{_WORD}left") != "1520"
                    or indent.get(f"{_WORD}right") != "160"
                    or indent.get(f"{_WORD}hanging") != "1280"
                    or any(
                        indent.get(f"{_WORD}{key}") is not None
                        for key in ("leftChars", "hangingChars", "firstLineChars")
                    )
                    or paragraph.find(f".//{_WORD}tab") is not None
                    or size is None
                    or size.get(f"{_WORD}val") != "32"
                    or bold is None
                    or bold.get(f"{_WORD}val", "1") not in {"1", "true"}
                    or keep_next is None
                    or keep_next.get(f"{_WORD}val", "1") not in {"1", "true"}
                ):
                    raise RuntimeError(f"gongwen_fields_attachment_layout_mismatch:{case.name}:{location}:{index}")
    else:
        expected_body = [] if case.enabled else [""] if case.empty else ["任务分工表 / 工作安排表"]
        expected_cell = [""] if case.empty else expected_body
        if body_attachments != expected_body or cell_attachments != expected_cell:
            raise RuntimeError(
                f"gongwen_fields_attachments_mismatch:{case.name}:body={body_attachments!r}:table={cell_attachments!r}"
            )

    expected_scalars = (
        "普通列表：普通甲 / 普通乙",
        "抄送：单位甲、单位乙" if case.enabled else "抄送：单位甲 / 单位乙",
        "（此件公开发布）" if case.enabled else "（（此件公开发布））",
        "成文：2026年10月20日" if case.enabled else "成文：2026-10-20",
    )
    if any(paragraphs.count(text) != 1 for text in expected_scalars):
        raise RuntimeError(f"gongwen_fields_scalar_mismatch:{case.name}:{paragraphs!r}")
    expected_empty_field = [] if case.enabled else ["："]
    if _between(paragraphs, _EMPTY_START, _EMPTY_END) != expected_empty_field:
        raise RuntimeError(f"gongwen_fields_empty_paragraph_mismatch:{case.name}")

    imprint_rows = tables[1].findall(f"{_WORD}tr")
    expected_rows = 1 if case.enabled and case.empty else 2
    if len(imprint_rows) != expected_rows:
        raise RuntimeError(f"gongwen_fields_empty_row_mismatch:{case.name}:{len(imprint_rows)}")
    header = [_paragraph_text(cell) for cell in imprint_rows[0].findall(f"{_WORD}tc")]
    if header != ["DOCWEN IMPRINT", "KEEP THIS ROW"]:
        raise RuntimeError(f"gongwen_fields_imprint_header_changed:{case.name}")
    if len(imprint_rows) == 2:
        actual = [_paragraph_text(cell) for cell in imprint_rows[1].findall(f"{_WORD}tc")]
        expected = ["", ""] if case.empty else ["测试机关", "2026年10月21日印发" if case.enabled else "2026-10-21"]
        if actual != expected:
            raise RuntimeError(f"gongwen_fields_imprint_mismatch:{case.name}:{actual!r}")
