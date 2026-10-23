"""Small real-file oracles for the macOS release package checks."""

from __future__ import annotations

import hashlib
import json
from dataclasses import dataclass
from pathlib import Path
from typing import Any

import fitz
import openpyxl
from docx import Document
from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation

TOKEN = "DOCWEN PROBE CONTENT"
MARKDOWN = f"# Probe title\n\n{TOKEN} 中文内容\n\n| Name | Value |\n| --- | --- |\n| Alpha | 17 |\n"


@dataclass(frozen=True)
class Case:
    name: str
    arguments: tuple[str, ...]
    oracle: str
    output: Path
    expected_exit: int = 0


def require(condition: Any, message: str) -> None:
    if not condition:
        raise ValueError(message)


def identity(path: Path) -> dict[str, Any]:
    return {"bytes": path.stat().st_size, "sha256": hashlib.sha256(path.read_bytes()).hexdigest()}


def make_inputs(root: Path, repo: Path) -> dict[str, Path]:
    root.mkdir()
    inputs = {suffix: root / f"中文 sample.{suffix}" for suffix in ("md", "docx", "xlsx", "csv", "rtf", "pptx")}
    inputs["md"].write_text(MARKDOWN, encoding="utf-8")
    inputs["csv"].write_text("Name,Value\nAlpha,17\n", encoding="utf-8")
    inputs["rtf"].write_text("{\\rtf1\\ansi " + TOKEN + "}", encoding="ascii")
    document = Document()
    document.add_heading("Probe title", 1)
    document.add_paragraph(TOKEN)
    table = document.add_table(rows=2, cols=2)
    for row, values in zip(table.rows, (("Name", "Value"), ("Alpha", "17")), strict=True):
        for cell, value in zip(row.cells, values, strict=True):
            cell.text = value
    document.save(inputs["docx"])
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    require(sheet is not None, "missing worksheet")
    sheet.append(["Name", "Value"])
    sheet.append([TOKEN, 17])
    inputs["xlsx-narrow"] = root / "narrow.xlsx"
    workbook.save(inputs["xlsx-narrow"])
    sheet.column_dimensions["A"].width = 40
    workbook.save(inputs["xlsx"])
    inputs["xlsx2"] = root / "second.xlsx"
    sheet.cell(2, 1, "SECOND WORKBOOK")
    workbook.save(inputs["xlsx2"])
    workbook.close()
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[5])
    slide.shapes.title.text = TOKEN
    presentation.save(inputs["pptx"])
    for name, pages in (
        ("first", ("FIRST PAGE",)),
        ("second", ("SECOND PAGE",)),
        ("two", ("FIRST PAGE", "SECOND PAGE")),
    ):
        path = root / f"{name}.pdf"
        with fitz.open() as pdf:
            for text in pages:
                pdf.new_page(width=300, height=200).insert_text((30, 60), text)
            pdf.save(path)
        inputs[name] = path
    for name, color in (("red", "red"), ("blue", "blue")):
        inputs[name] = root / f"{name}.png"
        Image.new("RGB", (100, 80), color).save(inputs[name])
    inputs["ocr"] = root / "ocr.png"
    image = Image.new("RGB", (1000, 200), "white")
    ImageDraw.Draw(image).text((35, 60), TOKEN, fill="black", font=ImageFont.load_default(size=45))
    image.save(inputs["ocr"])
    inputs["proof-md"] = root / "proof.md"
    inputs["proof-md"].write_text("# Review\n\n前文（missing closing bracket\n", encoding="utf-8")
    inputs["proof-docx"] = root / "proof.docx"
    document = Document()
    document.add_paragraph("前文（missing closing bracket")
    document.save(inputs["proof-docx"])
    # Use an existing public repository sample; private workspace files are never uploaded.
    inputs["corpus"] = root / "sample.md"
    inputs["corpus"].write_bytes((repo / "tests/fixtures/files/sample.md").read_bytes())
    return inputs


def make_cases(inputs: dict[str, Path], root: Path, *, scope: str) -> list[Case]:
    root.mkdir()
    cases: list[Case] = []

    def conversion(name: str, source: str, target: str, oracle: str, *extra: str, expected_exit: int = 0) -> None:
        output = root / name
        cases.append(
            Case(
                name,
                ("convert", str(inputs[source]), "--to", target, "--output-dir", str(output), *extra),
                oracle,
                output,
                expected_exit,
            )
        )

    if scope == "office":
        for source in ("docx", "xlsx", "md"):
            conversion(f"{source}-to-pdf", source, "pdf", "office-pdf")
        conversion("xlsx-narrow-to-pdf", "xlsx-narrow", "pdf", "office-pdf-observe")
        conversion("pptx-to-pdf", "pptx", "pdf", "unsupported-target", expected_exit=2)
        conversion("rtf-to-md", "rtf", "md", "text")
        return cases
    conversion("md-to-docx", "md", "docx", "docx")
    conversion("docx-to-md", "docx", "md", "text")
    conversion("md-to-xlsx", "md", "xlsx", "table")
    conversion("xlsx-to-md", "xlsx", "md", "text")
    conversion("csv-to-xlsx", "csv", "xlsx", "table")
    conversion("png-to-jpeg", "red", "jpg", "jpeg")
    conversion("pdf-to-png", "two", "png", "pages")
    conversion("image-ocr-to-md", "ocr", "md", "ocr", "--ocr", "--ocr-language", "english", "--no-recognize-tables")
    conversion("public-corpus-to-docx", "corpus", "docx", "corpus-docx")
    for suffix in ("md", "docx"):
        output = root / f"reviewed.{'json' if suffix == 'md' else 'docx'}"
        cases.append(
            Case(
                f"validate-{suffix}",
                ("validate", str(inputs[f"proof-{suffix}"]), "--check", "punct", "--report", str(output)),
                f"proof-{suffix}",
                output,
            )
        )
    output = root / "numbered.md"
    cases.append(
        Case(
            "number-markdown",
            ("number", "markdown", str(inputs["md"]), "--operation", "add", "--output", str(output)),
            "number",
            output,
        )
    )
    output = root / "merged.pdf"
    cases.append(
        Case(
            "merge-pdf",
            ("merge", "pdf", str(inputs["first"]), str(inputs["second"]), "--output", str(output)),
            "merge-pdf",
            output,
        )
    )
    output = root / "split"
    cases.append(
        Case(
            "split-pdf",
            ("split", "pdf", str(inputs["two"]), "--pages", "2", "--output-dir", str(output)),
            "split-pdf",
            output,
        )
    )
    output = root / "merged.xlsx"
    cases.append(
        Case(
            "merge-tables",
            ("merge", "tables", str(inputs["xlsx"]), str(inputs["xlsx2"]), "--output", str(output)),
            "merged-table",
            output,
        )
    )
    output = root / "merged.tif"
    cases.append(
        Case(
            "merge-images",
            ("merge", "images", str(inputs["red"]), str(inputs["blue"]), "--output", str(output)),
            "tiff",
            output,
        )
    )
    output = root / "preserve.pdf"
    output.write_bytes(b"KEEP EXISTING OUTPUT")
    cases.append(
        Case(
            "existing-output-refusal",
            ("merge", "pdf", str(inputs["first"]), str(inputs["second"]), "--output", str(output)),
            "conflict",
            output,
            7,
        )
    )
    return cases


def verify_case(case: Case, payload: dict[str, Any], inputs: dict[str, Path]) -> dict[str, Any]:
    """Inspect written content independently of CLI success reporting."""
    output = case.output
    if case.oracle == "unsupported-target":
        require(not output.exists(), "unsupported route created output")
        error = payload.get("error", {})
        require(
            error.get("code") == "invalid_input"
            and "No canonical runtime route for pptx -> pdf" in error.get("message", ""),
            "unexpected product route refusal",
        )
        return {"productRouteUnavailable": True}
    files = sorted(path for path in output.rglob("*") if path.is_file()) if output.is_dir() else [output]
    require(files and all(path.is_file() for path in files), "no output files")
    observations: dict[str, Any] = {}
    if case.oracle == "conflict":
        require(output.read_bytes() == b"KEEP EXISTING OUTPUT", "existing output changed")
        require(payload.get("error", {}).get("code") == "output_exists", "wrong conflict error")
    elif case.oracle in {"docx", "corpus-docx", "proof-docx"}:
        [path] = [path for path in files if path.suffix == ".docx"]
        document = Document(path)
        text = "\n".join(paragraph.text for paragraph in document.paragraphs)
        if case.oracle == "proof-docx":
            require(text == "前文（missing closing bracket", "proofread changed document text")
            require(len(document.comments) > 0, "missing proofread comments")
        else:
            require((TOKEN if case.oracle == "docx" else "Sample") in text, "missing DOCX content")
            if case.oracle == "docx":
                require(
                    any("Alpha" in cell.text for table in document.tables for row in table.rows for cell in row.cells),
                    "missing DOCX table",
                )
    elif case.oracle in {"text", "ocr", "number"}:
        text = "\n".join(path.read_text(encoding="utf-8-sig") for path in files if path.suffix == ".md")
        require(TOKEN in text, "missing Markdown content")
        if case.oracle == "number":
            require(text != inputs["md"].read_text(encoding="utf-8"), "numbering did not change headings")
            require("Probe title" in text and "Alpha" in text, "numbering lost content")
    elif case.oracle in {"table", "merged-table"}:
        [path] = [path for path in files if path.suffix == ".xlsx"]
        workbook = openpyxl.load_workbook(path)
        try:
            values = [cell for sheet in workbook for row in sheet.iter_rows(values_only=True) for cell in row]
            require(
                ("Alpha" in values and (17 in values or "17" in values))
                if case.oracle == "table"
                else any(TOKEN in str(value) for value in values)
                and any("SECOND WORKBOOK" in str(value) for value in values),
                "missing workbook values",
            )
        finally:
            workbook.close()
    elif case.oracle in {"merge-pdf", "split-pdf", "office-pdf", "office-pdf-observe"}:
        pages: list[str] = []
        for path in files:
            if path.suffix == ".pdf":
                with fitz.open(path) as pdf:
                    pages.extend(page.get_text() for page in pdf)
        if case.oracle.startswith("office-pdf"):
            text = " ".join(" ".join(pages).split())
            observations = {"extractedPdfText": text[:4000], "containsCompleteToken": TOKEN in text}
            require(pages and text, "empty Office PDF")
            if case.oracle == "office-pdf":
                require(TOKEN in text, f"missing Office PDF text; extracted={text[:1000]!r}")
        elif case.oracle == "merge-pdf":
            require(
                len(pages) == 2 and "FIRST PAGE" in pages[0] and "SECOND PAGE" in pages[1],
                "wrong merged PDF order/content",
            )
        else:
            require(
                len([path for path in files if path.suffix == ".pdf"]) == 2
                and len(pages) == 2
                and "SECOND PAGE" in pages[0]
                and "FIRST PAGE" in pages[1],
                "wrong split page order/content",
            )
    elif case.oracle == "proof-md":
        report = json.loads(output.read_text(encoding="utf-8"))
        require(report["schema"] == "docwen.proofread_report.v2" and report["issues"], "missing proofread issues")
        require(report["source"]["content_sha256"] == identity(inputs["proof-md"])["sha256"], "wrong proofread source")
    else:
        pictures = [path for path in files if path.suffix.lower() in {".jpg", ".jpeg", ".png", ".tif"}]
        require(len(pictures) == (2 if case.oracle == "pages" else 1), "wrong image count")
        for path in pictures:
            with Image.open(path) as picture:
                require(picture.width > 0 and picture.height > 0, "empty image")
                if case.oracle == "jpeg":
                    require(picture.format == "JPEG" and picture.size == (100, 80), "wrong JPEG")
                if case.oracle == "tiff":
                    require(picture.format == "TIFF" and picture.n_frames == 2, "wrong TIFF frames")
                    require(picture.convert("RGB").getpixel((0, 0)) == (255, 0, 0), "wrong first TIFF frame")
                    picture.seek(1)
                    require(picture.convert("RGB").getpixel((0, 0)) == (0, 0, 255), "wrong second TIFF frame")
    return {
        "files": {str(path.relative_to(output) if output.is_dir() else path.name): identity(path) for path in files},
        "observations": observations,
    }
